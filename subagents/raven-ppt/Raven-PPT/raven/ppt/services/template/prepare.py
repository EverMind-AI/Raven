"""Turning a template the user handed us into one an author can build in.

One operation, and it exists because of one measured fact: every template in a
sample of thirty ships example slides -- thirteen each, in all thirty. An author
that opens such a file and starts adding pages produces thirteen slides of somebody
else's content plus its own, which fails the page budget, leaves nothing mapping to
the first thirteen, and puts a stranger's placeholder text in the deck.

python-pptx has no API for removing a slide, so without this every author would
reinvent the same XML surgery, and get it wrong in the same way: dropping the entry
from the slide list without dropping the relationship leaves the part in the package
and PowerPoint repairs the file on open.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path

from raven.ppt.services.template.inventory import (
    PREPARED_FILE,
    TemplateInventory,
    inspect_template,
    template_dir,
)

__all__ = ["PREPARED_FILE", "Prepared", "prepare", "prepared_path", "strip_hidden"]


def prepared_path(project) -> Path:
    return template_dir(project) / PREPARED_FILE


@dataclass(frozen=True)
class Prepared:
    """A template ready to build in, and what had to be removed to get there."""

    path: Path
    inventory: TemplateInventory
    removed_slides: int


def prepare(source: Path, destination: Path) -> Prepared | None:
    """Copy the template, empty it of its example slides, and describe it.

    Returns None when there is nothing usable at `source` -- a missing file, or one
    python-pptx cannot open. The caller treats that as "no template", which is a
    normal state rather than an error: most decks do not have one.

    The original is never modified. It is the user's file, and a later run may want
    to prepare it again from a clean copy.
    """
    if not source.is_file():
        return None
    try:
        from pptx import Presentation
    except ImportError:  # pragma: no cover - python-pptx ships with the extra
        return None

    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    try:
        presentation = Presentation(str(destination))
    except Exception:  # noqa: BLE001 -- a file that will not open is not a template
        destination.unlink(missing_ok=True)
        return None

    removed = _drop_slides(presentation)
    try:
        presentation.save(str(destination))
    except OSError:
        destination.unlink(missing_ok=True)
        return None

    inventory = inspect_template(destination)
    if inventory is None:
        destination.unlink(missing_ok=True)
        return None
    return Prepared(path=destination, inventory=inventory, removed_slides=removed)


def strip_hidden(path: Path) -> int:
    """Drop the slides marked hidden from a template copy, and say how many.

    A hidden slide is in the file and not in the render: LibreOffice does not
    export it. Everything downstream numbers pages off the file, so the numbers
    then name a page the render does not have -- which is how one live run asked
    for thirteen pages of an eleven-page PDF, lost all eleven to a single
    out-of-range number, and told the author renders were unavailable on this
    machine. It then read every example page as code and never looked at a render
    again. Of a 193-template library 119 ship them, always the vendor's trailing
    advertisement, so this is the ordinary case and not the odd one.

    Takes a copy, never the user's own file. Returns 0 for a file that will not
    open, which the caller already treats as "no template".
    """
    try:
        from pptx import Presentation
    except ImportError:  # pragma: no cover - python-pptx ships with the extra
        return 0
    try:
        presentation = Presentation(str(path))
    except Exception:  # noqa: BLE001 -- a file that will not open is not a template
        return 0
    slide_ids = presentation.slides._sldIdLst
    hidden = [
        slide_id
        for slide_id, slide in zip(list(slide_ids), list(presentation.slides), strict=False)
        if slide._element.get("show") == "0"
    ]
    if not hidden:
        return 0
    for slide_id in hidden:
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    try:
        presentation.save(str(path))
    except OSError:
        return 0
    return len(hidden)


def _drop_slides(presentation) -> int:
    """Remove every slide, leaving the masters, layouts and theme in place.

    Both halves are required. Removing the id from the slide list without dropping
    the relationship leaves an orphaned part that PowerPoint offers to repair; the
    reverse leaves a dangling reference that will not open at all.
    """
    slide_ids = presentation.slides._sldIdLst
    removed = 0
    for slide_id in list(slide_ids):
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
        removed += 1
    return removed
