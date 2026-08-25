from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

# The name a template is stored under when it arrives without one of its own.
TEMPLATE_FILE = "source.pptx"
PREPARED_FILE = "prepared.pptx"
# The ground measured off a render of the template, recorded once beside it.
GROUND_FILE = "ground.txt"

# Placeholder roles worth naming for an author. The rest -- date, footer, slide
# number -- are the template's own furniture and not something a deck writes into.
_USEFUL_PLACEHOLDERS = {
    "TITLE",
    "CENTER_TITLE",
    "SUBTITLE",
    "BODY",
    "OBJECT",
    "PICTURE",
    "TABLE",
    "CHART",
}


def template_dir(project) -> Path:
    return project.root / "template"


def template_path(project) -> Path:
    """The template this project was given, under whatever the user called it.

    The name is kept rather than normalised because it is how the user refers to
    the file: a reply that says "built on source.pptx" is a reply about a file
    they never saw. There is one template per deck, so the lookup is "the .pptx
    in the slot that is not the prepared copy", and binding a second one replaces
    the first rather than leaving the lookup a choice to make.
    """
    folder = template_dir(project)
    found = sorted(p for p in folder.glob("*.pptx") if p.name != PREPARED_FILE) if folder.is_dir() else []
    return found[0] if found else folder / TEMPLATE_FILE


@dataclass(frozen=True)
class TemplateLayout:
    """One layout an author may build a page on."""

    index: int
    name: str
    placeholders: tuple[tuple[int, str], ...] = field(default_factory=tuple)

    def summary(self) -> str:
        if not self.placeholders:
            return f"[{self.index}] {self.name} — no placeholders, draw freely"
        named = ", ".join(f"{role.lower()}(idx {idx})" for idx, role in self.placeholders)
        return f"[{self.index}] {self.name} — {named}"


@dataclass(frozen=True)
class TemplateInventory:
    """What an author needs to know to build inside a template it cannot see."""

    path: Path
    width_in: float
    height_in: float
    layouts: tuple[TemplateLayout, ...] = field(default_factory=tuple)
    theme_colours: tuple[tuple[str, str], ...] = field(default_factory=tuple)
    colour_map: tuple[tuple[str, str], ...] = field(default_factory=tuple)
    """The master's `clrMap`: which theme slot each scheme name refers to.

    Load-bearing and easy to skip. A dark template is dark by mapping `bg1` to
    `dk1`, and its master paints `schemeClr val="bg1"` -- so reading the scheme
    straight off the theme says "background: white" about a black deck. Read
    without it, a derived palette comes out exactly inverted.
    """
    rendered_ground: str | None = None
    """The ground read off a render of the template, when one has been taken.

    Ahead of everything the file says about itself, because it is the only reading
    that cannot be wrong: measured across 119 templates, what they declared matched
    what they rendered 73 times, and eleven declared the inverse of what they draw.
    Written beside the template once, by whoever first renders it.
    """
    painted_ground: str | None = None
    """What the master paints its background with, when it paints one at all.

    A layer above `colour_map`, and the one that decides what a reader sees. A
    master may map `bg1` to `lt1` and then cover the whole page with `<p:bg>` in
    accent1 -- one template read as a #2F2F2F deck on that mapping while every page
    of it is #5E31FF purple, because the grey it declared is never painted. Held as
    the scheme name or hex the XML states; resolving it needs the palette and the
    map, which live with the theme.
    """
    fonts: tuple[str, ...] = field(default_factory=tuple)
    example_slides: int = 0

    def brief(self) -> str:
        """The inventory as the author reads it, in the reply to a build call."""
        lines = [f"Template: {self.path.name}, {self.width_in:g}x{self.height_in:g}in.", "Layouts:"]
        lines += [f"  {layout.summary()}" for layout in self.layouts]
        if self.theme_colours:
            shown = ", ".join(f"{name}={value}" for name, value in self.theme_colours)
            lines.append(f"Theme colours: {shown}")
        if self.fonts:
            lines.append(f"Theme fonts: {', '.join(self.fonts)}")
        if self.example_slides:
            lines.append(
                f"It ships {self.example_slides} example slide(s). Delete them before you add your own, "
                "or they count towards the deck and nothing will map to them."
            )
        lines.append(
            "Open it with Presentation(os.environ['PPT_TEMPLATE']) and build on these layouts. Take "
            "colours from the theme above rather than from ppt_theme -- the template's identity is the "
            "deck's identity now."
        )
        return "\n".join(lines)


def inspect_template(path: Path) -> TemplateInventory | None:
    """What the template offers, or None when there is no usable one there."""
    if not path.is_file():
        return None
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:  # pragma: no cover - python-pptx ships with the extra
        return None
    try:
        presentation = Presentation(str(path))
    except Exception:  # noqa: BLE001 -- any malformed file is simply not a template
        return None

    layouts = tuple(
        TemplateLayout(index=index, name=layout.name or f"layout {index}", placeholders=_placeholders(layout))
        for index, layout in enumerate(presentation.slide_layouts)
    )
    root = _theme_root(presentation)
    return TemplateInventory(
        path=path,
        width_in=round(Emu(presentation.slide_width or 0).inches, 3),
        height_in=round(Emu(presentation.slide_height or 0).inches, 3),
        layouts=layouts,
        theme_colours=_theme_colours(root),
        colour_map=_colour_map(presentation),
        painted_ground=_painted_ground(presentation),
        rendered_ground=read_ground(path.parent),
        fonts=_fonts(root),
        example_slides=len(presentation.slides),
    )


def read_ground(folder: Path) -> str | None:
    """The measured ground beside a template, or None before anything measured it."""
    try:
        value = (folder / GROUND_FILE).read_text(encoding="utf-8").strip()
    except OSError:
        return None
    return value if value.startswith("#") and len(value) == 7 else None


def write_ground(folder: Path, colour: str) -> None:
    """Record it beside the template, so later builds do not re-render to learn it."""
    (folder / GROUND_FILE).write_text(colour, encoding="utf-8")


def _placeholders(layout) -> tuple[tuple[int, str], ...]:
    found: list[tuple[int, str]] = []
    for shape in layout.placeholders:
        try:
            fmt = shape.placeholder_format
            role = str(fmt.type).split(" ")[0].upper()
            index = int(fmt.idx)
        except (AttributeError, TypeError, ValueError):
            continue
        if role in _USEFUL_PLACEHOLDERS:
            found.append((index, role))
    return tuple(found)


def _theme_root(presentation):
    """The theme part's XML, or None when the package has none."""
    try:
        theme = presentation.slide_masters[0].part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
        )
        root = getattr(theme, "_element", None)
        if root is None:
            from lxml import etree

            root = etree.fromstring(theme.blob)
    except Exception:  # noqa: BLE001 -- a theme is optional in a valid package
        return None
    return root


def _colour_map(presentation) -> tuple[tuple[str, str], ...]:
    """The master's `clrMap`, or PowerPoint's default when there is none."""
    try:
        master = presentation.slide_masters[0]._element
        node = master.find("{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap")
    except Exception:  # noqa: BLE001 -- a master is required, but be safe
        node = None
    if node is None:
        return (("bg1", "lt1"), ("tx1", "dk1"), ("bg2", "lt2"), ("tx2", "dk2"))
    return tuple(sorted(node.attrib.items()))


def _painted_ground(presentation) -> str | None:
    """The fill the master lays over the whole page, as the XML states it.

    Only a plain solid fill counts. A gradient or a picture background has no one
    colour to derive a palette from, and guessing an average of one is worse than
    falling back to what the colour map says.
    """
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    }
    try:
        master = presentation.slide_masters[0]._element
    except Exception:  # noqa: BLE001 -- a master is required, but be safe
        return None
    fill = master.find(".//p:bg//a:solidFill", ns)
    if fill is None:
        return None
    for colour in (fill.find("a:schemeClr", ns), fill.find("a:srgbClr", ns)):
        if colour is None or not colour.get("val"):
            continue
        # Only a colour stated plainly. DrawingML can alpha, tint, shade, lumMod or
        # saturate a fill, and resolving those needs whatever is behind it. One
        # template paints `tx1` at `alpha val="5000"` -- 5% black over white, a pale
        # grey page -- and taken at face value it read as a black deck and put black
        # body copy on it. What is not plain falls back to the colour map.
        if len(colour):
            return None
        val = colour.get("val")
        return val if colour.tag.endswith("schemeClr") else f"#{val.upper()}"
    return None


def _theme_colours(root) -> tuple[tuple[str, str], ...]:
    """The theme's colour scheme, read off the XML rather than guessed.

    An author given a template and no palette invents one, and a deck in invented
    colours inside somebody's template looks worse than either would alone.
    """
    if root is None:
        return ()
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    found: list[tuple[str, str]] = []
    for element in root.findall(".//a:clrScheme/*", ns):
        name = element.tag.rsplit("}", 1)[-1]
        srgb = element.find("a:srgbClr", ns)
        system = element.find("a:sysClr", ns)
        value = srgb.get("val") if srgb is not None else (system.get("lastClr") if system is not None else None)
        if value:
            found.append((name, f"#{value.upper()}"))
    return tuple(found)


def _fonts(root) -> tuple[str, ...]:
    """The theme's major and minor typefaces.

    Off the font scheme rather than off the layouts' placeholders: on thirty real
    templates every placeholder returned None, because the typeface is inherited
    from the theme rather than set on the shape. Reading the wrong place produced an
    empty list, which would have left an author choosing a font the template had
    already chosen.
    """
    if root is None:
        return ()
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    found: list[str] = []
    for which in ("majorFont", "minorFont"):
        for tag in ("latin", "ea"):
            element = root.find(f".//a:fontScheme/a:{which}/a:{tag}", ns)
            name = element.get("typeface") if element is not None else None
            if name and name not in found:
                found.append(name)
    return tuple(found)
