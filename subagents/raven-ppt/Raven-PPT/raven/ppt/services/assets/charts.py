"""The chart forms, projected as a module an author's build script imports.

There was no chart of any kind on this route. `add_chart` was never reachable --
the runner deliberately withholds matplotlib and says so -- and what an author
did instead was compute rectangles by eye, one page at a time: a live run spent
four requests moving a bar it had drawn 0.3in too long because its own arithmetic
divided the box before it subtracted the axis. Every deck that wanted a
comparison either drew one badly or gave up and wrote the numbers as prose.

So the mapping from a value to a length is code. Every form here is axis-aligned
rectangles and straight lines over a single linear scale, which is what `Box`
already divides and what `write` already sets type in.

Twenty-three forms is a vocabulary and not a ceiling, and for a while it was read
as one: a page wanting a shape none of them make had nothing to reach for but the
nearest of the twenty-three, because the scales, the axes, the shade lines and the
label fitting under them were all private. They are public now. `rect`, `disc`,
`ring`, `hline`, `vline`, `poly` and `write_label` are the ink; `span`, `snap` and
`linear` are the scale; `series_paints`, `stack_paints`, `shades`, `emphasis`,
`ink_on` and `contrast` are the palette discipline the forms obey, so a chart an
author writes obeys the same one rather than inventing a colour. What is still
out of reach is an arc: a pie, a donut, a gauge and a sunburst want a sector at a
computed angle and no primitive here makes one. A filled polygon is not on that
list any more -- `poly(fill=...)` closes and paints one -- so a radar, a band and
a ribbon are drawable by hand, and the skill says which of them are worth drawing.

The first fifteen were the forms a page could not do without. The eight after
them are the forms eight delivered decks went without: 160 pages carrying two
charts between them, and the arguments the rest made -- a series that moved, one
figure reported four ways, twelve capabilities against two competitors, a run of
stages losing people, dated events with no duration, a spread with one slow
request in it, a whole divided into very unequal parts, segments of very
different size -- each had a shape, and every one of them was set as prose or as
a plain table because no primitive here would draw it.

They live in their own module rather than in `ppt_layout` for one reason: the
grid is read by every page and this is read by the pages that carry data. Folding
900 lines of scales and axes into the module an author opens to look up
`split_left` would make the common case pay for the uncommon one, and both are
read in full when they are read at all. `ppt_charts` imports the grid, so nothing
is duplicated -- the measurement that sizes a table's columns is the same one
that decides whether a value fits over its bar.

`raven.ppt.services.assets.script_helpers` writes this as `ppt_charts.py` beside
the script, next to `ppt_layout.py`. Coordinates are inches, for the same reason
they are there.
"""

from __future__ import annotations

import ast
from functools import lru_cache

CHART_MODULE_FILENAME = "ppt_charts.py"

_CHART_MODULE = '''"""Charts drawn from rectangles, hairlines and labels.

    from ppt_charts import column, horizontal_bar, waterfall
    from ppt_layout import page
    from ppt_theme import THEMES

    T = THEMES[next(iter(THEMES))]        # one entry, and it is the template's
    frame = page()
    column(slide, frame.body, T, [("East", 185), ("South", 142), ("North", 128)],
           accent="East", unit="M")

Nothing here reaches for PowerPoint's chart object. A chart made of shapes stays
on the deck's palette, keeps the deck's face, and can be nudged by whoever opens
the file; `add_chart` arrives with Office's own six colours and its own
gridlines, which is as recognisable a tell as any stock template. It is also not
reachable on this route.

Every one of these obeys the same rules, so they are stated once here rather than
once for each below:

* **Length is the value.** A scale always contains zero, so a bar's length is
  proportional to what it says. `axis_max` only ever raises the top of a scale --
  it is there so two charts on one page can be read against each other -- and
  never truncates a bar. A value of exactly zero draws nothing; a value small
  enough to round away draws a sliver, because a bar that is not there reads as
  no data rather than as very little.
* **The label is on the mark.** Categories, values and axis ends are written
  beside the thing they belong to. The one construction that cannot do that -- a
  plot with several series in it -- gets a key of swatches, which is what a legend
  is for.
* **One thing is accented.** `accent` names what carries the page's claim, by
  index or by label (or several of either). It takes the theme's accent and
  everything else goes quiet in `muted`. Name nothing and a single series is
  drawn in `chart_series[0]`, the colour a theme reserves for the series a page
  is about.
* **Colour says which kind of thing it is.** Several bars standing side by side
  are several things, so they take `chart_series` in its own order -- it is
  spaced by hue and checked for colour vision. The parts of one stack, the roles
  of a waterfall and the two ends of a dumbbell are one thing divided, so they
  take depths of a single shade line instead. `chart_series` is ordered by
  prominence and not by lightness -- deep, mid, very pale, and round again -- so
  reading it in order for the second case comes out striped, which says
  "unrelated" about the parts of one quantity and loses the pale ones against the
  page.
* **Axes and baselines are hairlines**, under the thickness at which a filled
  shape stops being a rule and starts being a bar. There is no frame around a
  plot and no gridline that is not carrying a reading.
* **Nothing is invented.** These draw the numbers they are given. Where a value
  is missing the page has to say so; there is no placeholder here to pass.

`box` is a `Box` from `ppt_layout` -- inches, two corners -- and the chart fills
it. Give it a region out of the grid (`frame.body`, one of `body.rows(2)`) and
the margins, the gutters and the alignment come out right by construction.

There is no size, face or colour to pass. Type comes off the ramp in
`ppt_layout` and steps down one notch when a label does not fit, colour comes off
the theme, and the two together are what makes a deck of charts look like one
deck.

Every one of them hands back a `Drawn`. It *is* the plot's rectangle -- the box
left after the labels, the axis readings and the key have taken theirs, which is
the one region the caller cannot work out for itself -- and it carries two more
things: `where`, the chart's own map from a reading to an inch on that plot, and
what the chart had to give up to fit the box it was handed:

    drawn = scatter(slide, frame.body, T, points, axes=("Effort", "Impact"))
    drawn.names_not_written        # the points whose names had nowhere to go
    drawn.readings_not_written     # values the chart could not place
    drawn.marks_not_to_scale       # bubbles drawn at the floor, not to area
    drawn.type_pt                  # the step of the ramp its labels landed on

And every one of them can be asked before it is drawn.
`the_smallest_box_a_chart_needs` answers "how big does this have to be, with
*this* data" in inches, and `what_a_chart_will_do` answers the rest of it -- both
by running the chart against a slide that draws nothing, so what comes back is the
chart's own arithmetic rather than a second copy of it that drifts:

    room = the_smallest_box_a_chart_needs(column, T, rows)
    if room.w > cell.w or room.h > cell.h:
        horizontal_bar(slide, cell, T, rows)   # twelve categories, not enough width
    else:
        column(slide, cell, T, rows)

**And when none of the twenty-three is the shape of the argument, draw the shape
of the argument.** The twenty-three are shortcuts, not a ceiling and not a menu to
pick the nearest item off: the same base they are built on is public, so a form
that is not here is a page you write rather than a page you settle for.

    from ppt_charts import linear, rect, series_paints, span, write_label

    low, high = span(values)                       # a scale that contains zero
    at = linear(low, high, plot.y1, plot.y0)       # a reading -> an inch
    paints = series_paints(T, names, accent="Ours")

Ink: `rect`, `disc`, `ring`, `hline`, `vline`, `poly`, `write_label`. Scale:
`span`, `snap`, `linear`. Palette: `series_paints`, `stack_paints`, `shades`,
`emphasis`, `ink_on`, `contrast`. Fitting: `type_face`, `text_width`, `pick_size`,
`line_height`. Furniture: `key`, `scale_top`. Reading a value and writing one:
`number`, `fmt`. What a hand-drawn chart still owes the deck is every rule above
this line -- the scale contains zero where the mark is a length, the label sits on
the mark, one thing is accented, axes are hairlines, and the colour comes out of
`ppt_theme` through the palette calls rather than out of a string you typed. The
skill's `deck/build/references/charts.md` works one all the way through.
"""

import math
from datetime import date

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

# The grid, beside this module in the build directory. `_em_width` and `_rgb` are
# private to it and are borrowed rather than copied on purpose: they are the same
# two questions -- what is this string worth in inches, what colour is this -- and
# a second copy of either would drift, leaving a chart's labels and a table's
# columns disagreeing about the width of the same word.
from ppt_layout import KICKER_PT, LABEL_PT, Box, _em_width, _rgb, write

# A rule stops being a rule and starts being a bar at 4.5pt. Every axis, baseline,
# connector and tick here is drawn under that, so the page's own measurements read
# them as rules rather than as decoration to complain about.
HAIRLINE = 0.014

# The air between a bar and the value written at its end, and between an axis and
# what is written beside it.
LABEL_GAP = 0.07

# The air between a segment too narrow to hold its own number and the number
# written outside it. Tighter than LABEL_GAP because a leader crosses it, and the
# leader is what says which segment the number belongs to.
CALLOUT_GAP = 0.05

# The air two readings on one axis need between the type, rather than between the
# boxes the type is written in: a label's box carries TEXT_INSET of margin it does
# not set in, so two boxes touching is still a readable pair and two boxes a
# tenth of an inch apart can be "05-1606-30". About half an em at KICKER_PT, which
# is the space that makes two readings read as two.
TICK_GAP = 0.08

# How much of a category's slot the bar takes. The rest is the gap that makes six
# columns read as six rather than as a block.
BAR_SHARE = 0.62
# And how wide that share is allowed to get. A slot is the plot's width over the
# number of categories, so two categories in a half-page box ask for a bar 1.43in
# across -- measured on a delivered page, where the pair read as two slabs rather
# than as two readings and left the rest of the page empty, because the chart had
# spent its room on paint instead of on length. `horizontal_bar` and `stacked_bar`
# already cap a bar's thickness for the same reason; this is that cap stood upright.
# Five categories or more are already under it and do not move.
BAR_MAX = 0.62
# Wider when several series share the slot: the gap goes between groups, never
# inside one.
GROUP_SHARE = 0.80

# A mark's diameter, for the plots that place points rather than lengths.
DOT = 0.15

# The mark on a polyline or a row of positions. Smaller than DOT because a line's
# reading is where the line is: a mark wide enough to be a scatter point covers the
# two segments either side of it, and six of them down a row read as a bar.
STEP = 0.09

# A bubble's two ends. The largest magnitude on the plot is BUBBLE_SPAN across and
# every other mark is scaled from it by area, down to BUBBLE_FLOOR -- which is the
# smallest circle that still reads as a mark rather than as a speck of dust on the
# projector. Scaling by area, that floor binds at (BUBBLE_FLOOR / BUBBLE_SPAN)
# squared, about a fifty-sixth of the largest magnitude. Below it the area can no
# longer be the magnitude, and `_plot_points` says so rather than drawing a disc
# whose size is a lie.
BUBBLE_SPAN = 0.60
BUBBLE_FLOOR = 0.08

# A non-zero value never draws as nothing.
MIN_LENGTH = 0.014

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


# ------------------------------------------------------- asking before drawing


class TooSmall(ValueError):
    """A chart refusing a box, carrying how much more of each side it wants.

    `short` is the deficit in inches as (across, down), which is exactly what
    `the_smallest_box_a_chart_needs` adds to the box before asking again. `needs`
    and `had` are the two sizes it compared -- both for *this* data, not the constant
    pair a prose refusal quotes.
    """

    def __init__(self, what, had, needs):
        self.what = what
        self.had = (float(had[0]), float(had[1]))
        self.needs = (float(needs[0]), float(needs[1]))
        self.short = (max(0.0, self.needs[0] - self.had[0]), max(0.0, self.needs[1] - self.had[1]))
        super().__init__(
            f"{what} has {self.had[0]:.2f}x{self.had[1]:.2f}in and needs about "
            f"{self.needs[0]:.2f}x{self.needs[1]:.2f}in"
        )


class _Rehearsal:
    """A slide that draws nothing, which is how a chart is asked before it is drawn.

    Every decision a chart makes -- what the scale is, how much room the labels
    leave, which step of the ramp they set at, which of them there is finally no
    room for -- is arithmetic on the box and the data, and all of it happens on the
    way to the first rectangle. So the way to ask a chart what it would do is to
    run it, with a slide that swallows the shapes: `rect`, `disc`, `ring`,
    `poly` and `write_label` are the only five places in this module that touch `slide`,
    and all five return early on one of these. Which is also what makes the three
    measuring calls work on a chart an author wrote rather than only on the
    twenty-three: a form built out of the same five ink calls rehearses on the same
    slide, so it answers "will this fit" without a shape reaching the page.

    The alternative was a second body of sizing arithmetic beside each chart,
    answering the same question as the chart and free to disagree with it. That is
    the drift this module refuses everywhere else -- `_em_width` is borrowed from
    the grid rather than copied for exactly this reason -- and it would be worse
    here, because the whole value of the answer is that it is the one the chart
    will act on.
    """

    __slots__ = ()


class Drawn(Box):
    """The plot, and what the chart gave up to fit the box it was given.

    A `Box`, so everything that read the old return value still reads this one: the
    chart fills the region it was handed, and what comes back is the plot inside it
    -- the part carrying the marks, after the category labels, the axis readings
    and the key have taken theirs. That is the one rectangle the caller cannot
    derive, and it is what a second chart is aligned against or a note is hung off.

    On top of that:

    * `where` is the chart's own scale, as a function. `where(value)` is the inch a
      reading lands on for the charts with one value axis, `where(x, y)` the (x, y)
      pair for the two-dimensional ones, and each chart's docstring says which it
      is. Without it the plot is a rectangle with no way to put anything at a
      *reading* in it: an author wanting a rule at 80% could divide the box and
      hope it matched the chart's own scale, which is the arithmetic by eye this
      module exists to end.
    * `names_not_written`, `readings_not_written` and `marks_not_to_scale` are what
      the chart dropped or degraded. Each was a silent decision: a scatter whose
      names will not clear leaves them off, a stack segment too thin for its number
      loses it when there is no air beside it either, a bubble under the floor
      stops being drawn to area. All three were visible only by counting a render,
      which is the blind adjusting a primitive exists to remove.
    * `type_pt` is the smallest step of the ramp the chart set any of its own
      labels at. It is a reading, never a setting: there is still no way to hand a
      chart a size.
    """

    # No `__slots__`: a subtype of a tuple cannot have one, because the tuple's own
    # storage is already variable-length. The five below live in the instance dict.

    def __new__(cls, plot, where, type_pt, *, names=(), readings=(), floored=()):
        drawn = super().__new__(cls, plot.x0, plot.y0, plot.x1, plot.y1)
        drawn.where = where
        drawn.type_pt = float(type_pt)
        drawn.names_not_written = tuple(str(name) for name in names)
        drawn.readings_not_written = tuple(str(reading) for reading in readings)
        drawn.marks_not_to_scale = tuple(str(name) for name in floored)
        return drawn

    @property
    def box(self):
        """This, spelled the other way.

        A chart's `Drawn` *is* its plot box and `ppt_layout`'s is a shape carrying
        one, so the two answer the same question under two names -- and an author
        who has written `picture_fit(...).box` writes `line(...).box` next. It
        cost a build: `print("p11 chart", drawn.box, ...)` on the line after the
        chart was drawn, and `AttributeError` before the deck existed.
        """
        return Box(self.x0, self.y0, self.x1, self.y1)

    @property
    def nothing_was_dropped(self):
        """Whether every name, every reading and every magnitude reached the page."""
        return not (self.names_not_written or self.readings_not_written or self.marks_not_to_scale)


def _plain(shape):
    """Take the Office shape style off, which is the only way the shadow goes.

    `shape.shadow.inherit = False` writes an empty `a:effectLst`, and that is
    supposed to mean "no effects". It does not, because `add_shape` also stamps
    `p:style` on the shape and its `a:effectRef idx="2"` is a reference to the
    theme's second effect -- a drop shadow -- which a renderer resolves separately.
    Left alone, every bar, every baseline and every dot comes out with a grey shadow
    down its right side, which on a page of forty rectangles is the single loudest
    tell that nobody looked at the render.

    The style also carries a line, a fill and a font from the theme's accent, all
    of which are set explicitly here, so nothing is lost with it.
    """
    element = shape._element  # noqa: SLF001 -- python-pptx has no API for p:style
    for style in element.findall(f"{{{_P}}}style"):
        element.remove(style)
    return shape


def _solid(shape, colour, opacity):
    """Fill a shape solid, and make that fill translucent where it was asked to be.

    The alpha is a child element on the colour -- `<a:alpha val="50000"/>` under the
    `a:srgbClr` the fill just wrote, in thousandths -- and python-pptx exposes no
    API for it, so it is written here the way `_plain` writes its own fix-up.

    It works end to end: a translucent fill over an opaque one, written to pptx and
    converted by the renderer, comes out as the blend of the two with what is under it
    still legible -- so a radar is not a one-series form and a band between two series
    is drawable. What is true is only that a *theme token* is a pre-mixed solid:
    `accent_soft` is mixed at rest so it is the same colour whatever it lands on,
    which is a different decision and still the right one for a ground.
    """
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(colour)
    if opacity >= 1.0:
        return shape
    element = shape.fill.fore_color._xFill.find(f"{{{_A}}}srgbClr")  # noqa: SLF001 -- no API for a:alpha
    element.append(element.makeelement(f"{{{_A}}}alpha", {"val": str(int(round(opacity * 100000)))}))
    return shape


# The public helpers that take a slide and are not forms. "Takes a slide first" is
# how this module answers "is it a chart", and it was a complete answer while the
# only things here that touched a slide were the twenty-three. It stopped being one
# when the base under them was made public: a primitive listed as a form is one a
# page picks off the catalogue to make its argument with, and `rect` is not an
# argument. Read by `_drawable_names` below and by the catalogue in the service, so
# a base helper added later joins them here and nowhere else.
_NOT_A_FORM = (
    "rect",
    "disc",
    "ring",
    "hline",
    "vline",
    "poly",
    "write_label",
    "key",
    "scale_top",
)


# ----------------------------------------------------------------- the drawing


def rect(slide, box, colour, opacity=1.0):
    """A filled rectangle with no outline and no shadow: every bar, band and swatch.

    A box out of the grid or one worked out against a scale, painted. It is the
    whole of what a bar is, and drawing one with `add_shape` instead is how a page
    ends up with Office's drop shadow down the side of every column on it.

    `opacity` reads the three spellings a share does -- `0.35`, `35` and `"35%"` are
    the same third -- and 1.0, the default, is the opaque paint every form here
    draws. Below 1 the shape is layered rather than stacked: what it covers shows
    through it, which is what a band over a plot and a second filled series both
    need. It is not a way to soften a colour. A paler paint is `shades`, which
    stays the same colour whatever it lands on; a translucent one is a different
    colour over every different thing it crosses.
    """
    if box.w <= 0 or box.h <= 0 or isinstance(slide, _Rehearsal):
        return None
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *box.pptx())
    _solid(shape, colour, _share(opacity, what="a paint's opacity"))
    shape.line.fill.background()
    shape.shadow.inherit = False
    return _plain(shape)


def disc(slide, x, y, diameter, colour, opacity=1.0):
    """A filled circle centred on (x, y).

    `opacity` as on `rect`, and the case for it here is a scatter dense enough that
    marks land on each other: opaque discs make the pile-up read as one mark, and
    translucent ones make it read as a pile-up, which is the reading.
    """
    if isinstance(slide, _Rehearsal):
        return None
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x - diameter / 2), Inches(y - diameter / 2), Inches(diameter), Inches(diameter)
    )
    _solid(shape, colour, _share(opacity, what="a paint's opacity"))
    shape.line.fill.background()
    shape.shadow.inherit = False
    return _plain(shape)


def ring(slide, x, y, diameter, colour):
    """An open circle centred on (x, y): the mark that is not carrying an area.

    Filled discs on a bubble plot are read as areas, so the marks that are too
    small to be drawn to scale have to stop looking like them. An outline of the
    same paint at the floor size is the same mark, visibly not filled in.
    """
    if isinstance(slide, _Rehearsal):
        return None
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x - diameter / 2), Inches(y - diameter / 2), Inches(diameter), Inches(diameter)
    )
    shape.fill.background()
    shape.line.color.rgb = _rgb(colour)
    shape.line.width = Inches(HAIRLINE)
    shape.shadow.inherit = False
    return _plain(shape)


def hline(slide, x0, x1, y, colour, thickness=HAIRLINE):
    """A horizontal hairline with its top edge at `y`, so a bar sits on it."""
    return rect(slide, Box(min(x0, x1), y, max(x0, x1), y + thickness), colour)


def vline(slide, x, y0, y1, colour, thickness=HAIRLINE):
    """A vertical hairline centred on `x`."""
    return rect(slide, Box(x - thickness / 2, min(y0, y1), x + thickness / 2, max(y0, y1)), colour)


def poly(slide, points, colour, width=0.020, fill=None, opacity=1.0):
    """A polyline through `points`: the cumulative curve, a connector, a whisker.

    `fill` closes it and paints the inside, which is the one shape the twenty-three
    never make and the one this module said for a long time could not be made. It
    can: a freeform closes and takes a solid like any other autoshape, and what was
    actually missing was a caller with a reason. A radar's web, a band between two
    series, a ribbon that changes width along its length and a wedge are all
    polygons -- so the forms this module substitutes away from are drawable by hand
    even though no form here draws them, and the skill says which are worth drawing.

    Left off, the shape stays an outline, which is what every caller in this module
    wants: a curve read against an axis is a line, and filling the area under it
    would claim the area means something.

    `opacity` applies to `fill` and not to the outline, and that split is the whole
    of what makes several filled polygons readable at once: the fills layer, so a
    second one does not delete the first, while the edges stay solid so each shape
    is still a shape. Two filled radars at 1.0 are one radar -- whichever went down
    last. The same two with the upper at 0.45 are two, and where they cross is
    visibly where they cross.
    """
    if isinstance(slide, _Rehearsal):
        return None
    builder = slide.shapes.build_freeform(Inches(points[0][0]), Inches(points[0][1]))
    builder.add_line_segments([(Inches(x), Inches(y)) for x, y in points[1:]], close=fill is not None)
    shape = builder.convert_to_shape()
    if fill is None:
        shape.fill.background()
    else:
        _solid(shape, fill, _share(opacity, what="a paint's opacity"))
    shape.line.color.rgb = _rgb(colour)
    shape.line.width = Inches(width)
    shape.shadow.inherit = False
    return _plain(shape)


def write_label(slide, box, text, theme, *, size=LABEL_PT, colour=None, align="left", anchor="top", bold=False):
    """A label in the theme's own faces, which is the only way a mixed line sets.

    `ppt_layout.write` with both of the theme's faces named on the run, so a Han
    character in a category name is set in the CJK companion rather than in
    whatever the renderer falls back to. Every label these forms write goes through
    here, and one a page draws by hand should too.
    """
    if box.w <= 0 or box.h <= 0 or isinstance(slide, _Rehearsal):
        return None
    return write(
        slide,
        box,
        text,
        size=size,
        bold=bold,
        colour=colour or theme["foreground"],
        align=align,
        anchor=anchor,
        font=theme.get("font_family"),
        cjk_font=theme.get("cjk_font_family"),
    )


# ------------------------------------------------------------------ the paints


def _muted(theme):
    return theme.get("muted", theme["foreground"])


def _grid(theme):
    return theme.get("grid", theme.get("surface", theme["muted"]))


def _series(theme, index):
    """Series `index`, wrapping, so an extra series cannot crash a page.

    The order is the contract: `chart_series[0]` is the series a page is about and
    the tail is deliberately quiet on a light ground.
    """
    palette = list(theme.get("chart_series") or ())
    if not palette:
        palette = [theme["accent"]]
    return palette[index % len(palette)]


# The palest a filled area may go, as a contrast ratio against the page. WCAG's
# target for a graphic rather than for text, which is what a segment of a stack or
# a step of a waterfall is: under it the shape is on the page without being on it.
# It is not asked of `chart_series` -- six categorical colours that all clear 3:1
# and stay apart from each other is not a palette anyone ships, which is why the
# reviewed ones let their tail go quiet -- but a shade line is cut to the number of
# steps that are actually wanted, so it can hold.
AREA_CONTRAST = 3.0


def _mix(colour, other, weight):
    """`weight=0` is `colour`, `weight=1` is `other`, and the ends are exact."""
    start, end = str(colour).lstrip("#"), str(other).lstrip("#")
    channels = []
    for offset in (0, 2, 4):
        first = int(start[offset : offset + 2], 16)
        last = int(end[offset : offset + 2], 16)
        channels.append(int(round(first + (last - first) * weight)))
    return "#" + "".join(f"{channel:02X}" for channel in channels)


def _poles(theme, quiet):
    """The deep end of a shade line and its pale end, both theme tokens.

    The loud line runs from the paint a theme reserves for the series a page is
    about down to its pre-mixed soft accent. The quiet one runs from the secondary
    ink down to the grid tone, and is what a chart uses for everything it is not
    accenting, so the one accented mark is the only colour on the plot.
    """
    if quiet:
        return _muted(theme), _grid(theme)
    pale = theme.get("accent_soft") or _mix(theme["accent"], theme.get("background", "#FFFFFF"), 0.82)
    return _series(theme, 0), pale


def _shade(theme, share, quiet=False):
    """A point on the shade line: `share=0` is its pale end, `1` its deep end.

    This is the only paint in the module that is arithmetic rather than a lookup,
    and the exception is deliberate. `chart_series` is the one other place six
    ready paints sit, and it is ordered by *prominence*, not by lightness: it runs
    deep, mid, very pale, deep, mid, very pale, so anything that reads it in order
    comes out striped. A stripe says "these are unrelated" -- which is right for
    four bars side by side and wrong for the four parts of one bar, where the
    third part also happened to land on a tint the projector loses.

    Both ends are paints the theme already carries and only the distance between
    them is computed, which is the same arithmetic a theme used to derive its soft
    accent from its accent. Interpolation moves every channel one way, so the line
    cannot double back: it is monotone in lightness, which is the whole of what
    makes it a scale. Each step is a pre-mixed solid rather than a transparency,
    and that is a choice rather than a limit -- the export carries alpha, which is
    what `rect(opacity=...)` writes. A scale has to be the same colour on every
    page it appears on, and a translucent step is a different colour over every
    different thing it lands on, which is not a scale. Transparency is for
    layering two readings on purpose; a shade line is for one reading.
    """
    deep, pale = _poles(theme, quiet)
    return _mix(pale, deep, min(1.0, max(0.0, share)))


def _floor(theme, quiet=False):
    """How far along the line a filled area may still sit, as a share.

    The pale end itself is for a heatmap cell, where "nearly the page" is what a
    low value means and the scale printed under the grid says so. A segment of a
    stack has no such scale and, at the outer end of its bar, no neighbour to be
    read against either, so it stops where it still clears `AREA_CONTRAST` on the
    page. Searched rather than capped at a fixed share: the ten palettes start
    from ten different depths, and one cap would leave some of them invisible and
    the rest of them crowded.
    """
    deep, pale = _poles(theme, quiet)
    ground = theme.get("background", "#FFFFFF")
    if contrast(pale, ground) >= AREA_CONTRAST:
        return 0.0
    low, high = 0.0, 1.0
    for _ in range(24):
        middle = (low + high) / 2
        if contrast(_mix(pale, deep, middle), ground) >= AREA_CONTRAST:
            high = middle
        else:
            low = middle
    return high


def shades(theme, count, quiet=False):
    """`count` paints off one line, deepest first, evenly stepped and all readable.

    What everything ordered here is painted with: the parts of a stack, the roles
    of a waterfall, the two ends of a dumbbell. Even steps in mix weight, which on
    gamma-encoded channels is close enough to even steps in lightness that the
    widest and the narrowest step of a four-part stack are within a few units of
    each other on all ten reviewed palettes.
    """
    if count <= 1:
        return [_shade(theme, 1.0, quiet)]
    floor = _floor(theme, quiet)
    return [_shade(theme, 1.0 - (1.0 - floor) * index / (count - 1), quiet) for index in range(count)]


def _luminance(colour):
    text = str(colour).lstrip("#")
    channels = []
    for offset in (0, 2, 4):
        value = int(text[offset : offset + 2], 16) / 255.0
        channels.append(value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4)
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def contrast(one, other):
    """The WCAG ratio between two paints: 1.0 is the same colour, 21.0 is ink on paper.

    The same number the deck's own measurement reads off the render, so a paint
    checked here is a finding that does not arrive. Type wants 3:1 and refuses under
    2:1; a filled area against the page wants `AREA_CONTRAST`.
    """
    first, second = _luminance(one), _luminance(other)
    return (max(first, second) + 0.05) / (min(first, second) + 0.05)


def ink_on(ground, theme):
    """Whichever of the page's two inks is legible on `ground`.

    A number set inside a segment or a cell is on whatever colour the value put
    there, and picking one ink for all of them puts white on the light end of a
    tint. The deck's own contrast measurement reads the render at 3:1, so this is
    computed rather than guessed at.
    """
    dark = theme["foreground"]
    light = theme.get("background", "#FFFFFF")
    return dark if contrast(dark, ground) >= contrast(light, ground) else light


def _accented(labels, accent):
    """The indices `accent` names: an index, a label, or several of either."""
    if accent is None:
        return set()
    wanted = list(accent) if isinstance(accent, (list, tuple, set, frozenset)) else [accent]
    chosen = set()
    for entry in wanted:
        if isinstance(entry, bool):
            raise ValueError("accent names an index or a label, not a boolean")
        if isinstance(entry, int):
            index = entry if entry >= 0 else len(labels) + entry
            if not 0 <= index < len(labels):
                raise ValueError(f"accent names item {entry}, and there are {len(labels)}")
            chosen.add(index)
            continue
        matches = [index for index, label in enumerate(labels) if label == str(entry)]
        if not matches:
            raise ValueError(f"accent names {entry!r}, which is not one of: {', '.join(labels)}")
        chosen.update(matches)
    return chosen


# Where the quiet series start. The reviewed palettes run three colours of the
# accent's own hue and then three neutrals, so a series that is not the point can
# be grey and still be a different grey from the series beside it. Painting them
# all `muted` -- which is what this did first -- left a three-series key with two
# identical swatches on it.
_QUIET_SERIES = 3


def series_paints(theme, names, accent):
    """One paint per series: the accented one loud, the rest quiet but distinct."""
    chosen = _accented(names, accent)
    if not chosen:
        return [_series(theme, index) for index in range(len(names))]
    paints, quiet = [], 0
    for index in range(len(names)):
        if index in chosen:
            paints.append(theme["accent"])
            continue
        paints.append(_series(theme, _QUIET_SERIES + quiet))
        quiet += 1
    return paints


def stack_paints(theme, names, accent):
    """One paint per part of a stack: depths of one line, not the series palette.

    A stack is a single quantity cut up, so its parts are told apart by depth and
    still read as one bar. `chart_series` is for the other case -- several bars
    side by side, each its own thing -- and taking it in order here painted a
    three-part stack deep, mid, almost-white: three unrelated bands, the third of
    them barely on the page.

    Accenting a part takes it off the line: it gets the theme's accent and the
    rest go quiet, which is the module's one rule about emphasis applied to a
    stack rather than a second way of doing it.
    """
    chosen = _accented(names, accent)
    if not chosen:
        return shades(theme, len(names))
    quiet = shades(theme, max(1, len(names) - len(chosen)), quiet=True)
    paints, taken = [], 0
    for position in range(len(names)):
        if position in chosen:
            paints.append(theme["accent"])
            continue
        paints.append(quiet[taken])
        taken += 1
    return paints


def emphasis(theme, labels, accent):
    """(fills, inks) per item: what the bar is painted and what its label is set in.

    Nothing accented means one series in the theme's own series colour. Something
    accented means exactly that thing in the accent and the rest in `muted` --
    which is the whole of "only emphasise one", done by construction rather than
    left to the author's colour sense.
    """
    chosen = _accented(labels, accent)
    if not chosen:
        return [_series(theme, 0)] * len(labels), [theme["foreground"]] * len(labels)
    fills = [theme["accent"] if index in chosen else _muted(theme) for index in range(len(labels))]
    return fills, list(fills)


def _second_paint(theme, taken):
    """A paint for the second reading on a plot, told apart from the first.

    The theme's accent, which is what a curve riding over bars wants to be -- but
    only while the bars are not already painted it. A reviewed palette keeps
    `chart_series[0]` and `accent` apart; a theme read off a template need not, and
    the one this was rendered against has them the same blue. A rate drawn in the
    columns' own colour is not a second reading: on that render the growth curve
    was visible only in the gutters between the columns, and the two units the form
    exists to hold apart came out as one.
    """
    used = {str(paint).upper() for paint in taken}
    if str(theme["accent"]).upper() not in used:
        return theme["accent"]
    for step in range(1, 6):
        if str(_series(theme, step)).upper() not in used:
            return _series(theme, step)
    return _muted(theme)


# ------------------------------------------------------------------- the input


def number(value, label=None):
    """A reading from a number, or from a string with one in it: "48.3%" -> 48.3.

    Every value any of these charts is given comes through here, so this is where
    a value that cannot be a length is refused -- with `label`, whatever the row
    it came in on was called, in the message when the caller had one.

    A NaN or an infinity is that kind of value. Let past, it surfaces further in as
    `cannot convert float NaN to integer` out of the label formatter or `cannot
    convert float infinity to integer` out of the axis, with a traceback into a
    module the author did not write and no clue which of forty numbers caused it.
    They arrive by arithmetic rather than by typing -- `(new - old) / old` on a row
    whose `old` is zero -- and the author reading the error is the one who wrote
    that expression, so the message names the row and says the value is not
    finite. Refused at the door, once, for all of them.
    """
    where = f" for {label!r}" if label is not None else ""
    if isinstance(value, bool):
        raise ValueError(f"a chart value{where} is a number, not a boolean")
    if isinstance(value, (int, float)):
        reading = float(value)
    else:
        digits = ""
        for character in str(value).strip().replace(",", ""):
            if character.isdigit() or character == "." or (character in "+-" and not digits):
                digits += character
            elif digits:
                break
        try:
            reading = float(digits)
        except ValueError:
            raise ValueError(f"a chart value{where} is a number, not {value!r}") from None
    # Both paths, because a string of four hundred digits reads as an infinity too.
    if not math.isfinite(reading):
        raise ValueError(f"a chart value{where} is a finite number, not {value!r}")
    return reading


def fmt(value, unit="", sign=False):
    """A value as the shortest string that is still the value.

    Trailing zeros go (128.0 -> "128") and nothing else does: a chart that rounded
    its own labels would carry numbers the source it came from does not.

    Never in exponent form. `%g` puts anything over a million into scientific
    notation, so a revenue chart comes back labelled "1.85e+06" over a bar, which is
    a reading nobody in the room converts back.
    """
    reading = float(value)
    rounded = float(f"{reading:.10g}")
    if rounded == int(rounded) and abs(rounded) < 1e15:
        body = str(int(rounded))
    else:
        body = f"{rounded:f}".rstrip("0").rstrip(".")
    if sign and rounded >= 0:
        body = "+" + body
    return f"{body}{unit}"


def _pairs(data, what):
    """[(label, value)] from pairs or from a mapping."""
    items = list(data.items()) if hasattr(data, "items") else list(data)
    if not items:
        raise ValueError(f"{what} needs at least one value")
    rows = []
    for entry in items:
        try:
            label, value = entry
        except (TypeError, ValueError):
            raise ValueError(f"{what} takes (label, value) pairs or a mapping, not {entry!r}") from None
        rows.append((str(label), number(value, label)))
    return rows


def _triples(data, what, third):
    """[(label, first, second)] -- the shape a two-state chart reads."""
    rows = []
    for entry in list(data):
        try:
            label, first, second = entry
        except (TypeError, ValueError):
            raise ValueError(f"{what} takes (label, value, {third}) rows, not {entry!r}") from None
        rows.append((str(label), number(first, label), number(second, label)))
    if not rows:
        raise ValueError(f"{what} needs at least one row")
    return rows


def _series_rows(categories, series, what):
    """(names, columns): every series as long as `categories`."""
    names, columns = [], []
    items = list(series.items()) if hasattr(series, "items") else list(series)
    for entry in items:
        try:
            name, values = entry
        except (TypeError, ValueError):
            raise ValueError(f"{what} takes (name, values) series or a mapping, not {entry!r}") from None
        readings = [number(value, name) for value in values]
        if len(readings) != len(categories):
            raise ValueError(
                f"{what}: series {name!r} has {len(readings)} values and there are {len(categories)} categories"
            )
        names.append(str(name))
        columns.append(readings)
    if not names or not categories:
        raise ValueError(f"{what} needs at least one category and one series")
    return names, columns


def span(values, axis_max=None, axis_min=None):
    """The ends of a scale, which always contains zero.

    `axis_max` raises the top and never lowers it: a scale that cut a bar short
    would be the one failure this whole module exists to make impossible.
    """
    readings = list(values) or [0.0]
    low = min(0.0, min(readings)) if axis_min is None else min(number(axis_min), min(0.0, min(readings)))
    high = max(0.0, max(readings)) if axis_max is None else max(number(axis_max), max(0.0, max(readings)))
    if high - low < 1e-12:
        high = low + 1.0
    return low, high


def linear(low, high, near, far):
    """The map from a value to an inch: `low` lands on `near`, `high` on `far`."""
    span = (high - low) or 1.0

    def at(value):
        return near + (far - near) * (float(value) - low) / span

    return at


# -------------------------------------------------------------------- the type


def line_height(size):
    """The height one line at `size` needs in the box `write` builds for it."""
    return size * 1.15 / 72.0 + 0.06


def type_face(theme):
    """The face the deck is set in, which is what a width estimate needs.

    `_em_width` is a class average with a per-face scale on it, and the scale is
    the whole difference between a label that fits and one the render folds in
    half: the same "12M" that sets 0.378in in the face `Arial` resolves to sets
    0.446in in the one `Cambria` does. Every chart is handed a theme, so there is
    always a face to name and never a reason to fall back to the widest.
    """
    return theme.get("font_family")


def _widest(texts, size, face=None):
    return max([_em_width(str(text), size, face) for text in texts] or [0.0])


# What `write` spends on its own margins (0.04 a side) before a character is set,
# plus a hair. A column sized to the string alone is a column the string wraps in,
# which is the commonest geometry bug on a generated page and was the first thing
# every one of these charts did: "88%" came back as "88" over "%".
TEXT_INSET = 0.16


def text_width(texts, size, face=None):
    """How wide a box has to be for these strings to stay on one line in it."""
    return _widest(texts, size, face) + TEXT_INSET


def _fits(texts, room, size, face=None):
    return text_width(texts, size, face) <= room


def _ink(room, text, size, align="left", anchor="top", face=None):
    """The part of `room` the word written in it will actually cover.

    A text box is as wide as the region it was handed and the line in it is only
    as wide as the line. That difference does not matter until something has to
    keep clear of the label: a quadrant name is written into a whole quarter of a
    2x2 and an axis name into the full width of the plot, and handing those boxes
    to the label placer reserves half the chart for four words. What has to stay
    clear is the ink.
    """
    width = min(room.w, text_width([text], size, face))
    height = min(room.h, line_height(size))
    left = (room.x0 + room.x1 - width) / 2 if align == "center" else (room.x1 - width if align == "right" else room.x0)
    top = (room.y0 + room.y1 - height) / 2 if anchor == "middle" else (room.y1 - height if anchor == "bottom" else room.y0)
    return Box(left, top, left + width, top + height)


def pick_size(texts, room, largest=LABEL_PT, face=None):
    """The largest step of the ramp these labels fit in `room`, or None.

    Two steps and then nothing: 12pt still clears the absolute floor the
    measurements enforce and 10pt does not, so a label that will not fit at 12
    is dropped and the reading it carried is put somewhere it can be read --
    which is what `scale_top` is for.
    """
    for size in (largest, KICKER_PT):
        if _fits(texts, room, size, face):
            return size
    return None


# How far a category label may be wrapped before the answer stops being more lines
# and becomes a wider box. `_rows_needed` allows this many and the foot is built for
# exactly that many, so a name wanting one more is a name written past the region.
LABEL_ROWS = 2


def _rows_needed(texts, room, size, cap=LABEL_ROWS, face=None):
    """How many lines these labels wrap onto in `room`, measured, capped."""
    longest = _widest(texts, size, face)
    return max(1, min(cap, int(math.ceil(longest / max(room - 0.05, 0.05) - 1e-9))))


def _slot_needed(labels, face=None):
    """The narrowest a category's slot may be for its own name to fit under it.

    `_rows_needed` wraps a name onto at most `LABEL_ROWS` lines and the foot is cut
    for that many, so a name that wants a third line is a name set past the bottom
    of the region -- silently, because `write` neither shrinks to fit nor clips.
    Eight categories called "Manufacturing" and the like want 0.62in of slot each at
    12pt, so 4.9in of box; a flat 1.0in floor draws all eight anyway, at 0.125in a
    column.
    """
    return _widest(labels, KICKER_PT, face) / LABEL_ROWS + 0.05


def _rows_high(count):
    """The shallowest a chart of one row per item may be: a line of type each.

    Under this the rows are thinner than the smallest step of the ramp, so each
    label is set over its neighbour's. Every row chart here already steps from
    LABEL_PT down to KICKER_PT when a row is shorter than a line; this is that same
    rule carried past the last step, where the answer stops being a smaller size and
    becomes a taller box. Twelve rows want 3.02in; a flat 0.5in floor passes them.
    """
    return line_height(KICKER_PT) * count


def _column_needed(labels, face, share, step):
    """How wide a box must be for a name column capped at `share` of it to hold one line.

    Every row chart reserves `min(box.w * share, the labels + 0.08)` for its names
    and writes into that minus the 0.08. Where the cap is what binds, the name is
    wrapped inside a box one row tall, which sets its second line over the next
    row's -- so the cap binding at all is the box being too narrow, and this is the
    width at which it stops binding. `step` is the size the chart settled on for
    that column, which is the row height's business and already decided by here.
    """
    return (text_width(labels, step, face) + 0.08) / share


def _drawn(plot, where, sizes, *, names=(), readings=(), floored=()):
    """A chart's return: its plot, its scale, and what it could not put on the page."""
    steps = [size for size in sizes if size]
    return Drawn(plot, where, min(steps) if steps else KICKER_PT, names=names, readings=readings, floored=floored)


_SWATCH = 0.12
_KEY_GAP = 0.22


def _key_widths(names, size=KICKER_PT, face=None):
    """What each entry of a key costs across the page, swatch and gap included."""
    return [_SWATCH + 0.06 + _em_width(str(name), size, face) + TEXT_INSET + _KEY_GAP for name in names]


def _key_rows(names, width, size=KICKER_PT, face=None):
    """How many lines a key of these names needs in `width`."""
    rows, used = 1, 0.0
    for entry in _key_widths(names, size, face):
        if used and used + entry > width:
            rows += 1
            used = 0.0
        used += entry
    return rows


def key(slide, box, theme, names, paints, size=KICKER_PT):
    """Series names with their swatch, for the plot no direct label can reach.

    A legend is what you use when a label cannot reach its mark, and four series
    in a 0.3in column is that case. The swatch is a mark rather than a band: at
    0.12in square it is neither long enough nor thin enough to read as one.

    It wraps. The first version laid the entries out left to right and stopped
    measuring, so a four-series key on a 4in region put its last name off the
    page as a column of single letters -- which is the same failure it exists to
    prevent, and worse for being in the furniture rather than in the data.
    """
    widths = _key_widths(names, size, type_face(theme))
    line_h = box.h / max(1, _key_rows(names, box.w, size, type_face(theme)))
    swatch = min(_SWATCH, max(0.06, line_h * 0.5))
    x, top = box.x0, box.y0
    for name, paint, width in zip(names, paints, widths):
        if x > box.x0 and x + width > box.x1:
            x, top = box.x0, top + line_h
        rect(slide, Box.at(x, top + (line_h - swatch) / 2, w=swatch, h=swatch), paint)
        write_label(
            slide,
            Box(x + swatch + 0.06, top, min(box.x1, x + width - _KEY_GAP), top + line_h),
            str(name),
            theme,
            size=size,
            colour=_muted(theme),
            anchor="middle",
        )
        x += width
    return box


def scale_top(slide, plot, theme, high, unit, size=KICKER_PT):
    """The top of the scale, written where the reader can see it, on a hairline.

    What replaces the per-bar values when the bars are too narrow to hold them.
    Without it a grouped plot has lengths and no way to turn one into a number.
    """
    hline(slide, plot.x0, plot.x1, plot.y0, _grid(theme))
    room = text_width([fmt(high, unit)], size, type_face(theme))
    write_label(
        slide,
        Box(plot.x0, plot.y0 - line_height(size), plot.x0 + room, plot.y0),
        fmt(high, unit),
        theme,
        size=size,
        colour=_muted(theme),
        anchor="bottom",
    )


# Where the leader's foot may sit along the label: centred on the segment first,
# then off to one side, then hard against one end. Past that the number has slid
# far enough off its own segment that the leader stops explaining it.
_CALLOUT_ALONG = (0.5, 0.15, 0.85, 0.0, 1.0)


def _callout(slide, theme, reading, bands, taken):
    """A reading written outside the segment it belongs to, on a leader.

    A segment nine per cent of a bar wide cannot hold "9%" inside it, and dropping
    the number leaves a length nobody in the room can turn back into a reading --
    which is the one failure every chart here exists to prevent. So the number
    comes out of the segment and sits in the air beside the stack, joined to it by
    a hairline.

    `bands` are the strips it may use, in order of preference and each with the
    point on the segment its leader starts from: under a row then over it, right
    of a column then left of it. Within a band the label slides along until it
    clears everything already placed, but never so far that the leader's foot
    leaves the label -- a number that has slid off its own segment says less than
    no number. Where nothing clears, nothing is written: two readings on top of
    each other is not a reading.

    Returns whether it was written.
    """
    width, height = text_width([reading], KICKER_PT, type_face(theme)), line_height(KICKER_PT)
    for side, band, (x, y) in bands:
        if band.w < width - 1e-9 or band.h < height - 1e-9:
            continue
        for along in _CALLOUT_ALONG:
            if side in ("left", "right"):
                top = min(max(y - height * along, band.y0), band.y1 - height)
                room = (
                    Box(band.x0, top, band.x0 + width, top + height)
                    if side == "right"
                    else Box(band.x1 - width, top, band.x1, top + height)
                )
                if not room.y0 - 1e-9 <= y <= room.y1 + 1e-9:
                    continue
            else:
                left = min(max(x - width * along, band.x0), band.x1 - width)
                room = (
                    Box(left, band.y0, left + width, band.y0 + height)
                    if side == "below"
                    else Box(left, band.y1 - height, left + width, band.y1)
                )
                if not room.x0 - 1e-9 <= x <= room.x1 + 1e-9:
                    continue
            if not _clear(room, taken):
                continue
            if side == "right":
                hline(slide, x, room.x0, y - HAIRLINE / 2, _muted(theme))
            elif side == "left":
                hline(slide, room.x1, x, y - HAIRLINE / 2, _muted(theme))
            elif side == "below":
                vline(slide, x, y, room.y0, _muted(theme))
            else:
                vline(slide, x, room.y1, y, _muted(theme))
            write_label(
                slide,
                room,
                reading,
                theme,
                size=KICKER_PT,
                colour=_muted(theme),
                align={"right": "left", "left": "right"}.get(side, "center"),
                anchor="middle",
            )
            taken.append(room)
            return True
    return False


def _along_the_line(marks, thickness=0.05):
    """Little boxes covering a polyline, so a label can be measured off it.

    A reading written straight over a point on a rising line lands on the segment
    leaving it, and the segment is the one thing on the plot that has to stay
    unbroken. There is nothing to compare a label against unless the line has a
    footprint, so it is given one: a run of boxes along each segment, thick enough
    that `_clear` can see them.
    """
    boxes = []
    for index in range(len(marks) - 1):
        start, end = marks[index], marks[index + 1]
        steps = max(1, int(math.ceil(abs(end[0] - start[0]) / 0.25)))
        for step in range(steps):
            near = (start[0] + (end[0] - start[0]) * step / steps, start[1] + (end[1] - start[1]) * step / steps)
            far = (
                start[0] + (end[0] - start[0]) * (step + 1) / steps,
                start[1] + (end[1] - start[1]) * (step + 1) / steps,
            )
            boxes.append(
                Box(
                    min(near[0], far[0]),
                    min(near[1], far[1]) - thickness / 2,
                    max(near[0], far[0]),
                    max(near[1], far[1]) + thickness / 2,
                )
            )
    return boxes


def _beside(slide, theme, reading, at_x, at_y, size, colour, region, taken, push=0.0):
    """A reading placed in the first free spot around a mark; -> whether it was written.

    Six tries: over the mark, over it and off to either side, then the same three
    under it. `push` is how far off to the side the side tries go -- nothing on a
    line plot, where the label only has to step out of the slope's way, and half a
    column where it has to clear the column the mark is sitting inside.

    Where none of the six clears, nothing is written and the caller reports the
    reading: two numbers on top of each other is not a reading, and a number
    written across the line it belongs to is worse than one written nowhere.
    """
    width, height = text_width([reading], size, type_face(theme)), line_height(size)
    for down in (-1, 1):
        for across in (0, -1, 1):
            centre = at_x + across * (width / 2 + push)
            left = min(max(centre - width / 2, region.x0), region.x1 - width)
            top = at_y - STEP / 2 - height if down < 0 else at_y + STEP / 2
            room = Box(left, top, left + width, top + height)
            if room.y0 < region.y0 - 1e-9 or room.y1 > region.y1 + 1e-9 or not _clear(room, taken):
                continue
            write_label(slide, room, reading, theme, size=size, colour=colour, align="center", anchor="middle")
            taken.append(room)
            return True
    return False


def _room(box, what, across=0.8, down=0.5):
    """Refuse a box that is under what `what` needs, and say by how much.

    Constants alone -- 1.0x0.9in for a column chart whether it carries three
    categories or twelve -- pass the check for exactly the chart that cannot be read:
    twelve columns in an inch is twelve labels a fortieth of an inch apart. Every
    caller works its floor out of its own data first, and the constants are only the
    last word where the data asks for less.
    """
    if box.w < across - 1e-9 or box.h < down - 1e-9:
        raise TooSmall(what, (box.w, box.h), (across, down))


# ------------------------------------------------------------- category charts


def column(slide, box, theme, data, *, accent=None, unit="", axis_max=None):
    """One value per category, as length up from a shared baseline.

    `data` is [(label, value)] or {label: value}. Three to eight of them: a ninth
    column is narrower than its own label, and past that -- or where the labels
    are phrases rather than words -- `horizontal_bar` gives each label a line.

    The value is written over its column when it fits and the scale's top is
    written on the plot when it does not, so there is always a number to read a
    length against; `drawn.readings_not_written` is which values that cost.
    Negative values hang below the baseline with their labels under them.
    """
    pairs = _pairs(data, "column")
    labels = [label for label, _ in pairs]
    values = [value for _, value in pairs]
    texts = [fmt(value, unit) for value in values]
    _room(box, "a column chart", max(1.0, _slot_needed(labels, type_face(theme)) * len(pairs)), 0.9)
    low, high = span(values, axis_max)
    slot = box.w / len(pairs)
    value_size = pick_size(texts, slot - 0.06, face=type_face(theme))
    label_size = pick_size(labels, slot - 0.06, face=type_face(theme)) or KICKER_PT
    label_rows = _rows_needed(labels, slot, label_size, face=type_face(theme))
    head = line_height(value_size) if value_size else line_height(KICKER_PT)
    under = line_height(value_size) if value_size and min(values) < 0 else 0.0
    foot = under + line_height(label_size) * label_rows
    plot = Box(box.x0, box.y0 + head, box.x1, box.y1 - foot)
    _room(plot, "a column chart's plot", 1.0, 0.35)
    at = linear(low, high, plot.y1, plot.y0)
    base = at(0.0)
    fills, inks = emphasis(theme, labels, accent)
    width = min(slot * BAR_SHARE, BAR_MAX)
    if value_size is None:
        scale_top(slide, plot, theme, high, unit)
    for index, (label, value) in enumerate(pairs):
        centre = plot.x0 + slot * (index + 0.5)
        top, bottom = min(base, at(value)), max(base, at(value))
        if value:
            bottom = max(bottom, top + MIN_LENGTH)
            rect(slide, Box(centre - width / 2, top, centre + width / 2, bottom), fills[index])
        if value_size:
            if value >= 0:
                room = Box(centre - slot / 2, top - line_height(value_size), centre + slot / 2, top)
                anchor = "bottom"
            else:
                room = Box(centre - slot / 2, bottom, centre + slot / 2, bottom + line_height(value_size))
                anchor = "top"
            write_label(slide, room, texts[index], theme, size=value_size, colour=inks[index], align="center", anchor=anchor)
        write_label(
            slide,
            Box(centre - slot / 2, plot.y1 + under, centre + slot / 2, box.y1),
            label,
            theme,
            size=label_size,
            colour=_muted(theme),
            align="center",
        )
    hline(slide, plot.x0, plot.x1, base, _muted(theme))
    return _drawn(plot, at, (value_size, label_size), readings=() if value_size else texts)


def horizontal_bar(slide, box, theme, data, *, accent=None, unit="", axis_max=None):
    """A ranking: one row per item, length out from a shared baseline.

    `data` is [(label, value)] or {label: value}, in the order they should read --
    sort it yourself, because which end of a ranking goes on top is an argument
    and not a default. Five to twelve rows; the label gets a column of its own,
    which is why this and not `column` when the labels are phrases.
    """
    pairs = _pairs(data, "horizontal_bar")
    labels = [label for label, _ in pairs]
    values = [value for _, value in pairs]
    texts = [fmt(value, unit) for value in values]
    row = box.h / len(pairs)
    size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
    _room(
        box,
        "a horizontal bar chart",
        max(1.6, _column_needed(labels, type_face(theme), 0.42, size)),
        max(0.5, _rows_high(len(pairs))),
    )
    low, high = span(values, axis_max)
    label_w = min(box.w * 0.42, text_width(labels, size, type_face(theme)) + 0.08)
    value_text = text_width(texts, size, type_face(theme))
    plot = Box(box.x0 + label_w, box.y0, box.x1 - value_text - LABEL_GAP, box.y1)
    _room(plot, "a horizontal bar chart's plot", 0.6, 0.4)
    at = linear(low, high, plot.x0, plot.x1)
    base = at(0.0)
    fills, inks = emphasis(theme, labels, accent)
    # One thickness and one baseline for every bar in the series, which is both
    # what makes the lengths comparable and what tells the band measurement that
    # these are readings rather than decoration.
    thickness = min(row * 0.56, 0.44)
    for index, (label, value) in enumerate(pairs):
        middle = box.y0 + row * (index + 0.5)
        write_label(
            slide,
            Box(box.x0, middle - row / 2, box.x0 + label_w - 0.08, middle + row / 2),
            label,
            theme,
            size=size,
            colour=_muted(theme),
            anchor="middle",
        )
        left, right = min(base, at(value)), max(base, at(value))
        if value:
            right = max(right, left + MIN_LENGTH)
            rect(slide, Box(left, middle - thickness / 2, right, middle + thickness / 2), fills[index])
        write_label(
            slide,
            Box(right + LABEL_GAP, middle - row / 2, right + LABEL_GAP + value_text, middle + row / 2),
            texts[index],
            theme,
            size=size,
            colour=inks[index],
            anchor="middle",
        )
    vline(slide, base, plot.y0, plot.y1, _muted(theme))
    return _drawn(plot, at, (size,))


def dot_plot(slide, box, theme, categories, series, *, accent=None, unit="", axis_max=None):
    """One row per item, a mark per series, read by position on one shared scale.

    `categories` is the rows, in the order they should read -- sort them yourself,
    for `horizontal_bar`'s reason. `series` is [(name, values)] or {name: values},
    each as long as `categories`.

    What this draws that `grouped_bar` cannot is a dozen rows of it. A bar wants a
    slot deep enough to still be a bar and a mark wants one line of type, so four
    series across twelve categories is a grouped plot that does not fit the page
    and a dot plot that does. It is also the form for several readings of *one*
    quantity -- the same benchmark reported four different ways, the same figure
    from four sources -- where bars standing side by side say the four are four
    different things.

    `accent` names a *series*, as on `grouped_bar`: what a plot of positions argues
    is that one of them sits where the others do not.

    Nothing here is a length -- every mark is a position and the reading is where
    it sits -- so the scale is snapped out to round ends rather than pulled down to
    zero, as `dumbbell`'s is, and both ends are written under the plot. `axis_max`
    raises the top.
    """
    labels = [str(name) for name in categories]
    names, columns = _series_rows(labels, series, "dot_plot")
    face = type_face(theme)
    alone = len(names) == 1
    flat = [value for column_values in columns for value in column_values]
    low, high = snap(min(flat), max(flat))
    if axis_max is not None:
        high = max(high, number(axis_max))
    ends = [fmt(low, unit), fmt(high, unit)]
    texts = [fmt(value, unit) for value in flat]
    key_h = 0.0 if alone else line_height(KICKER_PT) * _key_rows(names, box.w, face=face)
    foot = line_height(KICKER_PT)
    row = max(box.h - key_h - foot, 0.01) / len(labels)
    size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
    mark = DOT if alone else STEP * 1.4
    reach = mark / 2 + (LABEL_GAP + text_width(texts, size, face) if alone else 0.02)
    _room(
        box,
        "a dot plot",
        max(1.8, _column_needed(labels, face, 0.34, size), text_width(labels, size, face) + 0.08 + reach + 0.8),
        max(0.6, key_h + foot + _rows_high(len(labels))),
    )
    body = Box(box.x0, box.y0 + key_h, box.x1, box.y1 - foot)
    row = body.h / len(labels)
    label_w = min(body.w * 0.34, text_width(labels, size, face) + 0.08)
    plot = Box(body.x0 + label_w + mark / 2, body.y0, body.x1 - reach, body.y1)
    _room(plot, "a dot plot's rows", 0.8, 0.3)
    at = linear(low, high, plot.x0, plot.x1)
    fills = series_paints(theme, names, accent)
    chosen = _accented(names, accent)
    if key_h:
        key(slide, Box(box.x0, box.y0, box.x1, box.y0 + key_h), theme, names, fills)
    for index, label in enumerate(labels):
        middle = body.y0 + row * (index + 0.5)
        write_label(
            slide,
            Box(box.x0, middle - row / 2, box.x0 + label_w - 0.08, middle + row / 2),
            label,
            theme,
            size=size,
            colour=_muted(theme),
            anchor="middle",
        )
        readings = [column_values[index] for column_values in columns]
        # The row's own leader, the width of the plot. It is the one rule this draws
        # and it carries a reading: which marks belong to this name, which a dozen
        # rows of unruled marks stop saying by about the fifth.
        #
        # The width of the plot rather than out to the row's furthest mark, which is
        # what it used to be. A rule that stops at the mark is a length, and this
        # chart's whole claim is that nothing on it is one -- read against a scale
        # that does not start at zero, a rule to the mark is a length measured from
        # an arbitrary place. Three rows of 92.3, 92.5 and 93.1 came out with the
        # first row's rule a seventh of the last one's, which reads as a sevenfold
        # difference between numbers eight tenths of a point apart.
        hline(slide, body.x0 + label_w, plot.x1, middle - HAIRLINE / 2, _grid(theme))
        for order, value in enumerate(readings):
            disc(slide, at(value), middle, mark, fills[order])
        if alone:
            # Past the end of the rule, not beside the mark. `reach` already holds
            # this column open to the right of the plot, and a reading written next
            # to its own mark now has the row's rule running under the digits.
            # Written here the readings also line up down one edge, which is the
            # column of numbers the rows are being compared in.
            write_label(
                slide,
                Box(plot.x1 + LABEL_GAP, middle - row / 2, body.x1, middle + row / 2),
                fmt(readings[0], unit),
                theme,
                size=size,
                colour=fills[0] if chosen else theme["foreground"],
                anchor="middle",
            )
    hline(slide, plot.x0, plot.x1, body.y1 - HAIRLINE, _muted(theme))
    for value in (low, high):
        width = text_width(ends, KICKER_PT, face)
        left = min(max(at(value) - width / 2, box.x0), box.x1 - width)
        write_label(
            slide,
            Box(left, body.y1, left + width, box.y1),
            fmt(value, unit),
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            align="center",
        )
    return _drawn(plot, at, (size, KICKER_PT))


def grouped_bar(slide, box, theme, categories, series, *, accent=None, unit="", axis_max=None):
    """Two to four series compared across the same categories.

    `categories` is the labels along the bottom; `series` is [(name, values)] or
    {name: values}, each as long as `categories`. The gap goes between groups and
    never inside one, which is what says the columns in a group belong together.

    `accent` names a *series* here rather than a category -- by index or by name --
    because what a grouped plot argues is that one of the series behaves
    differently. Name nothing and the series take `chart_series` in order.
    """
    labels = [str(name) for name in categories]
    names, columns = _series_rows(labels, series, "grouped_bar")
    _room(box, "a grouped bar chart", max(1.6, _slot_needed(labels, type_face(theme)) * len(labels)), 1.1)
    flat = [value for column_values in columns for value in column_values]
    low, high = span(flat, axis_max)
    fills = series_paints(theme, names, accent)
    key_h = line_height(KICKER_PT) * _key_rows(names, box.w, face=type_face(theme))
    slot = box.w / len(labels)
    width = slot * GROUP_SHARE / len(names)
    texts = [fmt(value, unit) for value in flat]
    value_size = pick_size(texts, width - 0.02, face=type_face(theme))
    label_size = pick_size(labels, slot - 0.06, face=type_face(theme)) or KICKER_PT
    label_rows = _rows_needed(labels, slot, label_size, face=type_face(theme))
    foot = line_height(label_size) * label_rows
    head = key_h + 0.06 + line_height(value_size or KICKER_PT)
    plot = Box(box.x0, box.y0 + head, box.x1, box.y1 - foot)
    _room(plot, "a grouped bar chart's plot", 1.0, 0.4)
    key(slide, Box(box.x0, box.y0, box.x1, box.y0 + key_h), theme, names, fills)
    at = linear(low, high, plot.y1, plot.y0)
    base = at(0.0)
    if value_size is None:
        scale_top(slide, plot, theme, high, unit)
    for group, label in enumerate(labels):
        left = plot.x0 + slot * group + slot * (1 - GROUP_SHARE) / 2
        for index, column_values in enumerate(columns):
            value = column_values[group]
            bar = Box(left + width * index, min(base, at(value)), left + width * (index + 1), max(base, at(value)))
            if value:
                rect(slide, Box(bar.x0, bar.y0, bar.x1, max(bar.y1, bar.y0 + MIN_LENGTH)), fills[index])
            if value_size:
                write_label(
                    slide,
                    Box(bar.x0, bar.y0 - line_height(value_size), bar.x1, bar.y0),
                    fmt(value, unit),
                    theme,
                    size=value_size,
                    colour=_muted(theme),
                    align="center",
                    anchor="bottom",
                )
        write_label(
            slide,
            Box(plot.x0 + slot * group, plot.y1, plot.x0 + slot * (group + 1), box.y1),
            label,
            theme,
            size=label_size,
            colour=_muted(theme),
            align="center",
        )
    hline(slide, plot.x0, plot.x1, base, _muted(theme))
    return _drawn(plot, at, (value_size, label_size, KICKER_PT), readings=() if value_size else texts)


def stacked_bar(slide, box, theme, categories, series, *, accent=None, unit="", share=False, direction="column"):
    """Category totals divided into their parts.

    `series` is [(name, values)] or {name: values}, each as long as `categories`,
    and the parts stack in the order given. `share=True` normalises every category
    to its own whole, which is the form for "what the parts of one whole are" --
    one category and `direction="bar"` is the full-width 100% bar that replaces a
    pie.

    `direction` is "column" (stacks upward, categories along the bottom) or "bar"
    (stacks rightward, one labelled row each). Parts are negative-free: a stack of
    a positive and a negative is two claims wearing one shape.

    The parts run from the theme's deepest data paint to its palest readable one,
    in the order given, so a stack reads as one quantity divided rather than as
    several bands that happen to be touching.

    A part too small to hold its own number keeps it: the reading moves outside the
    stack on a hairline leader rather than being dropped, because a nine per cent
    sliver with nothing written on it is a share the reader cannot recover. Where
    there is no air to put it in either -- several slivers in a row, a stack with
    no room around it -- the reading goes, and `drawn.readings_not_written` says
    which, so a page never has to be counted off its own render to find out.
    """
    if direction not in ("column", "bar"):
        raise ValueError(f'direction is "column" or "bar", not {direction!r}')
    labels = [str(name) for name in categories]
    names, columns = _series_rows(labels, series, "stacked_bar")
    if any(value < 0 for column_values in columns for value in column_values):
        raise ValueError("a stack has no room for a negative part; draw the gains and the losses as a waterfall")
    totals = [sum(column_values[index] for column_values in columns) for index in range(len(labels))]
    fills = stack_paints(theme, names, accent)
    key_h = line_height(KICKER_PT) * _key_rows(names, box.w, face=type_face(theme))
    if direction == "column":
        _room(
            box,
            "a stacked bar chart",
            max(1.6, _slot_needed(labels, type_face(theme)) * len(labels)),
            max(0.9, key_h + 0.06 + 0.4),
        )
    else:
        _room(
            box,
            "a stacked bar chart",
            max(1.6, _column_needed(labels, type_face(theme), 0.32, LABEL_PT)),
            max(0.9, key_h + 0.06 + _rows_high(len(labels))),
        )
    body = Box(box.x0, box.y0 + key_h + 0.06, box.x1, box.y1)
    high = 1.0 if share else max(totals + [0.0]) or 1.0
    inks = [ink_on(paint, theme) for paint in fills]
    if direction == "column":
        slot = body.w / len(labels)
        label_size = pick_size(labels, slot - 0.06, face=type_face(theme)) or KICKER_PT
        foot = line_height(label_size) * _rows_needed(labels, slot, label_size, face=type_face(theme))
        head = line_height(KICKER_PT) if not share else 0.0
        plot = Box(body.x0, body.y0 + head, body.x1, body.y1 - foot)
        _room(plot, "a stacked bar chart's plot", 1.0, 0.4)
        key(slide, Box(box.x0, box.y0, box.x1, box.y0 + key_h), theme, names, fills)
        at = linear(0.0, high, plot.y1, plot.y0)
        # Uncapped, unlike the charts whose slot holds one bar: a stack's slot holds
        # the whole composition, and a single-category stack is the one column on the
        # plot -- held to a bar's width it reads as a ribbon rather than as the thing
        # the page is about.
        width = slot * BAR_SHARE
        pitch, bar_h = slot, 0.0
    else:
        row = body.h / len(labels)
        label_size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
        label_w = min(body.w * 0.32, text_width(labels, label_size, type_face(theme)) + 0.08)
        plot = Box(body.x0 + label_w, body.y0, body.x1, body.y1)
        _room(plot, "a stacked bar chart's plot", 0.8, 0.3)
        key(slide, Box(box.x0, box.y0, box.x1, box.y0 + key_h), theme, names, fills)
        at = linear(0.0, high, plot.x0, plot.x1)
        thickness = min(row * 0.56, 0.62)
        pitch, bar_h = row, thickness
    taken, deferred, steps = [], [], [label_size]
    for group, label in enumerate(labels):
        whole = totals[group] or 1.0
        running = 0.0
        for index, column_values in enumerate(columns):
            value = column_values[group] / whole if share else column_values[group]
            reading = f"{column_values[group] / whole * 100:.0f}%" if share else fmt(column_values[group], unit)
            if direction == "column":
                centre = plot.x0 + slot * (group + 0.5)
                segment = Box(centre - width / 2, at(running + value), centre + width / 2, at(running))
            else:
                middle = body.y0 + row * (group + 0.5)
                segment = Box(at(running), middle - thickness / 2, at(running + value), middle + thickness / 2)
            running += value
            if not column_values[group]:
                continue
            rect(slide, segment, fills[index])
            size = pick_size(
                [reading], segment.w - TEXT_INSET, LABEL_PT if direction == "bar" else KICKER_PT, type_face(theme)
            )
            if size and line_height(size) <= segment.h:
                steps.append(size)
                write_label(slide, segment, reading, theme, size=size, colour=inks[index], align="center", anchor="middle")
            else:
                deferred.append((reading, _outside(segment, box, plot, body, group, direction, pitch, bar_h)))
        if direction == "column":
            if not share:
                centre = plot.x0 + slot * (group + 0.5)
                crown = Box(
                    centre - slot / 2, at(totals[group]) - line_height(KICKER_PT), centre + slot / 2, at(totals[group])
                )
                write_label(
                    slide,
                    crown,
                    fmt(totals[group], unit),
                    theme,
                    size=KICKER_PT,
                    align="center",
                    anchor="bottom",
                )
                taken.append(crown)
            write_label(
                slide,
                Box(plot.x0 + slot * group, plot.y1, plot.x0 + slot * (group + 1), body.y1),
                label,
                theme,
                size=label_size,
                colour=_muted(theme),
                align="center",
            )
        else:
            middle = body.y0 + row * (group + 0.5)
            write_label(
                slide,
                Box(body.x0, middle - row / 2, body.x0 + label_w - 0.08, middle + row / 2),
                label,
                theme,
                size=label_size,
                colour=_muted(theme),
                anchor="middle",
            )
    # After every segment, so nothing drawn later paints over a callout, and in
    # one pass over all of them, so the second sliver in a row knows where the
    # first one put its number.
    lost = [reading for reading, bands in deferred if not _callout(slide, theme, reading, bands, taken)]
    if direction == "column":
        hline(slide, plot.x0, plot.x1, at(0.0), _muted(theme))
    elif len(labels) > 1:
        # One row is a single whole divided up, and a whole has no axis to start
        # from. Several rows share one, and then the axis is what says so.
        vline(slide, plot.x0, plot.y0, plot.y1, _muted(theme))
    return _drawn(plot, at, steps + [KICKER_PT], readings=lost)


def _outside(segment, box, plot, body, group, direction, pitch, bar_h):
    """The strips a segment's number may be written in, best first.

    Beside a column and under a row, because that is where a stack leaves air:
    the gap that separates one column from the next, and the space a row's
    thickness does not use. Each strip runs right up to the neighbouring *bar* --
    not to the neighbouring slot -- because the air between two bars belongs to
    both of them and a strip that stopped halfway across it is half as tall as it
    could be, which is the difference between a four per cent sliver keeping its
    number and losing it. The leader is what says which of the two the number is
    about, and `_callout` keeps the two from ever landing on each other.
    """
    if direction == "column":
        gutter = pitch * (1 - BAR_SHARE) / 2
        after = min(box.x1, plot.x0 + pitch * (group + 1) + gutter)
        before = max(box.x0, plot.x0 + pitch * group - gutter)
        centre = (segment.y0 + segment.y1) / 2
        return (
            ("right", Box(segment.x1 + CALLOUT_GAP, plot.y0, after, plot.y1), (segment.x1, centre)),
            ("left", Box(before, plot.y0, segment.x0 - CALLOUT_GAP, plot.y1), (segment.x0, centre)),
        )
    under = min(body.y1, body.y0 + pitch * (group + 1.5) - bar_h / 2)
    over = max(body.y0, body.y0 + pitch * (group - 0.5) + bar_h / 2)
    centre = (segment.x0 + segment.x1) / 2
    return (
        ("below", Box(plot.x0, segment.y1 + CALLOUT_GAP, plot.x1, under), (centre, segment.y1)),
        ("above", Box(plot.x0, over, plot.x1, segment.y0 - CALLOUT_GAP), (centre, segment.y0)),
    )


def marimekko(slide, box, theme, categories, series, *, accent=None, unit=""):
    """A whole split two ways at once: columns as wide as they are big, stacked inside.

    `categories` is the columns and `series` is [(name, values)] or {name: values},
    each as long as `categories` -- exactly what `stacked_bar` takes. The width is
    the difference: a column is as wide a share of the plot as its own total is of
    everything on the page, and is then stacked to its own whole.

    So a segment's *area* is what that part of that category is worth against the
    grand total, which is the reading equal-width columns throw away: the same
    quarter of a column is a very different number in the widest column and the
    narrowest. It is the shape a market argument has -- which segments are big, and
    who holds what inside each -- and drawing it as equal columns is what makes a
    small segment's winner look like the winner overall.

    Each column's own total is written under its name, because a width nobody can
    turn back into a number is half a chart. Negative parts are refused, for
    `stacked_bar`'s reason.
    """
    labels = [str(name) for name in categories]
    names, columns = _series_rows(labels, series, "marimekko")
    if any(value < 0 for column_values in columns for value in column_values):
        raise ValueError("a marimekko divides a whole into parts, so none of them may be negative")
    totals = [sum(column_values[index] for column_values in columns) for index in range(len(labels))]
    whole = sum(totals)
    if whole <= 0:
        raise ValueError("a marimekko needs a total above zero to divide up")
    face = type_face(theme)
    fills = stack_paints(theme, names, accent)
    inks = [ink_on(paint, theme) for paint in fills]
    key_h = line_height(KICKER_PT) * _key_rows(names, box.w, face=face)
    foot = line_height(KICKER_PT) * 2
    _room(
        box,
        "a marimekko",
        max(2.0, _slot_needed(labels, face) * len(labels)),
        max(1.3, key_h + 0.06 + foot + 0.7),
    )
    body = Box(box.x0, box.y0 + key_h + 0.06, box.x1, box.y1)
    plot = Box(body.x0, body.y0, body.x1, body.y1 - foot)
    _room(plot, "a marimekko's columns", 1.2, 0.5)
    key(slide, Box(box.x0, box.y0, box.x1, box.y0 + key_h), theme, names, fills)
    cells, unnamed, unread, steps = [], [], [], []
    left = plot.x0
    for group, label in enumerate(labels):
        width = plot.w * totals[group] / whole
        running = 0.0
        for index, column_values in enumerate(columns):
            share = column_values[group] / (totals[group] or 1.0)
            top = plot.y0 + plot.h * running
            running += share
            segment = Box(left, top, left + width, plot.y0 + plot.h * running)
            cells.append(((group, index), segment))
            if not column_values[group]:
                continue
            # The gutter goes on the segment rather than between the columns, so
            # every part of one column stays flush with the parts above and below
            # it and only the columns come apart.
            rect(slide, segment.inset(0.010, 0.0), fills[index])
            reading = f"{share * 100:.0f}%"
            size = pick_size([reading], segment.w - TEXT_INSET, KICKER_PT, face)
            if size and line_height(size) <= segment.h:
                steps.append(size)
                write_label(slide, segment, reading, theme, size=size, colour=inks[index], align="center", anchor="middle")
            else:
                unread.append(reading)
        under = Box(left, plot.y1, left + width, plot.y1 + line_height(KICKER_PT))
        if _fits([label], width - 0.04, KICKER_PT, face):
            write_label(slide, under, label, theme, size=KICKER_PT, colour=_muted(theme), align="center")
            write_label(
                slide,
                Box(left, under.y1, left + width, body.y1),
                fmt(totals[group], unit),
                theme,
                size=KICKER_PT,
                colour=theme["foreground"],
                align="center",
            )
        else:
            unnamed.append(label)
            unread.append(fmt(totals[group], unit))
        left += width
    hline(slide, plot.x0, plot.x1, plot.y1, _muted(theme))
    # `where(category, part)` is the centre of that segment by index, as a
    # heatmap's is: a segment's reading is an area, and an area has no position.
    def where(category, part):
        for keys, segment in cells:
            if keys == (int(category), int(part)):
                return (segment.x0 + segment.x1) / 2, (segment.y0 + segment.y1) / 2
        raise ValueError(f"a marimekko of {len(labels)} columns and {len(names)} parts has no ({category}, {part})")

    return _drawn(plot, where, steps + [KICKER_PT], names=unnamed, readings=unread)


def butterfly(slide, box, theme, data, *, sides=(), accent=None, unit=""):
    """Two mirrored sides on one shared axis: A against B, cost against revenue.

    `data` is [(label, left, right)] and the labels run down the middle, which is
    what makes the two wings comparable -- a reader follows one row across rather
    than matching two charts by eye. `sides` names the wings, written over them.
    Both wings share one scale, so a bar on the left and a bar on the right of the
    same length are the same number.
    """
    rows = _triples(data, "butterfly", "right")
    labels = [label for label, _, _ in rows]
    lefts = [value for _, value, _ in rows]
    rights = [value for _, _, value in rows]
    high = max([abs(value) for value in lefts + rights] or [1.0]) or 1.0
    row = box.h / len(rows)
    size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
    head = line_height(KICKER_PT) if sides else 0.0
    _room(box, "a butterfly chart", 2.4, max(0.8, head + _rows_high(len(rows))))
    label_w = min(box.w * 0.30, text_width(labels, size, type_face(theme)) + 0.06)
    value_text = text_width([fmt(value, unit) for value in lefts + rights], size, type_face(theme))
    value_w = value_text + LABEL_GAP
    wing = (box.w - label_w) / 2 - value_w
    if wing <= 0.3:
        # The same refusal as `_room`, in the terms this chart is short in: it is
        # the wing rather than the box that has run out, and the box has to grow by
        # twice the shortfall for both of them to get it. Carried as a `TooSmall` so
        # `the_smallest_box_a_chart_needs` can answer for this chart too -- the
        # message is what an author reads and `short` is what a program acts on.
        raise TooSmall(
            f"a butterfly chart's wings have {wing:.2f}in; give it a wider box or shorter labels",
            (box.w, box.h),
            (box.w + 2 * (0.3 - wing) + 0.01, box.h),
        )
    middle_x = box.x0 + value_w + wing + label_w / 2
    body = Box(box.x0, box.y0 + head, box.x1, box.y1)
    row = body.h / len(rows)
    fills, inks = emphasis(theme, labels, accent)
    thickness = min(row * 0.56, 0.44)
    for index, (label, left_value, right_value) in enumerate(rows):
        centre = body.y0 + row * (index + 0.5)
        write_label(
            slide,
            Box(middle_x - label_w / 2, centre - row / 2, middle_x + label_w / 2, centre + row / 2),
            label,
            theme,
            size=size,
            colour=_muted(theme),
            align="center",
            anchor="middle",
        )
        for value, outward in ((left_value, -1), (right_value, 1)):
            edge = middle_x + outward * label_w / 2
            end = edge + outward * wing * abs(value) / high
            if value:
                bar = Box(min(edge, end), centre - thickness / 2, max(edge, end), centre + thickness / 2)
                rect(slide, Box(bar.x0, bar.y0, max(bar.x1, bar.x0 + MIN_LENGTH), bar.y1), fills[index])
            reading = fmt(value, unit)
            room = (
                Box(end - LABEL_GAP - value_text, centre - row / 2, end - LABEL_GAP, centre + row / 2)
                if outward < 0
                else Box(end + LABEL_GAP, centre - row / 2, end + LABEL_GAP + value_text, centre + row / 2)
            )
            write_label(
                slide,
                room,
                reading,
                theme,
                size=size,
                colour=inks[index],
                align="right" if outward < 0 else "left",
                anchor="middle",
            )
    for index, name in enumerate(list(sides)[:2]):
        outward = -1 if index == 0 else 1
        edge = middle_x + outward * label_w / 2
        room = Box(min(edge, edge + outward * wing), box.y0, max(edge, edge + outward * wing), box.y0 + head)
        write_label(
            slide,
            room,
            str(name),
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            align="center",
            anchor="middle",
            bold=True,
        )
    # `where(value)` puts a positive reading on the right wing and a negative one on
    # the left, which is the one reading of "where does this number sit" that a
    # mirrored pair of scales can answer.
    return _drawn(
        body,
        lambda value: middle_x + (1 if float(value) >= 0 else -1) * (label_w / 2 + wing * abs(float(value)) / high),
        (size, KICKER_PT if sides else None),
    )


# ------------------------------------------------------------ progress charts


def progress_bar(slide, box, theme, data, *, accent=None):
    """How far along each of a few items is: a track, the share filled, the percent.

    `data` is [(label, share)] or {label: share}, where a share reads as "76%",
    "0.76" or 76 -- all the same three quarters. A string is written back as it
    was given, so "~76%" stays approximate.

    The empty part of the track is the half that carries the meaning: a filled
    length says "three quarters" only when the reader can see the whole.
    """
    items = list(data.items()) if hasattr(data, "items") else list(data)
    if not items:
        raise ValueError("progress_bar needs at least one item")
    labels = [str(entry[0]) for entry in items]
    shares = [_share(entry[1], entry[0]) for entry in items]
    texts = [entry[1] if isinstance(entry[1], str) else f"{_share(entry[1], entry[0]) * 100:g}%" for entry in items]
    row = box.h / len(items)
    size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
    _room(
        box,
        "a progress chart",
        max(1.6, _column_needed(labels, type_face(theme), 0.42, size)),
        max(0.4, _rows_high(len(items))),
    )
    label_w = min(box.w * 0.42, text_width(labels, size, type_face(theme)) + 0.08)
    track = Box(box.x0 + label_w, box.y0, box.x1 - text_width(texts, size, type_face(theme)) - LABEL_GAP, box.y1)
    _room(track, "a progress chart's tracks", 0.6, 0.3)
    fills, inks = emphasis(theme, labels, accent)
    thickness = min(row * 0.42, 0.30)
    for index, label in enumerate(labels):
        middle = box.y0 + row * (index + 0.5)
        write_label(
            slide,
            Box(box.x0, middle - row / 2, box.x0 + label_w - 0.08, middle + row / 2),
            label,
            theme,
            size=size,
            colour=_muted(theme),
            anchor="middle",
        )
        rect(slide, Box(track.x0, middle - thickness / 2, track.x1, middle + thickness / 2), _grid(theme))
        if shares[index] > 0:
            end = track.x0 + max(track.w * shares[index], MIN_LENGTH)
            rect(slide, Box(track.x0, middle - thickness / 2, end, middle + thickness / 2), fills[index])
        write_label(
            slide,
            Box(track.x1 + LABEL_GAP, middle - row / 2, box.x1, middle + row / 2),
            texts[index],
            theme,
            size=size,
            colour=inks[index],
            anchor="middle",
        )
    # `where(share)` reads the same three spellings the tracks do, so a marker at
    # "80%" lands on the same inch the 80% track ends at.
    return _drawn(track, lambda value: track.x0 + track.w * _share(value), (size,))


def bullet(slide, box, theme, data, *, accent=None, unit="", axis_max=None):
    """KPIs against their targets: a thin actual on a longer track, the target ticked.

    `data` is [(label, actual, target)]. Every row shares one scale, so three KPIs
    in different units want three of these rather than one -- or a share each, and
    `progress_bar`. The tick is where the target is; the length is where you are.
    """
    rows = _triples(data, "bullet", "target")
    labels = [label for label, _, _ in rows]
    actuals = [value for _, value, _ in rows]
    targets = [value for _, _, value in rows]
    texts = [fmt(value, unit) for value in actuals]
    row = box.h / len(rows)
    size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
    _room(
        box,
        "a bullet chart",
        max(1.8, _column_needed(labels, type_face(theme), 0.36, size)),
        max(0.5, _rows_high(len(rows))),
    )
    low, high = span(actuals + targets, axis_max)
    label_w = min(box.w * 0.36, text_width(labels, size, type_face(theme)) + 0.08)
    plot = Box(box.x0 + label_w, box.y0, box.x1 - text_width(texts, size, type_face(theme)) - LABEL_GAP, box.y1)
    _room(plot, "a bullet chart's tracks", 0.6, 0.3)
    at = linear(low, high, plot.x0, plot.x1)
    fills, inks = emphasis(theme, labels, accent)
    thickness = min(row * 0.34, 0.24)
    for index, (label, actual, target) in enumerate(rows):
        middle = box.y0 + row * (index + 0.5)
        write_label(
            slide,
            Box(box.x0, middle - row / 2, box.x0 + label_w - 0.08, middle + row / 2),
            label,
            theme,
            size=size,
            colour=_muted(theme),
            anchor="middle",
        )
        band = min(row * 0.62, 0.44)
        rect(slide, Box(plot.x0, middle - band / 2, plot.x1, middle + band / 2), _grid(theme))
        if actual:
            rect(
                slide,
                Box(plot.x0, middle - thickness / 2, max(at(actual), plot.x0 + MIN_LENGTH), middle + thickness / 2),
                fills[index],
            )
        vline(slide, at(target), middle - band / 2, middle + band / 2, theme["foreground"], thickness=0.024)
        write_label(
            slide,
            Box(plot.x1 + LABEL_GAP, middle - row / 2, box.x1, middle + row / 2),
            texts[index],
            theme,
            size=size,
            colour=inks[index],
            anchor="middle",
        )
    return _drawn(plot, at, (size,))


def dumbbell(slide, box, theme, data, *, sides=(), accent=None, unit="", axis_max=None):
    """How far each item moved between two states: a mark at each, the gap between.

    `data` is [(label, before, after)]. Five to ten rows. What the reader follows
    is the gap -- which is why the two states are marks joined by a line rather
    than two bars, where the eye compares the lengths from the baseline instead.
    `sides` names the two states over the first row's own marks.

    The axis does not start at zero, and the reason is that nothing on it is a
    length: both ends are positions, and the reading is the distance between them.
    Forced to zero, five rows running 38% to 78% sit in the right-hand third of the
    page with the gaps too short to compare -- which loses exactly what the form is
    for. The axis is snapped out to round
    ends instead, and `axis_max` still raises the top. `line`, `dot_plot` and
    `box_plot` are the same case and are read the same way.
    """
    rows = _triples(data, "dumbbell", "after")
    labels = [label for label, _, _ in rows]
    firsts = [value for _, value, _ in rows]
    seconds = [value for _, _, value in rows]
    low, high = snap(min(firsts + seconds), max(firsts + seconds))
    if axis_max is not None:
        high = max(high, number(axis_max))
    texts = [fmt(value, unit) for value in firsts + seconds]
    row = box.h / len(rows)
    size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
    _room(
        box,
        "a dumbbell chart",
        max(2.0, _column_needed(labels, type_face(theme), 0.32, size)),
        max(0.6, (line_height(KICKER_PT) if sides else 0.0) + _rows_high(len(rows))),
    )
    label_w = min(box.w * 0.32, text_width(labels, size, type_face(theme)) + 0.08)
    value_text = text_width(texts, size, type_face(theme))
    reach = value_text + DOT / 2 + LABEL_GAP
    head = line_height(KICKER_PT) if sides else 0.0
    body = Box(box.x0, box.y0 + head, box.x1, box.y1)
    row = body.h / len(rows)
    # The plot is inset by a value's width at both ends, so the reading beside the
    # outermost mark is on the page rather than a tenth of an inch past its edge.
    plot = Box(box.x0 + label_w + reach, body.y0, box.x1 - reach, body.y1)
    _room(plot, "a dumbbell chart's plot", 0.8, 0.3)
    at = linear(low, high, plot.x0, plot.x1)
    chosen = _accented(labels, accent)
    # Two states of one thing, in order, so two depths of one line rather than two
    # entries of the series palette: the palette's pale rung is a 2:1 tint, and the
    # "before" mark drawn in it was not on the page at all.
    pair = tuple(reversed(shades(theme, 2)))
    for index, (label, first, second) in enumerate(rows):
        middle = body.y0 + row * (index + 0.5)
        ink = theme["accent"] if index in chosen else _muted(theme)
        write_label(
            slide,
            Box(box.x0, middle - row / 2, box.x0 + label_w - 0.08, middle + row / 2),
            label,
            theme,
            size=size,
            colour=ink,
            anchor="middle",
        )
        hline(slide, at(first), at(second), middle - 0.013, ink, thickness=0.026)
        for value, paint in ((first, pair[0]), (second, pair[1])):
            disc(slide, at(value), middle, DOT, theme["accent"] if index in chosen else paint)
        left, right = (first, second) if at(first) <= at(second) else (second, first)
        write_label(
            slide,
            Box(at(left) - DOT / 2 - LABEL_GAP - value_text, middle - row / 2, at(left) - DOT / 2 - LABEL_GAP, middle + row / 2),
            fmt(left, unit),
            theme,
            size=size,
            colour=ink,
            align="right",
            anchor="middle",
        )
        write_label(
            slide,
            Box(at(right) + DOT / 2 + LABEL_GAP, middle - row / 2, at(right) + DOT / 2 + LABEL_GAP + value_text, middle + row / 2),
            fmt(right, unit),
            theme,
            size=size,
            colour=ink,
            align="left",
            anchor="middle",
        )
    # The two states named over the first row's own marks, which is the only place
    # a name can sit and still say which end it belongs to. Each name gets a mark
    # of its own in front of it, because the pale end of the line is a mark's
    # contrast and not a word's, and the name has to say which colour it is about.
    for index, name in enumerate(list(sides)[:2]):
        anchor_x = at(firsts[0] if index == 0 else seconds[0])
        width = text_width([name], KICKER_PT, type_face(theme)) + DOT
        left = min(max(box.x0, anchor_x - width / 2), box.x1 - width)
        disc(slide, left + DOT / 2, box.y0 + head / 2, DOT * 0.7, pair[index])
        write_label(
            slide,
            Box(left + DOT, box.y0, left + width, box.y0 + head),
            str(name),
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            anchor="middle",
            bold=True,
        )
    return _drawn(plot, at, (size, KICKER_PT if sides else None))


# --------------------------------------------------------- sequence and shape


def line(slide, box, theme, categories, series, *, accent=None, unit="", axis_max=None):
    """Which way one or more series moved along an ordered axis.

    `categories` is what the marks are read against, in order -- quarters,
    releases, the size of the corpus -- and `series` is [(name, values)] or
    {name: values}, each as long as `categories`. A mark at every value it was
    given and a straight segment between them: nothing is interpolated, so a
    series with a hole in it is two calls rather than one line drawn straight
    through a number nobody measured.

    `accent` names a *series*, as on `grouped_bar`, because what several lines
    argue is that one of them behaves differently from the rest.

    The axis does not start at zero, for `dumbbell`'s reason: every mark is a
    position and nothing on the plot is a length, and four releases running 88 to
    94 pulled down to zero are one flat line across the top. Both ends of the
    scale are written beside it, so a position can still be turned back into a
    number, and `axis_max` raises the top.

    One series has its value written over every mark with room for it. Several
    have the last of each written past the end of its own line, which is the
    direct label a key cannot give: eighteen numbers over three crossing lines is
    not a reading anybody in the room takes.
    """
    labels = [str(name) for name in categories]
    names, columns = _series_rows(labels, series, "line")
    face = type_face(theme)
    alone = len(names) == 1
    flat = [value for column_values in columns for value in column_values]
    low, high = snap(min(flat), max(flat))
    if axis_max is not None:
        high = max(high, number(axis_max))
    scale = [fmt(low, unit), fmt(high, unit)]
    ends = [fmt(column_values[-1], unit) for column_values in columns]
    left = text_width(scale, KICKER_PT, face) + 0.06
    right = 0.0 if alone else text_width(ends, KICKER_PT, face) + LABEL_GAP
    key_h = 0.0 if alone else line_height(KICKER_PT) * _key_rows(names, box.w, face=face)
    _room(
        box,
        "a line chart",
        max(1.8, left + right + _slot_needed(labels, face) * len(labels)),
        max(0.9, key_h + 0.06 + 0.6),
    )
    body = Box(box.x0, box.y0 + (key_h + 0.06 if key_h else 0.0), box.x1, box.y1)
    slot = (body.w - left - right) / len(labels)
    label_size = pick_size(labels, slot - 0.06, face=face) or KICKER_PT
    label_rows = _rows_needed(labels, slot, label_size, face=face)
    texts = [fmt(value, unit) for value in flat]
    value_size = pick_size(texts, slot - 0.06, face=face) if alone else None
    head = line_height(value_size or KICKER_PT) + (STEP / 2 if value_size else 0.0)
    plot = Box(body.x0 + left, body.y0 + head, body.x1 - right, body.y1 - line_height(label_size) * label_rows)
    _room(plot, "a line chart's plot", 1.0, 0.5)
    at = linear(low, high, plot.y1, plot.y0)
    fills = series_paints(theme, names, accent)
    if key_h:
        key(slide, Box(box.x0, box.y0, box.x1, box.y0 + key_h), theme, names, fills)
    vline(slide, plot.x0, plot.y0, plot.y1, _muted(theme))
    hline(slide, plot.x0, plot.x1, plot.y1, _muted(theme))
    # Kept, so a reading placed beside a mark near the left edge cannot land on the
    # number that says where the scale ends -- which is `scatter`'s rule about its
    # own axis readings, applied to the other plot that writes them.
    ruled = []
    for value in (low, high):
        y = at(value)
        room = Box(body.x0, y - line_height(KICKER_PT) / 2, plot.x0 - 0.06, y + line_height(KICKER_PT) / 2)
        write_label(slide, room, fmt(value, unit), theme, size=KICKER_PT, colour=_muted(theme), align="right", anchor="middle")
        ruled.append(_ink(room, fmt(value, unit), KICKER_PT, "right", "middle", face))
    for index, label in enumerate(labels):
        write_label(
            slide,
            Box(plot.x0 + slot * index, plot.y1, plot.x0 + slot * (index + 1), body.y1),
            label,
            theme,
            size=label_size,
            colour=_muted(theme),
            align="center",
        )
    lost, taken = [], list(ruled)
    region = Box(box.x0, body.y0, box.x1, plot.y1)
    for order, column_values in enumerate(columns):
        marks = [(plot.x0 + slot * (index + 0.5), at(value)) for index, value in enumerate(column_values)]
        if len(marks) > 1:
            poly(slide, marks, fills[order])
        for x, y in marks:
            disc(slide, x, y, STEP, fills[order])
        taken.extend(_along_the_line(marks))
        taken.extend(Box(x - STEP / 2, y - STEP / 2, x + STEP / 2, y + STEP / 2) for x, y in marks)
        if alone:
            for index, (x, y) in enumerate(marks):
                reading = fmt(column_values[index], unit)
                if value_size is None or not _beside(
                    slide, theme, reading, x, y, value_size, theme["foreground"], region, taken
                ):
                    lost.append(reading)
            continue
        # Beside the last mark rather than at the plot's edge: the marks sit at the
        # middles of their slots, so a label hung off the edge stands half a slot
        # clear of the line it names and reads as belonging to neither.
        x, y = marks[-1]
        room = Box(x + STEP / 2 + LABEL_GAP, y - line_height(KICKER_PT) / 2, box.x1, y + line_height(KICKER_PT) / 2)
        if _clear(room, taken):
            write_label(slide, room, ends[order], theme, size=KICKER_PT, colour=fills[order], anchor="middle")
            taken.append(room)
        else:
            lost.append(ends[order])
    # `where(value)` is the inch that reading sits at, which is what a target line
    # or a second series drawn over the same plot needs; the categories are evenly
    # pitched and their own labels say which is which.
    return _drawn(plot, at, (value_size, label_size, KICKER_PT), readings=lost)


def combo(slide, box, theme, data, curve, *, sides=(), accent=None, unit="", curve_unit="", axis_max=None):
    """One quantity as columns and a second, in its own unit, as a line over them.

    `data` is [(label, value)] or {label: value} -- the columns, on a scale that
    contains zero, because a column is a length. `curve` is one value per column,
    read against a second scale written down the right-hand side: revenue and the
    rate it grew at, volume and the share of it that converted, a count and a
    percentage. `sides` names the two scales over their own edges of the plot.

    Two units on one plot is the one case where a second scale is not a trick,
    and the way it stops being one is that both are written: the ends of the
    columns' scale down the left, the ends of the curve's down the right, so
    neither length nor position is a number the reader has to take on trust.

    The columns carry their own numbers inside their tops rather than over them --
    the air over a column is where the line is, and a number written into it
    belongs to neither. A column with no room for its number loses it to
    `drawn.readings_not_written`, and the top of the column scale is written on
    the plot so a length is still readable.

    `accent` names a column, by index or by label. The curve is the theme's accent,
    as `pareto`'s is: it is the second reading on the page and the one thing here
    that is not a length. Where the columns are already painted that colour -- a
    theme read off a template can have its accent and its first series paint be one
    colour -- the curve steps to the next series paint the columns are not using,
    because two units sharing one paint is the failure this form exists to avoid.
    """
    pairs = _pairs(data, "combo")
    labels = [label for label, _ in pairs]
    values = [value for _, value in pairs]
    readings = [number(value, labels[index] if index < len(labels) else None) for index, value in enumerate(curve)]
    if len(readings) != len(pairs):
        raise ValueError(f"combo takes one curve value per column, and got {len(readings)} for {len(pairs)}")
    face = type_face(theme)
    low, high = span(values, axis_max)
    curve_low, curve_high = snap(min(readings), max(readings))
    bar_scale = [fmt(low, unit), fmt(high, unit)]
    curve_scale = [fmt(curve_low, curve_unit), fmt(curve_high, curve_unit)]
    left = text_width(bar_scale, KICKER_PT, face) + 0.06
    right = text_width(curve_scale, KICKER_PT, face) + 0.06
    head = line_height(KICKER_PT) if sides else 0.0
    _room(
        box,
        "a combo chart",
        max(2.2, left + right + _slot_needed(labels, face) * len(pairs)),
        max(1.1, head + 0.9),
    )
    body = Box(box.x0, box.y0 + head, box.x1, box.y1)
    slot = (body.w - left - right) / len(pairs)
    label_size = pick_size(labels, slot - 0.06, face=face) or KICKER_PT
    label_rows = _rows_needed(labels, slot, label_size, face=face)
    curve_texts = [fmt(value, curve_unit) for value in readings]
    curve_size = pick_size(curve_texts, slot - 0.06, face=face)
    crown = line_height(curve_size or KICKER_PT) + STEP / 2
    plot = Box(body.x0 + left, body.y0 + crown, body.x1 - right, body.y1 - line_height(label_size) * label_rows)
    _room(plot, "a combo chart's plot", 1.4, 0.6)
    at = linear(low, high, plot.y1, plot.y0)
    up = linear(curve_low, curve_high, plot.y1, plot.y0)
    fills, _ = emphasis(theme, labels, accent)
    curve_paint = _second_paint(theme, fills)
    width = min(slot * BAR_SHARE, BAR_MAX)
    base = at(0.0)
    taken, lost = [], []
    for index, (label, value) in enumerate(pairs):
        centre = plot.x0 + slot * (index + 0.5)
        top, bottom = min(base, at(value)), max(base, at(value))
        bar = Box(centre - width / 2, top, centre + width / 2, max(bottom, top + MIN_LENGTH))
        if value:
            rect(slide, bar, fills[index])
        reading = fmt(value, unit)
        size = pick_size([reading], bar.w - TEXT_INSET, LABEL_PT, face)
        if value:
            taken.append(bar)
        if value and size and line_height(size) <= bar.h:
            write_label(
                slide,
                Box(bar.x0, bar.y0, bar.x1, bar.y0 + line_height(size)),
                reading,
                theme,
                size=size,
                colour=ink_on(fills[index], theme),
                align="center",
            )
        else:
            lost.append(reading)
        write_label(
            slide,
            Box(plot.x0 + slot * index, plot.y1, plot.x0 + slot * (index + 1), body.y1),
            label,
            theme,
            size=label_size,
            colour=_muted(theme),
            align="center",
        )
    hline(slide, plot.x0, plot.x1, base, _muted(theme))
    for value, y in ((low, plot.y1), (high, plot.y0)):
        room = Box(body.x0, y - line_height(KICKER_PT) / 2, plot.x0 - 0.06, y + line_height(KICKER_PT) / 2)
        write_label(slide, room, fmt(value, unit), theme, size=KICKER_PT, colour=_muted(theme), align="right", anchor="middle")
        taken.append(_ink(room, fmt(value, unit), KICKER_PT, "right", "middle", face))
    for value, y in ((curve_low, plot.y1), (curve_high, plot.y0)):
        reading = fmt(value, curve_unit)
        room = Box(plot.x1 + 0.06, y - line_height(KICKER_PT) / 2, box.x1, y + line_height(KICKER_PT) / 2)
        write_label(slide, room, reading, theme, size=KICKER_PT, colour=curve_paint, anchor="middle")
        taken.append(_ink(room, reading, KICKER_PT, "left", "middle", face))
    marks = [(plot.x0 + slot * (index + 0.5), up(value)) for index, value in enumerate(readings)]
    if len(marks) > 1:
        poly(slide, marks, curve_paint)
    for x, y in marks:
        disc(slide, x, y, STEP, curve_paint)
    if curve_size:
        # The curve's own footprint and the columns are both in `taken`, so a rate
        # written over a mark that is sitting inside a column steps out into the
        # gutter beside it rather than being set in the accent on a filled bar,
        # which is where the first render of this put two of six.
        crossings = taken + _along_the_line(marks)
        region = Box(box.x0, body.y0, box.x1, plot.y1)
        for index, (x, y) in enumerate(marks):
            if not _beside(
                slide,
                theme,
                curve_texts[index],
                x,
                y,
                curve_size,
                curve_paint,
                region,
                crossings,
                push=slot * BAR_SHARE / 2,
            ):
                lost.append(curve_texts[index])
    else:
        lost.extend(curve_texts)
    for index, name in enumerate(list(sides)[:2]):
        room = (
            Box(box.x0, box.y0, plot.x0 + slot, box.y0 + head)
            if index == 0
            else Box(plot.x1 - slot, box.y0, box.x1, box.y0 + head)
        )
        write_label(
            slide,
            room,
            str(name),
            theme,
            size=KICKER_PT,
            colour=_muted(theme) if index == 0 else curve_paint,
            align="left" if index == 0 else "right",
            anchor="middle",
            bold=True,
        )
    # `where(value)` is the columns' own scale. The curve is read against the
    # right-hand pair of readings, which is what makes the second unit honest.
    return _drawn(plot, at, (label_size, curve_size, KICKER_PT), readings=lost)


def waterfall(slide, box, theme, data, *, accent=None, unit="", totals=()):
    """How a starting value became an ending one: gains, losses, connectors.

    `data` is [(label, value)] where each value is the step's own contribution.
    `totals` names the positions that are levels rather than steps -- the opening
    and closing balances -- by index, negatives counting from the end, so
    `totals=(0, -1)` is the usual shape.

    The three roles are three depths of one shade line -- the levels deepest, the
    rises under them, the falls palest -- so the chart reads as one quantity being
    moved rather than as four colours. Not red-up-green-down: no palette here
    carries a red or a green, and up is not good in every column anyway. Position
    already says which way a step went.
    """
    pairs = _pairs(data, "waterfall")
    labels = [label for label, _ in pairs]
    _room(box, "a waterfall chart", max(1.6, _slot_needed(labels, type_face(theme)) * len(pairs)), 1.0)
    marks = set()
    for index in totals:
        position = int(index) if int(index) >= 0 else len(pairs) + int(index)
        if not 0 <= position < len(pairs):
            raise ValueError(f"totals names step {index}, and there are {len(pairs)}")
        marks.add(position)
    steps, running = [], 0.0
    for index, (label, value) in enumerate(pairs):
        if index in marks:
            steps.append((label, 0.0, value, value, True))
            running = value
        else:
            steps.append((label, running, running + value, value, False))
            running += value
    low, high = span([edge for _, start, end, _, _ in steps for edge in (start, end)])
    texts = [fmt(value, unit, sign=not is_total) for _, _, _, value, is_total in steps]
    slot = box.w / len(steps)
    value_size = pick_size(texts, slot - 0.06, face=type_face(theme))
    label_size = pick_size(labels, slot - 0.06, face=type_face(theme)) or KICKER_PT
    label_rows = _rows_needed(labels, slot, label_size, face=type_face(theme))
    head = line_height(value_size) if value_size else line_height(KICKER_PT)
    foot = line_height(label_size) * label_rows + (line_height(value_size) if value_size else 0.0)
    plot = Box(box.x0, box.y0 + head, box.x1, box.y1 - foot)
    _room(plot, "a waterfall chart's plot", 1.0, 0.4)
    at = linear(low, high, plot.y1, plot.y0)
    chosen = _accented(labels, accent)
    # By role, not by position. Reading `chart_series` in order gave the rise and
    # both levels the same deep paint and the fall a lavender at 2:1 on white --
    # and the fall was what the page was about. Quiet depths when something is
    # accented, so the accented step is the only colour on the plot.
    levels, rises, falls = shades(theme, 3, quiet=bool(chosen))
    width = min(slot * BAR_SHARE, BAR_MAX)
    if value_size is None:
        scale_top(slide, plot, theme, high, unit)
    previous = None
    for index, (label, start, end, value, is_total) in enumerate(steps):
        centre = plot.x0 + slot * (index + 0.5)
        top, bottom = min(at(start), at(end)), max(at(start), at(end))
        if index in chosen:
            paint = theme["accent"]
        elif is_total:
            paint = levels
        else:
            paint = rises if value >= 0 else falls
        rect(slide, Box(centre - width / 2, top, centre + width / 2, max(bottom, top + MIN_LENGTH)), paint)
        if previous is not None:
            hline(slide, previous[0], centre - width / 2, previous[1], _grid(theme))
        previous = (centre + width / 2, at(end))
        if value_size:
            rising = end >= start
            room = (
                Box(centre - slot / 2, top - line_height(value_size), centre + slot / 2, top)
                if rising
                else Box(centre - slot / 2, bottom, centre + slot / 2, bottom + line_height(value_size))
            )
            write_label(
                slide,
                room,
                texts[index],
                theme,
                size=value_size,
                colour=theme["accent"] if index in chosen else _muted(theme),
                align="center",
                anchor="bottom" if rising else "top",
            )
        write_label(
            slide,
            Box(centre - slot / 2, box.y1 - line_height(label_size) * label_rows, centre + slot / 2, box.y1),
            label,
            theme,
            size=label_size,
            colour=_muted(theme),
            align="center",
        )
    hline(slide, plot.x0, plot.x1, at(0.0), _muted(theme))
    return _drawn(plot, at, (value_size, label_size), readings=() if value_size else texts)


def pareto(slide, box, theme, data, *, accent=None, unit="", threshold=0.8):
    """Which few items account for most of the total: bars, and the cumulative share.

    `data` is [(label, value)] or {label: value}, sorted here into descending
    order because a Pareto that is not sorted is not one. The polyline is the
    running share of the whole, read against the right-hand axis; `threshold` is
    the line drawn across it (`None` for none).
    """
    pairs = sorted(_pairs(data, "pareto"), key=lambda entry: -entry[1])
    if any(value < 0 for _, value in pairs):
        raise ValueError("a pareto chart adds its values up, so none of them may be negative")
    labels = [label for label, _ in pairs]
    values = [value for _, value in pairs]
    whole = sum(values) or 1.0
    cumulative, running = [], 0.0
    for value in values:
        running += value
        cumulative.append(running / whole)
    right = text_width(["100%"], KICKER_PT, type_face(theme)) + 0.06
    _room(box, "a pareto chart", max(2.0, right + _slot_needed(labels, type_face(theme)) * len(pairs)), 1.1)
    slot = (box.w - right) / len(pairs)
    label_size = pick_size(labels, slot - 0.06, face=type_face(theme)) or KICKER_PT
    label_rows = _rows_needed(labels, slot, label_size, face=type_face(theme))
    head = line_height(KICKER_PT)
    plot = Box(box.x0, box.y0 + head, box.x1 - right, box.y1 - line_height(label_size) * label_rows)
    _room(plot, "a pareto chart's plot", 1.2, 0.5)
    low, high = span(values)
    at = linear(low, high, plot.y1, plot.y0)
    up = linear(0.0, 1.0, plot.y1, plot.y0)
    fills, _ = emphasis(theme, labels, accent)
    width = min(slot * BAR_SHARE, BAR_MAX)
    if threshold is not None:
        level = up(float(threshold))
        hline(slide, plot.x0, plot.x1, level, _grid(theme))
        write_label(
            slide,
            Box(plot.x1 + 0.06, level - line_height(KICKER_PT) / 2, box.x1, level + line_height(KICKER_PT) / 2),
            f"{float(threshold) * 100:g}%",
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            anchor="middle",
        )
    for index, (label, value) in enumerate(pairs):
        centre = plot.x0 + slot * (index + 0.5)
        rect(slide, Box(centre - width / 2, at(value), centre + width / 2, at(0.0)), fills[index])
        write_label(
            slide,
            Box(centre - slot / 2, plot.y1, centre + slot / 2, box.y1),
            label,
            theme,
            size=label_size,
            colour=_muted(theme),
            align="center",
        )
    curve = [(plot.x0 + slot * (index + 0.5), up(share)) for index, share in enumerate(cumulative)]
    if len(curve) > 1:
        poly(slide, curve, theme["accent"])
    for x, y in curve:
        disc(slide, x, y, 0.09, theme["accent"])
    write_label(
        slide,
        Box(plot.x1 + 0.06, curve[-1][1] - line_height(KICKER_PT) / 2, box.x1, curve[-1][1] + line_height(KICKER_PT) / 2),
        f"{cumulative[-1] * 100:.0f}%",
        theme,
        size=KICKER_PT,
        colour=theme["accent"],
        anchor="middle",
    )
    write_label(
        slide,
        Box(plot.x0, box.y0, plot.x0 + text_width([fmt(high, unit)], KICKER_PT, type_face(theme)), plot.y0),
        fmt(high, unit),
        theme,
        size=KICKER_PT,
        colour=_muted(theme),
        anchor="bottom",
    )
    hline(slide, plot.x0, plot.x1, at(0.0), _muted(theme))
    # `where(value)` is the bars' own scale. The cumulative curve is read against the
    # right-hand percentages, and `up` is not what an author overlays anything on.
    return _drawn(plot, at, (label_size, KICKER_PT))


def funnel(slide, box, theme, data, *, accent=None, unit=""):
    """Where an ordered run of stages loses what it started with.

    `data` is [(label, value)] or {label: value} in the order the stages happen.
    Never sorted here: the order is the process, and re-ordering it would draw a
    different one. Each stage is a bar centred on one axis and as wide as its own
    value, so the shape of the narrowing is the shape of the loss -- and a stage
    that is nearly as wide as the one above it says "almost nobody was lost here"
    without a number being read at all.

    What each stage kept of the one before it is written in the air between the two
    bars, which is the only strip on the page that belongs to a pair of stages
    rather than to either one of them. A stage that grew reads with a plus.

    Negative values are refused: a funnel is a quantity being whittled down, and a
    bar of minus four hundred is a claim this shape cannot make.
    """
    pairs = _pairs(data, "funnel")
    labels = [label for label, _ in pairs]
    values = [value for _, value in pairs]
    if any(value < 0 for value in values):
        raise ValueError("a funnel narrows a quantity, so none of its stages may be negative")
    face = type_face(theme)
    texts = [fmt(value, unit) for value in values]
    changes = [""]
    for index in range(1, len(values)):
        before = values[index - 1]
        changes.append(f"{(values[index] / before - 1) * 100:+.0f}%" if before else "")
    row = box.h / len(pairs)
    size = LABEL_PT if line_height(LABEL_PT) <= row * 0.6 else KICKER_PT
    value_w = text_width(texts, size, face) + LABEL_GAP
    _room(
        box,
        "a funnel",
        max(2.0, _column_needed(labels, face, 0.30, size), text_width(labels, size, face) + 0.08 + value_w + 0.9),
        # The gap between two bars has to hold one line of type, because that is
        # where the drop between them is written and a drop nobody can read is a
        # funnel drawn as decoration.
        max(0.9, line_height(KICKER_PT) / 0.40 * len(pairs)),
    )
    label_w = min(box.w * 0.30, text_width(labels, size, face) + 0.08)
    plot = Box(box.x0 + label_w, box.y0, box.x1 - value_w, box.y1)
    _room(plot, "a funnel's stages", 0.8, 0.4)
    high = max(values) or 1.0
    centre_x = (plot.x0 + plot.x1) / 2
    # One quantity narrowing is one thing divided, so the stages step down a shade
    # line rather than taking the series palette, for `stacked_bar`'s reason.
    fills = stack_paints(theme, labels, accent)
    chosen = _accented(labels, accent)
    thickness = min(row * 0.60, 0.70)
    for index, (label, value) in enumerate(pairs):
        middle = box.y0 + row * (index + 0.5)
        write_label(
            slide,
            Box(box.x0, middle - row / 2, box.x0 + label_w - 0.08, middle + row / 2),
            label,
            theme,
            size=size,
            colour=_muted(theme),
            anchor="middle",
        )
        half = plot.w * (value / high) / 2
        if value:
            rect(
                slide,
                Box(
                    min(centre_x - half, centre_x - MIN_LENGTH / 2),
                    middle - thickness / 2,
                    max(centre_x + half, centre_x + MIN_LENGTH / 2),
                    middle + thickness / 2,
                ),
                fills[index],
            )
        write_label(
            slide,
            Box(plot.x1 + LABEL_GAP, middle - row / 2, box.x1, middle + row / 2),
            texts[index],
            theme,
            size=size,
            colour=theme["accent"] if index in chosen else theme["foreground"],
            anchor="middle",
        )
        if changes[index]:
            write_label(
                slide,
                Box(plot.x0, box.y0 + row * (index - 0.5) + thickness / 2, plot.x1, middle - thickness / 2),
                changes[index],
                theme,
                size=KICKER_PT,
                colour=_muted(theme),
                align="center",
                anchor="middle",
            )
    # `where(value)` is the half-width a stage of that size reaches out to on
    # either side of the axis, which is what a note hung off one stage needs.
    return _drawn(plot, lambda value: centre_x + plot.w * (number(value) / high) / 2, (size, KICKER_PT))


def histogram(slide, box, theme, data, *, accent=None, bins=None, unit=""):
    """How observations spread across numeric bins: contiguous, no gaps.

    `data` is either the observations themselves -- a flat sequence of numbers,
    binned here into equal-width bins -- or bins you have already counted, as
    [(label, count)]. `bins` sets how many; the default is about the square root
    of the count, between five and twelve.

    Contiguous rectangles rather than the separated columns of a category chart,
    because the axis underneath is continuous and a gap would say it is not.
    """
    items = list(data.items()) if hasattr(data, "items") else list(data)
    if not items:
        raise ValueError("histogram needs some observations")
    _room(box, "a histogram", 1.4, 0.9)
    if isinstance(items[0], (list, tuple)):
        pairs = _pairs(items, "histogram")
        labels = [label for label, _ in pairs]
        counts = [value for _, value in pairs]
        edges = None
    else:
        readings = sorted(number(value) for value in items)
        count = int(bins) if bins else max(5, min(12, int(round(math.sqrt(len(readings))))))
        low, high = readings[0], readings[-1]
        if high - low < 1e-12:
            high = low + 1.0
        step = (high - low) / count
        counts = [0] * count
        for value in readings:
            index = min(count - 1, int((value - low) / step))
            counts[index] += 1
        edges = [low + step * index for index in range(count + 1)]
        labels = [fmt(edge, unit) for edge in edges]
    _, top = span(counts)
    slot = box.w / len(counts)
    tick_size = KICKER_PT
    foot = line_height(tick_size)
    value_size = pick_size([fmt(value) for value in counts], slot - 0.04, face=type_face(theme))
    head = line_height(value_size) if value_size else line_height(tick_size)
    plot = Box(box.x0, box.y0 + head, box.x1, box.y1 - foot)
    _room(plot, "a histogram's plot", 1.0, 0.4)
    at = linear(0.0, top, plot.y1, plot.y0)
    fills, _ = emphasis(theme, [str(index) for index in range(len(counts))], accent)
    if value_size is None:
        scale_top(slide, plot, theme, top, "")
    for index, value in enumerate(counts):
        bar = Box(plot.x0 + slot * index, at(value), plot.x0 + slot * (index + 1), at(0.0))
        rect(slide, bar, fills[index])
        # A hairline in the ground colour between bins, so contiguous rectangles
        # still read as a count each without a gap that would break the axis.
        if index:
            vline(slide, bar.x0, bar.y0, bar.y1, theme.get("background", "#FFFFFF"), thickness=0.012)
        if value_size:
            write_label(
                slide,
                Box(bar.x0, bar.y0 - line_height(value_size), bar.x1, bar.y0),
                fmt(value),
                theme,
                size=value_size,
                colour=_muted(theme),
                align="center",
                anchor="bottom",
            )
    if edges is None:
        for index, label in enumerate(labels):
            write_label(
                slide,
                Box(plot.x0 + slot * index, plot.y1, plot.x0 + slot * (index + 1), box.y1),
                label,
                theme,
                size=tick_size,
                colour=_muted(theme),
                align="center",
            )
    else:
        # Edge labels, not bin labels: an edge belongs between two bars. Every
        # other one when they would touch, so a dense histogram still has a ruler.
        #
        # Each reading gets the width it sets and slides along the foot to stay
        # inside the region, which is what `dot_plot` and `box_plot` do with the
        # ends of their scales. Clipping it to the region instead -- which is what
        # this did -- left the first edge half a slot, and half a slot at this
        # module's floor is 0.06in, under one digit: a box no single character fits
        # is one the wrapper cannot break a word into, and the draw stopped
        # answering rather than crowding the label.
        needed = text_width(labels, tick_size, type_face(theme))
        every = max(1, int(math.ceil(needed / slot)))
        width = min(box.w, needed)
        for index, label in enumerate(labels):
            if index % every:
                continue
            left = min(max(plot.x0 + slot * index - width / 2, box.x0), box.x1 - width)
            write_label(
                slide,
                Box(left, plot.y1, left + width, box.y1),
                label,
                theme,
                size=tick_size,
                colour=_muted(theme),
                align="center",
            )
    hline(slide, plot.x0, plot.x1, at(0.0), _muted(theme))
    # `where(count)` is the height of that many observations, not a bin edge: the
    # bins are contiguous and their own labels already say where the edges are.
    return _drawn(
        plot, at, (value_size, tick_size), readings=() if value_size else [fmt(value) for value in counts]
    )


def _quantile(readings, share):
    """The reading `share` of the way through a sorted list, interpolated."""
    if len(readings) == 1:
        return readings[0]
    position = (len(readings) - 1) * share
    below = int(math.floor(position))
    above = min(below + 1, len(readings) - 1)
    return readings[below] + (readings[above] - readings[below]) * (position - below)


def _five(entry):
    """(group, low, q1, median, q3, high, outliers) from either shape a box takes."""
    parts = list(entry)
    if len(parts) == 2 and not isinstance(parts[1], (str, bytes)) and hasattr(parts[1], "__iter__"):
        label = str(parts[0])
        readings = sorted(number(value, label) for value in parts[1])
        if not readings:
            raise ValueError(f"box_plot was given no observations for {label!r}")
        first, middle, third = (_quantile(readings, share) for share in (0.25, 0.5, 0.75))
        # Tukey's fence. An outlier drawn as the end of the whisker is a
        # distribution that looks four times as wide as it is, which is the one
        # thing a reader takes off this shape without reading a number.
        fence = (third - first) * 1.5
        inside = [value for value in readings if first - fence <= value <= third + fence]
        low, high = (min(inside), max(inside)) if inside else (readings[0], readings[-1])
        return label, low, first, middle, third, high, [value for value in readings if value < low or value > high]
    if len(parts) == 6:
        label = str(parts[0])
        five = [number(value, label) for value in parts[1:]]
        if any(five[index] > five[index + 1] for index in range(4)):
            raise ValueError(f"a box for {label!r} runs low, q1, median, q3, high, and {five} does not")
        return (label, *five, [])
    raise ValueError(f"box_plot takes (group, [values]) or (group, low, q1, median, q3, high) rows, not {entry!r}")


def box_plot(slide, box, theme, data, *, accent=None, unit="", axis_max=None):
    """How a distribution sits per group: the middle half, the median, the reach.

    `data` is either the observations themselves as [(group, [values])], or the
    five numbers already worked out as [(group, low, q1, median, q3, high)] --
    the same two ways in that `histogram` takes its bins.

    Given observations, the quartiles are interpolated, the whiskers reach to the
    furthest observation within one and a half inter-quartile ranges of the box,
    and anything past that gets a mark of its own instead of being swallowed into
    the reach. Drawing an outlier as the end of a whisker makes a tight
    distribution look four times as wide as it is, which is exactly the reading
    this form exists to give and a bar of averages cannot.

    Nothing here is a length: the scale is snapped out to round ends, both of them
    are written under the plot, and `axis_max` raises the top.
    """
    rows = [_five(entry) for entry in list(data)]
    if not rows:
        raise ValueError("box_plot needs at least one group")
    labels = [row[0] for row in rows]
    face = type_face(theme)
    spread = [value for row in rows for value in row[1:6]] + [value for row in rows for value in row[6]]
    low, high = snap(min(spread), max(spread))
    if axis_max is not None:
        high = max(high, number(axis_max))
    ends = [fmt(low, unit), fmt(high, unit)]
    medians = [fmt(row[3], unit) for row in rows]
    foot = line_height(KICKER_PT)
    row_h = max(box.h - foot, 0.01) / len(rows)
    size = LABEL_PT if line_height(LABEL_PT) <= row_h else KICKER_PT
    value_w = text_width(medians, size, face) + LABEL_GAP
    _room(
        box,
        "a box plot",
        max(2.0, _column_needed(labels, face, 0.30, size), text_width(labels, size, face) + 0.08 + value_w + 0.9),
        max(0.7, foot + _rows_high(len(rows))),
    )
    label_w = min(box.w * 0.30, text_width(labels, size, face) + 0.08)
    body = Box(box.x0, box.y0, box.x1, box.y1 - foot)
    row_h = body.h / len(rows)
    plot = Box(body.x0 + label_w, body.y0, body.x1 - value_w, body.y1)
    _room(plot, "a box plot's rows", 0.9, 0.3)
    at = linear(low, high, plot.x0, plot.x1)
    fills, inks = emphasis(theme, labels, accent)
    thickness = min(row_h * 0.52, 0.42)
    for index, (label, least, first, middle_value, third, most, outliers) in enumerate(rows):
        middle = body.y0 + row_h * (index + 0.5)
        write_label(
            slide,
            Box(box.x0, middle - row_h / 2, box.x0 + label_w - 0.08, middle + row_h / 2),
            label,
            theme,
            size=size,
            colour=_muted(theme),
            anchor="middle",
        )
        hline(slide, at(least), at(most), middle - HAIRLINE / 2, _muted(theme))
        for reach in (least, most):
            vline(slide, at(reach), middle - thickness / 3, middle + thickness / 3, _muted(theme))
        # The median goes down before the box and again inside it. Drawn only
        # inside, in whichever ink reads on the fill, it comes out as a gap and the
        # box reads as two boxes; the nubs left standing either side of the fill
        # are what say it is one box with a rule across it.
        vline(
            slide,
            at(middle_value),
            middle - thickness / 2 - 0.05,
            middle + thickness / 2 + 0.05,
            theme["foreground"],
            thickness=0.024,
        )
        quarters = Box(at(first), middle - thickness / 2, at(third), middle + thickness / 2)
        rect(slide, Box(quarters.x0, quarters.y0, max(quarters.x1, quarters.x0 + MIN_LENGTH), quarters.y1), fills[index])
        vline(
            slide,
            at(middle_value),
            middle - thickness / 2,
            middle + thickness / 2,
            ink_on(fills[index], theme),
            thickness=0.024,
        )
        for outlier in outliers:
            disc(slide, at(outlier), middle, STEP, _muted(theme))
        write_label(
            slide,
            Box(plot.x1 + LABEL_GAP, middle - row_h / 2, box.x1, middle + row_h / 2),
            fmt(middle_value, unit),
            theme,
            size=size,
            colour=inks[index],
            anchor="middle",
        )
    hline(slide, plot.x0, plot.x1, body.y1 - HAIRLINE, _muted(theme))
    for value in (low, high):
        width = text_width(ends, KICKER_PT, face)
        left = min(max(at(value) - width / 2, box.x0), box.x1 - width)
        write_label(
            slide,
            Box(left, body.y1, left + width, box.y1),
            fmt(value, unit),
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            align="center",
        )
    return _drawn(plot, at, (size, KICKER_PT))


def gantt(slide, box, theme, data, *, accent=None, ticks=None):
    """When each task runs and for how long, on one shared time axis.

    `data` is [(task, start, end)] where the two ends are ISO dates ("2026-03-01")
    or plain numbers -- weeks, sprints, days from kick-off. `ticks` is how many
    marks the axis is asked for (up to eight); the labels are derived from the
    dates themselves rather than invented, and a mark whose label would land on its
    neighbour's is dropped, named in `drawn.readings_not_written`.

    The one place a vertical line is drawn: without them a bar's left edge is a
    position nobody can turn back into a date. They are hairlines in the grid
    tone, under the bars, and they carry a reading -- which is the test the rest
    of the module applies to a gridline.
    """
    rows = []
    dated = False
    for entry in list(data):
        try:
            label, start, end = entry
        except (TypeError, ValueError):
            raise ValueError(f"gantt takes (task, start, end) rows, not {entry!r}") from None
        first, is_date = _moment(start)
        last, also_date = _moment(end)
        dated = dated or is_date or also_date
        if last < first:
            raise ValueError(f"task {label!r} ends before it starts")
        rows.append((str(label), first, last))
    if not rows:
        raise ValueError("gantt needs at least one task")
    labels = [label for label, _, _ in rows]
    low = min(first for _, first, _ in rows)
    high = max(last for _, _, last in rows)
    if high - low < 1e-9:
        high = low + 1.0
    row = box.h / len(rows)
    size = LABEL_PT if line_height(LABEL_PT) <= row else KICKER_PT
    head = line_height(KICKER_PT)
    marks = _ticks(low, high, dated, ticks)
    # The axis reading is the one label on this chart with no slot of its own, so
    # the box has to be wide enough for the pitch between two marks to hold one.
    # The pitch is between marks, so it is one gap fewer than there are marks. That
    # sizes the middle of the ruler; the two clamped ends are `_thinned_ticks`.
    ruler = text_width([label for _, label in marks], KICKER_PT, type_face(theme)) * (len(marks) - 1)
    _room(
        box,
        "a gantt chart",
        max(
            2.0,
            _column_needed(labels, type_face(theme), 0.34, size),
            text_width(labels, size, type_face(theme)) + 0.08 + ruler,
        ),
        max(0.7, head + _rows_high(len(rows))),
    )
    label_w = min(box.w * 0.34, text_width(labels, size, type_face(theme)) + 0.08)
    body = Box(box.x0, box.y0 + head, box.x1, box.y1)
    row = body.h / len(rows)
    plot = Box(box.x0 + label_w, body.y0, box.x1, body.y1)
    _room(plot, "a gantt chart's plot", 0.9, 0.3)
    at = linear(low, high, plot.x0, plot.x1)
    ruled, unread = _thinned_ticks(marks, box, at, type_face(theme))
    for position, label, room in ruled:
        # The line goes with its label. A hairline nobody can turn back into a date
        # is the decoration this module refuses everywhere else, so a tick the
        # ruler had no room to letter is a tick that is not drawn either.
        vline(slide, at(position), plot.y0, plot.y1, _grid(theme))
        write_label(slide, room, label, theme, size=KICKER_PT, colour=_muted(theme), align="center", anchor="middle")
    fills, inks = emphasis(theme, labels, accent)
    thickness = max(0.06, min(max(row * 0.52, 0.23), row - 0.04, 0.42))
    for index, (label, first, last) in enumerate(rows):
        middle = body.y0 + row * (index + 0.5)
        write_label(
            slide,
            Box(box.x0, middle - row / 2, box.x0 + label_w - 0.08, middle + row / 2),
            label,
            theme,
            size=size,
            colour=inks[index],
            anchor="middle",
        )
        left, right = at(first), max(at(last), at(first) + MIN_LENGTH)
        rect(slide, Box(left, middle - thickness / 2, right, middle + thickness / 2), fills[index])
    # Inside the region, not on its edge: a hairline drawn with its top at the last
    # row's bottom hangs its own thickness past the box the author gave.
    hline(slide, plot.x0, plot.x1, plot.y1 - HAIRLINE, _muted(theme))
    # `where("2026-05-01")` reads the same two spellings the rows do, so a today
    # line goes on a schedule without the caller re-deriving the axis.
    return _drawn(plot, lambda moment: at(_moment(moment)[0]), (size, KICKER_PT), readings=unread)


def _moment(value):
    """A point on a time axis: a number as itself, an ISO date as its ordinal."""
    if isinstance(value, date):
        return float(value.toordinal()), True
    text = str(value).strip()
    try:
        return float(text), False
    except ValueError:
        pass
    try:
        return float(date.fromisoformat(text).toordinal()), True
    except ValueError:
        raise ValueError(f"a schedule runs between numbers or ISO dates, not {value!r}") from None


def _ticks(low, high, dated, count=None):
    """Up to eight marks along a time axis, labelled in the units it was given."""
    steps = max(2, min(8, int(count) if count else 5))
    marks = []
    for index in range(steps):
        position = low + (high - low) * index / (steps - 1)
        if dated:
            moment = date.fromordinal(int(round(position)))
            label = moment.isoformat()[:7] if high - low > 400 else moment.isoformat()[5:]
        else:
            label = fmt(position)
        marks.append((position, label))
    return marks


def _thinned_ticks(marks, box, at, face):
    """`marks` with the box each label goes in, minus the ones whose labels collide.

    A tick label is centred on its own tick and slid back inside the region -- a
    fixed 1.1in cropped where that leaves the box halves the last one and wraps it,
    so "09-01" comes out as "09-" over "01" on every serif theme. What the slide
    costs is the gap: the label on the last tick moves left by up to half its width
    and eats the air the even spacing had left it, and five marks in a 3.78in cell
    put "05-16" and "06-30" 0.275in into each other with neither of them reading.

    The pitch cannot see that -- it is the same for every pair, and only the marks
    at the ends are clamped -- so the thinning is measured where the labels will
    actually be set: `_em_width` of each of them, centred in the box it was slid
    into, against TICK_GAP. Measuring the boxes instead would drop a tick that
    still had a readable space in it, TEXT_INSET being margin rather than type.

    An axis loses a tick rather than the type: shrinking the two readings that
    happen to be at the end sets a ruler in two sizes, which reads as two rulers.
    Both ends survive where they can, being the readings the bars are drawn
    against, so a kept label the last one runs into is the one that goes -- and
    where even the two ends will not clear, only the first is written.

    -> ([(position, label, room)], the labels dropped), the second of which the
    caller reports as readings it could not write.
    """
    head = line_height(KICKER_PT)
    rooms, letters = [], []
    for position, label in marks:
        width = min(box.w, text_width([label], KICKER_PT, face))
        left = min(max(at(position) - width / 2, box.x0), box.x1 - width)
        rooms.append(Box(left, box.y0, left + width, box.y0 + head))
        set_in = _em_width(str(label), KICKER_PT, face)
        letters.append((left + (width - set_in) / 2, left + (width + set_in) / 2))
    last = len(marks) - 1
    kept = [0]
    for index in range(1, last):
        if letters[index][0] >= letters[kept[-1]][1] + TICK_GAP:
            kept.append(index)
    while len(kept) > 1 and letters[last][0] < letters[kept[-1]][1] + TICK_GAP:
        kept.pop()
    if last and letters[last][0] >= letters[kept[-1]][1] + TICK_GAP:
        kept.append(last)
    ruled = [(marks[index][0], marks[index][1], rooms[index]) for index in kept]
    written = set(kept)
    return ruled, tuple(label for index, (_, label) in enumerate(marks) if index not in written)


# The air between a mark on a timeline and the block of type written off it, with
# the stem crossing it. Tighter than LABEL_GAP for `_callout`'s reason.
STEM = 0.10


def milestone(slide, box, theme, data, *, accent=None):
    """Dated events on one line: what happened and when, with no duration.

    `data` is [(event, moment)] where the moment is an ISO date ("2026-03-01") or a
    plain number -- weeks, sprints, days from kick-off -- the same two spellings
    `gantt` reads. Sorted into time order here, because a timeline that is not in
    time order is not one.

    A `gantt` bar says how long something took. A milestone took no time, and drawn
    as a bar it comes out as the thinnest sliver the module will draw, which reads
    as a rounding error rather than as a date -- so it is a mark on the axis with
    its name and its date written off it instead. The blocks alternate above and
    below the line so two events a week apart do not have to share a strip, and
    where a name still will not clear, the mark keeps its place and the name is
    reported in `drawn.names_not_written`.
    """
    rows = []
    for entry in list(data):
        try:
            label, moment = entry
        except (TypeError, ValueError):
            raise ValueError(f"milestone takes (event, moment) rows, not {entry!r}") from None
        position, _ = _moment(moment)
        rows.append((str(label), position, str(moment)))
    if not rows:
        raise ValueError("milestone needs at least one event")
    rows.sort(key=lambda entry: entry[1])
    labels = [label for label, _, _ in rows]
    stamps = [stamp for _, _, stamp in rows]
    low = min(position for _, position, _ in rows)
    high = max(position for _, position, _ in rows)
    if high - low < 1e-9:
        high = low + 1.0
    face = type_face(theme)
    # Two strips, so the widest block only has to clear the one two events away.
    strips = max(1, int(math.ceil(len(rows) / 2)))
    name_size = pick_size(labels, box.w / strips - 0.06, face=face) or KICKER_PT
    width = max(text_width(labels, name_size, face), text_width(stamps, KICKER_PT, face))
    block = line_height(name_size) + line_height(KICKER_PT)
    _room(
        box,
        "a milestone timeline",
        max(2.0, width * strips),
        max(1.0, (block + STEM + DOT / 2) * 2),
    )
    # A timeline is a band and not a page: it is one line of marks with a block of
    # type either side of it, so a tall region is spent on more rows of blocks
    # rather than on air. Events crowd towards the axis first and step outwards
    # only where the row they wanted is taken.
    rungs = max(1, int((box.h / 2 - DOT / 2 - STEM) // block))
    axis_y = (box.y0 + box.y1) / 2
    plot = Box(box.x0 + width / 2, axis_y - DOT / 2, box.x1 - width / 2, axis_y + DOT / 2)
    _room(plot, "a milestone timeline's axis", 0.8, 0.0)
    at = linear(low, high, plot.x0, plot.x1)
    fills, inks = emphasis(theme, labels, accent)
    hline(slide, box.x0, box.x1, axis_y - HAIRLINE / 2, _muted(theme))
    taken, unnamed = [], []
    for index, (label, position, stamp) in enumerate(rows):
        x = at(position)
        disc(slide, x, axis_y, DOT, fills[index])
        left = min(max(x - width / 2, box.x0), box.x1 - width)
        sides = (True, False) if index % 2 == 0 else (False, True)
        for rung, above in ((rung, above) for rung in range(rungs) for above in sides):
            reach = DOT / 2 + STEM + block * rung
            top = axis_y - reach - block if above else axis_y + reach
            room = Box(left, top, left + width, top + block)
            if room.y0 < box.y0 - 1e-9 or room.y1 > box.y1 + 1e-9 or not _clear(room, taken):
                continue
            if above:
                vline(slide, x, room.y1, axis_y - DOT / 2, _muted(theme))
                name_room = Box(room.x0, room.y0, room.x1, room.y0 + line_height(name_size))
                stamp_room = Box(room.x0, room.y1 - line_height(KICKER_PT), room.x1, room.y1)
            else:
                vline(slide, x, axis_y + DOT / 2, room.y0, _muted(theme))
                stamp_room = Box(room.x0, room.y0, room.x1, room.y0 + line_height(KICKER_PT))
                name_room = Box(room.x0, room.y1 - line_height(name_size), room.x1, room.y1)
            write_label(slide, stamp_room, stamp, theme, size=KICKER_PT, colour=_muted(theme), align="center", anchor="middle")
            write_label(
                slide,
                name_room,
                label,
                theme,
                size=name_size,
                colour=inks[index],
                align="center",
                anchor="middle",
                bold=True,
            )
            taken.append(room)
            break
        else:
            unnamed.append(label)
    # `where("2026-05-01")` reads the same two spellings the events do, so a band
    # over one stretch of the line goes on without the caller re-deriving the axis.
    return _drawn(plot, lambda moment: at(_moment(moment)[0]), (name_size, KICKER_PT), names=unnamed)


# -------------------------------------------------------------- two dimensions


def _worst_row(row, short):
    """The worst aspect ratio in a row of areas laid along a side of length `short`."""
    total = sum(row)
    if total <= 0 or short <= 0 or min(row) <= 0:
        return float("inf")
    return max(short * short * max(row) / (total * total), (total * total) / (short * short * min(row)))


def _squarified(areas, room):
    """One rectangle per area, each with its share of `room`, kept as square as it can be.

    Bruls, Huizing and van Wijk's layout: largest first, a row at a time along
    whichever side of what is left is shorter, and a value joins the row it is on
    while that does not make the row's worst aspect ratio worse. Rows along the
    shorter side is the whole of why the tiles come out near-square rather than as
    the slivers slicing and dicing leaves -- and a sliver is a tile whose area
    nobody can read and whose name will not fit inside it.
    """
    order = sorted(range(len(areas)), key=lambda index: -areas[index])
    total = sum(areas)
    scale = (room.w * room.h) / total if total > 0 else 0.0
    sized = [areas[index] * scale for index in order]
    tiles = [None] * len(areas)
    left = Box(room.x0, room.y0, room.x1, room.y1)
    at = 0
    while at < len(sized):
        short = min(left.w, left.h)
        row = [sized[at]]
        end = at + 1
        while end < len(sized) and _worst_row(row + [sized[end]], short) <= _worst_row(row, short):
            row.append(sized[end])
            end += 1
        depth = sum(row) / short if short > 0 else 0.0
        offset = left.x0 if left.w <= left.h else left.y0
        for step, area in enumerate(row):
            along = area / depth if depth > 0 else 0.0
            if left.w <= left.h:
                tiles[order[at + step]] = Box(offset, left.y0, offset + along, left.y0 + depth)
            else:
                tiles[order[at + step]] = Box(left.x0, offset, left.x0 + depth, offset + along)
            offset += along
        left = (
            Box(left.x0, left.y0 + depth, left.x1, left.y1)
            if left.w <= left.h
            else Box(left.x0 + depth, left.y0, left.x1, left.y1)
        )
        at = end
    return tiles


def treemap(slide, box, theme, data, *, accent=None, unit=""):
    """How one whole divides into unequal parts, as area.

    `data` is [(label, value)] or {label: value}. Each part gets a rectangle whose
    *area* is its share, laid out largest first and kept as near square as the
    region allows, so twenty parts spanning three orders of magnitude are all on
    one page and the big ones are still the big ones. A 100% bar is the form for
    three or four parts of similar size; this is the form for the other case,
    where a bar's segments are too many and too uneven for any of them to carry
    its own name.

    Zero and negative parts are refused. A tile of no area is a part that is not
    on the page at all, and the reader has no way to tell it from one that was
    never passed.

    A tile too small for its name loses it, and a tile with room for a name but
    not for a number loses the number: `drawn.names_not_written` and
    `drawn.readings_not_written` say which, so a page never has to be counted off
    its own render.
    """
    pairs = _pairs(data, "treemap")
    labels = [label for label, _ in pairs]
    values = [value for _, value in pairs]
    if any(value <= 0 for value in values):
        raise ValueError("a treemap draws each part as an area, so none of them may be zero or negative")
    face = type_face(theme)
    _room(
        box,
        "a treemap",
        max(1.8, text_width(labels[:1], KICKER_PT, face) + 0.12),
        max(1.2, line_height(KICKER_PT) * 2 + 0.10),
    )
    tiles = _squarified(values, box)
    chosen = _accented(labels, accent)
    ranked = sorted(range(len(labels)), key=lambda index: -values[index])
    # One whole divided is one thing, so the parts step down a shade line by size
    # -- deepest for the biggest -- rather than taking the series palette, which
    # would say twenty unrelated things and lose the pale ones off the page.
    quiet = shades(theme, max(1, len(labels) - len(chosen)), quiet=True) if chosen else []
    loud = [] if chosen else shades(theme, len(labels))
    paints, taken = [None] * len(labels), 0
    for rank, index in enumerate(ranked):
        if index in chosen:
            paints[index] = theme["accent"]
        elif chosen:
            paints[index] = quiet[taken]
            taken += 1
        else:
            paints[index] = loud[rank]
    unnamed, unread, steps = [], [], []
    for index, tile in enumerate(tiles):
        rect(slide, tile.inset(0.010, 0.010), paints[index])
        room = tile.inset(0.06, 0.05)
        ink = ink_on(paints[index], theme)
        # A tile is as wide as its own share and no wider, so a name wraps before
        # it is given up. "EverOS stars" wanted 0.95in of a 0.83in tile and came
        # back unwritten -- on the one tile the page was accenting.
        size, lines = None, 1
        for step in (LABEL_PT, KICKER_PT):
            wrapped = _rows_needed([labels[index]], room.w, step, face=face)
            if (
                _widest([labels[index]], step, face) / wrapped <= room.w - TEXT_INSET
                and line_height(step) * wrapped <= room.h
            ):
                size, lines = step, wrapped
                break
        if size is None:
            unnamed.append(labels[index])
            unread.append(fmt(values[index], unit))
            continue
        steps.append(size)
        head = line_height(size) * lines
        write_label(slide, Box(room.x0, room.y0, room.x1, room.y0 + head), labels[index], theme, size=size, colour=ink)
        reading = fmt(values[index], unit)
        under = Box(room.x0, room.y0 + head, room.x1, room.y0 + head + line_height(KICKER_PT))
        if under.y1 <= room.y1 + 1e-9 and _fits([reading], room.w, KICKER_PT, face):
            write_label(slide, under, reading, theme, size=KICKER_PT, colour=ink)
        else:
            unread.append(reading)
    # `where(item)` is the centre of that part's tile, by index or by label, which
    # is the only reading a plot of areas has: the value is a size, not a position.
    def where(item):
        found = _accented(labels, item)
        if not found:
            raise ValueError(f"a treemap has no part called {item!r}")
        tile = tiles[min(found)]
        return (tile.x0 + tile.x1) / 2, (tile.y0 + tile.y1) / 2

    return _drawn(box, where, steps + [KICKER_PT], names=unnamed, readings=unread)


def heatmap(slide, box, theme, rows, columns, values, *, unit="", scale=None):
    """A value at every row-column intersection, as a tint.

    `rows` and `columns` are the labels; `values` is a row-major grid of numbers,
    one list per row. The tint is the module's own shade line -- the theme's soft
    accent at the low end, the paint it reserves for data at the high one -- and
    the ends of it are written under the grid once, which is what makes a tint
    readable at all. It is the one place the line runs all the way to its pale end:
    a cell near the low end of the scale should be nearly the page, and the legend
    is there to say so.
    `scale` fixes the (low, high) the tint is stretched over, for two heatmaps
    that must be read against each other.

    The number goes in the cell wherever it fits, because a tint is a ranking and
    a number is a reading, and a page usually wants both.
    """
    row_labels = [str(label) for label in rows]
    column_labels = [str(label) for label in columns]
    grid = []
    for index, line in enumerate(values):
        # By row label where there is one; the count is checked below, so a
        # grid with more rows than labels still reaches its own error.
        grid.append([number(value, row_labels[index] if index < len(row_labels) else None) for value in line])
    if not row_labels or not column_labels:
        raise ValueError("a heatmap needs at least one row and one column")
    if len(grid) != len(row_labels) or any(len(line) != len(column_labels) for line in grid):
        raise ValueError(
            f"a heatmap of {len(row_labels)} rows and {len(column_labels)} columns "
            f"needs that many values, and got {[len(line) for line in grid]}"
        )
    flat = [value for line in grid for value in line]
    low, high = (float(scale[0]), float(scale[1])) if scale else (min(flat), max(flat))
    if high - low < 1e-12:
        high = low + 1.0
    head = line_height(KICKER_PT)
    foot = line_height(KICKER_PT) + 0.10
    # A column head is written into one cell's width and neither shrinks nor wraps
    # to a second line the head was not cut for, so the grid has to be at least as
    # wide as its own headings; the row names have the same 0.30 cap every row chart
    # does. Six columns called "Q1 FY25" want 3.9in of grid, and the flat floor
    # asked for 1.8in of everything.
    _room(
        box,
        "a heatmap",
        max(
            1.8,
            _column_needed(row_labels, type_face(theme), 0.30, KICKER_PT),
            text_width(row_labels, KICKER_PT, type_face(theme))
            + 0.08
            + text_width(column_labels, KICKER_PT, type_face(theme)) * len(column_labels),
        ),
        max(1.0, head + foot + line_height(KICKER_PT) * len(row_labels)),
    )
    label_w = min(box.w * 0.30, text_width(row_labels, KICKER_PT, type_face(theme)) + 0.08)
    plot = Box(box.x0 + label_w, box.y0 + head, box.x1, box.y1 - foot)
    _room(plot, "a heatmap's grid", 0.8, 0.4)
    cell_w = plot.w / len(column_labels)
    cell_h = plot.h / len(row_labels)
    unwritten = []
    for index, label in enumerate(column_labels):
        write_label(
            slide,
            Box(plot.x0 + cell_w * index, box.y0, plot.x0 + cell_w * (index + 1), plot.y0),
            label,
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            align="center",
            anchor="bottom",
        )
    for r, line in enumerate(grid):
        write_label(
            slide,
            Box(box.x0, plot.y0 + cell_h * r, box.x0 + label_w - 0.08, plot.y0 + cell_h * (r + 1)),
            row_labels[r],
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            anchor="middle",
        )
        for c, value in enumerate(line):
            share = (value - low) / (high - low)
            cell = Box(plot.x0 + cell_w * c, plot.y0 + cell_h * r, plot.x0 + cell_w * (c + 1), plot.y0 + cell_h * (r + 1))
            tint = _shade(theme, share)
            rect(slide, cell.inset(0.010, 0.010), tint)
            reading = fmt(value, unit)
            size = pick_size([reading], cell.w - TEXT_INSET, KICKER_PT, type_face(theme))
            if size and line_height(size) <= cell.h:
                write_label(
                    slide,
                    cell,
                    reading,
                    theme,
                    size=size,
                    colour=ink_on(tint, theme),
                    align="center",
                    anchor="middle",
                )
            else:
                unwritten.append(reading)
    _legend(slide, Box(plot.x0, box.y1 - foot + 0.08, plot.x1, box.y1), theme, low, high, unit)
    # `where(column, row)` is the centre of that cell by index, which is what a ring
    # round the one cell a page is about needs and the only reading a grid of tints
    # has: the value itself is a colour, not a position.
    return _drawn(
        plot,
        lambda column, row: (plot.x0 + cell_w * (int(column) + 0.5), plot.y0 + cell_h * (int(row) + 0.5)),
        (KICKER_PT,),
        readings=unwritten,
    )


def _legend(slide, box, theme, low, high, unit):
    """The tint scale, shown once, with its two ends written under it."""
    steps = 5
    swatch = min(0.42, box.w / (steps * 3))
    height = min(0.11, box.h * 0.45)
    left = box.x1 - swatch * steps
    for index in range(steps):
        rect(
            slide,
            Box.at(left + swatch * index, box.y0, w=swatch, h=height),
            _shade(theme, index / (steps - 1)),
        )
    write_label(
        slide,
        Box(left - 1.2, box.y0 - 0.02, left - 0.06, box.y0 + height + 0.14),
        f"{fmt(low, unit)} to {fmt(high, unit)}",
        theme,
        size=KICKER_PT,
        colour=_muted(theme),
        align="right",
    )


def matrix_2x2(slide, box, theme, data, *, axes=(), quadrants=(), accent=None, limits=None):
    """Where items fall on two dimensions at once, with the quadrants named.

    `data` is [(label, x, y)]. `axes` names the two dimensions ("Impact",
    "Effort"); `quadrants` names the four grounds in reading order -- top-left,
    top-right, bottom-left, bottom-right. `limits` is (x_low, x_high, y_low,
    y_high) when the crossing point is a decided threshold rather than the middle
    of what happens to be plotted.

    The grounds alternate between the page's ground and the theme's surface, so
    the four quarters are visible without a colour being invented to tell them
    apart, and the cross is a hairline rather than a rule. Each name is written in
    the corner of its own quarter with the most room -- furthest from the crossing
    where the data leaves that corner free, which is where a name is least likely
    to be read as the next quadrant's, and out of the way of the marks where it
    does not.
    """
    rows = _triples(data, "matrix_2x2", "y")
    _room(box, "a 2x2 matrix", 2.0, 1.4)
    labels = [label for label, _, _ in rows]
    xs = [value for _, value, _ in rows]
    ys = [value for _, _, value in rows]
    x_low, x_high, y_low, y_high = _limits(limits, xs, ys)
    head = line_height(KICKER_PT) if len(axes) > 1 else 0.0
    foot = line_height(KICKER_PT) if axes else 0.0
    plot = Box(box.x0, box.y0 + head, box.x1, box.y1 - foot)
    _room(plot, "a 2x2 matrix's plot", 1.6, 1.0)
    under = Box(plot.x0, box.y1 - foot, plot.x1, box.y1) if foot else None
    over = Box(box.x0, box.y0, box.x1, box.y0 + head) if head else None
    at_x = linear(x_low, x_high, plot.x0, plot.x1)
    at_y = linear(y_low, y_high, plot.y1, plot.y0)
    middle_x, middle_y = at_x((x_low + x_high) / 2), at_y((y_low + y_high) / 2)
    surface = theme.get("surface", theme.get("background", "#FFFFFF"))
    corners = (
        Box(plot.x0, plot.y0, middle_x, middle_y),
        Box(middle_x, plot.y0, plot.x1, middle_y),
        Box(plot.x0, middle_y, middle_x, plot.y1),
        Box(middle_x, middle_y, plot.x1, plot.y1),
    )
    # Two of the four, in the theme's surface. The other two are the page's own
    # ground, left unpainted: a white rectangle on a white page is a shape the
    # measurements have to look at and a reader never sees.
    for index in (0, 3):
        rect(slide, corners[index], surface)
    # The marks, before any of them is drawn, because the quadrant names go down
    # first and have to know where the points are. `_plot_points` draws every one
    # of them at DOT here -- a matrix takes no magnitudes, so there is one diameter
    # on the plot and no bubble arithmetic to agree with.
    marks = [Box(at_x(x) - DOT / 2, at_y(y) - DOT / 2, at_x(x) + DOT / 2, at_y(y) + DOT / 2) for x, y in zip(xs, ys)]
    named = []
    for index, name in enumerate(list(quadrants)[:4]):
        room = corners[index].inset(0.10, 0.08)
        align, anchor = _quadrant_corner(room, str(name), (middle_x, middle_y), marks, type_face(theme))
        write_label(slide, room, str(name), theme, size=KICKER_PT, colour=_muted(theme), align=align, anchor=anchor, bold=True)
        named.append(_ink(room, str(name), KICKER_PT, align, anchor, type_face(theme)))
    vline(slide, middle_x, plot.y0, plot.y1, _muted(theme))
    hline(slide, plot.x0, plot.x1, middle_y, _muted(theme))
    named.extend(_axis_names(slide, theme, axes, under, over))
    unnamed, _ = _plot_points(slide, plot, box, theme, labels, xs, ys, None, at_x, at_y, accent, named)
    # `where(x, y)` is the pair of inches a reading lands on, which is what putting
    # a target, a threshold or a second cohort on the same two dimensions needs.
    return _drawn(plot, lambda x, y: (at_x(x), at_y(y)), (KICKER_PT,), names=unnamed)


def scatter(slide, box, theme, data, *, axes=(), accent=None, limits=None):
    """Two numeric variables by position, and optionally a third by size.

    `data` is [(label, x, y)], or [(label, x, y, magnitude)] for bubbles -- the
    third number becomes the mark's area, which is the only honest way to draw it
    (radius would show a doubling as a quadrupling). `axes` names the two;
    `limits` is (x_low, x_high, y_low, y_high) when the ends are decided rather
    than measured.

    Every point is labelled where the label reaches, which for a dozen named
    items is all of them, and `drawn.names_not_written` is the ones it did not.
    `drawn.marks_not_to_scale` is the bubbles too small to carry an area, drawn
    open at the floor instead. A scatter of hundreds is not this -- it is a shape,
    and a shape wants a hexbin nobody here can draw.
    """
    rows, sizes = [], []
    for entry in list(data):
        parts = list(entry)
        if len(parts) == 3:
            label, x, y = parts
            magnitude = None
        elif len(parts) == 4:
            label, x, y, magnitude = parts
            magnitude = number(magnitude, label)
        else:
            raise ValueError(f"scatter takes (label, x, y) or (label, x, y, magnitude), not {entry!r}")
        rows.append((str(label), number(x, label), number(y, label)))
        sizes.append(magnitude)
    if not rows:
        raise ValueError("scatter needs at least one point")
    _room(box, "a scatter plot", 1.8, 1.2)
    labels = [label for label, _, _ in rows]
    xs = [value for _, value, _ in rows]
    ys = [value for _, _, value in rows]
    x_low, x_high, y_low, y_high = _limits(limits, xs, ys)
    names = list(axes)
    head = line_height(KICKER_PT) if len(names) > 1 else 0.0
    ticks = line_height(KICKER_PT)
    foot = ticks + (line_height(KICKER_PT) if names else 0.0)
    left = text_width([fmt(y_low), fmt(y_high)], KICKER_PT, type_face(theme)) + 0.06
    plot = Box(box.x0 + left, box.y0 + head, box.x1, box.y1 - foot)
    _room(plot, "a scatter plot's frame", 1.2, 0.8)
    at_x = linear(x_low, x_high, plot.x0, plot.x1)
    at_y = linear(y_low, y_high, plot.y1, plot.y0)
    vline(slide, plot.x0, plot.y0, plot.y1, _muted(theme))
    hline(slide, plot.x0, plot.x1, plot.y1, _muted(theme))
    # The four axis readings, kept as boxes: a point at an end of either scale can
    # otherwise put its name over the number that says where the end is.
    read = []
    for value in (y_low, y_high):
        y = at_y(value)
        room = Box(box.x0, y - line_height(KICKER_PT) / 2, plot.x0 - 0.06, y + line_height(KICKER_PT) / 2)
        write_label(slide, room, fmt(value), theme, size=KICKER_PT, colour=_muted(theme), align="right", anchor="middle")
        read.append(_ink(room, fmt(value), KICKER_PT, "right", "middle", type_face(theme)))
    for value in (x_low, x_high):
        x = at_x(value)
        room = Box(max(box.x0, x - 0.8), plot.y1 + 0.02, min(box.x1, x + 0.8), plot.y1 + 0.02 + ticks)
        write_label(slide, room, fmt(value), theme, size=KICKER_PT, colour=_muted(theme), align="center")
        read.append(_ink(room, fmt(value), KICKER_PT, "center", "top", type_face(theme)))
    under = Box(plot.x0, plot.y1 + ticks, plot.x1, box.y1) if names else None
    over = Box(box.x0, box.y0, box.x1, box.y0 + head) if head else None
    read.extend(_axis_names(slide, theme, names, under, over))
    unnamed, floored = _plot_points(slide, plot, box, theme, labels, xs, ys, sizes, at_x, at_y, accent, read)
    # `where(x, y)` is the pair of inches a reading lands on, so a trend line or a
    # target box goes on the plot at its own coordinates rather than by eye.
    return _drawn(plot, lambda x, y: (at_x(x), at_y(y)), (KICKER_PT,), names=unnamed, floored=floored)


def _limits(limits, xs, ys):
    """The four ends of a two-dimensional plot, given or measured and rounded off.

    Measured ends padded by a twentieth and then snapped out to a round step, so
    an axis reads "20 to 100" rather than "23.04 to 94.96". The step is the axis
    and not the data: no reading moves, and nothing is added between two of them.
    """
    if limits:
        values = [float(value) for value in limits]
        if len(values) != 4:
            raise ValueError("limits is (x_low, x_high, y_low, y_high)")
        return values[0], values[1], values[2], values[3]
    ends = []
    for readings in (xs, ys):
        low, high = min(readings), max(readings)
        ends.extend(snap(low, high))
    return ends[0], ends[1], ends[2], ends[3]


def snap(low, high):
    """(low, high) widened to the nearest round step, with a little air first."""
    span = (high - low) or (abs(high) or 1.0)
    pad = span * 0.05
    low, high = low - pad, high + pad
    power = 10.0 ** math.floor(math.log10(span / 4.0)) if span > 0 else 1.0
    step = power
    for factor in (1.0, 2.0, 2.5, 5.0, 10.0):
        step = factor * power
        if span / step <= 6:
            break
    return math.floor(low / step) * step, math.ceil(high / step) * step


# Where a point's name may sit, as (across, down, how far out). The first four are
# the sides, inward first so a label near an edge turns back into the plot rather
# than off it; then the four corners, which is what a row of points along an axis
# leaves free; then the sides again at arm's length, for the point in the middle of
# a cluster. Twelve tries and then the name is not written -- see `_plot_points`.
_LABEL_SIDES = (
    (1, 0, 1.0),
    (-1, 0, 1.0),
    (0, -1, 1.0),
    (0, 1, 1.0),
    (1, -1, 0.7),
    (1, 1, 0.7),
    (-1, -1, 0.7),
    (-1, 1, 0.7),
    (1, 0, 2.4),
    (-1, 0, 2.4),
    (0, -1, 2.4),
    (0, 1, 2.4),
)


def _plot_points(slide, plot, region, theme, labels, xs, ys, sizes, at_x, at_y, accent, furniture=()):
    """The marks, then their labels, each in the first place that is still free.

    Every mark is drawn before any label, because a label has to know where the
    other marks are: two points a tenth of an inch apart put "Online" through
    "Offline" on the first render of this, which is the failure a legend does not
    have and is the price of labelling directly. Twelve positions are tried per
    point -- the four sides, the four corners, then the sides at arm's length --
    against every mark, every label already placed, and the plot's own furniture
    (`furniture`: axis readings, axis names, quadrant names), and the first that
    clears all of them inside `region` is the one used. `region` is the box the
    author gave rather than the plot, because a point sitting on an axis has half
    its name in the margin outside the plot and nothing else wrong with it -- what
    the name must not land on is a reading, and every reading is in `furniture`.

    **A name that clears nothing is not written.** Four positions were tried here
    once and the last was taken whether it was free or not, which put "Online"
    through "Partner" the moment five points landed in one corner -- two names
    neither of which can be read, where an unlabelled mark is at least still a
    reading of x and y. The point that carries the accent is placed first so the
    one the page is about keeps its name; where several go unnamed, the plot has
    more points than a directly-labelled plot can hold and wants either `limits`
    that spread them or fewer of them.

    Which ones those were comes back, in the order they were given, along with the
    marks that had to be drawn at the bubble floor rather than to their own area.
    Both were decided here and reported nowhere: the only way to find out that
    three of eleven names had gone was to count them off a render, which is a round
    trip for something the arithmetic knew before the first shape was written.
    """
    chosen = _accented(labels, accent)
    magnitudes = [value for value in (sizes or []) if value is not None]
    largest = max(magnitudes) if magnitudes else 0.0
    marks, taken, floored, unnamed = [], list(furniture), [], []
    for index, label in enumerate(labels):
        x, y = at_x(xs[index]), at_y(ys[index])
        scaled = None
        if sizes and sizes[index] is not None and largest > 0:
            # Area, not radius: a mark twice as wide is four times the ink, and a
            # reader reads the ink.
            scaled = BUBBLE_SPAN * math.sqrt(max(0.0, sizes[index]) / largest)
        diameter = DOT if scaled is None else max(scaled, BUBBLE_FLOOR)
        paint = theme["accent"] if index in chosen else (_muted(theme) if chosen else _series(theme, 0))
        # Under the floor the diameter is the floor and not the magnitude, so the
        # mark stops claiming an area: every one of them is drawn open, at one
        # size, which reads as "below the scale" instead of as a disc four times
        # the ink of the value beside it. It used to be a filled disc at the same
        # floor, and 4, 0.5 and 0.05 came out as three identical circles.
        if scaled is not None and scaled < BUBBLE_FLOOR:
            floored.append(label)
            ring(slide, x, y, diameter, paint)
        else:
            disc(slide, x, y, diameter, paint)
        marks.append((x, y, diameter, paint))
        taken.append(Box(x - diameter / 2, y - diameter / 2, x + diameter / 2, y + diameter / 2))
    order = sorted(range(len(labels)), key=lambda index: index not in chosen)
    for index in order:
        x, y, diameter, paint = marks[index]
        width, height = text_width([labels[index]], KICKER_PT, type_face(theme)), line_height(KICKER_PT)
        beside = diameter / 2 + 0.04
        inward = -1 if x > (plot.x0 + plot.x1) / 2 else 1
        for across, down, reach in _LABEL_SIDES:
            across *= inward
            centre_x = x + across * (beside * reach + width / 2)
            centre_y = y + down * (beside * reach + height / 2)
            room = Box(centre_x - width / 2, centre_y - height / 2, centre_x + width / 2, centre_y + height / 2)
            # A name centred over its mark may slide along the edge of the region
            # to stay inside it, as long as it stays over the mark: a point in a
            # corner had its name refused for four thousandths of an inch, which
            # costs the reader the name and buys nothing. A name offset to one
            # side may not slide that way -- that would move it to the other side.
            if across == 0:
                shift = max(0.0, region.x0 - room.x0) - max(0.0, room.x1 - region.x1)
                room = Box(room.x0 + shift, room.y0, room.x1 + shift, room.y1)
                if not room.x0 - 1e-9 <= x <= room.x1 + 1e-9:
                    continue
            if down == 0:
                shift = max(0.0, region.y0 - room.y0) - max(0.0, room.y1 - region.y1)
                room = Box(room.x0, room.y0 + shift, room.x1, room.y1 + shift)
                if not room.y0 - 1e-9 <= y <= room.y1 + 1e-9:
                    continue
            if room.x0 < region.x0 - 1e-9 or room.x1 > region.x1 + 1e-9:
                continue
            if room.y0 < region.y0 - 1e-9 or room.y1 > region.y1 + 1e-9:
                continue
            if not _clear(room, taken):
                continue
            taken.append(room)
            write_label(
                slide,
                room,
                labels[index],
                theme,
                size=KICKER_PT,
                colour=paint if index in chosen else theme["foreground"],
                align="left" if across > 0 else ("right" if across < 0 else "center"),
                anchor="middle",
            )
            break
        else:
            unnamed.append(index)
    return tuple(labels[index] for index in sorted(unnamed)), tuple(floored)


def _air(one, other):
    """How far two boxes stand from each other, negative by how deep they overlap.

    The overlap is reported along the shallower axis, which is how far one of them
    would have to move to be off the other -- so ranking two placements by this
    ranks them by how much of the mark is under the type.
    """
    return max(max(other.x0 - one.x1, one.x0 - other.x1), max(other.y0 - one.y1, one.y0 - other.y1))


# What the outer corner costs, off a render of four named quadrants: the top-right name
# was written across the mark at (4.1, 4.6) with 0.073 x 0.150in of the disc under the
# type, the disc being 0.150in across -- half the point covered.
def _quadrant_corner(room, name, crossing, marks, face):
    """Which corner of `room` a quadrant's name goes in: the emptiest it can have.

    The corner furthest from the crossing is the one no reader can mistake for the
    next quadrant's -- and is also exactly where the extreme point of that quadrant
    is, so a name nailed there is written across the mark whose quadrant it names.

    Between the two, the furniture is what moves: a quadrant name names a region,
    so it reads from any corner of that region, while a point is at its reading and
    nowhere else. Each of the four corners is measured against every mark and the
    ones that clear are preferred, furthest from the crossing first -- so a
    quadrant with nothing in its outer corner is still written in that corner, and
    one with a point there steps along an edge instead of onto it.
    Where all four are occupied the corner with the most air is used: there is no
    corner of a quadrant that is not inside it, so the name never wanders out of
    the region it belongs to, whatever the data does.
    """
    corners = []
    for align in ("left", "right"):
        for anchor in ("top", "bottom"):
            ink = _ink(room, name, KICKER_PT, align, anchor, face)
            air = min([_air(ink, mark) for mark in marks], default=room.w)
            away = math.hypot((ink.x0 + ink.x1) / 2 - crossing[0], (ink.y0 + ink.y1) / 2 - crossing[1])
            corners.append((air >= 0.0, away if air >= 0.0 else air, align, anchor))
    _, _, align, anchor = max(corners)
    return align, anchor


def _clear(room, taken, tolerance=0.015):
    """Whether `room` misses everything already placed."""
    return not any(
        room.x0 < other.x1 - tolerance
        and other.x0 < room.x1 - tolerance
        and room.y0 < other.y1 - tolerance
        and other.y0 < room.y1 - tolerance
        for other in taken
    )


def _axis_names(slide, theme, axes, under, over):
    """The two dimensions named, in the bands the plot left for them; -> their ink.

    The boxes come back so the point labels can be kept off them. They are drawn
    before the points for that reason -- neither can move, so whichever is placed
    second is the one that has to give way, and a point's name is the one of the
    two that has somewhere else to go.

    Both names run along the page rather than up it: python-pptx can rotate a
    text frame, and a rotated frame is measured on its unrotated box, so a
    vertical axis title is the one label a page can put off the edge without any
    measurement noticing. The y name goes over the plot instead.
    """
    names = list(axes)
    covered = []
    if names and under is not None:
        write_label(
            slide,
            under,
            str(names[0]),
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            align="center",
            anchor="bottom",
        )
        covered.append(_ink(under, str(names[0]), KICKER_PT, "center", "bottom", type_face(theme)))
    if len(names) > 1 and over is not None:
        write_label(
            slide,
            over,
            str(names[1]),
            theme,
            size=KICKER_PT,
            colour=_muted(theme),
            anchor="middle",
        )
        covered.append(_ink(over, str(names[1]), KICKER_PT, "left", "middle", type_face(theme)))
    return covered


def _share(value, label=None, what="a progress track"):
    """"76%" -> 0.76, "0.76" -> 0.76, 76 -> 0.76, and anything over the whole is 1.

    `what` names the thing being read, because the same three spellings are read for
    a progress track and for a paint's opacity and a refusal that says the wrong one
    sends the author looking at the wrong argument.
    """
    where = f" for {label!r}" if label is not None else ""
    text = str("" if value is None else value).strip().replace(",", "")
    try:
        reading = float(text.rstrip("%").strip())
    except ValueError:
        raise ValueError(f"{what}{where} needs a share, not {value!r}") from None
    # The same refusal `number` makes, repeated because a progress track does not
    # go through it: `float("nan")` parses, and the clamp below then turned a NaN
    # into an empty track and an infinity into a full one, silently. A wrong bar
    # nobody is told about is the one outcome worse than an error.
    if not math.isfinite(reading):
        raise ValueError(f"{what}{where} needs a finite share, not {value!r}")
    if text.endswith("%") or reading > 1:
        reading /= 100.0
    return min(1.0, max(0.0, reading))


# ---------------------------------------------------------- asking beforehand


def _chart_of(chart):
    """`chart` as the function, whether it arrived as the function or as its name.

    Unaccepted, `the_smallest_box_a_chart_needs("horizontal_bar", T, data)` raises
    `TypeError: 'str' object is not callable` out of the rehearsal, three frames down,
    with neither the name passed nor the word "name" in the message. The
    name is the natural thing to reach for -- it is what the chart is called in the
    catalogue and in the page's own plan -- so it is accepted, and a name that is not
    one is refused here saying so and listing the near misses.
    """
    if callable(chart):
        return chart
    known = _drawable_names()
    name = str(chart)
    if name in known:
        return globals()[name]
    import difflib

    close = difflib.get_close_matches(name, known, n=3, cutoff=0.5)
    hint = f"; closest: {', '.join(close)}" if close else f"; the charts are {', '.join(known)}"
    raise ValueError(f"no chart called {name!r}{hint}")


def _drawable_names():
    """The charts, read off the module rather than listed: first parameter is `slide`.

    A list is a thing somebody has to remember to add to, and a chart added without
    it would answer "no chart called that" about itself.
    """
    import inspect

    found = []
    for name, value in globals().items():
        if name.startswith("_") or not inspect.isfunction(value) or value.__module__ != __name__:
            continue  # `write` is imported from ppt_layout and also takes a slide first
        if name in _NOT_A_FORM:
            continue
        parameters = list(inspect.signature(value).parameters)
        if parameters and parameters[0] == "slide":
            found.append(name)
    return tuple(sorted(found))


def what_a_chart_will_do(chart, box, theme, *data, **knobs):
    """Everything `chart` would decide in `box`, decided without drawing any of it.

    `chart` is the chart function, or its name -- both work.

    Same arguments as the chart itself, minus the slide, and the same `Drawn` back:
    which readings it would have to drop, which point names would not clear, which
    bubbles fall to the floor, what step of the ramp its labels land on, and where
    its scale puts a value. It raises the same `TooSmall` a draw would, for the
    same box.

        will = what_a_chart_will_do(scatter, cell, T, points)
        if will.names_not_written:
            scatter(slide, cell, T, points, limits=(0, 100, 0, 100))
        else:
            scatter(slide, cell, T, points)

    Nothing is written to the slide, so this costs a page nothing to ask, and the
    answer cannot be wrong about the chart because it *is* the chart -- see
    `_Rehearsal`.
    """
    return _chart_of(chart)(_Rehearsal(), box, theme, *data, **knobs)


def whether_a_chart_fits(chart, box, theme, *data, **knobs):
    """Whether `chart` would take `box` at all: the refusal, asked as a question.

    A build script choosing between two forms wants a boolean and not a traceback,
    and the two-line try/except that turns one into the other is two lines every
    page that picks a chart would carry.
    """
    try:
        what_a_chart_will_do(chart, box, theme, *data, **knobs)
    except TooSmall:
        return False
    return True


def the_smallest_box_a_chart_needs(chart, theme, *data, **knobs):
    """The smallest box this chart takes *this* data in, as a `Box` at the origin.

    The question a page asks before it divides itself up: a chart's floor is not the
    pair of constants in its
    `_room` call -- it is those, and the width its category names want, and the
    height one line of type per row wants, and what a wrapped key or a column of
    axis readings takes out of both. All of that is arithmetic the chart does, so
    this asks the chart rather than restating it: a box that is refused comes back
    as a `TooSmall` carrying the shortfall, which is added and asked again, and
    then each side is tightened back down while the chart still takes it.

        room = the_smallest_box_a_chart_needs(column, T, rows)
        chart = column if room.w <= cell.w and room.h <= cell.h else horizontal_bar

    Twelve categories with names like "Manufacturing": `column` wants 8.05x0.90in
    and a quadrant of a 13.3in page has 5.90x2.50in, so the page that
    would have drawn twelve columns 0.49in apart draws a ranking instead -- decided
    before a shape is written, from the data alone.

    Both sides are minimal, one at a time: shrinking the width can raise the height
    a chart needs, so this is *a* smallest box and not the only one. Height is
    tightened first because that is the side a page has least of.
    """
    across, down, refusal = 0.0, 0.0, None
    for _ in range(80):
        try:
            what_a_chart_will_do(chart, Box(0.0, 0.0, across, down), theme, *data, **knobs)
        except TooSmall as refused:
            refusal = refused
            if refused.short == (0.0, 0.0):
                raise
            # A hair over the shortfall, because a floor met exactly to the last
            # bit of a float is a floor that fails again on the next comparison.
            across += refused.short[0] + 1e-4
            down += refused.short[1] + 1e-4
            continue
        break
    else:
        raise refusal
    # Growing overshoots wherever a floor eased on the way up -- a four-name key
    # wraps onto four lines in a box an inch wide and onto one in a box six inches
    # wide, and the height asked for at the first was carried all the way. So each
    # side comes back down while the chart still takes it.
    def takes(width, height):
        return whether_a_chart_fits(chart, Box(0.0, 0.0, width, height), theme, *data, **knobs)

    down = _tightest(lambda value: takes(across, value), down)
    across = _tightest(lambda value: takes(value, down), across)
    down = _tightest(lambda value: takes(across, value), down)
    return Box(0.0, 0.0, across, down)


def _tightest(takes, most):
    """The least of `most` that still `takes`, to a thousandth of an inch."""
    least = 0.0
    while most - least > 0.001:
        middle = (least + most) / 2
        if takes(middle):
            most = middle
        else:
            least = middle
    return most
'''


def chart_module_source() -> str:
    """Source of the `ppt_charts` module a build script imports."""
    return _CHART_MODULE


@lru_cache(maxsize=1)
def _catalogue() -> tuple[tuple[str, str], ...]:
    """Every chart above, with the line its own docstring opens on.

    Read off the source rather than typed out beside it. A hand-kept list of
    twenty-three names and twenty-three purposes is one more place to update when
    a chart is added, and the one that silently goes stale is the one nobody reads
    -- which is the whole failure this catalogue exists to answer.

    A chart is a public function that takes the slide first, minus the drawing base
    the module names in `_NOT_A_FORM`. The module also hands
    an author `what_a_chart_will_do` and `the_smallest_box_a_chart_needs`, which
    take a chart rather than a slide and are not charts: listing them here would
    put them in the brief as two more forms to draw, which is the opposite of what
    they are for. `rect` and the eight beside it are the same mistake from the other
    end -- a rectangle is what a form is made of, not a form to pick -- and the
    module's own tuple is read here rather than restated, so the two answers to
    "is this a chart" cannot come apart.
    """
    tree = ast.parse(_CHART_MODULE)
    base: tuple[str, ...] = ()
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(
            isinstance(target, ast.Name) and target.id == "_NOT_A_FORM" for target in node.targets
        ):
            base = tuple(ast.literal_eval(node.value))
    charts = []
    for node in tree.body:
        if not isinstance(node, ast.FunctionDef) or node.name.startswith("_"):
            continue
        if not node.args.args or node.args.args[0].arg != "slide" or node.name in base:
            continue
        purpose = (ast.get_docstring(node) or "").strip().split("\n", 1)[0].strip()
        charts.append((node.name, purpose))
    return tuple(charts)


def chart_names() -> tuple[str, ...]:
    """The name of every chart the module draws, in the order it defines them."""
    return tuple(name for name, _ in _catalogue())
