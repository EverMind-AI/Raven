"""Rewriting one page without breaking the other seventeen.

Every case here is a round that was lost. The prelude checks come from an edit
that took the deck from finished to unbuildable; the whole-write verification
comes from a round that built eighteen slides from seventeen mapped blocks.
"""

from __future__ import annotations

import textwrap

from raven.ppt.backends.script import (
    applied_lines,
    apply_verified,
    block_rejection,
    blocks_rejection,
    page_blocks,
    prelude_rejection,
    slide_creators,
    with_banner,
)

PRELUDE = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches
    from ppt_theme import THEMES, rgb

    TH = THEMES["ink-graphite"]
    ACC = rgb(TH["accent"])
    MUT = rgb(TH["muted"])
    FNT = "Source Serif 4"

    prs = Presentation()


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    def title(slide, text):
        slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1)).text_frame.text = text


    def footnote(slide, text):
        slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11), Inches(0.4)).text_frame.text = text
    """
).lstrip()

SCRIPT = PRELUDE + textwrap.dedent(
    """

    # SLIDE 1
    one = new_slide()
    title(one, "Unified video segmentation")

    # SLIDE 2
    two = new_slide()
    title(two, "Target queries")

    prs.save(os.environ["PPT_OUTPUT"])
    """
)

BLOCKS_TEXT = "one = new_slide()\ntitle(one, 'x')\ntwo = new_slide()\ntitle(two, 'y')\n"


def test_a_prelude_that_drops_a_helper_the_pages_call_is_refused() -> None:
    """One lost helper fails every page at once, at import time."""
    stripped = PRELUDE.replace("def title(slide, text):", "def caption(slide, text):")
    reason = prelude_rejection(stripped, PRELUDE, BLOCKS_TEXT)
    assert reason is not None and "title" in reason


def test_a_prelude_that_drops_a_helper_no_page_calls_is_allowed() -> None:
    """The check is what the pages call, not what the prelude happens to hold."""
    stripped = PRELUDE[: PRELUDE.index("def footnote")].rstrip() + "\n"
    assert "footnote" not in stripped
    assert prelude_rejection(stripped, PRELUDE, BLOCKS_TEXT) is None


def test_a_prelude_that_recolours_the_deck_is_refused_and_names_the_line() -> None:
    """The palette may be the user's choice rather than the deck's to revisit."""
    recoloured = PRELUDE.replace('ACC = rgb(TH["accent"])', 'ACC = rgb("FF0000")')
    reason = prelude_rejection(recoloured, PRELUDE, BLOCKS_TEXT)
    assert reason is not None
    assert "palette or font family" in reason
    assert "FF0000" in reason


def test_realigning_the_prelude_does_not_read_as_a_recolour() -> None:
    realigned = PRELUDE.replace('ACC = rgb(TH["accent"])', 'ACC  =  rgb(TH["accent"])')
    assert prelude_rejection(realigned, PRELUDE, BLOCKS_TEXT) is None


def test_reordering_two_different_names_is_reorganisation_not_a_recolour() -> None:
    reordered = PRELUDE.replace(
        'ACC = rgb(TH["accent"])\nMUT = rgb(TH["muted"])',
        'MUT = rgb(TH["muted"])\nACC = rgb(TH["accent"])',
    )
    assert prelude_rejection(reordered, PRELUDE, BLOCKS_TEXT) is None


def test_swapping_two_assignments_to_the_same_name_is_a_new_palette() -> None:
    """Which one wins changed, so the lines are the old ones but the deck is not."""
    doubled = PRELUDE.replace(
        'ACC = rgb(TH["accent"])',
        'ACC = rgb(TH["accent"])\nACC = rgb("00FF00")',
    )
    swapped = PRELUDE.replace(
        'ACC = rgb(TH["accent"])',
        'ACC = rgb("00FF00")\nACC = rgb(TH["accent"])',
    )
    assert prelude_rejection(swapped, doubled, BLOCKS_TEXT) is not None


def test_an_empty_prelude_is_refused() -> None:
    assert prelude_rejection("   \n", PRELUDE, BLOCKS_TEXT) == "the replacement prelude was empty"


def test_prelude_helpers_that_wrap_add_slide_count_as_slide_creators() -> None:
    """`title` and `footnote` do not create slides, so they must not be counted."""
    assert slide_creators(PRELUDE) == ["add_slide", "clone_page", "adapt", "new_slide"]


def test_a_page_cloned_from_the_template_counts_as_a_page() -> None:
    """`adapt` and `clone_page` create a slide -- through `add_slide`, inside -- and the
    text scan could not see that. With the structural pages cloned, the block that draws
    slide 1 holds an `adapt` call and no `add_slide`, so the whole design pass refused:
    measured across two live runs, every finished build, which is the pass the deck was
    counting on for its layout."""
    lines = (
        PRELUDE
        + "\n# SLIDE 1\nslide = adapt(prs, prototype(tpl, 1), texts={1: 'title'})\n"
        + "\n# SLIDE 2\nslide = new_slide()\n"
    ).splitlines(keepends=True)
    blocks = page_blocks(lines)

    assert blocks_rejection(lines, blocks, slide_count=2) is None


def test_the_ordinary_way_to_make_a_page_is_a_method_call() -> None:
    """`prs.slides.add_slide(LAY)` is how every drawn page starts, dots and all.

    The first fix for the cloned-page case excluded a leading dot -- to keep
    `ppt_template.prototype(tpl, 1)` from counting as a creator -- and that turned every
    drawn page into a block that draws nothing. Both real scripts measured refused on
    their first drawn page.
    """
    lines = (
        PRELUDE
        + "\n# SLIDE 1\nslide = prs.slides.add_slide(LAY)\ntitle(slide, 'drawn')\n"
        + "\n# SLIDE 2\nslide = adapt(prs, prototype(tpl, 2), items=[['01', 'cloned']])\n"
    ).splitlines(keepends=True)
    blocks = page_blocks(lines)

    assert blocks_rejection(lines, blocks, slide_count=2) is None


def test_a_block_count_that_disagrees_with_the_deck_is_refused() -> None:
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    reason = blocks_rejection(lines, blocks, slide_count=3)
    assert reason is not None and "2 page block(s)" in reason and "3 slide(s)" in reason


def test_the_advice_differs_by_where_the_mapping_came_from() -> None:
    """Execution is already exact, so a mismatch means the script has no per-page block."""
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    from_comments = blocks_rejection(lines, blocks, 3, derived="comments")
    from_execution = blocks_rejection(lines, blocks, 3, derived="execution")
    assert "# SLIDE <n>" in (from_comments or "")
    assert "shared loop" in (from_execution or "")


def test_pages_numbered_out_of_range_are_refused() -> None:
    script = SCRIPT.replace("# SLIDE 2", "# SLIDE 5")
    lines = script.splitlines(keepends=True)
    reason = blocks_rejection(lines, page_blocks(lines), slide_count=2)
    assert reason is not None and "not 1..2" in reason


def test_a_block_that_draws_two_slides_is_refused() -> None:
    script = SCRIPT.replace('title(two, "Target queries")', 'three = new_slide()\ntitle(three, "x")')
    lines = script.splitlines(keepends=True)
    reason = blocks_rejection(lines, page_blocks(lines), slide_count=2)
    assert reason is not None and "creates 2 slides" in reason


def test_a_returned_block_that_saves_the_deck_is_refused() -> None:
    assert "saving belongs at the end" in (block_rejection("prs.save('x')\n", "one = new_slide()\n") or "")


def test_an_empty_block_is_refused() -> None:
    assert block_rejection("# nothing\n", "one = new_slide()\n") == "the block came back with no code in it"


def test_a_block_that_lost_most_of_its_content_is_refused() -> None:
    original = "\n".join(f"line{i} = {i}" for i in range(20))
    assert "dropped content" in (block_rejection("line0 = 0\n", original) or "")


def test_a_missing_banner_is_restored_rather_than_bounced() -> None:
    """Mechanical discipline is the harness's job: the caller knows the page."""
    assert with_banner(3, "one = new_slide()\n").startswith("# SLIDE 3\n")
    assert with_banner(3, "# SLIDE 3\none = new_slide()\n") == "# SLIDE 3\none = new_slide()\n"


def test_a_banner_naming_a_different_page_is_refused_because_it_is_ambiguous() -> None:
    """It may be a typo, or another page's code in the wrong envelope."""
    reason = block_rejection("# SLIDE 7\nx = new_slide()\ny = 1\n", "# SLIDE 2\nx = new_slide()\ny = 1\n")
    assert reason is not None and "slide 2's block" in reason


def test_replacements_are_re_split_so_downstream_readers_see_every_line() -> None:
    """Returned as one multi-line string, a block hid all but its first line --
    and the banner count then read 0 where 18 stood, refusing every replacement."""
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    applied = applied_lines(lines, blocks, {1: "# SLIDE 1\nx = new_slide()\ntitle(x, 'y')\n"})
    assert all(line.count("\n") <= 1 for line in applied)
    assert len(page_blocks(applied)) == 2


def test_a_replacement_that_would_cost_the_script_a_page_is_dropped_alone() -> None:
    """One bad block costs itself, not the round."""
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    kept, dropped = apply_verified(
        lines,
        blocks,
        {1: "x = new_slide()\ntitle(x, 'no banner and no marker')\n", 2: "# SLIDE 2\ny = new_slide()\n"},
    )
    assert 1 in dropped and "fewer `# SLIDE` banners" in dropped[1]
    assert kept == {2: "# SLIDE 2\ny = new_slide()\n"}


def test_good_replacements_all_survive() -> None:
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    good = {1: "# SLIDE 1\na = new_slide()\ntitle(a, 'x')\n", 2: "# SLIDE 2\nb = new_slide()\ntitle(b, 'y')\n"}
    kept, dropped = apply_verified(lines, blocks, good)
    assert kept == good and dropped == {}


def test_replacing_the_highest_page_first_keeps_the_other_spans_valid() -> None:
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    applied = applied_lines(
        lines,
        blocks,
        {1: "# SLIDE 1\na = new_slide()\n" + "# filler\n" * 30, 2: "# SLIDE 2\nb = new_slide()\n"},
    )
    text = "".join(applied)
    assert text.index("# SLIDE 1") < text.index("# SLIDE 2")
    assert "prs.save(" in text


def test_a_block_that_rewrites_the_copy_is_refused() -> None:
    """The design pass may rearrange a page and may not rewrite it -- its own brief says
    so in the first paragraph -- and one round did both anyway. Measured page by page on
    a polished deck: 224 of 3420 characters gone across four pages, a whole bullet
    dropped, a chart's caption and a page's conclusion removed, and three surviving lines
    quietly reworded ("训练时用匈牙利匹配指派" came back as "训练时匈牙利匹配指派").

    Nothing measured it, because the words are in the program rather than in a place any
    check of the built file could compare against an earlier version.
    """
    original = (
        "# SLIDE 4\n"
        "write(slide, B(0.7, 1.0, 5, 0.4), '共同能力被抽象为一句话')\n"
        "write(slide, B(0.7, 2.0, 5, 1.2), ['· Qinst：实例数上界 I，训练时用匈牙利匹配指派', '· Qbg 由参考帧划成 4×4 网格'])\n"
    )

    # Merging two literals into one, reflowing, and moving a box are all rearrangement.
    merged = (
        "# SLIDE 4\n"
        "write(slide, B(6.9, 1.0, 5, 0.4), '共同能力被抽象为一句话')\n"
        "write(slide, B(6.9, 2.0, 5, 1.2), '· Qinst：实例数上界 I，训练时用匈牙利匹配指派 · Qbg 由参考帧划成 4×4 网格')\n"
    )
    assert block_rejection(merged, original) is None

    # Shaving one character out of a sentence is not.
    reworded = merged.replace("训练时用匈牙利", "训练时匈牙利")
    refusal = block_rejection(reworded, original)
    assert refusal is not None and "训练时用匈牙利匹配指派" in refusal

    # Nor is dropping a bullet.
    dropped = original.replace("'· Qbg 由参考帧划成 4×4 网格'", "''")
    refusal = block_rejection(dropped, original)
    assert refusal is not None and "Qbg" in refusal


def test_a_short_label_may_still_be_re_lettered() -> None:
    """ "01" to "1" is the pass's to decide; a sentence is not."""
    original = "# SLIDE 2\nchip(slide, '01')\nwrite(slide, box, '碎片化的代价')\n"
    relettered = "# SLIDE 2\nchip(slide, '1')\nwrite(slide, box, '碎片化的代价')\n"

    assert block_rejection(relettered, original) is None
