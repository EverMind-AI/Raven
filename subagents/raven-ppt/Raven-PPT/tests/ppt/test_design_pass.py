"""The design loop, driven with fakes so every branch is reachable.

The loop exists because a single pass can only *claim* to have improved a page.
These tests are mostly about what it refuses: a round that would break the
program, a prelude that would take every page down with it, and a content problem
it must not be handed because it is forbidden to solve it.
"""

from __future__ import annotations

import json
import textwrap
from pathlib import Path
from typing import Any

import pytest

from raven.ppt.backends.script import script_path
from raven.ppt.contracts import Audience, BuildOutcome, Finding, PageSource, Project, Severity
from raven.ppt.stages.design_pass import DesignPass, TypeFloors

PRELUDE = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from ppt_theme import THEMES, rgb

    TH = THEMES["ink-graphite"]
    ACC = rgb(TH["accent"])

    prs = Presentation()


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    def title(slide, text):
        slide.shapes.add_textbox(0, 0, 100, 50).text_frame.text = text
    """
).lstrip()

SCRIPT = PRELUDE + textwrap.dedent(
    """

    # SLIDE 1
    one = new_slide()
    title(one, "Unified video segmentation")

    # SLIDE 2
    two = new_slide()
    title(two, "Target queries")

    prs.save(os.environ["PPT_OUTPUT"])
    """
)


class FakeRenderer:
    def __init__(self, root: Path) -> None:
        self.root = root

    async def pages(self, pptx: Path, out_dir: Path, numbers) -> dict[int, Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        made = {}
        for number in numbers:
            path = out_dir / f"page-{number:02d}.png"
            path.write_bytes(b"\x89PNG fake")
            made[number] = path
        return made

    def contact_sheet(self, pngs, out: Path, columns: int = 3) -> Path:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(b"\x89PNG sheet")
        return out

    def data_uri(self, png: Path) -> str:
        return "data:image/png;base64,AAAA"


class FakeComposer:
    """Answers scripted per call, and records the briefs and parts it saw."""

    def __init__(self, replies: list[str]) -> None:
        self.replies = list(replies)
        self.seen: list[tuple[str, list[dict[str, Any]]]] = []

    async def ask(self, system: str, parts: list[dict[str, Any]], *, max_tokens: int) -> str:
        self.seen.append((system, parts))
        return self.replies.pop(0) if self.replies else json.dumps({"verdict": "ok"})

    def texts(self) -> str:
        return "\n".join(p.get("text", "") for _, parts in self.seen for p in parts)


@pytest.fixture()
def project(tmp_path: Path) -> Project:
    p = Project(workspace=tmp_path, slug="tarvis")
    p.build_dir.mkdir(parents=True)
    script_path(p).write_text(SCRIPT, encoding="utf-8")
    (p.build_dir / "deck.pptx").write_bytes(b"PK deck")
    return p


def _outcome(project: Project, pages: int = 2) -> BuildOutcome:
    lines = SCRIPT.splitlines(keepends=True)
    first = next(i for i, line in enumerate(lines) if "one = new_slide()" in line) - 1
    second = next(i for i, line in enumerate(lines) if "two = new_slide()" in line) - 1
    return BuildOutcome(
        ok=True,
        pptx_path=project.build_dir / "deck.pptx",
        pages=pages,
        sources=(
            PageSource(page=1, first_line=first, last_line=second),
            PageSource(page=2, first_line=second, last_line=len(lines) - 2),
        ),
        source_digest="x",
    )


def _measure(findings):
    async def measure(_project: Project, _pptx: Path, _outcome=None):
        return list(findings)

    return measure


def _pass(project: Project, replies: list[str], *, findings=(), builds=None, rounds: int = 1) -> DesignPass:
    outcomes = list(builds or [_outcome(project)])

    async def build(_project: Project) -> BuildOutcome:
        return outcomes.pop(0) if outcomes else _outcome(_project)

    return DesignPass(
        renderer=FakeRenderer(project.workspace),
        composer=FakeComposer(replies),
        build=build,
        measure=_measure(findings),
        floors=TypeFloors(),
        rounds=rounds,
        concurrency=2,
    )


@pytest.mark.asyncio
async def test_a_page_edit_is_written_and_the_deck_rebuilt(project: Project) -> None:
    block = "# SLIDE 1\none = new_slide()\ntitle(one, 'Unified video segmentation')\nACC\n"
    replies = [
        json.dumps({"verdict": "ok", "notes": ["setup is fine"], "vocabulary": ["keep the rail"]}),
        json.dumps({"slide": 1, "verdict": "edited", "notes": ["tightened the head"], "block": block}),
        json.dumps({"slide": 2, "verdict": "ok", "notes": []}),
    ]
    result = await _pass(project, replies).run(project, _outcome(project))

    assert result.ok, result.note
    assert "ACC" in script_path(project).read_text(encoding="utf-8")
    assert result.data["vocabulary"] == ["keep the rail"]
    assert result.data["changed"] is True


@pytest.mark.asyncio
async def test_the_design_pass_saves_before_and_after_rendered_artifacts(project: Project) -> None:
    result = await _pass(project, []).run(project, _outcome(project))

    assert result.ok
    for label in ("before", "after"):
        root = project.review_dir / "design_pass" / label
        assert (root / "deck.pptx").is_file()
        assert (root / "page-01.png").is_file()
        assert (root / "page-02.png").is_file()
        assert (root / "contact.png").is_file()
        assert (root / "manifest.json").is_file()
    assert result.data["artifacts"] == {
        "before": str(project.review_dir / "design_pass" / "before"),
        "after": str(project.review_dir / "design_pass" / "after"),
        "report": str(project.review_dir / "design_pass" / "report.json"),
    }
    report = json.loads((project.review_dir / "design_pass" / "report.json").read_text(encoding="utf-8"))
    assert report["status"] == "completed"
    assert report["changed"] is False
    assert report["rounds"][0]["pages"]["1"]["verdict"] == "ok"


@pytest.mark.asyncio
async def test_the_measured_floors_reach_both_briefs_from_the_measurement(project: Project) -> None:
    """Six copies of a number is five chances to be wrong about it."""
    design = _pass(project, [])
    design.floors = TypeFloors(body_pt=16.0, min_pt=12.0)
    await design.run(project, _outcome(project))
    briefs = "\n".join(system for system, _ in design.composer.seen)  # type: ignore[attr-defined]
    assert "12pt" in briefs and "16pt" in briefs
    assert "10.8pt" not in briefs and "14pt" not in briefs


@pytest.mark.asyncio
async def test_a_design_finding_reaches_the_page_it_is_about(project: Project) -> None:
    finding = Finding(
        kind="card_overflow",
        severity=Severity.WARNING,
        message="'VIPSeg' runs 0.27in past the edge of its card",
        page=2,
        audience=Audience.DESIGNER,
    )
    design = _pass(project, [], findings=[finding])
    await design.run(project, _outcome(project))
    assert "runs 0.27in past" in design.composer.texts()  # type: ignore[attr-defined]


@pytest.mark.asyncio
async def test_a_content_finding_is_withheld_from_the_design_pass(project: Project) -> None:
    """It may rearrange a page but never change what it says.

    Handed this, it returns the same page compartmented, however often it is
    redesigned -- which is what happened, measurably, before the split.
    """
    finding = Finding(
        kind="density",
        severity=Severity.WARNING,
        message="the page carries 295 words",
        page=2,
        audience=Audience.AUTHOR,
    )
    design = _pass(project, [], findings=[finding])
    await design.run(project, _outcome(project))
    assert "295 words" not in design.composer.texts()  # type: ignore[attr-defined]


@pytest.mark.asyncio
async def test_a_prelude_that_would_take_every_page_down_is_refused(project: Project) -> None:
    stripped = PRELUDE.replace("def title(slide, text):", "def caption(slide, text):")
    replies = [json.dumps({"verdict": "edited", "notes": ["renamed"], "prelude": stripped})]
    result = await _pass(project, replies).run(project, _outcome(project))

    assert result.ok
    assert "title" in (result.data["rounds"][0]["deck"].get("refused") or "")
    assert "def title" in script_path(project).read_text(encoding="utf-8")


def _died_in_page_one() -> BuildOutcome:
    """A build that failed on a line inside the block that draws page 1."""
    lines = SCRIPT.splitlines(keepends=True)
    marker = next(i for i, line in enumerate(lines) if "# SLIDE 1" in line)
    return BuildOutcome(
        ok=False,
        stderr=f'  File "/w/build/build.py", line {marker + 2}, in <module>\nNameError: nope\n',
    )


# The block keeps the page's copy -- a round that dropped it would be refused
# before it ever built (see `copy_rejection`), and these are about the build.
BREAKS_PAGE_ONE = json.dumps(
    {
        "slide": 1,
        "verdict": "edited",
        "notes": ["x"],
        "block": '# SLIDE 1\none = new_slide()\ntitle(one, "Unified video segmentation")\nnope\n',
    }
)


@pytest.mark.asyncio
async def test_the_page_that_will_not_run_costs_itself_and_not_the_round(project: Project) -> None:
    """One page's bad edit used to discard every other page's good one."""
    replies = [
        json.dumps({"verdict": "ok"}),
        BREAKS_PAGE_ONE,
        json.dumps(
            {
                "slide": 2,
                "verdict": "edited",
                "notes": ["y"],
                "block": '# SLIDE 2\ntwo = new_slide()\ntitle(two, "Target queries")\nACC  # page two stayed\n',
            }
        ),
    ]
    result = await _pass(project, replies, builds=[_died_in_page_one(), _outcome(project)]).run(
        project, _outcome(project)
    )

    written = script_path(project).read_text(encoding="utf-8")
    assert result.ok, result.note
    assert "nope" not in written, "page 1's edit had to go"
    assert "page two stayed" in written, "page 2's edit had nothing to do with it"
    round_one = result.data["rounds"][0]
    assert round_one["rewrote"] == [2]
    assert "NameError: nope" in round_one["would_not_run"][1]


@pytest.mark.asyncio
async def test_the_next_round_is_told_what_would_not_run(project: Project) -> None:
    """Otherwise round two makes the same call to the same helper again."""
    replies = [
        json.dumps({"verdict": "ok"}),
        BREAKS_PAGE_ONE,
        json.dumps({"slide": 2, "verdict": "ok"}),
        json.dumps({"slide": 1, "verdict": "ok"}),
        json.dumps({"slide": 2, "verdict": "ok"}),
    ]
    passes = _pass(project, replies, builds=[_died_in_page_one(), _outcome(project)], rounds=2)
    result = await passes.run(project, _outcome(project))

    assert result.ok, result.note
    assert "would not run and was dropped" in passes.composer.texts()
    assert "NameError: nope" in passes.composer.texts()


@pytest.mark.asyncio
async def test_a_failure_that_names_no_page_is_reverted_and_attributed(project: Project) -> None:
    """The next actor needs a deck to work from more than it needs this round."""
    lines = SCRIPT.splitlines(keepends=True)
    in_prelude = next(i for i, line in enumerate(lines) if "def new_slide" in line)
    broken_build = BuildOutcome(
        ok=False,
        stderr=f'  File "/w/build/build.py", line {in_prelude + 1}, in new_slide\nNameError: nope\n',
    )
    replies = [json.dumps({"verdict": "ok"}), BREAKS_PAGE_ONE, json.dumps({"slide": 2, "verdict": "ok"})]
    result = await _pass(project, replies, builds=[broken_build, _outcome(project)]).run(project, _outcome(project))

    assert not result.ok
    assert script_path(project).read_text(encoding="utf-8") == SCRIPT
    assert result.findings and result.findings[0].kind == "design_pass_broke_the_build"
    assert result.findings[0].audience is Audience.AUTHOR


@pytest.mark.asyncio
async def test_an_unparsable_reply_costs_that_page_and_not_the_round(project: Project) -> None:
    replies = [
        json.dumps({"verdict": "ok"}),
        '{"slide": 1, "verdict": "edited", "block": }',
        json.dumps({"slide": 2, "verdict": "edited", "notes": ["ok"], "block": "# SLIDE 2\ntwo = new_slide()\nACC\n"}),
    ]
    result = await _pass(project, replies).run(project, _outcome(project))

    assert result.ok
    assert "at character" in result.data["rounds"][0]["pages"][1]["error"]
    assert "ACC" in script_path(project).read_text(encoding="utf-8")


@pytest.mark.asyncio
async def test_the_loop_stops_early_when_a_round_changes_nothing(project: Project) -> None:
    design = _pass(project, [], rounds=3)
    result = await design.run(project, _outcome(project))
    assert result.ok
    assert len(result.data["rounds"]) == 1
    assert result.data["rounds"][0]["stopped"] == "nothing left to change"


@pytest.mark.asyncio
async def test_targeted_replay_skips_deck_setup_and_only_sends_selected_pages(project: Project) -> None:
    stage = _pass(project, [json.dumps({"slide": 2, "verdict": "ok"})])

    result = await stage.run(project, _outcome(project), pages=[2])

    assert result.ok
    assert len(stage.composer.seen) == 1  # type: ignore[attr-defined]
    assert "Slide 2 as it renders" in stage.composer.texts()  # type: ignore[attr-defined]
    assert "Slide 1 as it renders" not in stage.composer.texts()  # type: ignore[attr-defined]
    assert result.data["rounds"][0]["selected"] == [2]
    assert result.data["rounds"][0]["deck"]["skipped"] == "targeted replay keeps the shared setup"


@pytest.mark.asyncio
async def test_targeted_replay_can_rewrite_the_shared_setup_once(project: Project) -> None:
    stage = _pass(
        project,
        [json.dumps({"verdict": "ok"}), json.dumps({"slide": 2, "verdict": "ok"})],
    )

    result = await stage.run(project, _outcome(project), pages=[2], setup=True)

    assert result.ok
    assert len(stage.composer.seen) == 2  # type: ignore[attr-defined]
    assert result.data["rounds"][0]["deck"]["verdict"] == "ok"


@pytest.mark.asyncio
async def test_later_round_can_remove_copy_added_by_an_earlier_design_round(project: Project) -> None:
    first = '# SLIDE 1\none = new_slide()\ntitle(one, "Unified video segmentation")\ntitle(one, "temporary caption")\n'
    second = '# SLIDE 1\none = new_slide()\ntitle(one, "Unified video segmentation")\n'
    replies = [
        json.dumps({"verdict": "ok"}),
        json.dumps({"slide": 1, "verdict": "edited", "block": first}),
        json.dumps({"slide": 2, "verdict": "ok"}),
        json.dumps({"slide": 1, "verdict": "edited", "block": second}),
        json.dumps({"slide": 2, "verdict": "ok"}),
    ]

    result = await _pass(project, replies, rounds=2).run(project, _outcome(project))

    assert result.ok
    assert "temporary caption" not in script_path(project).read_text(encoding="utf-8")
    assert result.data["rounds"][1]["rewrote"] == [1]


@pytest.mark.asyncio
async def test_pages_that_cannot_be_told_apart_stop_the_pass_without_editing(project: Project) -> None:
    outcome = BuildOutcome(ok=True, pptx_path=project.build_dir / "deck.pptx", pages=5, sources=())
    result = await _pass(project, []).run(project, outcome)

    assert result.ok
    assert "2 page block(s)" in result.data["rounds"][0]["refused"]["blocks"]
    assert script_path(project).read_text(encoding="utf-8") == SCRIPT


@pytest.mark.asyncio
async def test_nothing_is_attempted_without_a_deck(project: Project) -> None:
    result = await _pass(project, []).run(project, BuildOutcome(ok=False, stderr="boom"))
    assert not result.ok and result.note == "no deck to look at"


async def test_a_deck_in_a_template_is_told_whose_house_it_is_in(project: Project, template_file):
    """The pass rewrites the shared setup and then every page, which is exactly
    where a template's palette and type would be redesigned away. A page that
    comes back in the pass's colours instead of the user's has lost the thing the
    user asked for."""
    from raven.ppt.services.template import bind

    assert bind(template_file(), project) is not None
    pass_ = _pass(project, [])

    await pass_.run(project, _outcome(project))

    briefs = [system for system, _ in pass_.composer.seen]
    assert briefs
    assert all("built inside a template the user supplied" in brief for brief in briefs)


async def test_a_deck_without_one_is_not(project: Project):
    pass_ = _pass(project, [])

    await pass_.run(project, _outcome(project))

    assert all("built inside a template" not in system for system, _ in pass_.composer.seen)


@pytest.mark.asyncio
async def test_a_page_that_adapts_a_template_page_is_told_so(project: Project) -> None:
    """The outline decided the prototype, so the pass that rewrites the page sees it.

    Without it a pass told to improve a page's design has no way to know the design is
    deliberately the template's, and redesigning it is the wrong move -- which is the
    whole reason the prototype is planned rather than left to the program.
    """
    from raven.ppt.contracts.outline import Outline, PagePlan, outline_path, write_outline

    write_outline(
        Outline(takeaway="one model, four tasks", pages=(PagePlan(page=1, claim="c", prototype=5),)),
        outline_path(project),
    )
    stage = _pass(project, [json.dumps({"verdict": "ok"}), json.dumps({"slide": 1, "verdict": "ok"})])

    await stage.run(project, _outcome(project))

    assert "started from the template's page 5" in stage.composer.texts()
    assert "designed reference, not an immutable form" in stage.composer.texts()


@pytest.mark.asyncio
async def test_each_page_receives_its_outline_during_design(project: Project) -> None:
    from raven.ppt.contracts.outline import Outline, PagePlan, outline_path, write_outline

    write_outline(
        Outline(
            takeaway="one model, four tasks",
            pages=(
                PagePlan(
                    page=1,
                    claim="one model handles four tasks",
                    carries="dominant architecture figure",
                    says=("shared representation", "task-specific queries"),
                    figures=("architecture-01",),
                    table_plan={
                        "columns": ("dimension", "system A", "system B"),
                        "rows": (("latency", "1s", "2s"),),
                        "reading": "compare on one baseline",
                    },
                ),
            ),
        ),
        outline_path(project),
    )
    stage = _pass(project, [json.dumps({"verdict": "ok"}), json.dumps({"slide": 1, "verdict": "ok"})])

    await stage.run(project, _outcome(project))

    seen = stage.composer.texts()
    assert "one model handles four tasks" in seen
    assert "dominant architecture figure" in seen
    assert "shared representation" in seen
    assert "architecture-01" in seen
    assert "planned table information shape" in seen
    assert "compare on one baseline" in seen


def test_the_house_numbers_reach_the_brief(tmp_path) -> None:
    """The generic template paragraph kept a pass from repainting a deck; it does not
    stop one from nudging a page title half an inch off the row the template puts it
    on. The numbers the page's author was given travel with it now."""
    from raven.ppt.stages._briefs import deck_brief, page_brief

    house = "- the page title goes at (0.72, 0.14) 11.88x0.98in, 28pt\n- its type ladder: body 18pt"
    for brief in (
        page_brief(body_pt=14, min_pt=10.8, house_style=True, house=house),
        deck_brief(body_pt=14, min_pt=10.8, house_style=True, house=house),
    ):
        assert "built inside a template the user supplied" in brief
        assert "11.88x0.98in, 28pt" in brief
        assert "cloned rather than drawn" in brief
    # And a deck with no template carries neither.
    assert "template the user supplied" not in page_brief(body_pt=14, min_pt=10.8)
    # A template whose style could not be measured still carries the paragraph.
    only_prose = page_brief(body_pt=14, min_pt=10.8, house_style=True)
    assert "built inside a template" in only_prose and "Measured off that template" not in only_prose
