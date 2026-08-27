"""Defects only the render shows, and the mixed ground truth that finds them.

Every threshold in `measure.rendered` is calibrated against one hand-built
reference deck that reports nothing under all three checks, so each is pinned
here twice: as its value, and as a pair of cases either side of it.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from raven.ppt.contracts.findings import Severity
from raven.ppt.services.measure.geometry import Rect
from raven.ppt.services.measure.rendered import (
    CARD_MIN_HEIGHT_PT,
    CARD_MIN_WIDTH_PT,
    CARD_SLOP_PT,
    COLLISION_SHARE,
    COLLISIONS_PER_PAGE,
    OUTSIDE_LINE_SHARE,
    OVERFLOWS_PER_PAGE,
    RULE_BOTTOM_SPARE,
    RULE_MAX_HEIGHT_PT,
    RULE_MIN_WIDTH_PT,
    RULE_TOP_SPARE,
    RULES_PER_PAGE,
    WORD_IN_CARD_SHARE,
    box_overflows,
    card_overflows,
    cards,
    hairline_rules,
    rule_strikes,
    word_collisions,
)
from raven.ppt.services.measure.words import WordBox, by_page, parse_bbox_xml, rect, words_from_pdf
from tests.ppt.conftest import DeckBuilder

pytest.importorskip("pptx")

_DRAWINGML = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _word(text: str, x0: float, y0: float, x1: float, y1: float, page: int = 1) -> WordBox:
    return WordBox(page=page, text=text, x0=x0, y0=y0, x1=x1, y1=y1)


# --- reading the render ------------------------------------------------------


def _bbox(*lines: str) -> str:
    return "\n".join(lines)


def _xml_word(x0: float, y0: float, x1: float, y1: float, text: str) -> str:
    return f'    <word xMin="{x0}" yMin="{y0}" xMax="{x1}" yMax="{y1}">{text}</word>'


def test_word_boxes_are_numbered_by_the_page_they_appear_under() -> None:
    words = parse_bbox_xml(
        _bbox(
            '  <page width="960" height="540">',
            _xml_word(100, 100, 160, 115, "Source:"),
            "  </page>",
            '  <page width="960" height="540">',
            _xml_word(50, 50, 90, 65, "clean"),
            "  </page>",
        )
    )

    assert [(word.page, word.text) for word in words] == [(1, "Source:"), (2, "clean")]
    assert rect(words[0]) == Rect(100.0, 100.0, 160.0, 115.0)
    assert sorted(by_page(words)) == [1, 2]


def test_a_word_before_any_page_is_dropped() -> None:
    """Nothing to attribute it to, and a page number is what a finding needs."""
    assert parse_bbox_xml(_xml_word(1, 1, 2, 2, "orphan")) == []


def test_no_pdftotext_is_no_signal_rather_than_a_clean_deck(monkeypatch: pytest.MonkeyPatch) -> None:
    """None and [] mean different things, and the callers act on the difference."""
    import shutil

    monkeypatch.setattr(shutil, "which", lambda _name: None)

    assert words_from_pdf(Path("nowhere.pdf")) is None


# --- words painted over words -----------------------------------------------


def test_words_sharing_a_place_collide_and_neighbours_do_not() -> None:
    findings = word_collisions(
        [
            _word("Source:", 100, 100, 160, 115),
            _word("OVIS", 105, 102, 150, 116),  # on top of Source:
            _word("material.pdf", 165, 100, 220, 115),  # its neighbour, clear of it
            _word("clean", 50, 50, 90, 65, page=2),
            _word("page", 200, 50, 240, 65, page=2),
        ]
    )

    assert len(findings) == 1
    assert findings[0].page == 1
    assert findings[0].kind == "word_collision"
    assert findings[0].severity is Severity.BLOCKING  # the deliberate exception to D2
    assert "'OVIS'" in findings[0].message


def test_the_overlap_share_is_four_tenths_of_the_smaller_box() -> None:
    """Below it, a generously-sized box grazes its neighbour on paper only.

    Both pairs below overlap. The one covering 38% of the smaller word is inside
    the tolerance; the one covering 42% is not.
    """
    assert COLLISION_SHARE == 0.4
    reported = []
    for shift in (37.2, 34.8):  # shares of 0.38 and 0.42 across a 60pt-wide word
        reported.append(
            len(
                word_collisions([_word("under", 100, 100, 160, 115), _word("over", 100 + shift, 100, 160 + shift, 115)])
            )
        )

    assert reported == [0, 1]


def test_bold_double_painting_is_discounted() -> None:
    """Some renderers paint bold as the same word twice in the same place.

    Any same-text pair is discounted rather than only exact stacks, because the
    second pass is sometimes offset by a pixel or two.
    """
    assert word_collisions([_word("Task", 100, 100, 140, 115), _word("Task", 101, 101, 141, 116)]) == []


def test_one_broken_card_is_reported_a_few_times_not_a_dozen() -> None:
    """The fix is the card, not the twelve pair-wise hits it produces."""
    assert COLLISIONS_PER_PAGE == 4
    words = []
    for index in range(10):
        words.append(_word(f"under{index}", 100, 100 + index, 160, 115 + index))
        words.append(_word(f"over{index}", 102, 101 + index, 158, 116 + index))

    assert len(word_collisions(words)) == COLLISIONS_PER_PAGE


def test_a_zero_area_word_cannot_collide() -> None:
    assert word_collisions([_word("", 100, 100, 100, 100), _word("x", 100, 100, 160, 115)]) == []


# --- rules through words ----------------------------------------------------


def test_a_rule_strikes_through_the_glyphs_and_not_their_edges() -> None:
    """The rule's position is exact in the .pptx; the word's is only true in the
    render, because the renderer grew the rows the rule was drawn between."""
    assert (RULE_TOP_SPARE, RULE_BOTTOM_SPARE) == (0.25, 0.15)
    word = _word("Clip-PanoFCN", 80, 138, 160, 152)  # 14pt tall: struck between 141.5 and 149.9

    reported = []
    for middle in (140.0, 145.0, 151.0):
        rule = Rect(60, middle - 0.5, 200, middle + 0.5)
        reported.append(len(rule_strikes({1: [rule]}, [word])))

    # A rule over the ascenders reads as a border; one under the baseline reads
    # as an underline. Only the one through the middle is a defect.
    assert reported == [0, 1, 0]


def test_a_rule_that_does_not_reach_the_word_is_not_a_strike() -> None:
    word = _word("Method", 80, 138, 160, 152)

    assert rule_strikes({1: [Rect(200, 144, 400, 145)]}, [word]) == []


def test_a_rule_is_reported_once_however_many_words_it_crosses() -> None:
    words = [_word(f"cell{index}", 80 + 90 * index, 138, 160 + 90 * index, 152) for index in range(5)]

    findings = rule_strikes({1: [Rect(60, 144, 600, 145)]}, words)

    assert len(findings) == 1
    assert findings[0].kind == "rule_strike"


def test_rules_are_reported_a_few_per_page() -> None:
    assert RULES_PER_PAGE == 3
    rules = [Rect(60, 144, 600, 145) for _ in range(5)]

    assert len(rule_strikes({1: rules}, [_word("struck", 80, 138, 160, 152)])) == RULES_PER_PAGE


def test_a_hairline_is_thin_and_long(deck: DeckBuilder) -> None:
    """Thicker than this is a bar, shorter than this is a tick or a bullet."""
    from pptx.util import Emu, Pt

    assert (RULE_MAX_HEIGHT_PT, RULE_MIN_WIDTH_PT) == (4.5, 36.0)
    page = deck.page()
    for height_pt, width_pt in ((4.5, 36.0), (4.6, 36.0), (4.5, 35.0)):
        shape = deck.panel(page, left=1.0, top=1.0, width=1.0, height=1.0)
        shape.height, shape.width = Pt(height_pt), Pt(width_pt)
        shape.top, shape.left = Emu(0), Emu(0)

    assert len(hairline_rules(deck.save())[1]) == 1


def test_a_rule_carrying_text_is_not_a_rule(deck: DeckBuilder) -> None:
    """A one-line label in a shallow box is copy, not a divider."""
    from pptx.util import Pt

    shape = deck.panel(deck.page(), left=1.0, top=1.0, width=4.0, height=1.0, text="Method")
    shape.height = Pt(4.0)

    assert hairline_rules(deck.save())[1] == []


# --- words escaping their card ----------------------------------------------


_CARD = Rect(72, 72, 360, 216)  # 1in,1in to 5in,3in


def test_a_word_running_past_its_card_is_reported() -> None:
    """Copy that outgrows its card lands on blank page, which no collision sees."""
    findings = card_overflows(
        {1: [_CARD]},
        [
            _word("inside", 100, 100, 180, 115),  # comfortably within the card
            _word("identical", 100, 210, 180, 226),  # bottom edge 10pt past the card's
            _word("Title", 500, 100, 580, 115),  # page furniture, in no card
        ],
    )

    assert len(findings) == 1
    assert findings[0].detail["words"] == ["identical"]
    assert findings[0].kind == "card_overflow"
    # The number to act on, not just the fact that something spilled: the copy
    # reaches 10pt past the card's bottom edge, so that is what it has to grow by.
    assert findings[0].detail["needs_height_in"] == pytest.approx((_CARD.height + 10 + CARD_SLOP_PT) / 72, abs=0.01)
    assert "do not shrink the type" in findings[0].message


def test_the_slop_is_three_points_of_border_and_antialiasing() -> None:
    assert CARD_SLOP_PT == 3.0
    reported = []
    for escape in (2.9, 3.1):
        word = _word("copy", 100, 200 + escape - 16, 180, 216 + escape)
        reported.append(len(card_overflows({1: [_CARD]}, [word])))

    assert reported == [0, 1]


def test_a_word_belongs_to_a_card_that_holds_a_third_of_it() -> None:
    """Deliberately not "the card contains the word's centre".

    Both words below have escaped far enough that their centre is outside the
    card, which is precisely the word most worth reporting -- a centre rule would
    drop both. The share rule keeps the one the card still holds a third of.
    """
    assert WORD_IN_CARD_SHARE == 0.3
    held = _word("held", 320, 200, 440, 216)  # a third inside, centre outside
    gone = _word("gone", 326, 200, 446, 216)  # 28% inside

    assert [finding.detail["words"] for finding in card_overflows({1: [_CARD]}, [held, gone])] == [["held"]]


def test_a_word_belongs_to_the_smallest_card_holding_it() -> None:
    """Cards nest: a stat sits in its own tile inside a band.

    Measured against the band, the word below is comfortably inside; against the
    tile it sits in, it has run out of the bottom.
    """
    outer, inner = Rect(72, 72, 600, 400), Rect(100, 100, 300, 200)

    findings = card_overflows({1: [outer, inner]}, [_word("overflowing", 110, 190, 200, 210)])

    assert len(findings) == 1


def test_one_card_is_one_finding_however_many_words_escape() -> None:
    """Eight words out of one card is one card too small, not eight problems.

    Reporting each word put nine findings on one deck for three cards, and the author
    read them as nine separate things to move.
    """
    words = [_word(f"word{index}", 100 + 10 * index, 210, 180 + 10 * index, 226) for index in range(8)]

    findings = card_overflows({1: [_CARD]}, words)

    assert len(findings) == 1
    assert len(findings[0].detail["words"]) == 6, "and it names a few of them, not all eight"


def test_overflowing_cards_are_reported_a_few_per_page_and_the_rest_are_counted() -> None:
    """The cap keeps a page of broken cards from burying the rest of the report, but
    a silent cap reads as "four cards overflow" when eight do."""
    assert OVERFLOWS_PER_PAGE == 4
    cards = [Rect(72 + 200 * index, 72, 172 + 200 * index, 172) for index in range(8)]
    words = [_word(f"word{index}", 80 + 200 * index, 166, 160 + 200 * index, 182) for index in range(8)]

    findings = card_overflows({1: cards}, words)

    listed = [finding for finding in findings if "unlisted_cards" not in finding.detail]
    rest = [finding for finding in findings if "unlisted_cards" in finding.detail]
    assert len(listed) == OVERFLOWS_PER_PAGE
    assert [finding.detail["unlisted_cards"] for finding in rest] == [8 - OVERFLOWS_PER_PAGE]


def test_a_page_with_no_cards_has_nothing_to_escape() -> None:
    assert card_overflows({1: []}, [_word("Title", 100, 100, 180, 115)]) == []


def test_a_card_is_a_filled_panel_big_enough_to_hold_copy(deck: DeckBuilder) -> None:
    from pptx.util import Emu, Pt

    assert (CARD_MIN_WIDTH_PT, CARD_MIN_HEIGHT_PT) == (72.0, 36.0)
    page = deck.page()
    for width_pt, height_pt, filled in (
        (72.0, 36.0, True),
        (71.0, 36.0, True),
        (72.0, 35.0, True),
        (72.0, 36.0, False),
    ):
        shape = deck.panel(page, left=1.0, top=1.0, width=1.0, height=1.0, filled=filled)
        shape.width, shape.height = Pt(width_pt), Pt(height_pt)
        shape.top, shape.left = Emu(0), Emu(0)

    assert len(cards(deck.save())[1]) == 1


def test_a_panel_carrying_copy_is_not_a_card_the_copy_can_escape(deck: DeckBuilder) -> None:
    """Its own text is inside it by construction; what escapes a card is other
    copy laid over it."""
    deck.panel(deck.page(), left=1.0, top=1.0, width=4.0, height=2.0, text="Results")

    assert cards(deck.save())[1] == []


# --- copy escaping its own text box ------------------------------------------


def _one_box(tmp_path: Path, *, text: str = "the copy this box was given", width: float = 3.0) -> Path:
    """One text box at (1.0, 1.0), `width` wide and 0.6in tall -- 72,72 to x,115.2 in points."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(width), Inches(0.6)).text_frame.text = text
    path = tmp_path / f"box-{width}-{len(text)}.pptx"
    presentation.save(str(path))
    return path


def test_copy_the_render_sets_below_its_box_is_reported(tmp_path: Path) -> None:
    """A live deck shipped a page whose last line rendered 0.30in under the box that
    holds it, over the template's corner ornament, and all sixteen findings on that deck
    were about something else: `off_page` compares against the canvas edge and the line
    was still on the canvas, `spilled_copy` is about a box with wrapping off, and
    `overset_copy` predicts the height from font metrics and under-read it."""
    deck = _one_box(tmp_path)
    words = [
        _word("the copy this box", 76, 76, 260, 94),
        _word("was given", 76, 118, 200, 136),  # wholly below the box, which ends at 115.2
    ]

    findings = box_overflows(deck, words)

    assert [f.kind for f in findings] == ["box_overflow"]
    assert findings[0].severity is Severity.WARNING
    assert findings[0].page == 1
    assert findings[0].detail["words"] == ["was given"]
    # The numbers to act on: 136 - 115.2 past a 0.6in box, so the copy took 0.89in.
    assert findings[0].detail["below_in"] == pytest.approx((136 - 115.2) / 72, abs=0.01)
    assert findings[0].detail["needs_in"] == pytest.approx(0.6 + (136 - 115.2) / 72, abs=0.01)
    assert "do not shrink the type to fit" in findings[0].message


def test_copy_inside_its_box_is_not(tmp_path: Path) -> None:
    deck = _one_box(tmp_path)

    assert box_overflows(deck, [_word("the copy this box was given", 76, 76, 260, 94)]) == []


def test_the_slack_is_half_the_escaping_line(tmp_path: Path) -> None:
    """A word's bbox is the font's box and not its ink, so the last line of a box sized
    exactly to its copy always ends a little below it -- 0.12 of the line's own height at
    worst on the deck this was measured against. Half a line is where that stops being
    the explanation."""
    assert OUTSIDE_LINE_SHARE == 0.5
    deck = _one_box(tmp_path)
    reported = []
    for below in (0.45 * 18, 0.55 * 18):
        tail = _word("was given", 76, 115.2 + below - 18, 200, 115.2 + below)
        reported.append(len(box_overflows(deck, [_word("the copy this box", 76, 76, 260, 94), tail])))

    assert reported == [0, 1]


def test_a_line_below_a_neighbouring_column_is_not_this_box_s_overflow(tmp_path: Path) -> None:
    """The tail of one column otherwise reads as the overflow of whatever box happens to
    sit above it in the other."""
    deck = _one_box(tmp_path)

    assert (
        box_overflows(
            deck,
            [
                _word("the copy this box", 76, 76, 260, 94),
                _word("another column", 500, 118, 620, 136),
            ],
        )
        == []
    )


def test_a_line_a_full_line_below_the_box_belongs_to_whatever_wrote_it(tmp_path: Path) -> None:
    """Copy flows, so the first line past a box starts within one line of where the box
    ended. Further down is the next block on the page, not this box's overflow."""
    deck = _one_box(tmp_path)

    assert (
        box_overflows(
            deck,
            [
                _word("the copy this box", 76, 76, 260, 94),
                _word("a separate block", 76, 140, 200, 158),
            ],
        )
        == []
    )


def test_a_box_the_render_put_nothing_in_owns_no_overflow(tmp_path: Path) -> None:
    deck = _one_box(tmp_path)

    assert box_overflows(deck, [_word("was given", 76, 118, 200, 136)]) == []


def test_copy_the_render_does_not_show_is_named() -> None:
    """A shape narrower than its own words clips instead of wrapping, and every other
    check passes: the words that did render sit exactly where they belong."""
    from raven.ppt.services.measure.rendered import clipped_copy

    declared = {1: "Backbone 提取帧级特征 Transformer Decoder 读出掩码 VOS 路径以首帧掩码为提示，四类任务共享同一主干"}
    shown = [
        _word("Bac", 10, 10, 30, 24),
        _word("提取帧级特征", 40, 10, 120, 24),
        _word("Tran", 130, 10, 160, 24),
        _word("读出掩码", 170, 10, 230, 24),
        _word("VOS", 240, 10, 270, 24),
        _word("路径以首帧掩码为提示，四类任务共享同一主干", 10, 40, 300, 54),
    ]
    findings = clipped_copy(Path("unused.pptx"), shown, page_texts=declared)

    assert [f.kind for f in findings] == ["clipped_copy"]
    assert findings[0].severity.value == "warning"
    assert "Backbone" in findings[0].message and "Transformer" in findings[0].message
    assert findings[0].detail["shown"] < 0.9


def test_a_page_the_render_shows_whole_is_not_reported() -> None:
    """The renderer breaks lines and splits runs, so the comparison has to survive
    the same words arriving in a different order and in different pieces."""
    from raven.ppt.services.measure.rendered import clipped_copy

    declared = {1: "Backbone 提取帧级特征，Transformer Decoder 读出掩码，四类任务共享同一主干与同一套查询"}
    shown = [
        _word("提取帧级特征，", 40, 40, 120, 54),
        _word("Backbone", 10, 10, 60, 24),
        _word("Transformer", 130, 10, 190, 24),
        _word("Decoder", 200, 10, 250, 24),
        _word("读出掩码，四类任务共享同一主干与同一套查询", 10, 70, 320, 84),
    ]

    assert clipped_copy(Path("unused.pptx"), shown, page_texts=declared) == []


def test_a_page_with_no_extracted_text_is_not_a_clipped_page() -> None:
    """No text layer is no signal. A deck rendered by something that draws type as
    curves would otherwise report every page as 0% shown."""
    from raven.ppt.services.measure.rendered import clipped_copy

    declared = {1: "这一页有足够长的正文，用来越过门禁的字符下限，但渲染结果里一个词都没有取出来。"}

    assert clipped_copy(Path("unused.pptx"), [], page_texts=declared) == []


def test_copy_touching_the_bottom_of_its_panel_is_reported() -> None:
    """Found by reading a polished deck page by page: two note cards ended exactly on
    their last line's descender while every other card on the same deck carried 0.20in
    of padding. Nothing overflowed, nothing collided, and the page read as cramped in a
    way no measurement had a name for."""
    from raven.ppt.services.measure.rendered import crowded_panels

    card = Rect(100, 100, 400, 200)
    words = [_word("inside", 120, 120, 200, 134), _word("flush", 120, 186, 200, 200)]

    findings = crowded_panels({1: [card]}, words)
    assert [f.kind for f in findings] == ["crowded_panel"]
    assert findings[0].detail["side"] == "bottom"
    assert findings[0].detail["gap_in"] == 0.0


def test_a_panel_with_padding_is_not() -> None:
    from raven.ppt.services.measure.rendered import crowded_panels

    assert crowded_panels({1: [Rect(100, 100, 400, 200)]}, [_word("comfortable", 120, 120, 240, 134)]) == []


def test_one_finding_per_panel_not_per_word() -> None:
    """A line of eight words against the rim is one cramped panel and one fix."""
    from raven.ppt.services.measure.rendered import crowded_panels

    words = [_word(f"w{index}", 110 + index * 30, 186, 135 + index * 30, 200) for index in range(8)]

    assert len(crowded_panels({1: [Rect(100, 100, 400, 200)]}, words)) == 1


def _blank_wide_deck(tmp_path: Path, *, panel: bool = False) -> Path:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if panel:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.5), Inches(5), Inches(3.5))
        shape.fill.solid()
    deck = tmp_path / ("panel.pptx" if panel else "blank.pptx")
    presentation.save(str(deck))
    return deck


def test_a_large_rendered_gap_between_body_groups_is_reported(tmp_path: Path) -> None:
    from raven.ppt.services.measure.rendered import excessive_whitespace

    deck = _blank_wide_deck(tmp_path)
    words = [
        _word("upper", 72, 110, 150, 128),
        _word("upper", 72, 136, 150, 154),
        _word("lower", 72, 300, 150, 318),
    ]

    findings = excessive_whitespace(deck, words)
    assert findings[0].kind == "excessive_whitespace"
    assert findings[0].detail["region"] == "between_groups"
    assert findings[0].detail["gap_in"] > 1.5


def test_a_page_that_uses_the_body_height_is_not_sparse(tmp_path: Path) -> None:
    from raven.ppt.services.measure.rendered import excessive_whitespace

    deck = _blank_wide_deck(tmp_path)
    words = [_word(f"line{top}", 72, top, 180, top + 18) for top in range(100, 481, 45)]

    assert excessive_whitespace(deck, words) == []
    assert excessive_whitespace(deck, [_word("small", 72, 110, 140, 128)], structural=[1]) == []


def test_a_large_panel_with_copy_only_at_the_top_is_reported(tmp_path: Path) -> None:
    from raven.ppt.services.measure.rendered import excessive_whitespace

    deck = _blank_wide_deck(tmp_path, panel=True)
    words = [_word("heading", 90, 125, 180, 145), _word("one line", 90, 155, 220, 175)]

    findings = excessive_whitespace(deck, words)
    assert "empty_panel" in {finding.detail["region"] for finding in findings}


def test_header_rows_that_are_too_far_apart_are_reported(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from raven.ppt.services.measure.rendered import excessive_whitespace

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(0.20), Inches(8), Inches(0.35)).text = "title"
    slide.shapes.add_textbox(Inches(1), Inches(0.95), Inches(8), Inches(0.30)).text = "explanation"
    deck = tmp_path / "loose-header.pptx"
    presentation.save(str(deck))
    words = [_word("title", 72, 18, 150, 38), _word("explanation", 72, 72, 180, 90)]

    findings = excessive_whitespace(deck, words)
    assert "loose_header" in {finding.detail["region"] for finding in findings}


def test_a_label_broken_one_character_short_is_reported(tmp_path: Path) -> None:
    """Four of eight labels on one delivered agenda page read "为什么要统 / 一". The words
    do not collide, the copy fits the box's height, the type is the right size, and the
    page reads as sloppy.

    Read off the render rather than predicted: the first version wrapped the copy with
    the font measurer and claimed a title broke that the render shows on one line.
    """
    from pptx import Presentation
    from pptx.util import Inches

    from raven.ppt.services.measure.rendered import orphan_lines

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.8))
    box.text_frame.text = "任务碎片化：为什么要统一"
    deck = tmp_path / "deck.pptx"
    presentation.save(str(deck))

    words = [
        _word("任务碎片化：为什么要统", 72, 72, 280, 90),
        _word("一", 72, 100, 90, 118),
    ]

    findings = orphan_lines(deck, words)
    assert [f.kind for f in findings] == ["orphan_line"]
    assert findings[0].detail["orphan"] == "一"


def test_a_label_that_did_not_break_is_not(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from raven.ppt.services.measure.rendered import orphan_lines

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(0.8))
    box.text_frame.text = "任务碎片化：为什么要统一"
    deck = tmp_path / "deck.pptx"
    presentation.save(str(deck))

    assert orphan_lines(deck, [_word("任务碎片化：为什么要统一", 72, 72, 380, 90)]) == []


def test_a_break_the_author_wrote_is_not_an_orphan(tmp_path: Path) -> None:
    """A two-line chevron label written as "提交并\n推送" used to report the same as a label
    the box was too narrow to hold, and one live run spent three consecutive iterations
    arguing back that its labels were meant to read that way instead of acting on
    anything. It was right: the first is typography and the second is a defect.

    The form pinned here is the one the file actually carries. `layout.write` assigns
    the paragraph's text, and python-pptx turns a `\n` in it into an `a:br` -- which
    reads back as `\v`, not as the `\n` that went in, so a check looking for the
    written character would never have found one.
    """
    from pptx import Presentation
    from pptx.util import Inches

    from raven.ppt.services.measure.rendered import orphan_lines

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1.4), Inches(1.0))
    box.text_frame.paragraphs[0].text = "提交并\n推送"
    deck = tmp_path / "written-break.pptx"
    presentation.save(str(deck))

    assert "\v" in box.text_frame.text
    assert box.text_frame.paragraphs[0]._p.findall(f"{{{_DRAWINGML}}}br")

    words = [_word("提交并", 76, 76, 130, 94), _word("推送", 76, 100, 112, 118)]

    assert orphan_lines(deck, words) == []


def test_a_genuine_orphan_in_a_box_that_also_carries_a_written_break_still_fires(tmp_path: Path) -> None:
    """The skip is per break and not per shape: a box that carries a written break
    somewhere in it is still checked everywhere else."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven.ppt.services.measure.rendered import orphan_lines

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1.2))
    box.text_frame.paragraphs[0].text = "提交并推送\n任务碎片化：为什么要统一"
    deck = tmp_path / "break-and-orphan.pptx"
    presentation.save(str(deck))

    words = [
        _word("提交并推送", 76, 76, 166, 94),
        _word("任务碎片化：为什么要统", 76, 100, 280, 118),
        _word("一", 76, 124, 94, 142),
    ]

    findings = orphan_lines(deck, words)

    assert [f.kind for f in findings] == ["orphan_line"]
    assert findings[0].detail["orphan"] == "一"


def _stacked_deck(tmp_path: Path, second_top_in: float) -> Path:
    """Two text boxes in one column: the first two lines deep, the second one line."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for top, height in ((1.0, 0.6), (second_top_in, 0.4)):
        box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(4), Inches(height))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "group"
        run.font.size = Pt(16)
    path = tmp_path / f"deck-{second_top_in}.pptx"
    presentation.save(str(path))
    return path


def test_two_groups_with_no_air_between_them_are_reported(tmp_path: Path) -> None:
    """The design brief asks for this in prose -- "a gap that is plainly wider than the
    gaps inside each group" -- and nothing measured it. Found by reading a deck page by
    page: four task definitions in one column ran together into a single grey mass."""
    from raven.ppt.services.measure.rendered import unseparated_blocks

    deck = _stacked_deck(tmp_path, 1.65)
    words = [
        _word("first", 72, 74, 140, 92),
        _word("line", 72, 96, 140, 114),
        _word("second", 72, 122, 140, 140),
    ]

    findings = unseparated_blocks(deck, words)
    assert [f.kind for f in findings] == ["unseparated_blocks"]
    assert findings[0].detail["gap_in"] < 0.4


def test_a_visible_gap_is_not(tmp_path: Path) -> None:
    from raven.ppt.services.measure.rendered import unseparated_blocks

    deck = _stacked_deck(tmp_path, 3.0)
    words = [
        _word("first", 72, 74, 140, 92),
        _word("line", 72, 96, 140, 114),
        _word("second", 72, 220, 140, 238),
    ]

    assert unseparated_blocks(deck, words) == []
