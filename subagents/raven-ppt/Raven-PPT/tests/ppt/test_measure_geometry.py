"""Where a shape actually is, which for anything grouped is not what it says.

Every geometric check in this package reads a rectangle off a shape, and for a
shape inside a group the numbers python-pptx hands back are in the group's own
coordinate space. Measured on a real template's agenda page: six section numbers
laid out two columns by three rows all reported at the same (7.67, 4.27) and 4.46in
wide -- so `covered_shape` said five of the six were completely hidden behind the
sixth, and refused to publish a deck whose page was fine.
"""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("pptx", reason="ppt extra not installed")

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from raven.ppt.services.measure.geometry import EMU_PER_INCH, page_box, shape_rect_emu

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _transform(group, *, at, size, child_at, child_size):
    """Give `group` the child space a designer's template gives it.

    Written after the children are in, not before: python-pptx recomputes a group's
    xfrm from its children every time one is added, and what it writes is the
    identity (off == chOff, ext == chExt). A real template's group is drawn at a
    different size from the space its children were laid out in, which is the whole
    reason a child's own numbers are not page numbers.
    """
    properties = group._element.find(f"{_P}grpSpPr")
    for existing in properties.findall(f"{_A}xfrm"):
        properties.remove(existing)
    frame = etree.SubElement(properties, f"{_A}xfrm")
    for tag, (x, y), keys in (
        ("off", at, ("x", "y")),
        ("ext", size, ("cx", "cy")),
        ("chOff", child_at, ("x", "y")),
        ("chExt", child_size, ("cx", "cy")),
    ):
        element = etree.SubElement(frame, f"{_A}{tag}")
        element.set(keys[0], str(int(x * EMU_PER_INCH)))
        element.set(keys[1], str(int(y * EMU_PER_INCH)))
    return group


def _deck():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    return presentation, presentation.slides.add_slide(presentation.slide_layouts[6])


def test_a_shape_at_the_top_level_is_where_it_says() -> None:
    _, slide = _deck()
    shape = slide.shapes.add_textbox(Inches(2), Inches(1.5), Inches(3), Inches(0.5))

    box = page_box(shape)

    assert (round(box.x0, 3), round(box.y0, 3)) == (2.0, 1.5)
    assert (round(box.x1 - box.x0, 3), round(box.y1 - box.y0, 3)) == (3.0, 0.5)


def test_a_grouped_shape_is_moved_and_scaled_by_its_group() -> None:
    """The group is drawn at half the size of the space its children were laid out in,
    so a child at (1, 1) of that space lands half an inch into the group."""
    _, slide = _deck()
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.5))
    _transform(group, at=(5.0, 3.0), size=(4.0, 2.0), child_at=(0.0, 0.0), child_size=(8.0, 4.0))

    box = page_box(child)

    assert (round(box.x0, 3), round(box.y0, 3)) == (5.5, 3.5)
    assert (round(box.x1 - box.x0, 3), round(box.y1 - box.y0, 3)) == (1.0, 0.25)
    # And the EMU form every gate reads goes through the same transform.
    emu = shape_rect_emu(child)
    assert round(emu.x0 / EMU_PER_INCH, 3) == 5.5


def test_a_child_space_with_an_offset_is_subtracted_first() -> None:
    """A template's group is as likely to start its child space at its own offset."""
    _, slide = _deck()
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Inches(3), Inches(2), Inches(1), Inches(1))
    _transform(group, at=(1.0, 1.0), size=(4.0, 4.0), child_at=(2.0, 2.0), child_size=(4.0, 4.0))

    box = page_box(child)

    assert (round(box.x0, 3), round(box.y0, 3)) == (2.0, 1.0)


def test_groups_inside_groups_compose() -> None:
    _, slide = _deck()
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    child = inner.shapes.add_textbox(Inches(4), Inches(0), Inches(2), Inches(2))
    _transform(inner, at=(2.0, 2.0), size=(4.0, 4.0), child_at=(0.0, 0.0), child_size=(8.0, 8.0))
    _transform(outer, at=(4.0, 0.0), size=(4.0, 4.0), child_at=(0.0, 0.0), child_size=(8.0, 8.0))

    box = page_box(child)

    # inner puts the child at (2+2, 2+0) of the outer space, outer halves it again.
    assert (round(box.x0, 3), round(box.y0, 3)) == (6.0, 1.0)
    assert round(box.x1 - box.x0, 3) == 0.5


def test_a_shape_with_no_geometry_has_no_box(tmp_path: Path) -> None:
    """A placeholder inheriting its position from the layout states none of its own."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    inherited = [shape for shape in slide.shapes if shape.left is None]
    if not inherited:  # pragma: no cover -- depends on the bundled default template
        pytest.skip("this python-pptx default states every placeholder's position")

    assert page_box(inherited[0]) is None
    assert shape_rect_emu(inherited[0]).area == 0
    del tmp_path


def _table_slide(rows: int, columns: int, *, at=(1.0, 1.0), size=(6.0, 3.0)):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_table(
        rows, columns, Inches(at[0]), Inches(at[1]), Inches(size[0]), Inches(size[1])
    )
    return presentation, slide, frame


def test_a_table_cell_reports_the_rectangle_it_covers() -> None:
    """A cell is not a shape, so nothing walking `iter_shapes` can ask where it is."""
    from raven.ppt.services.measure.geometry import cell_boxes

    _presentation, slide, frame = _table_slide(2, 3, at=(1.0, 1.0), size=(6.0, 2.0))
    table = frame.table
    for down in range(2):
        for across in range(3):
            table.cell(down, across).text = f"r{down}c{across}"

    boxes = {cell.text: box for cell, box in cell_boxes(slide)}

    assert len(boxes) == 6
    # Columns are equal thirds of 6in and rows halves of 2in, in points.
    first = boxes["r0c0"]
    assert (round(first.x0), round(first.y0)) == (72, 72)
    assert round(first.width) == 144 and round(first.height) == 72
    last = boxes["r1c2"]
    assert (round(last.x1), round(last.y1)) == (round(7.0 * 72), round(3.0 * 72))


def test_an_empty_cell_is_not_one_of_them() -> None:
    from raven.ppt.services.measure.geometry import cell_boxes

    _presentation, slide, frame = _table_slide(2, 2)
    frame.table.cell(0, 0).text = "held"

    assert [cell.text for cell, _box in cell_boxes(slide)] == ["held"]


def test_a_merged_cell_is_reported_once_over_the_span_it_covers() -> None:
    """From the origin, over the whole span: a band welded across three columns is one
    rectangle, and reading the origin's own column would say it is a third as wide.

    A spanned cell keeps a text frame of its own that a file can carry copy in -- the
    merge hides it and the render never draws it -- so it is skipped by what it is and
    not by being empty.
    """
    from raven.ppt.services.measure.geometry import cell_boxes

    _presentation, slide, frame = _table_slide(2, 3, at=(0.0, 0.0), size=(6.0, 2.0))
    table = frame.table
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).text = "one band"
    table.cell(1, 0).text = "under it"
    table.cell(0, 1).text_frame.text = "hidden by the merge"
    assert table.cell(0, 1).is_spanned

    boxes = {cell.text_frame.text: box for cell, box in cell_boxes(slide)}

    assert sorted(boxes) == ["one band", "under it"]
    assert round(boxes["one band"].width) == round(6.0 * 72), "the band spans all three columns"
    assert round(boxes["under it"].width) == round(2.0 * 72)


def test_a_table_inside_a_group_lands_where_the_group_puts_it() -> None:
    """The boundaries are fractions of the frame's own rectangle rather than a running
    sum of the declared widths, so the group's scale is carried without applying it
    twice."""
    from raven.ppt.services.measure.geometry import cell_boxes

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    frame = slide.shapes.add_table(1, 2, Inches(0), Inches(0), Inches(8), Inches(2))
    frame.table.cell(0, 0).text = "left"
    frame.table.cell(0, 1).text = "right"
    # python-pptx cannot add a table to a group, so the element is moved the way a
    # designer's template already carries one.
    group._element.append(frame._element)
    _transform(group, at=(2.0, 1.0), size=(4.0, 1.0), child_at=(0.0, 0.0), child_size=(8.0, 2.0))

    boxes = {cell.text: box for cell, box in cell_boxes(slide)}

    # The group halves the table, so each 4in column comes out 2in wide at x=2in.
    assert round(boxes["left"].x0) == round(2.0 * 72)
    assert round(boxes["left"].width) == round(2.0 * 72)
    assert round(boxes["right"].x0) == round(4.0 * 72)
