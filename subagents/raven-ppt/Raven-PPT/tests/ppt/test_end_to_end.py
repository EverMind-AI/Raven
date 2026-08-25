"""The whole route, with no model in it: ingest, build, measure, deliver.

Deterministic on purpose. A run with a model in it proves that a model can drive
these tools; this proves the tools themselves hold together -- that ingest writes
what the gates read, that the backend's page mapping survives to the design pass,
that a refusal actually refuses, and that a delivered file is the one that was
measured. Those are the seams where the predecessor's two routes drifted apart,
and none of them are visible from a unit test of either side.

The deck is written the way an author is told to write one: a shared prelude, then
one block per page, each opening with a `# SLIDE n` banner.
"""

from __future__ import annotations

import json
import textwrap
from pathlib import Path

import pytest

from raven.ppt.contracts import (
    Audience,
    DeckBrief,
    PageBudget,
    Project,
    Severity,
    brief_path,
    write_brief,
)
from raven.ppt.profiles import registry
from raven.ppt.services.render import available
from raven.ppt.stages._measure import DeckMeasurer
from raven.ppt.stages._views import DeckViews
from raven.ppt.stages.build import BuildStage
from raven.ppt.tools.ingest import PptIngestTool

pytest.importorskip("pptx")
pytest.importorskip("fitz")

MATERIAL = """\
# TarViS: A Unified Approach for Target-Based Video Segmentation

TarViS reaches 48.3 AP on YouTube-VIS 2021 with a Swin-L backbone.
It is trained jointly on four tasks and needs no task-specific fine-tuning.
Figure 1: the shared query interface across the four task definitions.
"""

DECK = textwrap.dedent(
    """
    import os

    from pptx import Presentation
    from pptx.util import Inches, Pt

    from ppt_theme import THEMES, rgb

    TH = THEMES["ink-graphite"]
    ACC = rgb(TH["accent"])
    FG = rgb(TH["foreground"])

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    def title(slide, text, size=30):
        box = slide.shapes.add_textbox(Inches(0.9), Inches(0.7), Inches(11.5), Inches(1.1))
        frame = box.text_frame
        frame.text = text
        run = frame.paragraphs[0].runs[0]
        run.font.size = Pt(size)
        run.font.color.rgb = FG


    def body(slide, text, top=2.2, size=18):
        box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(1.2))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = text
        run = frame.paragraphs[0].runs[0]
        run.font.size = Pt(size)
        run.font.color.rgb = FG


    # SLIDE 1
    cover = new_slide()
    title(cover, "TarViS: unified target-based video segmentation")
    body(cover, "One query interface. Four segmentation tasks. One jointly trained model.")

    # SLIDE 2
    claim = new_slide()
    title(claim, "One model, four task definitions")
    body(claim, "Trained jointly on four tasks, with no task-specific fine-tuning at inference.")

    # SLIDE 3
    result = new_slide()
    title(result, "48.3 AP on YouTube-VIS 2021")
    body(result, "Reached with a Swin-L backbone, from the same jointly trained model.")

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    materials = tmp_path / "materials"
    materials.mkdir()
    (materials / "paper.md").write_text(MATERIAL, encoding="utf-8")
    return tmp_path


async def _ingest(workspace: Path) -> dict:
    reply = await PptIngestTool(workspace).execute(project="tarvis", materials_dir="materials")
    return json.loads(reply)


def _stage(profile: str = "script_author") -> tuple[BuildStage, DeckViews]:
    from raven.ppt.backends.script import asset_helpers, run_script

    helpers = asset_helpers()

    async def backend(project, script):
        return await run_script(project, script, helpers=helpers)

    views = DeckViews(concurrency=1)
    return (
        BuildStage(
            backend=backend,
            measure=DeckMeasurer(views=views),
            profile=registry.get(profile),
            destination=lambda project: project.exports_dir / "TarViS.pptx",
        ),
        views,
    )


@pytest.mark.asyncio
async def test_the_route_runs_from_materials_to_a_delivered_deck(workspace: Path) -> None:
    ingested = await _ingest(workspace)
    assert ingested["ok"] is True

    project = Project(workspace=workspace, slug="tarvis")
    stage, _views = _stage()
    result = await stage.run(project, DECK)

    assert result.ok, [f.message for f in result.findings] or result.note
    delivered = Path(result.data["pptx_path"])
    assert delivered.is_file() and delivered.name == "TarViS.pptx"

    from pptx import Presentation

    assert len(Presentation(str(delivered)).slides) == 3
    # Page-to-code mapping survived from execution, which is what the design pass
    # needs and what nothing downstream can reconstruct.
    assert [source.page for source in result.data["outcome"].sources] == [1, 2, 3]


@pytest.mark.asyncio
async def test_a_deck_whose_pages_cannot_be_told_apart_is_refused(workspace: Path) -> None:
    await _ingest(workspace)
    project = Project(workspace=workspace, slug="tarvis")
    stage, _views = _stage()

    looped = textwrap.dedent(
        """
        import os

        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for text in ("one", "two", "three"):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1))
            box.text_frame.text = text
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        prs.save(os.environ["PPT_OUTPUT"])
        """
    ).lstrip()
    result = await stage.run(project, looped)

    assert not result.ok
    unmapped = [f for f in result.findings if f.kind == "unmapped_page"]
    assert unmapped and "# SLIDE 1" in unmapped[0].message


@pytest.mark.asyncio
async def test_the_delivered_file_is_the_one_that_was_measured(workspace: Path) -> None:
    """A deck edited between the checks and the delivery describes nothing."""
    await _ingest(workspace)
    project = Project(workspace=workspace, slug="tarvis")
    stage, _views = _stage()
    result = await stage.run(project, DECK)

    assert result.ok
    delivered = Path(result.data["pptx_path"]).read_bytes()
    built = result.data["outcome"].pptx_path.read_bytes()
    assert delivered == built


@pytest.mark.asyncio
async def test_type_under_the_floor_is_reported_and_still_delivered(workspace: Path) -> None:
    """Shrinking the copy would satisfy the measurement and make the page worse."""
    if not available().can_measure_words:
        pytest.skip("needs pypdfium2 or poppler to read the render")
    await _ingest(workspace)
    project = Project(workspace=workspace, slug="tarvis")
    stage, _views = _stage()

    small = DECK.replace("def body(slide, text, top=2.2, size=18):", "def body(slide, text, top=2.2, size=11):")
    result = await stage.run(project, small)

    assert result.ok, [f.message for f in result.findings]
    floors = [f for f in result.findings if f.kind == "type_floor"]
    assert floors, [f.kind for f in result.findings]
    assert all(f.severity is Severity.WARNING for f in floors)
    assert all(f.audience is Audience.DESIGNER for f in floors)
    assert Path(result.data["pptx_path"]).is_file()


@pytest.mark.asyncio
async def test_the_render_chain_produces_a_page_image_per_slide(workspace: Path) -> None:
    caps = available()
    if not (caps.can_convert and caps.can_rasterise):
        pytest.skip(f"needs LibreOffice and a PDF rasteriser: {caps.explain()}")
    await _ingest(workspace)
    project = Project(workspace=workspace, slug="tarvis")
    stage, views = _stage()
    result = await stage.run(project, DECK)
    assert result.ok

    renders = await views.pages(Path(result.data["pptx_path"]), project.review_dir, [1, 2, 3])
    assert sorted(renders) == [1, 2, 3]
    assert all(path.is_file() and path.stat().st_size > 0 for path in renders.values())


@pytest.mark.asyncio
async def test_a_deck_that_is_not_the_agreed_length_stops_delivery(workspace: Path) -> None:
    """The page budget is the user's decision, so it has to bind something.

    The predecessor's schema route enforced it in code and its script route
    enforced nothing, which made "16-20 slides" a sentence in a prompt.
    """
    await _ingest(workspace)
    project = Project(workspace=workspace, slug="tarvis")
    write_brief(
        DeckBrief(language="English", audience="a conference oral", pages=PageBudget(16, 20)),
        brief_path(project),
    )
    stage, _views = _stage()
    result = await stage.run(project, DECK)  # three pages against a sixteen-page brief

    assert not result.ok
    budget = [f for f in result.findings if f.kind == "page_budget"]
    assert budget and budget[0].severity is Severity.BLOCKING
    assert budget[0].audience is Audience.AUTHOR
    assert "16-20" in budget[0].message
    assert not (project.exports_dir / "TarViS.pptx").exists()


@pytest.mark.asyncio
async def test_a_deck_in_the_wrong_language_stops_delivery(workspace: Path) -> None:
    await _ingest(workspace)
    project = Project(workspace=workspace, slug="tarvis")
    write_brief(
        DeckBrief(language="中文", audience="内部技术评审", pages=PageBudget(1, 40)),
        brief_path(project),
    )
    stage, _views = _stage()
    result = await stage.run(project, DECK)

    assert not result.ok
    wrong = [f for f in result.findings if f.kind == "language"]
    assert wrong and "中文" in wrong[0].message
