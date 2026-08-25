"""Bringing something from the web into a deck's materials.

The tests that matter are about the bytes. A URL ending in `.png` that returns
HTML, and a `Content-Type: image/png` on something that is not an image, are both
ordinary on the open web -- and writing either into the materials directory would
fail later somewhere that reads like a bug in the ingest.
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import httpx
import pytest

from raven.ppt.contracts import Project
from raven.ppt.tools.fetch import MAX_BYTES, PptFetchTool

pytest.importorskip("PIL")


def _png(width: int = 8, height: int = 8) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, format="PNG")
    return buffer.getvalue()


def _pptx() -> bytes:
    """A real presentation, because a downloaded template is now prepared as it
    lands: the two ways one arrives -- by URL and by path -- have to end in the
    same state, and a zip that only looks like a deck no longer reaches it."""
    from pptx import Presentation

    buffer = io.BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(buffer)
    return buffer.getvalue()


def _not_a_deck() -> bytes:
    """A zip shaped like a .pptx that python-pptx will not open."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("ppt/presentation.xml", "<presentation/>")
    return buffer.getvalue()


@pytest.fixture()
def fetch(monkeypatch, tmp_path: Path):
    """A tool whose HTTP layer is a MockTransport.

    Patched at httpx rather than at the tool, so the code under test is the code
    that ships -- including the streaming size guard, which a fake download method
    would have skipped.
    """

    def make(body: bytes, *, headers: dict | None = None, status: int = 200):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(status, content=body, headers=headers or {})

        transport = httpx.MockTransport(handler)
        real = httpx.AsyncClient

        def factory(*args, **kwargs):
            kwargs.pop("proxy", None)
            return real(*args, transport=transport, **kwargs)

        monkeypatch.setattr(httpx, "AsyncClient", factory)
        monkeypatch.setattr("raven.security.network.validate_url_target", lambda url: None)
        return PptFetchTool(tmp_path)

    return make


async def _run(tool: PptFetchTool, **kw) -> dict:
    args = {"project": "tarvis", "url": "https://example.com/thing"}
    return json.loads(await tool.execute(**{**args, **kw}))


@pytest.mark.asyncio
async def test_an_image_lands_in_the_decks_own_sources(fetch, tmp_path: Path) -> None:
    """One place, always. A fetch that could land elsewhere started a second pile, and
    the ingest that followed replaced the index rather than adding to it."""
    body = await _run(fetch(_png()), url="https://example.com/figure")

    assert body["ok"] is True and body["kind"] == "image"
    landed = Path(body["path"])
    assert landed.parent == Project(workspace=tmp_path, slug="tarvis").sources_dir
    assert landed.suffix == ".png" and landed.read_bytes()


@pytest.mark.asyncio
async def test_the_url_is_recorded_beside_what_it_fetched(fetch, tmp_path: Path) -> None:
    """The attribution `documents.source_urls` exists to provide and never had:
    nothing wrote the manifest it reads, so every fetched figure reached a slide with
    no source at all."""
    from raven.ppt.services.ingest import fetched

    await _run(fetch(_png()), url="https://example.com/figure")

    held = fetched(Project(workspace=tmp_path, slug="tarvis"))
    assert [source.url for source in held] == ["https://example.com/figure"]


@pytest.mark.asyncio
async def test_the_suffix_comes_from_the_bytes_not_the_url(fetch, tmp_path: Path) -> None:
    """A URL ending in .png that returns HTML is ordinary on the open web."""
    body = await _run(fetch(b"<html><body>not an image</body></html>"), url="https://example.com/photo.png")

    assert body["ok"] is True and body["kind"] == "document"
    assert Path(body["path"]).suffix == ".html"


@pytest.mark.asyncio
async def test_a_lying_content_type_does_not_get_a_pass(fetch) -> None:
    body = await _run(fetch(b"\x00\x01\x02 not an image", headers={"content-type": "image/png"}))
    assert body["ok"] is False
    assert "not a PDF, a PowerPoint file, an image, or text" in body["error"]


@pytest.mark.asyncio
async def test_a_pdf_is_recognised_by_its_header(fetch) -> None:
    body = await _run(fetch(b"%PDF-1.7\n trailer"))
    assert body["kind"] == "document" and Path(body["path"]).suffix == ".pdf"


@pytest.mark.asyncio
async def test_a_pptx_lands_in_the_template_slot_not_the_materials(fetch, tmp_path: Path) -> None:
    """A deck is not a source to quote; it is a house style to work inside."""
    body = await _run(fetch(_pptx()), url="https://example.com/house-style.pptx")

    assert body["ok"] is True and body["kind"] == "template"
    project = Project(workspace=tmp_path, slug="tarvis")
    assert Path(body["path"]).parent == project.root / "template"
    assert "house style" in body["next_step"]


@pytest.mark.asyncio
async def test_a_downloaded_template_is_prepared_where_it_lands(fetch, tmp_path: Path) -> None:
    """The two ways a template arrives have to end in the same state. Downloading
    one and leaving it in the slot was the earlier behaviour, and it meant a
    template that came by URL was never prepared, so nothing downstream saw one."""
    from raven.ppt.services.template import bound

    body = await _run(fetch(_pptx()), url="https://example.com/house-style.pptx")

    assert body["ok"] is True
    assert bound(Project(workspace=tmp_path, slug="tarvis")) is not None
    assert body["template"] == "house-style.pptx"


@pytest.mark.asyncio
async def test_a_pptx_that_will_not_open_is_refused_rather_than_stored(fetch, tmp_path: Path) -> None:
    body = await _run(fetch(_not_a_deck()))

    assert body["ok"] is False
    assert "does not open as a presentation" in body["error"]
    assert not list((Project(workspace=tmp_path, slug="tarvis").root / "template").glob("*.pptx"))


@pytest.mark.asyncio
async def test_a_zip_that_is_not_a_presentation_is_refused(fetch) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hello")
    body = await _run(fetch(buffer.getvalue()))
    assert body["ok"] is False and "not a PowerPoint file" in body["error"]


@pytest.mark.asyncio
async def test_a_declared_size_over_the_limit_is_refused_before_the_body_arrives(fetch) -> None:
    body = await _run(fetch(b"%PDF-1.7", headers={"content-length": str(MAX_BYTES + 1)}))
    assert body["ok"] is False and "32MB" in body["error"]


@pytest.mark.asyncio
async def test_an_undeclared_oversized_body_is_refused_as_it_arrives(fetch) -> None:
    """A server may not declare a length, so the bytes are counted too."""
    body = await _run(fetch(b"%PDF-" + b"x" * (MAX_BYTES + 1)))
    assert body["ok"] is False and "32MB" in body["error"]


@pytest.mark.asyncio
async def test_a_failed_request_says_so_rather_than_writing_nothing_quietly(fetch) -> None:
    body = await _run(fetch(b"nope", status=404))
    assert body["ok"] is False and "download failed" in body["error"]


@pytest.mark.asyncio
async def test_there_is_no_materials_path_to_point_anywhere(fetch) -> None:
    """The parameter is gone, not defaulted: it existed to let a fetch land somewhere
    other than the deck's sources, and there was never a good reason to."""
    from raven.ppt.tools.fetch import PptFetchTool

    assert "materials_dir" not in PptFetchTool(Path("/tmp")).parameters["properties"]


@pytest.mark.asyncio
async def test_a_bad_project_name_is_refused_before_the_download(fetch) -> None:
    body = await _run(fetch(_png()), project="../etc")
    assert body["ok"] is False and "usable project name" in body["error"]


@pytest.mark.asyncio
async def test_a_hostile_filename_is_reduced_to_something_safe(fetch, tmp_path: Path) -> None:
    body = await _run(fetch(_png()), filename="../../etc/passwd")
    assert Path(body["path"]).parent == Project(workspace=tmp_path, slug="tarvis").sources_dir
    assert Path(body["path"]).name == "passwd.png"


@pytest.mark.asyncio
async def test_a_body_of_control_bytes_is_binary_however_well_it_decodes(fetch) -> None:
    """UTF-8 accepts control characters, so decoding is not a test for text.

    Without this, a truncated video saved as `.md` becomes the ingest's problem
    to explain.
    """
    body = await _run(fetch(bytes(range(1, 32)) * 40))
    assert body["ok"] is False and "not a PDF" in body["error"]


@pytest.mark.asyncio
async def test_an_empty_body_is_not_a_document(fetch) -> None:
    body = await _run(fetch(b"   \n\n  "))
    assert body["ok"] is False


@pytest.mark.asyncio
async def test_ordinary_prose_with_punctuation_and_cjk_is_still_text(fetch) -> None:
    body = await _run(fetch("# 标题\n\n正文 — with an em dash, 90% coverage.\n".encode()))
    assert body["ok"] is True and Path(body["path"]).suffix == ".md"


SVG = (
    b'<svg xmlns="http://www.w3.org/2000/svg" width="240" height="60" viewBox="0 0 240 60">'
    b'<rect width="240" height="60" fill="#155FFD"/><text x="12" y="40" fill="#fff">mem0</text></svg>'
)


def test_a_brand_mark_arrives_as_svg_and_is_kept_as_a_picture() -> None:
    """A logo saved as `.md` is a competitor analysis with no competitor's mark on it.

    Rasterising needs cairo's native library, which `cairosvg` cannot supply itself, so
    this asserts the conversion only where the conversion is possible. The degraded path
    is a case of its own below rather than a looser assertion here.
    """
    try:
        # Not `importorskip`: it catches ImportError only, and a missing libcairo
        # surfaces as OSError from cairocffi -- which would error this test rather
        # than skip it, the same trap the guard under test exists for.
        import cairosvg  # noqa: F401
    except (ImportError, OSError) as exc:
        pytest.skip(f"rasterising an SVG needs the native cairo library: {exc}")
    from raven.ppt.tools.fetch import _sniff

    payload, suffix, kind = _sniff(SVG)

    assert (suffix, kind) == (".png", "image")
    assert payload.startswith(b"\x89PNG"), "python-pptx places rasters only"


def test_an_svg_that_will_not_draw_is_still_the_text_it_is() -> None:
    from raven.ppt.tools.fetch import _sniff

    payload, suffix, kind = _sniff(b"<svg this is not really an svg at all")

    assert kind == "document"
    assert payload == b"<svg this is not really an svg at all"


def test_the_other_three_kinds_come_back_byte_for_byte() -> None:
    from raven.ppt.tools.fetch import _sniff

    for blob, expected in (
        (b"<html><body>hi</body></html>", ".html"),
        (b"# a markdown note", ".md"),
        (b"%PDF-1.4 trailer", ".pdf"),
    ):
        payload, suffix, _ = _sniff(blob)
        assert suffix == expected
        assert payload == blob


def test_a_machine_without_libcairo_keeps_the_svg_instead_of_failing(monkeypatch) -> None:
    """`import cairosvg` pulls in cairocffi, which raises OSError -- not ImportError --
    when the native library is absent, so the guard has to catch that too."""
    import builtins
    import sys

    from raven.ppt.tools.fetch import _sniff

    real_import = builtins.__import__

    def no_libcairo(name, *args, **kwargs):
        if name == "cairosvg":
            raise OSError("no library called 'cairo-2' was found")
        return real_import(name, *args, **kwargs)

    monkeypatch.delitem(sys.modules, "cairosvg", raising=False)
    monkeypatch.setattr(builtins, "__import__", no_libcairo)

    payload, suffix, kind = _sniff(SVG)

    assert kind == "document", "an SVG that cannot be drawn is still the text it is"
    assert payload == SVG
