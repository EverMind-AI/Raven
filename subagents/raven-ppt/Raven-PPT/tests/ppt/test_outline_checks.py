"""What the outline stage measures about the plan itself.

Both checks here were written against a real twelve-page plan and only after it was
read: three of its pages were section dividers announcing the claim of the page
behind them, and none of its twelve pages planned to show a figure while the
materials held thirteen.
"""

from __future__ import annotations

from raven.ppt.contracts.outline import Outline, PagePlan
from raven.ppt.tools.outline import (
    PLANNED_EVIDENCE,
    SAME_PROTOTYPE_IS_STRUCTURAL,
    STRUCTURAL_SHARE,
    _structural,
    _unplanned_figures,
)


class _State:
    def __init__(self, figures: int) -> None:
        self.figures = tuple(range(figures))


def _page(number: int, *, says: int = 3, prototype: int | None = None, figures: tuple[str, ...] = ()) -> PagePlan:
    return PagePlan(
        page=number,
        claim=f"claim {number}",
        says=tuple(f"point {index}" for index in range(says)),
        prototype=prototype,
        figures=figures,
    )


def test_the_shares_are_the_measured_ones() -> None:
    assert STRUCTURAL_SHARE == 0.25
    assert SAME_PROTOTYPE_IS_STRUCTURAL == 3
    assert PLANNED_EVIDENCE == 0.5


def test_three_dividers_in_twelve_pages_are_reported() -> None:
    """The live case: three pages adapting the same template page, one line each."""
    pages = tuple(
        _page(number, says=1 if number in (3, 5, 8) else 3, prototype=3 if number in (3, 5, 8) else number)
        for number in range(1, 13)
    )
    findings = _structural(Outline(takeaway="t", pages=pages))

    assert [f.kind for f in findings] == ["structural_pages"]
    assert findings[0].detail == {"pages": [3, 5, 8], "of": 12}


def test_one_divider_in_twelve_pages_is_not_worth_saying() -> None:
    pages = tuple(_page(number, says=1 if number == 5 else 3, prototype=number) for number in range(1, 13))

    assert _structural(Outline(takeaway="t", pages=pages)) == []


def test_a_plan_that_shows_nothing_while_the_materials_hold_figures() -> None:
    pages = tuple(_page(number) for number in range(1, 13))
    findings = _unplanned_figures(Outline(takeaway="t", pages=pages), _State(13))

    assert [f.kind for f in findings] == ["unplanned_figures"]
    assert findings[0].detail == {"pages_with_figures": 0, "pages": 12, "figures_held": 13}


def test_half_the_pages_showing_something_is_enough() -> None:
    pages = tuple(_page(number, figures=("fig_one",) if number % 2 else ()) for number in range(1, 13))

    assert _unplanned_figures(Outline(takeaway="t", pages=pages), _State(13)) == []


def test_materials_with_no_figures_are_not_a_gap() -> None:
    """A text-only source cannot be asked for pictures it does not have."""
    pages = tuple(_page(number) for number in range(1, 13))

    assert _unplanned_figures(Outline(takeaway="t", pages=pages), _State(0)) == []


class _Template:
    def __init__(self, source) -> None:
        self.source = source


class _Bound:
    def __init__(self, source) -> None:
        self.template = _Template(source)
        self.figures = ()


def _house_template(path):
    """A template that says what its own pages are for: cover, agenda, closing."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for heading in ("A deck about something", "Agenda", "A content page", "谢谢观看"):
        page = presentation.slides.add_slide(presentation.slide_layouts[6])
        page.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = heading
    presentation.save(str(path))
    return path


def test_the_cover_and_the_closing_page_have_to_be_the_templates(tmp_path) -> None:
    """The one requirement, not advice: a deck that draws its own cover reads as not
    the user's before a word of it is read."""
    from raven.ppt.tools.outline import _house_pages

    template = _house_template(tmp_path / "house.pptx")
    pages = tuple(_page(number, prototype=None) for number in range(1, 13))
    findings = _house_pages(Outline(takeaway="t", pages=pages), _Bound(template))

    roles = {f.detail["role"] for f in findings}
    assert roles == {"cover", "agenda", "closing"}
    assert all(f.severity.value == "blocking" for f in findings)


def test_a_plan_that_uses_them_passes(tmp_path) -> None:
    from raven.ppt.tools.outline import _house_pages

    template = _house_template(tmp_path / "house.pptx")
    pages = (
        _page(1, prototype=1),
        _page(2, prototype=2),
        *(_page(number, prototype=3) for number in range(3, 12)),
        _page(12, prototype=4),
    )

    assert _house_pages(Outline(takeaway="t", pages=pages), _Bound(template)) == []


def test_a_short_deck_needs_no_index(tmp_path) -> None:
    """Five pages do not need a table of contents."""
    from raven.ppt.tools.outline import _house_pages

    template = _house_template(tmp_path / "house.pptx")
    pages = (_page(1, prototype=1), _page(2), _page(3), _page(4), _page(5, prototype=4))
    findings = _house_pages(Outline(takeaway="t", pages=pages), _Bound(template))

    assert [f.detail["role"] for f in findings] == []


def test_no_template_means_no_requirement() -> None:
    from raven.ppt.tools.outline import _house_pages

    class NoTemplate:
        template = None
        figures = ()

    pages = tuple(_page(number) for number in range(1, 13))
    assert _house_pages(Outline(takeaway="t", pages=pages), NoTemplate()) == []
