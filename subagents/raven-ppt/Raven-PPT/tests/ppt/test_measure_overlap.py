"""Content behind something, which is the overlap no render check can see.

`word_collision` reads the render and finds a word painted over a word. These are
the pages where nothing is painted over anything and the content is still not there:
a figure under a card, a page whose background was drawn last.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from raven.ppt.services.measure.overlap import COVERED, overlap_findings

pytest.importorskip("pptx")
pytest.importorskip("PIL")


def _figure(path: Path) -> Path:
    from PIL import Image

    Image.new("RGB", (800, 600), "grey").save(path)
    return path


def _deck(tmp_path: Path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    return presentation, Inches


def _card(slide, inches, left: float, top: float, width: float, height: float):
    from pptx.enum.shapes import MSO_SHAPE

    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(left), inches(top), inches(width), inches(height))
    card.fill.solid()
    return card


def test_a_figure_under_two_cards_is_reported(tmp_path: Path) -> None:
    presentation, inches = _deck(tmp_path)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    figure = _figure(tmp_path / "fig.png")
    slide.shapes.add_picture(str(figure), inches(1), inches(1.5), width=inches(5), height=inches(3.75))
    _card(slide, inches, 1.0, 1.5, 2.5, 3.75)
    _card(slide, inches, 3.5, 1.5, 2.5, 3.75)
    built = tmp_path / "hidden.pptx"
    presentation.save(str(built))

    findings = overlap_findings(built)
    assert [f.kind for f in findings] == ["covered_shape"]
    assert findings[0].severity.value == "blocking"
    assert findings[0].detail["hidden"] >= COVERED
    assert findings[0].detail["under"] == [2, 3]


def test_the_same_cards_under_the_figure_are_a_design(tmp_path: Path) -> None:
    """Z-order is the whole measurement: a panel under a picture is how a page is built."""
    presentation, inches = _deck(tmp_path)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _card(slide, inches, 1.0, 1.5, 2.5, 3.75)
    _card(slide, inches, 3.5, 1.5, 2.5, 3.75)
    slide.shapes.add_picture(
        str(_figure(tmp_path / "fig.png")), inches(1), inches(1.5), width=inches(5), height=inches(3.75)
    )
    built = tmp_path / "layered.pptx"
    presentation.save(str(built))

    assert overlap_findings(built) == []


def test_copy_inside_a_card_is_not_covered(tmp_path: Path) -> None:
    """The commonest shape on any deck: a filled panel with text on top of it."""
    presentation, inches = _deck(tmp_path)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _card(slide, inches, 1.0, 1.5, 4.0, 2.0)
    box = slide.shapes.add_textbox(inches(1.2), inches(1.7), inches(3.6), inches(1.6))
    box.text_frame.text = "the card's own copy"
    built = tmp_path / "card.pptx"
    presentation.save(str(built))

    assert overlap_findings(built) == []


def test_a_background_drawn_last_hides_the_page(tmp_path: Path) -> None:
    """A ground cannot be hidden and can certainly hide: excluding big shapes from the
    check entirely made this page -- every word under a full-bleed rectangle -- clean."""
    presentation, inches = _deck(tmp_path)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(inches(1), inches(1), inches(6), inches(2))
    box.text_frame.text = "the copy this page is for"
    slide.shapes.add_picture(
        str(_figure(tmp_path / "fig.png")), inches(1), inches(3.5), width=inches(5), height=inches(3)
    )
    _card(slide, inches, 0.0, 0.0, 13.333, 7.5)
    built = tmp_path / "buried.pptx"
    presentation.save(str(built))

    findings = overlap_findings(built)
    assert {f.kind for f in findings} == {"covered_shape"}
    assert len(findings) == 2, "the text and the figure are both under it"


def test_a_mark_is_not_measured(tmp_path: Path) -> None:
    """A bullet, a rule, a step number's bubble: too small to be content, and a page
    is full of them sitting on top of each other by design."""
    presentation, inches = _deck(tmp_path)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    mark = slide.shapes.add_textbox(inches(1), inches(1), inches(0.2), inches(0.2))
    mark.text_frame.text = "1"
    _card(slide, inches, 0.9, 0.9, 0.4, 0.4)
    built = tmp_path / "marks.pptx"
    presentation.save(str(built))

    assert overlap_findings(built) == []
