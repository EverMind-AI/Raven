"""Every element this route writes, in an order PowerPoint will accept.

The one class of defect a pipeline that judges decks by rendering them cannot see.
ECMA-376 fixes the order of a great many elements' children; LibreOffice does not
check it and PowerPoint does, dropping what it cannot parse. So a deck can render
perfectly here, measure clean on every gate, and open in the room with its tables
restyled -- which is exactly what shipped: 175 cells across five tables in two
delivered decks, all carrying `lnB, lnT, lnR, lnL` where the schema says
`lnL, lnR, lnT, lnB`.

Scanned over a deck built only from these helpers, so a violation is this code's
and not some user template's.
"""

from __future__ import annotations

import textwrap
import zipfile
from collections import Counter
from pathlib import Path

import pytest

pytest.importorskip("pptx", reason="ppt extra not installed")

from lxml import etree

from raven.ppt.backends.script import run_script
from raven.ppt.backends.script.workspace import asset_helpers
from raven.ppt.contracts import Project

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

# The fill group: any one of these, in the one position the group occupies.
_FILLS = ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")

# The sequences, as ordered groups of alternatives. Only the elements these helpers
# actually write are listed -- a table of the whole schema would be a copy of the
# standard, and what this guards is our own output.
SEQUENCES: dict[str, list[tuple[str, ...]]] = {
    f"{_A}rPr": [
        ("ln",),
        _FILLS,
        ("effectLst", "effectDag"),
        ("highlight",),
        ("uLnTx", "uLn"),
        ("uFillTx", "uFill"),
        ("latin",),
        ("ea",),
        ("cs",),
        ("sym",),
        ("hlinkClick",),
        ("hlinkMouseOver",),
        ("rtl",),
        ("extLst",),
    ],
    f"{_A}tcPr": [
        ("lnL",),
        ("lnR",),
        ("lnT",),
        ("lnB",),
        ("lnTlToBr",),
        ("lnBlToTr",),
        ("cell3D",),
        _FILLS,
        ("headers",),
        ("extLst",),
    ],
    f"{_A}bodyPr": [
        ("prstTxWarp",),
        ("noAutofit", "normAutofit", "spAutoFit"),
        ("scene3d",),
        ("sp3d", "flatTx"),
        ("extLst",),
    ],
    f"{_P}spPr": [
        ("xfrm",),
        ("custGeom", "prstGeom"),
        _FILLS,
        ("ln",),
        ("effectLst", "effectDag"),
        ("scene3d",),
        ("sp3d",),
        ("extLst",),
    ],
    f"{_A}blipFill": [("blip",), ("srcRect",), ("tile", "stretch")],
    f"{_A}pPr": [
        ("lnSpc",),
        ("spcBef",),
        ("spcAft",),
        ("buClrTx", "buClr"),
        ("buSzTx", "buSzPct", "buSzPts"),
        ("buFontTx", "buFont"),
        ("buNone", "buAutoNum", "buChar"),
        ("tabLst",),
        ("defRPr",),
        ("extLst",),
    ],
}
_RANK = {
    tag: {name: index for index, group in enumerate(groups) for name in group} for tag, groups in SEQUENCES.items()
}


# One page per primitive, so a violation names the helper that wrote it.
EVERY_PRIMITIVE = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches
    from ppt_icons import add_icon
    from ppt_layout import Box, GUTTER, card, formula, page, plane, rule, table, write
    from ppt_theme import THEMES

    T = THEMES["ink-graphite"]
    F, C = T["font_family"], T["cjk_font_family"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # SLIDE 1
    one = prs.slides.add_slide(prs.slide_layouts[6])
    frame = page()
    write(one, frame.kicker, "方法", size=12, colour=T["accent"], font=F, cjk_font=C)
    write(one, frame.title, "架构：Backbone 到共享解码器", size=30, bold=True, colour=T["foreground"], font=F, cjk_font=C)
    rule(one, frame.title, T)
    plane(one, frame.body.rows(2)[1], T, radius=True)
    write(one, frame.footer, "来源：Table 1", size=12, colour=T["muted"], font=F, cjk_font=C)
    add_icon(one, "target", Inches(0.72), Inches(2.1), Inches(0.42), T["accent"])

    # SLIDE 2
    two = prs.slides.add_slide(prs.slide_layouts[6])
    body = page().body
    table(two, body.rows(2)[0], [["方法", "AP"], ["TarViS", "48.3"], ["VITA", "45.7"]], T, numeric_from=1)
    for box, name in zip(body.rows(2)[1].columns(3, gutter=GUTTER), ("target", "clock", "chart-bar")):
        card(two, box, T, icon=name, title="语义查询是必要的", body="去掉后全线下降")

    # SLIDE 3
    three = prs.slides.add_slide(prs.slide_layouts[6])
    formula(three, page().body.rows(3)[0], "掩码 logits = (F_4, Q'_{inst})；分类 = concat(Q'_{sem}, Q'_{bg})", T)

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


def _violations(path: Path) -> Counter:
    found: Counter = Counter()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            for element in etree.fromstring(archive.read(name)).iter():
                ranks = _RANK.get(element.tag)
                if ranks is None:
                    continue
                written = [ranks[child.tag.split("}")[-1]] for child in element if child.tag.split("}")[-1] in ranks]
                if written != sorted(written):
                    order = tuple(child.tag.split("}")[-1] for child in element)
                    found[(element.tag.split("}")[-1], order)] += 1
    return found


@pytest.mark.asyncio
async def test_nothing_this_route_writes_is_out_of_the_schema_s_order(tmp_path: Path) -> None:
    project = Project(workspace=tmp_path, slug="orders")
    outcome = await run_script(project, EVERY_PRIMITIVE, helpers=asset_helpers(), timeout_s=120.0)
    assert outcome.ok, outcome.stderr

    found = _violations(outcome.pptx_path)
    assert found == Counter(), "elements out of the schema's order: " + "; ".join(
        f"{count}x {tag} {list(order)}" for (tag, order), count in found.most_common()
    )


def test_the_check_would_catch_the_defect_that_shipped(tmp_path: Path) -> None:
    """The scan itself, against the XML two delivered decks actually carried."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(1.5)).table
    for row in table.rows:
        for cell in row.cells:
            properties = cell._tc.get_or_add_tcPr()
            for edge in ("lnL", "lnR", "lnT", "lnB"):
                element = properties.makeelement(f"{_A}{edge}", {})
                element.append(element.makeelement(f"{_A}noFill", {}))
                properties.insert(0, element)  # what shipped
    built = tmp_path / "reversed.pptx"
    presentation.save(str(built))

    found = _violations(built)
    assert sum(found.values()) == 4, found
    assert next(iter(found))[1][:4] == ("lnB", "lnT", "lnR", "lnL")
