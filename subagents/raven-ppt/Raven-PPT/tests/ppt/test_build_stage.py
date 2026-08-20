"""The order of build, measure, design, deliver -- and why it is that order."""

from __future__ import annotations

import textwrap
from pathlib import Path

import pytest

from raven.ppt.backends.script import script_path
from raven.ppt.contracts import Audience, BuildOutcome, Finding, PageSource, Project, Severity, StageResult
from raven.ppt.profiles import registry
from raven.ppt.stages.build import BuildStage

SCRIPT = textwrap.dedent(
    """
    import os
    from pptx import Presentation

    prs = Presentation()


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    # SLIDE 1
    one = new_slide()

    # SLIDE 2
    two = new_slide()

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.fixture()
def project(tmp_path: Path) -> Project:
    p = Project(workspace=tmp_path, slug="tarvis")
    p.build_dir.mkdir(parents=True)
    p.state_dir.mkdir(parents=True)
    script_path(p).write_text(SCRIPT, encoding="utf-8")
    return p


def _outcome(project: Project, pages: int = 2, mapped: bool = True) -> BuildOutcome:
    deck = project.build_dir / "deck.pptx"
    deck.write_bytes(b"PK deck")
    lines = SCRIPT.splitlines(keepends=True)
    first = next(i for i, line in enumerate(lines) if "# SLIDE 1" in line)
    second = next(i for i, line in enumerate(lines) if "# SLIDE 2" in line)
    sources = (
        (
            PageSource(page=1, first_line=first, last_line=second),
            PageSource(page=2, first_line=second, last_line=len(lines) - 2),
        )
        if mapped
        else ()
    )
    return BuildOutcome(ok=True, pptx_path=deck, pages=pages, sources=sources, source_digest="x")


class FakeDesignPass:
    def __init__(self, result: StageResult | None = None) -> None:
        self.result = result
        self.calls = 0

    async def run(self, project: Project, outcome: BuildOutcome) -> StageResult:
        self.calls += 1
        return self.result or StageResult(ok=True, data={"outcome": outcome, "rounds": []})


def _measure(findings):
    async def measure(_project: Project, _pptx: Path, _outcome=None) -> list[Finding]:
        return list(findings)

    return measure


def _stage(project: Project, *, findings=(), design=None, outcome=None, profile="script_author") -> BuildStage:
    measured = list(findings)

    async def backend(_project: Project, _script: str | None) -> BuildOutcome:
        return outcome or _outcome(_project)

    return BuildStage(
        backend=backend,
        measure=_measure(measured),
        profile=registry.get(profile),
        design_pass=design,
    )


def _fact() -> Finding:
    return Finding(kind="fact", severity=Severity.BLOCKING, message="48.3 is not in the sources", page=11)


def _warning(kind: str = "type_floor", audience: Audience = Audience.DESIGNER) -> Finding:
    return Finding(kind=kind, severity=Severity.WARNING, message="11.5pt body", page=3, audience=audience)


@pytest.mark.asyncio
async def test_a_clean_deck_is_measured_designed_and_delivered(project: Project) -> None:
    design = FakeDesignPass()
    result = await _stage(project, design=design).run(project)

    assert result.ok, result.note
    assert design.calls == 1
    assert Path(result.data["pptx_path"]).is_file()


@pytest.mark.asyncio
async def test_an_author_refusal_no_longer_holds_up_the_design_pass(project: Project) -> None:
    """The author's list and the designer's are different axes, not a queue.

    This gated on blocking findings twice and both were wrong. Any blocking finding at
    all suppressed the one stage whose job is layout exactly when the layout was wrong;
    then any blocking finding not addressed to the designer, which is the same mistake
    one step back -- a deck with 29 unreplaced placeholders (the author's to fix) also
    had 25 pairs of overlapping words (the pass's), and the second list waited on the
    first. Across three complete live runs the pass ran zero times.

    The deck still does not publish while the author's finding stands; that is what the
    refusal is for. It is the polishing that no longer waits on it.
    """
    design = FakeDesignPass()
    result = await _stage(project, findings=[_fact()], design=design).run(project)

    assert not result.ok, "a claim the materials do not make still refuses the export"
    assert design.calls == 1, "and the layout is still looked at"
    assert not (project.exports_dir / "deck.pptx").exists()


@pytest.mark.asyncio
async def test_a_layout_refusal_still_gets_the_design_pass(project: Project) -> None:
    """The stage that fixes layout has to run when layout is what was refused.

    Polishing was skipped on any blocking finding at all, and three live runs paid
    for it: two of them refused on a band gate that was itself mistaken -- eight
    kicker rules in one, five planes carrying copy in the other -- and in both the
    design pass never ran for a single build of the whole run. The one run with no
    blocking finding is the one that came back composed.
    """
    design = FakeDesignPass()
    collision = Finding(
        kind="word_collision",
        severity=Severity.BLOCKING,
        message="'31.1' is painted over 'VPS-only'",
        page=16,
        audience=Audience.DESIGNER,
    )
    result = await _stage(project, findings=[collision], design=design).run(project)

    assert design.calls == 1
    assert not result.ok  # and the export still refuses
    assert not (project.exports_dir / "deck.pptx").exists()


@pytest.mark.asyncio
async def test_a_warning_never_stops_the_deck(project: Project) -> None:
    """Shrinking the copy would satisfy the measurement and make the page worse."""
    result = await _stage(project, findings=[_warning()], design=FakeDesignPass()).run(project)

    assert result.ok
    assert [f.kind for f in result.findings] == ["type_floor"]


@pytest.mark.asyncio
async def test_the_deck_is_measured_again_after_the_design_pass(project: Project) -> None:
    """The pass rewrote the pages it had just looked at."""
    seen: list[int] = []

    async def backend(_project: Project, _script: str | None) -> BuildOutcome:
        return _outcome(_project)

    async def measure(_p: Project, _pptx: Path, _outcome=None) -> list[Finding]:
        seen.append(1)
        return []

    stage = BuildStage(
        backend=backend, measure=measure, profile=registry.get("script_author"), design_pass=FakeDesignPass()
    )
    await stage.run(project)
    assert len(seen) == 2


@pytest.mark.asyncio
async def test_a_design_pass_that_broke_the_build_reports_and_does_not_deliver(project: Project) -> None:
    broke = Finding(
        kind="design_pass_broke_the_build",
        severity=Severity.BLOCKING,
        message="the design pass left the program unable to run",
        audience=Audience.AUTHOR,
    )
    design = FakeDesignPass(StageResult(ok=False, findings=(broke,), data={"outcome": _outcome(project)}))
    result = await _stage(project, design=design).run(project)

    assert not result.ok
    assert any(f.kind == "design_pass_broke_the_build" for f in result.findings)
    assert "pptx_path" not in result.data


@pytest.mark.asyncio
async def test_pages_that_cannot_be_told_apart_are_blocking_and_carry_a_skeleton(project: Project) -> None:
    """Refused without one, an author rewrote the same shim three times and left."""
    looped = textwrap.dedent(
        """
        import os
        from pptx import Presentation
        prs = Presentation()
        for _ in range(5):
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(os.environ["PPT_OUTPUT"])
        """
    ).lstrip()
    script_path(project).write_text(looped, encoding="utf-8")
    design = FakeDesignPass()
    result = await _stage(project, outcome=_outcome(project, pages=5, mapped=False), design=design).run(project)

    assert not result.ok
    unmapped = [f for f in result.findings if f.kind == "unmapped_page"]
    assert unmapped and unmapped[0].audience is Audience.AUTHOR
    assert "# SLIDE 1" in unmapped[0].message
    # The pass still looks: a deck whose pages cannot be told apart is the author's to
    # fix, and the layout of the pages that did map is worth a look either way.
    assert design.calls == 1


@pytest.mark.asyncio
async def test_a_mapped_deck_produces_no_mapping_finding(project: Project) -> None:
    result = await _stage(project).run(project)
    assert not any(f.kind == "unmapped_page" for f in result.findings)


@pytest.mark.asyncio
async def test_a_build_that_failed_stops_before_anything_else(project: Project) -> None:
    design = FakeDesignPass()
    result = await _stage(project, outcome=BuildOutcome(ok=False, stderr="boom"), design=design).run(project)

    assert not result.ok
    assert design.calls == 0
    assert result.data["outcome"].stderr == "boom"


@pytest.mark.asyncio
async def test_the_route_decides_which_warning_kinds_refuse_delivery(project: Project) -> None:
    """A filled colour bar is fatal where prose never stopped it coming back."""
    band = Finding(kind="band", severity=Severity.WARNING, message="a filled bar across page 4", page=4)
    result = await _stage(project, findings=[band], design=FakeDesignPass()).run(project)
    assert not result.ok

    other = Finding(kind="wrapped_label", severity=Severity.WARNING, message="a label wrapped", page=4)
    assert (await _stage(project, findings=[other], design=FakeDesignPass()).run(project)).ok


@pytest.mark.asyncio
async def test_polish_false_skips_the_pass_that_is_configured(project: Project) -> None:
    design = FakeDesignPass()
    result = await _stage(project, design=design).run(project, polish=False)
    assert result.ok and design.calls == 0


@pytest.mark.asyncio
async def test_the_design_pass_can_be_absent_and_says_so(project: Project) -> None:
    """Absent is a real configuration, and silence about it is not.

    Three live runs never reached the pass and no reply mentioned it, so the absence
    read as "the layout was fine".
    """
    result = await _stage(project, design=None).run(project)
    assert result.ok
    assert result.data["design_pass"] == {
        "skipped": "the design pass is off in this install -- tools.ppt.designer.enabled turns it on"
    }


@pytest.mark.asyncio
async def test_a_draft_is_not_polished_and_the_reply_says_why(project: Project) -> None:
    """A draft goes back to the author with every measurement; the pass sees finished
    decks. Asking the pass -- which may only rearrange -- to fix what a half-written
    program got wrong asks the weaker tool to do the stronger one's work."""
    design = FakeDesignPass()
    result = await _stage(project, design=design).run(project, draft=True)

    assert result.ok
    assert design.calls == 0
    assert "yours to answer" in result.data["design_pass"]["skipped"]
