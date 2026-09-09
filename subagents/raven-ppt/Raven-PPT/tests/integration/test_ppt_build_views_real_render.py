"""How many renders one `ppt_build` reply really carries, through the real renderer.

Everything under `tests/ppt` builds the tool with a fake `views`, so a batch is a
list the fixture wrote down rather than pictures a model can look at. Two constants
carried this number once and drifted to 1 and 3 without a test noticing, and the
unit tests could not notice: they read the same list the fake produced. So this one
converts a real deck with LibreOffice and counts the image blocks in the reply.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

pytest.importorskip("pptx", reason="ppt extra not installed")

from raven.ppt.contracts import (
    DeckBrief,
    IntakePlan,
    Outline,
    PageBudget,
    PagePlan,
    Project,
    brief_path,
    intake_path,
    outline_path,
    write_brief,
    write_outline,
    write_plan,
)
from raven.ppt.stages.build import BATCH_VIEWS
from raven.ppt.tools.assembly import build_ppt_tools

PAGES = 10

# A deck with nothing on its pages but their own numbers: the point is how many come
# back, and a page has to carry something for the render to be a page at all.
SCRIPT = """
import os
from pptx import Presentation
from pptx.util import Inches, Pt

presentation = Presentation()
presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
for number in range(1, %d + 1):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(10), Inches(1.5)).text_frame.paragraphs[0]
    paragraph.text = "PAGE %%d" %% number
    paragraph.runs[0].font.size = Pt(80)
presentation.save(os.environ["PPT_OUTPUT"])
""" % PAGES


def _project(root: Path) -> Project:
    project = Project(workspace=root, slug="views")
    write_brief(DeckBrief(language="English", audience="a review", pages=PageBudget(1, 40)), brief_path(project))
    write_plan(IntakePlan(topic="a deck", digest="x"), intake_path(project))
    write_outline(
        Outline(
            takeaway="it works",
            pages=tuple(PagePlan(page=number, claim=f"page {number}") for number in range(1, PAGES + 1)),
        ),
        outline_path(project),
    )
    return project


async def _one_build(root: Path, **kwargs) -> tuple[int, dict, int]:
    """(images in the reply, the reply body, the tool's own `slides` cap)."""
    _project(root)
    build = next(tool for tool in build_ppt_tools(root, **kwargs) if tool.name == "ppt_build")
    result = await build.execute(project="views", script=SCRIPT, draft=True)
    assert not isinstance(result, str), f"the build failed: {result}"
    body = json.loads(result.model_text)
    # Not `ok`: a draft build reports what would block publication, and this deck's
    # pages carry nothing but their numbers. What is asserted is that the deck was
    # really built and that the reply really carried the renders.
    assert body["slides"] == PAGES, body
    images = [block for block in (result.blocks or []) if block.get("type") == "image_url"]
    return len(images), body, build.parameters["properties"]["slides"]["maxItems"]


@pytest.mark.asyncio
async def test_one_build_carries_the_default_batch_of_renders(tmp_path: Path) -> None:
    images, body, cap = await _one_build(tmp_path)

    assert images == BATCH_VIEWS, "the reply carries as many pictures as the budget allows"
    assert cap == BATCH_VIEWS, "and the `slides` cap is the same number, not a second one"
    assert body["pages_shown"] == ", ".join(str(number) for number in range(1, BATCH_VIEWS + 1))
    assert body["pages_not_shown"] == PAGES - BATCH_VIEWS


@pytest.mark.asyncio
async def test_the_configured_budget_is_what_the_reply_carries(tmp_path: Path) -> None:
    """`tools.ppt.viewsPerCall` reaches the pictures, not just the schema."""
    wanted = 5
    images, body, cap = await _one_build(tmp_path, views_per_call=wanted)

    assert images == wanted and cap == wanted
    assert body["pages_not_shown"] == PAGES - wanted


@pytest.mark.asyncio
async def test_a_named_page_comes_back_alone(tmp_path: Path) -> None:
    """The other direction: naming fewer pages than the budget renders only those.
    `slides` used to be filled in with `[page_from]` when it was empty, which made the
    walk unreachable -- so the two cases have to be asserted against each other."""
    _project(tmp_path)
    build = next(tool for tool in build_ppt_tools(tmp_path) if tool.name == "ppt_build")

    result = await build.execute(project="views", script=SCRIPT, slides=[7], draft=True)
    body = json.loads(result.model_text)
    images = [block for block in (result.blocks or []) if block.get("type") == "image_url"]

    assert len(images) == 1
    assert body["pages_shown"] == "7"
