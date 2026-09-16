"""The formula picture the hand-built deck lane places: true size, display-style fractions, the refusals."""

from __future__ import annotations

import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from raven_ppt.services.assets import formulas

# The first render imports matplotlib and builds its font cache, which is the whole
# of the wall clock here.
pytestmark = pytest.mark.slow

EWC = r"L(\theta)=L_{\mathrm{new}}(\theta)+\frac{\lambda}{2}\sum_i F_i\,(\theta_i-\theta^{*}_{i})^{2}"


def _slide():
    deck = Presentation()
    return deck.slides.add_slide(deck.slide_layouts[6])


def _ink_rows(png: bytes) -> int:
    """How many pixel rows carry ink: the glyph itself, not the line box around it."""
    with Image.open(io.BytesIO(png)) as image:
        alpha = image.getchannel("A")
        rows = [y for y in range(alpha.height) if any(alpha.getpixel((x, y)) > 64 for x in range(alpha.width))]
    return max(rows) - min(rows) + 1


def test_the_picture_is_true_size_and_grows_with_the_point_size() -> None:
    small = formulas.formula_png("L", size_pt=18)
    large = formulas.formula_png("L", size_pt=36)
    with Image.open(io.BytesIO(small.png)) as image:
        assert (image.width / formulas.DPI, image.height / formulas.DPI) == (small.width_in, small.height_in)
        assert image.mode == "RGBA"
    # A capital stands about 0.72 of its point size: 18pt at 300 dpi is 54 rows of ink.
    assert 48 <= _ink_rows(small.png) <= 60
    assert 1.85 < _ink_rows(large.png) / _ink_rows(small.png) < 2.15, "twice the points is twice the ink"


def test_a_serif_formula_is_sized_up_to_stand_level_with_the_body() -> None:
    sans = formulas.formula_png("L", size_pt=18, serif=False)
    serif = formulas.formula_png("L", size_pt=18, serif=True)
    assert serif.size_pt == pytest.approx(18 * formulas.SERIF_SCALE)
    assert 0.95 < _ink_rows(serif.png) / _ink_rows(sans.png) < 1.15, "STIX at 1.2x stands as tall as DejaVu at 1x"


def test_a_fraction_is_set_display_style_and_a_display_fraction_is_left_alone() -> None:
    assert formulas.display_style(r"\frac{\lambda}{2}") == r"\dfrac{\lambda}{2}"
    assert formulas.display_style(r"\dfrac{a}{b}+\tfrac{c}{d}") == r"\dfrac{a}{b}+\tfrac{c}{d}"
    inline = formulas.formula_png(r"\dfrac{\lambda}{2}", size_pt=18)
    promoted = formulas.formula_png(r"\frac{\lambda}{2}", size_pt=18)
    assert promoted.height_in == pytest.approx(inline.height_in), "the promoted fraction is the display one"


def test_cjk_and_unknown_constructs_are_refused_with_the_way_round() -> None:
    with pytest.raises(ValueError, match="outside"):
        formulas.formula_png(r"\frac{收入}{成本}")
    with pytest.raises(ValueError, match="did not typeset"):
        formulas.formula_png(r"\begin{align}a&=b\end{align}")


def test_add_formula_places_the_picture_at_true_size_and_names_it() -> None:
    slide = _slide()
    rendered = formulas.formula_png(EWC, size_pt=18, colour="#1F2A44")
    shape = formulas.add_formula(slide, EWC, 0.8, 2.6, size_pt=18, colour="1F2A44")
    assert shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == "formula"
    assert shape.left == Inches(0.8) and shape.top == Inches(2.6)
    assert shape.width == Inches(rendered.width_in)
    assert 1.0 < rendered.width_in < 5.0, "one line of mathematics at body size"


def test_add_formula_scales_a_long_expression_down_to_its_column() -> None:
    slide = _slide()
    shape = formulas.add_formula(slide, EWC, 0.8, 2.6, size_pt=18, max_width_in=2.0)
    assert shape.width == Inches(2.0)
    assert shape.height < Inches(0.4), "height follows the width, the aspect holds"


def test_math_runs_sets_real_scripts_and_italic_variables_in_prose() -> None:
    slide = _slide()
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    runs = formulas.math_runs(
        box.text_frame.paragraphs[0], "旧后验 p(θ | D_A) 由 F_{new} 和 θ* 决定，softmax 不变", size_pt=16
    )
    text = "".join(run.text for run in runs)
    assert text == "旧后验 p(θ | DA) 由 Fnew 和 θ* 决定，softmax 不变", "the markers are consumed, the letters stay"
    baseline = {run.text: run._r.rPr.get("baseline") for run in runs if run._r.rPr is not None}
    assert baseline.get("A") == formulas.SUBSCRIPT_BASELINE
    assert baseline.get("new") == formulas.SUBSCRIPT_BASELINE
    assert baseline.get("*") == formulas.SUPERSCRIPT_BASELINE
    italic = {run.text: run.font.italic for run in runs}
    assert italic["p"] is True and italic["F"] is True and italic["A"] is True
    assert italic["softmax"] is False
    assert not any(run.font.italic for run in runs if "θ" in run.text), "Greek stays upright with the prose"
    assert all(run.font.size == Pt(16) for run in runs), "scripts are set at the line's own size"


def test_math_runs_leaves_a_plain_sentence_as_one_line_of_text() -> None:
    slide = _slide()
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    runs = formulas.math_runs(box.text_frame.paragraphs[0], "每个任务存一份对角 F 与 θ*：共 2×N 个浮点数。")
    assert "".join(run.text for run in runs) == "每个任务存一份对角 F 与 θ*：共 2×N 个浮点数。"
    assert not any((run._r.rPr is not None and run._r.rPr.get("baseline")) for run in runs if run.text != "*")


def test_math_runs_keeps_english_words_upright_and_stars_only_after_letters() -> None:
    slide = _slide()
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    runs = formulas.math_runs(box.text_frame.paragraphs[0], "a posterior p(x | D_A) is what I keep, and (a+b)*c too")
    texts = [run.text for run in runs]
    assert (
        runs[texts.index("p")].font.italic and runs[texts.index("x")].font.italic and runs[texts.index("A")].font.italic
    )
    assert not runs[0].font.italic and not runs[texts.index("I")].font.italic, "English words stay upright"
    assert runs[texts.index("b")].font.italic, "a letter beside an operator is a variable"
    star = [run for run in runs if "*" in run.text]
    assert star and all(run._r.rPr is None or not run._r.rPr.get("baseline") for run in star), (
        "an operator star stays on the line"
    )
