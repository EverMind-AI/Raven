"""Working inside a user's template: preparing it, reading it, reusing its pages.

The measurements behind these tests were taken on the 193 templates the user
supplied, and two of them decide the design:

85% of a template's visual elements live on its example slides rather than on its
layouts, so those slides are the reference and `add_slide(layout)` is not. And two
thirds of those pages hold something python-pptx cannot write -- custom geometry,
a gradient, a fill at 60% opacity -- so cloning a page is not a convenience beside
the decompiler, it is the only way that part of a template reaches a deck at all.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

import pytest

pytest.importorskip("pptx")

from raven_ppt.services.template import (  # noqa: E402
    clone_page,
    decompile,
    drop_shape,
    inspect_template,
    prepare,
    replace_picture,
    replace_text,
    strip_hidden,
)
from tests._ppt_engine_fixtures import deck, image, noise_image, noise_png, product_page, template_file  # noqa: F401


@pytest.fixture
def template(template_file) -> Path:
    return template_file()


# --- preparing ------------------------------------------------------------


def test_prepare_empties_the_example_slides_and_keeps_the_design(template: Path, tmp_path: Path):
    """The whole point: the master, theme, layouts and canvas come across; the
    stranger's content does not. Every template measured ships example slides, so
    an author that opened the original would produce them plus its own pages."""
    from pptx import Presentation

    before = inspect_template(template)
    prepared = prepare(template, tmp_path / "prepared.pptx")

    assert prepared is not None
    assert prepared.removed_slides == 2
    assert len(Presentation(str(prepared.path)).slides) == 0
    assert prepared.inventory.width_in == pytest.approx(13.333, abs=0.01)
    assert [layout.name for layout in prepared.inventory.layouts] == [layout.name for layout in before.layouts]


def test_prepare_leaves_the_users_file_alone(template: Path, tmp_path: Path):
    """It is their file, and a later run wants to prepare it again from clean."""
    from pptx import Presentation

    original = template.read_bytes()
    prepare(template, tmp_path / "prepared.pptx")

    assert template.read_bytes() == original
    assert len(Presentation(str(template)).slides) == 2


def test_prepare_declines_what_is_not_a_template(tmp_path: Path):
    """No template is a normal state -- most decks have none -- so this returns
    None rather than raising, and the destination is not left half-written."""
    (tmp_path / "notes.txt").write_text("this is not a deck")

    assert prepare(tmp_path / "notes.txt", tmp_path / "out.pptx") is None
    assert prepare(tmp_path / "missing.pptx", tmp_path / "out.pptx") is None
    assert not (tmp_path / "out.pptx").exists()


def test_strip_hidden_removes_the_pages_a_render_will_not_have(template: Path):
    """LibreOffice does not export a hidden slide, so it is in the file and not in
    the PDF. Every reader downstream numbers pages off the file, and one live run
    asked for thirteen pages of an eleven-page render, lost all eleven to the one
    out-of-range number, and told the author renders were unavailable on the host.
    119 of the 193 templates measured ship them."""
    from pptx import Presentation

    presentation = Presentation(str(template))
    presentation.slides[0]._element.set("show", "0")
    presentation.save(str(template))

    assert strip_hidden(template) == 1

    after = Presentation(str(template))
    assert len(after.slides) == 1
    assert "Section title" not in _texts(after.slides[0]), "the hidden page goes, not the last one"


def test_strip_hidden_leaves_a_template_without_any_alone(template: Path):
    """The common case is still a template with none, and it is not worth a rewrite."""
    before = template.read_bytes()

    assert strip_hidden(template) == 0
    assert template.read_bytes() == before


# --- reading it -----------------------------------------------------------


def test_inventory_describes_what_an_author_cannot_see(template: Path):
    inventory = inspect_template(template)

    assert inventory is not None
    assert inventory.example_slides == 2
    assert inventory.layouts
    assert inventory.theme_colours
    assert "Layouts:" in inventory.brief()


# --- decompiling ----------------------------------------------------------


def _run(source: str, prepared: Path, workdir: Path):
    """Execute the emitted reference the way an author would, and return the page.

    Against the prepared template rather than a blank presentation, because the
    reference opens with the `add_slide(prs.slide_layouts[n])` that most of the
    page's design actually arrives on -- an index that only means anything in the
    template it was read out of.
    """
    import os

    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    env = {
        "prs": Presentation(str(prepared)),
        "Inches": Inches,
        "Pt": Pt,
        "RGBColor": RGBColor,
        "MSO_SHAPE": MSO_SHAPE,
        "MSO_CONNECTOR": MSO_CONNECTOR,
        "MSO_ANCHOR": MSO_ANCHOR,
        "PP_ALIGN": PP_ALIGN,
    }
    here = os.getcwd()
    os.chdir(workdir)
    try:
        exec(source, env)  # noqa: S102 -- executing the reference is the assertion
    finally:
        os.chdir(here)
    return env["slide"]


def test_the_reference_runs(template: Path, tmp_path: Path):
    """The one assertion that matters. Four bugs in this decompiler -- a group's
    scale, a paragraph fetched before `frame.text` rebuilt the list, a colour read
    without the master's map, and a page whose whole design was on its layout --
    were all found by replaying the output and none by reading it."""
    prepared = prepare(template, tmp_path / "prepared.pptx")
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert prepared is not None and source is not None
    slide = _run(source.source, prepared.path, tmp_path / "img")
    assert len(slide.shapes) >= 4


def test_the_reference_opens_on_the_layout_the_page_was_drawn_on(template: Path, tmp_path: Path):
    """Where a template's cover keeps its illustration. None of a layout's own
    decoration reaches `slide.shapes`, so a reference that listed only the page's
    shapes handed the author a blank background and no way to know why."""
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert source is not None
    assert source.source.splitlines()[0].startswith("slide = prs.slides.add_slide(prs.slide_layouts[")


def test_type_the_template_states_on_its_layout_reaches_the_reference(tmp_path: Path, image):
    """A template that sets its heading once, on the layout, states nothing on the
    page -- so the page's own runs carry no size at all. Read only those and the
    reference says the cover heading has no type, beside a render where it is
    40pt."""
    from pptx import Presentation
    from pptx.util import Pt

    presentation = Presentation()
    layout = presentation.slide_layouts[0]
    layout.placeholders[0].text_frame.paragraphs[0].font.size = Pt(40)
    page = presentation.slides.add_slide(layout)
    page.placeholders[0].text_frame.text = "Inherited heading"
    path = tmp_path / "inheriting.pptx"
    presentation.save(str(path))

    source = decompile(path, 0)

    assert source is not None
    assert "Pt(40)" in source.source


def test_the_reference_states_what_it_cannot_draw(template: Path, tmp_path: Path):
    """Two thirds of real template pages hold something python-pptx cannot write.
    Approximating it silently would produce a page that looks finished and is
    wrong, with no way for the author to tell which elements were guesses."""
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert source is not None
    assert "a custom-drawn shape" in source.unredrawable
    assert "clone the page" in source.summary()


def _charted(path: Path, *, gradient: bool = False) -> Path:
    """A page holding a native chart, the way a real template holds one.

    python-pptx writes a chart even though it cannot rewrite one, which is exactly
    the asymmetry these tests are about.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    page.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1)).text_frame.text = "Quarterly volume"
    data = CategoryChartData()
    data.categories = ["Q1 of the template", "Q2 of the template"]
    data.add_series("the template's own series", (19.2, 21.5))
    page.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(2), Inches(2.5), Inches(6), Inches(3.5), data)
    if gradient:
        panel = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9), Inches(2.5), Inches(3), Inches(3.5))
        panel.fill.gradient()
    presentation.save(str(path))
    return path


def test_a_chart_is_the_authors_to_draw_rather_than_the_clones_to_keep(tmp_path: Path):
    """A live run was told to clone a chart page and replace_text it, got a chart it
    could not fill, and spent 73 minutes in a shell measuring the template's own axis
    type so it could rebuild the chart by hand. Nothing in this engine writes chart
    data and replace_text cannot reach a chart's labels, so a clone delivers the
    template's numbers and no route empties them; the layout is still worth cloning,
    and only the chart is the author's to draw."""
    source = decompile(_charted(tmp_path / "charted.pptx"), 0, images_dir=tmp_path / "img")

    assert source is not None
    assert "a chart" not in source.unredrawable, "cloning is not what gets a chart into a deck"
    assert source.redraw_yourself == ("Inches(2.00), Inches(2.50), Inches(6.00), Inches(3.50)",)
    said = source.summary()
    assert "the chart must be redrawn with ppt_charts from your own data" in said
    assert "placed in the same position" in said
    assert "drop_shape(shape_near(slide, 2.00, 2.50))" in said
    assert "Box.at(2.00, 2.50, w=6.00, h=3.50)" in said
    assert "kept by cloning" not in said, "the line under the chart said the one thing that is untrue of it"


def test_a_page_with_a_chart_and_a_gradient_says_both_without_contradicting_itself(tmp_path: Path):
    """The two answers are opposite -- clone that, draw this -- and a page can want
    both. Said as one note, the chart an exception to the clone, because two notes
    read as two instructions and the clone note ends on "replace its text and
    pictures", which is exactly what a chart does not answer to."""
    source = decompile(_charted(tmp_path / "both.pptx", gradient=True), 0, images_dir=tmp_path / "img")

    assert source is not None
    assert "a gradient fill" in source.unredrawable
    assert source.redraw_yourself
    notes = [line for line in source.summary().splitlines() if "python-pptx cannot write" in line]
    assert len(notes) == 1, notes
    assert "clone the page" in notes[0]
    assert "A chart on it is the exception" in notes[0]
    assert "the chart must be redrawn with ppt_charts" in notes[0]


def test_a_run_of_custom_drawn_shapes_is_said_once(tmp_path: Path):
    """A Bauhaus contents page carries a hundred freeforms as pattern; two lines each
    made one page 12,800 characters of a reply that cuts at 16,000. Consecutive ones
    fold into a single line that keeps their ordinals and the box they span."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "title"
    for n in range(3):
        builder = slide.shapes.build_freeform(Inches(1 + n), Inches(3), scale=914400)
        builder.add_line_segments([(1, 0), (1, 1), (0, 1)])
        builder.convert_to_shape()
    slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = "after"
    path = tmp_path / "pattern.pptx"
    prs.save(path)

    source = decompile(path, 0)

    assert source is not None
    lines = source.source.splitlines()
    folded = [line for line in lines if "custom-drawn shapes" in line]
    assert len(folded) == 1 and folded[0].startswith("# [2]-[4] 3 custom-drawn shapes between")
    assert "# [5]" in lines, "the shape after the run keeps its own ordinal"
    assert sum("custom-drawn shape at" in line for line in lines) == 0


def test_the_reference_names_the_imports_its_code_needs(template: Path, tmp_path: Path):
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert source is not None
    head = source.summary()
    assert "from pptx.util import Inches" in head
    assert "from pptx.enum.text import MSO_ANCHOR" in head


def test_the_reference_keeps_the_typography_the_template_set(template: Path, tmp_path: Path):
    """Anchor and alignment are 30% and 14% of the shapes measured. Dropped, every
    label sits at the top-left of a box the template centres in, and the page reads
    as slightly fallen rather than as wrong."""
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert source is not None
    assert "frame.vertical_anchor = MSO_ANCHOR.MIDDLE" in source.source
    assert "para.alignment = PP_ALIGN.CENTER" in source.source
    assert "Pt(32)" in source.source


def test_type_the_template_states_only_on_its_master_reaches_the_reference(tmp_path: Path):
    """The last rung of the chain, and on the templates measured it is where the
    other 38% of the type is: not one text shape on any page of any of them states
    its own size, 62% take it from the layout's placeholder and the rest from
    here."""
    from pptx import Presentation

    presentation = Presentation()
    master = presentation.slide_masters[0]._element
    title = master.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}txStyles/"
        "{http://schemas.openxmlformats.org/presentationml/2006/main}titleStyle/"
        "{http://schemas.openxmlformats.org/drawingml/2006/main}lvl1pPr/"
        "{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr"
    )
    assert title is not None
    title.set("sz", "5400")
    page = presentation.slides.add_slide(presentation.slide_layouts[5])
    page.placeholders[0].text_frame.text = "From the master"
    path = tmp_path / "mastered.pptx"
    presentation.save(str(path))

    source = decompile(path, 0)

    assert source is not None
    assert "Pt(54)" in source.source


def test_the_reference_keeps_rotation_and_connectors(template: Path, tmp_path: Path):
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert source is not None
    assert "panel.rotation = 15" in source.source
    assert "add_connector(MSO_CONNECTOR.STRAIGHT" in source.source


def test_a_freeform_is_named_rather_than_drawn_as_a_rectangle(template: Path, tmp_path: Path):
    """18% of the shapes measured are freeforms -- the swooshes and cut corners a
    design is recognisable by. `add_shape` would put a rectangle there, which reads
    as a bug in the deck rather than as a limit of the reference."""
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert source is not None
    assert "# a custom-drawn shape at" in source.source
    assert "MSO_SHAPE.RECTANGLE" not in source.source


def test_a_picture_comes_out_as_a_file_the_code_can_name(template: Path, tmp_path: Path):
    source = decompile(template, 0, images_dir=tmp_path / "img")

    assert source is not None
    assert source.picture_files
    assert (tmp_path / "img" / source.picture_files[0]).is_file()
    assert "add_picture(" in source.source


def test_a_group_reaches_page_coordinates(tmp_path: Path):
    """A group declares its own child space -- an offset *and* an extent. Applying
    only the offset put shapes at `Inches(21.61)` on a 13.33-inch canvas."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = page.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    inner.text_frame.text = "inside a group"
    path = tmp_path / "grouped.pptx"
    presentation.save(str(path))

    source = decompile(path, 0)

    assert source is not None
    assert "inside a group" in source.source
    placed = next(line for line in source.source.splitlines() if line.startswith("box = slide.shapes.add_textbox"))
    numbers = [float(value) for value in re.findall(r"Inches\(([-\d.]+)\)", placed)]
    assert numbers and all(0 <= value <= 13.5 for value in numbers)


def test_decompiling_what_is_not_there(template: Path):
    assert decompile(template, 99) is None
    assert decompile(Path("/nonexistent.pptx"), 0) is None


# --- reusing a page -------------------------------------------------------


def test_cloning_a_page_carries_its_pictures(template: Path, tmp_path: Path):
    """Copied shape XML refers to relationships by id, and the new slide has none
    of them, so every picture on a naive copy resolves to nothing -- verified as
    `no relationship with key 'rId4'`. This is why the operation exists."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    source = Presentation(str(template))
    prepared = prepare(template, tmp_path / "prepared.pptx")
    assert prepared is not None
    target = Presentation(str(prepared.path))

    clone = clone_page(target, source.slides[0])
    out = tmp_path / "cloned.pptx"
    target.save(str(out))

    assert len(clone.shapes) == len(source.slides[0].shapes)
    reopened = Presentation(str(out)).slides[0]
    pictures = [s for s in reopened.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert pictures and pictures[0].image.blob


def test_cloning_across_two_files_writes_one_package(template: Path, tmp_path: Path):
    """The normal case -- a page from the user's original into the prepared copy --
    and the one that quietly produces a broken file. Handing `add_slide` a layout
    another package owns writes that layout, its master and its theme into the zip
    a second time under names it already holds, and PowerPoint offers to repair
    what opens."""
    import zipfile

    from pptx import Presentation

    prepared = prepare(template, tmp_path / "prepared.pptx")
    assert prepared is not None
    target = Presentation(str(prepared.path))

    clone_page(target, Presentation(str(template)).slides[0])
    out = tmp_path / "cloned.pptx"
    target.save(str(out))

    names = zipfile.ZipFile(out).namelist()
    assert len(names) == len(set(names))
    assert len(Presentation(str(out)).slides) == 1


def test_replacing_a_picture_keeps_the_frame_the_template_chose(template: Path, tmp_path: Path, image):
    """Deleting the frame and adding another loses the crop, the outline, the
    shadow and the z-order. The frame is the design; only the pixels are content."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(template))
    page = presentation.slides[0]
    picture = next(s for s in page.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    where = (picture.left, picture.top, picture.width, picture.height)
    before = picture.image.blob

    other = image("swap.png", (200, 40, 40))
    replace_picture(picture, other, fit="cover")
    out = tmp_path / "swapped.pptx"
    presentation.save(str(out))

    after = next(s for s in Presentation(str(out)).slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert (after.left, after.top, after.width, after.height) == where
    assert after.image.blob != before
    assert after.image.blob == other.read_bytes()


def test_replacing_text_keeps_how_the_template_set_it(template: Path, tmp_path: Path):
    """Assigning to `.text` drops every run property, so a heading comes back at
    body size in body colour -- the page keeps its geometry and loses its
    typography, which reads as worse than a missing page because it looks
    deliberate."""
    from pptx import Presentation

    presentation = Presentation(str(template))
    heading = next(s for s in presentation.slides[0].shapes if getattr(s, "has_text_frame", False))

    replace_text(heading, "Our results")
    out = tmp_path / "retitled.pptx"
    presentation.save(str(out))

    written = next(s for s in Presentation(str(out)).slides[0].shapes if getattr(s, "has_text_frame", False))
    run = written.text_frame.paragraphs[0].runs[0]
    assert written.text_frame.text == "Our results"
    assert run.font.size.pt == 32
    assert run.font.bold


def test_replacing_text_with_more_lines_than_the_template_had(template: Path, tmp_path: Path):
    """A template drawn with one line has to take three. A new paragraph is copied
    from the last rather than added blank, so it inherits the template's list style
    rather than a python-pptx default."""
    from pptx import Presentation

    presentation = Presentation(str(template))
    heading = next(s for s in presentation.slides[0].shapes if getattr(s, "has_text_frame", False))

    replace_text(heading, "one\ntwo\nthree")

    assert [p.text for p in heading.text_frame.paragraphs] == ["one", "two", "three"]
    assert heading.text_frame.paragraphs[2].runs[0].font.size.pt == 32


def test_replacing_text_leaves_no_line_break_of_the_templates_behind(template: Path):
    """A placeholder written over two lines is run, `<a:br/>`, run. Replacing it drops
    the second run and used to keep the break, so the new text carried a trailing
    empty line: one cover's title sat a line high inside a box that had grown a line
    taller than anything visible in it, and the text read back as "...方法\x0b".
    """
    from pptx import Presentation
    from pptx.oxml.ns import qn

    presentation = Presentation(str(template))
    heading = next(s for s in presentation.slides[0].shapes if getattr(s, "has_text_frame", False))
    paragraph = heading.text_frame.paragraphs[0]
    # The template's own two-line prompt: a break between two runs.
    paragraph._p.append(paragraph._p.makeelement(qn("a:br"), {}))
    second = copy.deepcopy(paragraph.runs[0]._r)
    paragraph._p.append(second)
    assert "\x0b" in heading.text_frame.text

    replace_text(heading, "TarViS：面向目标的视频分割统一方法")

    assert heading.text_frame.text == "TarViS：面向目标的视频分割统一方法"
    assert paragraph._p.findall(qn("a:br")) == []


def test_replacing_text_refuses_a_shape_that_holds_none(template: Path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(template))
    picture = next(s for s in presentation.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)

    with pytest.raises(ValueError, match="holds no text"):
        replace_text(picture, "anything")


def test_dropping_what_the_page_does_not_need(template: Path):
    """The commonest edit after text: the template's page has six cards and this
    one makes four points."""
    from pptx import Presentation

    presentation = Presentation(str(template))
    page = presentation.slides[0]
    before = len(page.shapes)

    drop_shape(page.shapes[0])

    assert len(page.shapes) == before - 1


# --- repeating units, which is what a template page is made of ------------


def _card_page(path: Path, slots: int = 4, shapes_per_slot: int = 2):
    """A page built the way real templates are: one unit repeated in a row.

    Measured over 119 templates and their 1563 example pages: 77% of pages have
    groups, 75% have a unit repeated at least twice, and the most common shapes of
    it are 3x2, 2x3, 4x2 and 4x3. So this is the page shape the operations below
    exist for, not a contrivance.
    """
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = page.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(9.0), Inches(0.8))
    title.text_frame.text = "单击此处添加页面标题"
    for index in range(slots):
        group = page.shapes.add_group_shape()
        number = group.shapes.add_textbox(Inches(0.7 + index * 3.0), Inches(2.0), Inches(0.6), Inches(0.6))
        number.text_frame.text = f"0{index + 1}"
        heading = group.shapes.add_textbox(Inches(0.7 + index * 3.0), Inches(2.8), Inches(2.6), Inches(0.5))
        heading.text_frame.text = "单击添加小标题"
        if shapes_per_slot > 2:
            body = group.shapes.add_textbox(Inches(0.7 + index * 3.0), Inches(3.4), Inches(2.6), Inches(1.2))
            body.text_frame.text = "单击此处添加文本"
    presentation.save(str(path))
    return path


def test_a_unit_repeated_on_a_page_is_found(tmp_path: Path):
    from pptx import Presentation

    from raven_ppt.services.template import units

    page = Presentation(str(_card_page(tmp_path / "cards.pptx"))).slides[0]
    runs = units(page)
    assert [len(run) for run in runs] == [4]


def test_a_flat_page_repeats_nothing(tmp_path: Path):
    """23% of real example pages are flat, and on those `texts` is the whole story."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import units

    presentation = Presentation()
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    page.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "alone"
    path = tmp_path / "flat.pptx"
    presentation.save(str(path))
    assert units(Presentation(str(path)).slides[0]) == []


def test_filling_fewer_items_than_slots_deletes_the_spares(tmp_path: Path):
    """The failure this closes: a live deck wrote "" into two agenda slots, which
    emptied their text and left two numbered bubbles sitting on the page."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, replace_text, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx")))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    replace_text(slide, "单击此处添加页面标题", "四类任务本质相同")
    fill(max(units(slide), key=len), [["01", "视频实例分割"], ["02", "视频全景分割"], ["03", "半监督分割"]])

    assert [len(run) for run in units(slide)] == [3], "the fourth unit is gone, not emptied"
    said = _texts(slide)
    assert "视频实例分割" in said and "半监督分割" in said
    assert "单击添加小标题" not in said, "no slot keeps its placeholder"
    assert "04" not in said


def test_more_items_than_slots_grows_the_run_rather_than_dropping_content(tmp_path: Path):
    """This used to refuse. Ten measured builds died on the refusal, nine of them one or
    two items over, and the authors then cloned units by hand; a regular row grows."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx")))
    out = Presentation(str(tmp_path / "cards.pptx"))

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["0%d" % n, "点 %d" % n] for n in range(1, 7)])

    assert [len(run) for run in units(slide)] == [6]
    assert {"点 1", "点 6", "06"} <= _texts(slide), "nothing dropped, the new slots filled and numbered"


def test_text_inside_a_group_is_replaced_and_unnamed_text_is_left_standing(tmp_path: Path):
    """77% of real example pages have groups, and `slide.shapes` does not descend
    into them -- so this was the reason a deck shipped with the strings its author had
    named still on the page as placeholders.

    And what no call names is still the template's, which is the contract the
    gate reads: an unfilled frame says the template's words out loud, where
    `placeholder_copy` refuses it by name, rather than going blank where nothing can
    tell it from a page the design meant to leave empty.
    """
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, replace_text, shape_at

    source = Presentation(str(_card_page(tmp_path / "cards.pptx", slots=2, shapes_per_slot=3)))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    replace_text(slide, "单击此处添加页面标题", "标题")
    replace_text(shape_at(slide, 3), "组内被按序号替换")
    said = _texts(slide)
    assert "组内被按序号替换" in said, "an index reaches a shape inside a group"
    assert "单击添加小标题" in said, "a frame no call named still says what the template wrote"


def test_a_key_that_matches_nothing_says_what_the_page_holds(tmp_path: Path):
    from pptx import Presentation

    from raven_ppt.services.template import clone_page

    source = Presentation(str(_card_page(tmp_path / "cards.pptx")))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    with pytest.raises(KeyError) as caught:
        replace_text(slide, "没有这段文字", "x")
    assert "单击此处添加页面标题" in str(caught.value), "the refusal lists what is there"


# --- what a failed lookup says -------------------------------------------------------
#
# The class of failure these cover, measured on one live run: six of its seven builds
# died naming a string the page it was editing did not hold, four of them on the same
# key in four consecutive builds. The answer was in every one of those messages -- shape
# 25 of 27 held '数字健康兴起' against the '数字健康崛起' that was asked for, one glyph
# apart and the same meaning -- and the message never said so, so the author read the
# inventory, fixed something else, and rebuilt. The programs differ between attempts, so
# this was never a frozen loop: nothing pointed at the line.


def _page_of_mostly_drawings(path: Path):
    """A page shaped like the one that raised four times: a row of cards whose text is
    outnumbered by wordless boxes, so a hint must not be built out of an empty string."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    heading = page.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(9.0), Inches(0.8))
    heading.text_frame.text = "生物科技如何改变医疗健康行业"
    for index, (label, blurb) in enumerate(
        (
            ("基因编辑革新", "定制化治疗方案成为现实"),
            ("再生医学进展", "器官再生与移植迎来新希望"),
            ("精准医疗崛起", "基于个体基因特征的精准治疗"),
            ("数字健康兴起", "利用数据科技优化医疗服务"),
        )
    ):
        left = 0.7 + index * 3.0
        for spare in range(3):
            page.shapes.add_textbox(Inches(left + spare * 0.1), Inches(1.6), Inches(0.3), Inches(0.3))
        page.shapes.add_textbox(Inches(left), Inches(2.4), Inches(0.6), Inches(0.5)).text_frame.text = f"0{index + 1}"
        page.shapes.add_textbox(Inches(left), Inches(3.0), Inches(2.6), Inches(0.5)).text_frame.text = label
        page.shapes.add_textbox(Inches(left), Inches(3.6), Inches(2.6), Inches(1.0)).text_frame.text = blurb
    presentation.save(str(path))
    return path


def _cloned(tmp_path: Path, source_name: str = "cards.pptx", page: int = 1):
    """The same page these cases always used, cloned the way the copy route clones it."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import prototype

    source = Presentation(str(tmp_path / source_name))
    out = Presentation()
    out.slide_width, out.slide_height = Inches(13.333), Inches(7.5)
    return clone_page(out, prototype(source, page))


def test_a_key_one_glyph_out_is_told_which_shape_it_meant(tmp_path: Path):
    """The decisive live case, replayed: 崛起 for 兴起, near-synonyms one glyph apart.

    Through `replace_text`, which is the lookup that survived `adapt`'s removal. The
    hint and the shape number are the same; what changed is the call and the phrase
    the inventory opens with, because the old one belonged to `adapt`'s own advice.
    """
    _page_of_mostly_drawings(tmp_path / "cards.pptx")
    slide = _cloned(tmp_path)

    with pytest.raises(KeyError) as refused:
        replace_text(slide, "数字健康崛起", "Digital health")

    message = str(refused.value)
    hint, _, inventory = message.partition("The page holds")
    assert "数字健康兴起" in hint, "the near miss is named before the inventory, not inside it"
    assert "数字健康崛起" in hint, "and beside what was actually asked for"
    assert inventory, "the inventory still follows it -- it is the answer when nothing is near"
    assert "'数字健康兴起'" in hint


def test_a_hint_is_never_built_out_of_a_wordless_shape(tmp_path: Path):
    """12 of the 27 shapes on the page that raised four times hold no text at all, and an
    empty string scores against anything short."""
    _page_of_mostly_drawings(tmp_path / "cards.pptx")
    slide = _cloned(tmp_path)

    with pytest.raises(KeyError) as refused:
        replace_text(slide, "完全无关的一段文字", "x")

    hint = str(refused.value).partition("The page holds")[0]
    assert "Did you mean" not in hint, "nothing on the page is close, so nothing is offered"
    assert "''" not in hint


def test_a_key_absent_from_the_page_but_present_on_another_is_traced_to_it(tmp_path, monkeypatch):
    """The live run's other two failures: the author named page 2's copy while writing
    into page 1, so no shape on the page is near it and the inventory cannot explain a
    string the page never held. Word for word only -- the page number is right or there
    is no clause.

    The clause needs the file the page was cloned out of. `adapt` held it and is gone;
    `replace_text` is handed a page already in the deck being built, so the template is
    reached where the runner puts it, `PPT_TEMPLATE_SOURCE`. Bound here the way the
    runner binds it.
    """
    from pptx import Presentation
    from pptx.util import Inches

    _page_of_mostly_drawings(tmp_path / "cards.pptx")
    source = Presentation(str(tmp_path / "cards.pptx"))
    second = source.slides.add_slide(source.slide_layouts[6])
    second.shapes.add_textbox(
        Inches(1.0), Inches(1.0), Inches(6.0), Inches(0.8)
    ).text_frame.text = "新技术、新产品及新服务在行业中的应用"
    source.save(str(tmp_path / "two.pptx"))
    monkeypatch.setenv("PPT_TEMPLATE_SOURCE", str(tmp_path / "two.pptx"))
    slide = _cloned(tmp_path, "two.pptx", 1)

    with pytest.raises(KeyError) as refused:
        replace_text(slide, "新技术、新产品及新服务在行业中的应用", "x")

    hint = str(refused.value).partition("The page holds")[0]
    assert "is on page 2 of this template" in hint


def test_a_key_absent_with_no_template_bound_says_nothing_extra(tmp_path, monkeypatch):
    """And the same call with no template in the environment: one clause fewer, never a
    second error on top of the refusal being raised."""
    monkeypatch.delenv("PPT_TEMPLATE_SOURCE", raising=False)
    _page_of_mostly_drawings(tmp_path / "cards.pptx")
    slide = _cloned(tmp_path)

    with pytest.raises(KeyError) as refused:
        replace_text(slide, "新技术、新产品及新服务在行业中的应用", "x")

    hint = str(refused.value).partition("The page holds")[0]
    assert "of this template" not in hint


def test_a_key_too_short_to_judge_is_offered_nothing(tmp_path: Path):
    """A page's 01/02/03/04 row scores 0.5 against every member of itself, so under four
    characters the measure is turned off rather than tuned."""
    _page_of_mostly_drawings(tmp_path / "cards.pptx")
    slide = _cloned(tmp_path)

    with pytest.raises(KeyError) as refused:
        replace_text(slide, "09", "x")

    hint = str(refused.value).partition("The page holds")[0]
    assert "Did you mean" not in hint


def test_what_matches_is_untouched_by_what_the_failure_says(tmp_path: Path):
    """The change is to the refusal, not to eligibility: a substring still resolves and a
    key that resolved before still resolves."""
    _page_of_mostly_drawings(tmp_path / "cards.pptx")
    slide = _cloned(tmp_path)

    replace_text(slide, "数字健康兴起", "Digital health")
    replace_text(slide, "利用数据", "Data science")

    said = [s.text_frame.text for s in slide.shapes if getattr(s, "has_text_frame", False)]
    assert "Digital health" in said
    assert "Data science" in said


def test_the_replace_text_refusal_names_the_copy_it_meant(tmp_path: Path):
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, replace_text

    source = Presentation(str(_page_of_mostly_drawings(tmp_path / "cards.pptx")))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])

    with pytest.raises(KeyError) as refused:
        replace_text(slide, "数字健康崛起", "Digital health")

    hint = str(refused.value).partition("The page holds")[0]
    assert "数字健康兴起" in hint
    assert "The page holds" in str(refused.value), "the inventory is kept"


def test_the_unit_refusal_names_the_slot_it_meant(tmp_path: Path):
    """A dict item is keyed by the copy a slot holds now, so the same slip lands here."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx")))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    run = max(units(slide), key=len)

    with pytest.raises(KeyError) as refused:
        fill(run, [{"单击添加小标提": "Digital health"}])

    message = str(refused.value)
    assert "Did you mean" in message.partition("It holds")[0], message
    assert "单击添加小标题" in message.partition("It holds")[0]
    assert "It holds" in message, "the unit's own copy is still listed"


def test_the_prefix_refusal_names_the_copy_it_meant(tmp_path: Path):
    """`shape_saying` matches a prefix, so the key is weighed against each candidate's
    opening of the same length -- scoring six characters against a whole paragraph finds
    nothing."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import shape_saying

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(6.0), Inches(0.6))
    box.text_frame.text = "数字健康兴起：数据科技如何改变医疗服务的交付方式"

    with pytest.raises(KeyError) as refused:
        shape_saying(slide, "数字健康崛起")

    message = str(refused.value)
    assert "Did you mean" in message.partition("Its copy reads")[0]
    assert "Its copy reads" in message


def test_the_arrangement_of_a_run_is_reported_not_reflowed(tmp_path: Path):
    """29% of real runs follow no grid, so re-flowing on the engine's own initiative
    would destroy the arrangement a template was drawn with. It reports; `place` moves."""
    from pptx import Presentation

    from raven_ppt.services.template import arrangement, boxes, clone_page, place, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx")))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    run = max(units(slide), key=len)
    assert arrangement(run) == ("row", 1, 4)
    place(run[0], (1.0, 5.0, 2.0, 1.0))
    assert [round(value, 2) for value in boxes(run)[0]] == [1.0, 5.0, 2.0, 1.0]


def _layout_box(x0: float, y0: float, x1: float, y1: float):
    """A `ppt_layout.Box`, taken from the projection an author's script imports.

    The grid module reaches the author as a file beside their script rather than as an
    import of this package, so the class that turns up at `place` is the projected one --
    which is the reason `place` recognises a box by its corners and not by its type.
    """
    from raven_ppt.services.assets.layout import layout_module_source

    namespace: dict = {}
    exec(compile(layout_module_source(), "ppt_layout.py", "exec"), namespace)
    return namespace["Box"].corners(x0, y0, x1, y1)


def _layout_box_at(x: float, y: float, w: float, h: float):
    """The same projected `Box`, built from the corner and size a reference prints."""
    from raven_ppt.services.assets.layout import layout_module_source

    namespace: dict = {}
    exec(compile(layout_module_source(), "ppt_layout.py", "exec"), namespace)
    return namespace["Box"].at(x, y, w=w, h=h)


def _run_of(path: Path):
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, units

    source = Presentation(str(path))
    out = Presentation(str(path))
    return max(units(clone_page(out, source.slides[0])), key=len)


def test_a_layout_box_reaching_place_is_the_two_corners_it_is(tmp_path: Path):
    """One word, two rectangles, and the wrong reading drew a frame off the page.

    `place` takes (left, top, width, height) and `ppt_layout.Box` is (x0, y0, x1, y1),
    and both are called `box`. Unpacked as a size,
    `place(unit, Box.corners(0.72, 1.24, 12.6, 6.7))` made a 12.6x6.7in frame on a
    13.33x7.5in page where those corners name an 11.88x5.46in one -- and it drew that
    frame in silence, which no render reports, because a box that is wrong is still a box.

    A Box says which of the two rectangles it is, so it is converted. Four bare numbers
    cannot, so they stay a size -- which is what every python-pptx call beside them takes.
    """
    from raven_ppt.services.template import boxes, place

    run = _run_of(_card_page(tmp_path / "cards.pptx"))

    place(run[0], _layout_box(0.72, 1.24, 12.6, 6.7))
    assert [round(value, 2) for value in boxes(run)[0]] == [0.72, 1.24, 11.88, 5.46]

    place(run[1], (0.72, 1.24, 12.6, 6.7))
    assert [round(value, 2) for value in boxes(run)[1]] == [0.72, 1.24, 12.6, 6.7]


def test_a_rectangle_in_emu_is_refused_rather_than_placed(tmp_path: Path):
    """`box.pptx()` is the third spelling of a rectangle and the only unreadable one.

    Its four numbers are python-pptx lengths, so `place(unit, box.pptx())` set a 12.6in
    frame to 11521440 inches and said nothing. A shape's own `.left` and `.width` arrive
    the same way, which is how an author copies one unit's geometry onto another.
    """
    from raven_ppt.services.template import place

    run = _run_of(_card_page(tmp_path / "cards.pptx"))

    with pytest.raises(ValueError) as refused:
        place(run[0], _layout_box(0.72, 1.24, 12.6, 6.7).pptx())
    said = str(refused.value)
    assert "EMU" in said and "914400" in said
    assert "0.72, 1.24, 11.88, 5.46" in said, "the inches those lengths stand for"

    other = run[1]
    with pytest.raises(ValueError, match="EMU"):
        place(run[0], (other.left, other.top, other.width, other.height))

    for neither in ((0.72, 1.24, 12.6), 5, ("a", "b", "c", "d")):
        with pytest.raises(ValueError, match="ppt_layout Box"):
            place(run[0], neither)


def test_the_boxes_of_a_run_are_sizes_and_say_so(tmp_path: Path):
    """`boxes` hands back what `place` takes, and the docstring is where that is settled.

    Read as corners, a 2.0in-wide unit at x=1.0 is a unit ending at x=2.0 -- 1.0in wide,
    half of what is on the page -- and the two readings are the same four numbers. So the
    order is written down where an author reads it rather than inferred from the values.
    """
    from raven_ppt.services.template import boxes, place

    run = _run_of(_card_page(tmp_path / "cards.pptx"))
    place(run[0], (1.0, 5.0, 2.0, 1.0))

    spot = boxes(run)[0]
    assert (round(spot[2], 2), round(spot[3], 2)) == (2.0, 1.0)
    assert round(run[0].width / 914400, 2) == 2.0, "the third number is a width, not a far edge"
    assert "(left, top, width, height)" in boxes.__doc__
    assert "ppt_layout.Box" in boxes.__doc__, "the other rectangle the word names"


def test_a_pictures_box_reshapes_the_frame_in_either_spelling(tmp_path: Path, image):
    """The reshape a fitting refusal offers, given as a size and as two corners.

    `replace_picture(shape, image, box=...)` exists because a live run told its landscape
    figure could not go in a portrait frame answered by permuting the shape number three
    times. The box goes through the same reading as `place`: a Box is converted, four
    numbers are a size. Stretch, so the frame stays exactly where it was put and the
    assertion is about the box rather than about the fit.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page

    path = _card_page(tmp_path / "cards.pptx")
    once = Presentation(str(path))
    once.slides[0].shapes.add_picture(str(image("shot.png", (90, 90, 90))), Inches(1), Inches(4), width=Inches(3))
    once.save(str(path))

    def framed(box):
        source, out = Presentation(str(path)), Presentation(str(path))
        slide = clone_page(out, source.slides[0])
        frame = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        replace_picture(frame, str(image("fig.png", (10, 20, 30))), "stretch", box=box)
        return [round(value / 914400, 2) for value in (frame.left, frame.top, frame.width, frame.height)]

    assert framed(_layout_box(0.8, 1.6, 8.2, 5.8)) == [0.8, 1.6, 7.4, 4.2]
    assert framed((0.8, 1.6, 7.4, 4.2)) == [0.8, 1.6, 7.4, 4.2]

    with pytest.raises(ValueError, match="EMU"):
        framed(_layout_box(0.8, 1.6, 8.2, 5.8).pptx())


def _texts(slide) -> set[str]:
    from raven_ppt.services.measure.geometry import iter_shapes

    return {
        " ".join(shape.text_frame.text.split())
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
    }


def test_the_reference_numbering_is_the_numbering_shape_at_takes(tmp_path: Path):
    """One numbering for the page, printed where the author reads it.

    A live author read the reference, counted, and named shape 15 on a page whose text
    frames stopped at fourteen: the reference numbered shapes and the route since
    removed numbered text frames. The two orders have to be the same walk.
    """
    import re

    from pptx import Presentation

    from raven_ppt.services.template import shape_at
    from raven_ppt.services.template.compose import _all_shapes
    from raven_ppt.services.template.decompile import _flatten

    path = _card_page(tmp_path / "cards.pptx", slots=3, shapes_per_slot=3)
    source = decompile(path, 0)
    assert source is not None
    printed = [int(number) for number in re.findall(r"^# \[(\d+)\]", source.source, re.M)]
    assert printed == list(range(1, len(printed) + 1)), "the ordinals run 1..n with no gaps"

    page = Presentation(str(path)).slides[0]
    walked = list(_all_shapes(page.shapes))
    assert len(walked) == len(list(_flatten(page.shapes))) == len(printed)

    # And the ordinal reaches the shape the reference showed at that ordinal.
    wanted = next(index for index, shape in enumerate(walked, start=1) if shape.text_frame.text == "02")
    out = Presentation(str(path))
    slide = clone_page(out, page)
    replace_text(shape_at(slide, wanted), "第二张卡")
    assert "第二张卡" in _texts(slide)


def test_an_index_pointing_at_the_wrong_kind_of_shape_is_refused(tmp_path: Path, image):
    """A number that lands on copy is a miscount, and the refusal lists the page.

    Two models each aimed at a page's second shape when its picture was the third. The
    check used to sit in the call that resolved the number and moved onto this one with
    it (D41), because the numbering is what makes the mistake worth a refusal: without
    it a card's heading silently became a photograph.
    """
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, prototype, replace_picture, shape_at

    photo = image("shot.png", (90, 90, 90))
    path = _card_page(tmp_path / "cards.pptx")
    twice = Presentation(str(path))
    page = twice.slides[0]
    page.shapes.add_picture(str(photo), Inches(1), Inches(4), width=Inches(3))
    page.shapes.add_picture(str(photo), Inches(6), Inches(4), width=Inches(3))
    twice.save(str(path))

    source, out = Presentation(str(path)), Presentation(str(path))
    slide = clone_page(out, prototype(source, 1))
    with pytest.raises(ValueError, match="picture cannot stand in for"):
        replace_picture(shape_at(slide, 2), str(photo))


def test_a_landscape_figure_in_a_portrait_slot_is_placed_with_a_warning(tmp_path: Path, image) -> None:
    """A layout decision, handed back rather than hidden.

    A template's portrait photo slot runs about 0.6 wide-to-tall and a paper's
    architecture figure about 2.4. Contained, the figure became a strip a quarter of the
    frame's height with empty space above and below; cropped, it would have lost its
    outer columns. A live page shipped the first. Only the author can decide which way
    the page should go, so the numbers go back with the ways out.
    """
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import replace_picture

    wide = tmp_path / "figure.png"
    Image.new("RGB", (2400, 1000), (40, 60, 200)).save(wide)
    presentation = Presentation()
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = page.shapes.add_picture(str(image("slot.png", (9, 9, 9))), Inches(1), Inches(1), Inches(3.5), Inches(5.8))

    with pytest.warns(UserWarning, match="apart") as caught:
        replace_picture(frame, wide)

    said = str(caught[0].message)
    assert "Placed as asked" in said
    assert "sits as a strip" in said, "contained is the default fit, and the warning says what that does"
    assert frame.image is not None, "the picture was placed all the same"


def test_a_figure_near_its_frames_proportions_is_fitted(tmp_path: Path, image) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import replace_picture

    near = tmp_path / "near.png"
    Image.new("RGB", (1600, 900), (40, 60, 200)).save(near)
    presentation = Presentation()
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = page.shapes.add_picture(str(image("slot.png", (9, 9, 9))), Inches(1), Inches(1), Inches(6.0), Inches(4.0))

    replace_picture(frame, near)
    # contain gave way on the height and kept the centre
    assert round(frame.width / frame.height, 2) == round(1600 / 900, 2)


def test_replace_picture_swaps_a_photograph_used_as_a_shape_fill(tmp_path: Path) -> None:
    """`template_picture` tells an author to replace the template's photograph, and on a
    picture-*filled* shape that advice used to end in `AttributeError: blipFill`: the two
    spellings differ (`p:pic/p:blipFill` against `p:sp/p:spPr/a:blipFill`) and only the
    first was known. The fill also has no frame to shrink, so the swap crops through the
    source rectangle instead -- without which the template's own stretch survives and a
    figure comes out distorted.
    """
    from lxml import etree
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_picture

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    rels = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    first = tmp_path / "one.png"
    second = tmp_path / "two.png"
    Image.new("RGB", (800, 600), (20, 20, 20)).save(first)
    Image.new("RGB", (1600, 900), (200, 200, 200)).save(second)

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8), Inches(1), Inches(4), Inches(3))
    _, relationship = shape.part.get_or_add_image_part(str(first))
    fill = etree.SubElement(shape._element.spPr, f"{{{namespace}}}blipFill")
    etree.SubElement(fill, f"{{{namespace}}}blip").set(f"{{{rels}}}embed", relationship)

    replace_picture(shape, second, fit="cover")

    blip = shape._element.spPr.find(f"{{{namespace}}}blipFill/{{{namespace}}}blip")
    assert blip.get(f"{{{rels}}}embed") != relationship, "the fill points at the new image"
    crop = shape._element.spPr.find(f"{{{namespace}}}blipFill/{{{namespace}}}srcRect")
    assert crop is not None, "the source is cropped rather than stretched"
    assert {side for side in ("l", "r", "t", "b") if crop.get(side)} == {"l", "r"}, (
        "a 1.78 picture in a 1.33 frame is trimmed on its sides"
    )


def test_replace_picture_drops_the_inset_the_template_fitted_its_own_photo_with(tmp_path: Path) -> None:
    """The crop is stated against the frame, so a surviving `fillRect` contradicts it.

    Measured on `warm_bauhaus_quarterly_review` page 4: the replacement got a 23.9%
    crop for the frame and kept the template's -32% inset for a box half again as
    wide, and the two together read as a 1.64x stretch.
    """
    from lxml import etree
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_picture

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    rels = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    first, second = tmp_path / "one.png", tmp_path / "two.png"
    Image.new("RGB", (800, 600), (20, 20, 20)).save(first)
    Image.new("RGB", (800, 458), (200, 200, 200)).save(second)

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0), Inches(1.49), Inches(5.49), Inches(6.01))
    _, relationship = shape.part.get_or_add_image_part(str(first))
    fill = etree.SubElement(shape._element.spPr, f"{{{namespace}}}blipFill")
    etree.SubElement(fill, f"{{{namespace}}}blip").set(f"{{{rels}}}embed", relationship)
    stretch = etree.SubElement(fill, f"{{{namespace}}}stretch")
    inset = etree.SubElement(stretch, f"{{{namespace}}}fillRect")
    inset.set("l", "-32200")
    inset.set("r", "-32000")

    replace_picture(shape, second, fit="cover")

    fill = shape._element.spPr.find(f"{{{namespace}}}blipFill")
    assert fill.find(f"{{{namespace}}}srcRect") is not None, "the crop states the fit"
    assert fill.find(f"{{{namespace}}}stretch/{{{namespace}}}fillRect") is None, (
        "the template's inset went with the image it was cut for"
    )


def test_a_shape_in_a_group_is_found_at_the_position_the_page_shows_it(tmp_path: Path) -> None:
    """`shape.left` inside a group is in the group's own space, scaled by its extents.

    Three of four live authors compared it against a page coordinate and got
    `no text shape near (1.56, 2.47)` for a shape that was exactly there.
    """
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    from raven_ppt.services.template.compose import page_position, shape_near

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(2.0), Inches(1.0), Inches(2.0), Inches(0.5))
    first.text_frame.text = "inside the group"
    second = slide.shapes.add_textbox(Inches(5.0), Inches(1.0), Inches(2.0), Inches(0.5))
    second.text_frame.text = "also inside"
    group = slide.shapes.add_group_shape([first, second])
    # A group whose children are drawn at half scale and offset a page inch down.
    group.left, group.top = Inches(1.0), Inches(3.0)
    group.width, group.height = Emu(group.width // 2), Emu(group.height // 2)

    at_left, at_top = page_position(first)

    assert (at_left, at_top) != pytest.approx((2.0, 1.0)), "the declared numbers are the group's"
    # python-pptx builds a fresh proxy per access, so identity is the element.
    assert shape_near(slide, at_left, at_top, with_text=True)._element is first._element
    assert Pt  # the import is what the author's own script does


def test_a_position_that_matches_nothing_says_what_the_page_holds(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import shape_near

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(2.0), Inches(0.5))
    box.text_frame.text = "the only copy"

    with pytest.raises(KeyError) as refused:
        shape_near(slide, 9.0, 6.0)

    assert "the page holds" in str(refused.value).casefold()
    assert "(1.00, 1.00)" in str(refused.value)


def _positions_page(spots):
    """A page whose shapes sit at given inch positions, for weighing a failed lookup."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for left, top, text in spots:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(2.0), Inches(0.5))
        if text:
            box.text_frame.text = text
    return slide


def test_a_point_on_the_asked_for_axis_is_named_as_a_row_counted_one_off():
    """A live page failure, replayed. Asked for (7.39, 2.3): the nearest shape is a
    drawing 0.72in away and the one the author meant is the text box 0.87in away on
    exactly the x it asked for -- so distance alone names the wrong one, and the axis
    that matches is what says the mistake was a row counted off by one. Both are
    offered, each with its reason, rather than one of them guessed at."""
    from raven_ppt.services.template.compose import shape_near

    slide = _positions_page(
        [
            (0.72, 0.14, "成都夜间消费多中心格局"),
            (6.98, 1.71, None),
            (7.39, 1.43, "存量更新片区"),
            (1.23, 2.23, "核心商圈带"),
        ]
    )

    with pytest.raises(KeyError) as refused:
        shape_near(slide, 7.39, 2.30, 0.2)

    hint = str(refused.value).partition("Positions are on the page")[0]
    assert "0.72in away" in hint, "the nearest shape is named, with the distance"
    assert "存量更新片区" in hint, "and so is the one sharing the x that was asked for"
    assert "on the x you asked for" in hint
    assert "a row counted one off" in hint
    assert "The page holds" in str(refused.value), "the inventory is kept"


def test_the_nearest_shape_carries_the_axis_note_when_it_is_the_same_shape():
    """The second live failure: one shape is both the nearest and the one on the asked-for
    x, so it is one clause rather than two."""
    from raven_ppt.services.template.compose import shape_near

    slide = _positions_page([(0.72, 0.14, "业态规划与运营主体"), (6.56, 3.78 + 0.48, None), (1.10, 2.36, "业态配比")])

    with pytest.raises(KeyError) as refused:
        shape_near(slide, 6.62, 3.78, 0.2)

    hint = str(refused.value).partition("Positions are on the page")[0]
    assert hint.count("Nearest is") == 1
    assert "0.48in away, on the x you asked for" in hint
    assert "a row counted one off" in hint


def test_a_point_in_empty_space_is_pointed_nowhere():
    """Over the ten bundled templates no grid point further than 3in from every shape
    draws a clause at all, and none anywhere draws one naming something over 2in away --
    a hint that sends the author across the page is worse than the silence it replaces."""
    from raven_ppt.services.template.compose import shape_near

    slide = _positions_page([(0.5, 0.5, "top left"), (0.5, 1.5, "under it")])

    with pytest.raises(KeyError) as refused:
        shape_near(slide, 11.0, 6.5, 0.2)

    hint = str(refused.value).partition("Positions are on the page")[0]
    assert "Nearest is" not in hint
    assert "counted one off" not in hint
    assert "The page holds" in str(refused.value), "which is what makes it actionable instead"


def test_the_tolerance_that_matches_is_not_widened_by_what_the_failure_says():
    """The 0.2in radius is the contract. A shape 0.5in out still refuses, and a shape
    inside the radius still resolves -- the change is only to the wording of the refusal."""
    from raven_ppt.services.template.compose import shape_near

    slide = _positions_page([(3.0, 3.0, "wanted"), (3.5, 3.0, "beside it")])

    assert shape_near(slide, 3.05, 3.05, 0.2).text_frame.text == "wanted"
    with pytest.raises(KeyError):
        shape_near(slide, 3.0, 3.5, 0.2)
    # And a tolerance the caller narrows is what "on the x you asked for" means, so a
    # shape 0.08in off the x cannot be called aligned under tol=0.05.
    with pytest.raises(KeyError) as refused:
        shape_near(slide, 3.08, 4.0, 0.05)
    assert "on the x you asked for" not in str(refused.value)


def test_copy_the_template_states_under_the_floor_is_lifted_and_stops_refitting(tmp_path: Path) -> None:
    """Two live authors wrote this for themselves, same name, same 14pt default."""
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    from raven_ppt.services.template.compose import raise_type

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    body = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(4.0), Inches(0.6))
    run = body.text_frame.paragraphs[0].add_run()
    run.text = "a paragraph long enough to be body copy rather than a label"
    run.font.size = Pt(12)
    body.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    label = slide.shapes.add_textbox(Inches(6.0), Inches(1.0), Inches(1.0), Inches(0.4))
    tiny = label.text_frame.paragraphs[0].add_run()
    tiny.text = "01"
    tiny.font.size = Pt(10)

    assert raise_type(slide) == 1

    assert body.text_frame.paragraphs[0].runs[0].font.size == Pt(14)
    assert body.text_frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT, (
        "the autofit is what shrank it, and turning it off alone leaves the lifted copy running "
        "out of a box drawn for the smaller type"
    )
    assert label.text_frame.paragraphs[0].runs[0].font.size == Pt(10), "a two-character label is set small on purpose"


def test_a_shape_is_found_by_the_copy_it_starts_with(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import shape_saying

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    wanted = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(4.0), Inches(0.6))
    wanted.text_frame.text = "Method: three steps"

    assert shape_saying(slide, "Method")._element is wanted._element
    with pytest.raises(KeyError, match="[Ii]ts copy reads"):
        shape_saying(slide, "Results")


def test_the_templates_own_wording_is_still_there_after_a_replacement(tmp_path: Path) -> None:
    """The failure this used to describe cannot happen any more.

    Three builds of one live run died looking for a template string the route since
    removed had emptied one line earlier, and a fourth wrote its own version of
    `shape_saying` without the raise and shipped eight blank pages. Nothing empties a
    frame now, so the block the author is reaching for is exactly where the template
    left it -- and a prefix that still matches nothing gets a refusal about the prefix,
    not about a pass that is gone.
    """
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import prototype, shape_saying

    template = Presentation()
    template.slide_width, template.slide_height = Inches(13.333), Inches(7.5)
    page = template.slides.add_slide(template.slide_layouts[6])
    for index, said in enumerate(("趋势展望", "在演示中，简洁清晰的逻辑")):
        box = page.shapes.add_textbox(Inches(1.0), Inches(1.0 + index), Inches(6.0), Inches(0.6))
        box.text_frame.text = said

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    slide = clone_page(built, prototype(template, 1))
    replace_text(slide, "趋势展望", "Momentum")

    assert shape_saying(slide, "在演示中").text_frame.text == "在演示中，简洁清晰的逻辑"
    with pytest.raises(KeyError) as raised:
        shape_saying(slide, "没有这一段")
    assert "empties every text" not in str(raised.value)
    assert "`replace_text(slide, old, new)`" in str(raised.value)


def test_an_empty_value_over_a_number_restates_it_instead_of_blanking_it(tmp_path: Path):
    """A template numbers its slots 01..08 and the deck has fewer sections.

    Measured on a live run: the author spent three requests working out which of the
    unit's two shapes was the number -- `["...", "01"]`, then `["...", None]`, then
    `["...", ""]` -- and shipped the one that blanked all six, so the page came out
    with six empty circles where the template had numbers. An empty string over a
    number now means "this unit's number", which is the only thing it could sensibly
    mean.
    """
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, replace_text, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx", slots=6)))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    replace_text(slide, "单击此处添加页面标题", "四类任务本质相同")
    fill(max(units(slide), key=len), [["", f"第 {n} 条"] for n in range(1, 5)])

    said = _texts(slide)
    assert [n for n in ("01", "02", "03", "04") if n in said] == ["01", "02", "03", "04"]
    assert "05" not in said and "06" not in said, "the spare units are gone, numbers and all"


def test_a_single_digit_template_keeps_a_single_digit(tmp_path: Path):
    """The padding is the template's, not this function's."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    path = _card_page(tmp_path / "cards.pptx", slots=4)
    source = Presentation(str(path))
    for group in source.slides[0].shapes:
        for shape in getattr(group, "shapes", []):
            if shape.text_frame.text.strip().startswith("0"):
                shape.text_frame.text = shape.text_frame.text.strip().lstrip("0")
    source.save(str(path))

    source = Presentation(str(path))
    out = Presentation(str(path))
    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["", "甲"], ["", "乙"]])

    said = _texts(slide)
    assert "1" in said and "2" in said
    assert "01" not in said


def test_an_empty_value_over_words_still_empties_them(tmp_path: Path):
    """Only a number is restated. Everything else means what it says."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx", slots=3)))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["01", ""], ["02", ""]])

    assert "单击添加小标题" not in _texts(slide)


def test_a_number_kept_with_none_is_the_templates_own(tmp_path: Path):
    """`None` has not changed: it leaves the shape exactly as the template wrote it."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx", slots=4)))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [[None, "甲"], [None, "乙"]])

    said = _texts(slide)
    assert "01" in said and "02" in said


def test_a_short_item_restates_the_number_and_leaves_the_rest_to_the_gate(tmp_path: Path):
    """What a short list keeps, and who answers for the rest.

    Taking the numbers off a real template's agenda: the unit holds the number beside
    the copy, the author gave fewer values than shapes, and the pass that used to empty
    unnamed text emptied all six of them, so the page shipped blank folders where the
    template had 01 to 08. The number is still restated for its new position, because a
    unit moved up the run cannot keep the old one.

    Everything else a short list runs out before is now left saying what the template
    wrote. That is not a good page either -- a delivered deck's four cards each kept
    "单击添加小标题" under the author's own heading -- but it is a page
    `placeholder_copy` refuses by name, which a blank one was not.
    """
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    source = Presentation(str(_card_page(tmp_path / "cards.pptx", slots=4, shapes_per_slot=3)))
    out = Presentation(str(tmp_path / "cards.pptx"))
    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["01", "甲"], ["02", "乙"]])

    said = _texts(slide)
    assert "01" in said and "甲" in said
    assert "单击此处添加文本" in said, "a placeholder past the end of the list stands for the gate to refuse"


def test_a_short_item_restates_a_number_it_never_reached(tmp_path: Path):
    """The agenda case on its own: the number is the tail, and it survives.

    `_card_page` puts the number first, so the test above gives it a value. Here the
    unit is walked so the number is what the list runs out before -- the shape the
    author had no opinion about -- and it has to come back renumbered for its new
    position rather than emptied.
    """
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, fill, units

    # Built here rather than from `_card_page`, which puts the number first: the unit
    # this protects is the one the docstring above describes, a label with the number
    # after it, so a one-value list runs out exactly before the number.
    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    page = built.slides.add_slide(built.slide_layouts[6])
    for index in range(4):
        group = page.shapes.add_group_shape()
        label = group.shapes.add_textbox(Inches(0.7 + index * 3.0), Inches(2.8), Inches(2.6), Inches(0.5))
        label.text_frame.text = "单击添加小标题"
        number = group.shapes.add_textbox(Inches(0.7 + index * 3.0), Inches(2.0), Inches(0.6), Inches(0.6))
        number.text_frame.text = f"0{index + 1}"
    built.save(str(tmp_path / "agenda.pptx"))

    source = Presentation(str(tmp_path / "agenda.pptx"))
    out = Presentation(str(tmp_path / "agenda.pptx"))
    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["甲"], ["乙"], ["丙"]])

    said = _texts(slide)
    assert {"甲", "乙", "丙"} <= said
    assert {"01", "02", "03"} <= said, "a number past the end of the list is the template's own"
    assert "04" not in said, "the fourth unit was not filled, so it is gone rather than emptied"
    assert "单击添加小标题" not in said


def test_a_hidden_page_keeps_its_number_but_is_not_offered(tmp_path) -> None:
    """The bundled templates each shipped two hidden pages of the vendor's own
    advertising, and nothing stopped an author naming one as a prototype.

    Numbering has to survive the fix: a page is asked for by its place in the file,
    so hiding page 2 must not turn page 3 into page 2.
    """
    from pptx import Presentation

    from raven_ppt.services.template.menu import menu

    presentation = Presentation()
    for _ in range(3):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides[1].element.set("show", "0")
    path = tmp_path / "with_a_hidden_page.pptx"
    presentation.save(path)

    listing = menu(path)

    assert [entry.number for entry in listing] == [1, 2, 3]
    assert [entry.hidden for entry in listing] == [False, True, False]
    assert "hidden in the file" in listing[1].line()


def test_the_layout_the_template_named_outranks_what_the_page_happens_to_say() -> None:
    """The precedence that decides which page a deck closes on.

    Read heading-first, `gold_panel` page 14 -- "annual reflections and thanks" on a
    layout called `Section Header` -- was that template's closing page, and page 25, on
    a layout called `Closing`, was never offered as one. `roles()` takes the first
    match, so the wrong page won by sitting eleven pages earlier.
    """
    from raven_ppt.services.template.menu import AGENDA, CLOSING, COVER, SECTION, _role

    assert _role(14, "Section Header", "年度感悟与感谢", blocks=3, longest=22) == SECTION
    assert _role(25, "Closing", "演示结束", blocks=3, longest=13) == CLOSING
    assert _role(1, "Title Slide", "A deck about something", blocks=4, longest=30) == COVER
    assert _role(2, "Title and Content", "Agenda", blocks=6, longest=20, says=("Agenda",)) == AGENDA
    # Position still beats the shape of the page: a first page on an unnamed layout
    # carrying one short line is the cover, not a divider.
    assert _role(1, "Blank", "Something", blocks=1, longest=9) == COVER


def test_a_page_that_thanks_the_reader_over_six_cards_is_a_content_page() -> None:
    """`gold_panel` page 18 is a six-card content page headed "thanks and outlook".

    The words alone made it the closing page, which took it out of the offered content
    list -- and seven of the measured answers named it anyway, off the render. A closing
    page carries a farewell and little else, so the shape decides.
    """
    from raven_ppt.services.template.menu import _CLOSING_BLOCKS, CLOSING, _role

    assert _role(18, "Title Only", "感谢与展望", blocks=25, longest=40) == ""
    assert _role(18, "Title Only", "感谢与展望", blocks=_CLOSING_BLOCKS, longest=40) == CLOSING


def test_an_agenda_names_itself_somewhere_other_than_its_first_line(tmp_path: Path) -> None:
    """Two templates reported no agenda page, for two versions of the same reason.

    `_heading` takes the first text line in document order: on `beige_geometric` page 2
    that is "01" and the word "Agenda" is in the last shape on the page. `gold_panel`
    page 2 sets the same word to wrap, so it arrives as "AG" and "ENDA" and no substring
    test finds it. Both pages are the template's agenda.
    """
    from raven_ppt.services.template.menu import AGENDA, _role, _says

    class _Frame:
        def __init__(self, text: str) -> None:
            self.text = text

    class _Shape:
        def __init__(self, text: str) -> None:
            self.text_frame = _Frame(text)

    listed = [_Shape("01"), _Shape("Review of the work content"), _Shape("Agenda")]
    wrapped = [_Shape("<"), _Shape("年度回顾与成长"), _Shape("AG\nENDA")]

    assert _role(2, "Blank", "01", blocks=13, longest=46, says=_says(listed)) == AGENDA
    assert _role(2, "Blank", "<", blocks=15, longest=36, says=_says(wrapped)) == AGENDA
    # And the line that nearly made it an agenda for the wrong reason: body copy is not
    # what a page calls itself, so a long line is not read for role words.
    assert "Review of the work content" not in _says(listed)
    assert _role(4, "Title Only", "Something", blocks=9, longest=46, says=_says(listed[:2])) == ""


def _unit_page(tmp_path: Path, name: str, build) -> Path:
    """A page of repeated units laid out by `build(group, index)`, saved and returned."""
    from pptx import Presentation
    from pptx.util import Inches

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    page = built.slides.add_slide(built.slide_layouts[6])
    for index in range(4):
        build(page.shapes.add_group_shape(), index)
    path = tmp_path / name
    built.save(str(path))
    return path


def _stat_card(group, index: int) -> None:
    """A card whose number is drawn last: heading, body, then the 60pt figure above them."""
    from pptx.util import Inches, Pt

    left = Inches(0.9 + index * 3.0)
    heading = group.shapes.add_textbox(left, Inches(4.5), Inches(2.6), Inches(0.4))
    heading.text_frame.text = "单击添加小标题"
    body = group.shapes.add_textbox(left, Inches(5.3), Inches(2.6), Inches(1.0))
    body.text_frame.text = "单击此处添加文本"
    figure = group.shapes.add_textbox(left, Inches(2.9), Inches(2.6), Inches(1.2))
    figure.text_frame.text = "68%"
    figure.text_frame.paragraphs[0].runs[0].font.size = Pt(60)


def test_items_are_counted_in_reading_order_not_file_order(tmp_path: Path) -> None:
    """The number box a reader sees first is stored last on the reference page this
    was measured on, and a positional item counted the way a reader counts put the body
    copy into the 60pt figure box and the figure into the heading slot."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    path = _unit_page(tmp_path, "stats.pptx", _stat_card)
    source, out = Presentation(str(path)), Presentation(str(path))

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["72%", "过夜访客", "约七成过夜。"]] * 4)

    first = min(units(slide)[0], key=lambda unit: unit.left)
    by_top = sorted((s.top, s.text_frame.text) for s in first.shapes if getattr(s, "has_text_frame", False))
    assert [text for _, text in by_top] == ["72%", "过夜访客", "约七成过夜。"]


def test_units_match_siblings_whatever_order_their_shapes_were_saved_in(tmp_path: Path) -> None:
    """A reference page's fourth row is the same four shapes as the three above it,
    saved icon-first; read as an ordered signature it was not a sibling, its texts were
    emptied with everyone else's and its icon was left standing beside nothing."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import units

    def row(group, index: int) -> None:
        top = Inches(1.6 + index * 1.2)
        parts = [
            lambda: group.shapes.add_textbox(Inches(6.5), top, Inches(3.0), Inches(0.4)),
            lambda: group.shapes.add_shape(9, Inches(5.9), top, Inches(0.5), Inches(0.5)),
        ]
        if index == 3:
            parts.reverse()
        for make in parts:
            shape = make()
            if getattr(shape, "has_text_frame", False) and shape.width > Inches(1):
                shape.text_frame.text = f"第 {index + 1} 条"

    path = _unit_page(tmp_path, "rows.pptx", row)
    runs = units(Presentation(str(path)).slides[0])

    assert [len(run) for run in runs] == [4]


def test_the_survivors_of_a_row_share_its_original_width(tmp_path: Path) -> None:
    """Three items on a four-card row used to leave three cards left-aligned with a
    card-sized hole on the right, which every author closed by hand or shipped."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    path = _unit_page(tmp_path, "row.pptx", _stat_card)
    source, out = Presentation(str(path)), Presentation(str(path))
    before = units(source.slides[0])[0]
    left0 = min(unit.left for unit in before)
    right0 = max(unit.left + unit.width for unit in before)

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["1", "甲", "一"], ["2", "乙", "二"], ["3", "丙", "三"]])

    after = sorted(units(slide)[0], key=lambda unit: unit.left)
    assert len(after) == 3
    assert after[0].left == left0, "the first card keeps the row's left edge"
    assert abs((after[-1].left + after[-1].width) - right0) < 12700, "the last card reaches the row's right edge"
    gaps = [after[i + 1].left - (after[i].left + after[i].width) for i in range(2)]
    assert abs(gaps[0] - gaps[1]) < 12700, "the survivors are evenly spaced"


def test_a_short_item_skips_the_number_that_reads_first(tmp_path: Path) -> None:
    """`["甲"]` on a `[01, label]` unit keeps the 01 and writes 甲 into the label, whichever
    of the two the file stored first -- the failure this closes wrote 甲 over the 01."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, fill, units

    def numbered(group, index: int) -> None:
        number = group.shapes.add_textbox(Inches(0.7 + index * 3.0), Inches(2.0), Inches(0.6), Inches(0.6))
        number.text_frame.text = f"0{index + 1}"
        label = group.shapes.add_textbox(Inches(1.4 + index * 3.0), Inches(2.05), Inches(2.0), Inches(0.5))
        label.text_frame.text = "单击添加小标题"

    path = _unit_page(tmp_path, "numbered.pptx", numbered)
    source, out = Presentation(str(path)), Presentation(str(path))

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["甲"], ["乙"]])

    said = _texts(slide)
    assert {"甲", "乙", "01", "02"} <= said
    assert "03" not in said and "单击添加小标题" not in said


def test_a_full_item_addresses_every_shape_including_the_number(tmp_path: Path) -> None:
    """The author who wants their own numbering writes one value per shape."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, fill, units

    def numbered(group, index: int) -> None:
        number = group.shapes.add_textbox(Inches(0.7 + index * 3.0), Inches(2.0), Inches(0.6), Inches(0.6))
        number.text_frame.text = f"0{index + 1}"
        label = group.shapes.add_textbox(Inches(1.4 + index * 3.0), Inches(2.05), Inches(2.0), Inches(0.5))
        label.text_frame.text = "单击添加小标题"

    path = _unit_page(tmp_path, "quarters.pptx", numbered)
    source, out = Presentation(str(path)), Presentation(str(path))

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["Q1", "甲"], ["Q2", "乙"]])

    assert {"Q1", "Q2", "甲", "乙"} <= _texts(slide)
    assert not {"01", "02"} & _texts(slide)


def test_an_irregular_run_is_not_moved(tmp_path: Path) -> None:
    """Pills along a path: deleting the spare must not re-space the rest, because any
    move is a guess at the design."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, fill, units

    spots = [(1.0, 4.5), (4.0, 6.0), (7.0, 2.0), (10.0, 3.5)]

    def pill(group, index: int) -> None:
        left, top = spots[index]
        shape = group.shapes.add_shape(1, Inches(left), Inches(top), Inches(1.8), Inches(0.5))
        shape.text_frame.text = "单击添加小标题"

    path = _unit_page(tmp_path, "path.pptx", pill)
    source, out = Presentation(str(path)), Presentation(str(path))

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["甲"], ["乙"], ["丙"]])

    kept = sorted((round(u.left / 914400, 2), round(u.top / 914400, 2)) for u in units(slide)[0])
    assert kept == sorted(spots[i] for i in (2, 0, 3)) or len(kept) == 3
    assert all(spot in [(l, t) for l, t in spots] for spot in kept), "every survivor stands where the template put it"


def test_bundled_opens_a_template_by_stem_and_names_the_rest_when_wrong(tmp_path: Path, monkeypatch) -> None:
    from pptx import Presentation

    from raven_ppt.services.template.compose import bundled

    built = Presentation()
    built.slides.add_slide(built.slide_layouts[6])
    (tmp_path / "shelf").mkdir()
    built.save(str(tmp_path / "shelf" / "one_template.pptx"))
    monkeypatch.setenv("PPT_BUNDLED_TEMPLATES", str(tmp_path / "shelf"))

    assert len(bundled("one_template").slides) == 1
    assert len(bundled("one_template.pptx").slides) == 1, "the suffix is forgiven"
    with pytest.raises(FileNotFoundError, match="one_template"):
        bundled("two_template")
    monkeypatch.delenv("PPT_BUNDLED_TEMPLATES")
    with pytest.raises(RuntimeError, match="PPT_BUNDLED_TEMPLATES"):
        bundled("one_template")


def test_a_page_borrowed_across_templates_lands_on_the_deck_s_own_layout_and_theme(tmp_path: Path) -> None:
    """Measured on four such clones rendered beside their sources: the arrangement comes
    across, the colours and the master are the deck's. Here the two facts that make
    that so -- the layout is matched by name inside the deck's package, and a scheme
    colour is left a scheme colour -- and the file holds each layout part once."""
    import zipfile

    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.util import Inches

    from raven_ppt.services.template.compose import clone_page, replace_text

    def deck(path: Path, theme_hint: str):
        built = Presentation()
        built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
        page = built.slides.add_slide(built.slide_layouts[5])
        card = page.shapes.add_shape(1, Inches(1), Inches(2), Inches(3), Inches(2))
        card.fill.solid()
        card.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        card.text_frame.text = theme_hint
        built.save(str(path))
        return path

    lender = Presentation(str(deck(tmp_path / "lender.pptx", "借来的卡")))
    target = Presentation(str(deck(tmp_path / "target.pptx", "自己的卡")))

    slide = clone_page(target, lender.slides[0])
    replace_text(slide, "借来的卡", "换了字")
    target.save(str(tmp_path / "out.pptx"))

    assert slide.slide_layout.name == lender.slides[0].slide_layout.name
    assert slide.slide_layout.part.package is target.part.package, "the deck's own layout, not the lender's"
    card = next(s for s in slide.shapes if getattr(s, "has_text_frame", False) and s.text_frame.text == "换了字")
    assert card.fill.fore_color.theme_color == MSO_THEME_COLOR.ACCENT_1, "a theme colour stays one, so it re-themes"
    names = [n for n in zipfile.ZipFile(tmp_path / "out.pptx").namelist() if "slideLayout" in n and n.endswith(".xml")]
    assert len(names) == len(set(names)), "no layout part written twice"
    assert not isinstance(RGBColor, str)


def test_replace_text_takes_a_list_of_strings_and_nested_lists_as_paragraphs(tmp_path: Path) -> None:
    """`[["选址评估"], ["六维模型"]]` crashed a live build with `'list' object has no
    attribute 'replace'`: two paragraphs, each a list holding one plain string."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_text

    built = Presentation()
    slide = built.slides.add_slide(built.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    box.text_frame.text = "旧"

    replace_text(box, [["选址评估"], ["六维模型"]])
    assert [p.text for p in box.text_frame.paragraphs] == ["选址评估", "六维模型"]

    replace_text(box, ["第一段", "第二段\n第三段"])
    assert [p.text for p in box.text_frame.paragraphs] == ["第一段", "第二段", "第三段"]

    with pytest.raises(TypeError, match="each paragraph"):
        replace_text(box, [{"not": "a paragraph"}])


def test_replace_text_takes_the_shape_whose_whole_copy_is_the_key() -> None:
    """A delivered page lost its title to a label one group down.

    `_all_shapes` reaches a group's members before the top-level shape drawn over
    them, so the first match for `自动化与人工智能` was an arrow label reading
    `自动化与人工智能助力产业升级`, and the title -- that string and nothing else, last
    in the z-order because it sits on top -- was never written. Replaying the program
    that shipped it, two of its 48 copy-route calls were landing this way.
    """
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_text

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    slide = built.slides.add_slide(built.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    label = group.shapes.add_textbox(Inches(9.5), Inches(3.2), Inches(3.0), Inches(0.5))
    label.text_frame.text = "自动化与人工智能助力产业升级"
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.1), Inches(8.0), Inches(0.8))
    title.text_frame.text = "自动化与人工智能"

    replace_text(slide, "自动化与人工智能", "System prices fell 31% in one year")

    assert title.text_frame.text == "System prices fell 31% in one year"
    assert label.text_frame.text == "自动化与人工智能助力产业升级"


def test_replace_text_still_names_a_block_by_its_opening_words() -> None:
    """Naming a block by its opening words is the contract, and it is unchanged.

    Only the exact hit is promoted. With no shape whose whole copy is the key, the
    page's own order decides as it always has -- so a key that is nobody's whole
    text still reaches the first shape carrying it.
    """
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_text

    built = Presentation()
    slide = built.slides.add_slide(built.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    first = group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(0.5))
    first.text_frame.text = "自动化与人工智能助力产业升级"
    second = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    second.text_frame.text = "自动化与人工智能的未来"

    replace_text(slide, "自动化与人工", "the opening words still name a block")

    assert first.text_frame.text == "the opening words still name a block"
    assert second.text_frame.text == "自动化与人工智能的未来"


def test_replace_text_refuses_a_key_no_shape_holds() -> None:
    """The refusal still lists the page's copy, which is what teaches the contract."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_text

    built = Presentation()
    slide = built.slides.add_slide(built.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    box.text_frame.text = "自动化与人工智能"

    with pytest.raises(KeyError, match="no text on this page matches"):
        replace_text(slide, "人工智能与自动化", "never written")


def test_place_takes_a_page_box_for_a_shape_inside_a_scaled_group(tmp_path: Path) -> None:
    """A group's children keep their numbers in the group's child space. A live
    program found `place` equivalent to assigning `.top` and wrote its own
    conversion; the box an author gives is on the page, so it is converted here."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import page_position, place

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    slide = built.slides.add_slide(built.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_shape(1, Inches(1.0), Inches(1.0), Inches(2.0), Inches(1.0))
    # Shrink the group on the page to half its child extent, and move it: the
    # child's numbers now mean something else on the page.
    group.left, group.top = Inches(4.0), Inches(3.0)
    group.width, group.height = Inches(1.0), Inches(0.5)

    place(child, (5.0, 3.5, 0.5, 0.25))

    x, y = page_position(child)
    assert (round(x, 3), round(y, 3)) == (5.0, 3.5), "the child sits where the page box said"
    assert child.width == Inches(1.0) and child.height == Inches(0.5), "its size is scaled into the group's space"


def _ids_on(slide) -> list[str]:
    return [
        node.get("id")
        for node in slide._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    ]


def test_more_items_than_slots_grows_a_row_within_its_width(tmp_path: Path) -> None:
    """Ten measured builds died on 'this page repeats N units and N+1 items were given',
    and the authors then wrote their own clone helper. Five items on a four-card row are
    five cards across the same width, shrunk alike, and the fifth is filled."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    path = _unit_page(tmp_path, "grow.pptx", _stat_card)
    source, out = Presentation(str(path)), Presentation(str(path))
    before = units(source.slides[0])[0]
    left0 = min(u.left for u in before)
    right0 = max(u.left + u.width for u in before)
    width0 = before[0].width

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [[f"{n}0%", f"第 {n} 项", "正文"] for n in range(1, 6)])

    grown = sorted(units(slide)[0], key=lambda u: u.left)
    assert len(grown) == 5
    assert grown[0].left == left0 and abs((grown[-1].left + grown[-1].width) - right0) < 12700
    assert all(u.width == grown[0].width for u in grown), "shrunk alike"
    assert grown[0].width < width0
    assert {"10%", "50%", "第 5 项"} <= _texts(slide), "the fifth unit is a slot like the others"
    ids = _ids_on(slide)
    assert len(ids) == len(set(ids)), "no two shapes share an id"


def test_add_unit_on_a_column_grows_it_down_its_own_extent(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import add_unit, units

    def row(group, index: int) -> None:
        top = Inches(1.6 + index * 1.3)
        badge = group.shapes.add_shape(9, Inches(6.0), top, Inches(0.5), Inches(0.5))
        badge.text_frame.text = f"0{index + 1}"
        label = group.shapes.add_textbox(Inches(6.7), top, Inches(4.0), Inches(0.5))
        label.text_frame.text = "单击添加小标题"

    path = _unit_page(tmp_path, "column.pptx", row)
    deck = Presentation(str(path))
    slide = deck.slides[0]
    run = units(slide)[0]
    top0 = min(u.top for u in run)
    bottom0 = max(u.top + u.height for u in run)

    added = add_unit(run, 1)

    assert len(added) == 1
    grown = sorted(units(slide)[0], key=lambda u: u.top)
    assert len(grown) == 5
    assert grown[0].top == top0 and abs((grown[-1].top + grown[-1].height) - bottom0) < 12700
    gaps = {round((grown[i + 1].top - grown[i].top - grown[i].height) / 914400, 2) for i in range(4)}
    assert len(gaps) == 1, f"even gaps, not {gaps}"


def test_a_grid_gains_a_row_at_its_own_pitch(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import add_unit, units

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    page = built.slides.add_slide(built.slide_layouts[6])
    for index in range(6):
        group = page.shapes.add_group_shape()
        col, row = index % 3, index // 3
        card = group.shapes.add_shape(1, Inches(0.7 + col * 4.1), Inches(1.5 + row * 1.8), Inches(3.8), Inches(1.5))
        card.text_frame.text = f"卡 {index + 1}"
    built.save(str(tmp_path / "grid.pptx"))
    deck = Presentation(str(tmp_path / "grid.pptx"))
    slide = deck.slides[0]

    added = add_unit(units(slide)[0], 2)

    tops = sorted({round(u.top / 914400, 2) for u in units(slide)[0]})
    assert tops == [1.5, 3.3, 5.1], "a third row at the grid's own pitch"
    lefts_last_row = sorted(round(u.left / 914400, 2) for u in added)
    assert lefts_last_row == [2.75, 6.85], "the short last row is centred"


def test_growth_refuses_where_a_unit_could_no_longer_carry_copy(tmp_path: Path) -> None:
    from pptx import Presentation

    from raven_ppt.services.template import add_unit, units

    path = _unit_page(tmp_path, "toomany.pptx", _stat_card)
    slide = Presentation(str(path)).slides[0]

    with pytest.raises(ValueError, match="slots"):
        add_unit(units(slide)[0], 6)


def test_an_irregular_run_is_not_grown(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, fill, units

    spots = [(1.0, 4.5), (4.0, 6.0), (7.0, 2.0), (10.0, 3.5)]

    def pill(group, index: int) -> None:
        left, top = spots[index]
        shape = group.shapes.add_shape(1, Inches(left), Inches(top), Inches(1.8), Inches(0.5))
        shape.text_frame.text = "单击添加小标题"

    path = _unit_page(tmp_path, "curve.pptx", pill)
    source, out = Presentation(str(path)), Presentation(str(path))
    slide = clone_page(out, source.slides[0])

    with pytest.raises(ValueError, match="no row, column or grid") as refused:
        fill(max(units(slide), key=len), [["甲"], ["乙"], ["丙"], ["丁"], ["戊"]])
    # The way out is named, with the geometry to take it: two measured runs met this
    # refusal on a template's diagonal pair and had only "another prototype" to go on.
    assert "clone_shape(run[-1], (left, top, width, height))" in str(refused.value)
    assert "(1.00, 4.50, 1.80, 0.50)" in str(refused.value)


def test_remove_unit_closes_the_gap_and_clone_shape_lands_at_its_box(tmp_path: Path) -> None:
    from pptx import Presentation

    from raven_ppt.services.template import clone_shape, remove_unit, units
    from raven_ppt.services.template.compose import page_position

    path = _unit_page(tmp_path, "remove.pptx", _stat_card)
    slide = Presentation(str(path)).slides[0]
    run = sorted(units(slide)[0], key=lambda u: u.left)
    left0, right0 = run[0].left, run[-1].left + run[-1].width

    remove_unit(run[1])

    left = sorted(units(slide)[0], key=lambda u: u.left)
    assert len(left) == 3
    assert left[0].left == left0 and abs((left[-1].left + left[-1].width) - right0) < 12700

    copy_ = clone_shape(left[0], (0.5, 6.5, 2.0, 0.6))
    assert (round(page_position(copy_)[0], 2), round(page_position(copy_)[1], 2)) == (0.5, 6.5)
    ids = _ids_on(slide)
    assert len(ids) == len(set(ids))


# --- a picture behind the page, and the pictures a layout carries -----------------------


def _canvas(tmp_path: Path, picture_size=(800, 600), shade=(20, 20, 20)):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image = tmp_path / "wash.png"
    Image.new("RGB", picture_size, shade).save(image)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "the title"
    return presentation, slide, image


def test_backdrop_sits_behind_the_page_cover_cropped_and_washed(tmp_path: Path) -> None:
    """The one generated picture that never poses as evidence: full-bleed, first in the
    z-order so the title already there stays over it, cropped rather than stretched to
    the canvas, and washed through its own `alphaModFix` so the page's ground shows."""
    from raven_ppt.services.template.compose import backdrop

    presentation, slide, image = _canvas(tmp_path)
    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"

    picture = backdrop(slide, image, alpha=0.3, scrim=None)

    tree = slide.shapes._spTree
    assert list(tree).index(picture._element) == 2, "behind everything: first drawn shape in the tree"
    assert picture.name == "backdrop"
    assert (picture.left, picture.top) == (0, 0)
    assert (picture.width, picture.height) == (presentation.slide_width, presentation.slide_height)
    assert picture.crop_top > 0 and picture.crop_bottom > 0 and picture.crop_left == 0, (
        "a 4:3 picture behind a 16:9 page loses its top and bottom, not its proportions"
    )
    fix = picture._element.find(f".//{{{namespace}}}blip/{{{namespace}}}alphaModFix")
    assert fix is not None and fix.get("amt") == "30000"
    assert slide.shapes[-1].text_frame.text == "the title", "the page's own shapes are untouched"


def test_backdrop_by_default_dims_the_photograph_under_a_plane_of_ink_and_sets_the_type_light(tmp_path: Path) -> None:
    """The form a reference cover uses: the photograph at full strength, a plane of the
    theme's ink over it, white type -- against a live cover whose photograph was washed to
    30% on a white page under black type and read as fog. A run already in a saturated
    accent keeps its colour: it is the one thing the page says with colour."""
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    from raven_ppt.services.template.compose import backdrop

    presentation, slide, image = _canvas(tmp_path)
    kicker = slide.shapes.add_textbox(Inches(1), Inches(0.4), Inches(4), Inches(0.5))
    kicker.text_frame.text = "01 / 06"
    kicker.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0x6A, 0x00)
    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"

    picture = backdrop(slide, image)

    tree = list(slide.shapes._spTree)
    assert tree.index(picture._element) == 2
    fix = picture._element.find(f".//{{{namespace}}}blip/{{{namespace}}}alphaModFix")
    assert fix is not None and fix.get("amt") == "100000", "the photograph itself is at full strength"
    plane = slide.shapes[1]
    assert tree.index(plane._element) == 3 and plane.name == "backdrop scrim"
    assert (plane.width, plane.height) == (presentation.slide_width, presentation.slide_height)
    colour = plane._element.spPr.find(f"{{{namespace}}}solidFill/{{{namespace}}}srgbClr")
    assert colour.get("val") == "000000", "the theme's dk1 is the ink"
    assert colour.find(f"{{{namespace}}}alpha").get("val") == "62000"
    assert plane.line.fill.type is None or plane.line.fill.type == 5, "no outline"  # 5 = MSO_FILL.BACKGROUND
    title_run = slide.shapes[-2].text_frame.paragraphs[0].runs[0]
    assert str(title_run.font.color.rgb) == "FFFFFF", "the title is set to the theme's lt1"
    kicker_run = slide.shapes[-1].text_frame.paragraphs[0].runs[0]
    assert str(kicker_run.font.color.rgb) == "FF6A00", "a saturated accent stays"


def test_drop_shape_refuses_a_layout_shape_and_names_the_way_round(tmp_path: Path) -> None:
    """A live cover deleted the layout's whole design group to make room for a washed
    photograph; the layout is the template's and shared by every page on it."""
    from raven_ppt.services.template.compose import drop_shape

    presentation, slide, _ = _canvas(tmp_path)
    on_layout = presentation.slide_layouts[0].shapes[0]

    with pytest.raises(ValueError, match="on a layout, not on the page"):
        drop_shape(on_layout)
    assert len(presentation.slide_layouts[0].shapes) >= 1, "nothing was removed"
    drop_shape(slide.shapes[-1])
    assert len(slide.shapes) == 0, "a page's own shape still goes"


def _region_page(tmp_path: Path, name: str = "region.pptx"):
    """A page shaped like the one the live failure ran on.

    Every element the author's own two sweeps kept: an arrow under the height bar it
    tested, a number label that is not Chinese, a horizontal connector with no height
    at all, and a card that the region sits inside rather than the other way round.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    page.shapes.add_textbox(Inches(0.7), Inches(0.2), Inches(11), Inches(0.8)).text_frame.text = "The page title"
    card = page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(1.5), Inches(8), Inches(4.5))
    arrow = page.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3), Inches(3), Inches(1.27), Inches(0.68))
    label = page.shapes.add_textbox(Inches(4), Inches(3), Inches(1.32), Inches(0.5))
    label.text_frame.text = "01"
    rule = page.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(4), Inches(8), Inches(4))
    outside = page.shapes.add_textbox(Inches(0.2), Inches(6.6), Inches(4), Inches(0.4))
    outside.text_frame.text = "the source line"
    path = tmp_path / name
    presentation.save(str(path))
    return presentation, page, {"card": card, "arrow": arrow, "label": label, "rule": rule, "outside": outside}


def test_clear_region_takes_what_a_hand_written_sweep_kept(tmp_path: Path) -> None:
    """A live program had no way to say "clear this box", so it swept the page by content
    type instead: one sweep on whether the text was Chinese, one on whether the shape was
    over 1.0x0.7in. Against the page it ran on that kept every arrow, every number label
    and -- a rule having one zero dimension -- every connector, and it drew two charts on
    top of them."""
    from raven_ppt.services.template.compose import clear_region

    _, page, made = _region_page(tmp_path)
    before = len(list(page.shapes))

    cleared = clear_region(page, _layout_box(2.5, 2.5, 9.5, 5.5))

    said = "\n".join(cleared)
    assert "a shape" in said or "text" in said, said
    assert any("'01'" in line for line in cleared), "the label a Chinese-text sweep kept"
    assert any("a connector" in line for line in cleared), "a rule has one zero dimension"
    assert len(cleared) == 3, cleared
    assert len(list(page.shapes)) == before - 3
    assert made["outside"]._element.getparent() is not None, "a shape outside the box stays"


def test_clear_region_keeps_the_card_the_box_sits_in(tmp_path: Path) -> None:
    """On a template the chart is drawn inside a card, and the card is the arrangement
    the page was cloned for. Measured over the region a chart occupies on five template
    pages: the card covers 0.97 to 1.00 of the box while the box covers only 0.42 to 0.69
    of the card, where the chart covers 1.00 of both."""
    from raven_ppt.services.template.compose import clear_region

    _, page, made = _region_page(tmp_path)

    cleared = clear_region(page, _layout_box(2.5, 2.5, 9.5, 5.5))

    assert made["card"]._element.getparent() is not None, "the card the box sits inside stays"
    assert any("at (2.00, 1.50)" in line for line in cleared.left_standing), cleared.left_standing


def test_clear_region_says_what_it_left_over_the_box(tmp_path: Path) -> None:
    """The sweep this replaces failed invisibly: the author read the render three builds
    later. Anything still lying over the box is named in the return value and warned
    about, because the share that decides it is one an author may have to lower."""
    from raven_ppt.services.template.compose import clear_region

    _, page, _made = _region_page(tmp_path)

    with pytest.warns(UserWarning, match="left 1 shape"):
        cleared = clear_region(page, _layout_box(2.5, 2.5, 9.5, 5.5))

    assert len(cleared.left_standing) == 1
    assert "Still over it" in str(cleared)
    assert str(cleared).startswith("cleared (2.5, 2.5)-(9.5, 5.5): removed 3")


def test_clear_region_spares_what_you_name(tmp_path: Path) -> None:
    """The author's own replaced copy can sit inside the box it wants cleared under it."""
    from raven_ppt.services.template.compose import clear_region

    _, page, made = _region_page(tmp_path)

    cleared = clear_region(page, _layout_box(2.5, 2.5, 9.5, 5.5), keep=("01",))

    assert made["label"]._element.getparent() is not None
    assert any("kept, you named it" in line for line in cleared.left_standing), cleared.left_standing
    assert len(cleared) == 2


def test_shapes_in_answers_by_place_where_the_others_answer_by_number_or_words(tmp_path: Path) -> None:
    """`shape_at` takes an ordinal and `shape_saying` takes text; a runtime count over the
    live program shows drop_shape 30 times and shape_saying 33, with shape_at never -- it
    needed to find by place and hand-rolled a walker instead. `share=0` asks the other
    question, everything that touches the box at all."""
    from raven_ppt.services.template.compose import page_box, shapes_in

    _, page, made = _region_page(tmp_path)

    inside = shapes_in(page, _layout_box(2.5, 2.5, 9.5, 5.5))
    assert made["arrow"] in inside and made["label"] in inside and made["rule"] in inside
    assert made["card"] not in inside, "the box sits in the card"
    assert made["outside"] not in inside

    touching = shapes_in(page, _layout_box(2.5, 2.5, 9.5, 5.5), share=0)
    assert made["card"] in touching, "share=0 asks what touches the box at all"

    # And a shape's own box is a region, so clearing where something was composes.
    assert shapes_in(page, page_box(made["arrow"])) == [made["arrow"]]


def test_page_box_reports_the_size_a_group_scales_and_not_the_declared_one(tmp_path: Path) -> None:
    """A group scales what is inside it. Both hand-written sweeps read `shape.width`
    directly, so their size tests were about a rectangle that is not on the page."""
    from pptx.util import Inches

    from raven_ppt.services.template.compose import page_box

    _, page, _made = _region_page(tmp_path, "grouped.pptx")
    group = page.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    inner.text_frame.text = "inside a group that scales it"
    # python-pptx sizes a fresh group to its children, so off == chOff and the scale is
    # 1.0. Doubling the frame's extent against the child extent is what a template does.
    frame = group._element.grpSpPr.xfrm
    frame.off.x, frame.off.y = Inches(4), Inches(2)
    frame.ext.cx, frame.ext.cy = Inches(4), Inches(2)

    drawn = page_box(inner)

    assert (drawn.x0, drawn.y0) == pytest.approx((4.0, 2.0)), drawn
    assert (drawn.w, drawn.h) == pytest.approx((4.0, 2.0)), "the group scaled 2x1in to 4x2in"
    assert (inner.width / 914400, inner.height / 914400) == pytest.approx((2.0, 1.0)), "declared, not drawn"
    assert (drawn.x1, drawn.y1) == pytest.approx((drawn.x0 + drawn.w, drawn.y0 + drawn.h))


def test_a_region_refuses_four_numbers_that_cannot_say_which_reading_they_are(tmp_path: Path) -> None:
    """`ppt_layout.Box` is two corners and every python-pptx call in the same script is a
    corner and a size, and four bare numbers are both. This used to guess, reading a bare
    tuple as corners unless its last two numbers were smaller than its first two, on the
    theory that a size always is; on a 13.33x7.5in canvas most sizes are larger than most
    origins, so the reading it treated as the exception was the common one."""
    from raven_ppt.services.template.compose import _as_region

    assert _as_region(_layout_box(1.0, 2.0, 5.0, 6.0)) == (1.0, 2.0, 5.0, 6.0)
    assert _as_region(_layout_box_at(1.0, 2.0, 4.0, 4.0)) == (1.0, 2.0, 5.0, 6.0)

    class Boxish:
        x0, y0, x1, y1 = 0.5, 0.5, 3.0, 3.0

    assert _as_region(Boxish()) == (0.5, 0.5, 3.0, 3.0)

    # The numbers the guess got wrong: as a size this is the region (1,1)-(5,4), and it
    # became (1,1)-(4,3) with nothing said. Both readings are named, because the whole
    # point is that the caller is the only one who knows which was meant.
    with pytest.raises(ValueError, match="says which of its two readings") as refused:
        _as_region((1.0, 1.0, 4.0, 3.0))
    assert "(1, 1) to (4, 3)" in str(refused.value), "the corner reading"
    assert "(1, 1) to (5, 4)" in str(refused.value), "the size reading"
    assert "Box.at(x, y, w=, h=)" in str(refused.value), "and a next step that runs"

    with pytest.raises(ValueError, match="says which of its two readings"):
        _as_region((1.0, 2.0, 3.0))


def test_clear_region_refuses_the_size_a_reference_prints_rather_than_emptying_the_wrong_box(
    tmp_path: Path,
) -> None:
    """The defect the guess left behind, on the page the live failure ran on. A template
    reference prints `(left, top, width, height)`, so the author of that region holds
    `(2.5, 2.5, 7, 3)` for the box (2.5, 2.5)-(9.5, 5.5). Read as corners that is
    (2.5, 2.5)-(7, 3), which contains none of the furniture: the call reported "nothing
    was in it", left the arrow, the label and the connector standing, named none of
    them, and the author drew two charts on top of them."""
    from raven_ppt.services.template.compose import clear_region

    _, page, made = _region_page(tmp_path)

    with pytest.raises(ValueError, match="says which of its two readings") as refused:
        clear_region(page, (2.5, 2.5, 7, 3))
    assert "(2.5, 2.5) to (7, 3)" in str(refused.value), "the corner reading, which emptied nothing"
    assert "(2.5, 2.5) to (9.5, 5.5)" in str(refused.value), "the region the author had"
    for what in ("arrow", "label", "rule"):
        assert made[what]._element.getparent() is not None, f"{what} is untouched by a refusal"

    # Said as a size, the same four numbers empty the box the author meant.
    cleared = clear_region(page, _layout_box_at(2.5, 2.5, 7, 3))

    assert str(cleared).startswith("cleared (2.5, 2.5)-(9.5, 5.5): removed 3")
    for what in ("arrow", "label", "rule"):
        assert made[what]._element.getparent() is None, f"{what} went"
    assert made["card"]._element.getparent() is not None, "the card the box sits inside stays"


def test_backdrop_takes_a_box_and_refuses_an_alpha_that_is_not_a_wash(tmp_path: Path) -> None:
    from pptx.util import Inches

    from raven_ppt.services.template.compose import backdrop

    _, slide, image = _canvas(tmp_path)

    picture = backdrop(slide, image, alpha=1.0, box=(0, 3.5, 13.333, 4))
    assert (picture.left, picture.top) == (0, Inches(3.5))
    assert picture.height == Inches(4)

    with pytest.raises(ValueError, match="alpha=0 "):
        backdrop(slide, image, alpha=0)
    with pytest.raises(ValueError, match="alpha=1.5"):
        backdrop(slide, image, alpha=1.5)
    with pytest.raises(ValueError, match="not one"):
        backdrop(slide, image.with_name("missing.png"))


def test_backdrop_in_a_box_lightens_only_the_type_the_plane_covers(tmp_path: Path) -> None:
    """A boxed photograph across the lower half of the page turned a title at y=1 white,
    with nothing dark behind it: the type on the page's own light ground would vanish.
    The plane is where the page went dark, and only the type it covers follows."""
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    from raven_ppt.services.template.compose import backdrop

    _, slide, image = _canvas(tmp_path)
    title = slide.shapes[0]
    covered = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    covered.text_frame.text = "a caption over the photograph"
    straddling = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    straddling.text_frame.text = "a line the plane's edge crosses"
    kicker = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(0.5))
    kicker.text_frame.text = "01 / 06"
    kicker.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0x6A, 0x00)

    backdrop(slide, image, box=(0, 3.5, 13.333, 4))

    def colour_of(shape):
        run = shape.text_frame.paragraphs[0].runs[0]
        return str(run.font.color.rgb) if run.font.color.type is not None else None

    assert colour_of(title) is None, "the title at y=1 is outside the plane and keeps the template's ink"
    assert colour_of(covered) == "FFFFFF"
    assert colour_of(straddling) == "FFFFFF", "a shape the plane reaches at all is over the dark part"
    assert colour_of(kicker) == "FF6A00", "a saturated accent stays"


def test_backdrop_in_a_box_reads_grouped_type_in_page_space(tmp_path: Path) -> None:
    """A group draws its children in a space of its own and places that space on the
    page, so a child's frame is not where it sits: a caption grouped at page (1, 5)
    carried child-space numbers off the page and stayed dark over a lower-half scrim,
    while a grouped title whose child space put it at y=5 sat at y=3 on the page and
    went white over the light ground. Both frames are carried through the group's
    transform before the test."""
    from pptx.util import Inches

    from raven_ppt.services.template.compose import backdrop

    _, slide, image = _canvas(tmp_path)
    captioned = slide.shapes.add_group_shape()
    caption = captioned.shapes.add_textbox(Inches(10), Inches(10), Inches(4), Inches(1))
    caption.text_frame.text = "a grouped caption over the photograph"
    # python-pptx sizes the group to its children (off == chOff): move the group to
    # page (1, 5) and leave its child space at (10, 10), where the plane is not.
    placed = captioned._element.grpSpPr.xfrm
    placed.off.x, placed.off.y = Inches(1), Inches(5)
    titled = slide.shapes.add_group_shape()
    title = titled.shapes.add_textbox(0, Inches(5), Inches(6), Inches(0.4))
    title.text_frame.text = "a grouped title above the photograph"
    # The group's frame reaches the plane (page 3..5); its child, at the top of a
    # child space two inches tall, lands at page 3..3.4, above the plane's 3.5.
    placed = titled._element.grpSpPr.xfrm
    placed.off.x, placed.off.y = Inches(1), Inches(3)
    placed.ext.cy = placed.chExt.cy = Inches(2)

    backdrop(slide, image, box=(0, 3.5, 13.333, 4))

    def colour_of(shape):
        run = shape.text_frame.paragraphs[0].runs[0]
        return str(run.font.color.rgb) if run.font.color.type is not None else None

    assert colour_of(caption) == "FFFFFF", "grouped, at page (1, 5): under the plane"
    assert colour_of(title) is None, "grouped, at page (1, 3): above the plane, on the light ground"


def test_backdrop_in_a_box_follows_a_group_that_is_mirrored_or_turned(tmp_path: Path) -> None:
    """A group mirrors (`flipH`) and turns (`rot`) its frame about its centre after placing
    its children, so a child's numbers can put it on the other side of the page from
    where it is drawn -- the beige template groups text this way. A mirrored group
    spanning x=1..7: its child at local x=0..1 is drawn at x=6..7, under a scrim at
    x=4..7.5, and its child at local x=5..6 is drawn at x=1..2, on the light ground.
    A group turned a quarter clockwise: the child along its bottom edge stands up its
    right side, above the plane; the child at its right edge lies along the bottom,
    under it."""
    from pptx.util import Inches

    from raven_ppt.services.template.compose import backdrop

    _, slide, image = _canvas(tmp_path)
    mirrored = slide.shapes.add_group_shape()
    drawn_right = mirrored.shapes.add_textbox(0, Inches(5), Inches(1), Inches(1))
    drawn_right.text_frame.text = "at local x=0, drawn at x=6"
    drawn_left = mirrored.shapes.add_textbox(Inches(5), Inches(5), Inches(1), Inches(1))
    drawn_left.text_frame.text = "at local x=5, drawn at x=1"
    frame = mirrored._element.grpSpPr.xfrm
    frame.off.x = Inches(1)
    frame.set("flipH", "1")
    turned = slide.shapes.add_group_shape()
    stood_up = turned.shapes.add_textbox(Inches(2), Inches(4.5), Inches(1), Inches(0.5))
    stood_up.text_frame.text = "along the bottom edge, turned onto the right side"
    laid_down = turned.shapes.add_textbox(Inches(5), Inches(3), Inches(1), Inches(0.5))
    laid_down.text_frame.text = "at the right edge, turned onto the bottom"
    frame = turned._element.grpSpPr.xfrm
    # The frame page (2, 3) to (6, 5), centre (4, 4), a quarter turn clockwise: a point
    # (x, y) in it goes to (8 - y, x).
    frame.off.x, frame.off.y = Inches(2), Inches(3)
    frame.ext.cx = frame.chExt.cx = Inches(4)
    frame.ext.cy = frame.chExt.cy = Inches(2)
    frame.chOff.x, frame.chOff.y = Inches(2), Inches(3)
    frame.set("rot", str(90 * 60000))

    backdrop(slide, image, box=(4, 3.5, 3.5, 4))

    def colour_of(shape):
        run = shape.text_frame.paragraphs[0].runs[0]
        return str(run.font.color.rgb) if run.font.color.type is not None else None

    assert colour_of(drawn_right) == "FFFFFF", "mirrored under the scrim"
    assert colour_of(drawn_left) is None, "mirrored onto the light ground"
    assert colour_of(laid_down) == "FFFFFF", "turned onto the bottom, page (4.5..5, 5..6), under the scrim"
    assert colour_of(stood_up) is None, "turned onto the right side, page (3..3.5, 2..3), above the plane"


def test_backdrop_in_a_box_visits_a_turned_group_whose_declared_frame_misses_the_plane(tmp_path: Path) -> None:
    """A group's declared frame is the untransformed one: turned a quarter, the group
    lies somewhere else on the page. Pruned on that frame before its turn was applied,
    a group at x=1..3, y=1..7 missed a scrim at x=3.5..5.5, y=2.5..5.5 and its top-edge
    child, drawn under the scrim at x=4..5, y=3..5, stayed dark. A group is never
    pruned on its own frame; its children are judged where the turn puts them."""
    from pptx.util import Inches

    from raven_ppt.services.template.compose import backdrop

    _, slide, image = _canvas(tmp_path)
    group = slide.shapes.add_group_shape()
    top_edge = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    top_edge.text_frame.text = "along the top edge, turned under the scrim"
    bottom_edge = group.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1))
    bottom_edge.text_frame.text = "along the bottom edge, turned onto the light ground"
    # The frame page (1, 1) to (3, 7), centre (2, 4), a quarter turn clockwise: a point
    # (x, y) in it goes to (6 - y, x + 2). The top-edge child lands at x=4..5, y=3..5,
    # the bottom-edge child at x=-1..0, y=3..5; the frame itself, untransformed, is
    # x=1..3 and never touches a scrim at x=3.5..5.5.
    frame = group._element.grpSpPr.xfrm
    frame.set("rot", str(90 * 60000))

    backdrop(slide, image, box=(3.5, 2.5, 2, 3))

    def colour_of(shape):
        run = shape.text_frame.paragraphs[0].runs[0]
        return str(run.font.color.rgb) if run.font.color.type is not None else None

    assert colour_of(top_edge) == "FFFFFF", "turned to x=4..5, y=3..5: under the scrim"
    assert colour_of(bottom_edge) is None, "turned to x=-1..0, y=3..5: off the plane"


def test_layout_pictures_names_the_layouts_photographs_and_replace_picture_swaps_them(tmp_path: Path) -> None:
    """A template's cover photograph is as often on the layout as on the page, where
    nothing on the page can be handed to `replace_picture`. These are those shapes, and
    the swap reaches every page built on the layout."""
    from PIL import Image

    from raven_ppt.services.template.compose import layout_pictures, replace_picture
    from tests._ppt_engine_fixtures import layout_picture

    presentation, slide, image = _canvas(tmp_path)
    rels = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    layout = slide.slide_layout
    small = layout_picture(layout, image, 1, 1, 2, 1.5)
    large = layout_picture(layout, image, 6, 0, 7.3, 7.5)
    other = presentation.slides.add_slide(layout)

    found = layout_pictures(slide)

    assert [shape.shape_id for shape in found] == [large.shape_id, small.shape_id], "largest first"
    assert layout_pictures(other)[0].shape_id == large.shape_id, "the same shapes, seen from any page on the layout"

    replacement = tmp_path / "new.png"
    Image.new("RGB", (1600, 900), (200, 200, 200)).save(replacement)
    before = large._element.blipFill.blip.get(f"{{{rels}}}embed")
    replace_picture(found[0], replacement, "cover")
    assert large._element.blipFill.blip.get(f"{{{rels}}}embed") != before
    assert large.crop_left > 0, "a 16:9 picture in a portrait frame is trimmed on its sides"


@pytest.mark.skipif(
    not __import__("raven_ppt.services.render", fromlist=["available"]).available().can_convert,
    reason="needs LibreOffice to render",
)
def test_a_washed_backdrop_renders_as_a_blend_not_a_slab(tmp_path: Path) -> None:
    """`alphaModFix` is the transparency the renderer honours: a black picture at 0.3
    over a white page has to come out grey, or the wash exists only in the XML."""
    import asyncio

    from PIL import Image

    from raven_ppt.services.render import available
    from raven_ppt.services.template.compose import backdrop
    from raven_ppt.stages._views import DeckViews

    if not available().can_rasterise:
        pytest.skip("needs a PDF rasteriser")
    presentation, slide, image = _canvas(tmp_path, shade=(0, 0, 0))
    backdrop(slide, image, alpha=0.3, scrim=None)
    deck = tmp_path / "deck.pptx"
    presentation.save(str(deck))

    views = DeckViews(dpi=48)
    rendered = asyncio.run(views.pages(deck, tmp_path / "render", [1]))
    page = Image.open(rendered[1]).convert("RGB")
    sample = page.getpixel((page.width // 2, int(page.height * 0.8)))
    assert 150 <= sample[0] <= 200, f"black at alpha 0.3 over white should read around 178, not {sample}"


def test_replace_picture_washes_the_new_picture_when_asked(tmp_path: Path) -> None:
    """A layout picture the size of the page is the page's background; a photograph
    swapped in at full strength drowns the type, so the swap takes the same `alpha`
    a backdrop does -- on a picture frame and on a picture-filled shape alike."""
    from lxml import etree
    from PIL import Image
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_picture

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    rels = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    _, slide, image = _canvas(tmp_path)
    replacement = tmp_path / "new.png"
    Image.new("RGB", (1600, 900), (200, 200, 200)).save(replacement)
    frame = slide.shapes.add_picture(str(image), Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    filled = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    _, relationship = filled.part.get_or_add_image_part(str(image))
    fill = etree.SubElement(filled._element.spPr, f"{{{namespace}}}blipFill")
    etree.SubElement(fill, f"{{{namespace}}}blip").set(f"{{{rels}}}embed", relationship)

    replace_picture(frame, replacement, "cover", alpha=0.25)
    replace_picture(filled, replacement, "cover", alpha=0.25)

    for shape in (frame, filled):
        fix = shape._element.find(f".//{{{namespace}}}blip/{{{namespace}}}alphaModFix")
        assert fix is not None and fix.get("amt") == "25000", shape.name
    replace_picture(frame, replacement, "cover", alpha=0.6)
    assert len(frame._element.findall(f".//{{{namespace}}}alphaModFix")) == 1, "a second wash replaces the first"
    replace_picture(frame, replacement, "cover")
    assert frame._element.find(f".//{{{namespace}}}alphaModFix").get("amt") == "60000", "None leaves the wash alone"
    with pytest.raises(ValueError, match="alpha=2"):
        replace_picture(frame, replacement, "cover", alpha=2)


def test_surplus_empty_values_are_dropped_rather_than_refused(tmp_path: Path) -> None:
    """`["", title, sub]` against a two-frame unit: the empty string is the placeholder an
    author writes for a number tile it means to leave alone, and it would have written
    nothing. A live build ended on the refusal over exactly that. Values that would
    have written something are still refused when they outnumber the frames."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, fill, units

    def two_lines(group, index: int) -> None:
        left = Inches(0.9 + index * 3.0)
        title = group.shapes.add_textbox(left, Inches(2.0), Inches(2.6), Inches(0.5))
        title.text_frame.text = "工作内容回顾"
        sub = group.shapes.add_textbox(left, Inches(2.6), Inches(2.6), Inches(0.4))
        sub.text_frame.text = "Review of the work content"

    path = _unit_page(tmp_path, "agenda.pptx", two_lines)
    source, out = Presentation(str(path)), Presentation(str(path))

    slide = clone_page(out, source.slides[0])
    fill(max(units(slide), key=len), [["", "标杆案例调研", "国内外四个样本"]] * 4)

    first = min(units(slide)[0], key=lambda unit: unit.left)
    texts = [
        s.text_frame.text for s in sorted(first.shapes, key=lambda s: s.top) if getattr(s, "has_text_frame", False)
    ]
    assert texts == ["标杆案例调研", "国内外四个样本"]

    again = clone_page(Presentation(str(path)), source.slides[0])
    with pytest.raises(ValueError, match="2 text shape\\(s\\) a unit on this page holds"):
        fill(max(units(again), key=len), [["甲", "乙", "丙"]] * 4)


def test_wash_sets_any_pictures_transparency_and_refuses_a_solid_fill(tmp_path: Path) -> None:
    """The share `backdrop` and `replace_picture(alpha=)` take, on a picture already on
    the page -- a frame, or a shape filled with one -- and a refusal for a shape that
    shows no picture, whose solid fill has a transparency of its own."""
    from lxml import etree
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    from raven_ppt.services.template.compose import wash

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    rels = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    _, slide, image = _canvas(tmp_path)
    frame = slide.shapes.add_picture(str(image), Inches(1), Inches(2), Inches(4), Inches(3))
    filled = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(2), Inches(4), Inches(3))
    _, relationship = filled.part.get_or_add_image_part(str(image))
    fill = etree.SubElement(filled._element.spPr, f"{{{namespace}}}blipFill")
    etree.SubElement(fill, f"{{{namespace}}}blip").set(f"{{{rels}}}embed", relationship)
    solid = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(6), Inches(2), Inches(1))

    assert wash(frame, 0.3) is frame
    wash(filled, 0.5)
    wash(frame, 0.6)

    assert frame._element.find(f".//{{{namespace}}}blip/{{{namespace}}}alphaModFix").get("amt") == "60000"
    assert len(frame._element.findall(f".//{{{namespace}}}alphaModFix")) == 1, "a second wash replaces the first"
    assert filled._element.find(f".//{{{namespace}}}blip/{{{namespace}}}alphaModFix").get("amt") == "50000"
    with pytest.raises(ValueError, match="shows none"):
        wash(solid, 0.3)
    with pytest.raises(ValueError, match="alpha=0 "):
        wash(frame, 0)


def test_a_shaped_picture_frame_is_covered_not_shrunk(tmp_path: Path) -> None:
    """A frame cut to a curve is the page's design. Contain shrank a 13.35in wave-edged
    frame to 7.56in and left the photograph in a plain rectangle beside the panel it
    was drawn to complete; a shaped frame takes cover unless the author says otherwise."""
    from PIL import Image
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_picture

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    _, slide, image = _canvas(tmp_path, picture_size=(1200, 800))
    frame = slide.shapes.add_picture(str(image), Inches(0), Inches(3.25), Inches(13.35), Inches(4.25))
    geometry = frame._element.find(f".//{{{namespace}}}prstGeom")
    geometry.set("prst", "wave")
    replacement = tmp_path / "photo.png"
    Image.new("RGB", (1600, 1000), (60, 60, 60)).save(replacement)

    with pytest.warns(UserWarning, match="shaped picture frame"):
        replace_picture(frame, replacement)

    assert frame.width == Inches(13.35) and frame.left == 0, "the frame keeps its size and place"
    assert frame.crop_left > 0 or frame.crop_top > 0, "the picture is cropped into it"

    plain = slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(4), Inches(3))
    replace_picture(plain, replacement)
    assert plain.height < Inches(3), "a plain rectangle still gives way to the picture under contain"


def _drawing_page(tmp_path: Path):
    """A page whose illustration is a group of freeforms with a loose star beside it, under a title."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    title.text_frame.text = "Where the market sits"
    cartoon = slide.shapes.add_group_shape()
    cartoon.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8), Inches(2), Inches(2), Inches(2))
    cartoon.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.5), Inches(4), Inches(1), Inches(2))
    cartoon.name = "cartoon"
    star = slide.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(10.2), Inches(2.2), Inches(1), Inches(1))
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(6), Inches(4))
    panel.text_frame.text = "A panel of copy the picture must not take"
    return presentation, slide, cartoon, star, panel


def test_replace_picture_puts_a_picture_where_a_drawn_illustration_was(tmp_path: Path, image) -> None:
    """A section page's cartoon is a group of freeforms, not a picture, and `replace_picture`
    used to refuse it ("no image to replace") -- so the template's illustration stayed on
    every borrowed page. The picture takes the drawing's box and depth, and the drawing goes."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from raven_ppt.services.template.compose import replace_picture

    presentation, slide, cartoon, star, panel = _drawing_page(tmp_path)
    tree = cartoon._element.getparent()
    depth = tree.index(cartoon._element)
    where = (cartoon.left, cartoon.top, cartoon.width, cartoon.height)
    figure = image("cartoon.png", (30, 120, 120))

    picture = replace_picture(cartoon, figure)
    out = tmp_path / "swapped.pptx"
    presentation.save(str(out))

    page = Presentation(str(out)).slides[0]
    kinds = [shape.shape_type for shape in page.shapes]
    assert MSO_SHAPE_TYPE.GROUP not in kinds, "the drawing is gone"
    assert kinds.count(MSO_SHAPE_TYPE.PICTURE) == 1
    assert tree.index(picture._element) == depth, "the picture sits where the drawing sat in the z-order"
    frame = next(shape for shape in page.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert frame.image.blob == figure.read_bytes()
    assert frame.left >= where[0] and frame.top >= where[1], "contain keeps the picture inside the drawing's box"
    assert frame.left + frame.width <= where[0] + where[2] + 1 and frame.top + frame.height <= where[1] + where[3] + 1
    assert frame.name == "cartoon"


def test_replace_picture_takes_a_group_member_with_its_whole_group(tmp_path: Path, image) -> None:
    """`shape_at` numbers a group's members too, so an author points at the oval inside
    the cartoon; a group is one drawing, and half a cartoon left behind is worse than none."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from raven_ppt.services.template.compose import replace_picture

    presentation, slide, cartoon, star, panel = _drawing_page(tmp_path)
    oval = cartoon.shapes[0]

    replace_picture(oval, image("cartoon.png", (30, 120, 120)))

    kinds = [shape.shape_type for shape in slide.shapes]
    assert MSO_SHAPE_TYPE.GROUP not in kinds and kinds.count(MSO_SHAPE_TYPE.PICTURE) == 1
    assert any(shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX for shape in slide.shapes), "the title is untouched"


def test_replace_picture_swaps_a_drawing_and_several_loose_shapes_for_one_picture(tmp_path: Path, image) -> None:
    """A key naming the drawn illustration on a page used to come back as missed, so the
    author fell back to drawing over it. `shape_at` numbers a group's members and not the
    group, so a member stands for its whole cartoon; a list of shapes names the loose
    parts of one drawing together, and the picture spans their union."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from raven_ppt.services.template.compose import _all_shapes, clone_page, replace_text, shape_at

    presentation, slide, cartoon, star, panel = _drawing_page(tmp_path)
    every = list(_all_shapes(slide.shapes))
    oval_index = every.index(cartoon.shapes[0]) + 1
    star_index = every.index(star) + 1
    figure = image("scene.png", (200, 120, 40))

    page = clone_page(presentation, slide)
    replace_text(page, "Where the market sits", "Night market")
    replace_picture([shape_at(page, oval_index), shape_at(page, star_index)], figure)

    kinds = [shape.shape_type for shape in page.shapes]
    assert MSO_SHAPE_TYPE.GROUP not in kinds and kinds.count(MSO_SHAPE_TYPE.PICTURE) == 1
    assert not any(shape.name == star.name for shape in page.shapes), "the loose star went with the group"
    frame = next(shape for shape in page.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert frame.left >= cartoon.left and frame.left + frame.width <= star.left + star.width + 1, (
        "the picture spans the union of the drawing and the star"
    )
    assert any("Night market" in shape.text_frame.text for shape in page.shapes if shape.has_text_frame)


def test_replace_picture_does_not_read_a_text_panel_as_an_illustration(tmp_path: Path, image) -> None:
    """A shape that holds copy is a miscount, not a drawing: the copy panel is refused
    rather than giving way to a picture."""
    import pytest

    from raven_ppt.services.template.compose import replace_picture

    _presentation, _slide, _cartoon, _star, panel = _drawing_page(tmp_path)

    with pytest.raises(ValueError, match="not an illustration"):
        replace_picture(panel, image("x.png", (1, 2, 3)))


def test_a_drawing_inside_a_card_takes_only_the_wordless_group_around_it(tmp_path: Path, image) -> None:
    """Template pages keep everything one group down -- the card, its icon, its copy in one
    group -- and climbing to the top swapped a whole page's content for one picture (the
    blue template's four-card page). The icon's own group goes; the card and its words stay."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_picture

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    content = slide.shapes.add_group_shape()
    card = content.shapes.add_group_shape()
    body = card.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(4))
    body.text_frame.text = "Sales up"
    icon = card.shapes.add_group_shape()
    icon.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.5), Inches(2.5), Inches(1), Inches(1))
    icon.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.8), Inches(2.8), Inches(0.4), Inches(0.4))
    icon.name = "icon"

    picture = replace_picture(icon.shapes[1], image("icon.png", (10, 10, 200)))

    assert picture.name == "icon"
    assert [shape.shape_type for shape in card.shapes] == [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PICTURE]
    assert card.shapes[0].text_frame.text == "Sales up"
    assert len(slide.shapes) == 1 and slide.shapes[0].shape_type == MSO_SHAPE_TYPE.GROUP


def test_the_menu_lists_picture_filled_shapes_and_drawings_as_picture_slots(tmp_path: Path, image) -> None:
    """The amber template draws every photograph as a rounded rectangle filled with one and
    a section page's cartoon as a group of freeforms; counting `p:pic` alone told an author
    those pages held no picture, and the frame it named was refused on three pages of
    one live run. Both are slots now, numbered as `shape_at` numbers; a card icon is not."""
    from lxml import etree
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    from raven_ppt.services.template.menu import menu

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    rels = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(8), Inches(1))
    title.text_frame.text = "Where the market sits"
    photo = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5), Inches(5.4), Inches(3.6))
    _, relationship = photo.part.get_or_add_image_part(str(image("photo.png", (90, 90, 90))))
    fill = etree.SubElement(photo._element.spPr, f"{{{namespace}}}blipFill")
    etree.SubElement(fill, f"{{{namespace}}}blip").set(f"{{{rels}}}embed", relationship)
    cartoon = slide.shapes.add_group_shape()
    cartoon.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7), Inches(2), Inches(2.4), Inches(2.4))
    cartoon.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.5), Inches(4.4), Inches(1.4), Inches(1.1))
    cartoon.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8), Inches(4.4), Inches(0.3), Inches(1.1))
    band = slide.shapes.add_group_shape()
    band.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.4), Inches(13.333), Inches(1.1))
    band.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(6.5), Inches(1.5), Inches(1))
    card = slide.shapes.add_group_shape()
    body = card.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10), Inches(2), Inches(2.9), Inches(4))
    body.text_frame.text = "Sales up"
    icon = card.shapes.add_group_shape()
    icon.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.2), Inches(2.2), Inches(0.7), Inches(0.7))
    icon.shapes.build_freeform(Inches(10.4), Inches(2.4)).add_line_segments(
        [(Inches(10.6), Inches(2.4)), (Inches(10.5), Inches(2.7))]
    ).convert_to_shape()
    path = tmp_path / "slots.pptx"
    presentation.save(str(path))

    (entry,) = menu(path)

    assert entry.pictures == 1
    assert entry.picture_slots == ("[2] 5.4x3.6in photo", "[3] 2.4x3.5in drawing"), (
        "the photograph, the cartoon by its first member; not the page-wide band, not the card, not its icon"
    )
    assert "picture slots [2] 5.4x3.6in photo, [3] 2.4x3.5in drawing" in entry.line()


def test_a_missed_index_names_a_drawing_a_picture_may_stand_in_for(tmp_path: Path) -> None:
    """The refusal used to list every wordless shape as `shape`, so an author reading it could
    not tell the cartoon from a band; it now says which shapes a picture may take over."""
    import pytest

    from raven_ppt.services.template.compose import shape_at

    presentation, slide, cartoon, star, panel = _drawing_page(tmp_path)

    with pytest.raises(IndexError) as caught:
        shape_at(slide, 99)

    assert "[2] drawing 2.0x4.0in (a picture may take its place)" in str(caught.value), "a member is named by its whole"


def test_a_role_label_is_the_whole_string_and_not_a_word_inside_it() -> None:
    """What `role_named` has to keep out. `_role` calls a page a divider when its heading
    merely mentions one, which is right for a page and would be a hole in a gate: the
    instruction to fill a divider's title slot names the role and is still an unfilled
    slot, and `Trends` carries `end`. Only a string that is nothing but the role's own
    name is the label a deck inherits."""
    from raven_ppt.services.template.menu import AGENDA, CLOSING, role_named

    assert role_named("目录") == AGENDA
    assert role_named("Agenda") == AGENDA, "the two halves of one label, graded the same"
    assert role_named("AGENDA") == AGENDA
    assert role_named("谢谢") == CLOSING
    assert role_named("单击此处添加章节标题") == ""
    assert role_named("Trends") == ""
    assert role_named("工作感悟") == "", "one of the two live misses the length floor let through"


def test_the_menu_names_a_small_picture_an_icon_slot(tmp_path: Path) -> None:
    """The red template's page 4 carries three 1.7in seals over its three cards, and a live
    deck kept them as the marks on three phases of its own; the menu had called them
    photos. A picture no longer than ICON_MAX_IN a side is an icon slot, one per unit."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.measure.geometry import ICON_MAX_IN
    from raven_ppt.services.template.menu import ICON_SLOT_NOTE, menu

    seal = tmp_path / "seal.png"
    Image.new("RGB", (120, 120), (200, 30, 40)).save(seal)
    photo = tmp_path / "photo.png"
    Image.new("RGB", (400, 300), (90, 90, 90)).save(photo)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(seal), Inches(1.6), Inches(2.5), Inches(1.7), Inches(1.7))
    slide.shapes.add_picture(str(photo), Inches(5.0), Inches(2.5), Inches(5.3), Inches(2.9))
    path = tmp_path / "slots.pptx"
    presentation.save(str(path))

    (entry,) = menu(path)

    assert ICON_MAX_IN == 1.8
    assert entry.picture_slots == ("[1] 1.7x1.7in icon", "[2] 5.3x2.9in photo")
    assert "one per unit" in ICON_SLOT_NOTE and "swap_icon" in ICON_SLOT_NOTE and "drop_shape" in ICON_SLOT_NOTE


def test_a_small_transparent_glyph_is_an_icon_slot_and_the_same_glyph_grown_is_a_cut_out(tmp_path: Path) -> None:
    """Size before transparency: a transparent glyph within ICON_MAX_IN is a mark whichever
    way it is drawn, and a slot named `cut-out` would not carry ICON_SLOT_NOTE's one-per-unit
    ask. The same PNG past that size is the floating illustration a cut-out slot is for."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.menu import menu

    glyph = tmp_path / "glyph.png"
    canvas = Image.new("RGBA", (200, 200), (0, 0, 0, 0))
    canvas.paste((30, 120, 120, 255), (60, 60, 140, 140))
    canvas.save(glyph)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(glyph), Inches(1.0), Inches(2.5), Inches(1.5), Inches(1.5))
    slide.shapes.add_picture(str(glyph), Inches(6.0), Inches(2.0), Inches(3.0), Inches(3.0))
    path = tmp_path / "glyphs.pptx"
    presentation.save(str(path))

    (entry,) = menu(path)

    assert entry.picture_slots == ("[1] 1.5x1.5in icon", "[2] 3.0x3.0in cut-out")


def test_the_menu_names_a_transparent_illustration_a_cut_out(tmp_path: Path) -> None:
    """The teal template's cartoons are PNGs two thirds transparent, floating on the page's
    ground with boxes that run into the title row; the menu called them photos, and an
    author put a photograph in one -- which hugged the title on the live page."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.menu import menu

    cut = tmp_path / "cartoon.png"
    canvas = Image.new("RGBA", (400, 300), (0, 0, 0, 0))
    canvas.paste((30, 120, 120, 255), (100, 60, 300, 240))
    canvas.save(cut)
    photo = tmp_path / "photo.png"
    Image.new("RGB", (400, 300), (90, 90, 90)).save(photo)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(cut), Inches(7.1), Inches(1.0), Inches(5.0), Inches(3.3))
    slide.shapes.add_picture(str(photo), Inches(0.7), Inches(3.3), Inches(5.7), Inches(3.5))
    path = tmp_path / "slots.pptx"
    presentation.save(str(path))

    (entry,) = menu(path)

    assert entry.picture_slots == ("[1] 5.0x3.3in cut-out", "[2] 5.7x3.5in photo")


def test_an_opaque_picture_in_a_cut_outs_box_is_warned_about_and_a_cut_out_is_not(tmp_path: Path) -> None:
    import warnings

    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import replace_picture

    cut = tmp_path / "cartoon.png"
    canvas = Image.new("RGBA", (400, 300), (0, 0, 0, 0))
    canvas.paste((30, 120, 120, 255), (100, 60, 300, 240))
    canvas.save(cut)
    photo = tmp_path / "photo.png"
    Image.new("RGB", (400, 300), (90, 90, 90)).save(photo)
    another = tmp_path / "stall.png"
    canvas.save(another)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_picture(str(cut), Inches(7.1), Inches(1.0), Inches(5.0), Inches(3.3))

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        replace_picture(frame, photo, "cover")
    said = [str(w.message) for w in caught if "cut-out" in str(w.message)]
    assert len(said) == 1 and "5.0x3.3in at 7.10, 1.00" in said[0] and "transparent=true" in said[0]

    frame2 = slide.shapes.add_picture(str(cut), Inches(1), Inches(1), Inches(5.0), Inches(3.3))
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        replace_picture(frame2, another, "contain")
    assert not [w for w in caught if "cut-out" in str(w.message)], "a cut-out for a cut-out is what the slot wants"


def _bottom_anchored_title(anchor: str = "b", *, own_autofit: bool = False):
    """A page whose title placeholder takes its anchor from the layout, as templates do.

    Not one built by hand on a blank layout: the measurement this exists for is that
    the anchor is written on the layout and the shape answers nothing, which is how
    every closing page in the bundled templates is drawn.
    """
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches, Pt

    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    layout = built.slide_layouts[0]
    holder = next(p for p in layout.placeholders if p.placeholder_format.idx == 0)
    body = holder.text_frame._txBody.find(f"{ns}bodyPr")
    body.set("anchor", anchor)
    etree.SubElement(body, f"{ns}normAutofit")
    slide = built.slides.add_slide(layout)
    title = slide.shapes.title
    title.left, title.top = Inches(0.72), Inches(1.24)
    title.width, title.height = Inches(5.58), Inches(3.0)
    frame = title.text_frame
    if own_autofit:
        etree.SubElement(frame._txBody.find(f"{ns}bodyPr"), f"{ns}normAutofit")
    frame.text = "Thank you\nfor watching"
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(72)
    return built, title


def test_replace_text_reports_a_display_headline_growing_up_off_the_canvas(tmp_path: Path) -> None:
    """The closing page of a live medium-tier deck, and the one defect that destroyed
    a page rather than spoiling it: `replace_text` keeps the prototype's geometry, the
    frame is bottom-anchored, and a 4-line headline at the template's display size grew
    upward until its first line was cut off by the top edge of the slide. Verified
    against the render of that build."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    _, title = _bottom_anchored_title()
    replace_text(title, "Buy the cycle.\nPrice the tail risks.")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()

    assert len(caught) == 1
    said = str(caught[0].message)
    assert "page 1:" in said
    assert "needs 4 lines at 72pt in this 5.38in column and the box shows 2" in said
    assert "the copy it replaces took 2" in said
    assert "Bottom-anchored: it grows upward, beginning" in said
    assert "above the top of the slide, and its first line is cut off" in said


def test_replace_text_says_nothing_when_the_copy_is_the_templates_own(tmp_path: Path) -> None:
    """The one case that must never report. The box's height in lines is arithmetic on
    an assumed line-height factor and can read under what the designer put in it, so
    the prototype's own copy is the floor -- which makes writing it back silent by
    construction rather than by a threshold that happens to hold. Replayed over the ten
    bundled templates, 2,084 boxes on 211 pages: no reports."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    _, title = _bottom_anchored_title()
    held = title.text_frame.text
    replace_text(title, held)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    assert not caught, [str(w.message) for w in caught]


def test_replace_text_reports_a_shrink_where_the_frame_itself_autofits(tmp_path: Path) -> None:
    """`normAutofit` on the shape and on the layout are not the same thing to the
    renderer. The bundled card labels carry it on the shape and a live deck's copy came
    back at 17pt where the template set 20pt; the closing headline inherits it from the
    layout alone and was drawn at the full 72pt and clipped. So the shrink is read off
    the shape, and how far it shrinks is not predicted -- an earlier version put the
    scale at shows/needs and promised 10pt where the render gave 17pt."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, _shrinks_to_fit, replace_text

    _, inherited = _bottom_anchored_title()
    assert not _shrinks_to_fit(inherited), "the layout's autofit is not the shape's"
    _, own = _bottom_anchored_title(own_autofit=True)
    assert _shrinks_to_fit(own)

    replace_text(own, "Buy the cycle.\nPrice the tail risks.")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    said = str(caught[0].message)
    assert "shrinks its text to fit, so this box comes back under the 72pt" in said
    assert "grows upward" not in said, "a frame that shrinks does not spill"
    assert "10pt" not in said, "how far it shrinks is the renderer's arithmetic"


def test_replace_text_reports_copy_running_off_the_bottom_of_the_page(tmp_path: Path) -> None:
    """The other direction: a top-anchored frame low on the page grows down, and the
    report says how far past the canvas rather than only how far past the box."""
    import warnings

    from pptx import Presentation
    from pptx.util import Inches, Pt

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    slide = built.slides.add_slide(built.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.72), Inches(6.6), Inches(3.0), Inches(0.6))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = "A short note"
    frame.paragraphs[0].runs[0].font.size = Pt(18)

    replace_text(
        box,
        "A source note long enough to need four lines of this narrow column, "
        "which the page has no room below the box to give it",
    )
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    said = str(caught[0].message)
    assert "It grows downward" in said and "past the bottom edge of the slide" in said


def test_replace_text_keeps_the_prototypes_copy_across_a_second_write(tmp_path: Path) -> None:
    """The floor is the template's copy, not the author's own first draft. A page
    written twice measured the second string against the first, and the box's certified
    capacity went with it."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    _, title = _bottom_anchored_title()
    replace_text(title, "Buy")
    replace_text(title, "Buy the cycle.\nPrice the tail risks.")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    assert len(caught) == 1
    assert "the copy it replaces took 2" in str(caught[0].message), "the prototype's, not 'Buy'"


# --- one door: what an unfilled frame does, and what the refusals say ------------

_BUNDLED = Path(__file__).resolve().parents[1] / "plugins-dist/ppt-engine/raven_ppt/assets/templates"
_needs_payload = pytest.mark.skipif(
    not any(_BUNDLED.glob("*.pptx")),
    reason="the template payload is fetched, not tracked; run plugins-dist/ppt-engine/fetch_templates.py",
)


def _all_of(container):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in container.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _all_of(shape)
        else:
            yield shape


@_needs_payload
def test_a_frame_no_call_named_is_refused_by_the_gate(tmp_path: Path) -> None:
    """The whole of D38 in one page, on a real template rather than a fixture.

    The route since removed emptied the frames a call did not name, and the resulting
    page was unnameable: eight of them shipped out of one live deck with every check
    green, each holding the two lines its call had written over thirteen to twenty-six
    blank boxes. Keeping the template's words instead makes the same mistake loud -- the
    page says the template's own copy out loud and `placeholder_copy` refuses to publish
    it, naming the page and the line. That is the trade: the failure is not prevented,
    it is made impossible to ship.
    """
    from pptx import Presentation

    from raven_ppt.services.measure.adherence import placeholder_copy
    from raven_ppt.services.template import clone_page, prototype

    house = _BUNDLED / "amber_wave_quarterly_summary.pptx"
    template = Presentation(str(house))
    built = Presentation(str(house))
    for slide in list(built.slides._sldIdLst):
        built.slides._sldIdLst.remove(slide)

    page = clone_page(built, prototype(template, 4))
    named = next(
        shape for shape in _all_of(page) if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
    )
    replace_text(named, "本页自己的标题")
    kept = [
        shape.text_frame.text.strip()
        for shape in _all_of(page)
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
    ]
    assert len(kept) > 1, f"a frame no call named still holds the template's own copy: {kept}"

    out = tmp_path / "one-page.pptx"
    built.save(str(out))
    refused = placeholder_copy(out, house)
    assert refused, f"the gate has to see what the clone left standing: {kept}"
    assert all(finding.page == 1 for finding in refused)
    assert "placeholder text" in refused[0].message


def _ascending_run(tmp_path: Path, name: str, boxes, frames_per_unit: int) -> Path:
    """A page whose units climb a diagonal, which is the arrangement that refuses growth.

    The boxes are the real ones off the page each recorded crash was on, so the run is
    "irregular" here for the same reason it was there: no row, no column, no grid, and
    where a fifth unit would go is the designer's call and not this code's.
    """
    from pptx import Presentation
    from pptx.util import Inches

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    page = built.slides.add_slide(built.slide_layouts[6])
    for index, (left, top, width, height) in enumerate(boxes):
        group = page.shapes.add_group_shape()
        said = ("跨界合作的意义", "资源共享，优势互补", "单击此处添加文本")
        for row in range(frames_per_unit):
            box = group.shapes.add_textbox(
                Inches(left), Inches(top + row * (height / frames_per_unit)), Inches(width), Inches(0.4)
            )
            box.text_frame.text = said[row] if row < len(said) else f"0{index + 1}"
    path = tmp_path / name
    built.save(str(path))
    return path


# The three build crashes of the newest recorded run, by the prototype each was on and
# the geometry that page really has. Two are the run refusing to grow, one is a unit
# refusing an entry; all three used to print `frames` -- the prototype's own shapes --
# where the author's values belonged, so a probe passing ['a', 'b', 'c'] read the
# template's Chinese back at itself and changed the prototype instead of the list.
_STAIRS = ((0.74, 4.03, 2.65, 3.40), (3.79, 3.24, 2.67, 4.18), (6.86, 2.45, 2.67, 4.97), (9.93, 1.66, 2.67, 5.76))
_BRANCHES = ((5.50, 1.86, 2.38, 1.22), (3.90, 3.53, 2.38, 1.12), (7.06, 3.53, 2.38, 1.12), (7.06, 5.22, 2.38, 1.12))


@pytest.mark.parametrize(
    ("label", "boxes", "frames", "items", "expected"),
    [
        (
            "prototype 10: 5 items to a 4-unit page",
            _STAIRS,
            3,
            [["01", "Now: Blackwell Ultra", "shipping in volume"]] * 5,
            ("cannot take 5", "The 5 entries you gave", "'Now: Blackwell Ultra'"),
        ),
        (
            "prototype 6: 7 items to a 4-unit page",
            _BRANCHES,
            3,
            [["Revenue", "up 65% YoY", "the four-year trail"]] * 7,
            ("cannot take 7", "The 7 entries you gave", "'Revenue'"),
        ),
        (
            "prototype 11: 5 items of 3 values, unit holds 2",
            _STAIRS + ((9.93, 0.5, 2.67, 1.0),),
            2,
            [["Accelerators", "AMD and Intel compete", "third value"]] * 5,
            (
                "2 text shape(s) a unit on this page holds",
                "The 5 entries you gave",
                "'Accelerators'",
                "'跨界合作的意义'",
            ),
        ),
    ],
)
def test_an_arity_refusal_names_the_callers_values_and_the_pages_real_shape(
    tmp_path: Path, label, boxes, frames, items, expected
) -> None:
    """What the message is for. Ten builds across the measured runs died on an arity
    refusal, and the sentence they got quoted the shapes already on the prototype on
    both sides of "N values were given"."""
    from pptx import Presentation

    from raven_ppt.services.template import clone_page, fill, units

    path = _ascending_run(tmp_path, "stairs.pptx", boxes, frames)
    source, out = Presentation(str(path)), Presentation(str(path))
    slide = clone_page(out, source.slides[0])

    with pytest.raises(ValueError) as raised:
        fill(max(units(slide), key=len), items)
    said = str(raised.value)
    for phrase in expected:
        assert phrase in said, f"{label}: {phrase!r} missing from: {said}"


@_needs_payload
def test_replace_text_carries_a_run_list_on_a_cloned_page(tmp_path: Path) -> None:
    """The route since removed stringified the text it was handed, so a page wanting one
    word in the accent had to be written twice: one call for the cards, then a second
    pass with `replace_text` for the heading. Three recorded scripts did exactly that,
    twenty call sites between them."""
    from pptx import Presentation

    from raven_ppt.services.assets import script_helpers
    from raven_ppt.services.template import clone_page, prototype

    namespace: dict = {}
    exec(compile(script_helpers.script_helper_files()["ppt_layout.py"], "ppt_layout", "exec"), namespace)  # noqa: S102
    Run = namespace["Run"]

    house = _BUNDLED / "amber_wave_quarterly_summary.pptx"
    template = Presentation(str(house))
    built = Presentation(str(house))
    for slide in list(built.slides._sldIdLst):
        built.slides._sldIdLst.remove(slide)

    page = clone_page(built, prototype(template, 4))
    named = next(
        shape for shape in _all_of(page) if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
    )
    replace_text(named, [Run("访客中约 "), Run("84%", bold=True), Run(" 到访过")])
    wrote = [
        shape for shape in _all_of(page) if getattr(shape, "has_text_frame", False) and "84%" in shape.text_frame.text
    ]
    assert wrote, "the run list reached the page"
    runs = [run for paragraph in wrote[0].text_frame.paragraphs for run in paragraph.runs]
    assert [run.text for run in runs] == ["访客中约 ", "84%", " 到访过"], "one run per piece, not one line"
    assert runs[1].font.bold is True


def test_replace_picture_takes_the_wash_and_the_trim_beside_the_fit(tmp_path: Path, image) -> None:
    """The route since removed carried image, box and fit and stopped there, so a page
    that wanted the wash `backdrop` takes, or a trim off one edge, had to reach the frame
    again by index on the returned slide -- five call sites in one recorded script."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template import clone_page, shape_at

    old, new = image("old.png", (20, 90, 140)), image("new.png", (200, 40, 40))
    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    page = built.slides.add_slide(built.slide_layouts[6])
    page.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(9), Inches(0.8)).text_frame.text = "单击此处添加页面标题"
    page.shapes.add_picture(str(old), Inches(1), Inches(2), Inches(5), Inches(3))
    path = tmp_path / "photo.pptx"
    built.save(str(path))

    source = Presentation(str(path))
    washed = clone_page(Presentation(str(path)), source.slides[0])
    replace_picture(shape_at(washed, 2), new, "cover", alpha=0.3)
    assert "alphaModFix" in washed._element.xml, "alpha reached replace_picture"
    trimmed = clone_page(Presentation(str(path)), source.slides[0])
    replace_picture(shape_at(trimmed, 2), new, "cover", trim=(0, 0, 0, 0.2))
    assert "srcRect" in trimmed._element.xml, "trim reached replace_picture"

    with pytest.raises(TypeError, match="image"):
        replace_picture(shape_at(clone_page(Presentation(str(path)), source.slides[0]), 2))
    with pytest.raises(TypeError, match="nope"):
        replace_picture(shape_at(clone_page(Presentation(str(path)), source.slides[0]), 2), new, "cover", nope=1)


def _bottom_anchored_title(anchor: str = "b", *, own_autofit: bool = False):
    """A page whose title placeholder takes its anchor from the layout, as templates do.

    Not one built by hand on a blank layout: the measurement this exists for is that
    the anchor is written on the layout and the shape answers nothing, which is how
    every closing page in the bundled templates is drawn.
    """
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches, Pt

    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    layout = built.slide_layouts[0]
    holder = next(p for p in layout.placeholders if p.placeholder_format.idx == 0)
    body = holder.text_frame._txBody.find(f"{ns}bodyPr")
    body.set("anchor", anchor)
    etree.SubElement(body, f"{ns}normAutofit")
    slide = built.slides.add_slide(layout)
    title = slide.shapes.title
    title.left, title.top = Inches(0.72), Inches(1.24)
    title.width, title.height = Inches(5.58), Inches(3.0)
    frame = title.text_frame
    if own_autofit:
        etree.SubElement(frame._txBody.find(f"{ns}bodyPr"), f"{ns}normAutofit")
    frame.text = "Thank you\nfor watching"
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(72)
    return built, title


def test_replace_text_reports_a_display_headline_growing_up_off_the_canvas(tmp_path: Path) -> None:
    """The closing page of a live medium-tier deck, and the one defect that destroyed
    a page rather than spoiling it: `replace_text` keeps the prototype's geometry, the
    frame is bottom-anchored, and a 4-line headline at the template's display size grew
    upward until its first line was cut off by the top edge of the slide. Verified
    against the render of that build."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    _, title = _bottom_anchored_title()
    replace_text(title, "Buy the cycle.\nPrice the tail risks.")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()

    assert len(caught) == 1
    said = str(caught[0].message)
    assert "page 1:" in said
    assert "needs 4 lines at 72pt in this 5.38in column and the box shows 2" in said
    assert "the copy it replaces took 2" in said
    assert "Bottom-anchored: it grows upward, beginning" in said
    assert "above the top of the slide, and its first line is cut off" in said


def test_replace_text_says_nothing_when_the_copy_is_the_templates_own(tmp_path: Path) -> None:
    """The one case that must never report. The box's height in lines is arithmetic on
    an assumed line-height factor and can read under what the designer put in it, so
    the prototype's own copy is the floor -- which makes writing it back silent by
    construction rather than by a threshold that happens to hold. Replayed over the ten
    bundled templates, 2,084 boxes on 211 pages: no reports."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    _, title = _bottom_anchored_title()
    held = title.text_frame.text
    replace_text(title, held)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    assert not caught, [str(w.message) for w in caught]


def test_replace_text_reports_a_shrink_where_the_frame_itself_autofits(tmp_path: Path) -> None:
    """`normAutofit` on the shape and on the layout are not the same thing to the
    renderer. The bundled card labels carry it on the shape and a live deck's copy came
    back at 17pt where the template set 20pt; the closing headline inherits it from the
    layout alone and was drawn at the full 72pt and clipped. So the shrink is read off
    the shape, and how far it shrinks is not predicted -- an earlier version put the
    scale at shows/needs and promised 10pt where the render gave 17pt."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, _shrinks_to_fit, replace_text

    _, inherited = _bottom_anchored_title()
    assert not _shrinks_to_fit(inherited), "the layout's autofit is not the shape's"
    _, own = _bottom_anchored_title(own_autofit=True)
    assert _shrinks_to_fit(own)

    replace_text(own, "Buy the cycle.\nPrice the tail risks.")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    said = str(caught[0].message)
    assert "shrinks its text to fit, so this box comes back under the 72pt" in said
    assert "grows upward" not in said, "a frame that shrinks does not spill"
    assert "10pt" not in said, "how far it shrinks is the renderer's arithmetic"


def test_replace_text_reports_copy_running_off_the_bottom_of_the_page(tmp_path: Path) -> None:
    """The other direction: a top-anchored frame low on the page grows down, and the
    report says how far past the canvas rather than only how far past the box."""
    import warnings

    from pptx import Presentation
    from pptx.util import Inches, Pt

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    built = Presentation()
    built.slide_width, built.slide_height = Inches(13.333), Inches(7.5)
    slide = built.slides.add_slide(built.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.72), Inches(6.6), Inches(3.0), Inches(0.6))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = "A short note"
    frame.paragraphs[0].runs[0].font.size = Pt(18)

    replace_text(
        box,
        "A source note long enough to need four lines of this narrow column, "
        "which the page has no room below the box to give it",
    )
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    said = str(caught[0].message)
    assert "It grows downward" in said and "past the bottom edge of the slide" in said


def test_replace_text_keeps_the_prototypes_copy_across_a_second_write(tmp_path: Path) -> None:
    """The floor is the template's copy, not the author's own first draft. A page
    written twice measured the second string against the first, and the box's certified
    capacity went with it."""
    import warnings

    from raven_ppt.services.template.compose import _say_what_did_not_fit, replace_text

    _, title = _bottom_anchored_title()
    replace_text(title, "Buy")
    replace_text(title, "Buy the cycle.\nPrice the tail risks.")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        _say_what_did_not_fit()
    assert len(caught) == 1
    assert "the copy it replaces took 2" in str(caught[0].message), "the prototype's, not 'Buy'"


def test_replace_text_writes_into_a_frame_that_has_no_paragraph_at_all() -> None:
    """Eight card panels in the bundled templates carry a `<p:txBody>` of nothing but
    `<a:bodyPr/>` and an empty `<a:lstStyle/>` -- `black_circuit_tech_launch` page 25
    and `green_aurora_tech_trends` page 20, four each. Copying the last paragraph's
    properties raised `IndexError` on the empty tuple, and a build is one program, so
    the raise took every page after it."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import _A, replace_text

    built = Presentation()
    slide = built.slides.add_slide(built.slide_layouts[6])
    panel = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(4), Inches(2))
    body = panel.text_frame._txBody
    for paragraph in list(body.findall(f"{{{_A}}}p")):
        body.remove(paragraph)
    assert panel.text_frame.paragraphs == (), "the shape the templates hold: a text body with no paragraph"

    replace_text(panel, "Grid-scale storage fell to $117/kWh")

    assert [p.text for p in panel.text_frame.paragraphs] == ["Grid-scale storage fell to $117/kWh"]
    written = panel.text_frame.paragraphs[0]
    assert written._p.find(f"{{{_A}}}pPr") is None, "nothing invented: the paragraph states nothing of its own"
    assert written.runs[0]._r.find(f"{{{_A}}}rPr") is None, "nor does its run -- the file's own defaults answer"

    # Several lines still land as several paragraphs, the second copied off the first.
    replace_text(panel, "first\nsecond")
    assert [p.text for p in panel.text_frame.paragraphs] == ["first", "second"]


def test_replace_text_writing_nothing_back_leaves_a_paragraphless_frame_empty() -> None:
    """Writing a paragraphless frame's own copy back is writing `''`. It must not raise
    and must not put anything on the page: the ten templates' 3157 text frames all
    written back are 0 raises and no visible change to either affected page."""
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import _A, replace_text

    built = Presentation()
    slide = built.slides.add_slide(built.slide_layouts[6])
    panel = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(4), Inches(2))
    body = panel.text_frame._txBody
    for paragraph in list(body.findall(f"{{{_A}}}p")):
        body.remove(paragraph)

    replace_text(panel, panel.text_frame.text)

    assert panel.text_frame.text == ""
    assert len(panel.text_frame.paragraphs) == 1, "the one paragraph the format requires, and no run in it"
    assert panel.text_frame.paragraphs[0].runs == ()
