"""The type census, and the two numbers a built page has to clear."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from raven_ppt.contracts.findings import Severity
from raven_ppt.services.measure.type_size import (
    BODY_FLOOR_PT,
    BODY_PT,
    LABEL_PT,
    MIN_FLOOR_PT,
    Span,
    census,
    drift_findings,
    outranked_findings,
    row_findings,
    scale_findings,
    type_findings,
    type_floors,
)
from tests._ppt_engine_fixtures import (  # noqa: F401
    DeckBuilder,
    deck,
    image,
    noise_image,
    noise_png,
    product_page,
    template_file,
)

pytest.importorskip("pptx")


def test_the_floors_are_two_constants(deck: DeckBuilder) -> None:
    """Every PowerPoint canvas is 7.5in tall, so the floors do not scale.

    The page height is still accepted, and still ignored: the day a 5.625in
    canvas turns up, this test is where the exception gets written.
    """
    assert (BODY_FLOOR_PT, MIN_FLOOR_PT) == (14.0, 10.8)
    assert type_floors() == (14.0, 10.8)
    assert type_floors(7.5) == (14.0, 10.8)
    assert type_floors(5.625) == (14.0, 10.8)


def test_body_size_is_what_most_of_the_copy_runs_at(deck: DeckBuilder) -> None:
    """A 32pt title does not make a page of 12pt copy a 32pt page."""
    page = deck.page()
    deck.text(page, ("Title of the page", 32.0), ("body copy here" * 6, 12.0))

    measured = census(deck.save())[0]

    assert measured.body_pt == 12.0
    assert measured.under_floor is True


def test_table_cells_count_toward_the_census(deck: DeckBuilder) -> None:
    """Results decks put their smallest type in tables, so tables are read."""
    from pptx.util import Pt

    page = deck.page()
    deck.text(page, ("a heading that clears the floor", 18.0), height=1.0)
    table = deck.table(page, 2, 2, top=2.0)
    for row in table.table.rows:
        for cell in row.cells:
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = "a cell of table copy"
            run.font.size = Pt(9.0)
    built = deck.save()

    measured = census(built)[0]

    assert measured.smallest_pt == 9.0
    assert measured.below_hard_floor > 0
    assert [finding.detail["smallest_pt"] for finding in type_findings(built)] == [9.0]


def test_short_marks_are_not_held_to_the_body_floor(deck: DeckBuilder) -> None:
    """Axis labels and page numbers are set small on purpose."""
    page = deck.page()
    deck.text(
        page,
        ("Findings across the three benchmarks", 20.0),
        ("body copy that carries the page" * 3, 16.0),
        ("7", 9.0),
    )
    built = deck.save()

    measured = census(built)[0]

    assert measured.body_pt == 16.0
    assert measured.smallest_pt == 16.0  # the 9pt page number never entered the census
    assert type_findings(built) == []


def test_findings_are_reported_per_page(deck: DeckBuilder) -> None:
    """A deck is rarely wrong everywhere, and the fix is per page."""
    for text, size in (
        ("copy that clears the floor comfortably" * 3, 16.0),
        ("copy set too small to project" * 3, 11.5),
        ("copy that clears the floor comfortably" * 3, 15.0),
    ):
        deck.text(deck.page(), (text, size))

    findings = type_findings(deck.save())

    assert [finding.page for finding in findings] == [2]
    assert findings[0].detail["body_pt"] == 11.5
    assert findings[0].detail["body_floor_pt"] == 14.0
    assert "under the 14.0pt floor" in findings[0].message


def test_the_floor_is_a_warning_and_not_a_refusal(deck: DeckBuilder) -> None:
    """Raising a size costs room, and the room has to come from somewhere.

    A warning rather than a refusal because the room comes from the copy: a gate
    that blocked publication until the floor was met could be answered by
    shrinking the copy back, which is the oscillation D2 describes.
    """
    deck.text(deck.page(), ("copy set too small to project" * 3, 11.0))

    finding = type_findings(deck.save())[0]

    assert finding.kind == "type_floor"
    assert finding.severity is Severity.WARNING
    assert "Do not shrink it back to fit" in finding.message


def test_a_deck_that_clears_the_floor_reports_nothing(deck: DeckBuilder) -> None:
    for _ in range(3):
        deck.text(deck.page(), ("copy that clears the floor" * 4, 15.0))

    assert type_findings(deck.save()) == []


def test_a_page_with_no_sized_copy_is_not_a_finding(deck: DeckBuilder) -> None:
    """A page of pictures has no body size to be under a floor."""
    deck.page()

    built = deck.save()

    assert census(built)[0].body_pt is None
    assert type_findings(built) == []


def test_a_source_line_is_not_held_to_the_body_floor(tmp_path: Path) -> None:
    """ "来源：TarViS 原论文（CVPR 2023）" at 11pt is legible, deliberate, and 24
    characters long, so no length rule tells it from copy. It appeared on seven pages
    of one delivered deck as the same finding, which is how a check teaches an author
    to stop reading it."""
    from raven_ppt.services.measure.type_size import Span, type_findings

    deck = _deck(
        tmp_path,
        [
            ("来源：TarViS 原论文（CVPR 2023），Table 2", 1.0, 6.9, 6.0, 0.3),
            ("这一段是页面的正文，长度足够被当作正文而不是标记来判断", 1.0, 2.0, 6.0, 1.0),
        ],
    )
    spans = [
        Span(page=1, size_pt=11.0, text="来源：TarViS 原论文（CVPR 2023），Table 2", x0=75, y0=500, x1=400, y1=515),
        Span(
            page=1,
            size_pt=17.0,
            text="这一段是页面的正文，长度足够被当作正文而不是标记来判断",
            x0=75,
            y0=150,
            x1=460,
            y1=170,
        ),
    ]

    assert type_findings(deck, spans) == []


def test_body_copy_under_the_floor_still_reports(tmp_path: Path) -> None:
    from raven_ppt.services.measure.type_size import Span, type_findings

    deck = _deck(tmp_path, [("这一段正文被框压到了读者看不清的字号，需要报出来给设计环", 1.0, 2.0, 6.0, 1.0)])
    spans = [
        Span(
            page=1,
            size_pt=10.8,
            text="这一段正文被框压到了读者看不清的字号，需要报出来给设计环",
            x0=75,
            y0=150,
            x1=460,
            y1=165,
        )
    ]

    findings = type_findings(deck, spans)
    assert [f.kind for f in findings] == ["type_floor"]
    assert findings[0].detail["sizes_pt"] == [10.8]


def _deck(
    tmp_path: Path,
    boxes: list[tuple[str, float, float, float, float]],
    *,
    name: str = "deck",
    size: float | None = None,
) -> Path:
    """One page of boxes at stated positions. `size` states the runs' own point size.

    Without it the boxes declare nothing, which is the state a template placeholder is
    in and what the census is silent about."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for text, left, top, width, height in boxes:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        if size is None:
            box.text_frame.text = text
            continue
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
    path = tmp_path / f"{name}.pptx"
    presentation.save(str(path))
    return path


# The band between the floor and the body size, which is where a hand-picked integer
# lands. Its own tests, because the claim is different from the floor's: not "a reader
# cannot read this" but "nobody chose this size".
COPY = "这一段是页面的正文，长度足够被当作正文而不是一个标记来判断"


def test_copy_between_the_floor_and_the_body_size_is_reported(deck: DeckBuilder) -> None:
    """15pt clears every floor there is and is still not a step of the ramp.

    The size 15 of 43 copy blocks across 34 generated decks came out at -- more than
    any other -- and the one `type_floor` is structurally unable to see, because the
    tier that applies to copy is 14pt and 15 clears it.
    """
    deck.text(deck.page(), (COPY, 15.0), top=2.0, height=1.0)

    findings = scale_findings(deck.save())

    assert [finding.kind for finding in findings] == ["type_scale"]
    assert findings[0].page == 1
    assert findings[0].detail["sizes_pt"] == [15.0]
    assert findings[0].detail["reaches_body_pt"] is False
    assert "not a step of the ramp" in findings[0].message
    assert "under BODY_PT (16pt)" in findings[0].message
    assert "`size=BODY_PT`" in findings[0].message


def test_the_scale_reports_and_never_refuses(deck: DeckBuilder) -> None:
    """Invariant 3. Raising a size costs room and the room comes from the copy, so a
    refusal here could be answered by cutting the page back -- the oscillation of
    design doc D2, which the floor above declines for the same reason."""
    deck.text(deck.page(), (COPY, 15.0), top=2.0, height=1.0)

    assert scale_findings(deck.save())[0].severity is Severity.WARNING


def test_a_step_of_the_ramp_is_not_a_finding(deck: DeckBuilder) -> None:
    """`BODY_PT` and the one step under it are the two sizes copy may be set at.

    14pt is in the ramp on purpose -- the comment beside it says body has room to drop
    one step and still clear the floor -- so a card body at 14pt is a decision and not
    a drift. Reporting it would have fired on 7 more of the 43 copy blocks measured.
    """
    deck.text(deck.page(), (COPY, 16.0), top=2.0, height=1.0)
    deck.text(deck.page(), (COPY, 14.0), top=2.0, height=1.0)
    deck.text(deck.page(), (COPY, 20.0), top=2.0, height=1.0)

    assert scale_findings(deck.save()) == []


def test_copy_under_the_floor_is_left_to_the_floor(deck: DeckBuilder) -> None:
    """Two findings on one box is a finding an author learns to skip.

    12pt copy is already reported, by name and with the same move -- bring it up. So
    this check starts at the floor and says nothing below it.
    """
    deck.text(deck.page(), (COPY, 12.0), top=2.0, height=1.0)
    built = deck.save()

    assert [finding.kind for finding in type_findings(built)] == ["type_floor"]
    assert scale_findings(built) == []


def test_a_short_label_off_the_ramp_is_not_copy(deck: DeckBuilder) -> None:
    """The 20-character line this file already draws for its floor tier, reused.

    A byline, a unit, a chart's axis: "公司内部技术评审" at 12pt is a label, and both
    models fought the floor over exactly that until the tier existed.
    """
    deck.text(deck.page(), ("公司内部技术评审", 15.0), top=2.0, height=0.4)

    assert scale_findings(deck.save()) == []


def test_a_source_line_and_a_footer_are_not_copy(deck: DeckBuilder) -> None:
    """A deck sets its credits smaller on purpose, which is why `_is_caption` exists."""
    page = deck.page()
    deck.text(page, ("来源：TarViS 原论文（CVPR 2023），Table 2", 15.0), top=2.0, height=0.4)
    deck.text(page, ("TarViS · CVPR 2023 · arXiv:2301.02657 · 第 4 页", 15.0), top=7.0, height=0.3)

    assert scale_findings(deck.save()) == []


def test_a_page_of_pictures_is_not_a_finding(deck: DeckBuilder) -> None:
    deck.page()

    assert scale_findings(deck.save()) == []


def test_the_templates_own_copy_is_not_the_authors_choice(tmp_path: Path) -> None:
    """A deck inside a template inherits sizes it cannot name the ramp for.

    The ten bundled templates set 335 of their 356 copy blocks under 16pt, 269 of them
    at 12pt, so a check held to our ramp with no notion of a clone would report every
    page of every templated deck -- and ask for a fix that means abandoning the
    template, which is the dead end invariant 6 forbids. A cloned page keeps its
    prototype's positions exactly, so position is the answer.

    Asserted in both directions: the same deck fires when nothing says the box came
    from a template, so it is the scope suppressing it and not the band.
    """
    template = _deck(tmp_path, [(COPY, 1.0, 2.0, 6.0, 1.0)], name="template", size=15.0)
    cloned = _deck(tmp_path, [(COPY, 1.0, 2.0, 6.0, 1.0)], name="built", size=15.0)

    assert scale_findings(cloned, template) == []
    assert [finding.kind for finding in scale_findings(cloned)] == ["type_scale"]


def test_a_box_the_author_added_inside_a_template_is_still_the_authors(tmp_path: Path) -> None:
    """The mix a real templated deck is: the template's frame, the author's body.

    Per shape rather than per page, because `clone_page` copies a prototype and the author
    then composes inside it -- scoping by page would have excluded the copy they wrote.
    """
    template = _deck(tmp_path, [(COPY, 1.0, 2.0, 6.0, 1.0)], name="template", size=15.0)
    built = _deck(
        tmp_path,
        [(COPY, 1.0, 2.0, 6.0, 1.0), (COPY, 7.0, 2.0, 5.0, 1.0)],
        name="built",
        size=15.0,
    )

    findings = scale_findings(built, template)

    assert [finding.detail["sizes_pt"] for finding in findings] == [[15.0]]


def test_the_ramp_is_the_one_ppt_layout_hands_the_author(tmp_path: Path) -> None:
    """The two constants, pinned against the module the author actually imports.

    `ppt_layout` reaches the author as text -- `script_helpers` writes it beside the
    build script -- so the engine cannot import it and the sizes are stated twice. This
    is the only place the two copies can be held to each other.
    """
    import sys

    from raven_ppt.services.assets.script_helpers import script_helper_files

    for name, text in script_helper_files().items():
        (tmp_path / name).write_text(text, encoding="utf-8")
    sys.path.insert(0, str(tmp_path))
    try:
        for stale in ("ppt_layout",):
            sys.modules.pop(stale, None)
        import ppt_layout
    finally:
        sys.path.remove(str(tmp_path))

    assert (BODY_PT, LABEL_PT) == (float(ppt_layout.BODY_PT), float(ppt_layout.LABEL_PT))
    # And that the band this check owns is the gap the ramp leaves: the only step
    # between the floor and the body size is `LABEL_PT` itself.
    assert [step for step in ppt_layout._RAMP if BODY_FLOOR_PT <= step < BODY_PT] == [LABEL_PT]


def test_it_fires_on_a_quarter_of_a_corpus_and_not_on_the_rest(tmp_path: Path) -> None:
    """The sweep, as one deck: nine pages shaped like the ones it was measured over.

    Run over the 45 generated pages in /tmp/ab it reported 11, and 10 of those 11 were
    pages `type_floor` said nothing about at all. This is that shape held still -- a
    check that fires everywhere reports nothing, and the pages below are the ones it
    has to stay quiet on.
    """
    pages: list[tuple[str, list[tuple[str, float, float, float, float]], float]] = [
        # Fires: the body copy one point under what the deck calls body.
        ("15pt body", [(COPY, 1.0, 2.0, 6.0, 1.0)], 15.0),
        # Fires: four card bodies at once, one finding for the page.
        ("four cards at 15pt", [(COPY, 1.0 + n * 3.0, 2.0, 2.8, 1.4) for n in range(4)], 15.0),
        # Quiet: the ramp's own two steps for copy.
        ("16pt body", [(COPY, 1.0, 2.0, 6.0, 1.0)], 16.0),
        ("14pt body", [(COPY, 1.0, 2.0, 6.0, 1.0)], 14.0),
        ("20pt lead", [(COPY, 1.0, 2.0, 6.0, 1.0)], 20.0),
        # Quiet: under the floor, which `type_floor` reports by name.
        ("13pt body", [(COPY, 1.0, 2.0, 6.0, 1.0)], 13.0),
        ("9pt labels", [("flowChartAlternateProcess", 1.0 + n, 2.0, 0.9, 0.3) for n in range(6)], 9.0),
        # Quiet: a chart's marks and a page's kicker, set small on purpose.
        ("chart marks", [(f"{n + 1}月", 1.0 + n * 0.4, 5.0, 0.35, 0.3) for n in range(12)], 12.0),
        # Quiet: the two things a deck sets smallest on purpose.
        ("a source line", [("来源：TarViS 原论文（CVPR 2023），Table 2", 1.0, 2.0, 6.0, 0.4)], 15.0),
    ]
    built = _multi(tmp_path / "corpus", [boxes for _, boxes, _ in pages], [size for *_, size in pages])

    findings = scale_findings(built)

    assert [finding.page for finding in findings] == [1, 2]
    assert [len(finding.detail["sizes_pt"]) for finding in findings] == [1, 4]
    named = [name for name, *_ in pages]
    assert [named[finding.page - 1] for finding in findings] == ["15pt body", "four cards at 15pt"]


def _multi(stem: Path, pages: list[list[tuple[str, float, float, float, float]]], sizes: list[float]) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for boxes, size in zip(pages, sizes):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for text, left, top, width, height in boxes:
            box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = text
            run.font.size = Pt(size)
    path = stem.with_suffix(".pptx")
    presentation.save(str(path))
    return path


# The two pair readings: a row against itself, and a page's title against the line under
# it. Their own helpers, because both need a render's spans and one needs repeating units,
# and neither is expressible with `DeckBuilder`'s flat pages.
TITLE = "Network Status and the Global Fleet"
CAPTION = "Systems in service, route kilometres, landings, and the ageing of the fleet"


def _titled(
    tmp_path: Path,
    caption: str,
    *,
    caption_pt: float | None,
    caption_width: float = 9.0,
    title_pt: float | None = None,
    name: str = "deck",
) -> tuple[Path, Any]:
    """One page: the layout's own title placeholder, and a line under it.

    The title states no size unless `title_pt` says so, which is the state a cloned
    template page is in -- python-pptx's default master sets 44pt for a title, and that
    is the number nothing in an author's program shows it.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(0.5), Inches(0.3), Inches(9.0), Inches(1.0)
    if title_pt is None:
        title.text_frame.text = TITLE
    else:
        run = title.text_frame.paragraphs[0].add_run()
        run.text = TITLE
        run.font.size = Pt(title_pt)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(caption_width), Inches(0.6))
    if caption_pt is None:
        box.text_frame.text = caption
    else:
        run = box.text_frame.paragraphs[0].add_run()
        run.text = caption
        run.font.size = Pt(caption_pt)
    path = tmp_path / f"{name}.pptx"
    presentation.save(str(path))
    return path, presentation


def _pair_spans(title_pt: float, caption_pt: float, caption: str = CAPTION) -> list[Span]:
    """What the renderer set the two lines at, inside the boxes `_titled` drew."""
    return [
        Span(page=1, size_pt=title_pt, text=TITLE, x0=40, y0=25, x1=400, y1=25 + title_pt),
        Span(page=1, size_pt=caption_pt, text=caption, x0=40, y0=115, x1=600, y1=115 + caption_pt),
    ]


def test_a_title_and_the_line_under_it_at_one_size_is_reported(tmp_path: Path) -> None:
    """The state a reader calls two competing headings, with the caption dominant."""
    path, _ = _titled(tmp_path, CAPTION, caption_pt=44.0)

    found = outranked_findings(path, _pair_spans(44.0, 44.0))

    assert [finding.kind for finding in found] == ["outranked_title"]
    assert found[0].severity is Severity.WARNING
    assert found[0].page == 1
    assert found[0].detail["title_pt"] == 44.0
    assert found[0].detail["under_pt"] == 44.0
    assert found[0].detail["under_chars"] > found[0].detail["title_chars"]


def test_the_finding_states_the_size_the_author_could_not_see(tmp_path: Path) -> None:
    """The whole reason this warns rather than refuses.

    The title box states nothing; its 44pt comes off the master. An author matching it
    is matching a number no line of its program mentions, which is a knowledge gap and
    not a decision to overrule -- so the number is in the message, not just the detail.
    """
    path, _ = _titled(tmp_path, CAPTION, caption_pt=44.0)

    found = outranked_findings(path, _pair_spans(44.0, 44.0))

    assert found[0].detail["inherited_pt"] == 44.0
    assert "states no size of its own" in found[0].message
    assert "44pt through the layout and the master" in found[0].message


def test_a_line_the_author_deliberately_set_larger_is_left_alone(tmp_path: Path) -> None:
    """A pull-quote, a statistic, a section numeral -- four template pages do this."""
    path, _ = _titled(tmp_path, "662+", caption_pt=54.0)

    assert outranked_findings(path, _pair_spans(44.0, 54.0, "662+")) == []


def test_a_label_beside_the_title_is_not_a_second_heading_row(tmp_path: Path) -> None:
    """A narrow gloss under a full-width title reads as one heading pair, not two."""
    path, _ = _titled(tmp_path, "Agenda", caption_pt=44.0, caption_width=2.5)

    assert outranked_findings(path, _pair_spans(44.0, 44.0, "Agenda")) == []


def test_two_lines_of_body_type_are_not_a_heading_hierarchy(tmp_path: Path) -> None:
    """A page with no title at all: `_heading_rows` infers one, and a 12pt line is not it.

    One delivered page put a 12pt chart footnote 0.04in above its 12pt caption, and read
    as a title being outranked by the line under it.
    """
    path, _ = _titled(tmp_path, CAPTION, caption_pt=12.0, title_pt=12.0)

    assert outranked_findings(path, _pair_spans(12.0, 12.0)) == []


def test_the_advice_turns_on_whose_box_carries_the_size(tmp_path: Path) -> None:
    """Same tie, opposite advice, which is what `template` is for.

    Every tie in the evidence tree is on a box the template drew, so telling the author
    to stop choosing 44pt would be telling it about a choice it never made. A box the
    author drew is the other case and the cheap one: two numbers it can both see.
    """
    path, _ = _titled(tmp_path, CAPTION, caption_pt=44.0)
    spans = _pair_spans(44.0, 44.0)

    theirs = outranked_findings(path, spans, path)
    mine = outranked_findings(path, spans, _deck(tmp_path, [("elsewhere", 6.0, 6.0, 2.0, 0.4)], name="other"))

    assert theirs[0].detail["under_is_the_templates_box"] is True
    assert "do not touch the title" in theirs[0].message
    assert mine[0].detail["under_is_the_templates_box"] is False
    assert "You drew this box and chose this size" in mine[0].message


def test_without_a_render_neither_pair_reading_answers(tmp_path: Path) -> None:
    """No spans is no signal, and a check with no signal says nothing."""
    path, _ = _titled(tmp_path, CAPTION, caption_pt=44.0)

    assert outranked_findings(path, None) == []
    assert outranked_findings(path, []) == []
    assert row_findings(path, None) == []


CARDS = [
    "Systems in service",
    "Route kilometres",
    "Aging fleet",
]


def _row(tmp_path: Path, sizes: list[float] | None, *, declared: float | None = 20.0, name: str = "row") -> Path:
    """One page carrying three copies of a one-box unit, which is what `units` groups.

    `sizes` is unused here and named in `_row_spans`: the file states one size for all
    three, and what differs is what the renderer did with it.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, text in enumerate(CARDS):
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Inches(1.0 + 3.0 * index), Inches(4.0), Inches(2.0), Inches(0.6))
        if declared is None:
            box.text_frame.text = text
            continue
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(declared)
    path = tmp_path / f"{name}.pptx"
    presentation.save(str(path))
    return path


def _row_spans(sizes: list[float]) -> list[Span]:
    """One span per card, landing inside the box `_row` drew for it."""
    found = []
    for index, (text, size) in enumerate(zip(CARDS, sizes)):
        left = 72.0 * (1.0 + 3.0 * index) + 6
        found.append(Span(page=1, size_pt=size, text=text, x0=left, y0=295, x1=left + 100, y1=295 + size))
    return found


def test_a_row_whose_members_came_out_at_different_sizes_is_reported(tmp_path: Path) -> None:
    """The finding is the row: every box behaved, and the row still reads wrong."""
    path = _row(tmp_path, None)

    found = row_findings(path, _row_spans([16.8, 16.8, 20.0]))

    assert [finding.kind for finding in found] == ["row_type_drift"]
    assert found[0].severity is Severity.WARNING
    assert found[0].detail["sizes_pt"] == [16.8, 16.8, 20.0]
    assert found[0].detail["units"] == 3
    assert found[0].detail["spread"] == round(20.0 / 16.8, 3)
    assert "the row is what reads wrong" in found[0].message


def test_a_row_that_shrank_evenly_is_not_a_row_a_reader_complains_about(tmp_path: Path) -> None:
    """Where this reading and `type_drift` part company, on one delivered page each.

    Three boxes drawn at 20pt and all rendered at 17pt are a row nobody can tell apart;
    `type_drift` reports all three of them against the size the slot was drawn at, which
    is the question "did this box shrink" and not the question a reader is asking.
    """
    path = _row(tmp_path, None)
    spans = _row_spans([17.0, 17.0, 17.0])

    assert row_findings(path, spans) == []
    assert [finding.kind for finding in drift_findings(path, spans)] == ["type_drift"]


def test_a_member_reading_larger_than_its_own_size_is_a_misattribution(tmp_path: Path) -> None:
    """Autofit only shrinks, so a box cannot render above the size its runs state.

    One delivered page's 14pt body box read 30pt, off the number tile stacked beside it,
    and the row came back with a spread of 2.1.
    """
    path = _row(tmp_path, None)

    assert row_findings(path, _row_spans([20.0, 20.0, 30.0])) == []


def test_a_page_with_nothing_repeating_has_no_row_to_read(tmp_path: Path) -> None:
    found = row_findings(
        _deck(tmp_path, [("a line of copy on its own", 1.0, 1.0, 4.0, 0.5)], size=20.0),
        [Span(page=1, size_pt=20.0, text="a line of copy on its own", x0=80, y0=80, x1=300, y1=100)],
    )

    assert found == []


# What `_heading_rows` answers is the role the template named, not where it sits: a
# SUBTITLE placeholder set above the title comes back as the subtitle on purpose, because
# a kicker is still the row the template called that. So the pair reading has to ask about
# the relation itself, and both halves of the relation are in the message it prints.
KICKER = "Q3"
SHORT = "Q3 2026 review"


def _kickered(tmp_path: Path, *, kicker: str = KICKER, name: str = "kicker") -> Path:
    """One page whose SUBTITLE placeholder sits above its title, both full width.

    The template idiom the geometry guard is about: a section kicker over the title,
    which `_heading_rows` returns as the subtitle and which is not the line under it.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(0.5), Inches(2.0), Inches(9.0), Inches(1.0)
    run = title.text_frame.paragraphs[0].add_run()
    run.text = TITLE
    run.font.size = Pt(44.0)
    above = next(shape for shape in slide.placeholders if shape.placeholder_format.idx == 1)
    above.left, above.top, above.width, above.height = Inches(0.5), Inches(0.5), Inches(9.0), Inches(0.8)
    kick = above.text_frame.paragraphs[0].add_run()
    kick.text = kicker
    kick.font.size = Pt(44.0)
    path = tmp_path / f"{name}.pptx"
    presentation.save(str(path))
    return path


def _kicker_spans(kicker: str = KICKER) -> list[Span]:
    return [
        Span(page=1, size_pt=44.0, text=TITLE, x0=40, y0=150, x1=560, y1=194),
        Span(page=1, size_pt=44.0, text=kicker, x0=40, y0=40, x1=90, y1=84),
    ]


def test_a_kicker_set_over_the_title_is_not_the_line_under_it(tmp_path: Path) -> None:
    """The finding is about the line *under* the title, and this one is above it.

    Read only for width and size, a full-width two-character kicker over a title at one
    size reported the title as outranked -- and said the kicker was "the line under it"
    and "the longer of the two (2 characters against 24)" in one sentence.
    """
    assert outranked_findings(_kickered(tmp_path), _kicker_spans()) == []


def test_a_title_longer_than_the_line_under_it_has_not_been_outranked(tmp_path: Path) -> None:
    """At one size the eye goes to the longer line, so a longer title still reads as one.

    Two pages of the evidence tree are this shape: an 83-character title over a
    51-character caption, both at 28pt.
    """
    path, _ = _titled(tmp_path, SHORT, caption_pt=44.0)

    assert outranked_findings(path, _pair_spans(44.0, 44.0, SHORT)) == []


HEADINGS = ["Systems in service", "Aging fleet"]
BODIES = ["Route kilometres and landings across the network", "Average airframe age against the benchmark"]


def _two_level(tmp_path: Path, *, name: str = "two_level") -> Path:
    """Two cards, each holding a 3x0.6in heading over a separate 3x0.6in body, both 20pt.

    One box shape and one stated size across two levels of one unit, which is the
    ordinary way a card is drawn rather than a contrived one.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, (heading, body) in enumerate(zip(HEADINGS, BODIES)):
        group = slide.shapes.add_group_shape()
        for row, text in ((0, heading), (1, body)):
            box = group.shapes.add_textbox(Inches(1.0 + 4.0 * index), Inches(3.0 + 1.0 * row), Inches(3.0), Inches(0.6))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = text
            run.font.size = Pt(20.0)
    path = tmp_path / f"{name}.pptx"
    presentation.save(str(path))
    return path


def _two_level_spans(headings: list[float], bodies: list[float]) -> list[Span]:
    found = []
    for index, (heading, body) in enumerate(zip(HEADINGS, BODIES)):
        left = 72.0 * (1.0 + 4.0 * index) + 6
        found.append(Span(page=1, size_pt=headings[index], text=heading, x0=left, y0=220, x1=left + 120, y1=240))
        found.append(Span(page=1, size_pt=bodies[index], text=body, x0=left, y0=292, x1=left + 120, y1=306))
    return found


def test_two_levels_of_one_unit_drawn_alike_are_still_two_rows(tmp_path: Path) -> None:
    """The headings agree with each other and the bodies agree with each other.

    Keyed on the box and the stated size alone, all four boxes merged into one level and
    reported a 1.429 spread over a run where neither row drifts at all.
    """
    path = _two_level(tmp_path)

    assert row_findings(path, _two_level_spans([20.0, 20.0], [14.0, 14.0])) == []


def test_the_level_that_drifts_is_the_one_named_when_two_are_drawn_alike(tmp_path: Path) -> None:
    """And the split does not cost the reading anything: the uneven level still reports."""
    path = _two_level(tmp_path)

    found = row_findings(path, _two_level_spans([20.0, 20.0], [14.0, 10.0]))

    assert [finding.kind for finding in found] == ["row_type_drift"]
    assert found[0].detail["sizes_pt"] == [10.0, 14.0]
    assert found[0].detail["slot_at"] == 2
    assert found[0].detail["spread"] == 1.4
    assert "2nd slot down each unit" in found[0].message


def test_both_readings_record_the_boxes_they_are_about(tmp_path: Path) -> None:
    """The one identity `quiet` can compare across the two type baselines.

    A level here is a slot inside a repeating unit and `type_drift`'s group is a shape
    repeated anywhere in the deck, so neither one's own key means anything to the other.
    """
    path = _two_level(tmp_path)
    spans = _two_level_spans([20.0, 20.0], [14.0, 10.0])

    row = row_findings(path, spans)[0]
    drift = drift_findings(path, spans)[0]

    assert len(row.detail["boxes"]) == 2
    assert all(len(box) == 4 for box in row.detail["boxes"])
    assert row.detail["boxes"] == sorted(row.detail["boxes"])
    assert drift.detail["boxes"] == row.detail["boxes"]


# And the other direction, which pulls against the two above: a level has to keep its
# position when a unit leaves an optional slot unspoken. `adapt` empties every frame it
# was not told about, so a card that supplied only its body arrives holding an empty
# label frame the card beside it filled.
OPTIONAL = "Optional label"
LOOSE_BODIES = ["Route kilometres flown", "Average airframe age against the industry benchmark"]


def _optional_row(tmp_path: Path, *, name: str = "optional") -> Path:
    """Two cards, an optional 3x0.4in label over an identical 3x0.6in body slot.

    The first card supplied only its body; the second supplied both.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, body in enumerate(LOOSE_BODIES):
        group = slide.shapes.add_group_shape()
        label = group.shapes.add_textbox(Inches(1.0 + 4.0 * index), Inches(2.6), Inches(3.0), Inches(0.4))
        if index:
            spoken = label.text_frame.paragraphs[0].add_run()
            spoken.text = OPTIONAL
            spoken.font.size = Pt(14.0)
        box = group.shapes.add_textbox(Inches(1.0 + 4.0 * index), Inches(3.2), Inches(3.0), Inches(0.6))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = body
        run.font.size = Pt(20.0)
    path = tmp_path / f"{name}.pptx"
    presentation.save(str(path))
    return path


def _optional_spans(bodies: list[float]) -> list[Span]:
    found = [Span(page=1, size_pt=14.0, text=OPTIONAL, x0=366, y0=192, x1=452, y1=206)]
    for index, (body, size) in enumerate(zip(LOOSE_BODIES, bodies)):
        left = 72.0 * (1.0 + 4.0 * index) + 6
        found.append(Span(page=1, size_pt=size, text=body, x0=left, y0=238, x1=left + 120, y1=238 + size))
    return found


def test_an_optional_slot_one_unit_left_empty_keeps_the_row_comparable(tmp_path: Path) -> None:
    """The position is the slot's, not the copy's, or the row stops being a row.

    Ranked over the frames holding words, the first card's body came second in its unit
    and the second card's body third, so a row running 20pt against 14pt reported
    nothing -- a false silence on a layout `adapt` produces routinely.
    """
    path = _optional_row(tmp_path)

    found = row_findings(path, _optional_spans([20.0, 14.0]))

    assert [finding.kind for finding in found] == ["row_type_drift"]
    assert found[0].detail["sizes_pt"] == [14.0, 20.0]
    assert found[0].detail["slot_at"] == 2
    assert found[0].detail["spread"] == round(20.0 / 14.0, 3)


def test_an_optional_slot_left_empty_everywhere_is_furniture_and_holds_no_position(
    tmp_path: Path,
) -> None:
    """What keeps the guard above from counting an icon's container as a slot.

    A shape no unit of the run puts copy in is a spacer or an icon frame, and it never
    had a position. One template's zigzag timeline draws its icon container above the
    heading in one of five units and below it in the other four; counted as a position,
    the odd card falls out of that row and the finding is left reporting four of five.
    """
    path = _optional_row(tmp_path)
    spans = [span for span in _optional_spans([20.0, 14.0]) if span.text != OPTIONAL]
    from pptx import Presentation

    presentation = Presentation(path)
    for group in presentation.slides[0].shapes:
        for shape in group.shapes:
            if shape.text_frame.text.strip() == OPTIONAL:
                shape.text_frame.paragraphs[0].runs[0].text = ""
    emptied = tmp_path / "furniture.pptx"
    presentation.save(str(emptied))

    found = row_findings(emptied, spans)

    assert [finding.kind for finding in found] == ["row_type_drift"]
    assert found[0].detail["slot_at"] == 1
    assert found[0].detail["sizes_pt"] == [14.0, 20.0]
