"""What the model sees when the user hands over a template.

Two calls, because the author has two questions. "What does this design look
like" is answered with pictures -- 85% of a template's visual elements are on its
example pages, so a list of layout names answers it barely at all. "How do I draw
a page in it" is answered with code, because an author given the loop that drew
six cards can draw eight, and an author given a description of them cannot.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")

from raven.contracts.tool import ToolResult  # noqa: E402
from raven_ppt.contracts import Project  # noqa: E402
from raven_ppt.services.render import RenderError, sheet  # noqa: E402
from raven_ppt.services.template.defaults import reference_pages  # noqa: E402
from raven_ppt.tools.template import PptTemplateTool  # noqa: E402
from tests._ppt_engine_fixtures import deck, image, noise_image, noise_png, product_page, template_file  # noqa: F401

_TEMPLATES = Path(__file__).resolve().parents[1] / "plugins-dist" / "ppt-engine" / "raven_ppt" / "assets" / "templates"
_needs_templates = pytest.mark.skipif(
    not any(_TEMPLATES.glob("*.pptx")),
    reason="the template payload is fetched, not tracked; see plugins-dist/ppt-engine/templates.manifest.json",
)


class FakeViews:
    """Renders without LibreOffice. `pages` is asked for the template's own file.

    The pages are real PNGs and the tiling is the real tiler: what this reply carries
    is one sheet of them, so a stand-in that could not be composed would test the
    fallback and nothing else. `tiling=False` is how the fallback is asked for.
    """

    def __init__(self, available: bool = True, tiling: bool = True) -> None:
        self.available = available
        self.tiling = tiling
        self.rendered: list[Path] = []
        self.sheets: list[Path] = []

    async def pages(self, pptx: Path, out_dir: Path, numbers) -> dict[int, Path]:
        self.rendered.append(pptx)
        return await self.pages_of(pptx, out_dir, numbers)

    async def pdf(self, pptx: Path, out_dir: Path) -> Path | None:
        """The conversion the tool now pays for once and reads twice: pictures for the
        author, type sizes for the house style."""
        self.rendered.append(pptx)
        if not self.available:
            return None
        out_dir.mkdir(parents=True, exist_ok=True)
        pdf = out_dir / f"{pptx.stem}.pdf"
        pdf.write_bytes(b"%PDF-1.4")
        return pdf

    async def pages_of(self, pdf: Path, out_dir: Path, numbers) -> dict[int, Path]:
        if not self.available:
            return {}
        out_dir.mkdir(parents=True, exist_ok=True)
        made = {}
        for number in numbers or ():
            path = out_dir / f"template-{number:02d}.png"
            _page_png(path)
            made[number] = path
        return made

    def contact_sheet(self, pngs, out: Path, columns: int = 4, *, labels=None) -> Path:
        if not self.tiling:
            raise RenderError("this host cannot tile")
        out.parent.mkdir(parents=True, exist_ok=True)
        made = sheet.contact_sheet(list(pngs), out, columns, labels=labels)
        self.sheets.append(made)
        return made

    def data_uri(self, png: Path, budget: int | None = None) -> str:
        return f"data:image/png;base64,{png.stem}"

    def sheet_uri(self, png: Path) -> str:
        return f"data:image/jpeg;base64,{png.stem}"


def _page_png(path: Path, size: tuple[int, int] = (320, 180)) -> Path:
    image = pytest.importorskip("PIL.Image", reason="Pillow tiles the pages into a sheet")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.new("RGB", size, (210, 210, 215)).save(path, "PNG")
    return path


@pytest.fixture
def workspace(tmp_path: Path) -> Path:
    (tmp_path / "uploads").mkdir()
    return tmp_path


@pytest.fixture
def house(workspace: Path, template_file) -> Path:
    return template_file(where=workspace / "uploads")


def _body(result) -> dict:
    return json.loads(result.model_text if isinstance(result, ToolResult) else result)


async def test_binding_a_template_by_path(workspace: Path, house: Path):
    """The user's entry point: they say "use this deck", and this is where the
    path lands. Everything downstream then asks one question -- is there a
    template -- rather than where it came from."""
    from raven_ppt.services.template import bound

    tool = PptTemplateTool(workspace, FakeViews())

    body = _body(await tool.execute(project="talk", path="uploads/house-style.pptx"))

    assert body["ok"]
    assert body["template"] == "house-style.pptx"
    assert body["example_pages"] == 2
    assert body["canvas_in"].startswith("13.33")
    assert body["layouts"]
    assert bound(Project(workspace=workspace, slug="talk")) is not None


async def test_binding_refreshes_the_script_workspace(workspace: Path, house: Path):
    marker = workspace / "provisioned"
    tool = PptTemplateTool(workspace, FakeViews(), provision=lambda _project: marker.write_text("ready"))

    await tool.execute(project="talk", path="uploads/house-style.pptx")

    assert marker.read_text() == "ready"


async def test_the_users_file_is_copied_not_taken(workspace: Path, house: Path):
    before = house.read_bytes()
    tool = PptTemplateTool(workspace, FakeViews())

    await tool.execute(project="talk", path="uploads/house-style.pptx")

    assert house.read_bytes() == before


async def test_all_example_pages_come_back_as_pictures(workspace: Path, house: Path):
    """Every example page is available so the model can choose and adapt a close fit.

    On one sheet, with the page numbers on the cells and the shape of every page in the
    text above it. Measured on the model this engine runs: the sheet and 20 separate
    renders both score 90% top-1 at naming the page whose arrangement fits a need, and
    the sheet does it with one image pinned in the request afterwards instead of 20 --
    the cost a recorded run died of.
    """
    views = FakeViews()
    tool = PptTemplateTool(workspace, views)

    result = await tool.execute(project="talk", path="uploads/house-style.pptx")

    assert isinstance(result, ToolResult)
    assert [block["type"] for block in result.blocks[:2]] == ["text", "image_url"]
    assert result.blocks[1]["image_url"]["url"].startswith("data:image/jpeg;base64,"), (
        "the sheet goes through the sheet door"
    )
    made = [path.name for path in views.sheets if path.name.startswith("pages_sheet_")]
    assert made == ["pages_sheet_2.png"], "two example pages arrive as one picture, not two"
    legend = result.blocks[0]["text"]
    assert "The 2 example pages are one picture below, 4 to a row" in legend
    assert "Template page 1 -- the template's cover" in legend
    assert "Template page 2 -- a content example" in legend
    assert "prototype(tpl, 2)" in legend, "the call that takes the page is in the key to the sheet"
    body = _body(result)
    assert body["house_pages"] == {"cover": 1}
    assert len(body["template_pages"]) == 2
    next_step = body["next_step"]
    assert next_step.startswith("clone the template's structural pages")
    # Named page by page and by the shape each page holds. Prose about "adaptable
    # prototypes" left a live run cloning only the pages the reply named by role.
    assert body["content_pages"].startswith("1 of the template's pages are content examples")
    assert "Clone the one whose arrangement matches" in body["content_pages"]


async def test_a_machine_that_cannot_render_still_binds(workspace: Path, house: Path):
    """LibreOffice is optional everywhere else in this package and stays optional
    here: without it the author reads the pages as code instead of looking at
    them, which is a worse workflow and not a broken one."""
    tool = PptTemplateTool(workspace, FakeViews(available=False))

    body = _body(await tool.execute(project="talk", path="uploads/house-style.pptx"))

    assert body["ok"]
    assert "code" in body["renders"]


async def test_the_binding_holds_without_the_path(workspace: Path, house: Path):
    tool = PptTemplateTool(workspace, FakeViews())
    await tool.execute(project="talk", path="uploads/house-style.pptx")

    body = _body(await tool.execute(project="talk"))

    assert body["ok"]
    assert body["template"] == "house-style.pptx"


async def test_asking_for_a_template_nobody_bound(workspace: Path):
    tool = PptTemplateTool(workspace, FakeViews())

    body = _body(await tool.execute(project="talk"))

    assert not body["ok"]
    assert "no template is bound" in body["error"]
    assert "ppt_template with the path" in body["hint"]


async def test_a_path_outside_the_workspace_is_refused(workspace: Path, tmp_path: Path, template_file):
    """Every path that arrives from a model is checked. This one is a file the
    user names, which makes it exactly the path worth checking."""
    outside = template_file(name="elsewhere.pptx", where=tmp_path.parent)
    tool = PptTemplateTool(workspace, FakeViews())

    body = _body(await tool.execute(project="talk", path=f"../{outside.name}"))

    assert not body["ok"]
    assert "could not be used as a template" in body["error"]


async def test_something_that_is_not_a_deck_is_refused(workspace: Path):
    (workspace / "uploads" / "notes.txt").write_text("not a deck")
    tool = PptTemplateTool(workspace, FakeViews())

    body = _body(await tool.execute(project="talk", path="uploads/notes.txt"))

    assert not body["ok"]
    assert ".pptx" in body["hint"]


async def test_pages_come_back_as_source_the_author_can_paste(workspace: Path, house: Path):
    tool = PptTemplateTool(workspace, FakeViews())
    await tool.execute(project="talk", path="uploads/house-style.pptx")

    result = await tool.execute(project="talk", pages=[1])

    assert isinstance(result, ToolResult)
    code = result.blocks[0]["text"]
    assert "# page 0 of the template" in code
    assert "slide.shapes.add_textbox" in code
    assert _body(result)["pages_read"] == [1]


async def test_a_page_read_carries_the_pages_and_not_the_roster_again(workspace: Path, house: Path):
    tool = PptTemplateTool(workspace, FakeViews())
    bind = _body(await tool.execute(project="talk", path="uploads/house-style.pptx"))
    assert bind["layouts"] and bind["canvas_in"]

    read = _body(await tool.execute(project="talk", pages=[1]))

    assert read["pages_read"] == [1]
    for key in ("layouts", "canvas_in", "example_pages", "build_from", "theme_colours", "fonts"):
        assert key not in read, key


async def test_a_call_that_binds_and_reads_in_one_go_keeps_the_roster(workspace: Path, house: Path):
    tool = PptTemplateTool(workspace, FakeViews())

    body = _body(await tool.execute(project="talk", path="uploads/house-style.pptx", pages=[1]))

    assert body["pages_read"] == [1]
    assert body["layouts"] and body["canvas_in"] and body["example_pages"]


async def test_a_pages_pictures_land_where_its_code_looks_for_them(workspace: Path, house: Path):
    """`add_picture("template_00.png", ...)` in the reference has to be a line the
    author can paste and run, and the program runs in the build directory."""
    tool = PptTemplateTool(workspace, FakeViews())
    await tool.execute(project="talk", path="uploads/house-style.pptx")

    result = await tool.execute(project="talk", pages=[1])

    build = Project(workspace=workspace, slug="talk").build_dir
    named = [line for line in result.blocks[0]["text"].splitlines() if "add_picture(" in line]
    assert named
    assert (build / named[0].split('"')[1]).is_file()


async def test_a_page_that_cannot_be_redrawn_says_so_and_names_the_way_round(workspace: Path, house: Path):
    """Two thirds of real template pages hold something python-pptx cannot write.
    The reply names those pages and the operation that gets them into the deck
    anyway, because an author that only sees code assumes code is enough."""
    tool = PptTemplateTool(workspace, FakeViews())
    await tool.execute(project="talk", path="uploads/house-style.pptx")

    body = _body(await tool.execute(project="talk", pages=[1]))

    assert body["cannot_be_redrawn"]["1"]
    # And the operation it names is the one the author is taught everywhere else. Two
    # doors onto a template page cost a live deck eight blank slides, so a reply that
    # answered "you cannot redraw this" with a second vocabulary was the last place a
    # second route could hide.
    assert "clone_page(prs, prototype(tpl, N))" in body["next_step"]
    assert "adapt(prs" not in body["next_step"]


def _with_a_chart(path: Path) -> Path:
    """The fixture template with a native chart added to its first example page."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation(str(path))
    data = CategoryChartData()
    data.categories = ["Q1 of the template", "Q2 of the template"]
    data.add_series("the template's own series", (19.2, 21.5))
    presentation.slides[0].shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(4.5), Inches(6), Inches(2.5), data
    )
    presentation.save(str(path))
    return path


async def test_a_template_chart_is_the_authors_to_draw_and_says_where(workspace: Path, house: Path):
    """The advice this replaced sent a run to clone a chart page and replace_text it.
    Nothing writes chart data here, so the clone arrived holding the template's own
    categories; the run then spent 73 minutes in a shell measuring the template's axis
    type by hand. The arrangement is still what the page is worth cloning for, so the
    page stays on offer and the offer carries the box the chart leaves behind."""
    tool = PptTemplateTool(workspace, FakeViews())
    await tool.execute(project="talk", path="uploads/house-style.pptx")
    _with_a_chart(house)
    await tool.execute(project="talk", path="uploads/house-style.pptx")

    body = _body(await tool.execute(project="talk", pages=[1]))

    assert body["charts_to_draw"]["1"] == ["Inches(6.00), Inches(4.50), Inches(6.00), Inches(2.50)"]
    assert "a chart" not in body.get("cannot_be_redrawn", {}).get("1", [])
    said = body["next_step"]
    assert "the chart must be redrawn with ppt_charts from your own data" in said
    assert "placed in the same position" in said
    assert "ppt_build" in said, "the sentence that told an author to write code named no tool"


async def test_the_page_the_author_chooses_from_says_the_chart_is_its_own(workspace: Path, house: Path):
    """The prototype is chosen off the renders, so the line beside the render is where
    the offer is made. A page named there as an arrangement and nothing else is cloned
    whole, chart included."""
    _with_a_chart(house)
    tool = PptTemplateTool(workspace, FakeViews())

    result = await tool.execute(project="talk", path="uploads/house-style.pptx")

    legend = "\n".join(block["text"] for block in result.blocks if block.get("type") == "text")
    assert "the chart must be redrawn with `ppt_charts` from your own data" in legend
    assert "Inches(6.00), Inches(4.50), Inches(6.00), Inches(2.50)" in legend
    assert "The chart on it is the template's own" in legend


async def test_asking_for_pages_the_template_does_not_have(workspace: Path, house: Path):
    tool = PptTemplateTool(workspace, FakeViews())
    await tool.execute(project="talk", path="uploads/house-style.pptx")

    body = _body(await tool.execute(project="talk", pages=[40]))

    assert not body["ok"]
    assert "2 example pages" in body["hint"]


async def test_a_project_name_that_is_not_one(workspace: Path):
    tool = PptTemplateTool(workspace, FakeViews())

    body = _body(await tool.execute(project="Not A Slug", path="uploads/x.pptx"))

    assert not body["ok"]
    assert "not a usable project name" in body["error"]


async def test_the_reply_names_the_layouts_that_carry_pictures(workspace: Path, house: Path, image):
    """A photograph on the layout shows on every page and is on none of them, so the
    author told to replace the template's pictures found nothing on the page to replace.
    The reply names the layout and the call that reaches it."""
    from pptx import Presentation

    from tests._ppt_engine_fixtures import layout_picture

    presentation = Presentation(str(house))
    layout = presentation.slides[0].slide_layout
    layout_picture(layout, image("cover.png", (60, 60, 60)), 0, 0, 6.4, 4.7)
    presentation.save(str(house))

    tool = PptTemplateTool(workspace, views=FakeViews())
    result = await tool.execute(project="deck", path=str(house))

    body = _body(result)
    assert body["layout_pictures"] == [
        f"layout '{layout.name}' carries 1 picture(s) (6.4x4.7in), under example page(s) 1, 2 -- on the page you "
        "build from one of those: `replace_picture(layout_pictures(slide)[0], FIGURES / 'x.png', 'cover')`"
    ], "a picture short of the page is swapped at full strength; a page-sized one is a background at alpha=0.1"

    assert "layout_pictures(slide)[0]" in body["next_step"]


async def test_a_page_larger_than_the_whole_budget_is_still_bounded(workspace: Path, house: Path, monkeypatch):
    """The budget used to apply from the second page on, so a single page of 19,000
    characters crossed the host's 16,000-character mark and came back cut mid-line,
    reading as a complete page. It stops at a whole line inside the budget, says so
    where it stops, and the reply names the page as cut."""
    from raven_ppt.tools import template as module

    monkeypatch.setattr(module, "SOURCE_BUDGET_CHARS", 400)
    tool = PptTemplateTool(workspace, FakeViews())
    await tool.execute(project="talk", path="uploads/house-style.pptx")

    result = await tool.execute(project="talk", pages=[1])

    code = result.blocks[0]["text"]
    body = _body(result)
    assert body["pages_read"] == [1]
    assert body["pages_cut"]["1"]["lines_withheld"] > 0
    assert "# -- cut:" in code and "clone" in code
    kept = code.split("# -- cut:")[0]
    assert len(kept) <= 400 + len("\n".join(f"# {line}" for line in ["x"] * 8)) + 300, (
        "the page stops inside the budget, plus the imports"
    )
    assert "clone" in body["next_step"]


async def test_rebinding_drops_the_band_grid_read_off_the_previous_template(workspace: Path, house: Path):
    """`bands.json` is a reading of one template's pages, kept beside it the way the
    palette is; a rebind dropped the palette and not the grid, so every band check then
    judged the new deck against the previous template's rows (a 37% error in the body
    area on two synthetic templates differing only in title-row height)."""
    from raven_ppt.services.template.bands import bands_path

    tool = PptTemplateTool(workspace, FakeViews())
    project = Project(workspace=workspace, slug="talk")

    await tool.execute(project="talk", path="uploads/house-style.pptx")
    bands_path(project).write_text('{"schema": "stale", "title_bottom": 1.0}', encoding="utf-8")

    body = _body(await tool.execute(project="talk", path="uploads/house-style.pptx"))

    assert body["ok"]
    assert not bands_path(project).exists(), "the grid goes with the template it was read from"


def test_the_render_caption_names_the_picture_slots_and_leaves_the_fill_to_the_author() -> None:
    """A template page's photograph and cartoon are placeholders; the caption says where
    they are in `shape_at`'s numbering and what may go there, and the starter call carries
    `pictures={...}`, so the choice of picture is made at the render, not after a refusal."""
    from types import SimpleNamespace

    from raven_ppt.tools.template import _render_label

    entry = SimpleNamespace(
        arrangement="a photograph beside three cards",
        slots=3,
        picture_slots=("[2] 5.4x3.6in photo", "[7] 2.4x2.5in drawing"),
    )

    said = _render_label(9, None, entry)

    assert "Picture slots: [2] 5.4x3.6in photo, [7] 2.4x2.5in drawing" in said
    assert "transparent=true" in said and "yours to decide" in said
    assert "`remove_unit` for the slots it does not fill" in said
    assert "`replace_picture` for the figure" in said

    bare = _render_label(4, None, SimpleNamespace(arrangement="two columns", slots=0, picture_slots=()))
    assert "Picture slots" not in bare and "pictures={...}" not in bare

    marked = _render_label(
        4, None, SimpleNamespace(arrangement="three cards", slots=3, picture_slots=("[10] 1.7x1.7in icon",))
    )
    assert "An icon slot" in marked and "one per unit" in marked and "swap_icon(slide, shape_at(slide, n)" in marked
    assert "An icon slot" not in said, "said beside the page that has one, and nowhere else"


async def test_a_host_that_cannot_tile_gets_one_picture_per_page(workspace: Path, house: Path):
    """The sheet is the cheaper reply, not the only one.

    Pillow composes the sheet and a host without it still has pages to show. A reply
    with no renders in it is the one outcome to avoid: one live run read every example
    page as code and never asked for a picture again, on a host that could render.
    """
    views = FakeViews(tiling=False)
    tool = PptTemplateTool(workspace, views)

    result = await tool.execute(project="talk", path="uploads/house-style.pptx")

    assert isinstance(result, ToolResult)
    assert [block["type"] for block in result.blocks[:4]] == ["text", "image_url", "text", "image_url"]
    assert result.blocks[1]["image_url"]["url"].startswith("data:image/png;base64,"), (
        "a page comes through the page door"
    )
    assert result.blocks[0]["text"].startswith("Template page 1 -- the template's cover")
    assert result.blocks[2]["text"].startswith("Template page 2 -- a content example")


@_needs_templates
async def test_a_complete_borrow_sheet_is_found_before_anything_is_rasterised(workspace: Path, house: Path):
    """The ordinary follow-up call must not pay for the pictures again.

    Naming the sheet after what rendered put the cache lookup after the render, and
    `pages_of` rasterises every page it is handed whether or not a sheet already holds
    them -- 9.9s of the 41.3s measured over the seven templates, on the path a
    palette-only follow-up takes. The complete sheet's name is knowable up front, so
    it is looked for there; a sheet short of the offer deliberately has no fast path,
    because falling through is what asks the failed lender again.
    """
    views = FakeViews()
    tool = PptTemplateTool(workspace, views)
    lenders = {stem for stem, _, _ in PptTemplateTool._borrowable(house)}

    await tool.execute(project="talk", path="uploads/house-style.pptx")
    borrowed = [path for path in views.rendered if path.stem in lenders]
    assert borrowed, "the first call renders the lenders"
    composed = [path.name for path in views.sheets if path.name.startswith("sheet_")]
    assert composed, "and composes their sheet"

    await tool.execute(project="talk", path="uploads/house-style.pptx")

    assert [path for path in views.rendered if path.stem in lenders] == borrowed, (
        "the second call rasterised the lenders again"
    )
    assert [path.name for path in views.sheets if path.name.startswith("sheet_")] == composed, (
        "and composed their sheet again"
    )


@_needs_templates
async def test_a_lender_that_will_not_render_does_not_leave_a_sheet_claiming_its_page(
    workspace: Path, house: Path, monkeypatch: pytest.MonkeyPatch
):
    """A partial sheet used to be cached under the whole offer's name.

    One lender that cannot be converted leaves the other lenders' cells intact, and
    that sheet is worth showing -- but stored as the sheet of every offered page it
    claimed cells it did not have, both sentences beside it counted the offer rather
    than the picture, and every later call read it out of the cache, so the lender that
    failed was never asked again. Keyed on what rendered instead: the reply counts the
    cells it has, and the call after the lender comes back names a different sheet.
    """
    from raven_ppt.tools import template as template_module

    views = FakeViews()
    tool = PptTemplateTool(workspace, views)
    offers = template_module.PptTemplateTool._borrowable(house)
    stems = list(dict.fromkeys(stem for stem, _, _ in offers))
    assert len(stems) > 1, "the fixture must offer pages from more than one template"
    broken, working = stems[0], stems[1:]
    offered_by_broken = [(stem, number) for stem, number, _ in offers if stem == broken]
    assert offered_by_broken, "the lender chosen to fail must have an offer"

    real_pdf = views.pdf
    failing = {broken}

    async def pdf(pptx: Path, out_dir: Path):
        return None if pptx.stem in failing else await real_pdf(pptx, out_dir)

    monkeypatch.setattr(views, "pdf", pdf)
    monkeypatch.setattr(
        template_module.PptTemplateTool,
        "_borrowable",
        staticmethod(lambda source: offers),
    )

    result = await tool.execute(project="talk", path="uploads/house-style.pptx")

    body = _body(result)
    pictured = len(offers) - len(offered_by_broken)
    borrow = [path.name for path in views.sheets if path.name.startswith("sheet_")]
    assert borrow == [f"sheet_{pictured}_" + borrow[0].split("_")[2]], borrow
    assert not borrow[0].startswith(f"sheet_{len(offers)}_"), "a partial sheet is not the whole offer's sheet"
    assert f"shows {pictured} of them" in body["next_step"]
    assert f"the other {len(offered_by_broken)} did not render" in body["next_step"]
    assert f"shows all {len(offers)}" not in body["next_step"]
    said = next(
        block["text"] for block in result.blocks if "pages offered from the other bundled" in block.get("text", "")
    )
    assert f"And {pictured} of the {len(offers)} pages offered" in said
    assert "the rest did not render here" in said
    assert all(stem in str(views.rendered) for stem in working[:1]), "the lenders that work are still rendered"

    # The lender comes back: the sheet it could not be part of is built now, under its
    # own name, rather than served from the cache of the partial one.
    failing.clear()
    views.sheets.clear()
    again = await tool.execute(project="talk", path="uploads/house-style.pptx")

    names = [path.name for path in views.sheets if path.name.startswith("sheet_")]
    assert names and names[0].startswith(f"sheet_{len(offers)}_"), names
    assert f"shows all {len(offers)} of them" in _body(again)["next_step"]


@_needs_templates
async def test_the_pages_worth_borrowing_arrive_as_a_picture_of_their_own(workspace: Path, house: Path):
    """The offer the product could not make at any price.

    The borrowable reference pages have always been text -- "<stem> page N is <shape>"
    -- and on the model this engine runs that text answers 27% of a need list the same
    pages on one sheet answer 100%. Their own sheet rather than more cells on the
    template's, because the two answer different questions and a cell whose provenance
    the author cannot read is worse than no cell.
    """
    views = FakeViews()
    tool = PptTemplateTool(workspace, views)

    result = await tool.execute(project="talk", path="uploads/house-style.pptx")

    assert isinstance(result, ToolResult)
    made = sorted(path.name for path in views.sheets)
    assert len(made) == 2, "the template's pages and the borrowable pages are two sheets"
    assert made[0] == "pages_sheet_2.png"
    # Named after which pages are on it: a deck rebound to another bundled template is
    # offered a different set of the same size, and a name keyed on the count alone
    # would serve it the first template's sheet out of the cache. The count comes from
    # the table rather than a literal, because a branch that bundles another template
    # adds rows to it and the shape of the name is what this pins.
    # An uploaded template is bound here, so no bundled stem is excluded from the offer.
    offered = len(reference_pages(except_stem=""))
    assert made[1].startswith(f"sheet_{offered}_") and made[1].endswith(".png"), made[1]
    assert made[1] != f"sheet_{offered}_.png"
    said, shown = result.blocks[-2], result.blocks[-1]
    assert said["type"] == "text" and shown["type"] == "image_url"
    assert "6 to a row" in said["text"] and "their arrangement comes across" in said["text"]
    assert shown["image_url"]["url"].startswith("data:image/jpeg;base64,")
    body = _body(result)
    offered = body["borrowable_pages"]
    assert offered and all(" = " in line for line in offered), "every offer names the cell it is a key to"
    assert any(line.startswith("amber 4 = amber_wave_quarterly_summary page 4 is ") for line in offered)
    assert f"The picture below shows all {len(offered)} of them" in body["next_step"]
