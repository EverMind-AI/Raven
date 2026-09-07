"""What the outline stage measures about the plan itself.

Both checks here were written against a real twelve-page plan and only after it was
read: three of its pages were section dividers announcing the claim of the page
behind them, and none of its twelve pages planned to show a figure while the
materials held thirteen.

The table-plan half at the end comes from a different source -- the column-squeeze
measurement the built deck is already held to -- and its last case draws the two
plans it disagrees about and measures them, so the derivation is run rather than
argued.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from raven_ppt.contracts.outline import Outline, PagePlan
from raven_ppt.tools.outline import (
    SAME_PROTOTYPE_IS_STRUCTURAL,
    STRUCTURAL_SHARE,
    _composition,
    _invented_layouts,
    _layout_spread,
    _references,
    _structural,
    _thin_pages,
    catalogue_ids,
)
from tests._ppt_engine_fixtures import deck, image, noise_image, noise_png, product_page, template_file  # noqa: F401


class _State:
    def __init__(self, figures: int) -> None:
        self.figures = tuple(type("Figure", (), {"figure_id": f"fig_{index}"})() for index in range(figures))


def _page(number: int, *, says: int = 3, prototype: int | None = None, figures: tuple[str, ...] = ()) -> PagePlan:
    return PagePlan(
        page=number,
        claim=f"claim {number}",
        says=tuple(f"point {index}" for index in range(says)),
        prototype=prototype,
        figures=figures,
    )


def test_the_shares_are_the_measured_ones() -> None:
    assert STRUCTURAL_SHARE == 0.25
    assert SAME_PROTOTYPE_IS_STRUCTURAL == 3


def test_three_dividers_in_twelve_pages_are_reported() -> None:
    """The live case: three pages adapting the same template page, one line each."""
    pages = tuple(
        _page(number, says=1 if number in (3, 5, 8) else 3, prototype=3 if number in (3, 5, 8) else number)
        for number in range(1, 13)
    )
    findings = _structural(Outline(takeaway="t", pages=pages))

    assert [f.kind for f in findings] == ["structural_pages"]
    assert findings[0].detail == {"pages": [3, 5, 8], "of": 12}


def test_one_divider_in_twelve_pages_is_not_worth_saying() -> None:
    pages = tuple(_page(number, says=1 if number == 5 else 3, prototype=number) for number in range(1, 13))

    assert _structural(Outline(takeaway="t", pages=pages)) == []


class _Template:
    def __init__(self, source) -> None:
        self.source = source


class _Bound:
    def __init__(self, source) -> None:
        self.template = _Template(source)
        self.figures = ()


def _house_template(path):
    """A template that says what its own pages are for: cover, agenda, closing."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for heading in ("A deck about something", "Agenda", "A content page", "谢谢观看"):
        page = presentation.slides.add_slide(presentation.slide_layouts[6])
        page.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = heading
    presentation.save(str(path))
    return path


def test_the_cover_and_the_closing_page_have_to_be_the_templates(tmp_path) -> None:
    """The one requirement, not advice: a deck that draws its own cover reads as not
    the user's before a word of it is read."""
    from raven_ppt.tools.outline import _house_pages

    template = _house_template(tmp_path / "house.pptx")
    pages = tuple(_page(number, prototype=None) for number in range(1, 13))
    findings = _house_pages(Outline(takeaway="t", pages=pages), _Bound(template))

    roles = {f.detail["role"] for f in findings}
    assert roles == {"cover", "agenda", "closing"}
    assert all(f.severity.value == "blocking" for f in findings)


def test_a_plan_that_uses_them_passes(tmp_path) -> None:
    from raven_ppt.tools.outline import _house_pages

    template = _house_template(tmp_path / "house.pptx")
    pages = (
        _page(1, prototype=1),
        _page(2, prototype=2),
        *(_page(number, prototype=3) for number in range(3, 12)),
        _page(12, prototype=4),
    )

    assert _house_pages(Outline(takeaway="t", pages=pages), _Bound(template)) == []


def test_a_short_deck_needs_no_index(tmp_path) -> None:
    """Five pages do not need a table of contents."""
    from raven_ppt.tools.outline import _house_pages

    template = _house_template(tmp_path / "house.pptx")
    pages = (_page(1, prototype=1), _page(2), _page(3), _page(4), _page(5, prototype=4))
    findings = _house_pages(Outline(takeaway="t", pages=pages), _Bound(template))

    assert [f.detail["role"] for f in findings] == []


def test_no_template_means_no_requirement() -> None:
    from raven_ppt.tools.outline import _house_pages

    class NoTemplate:
        template = None
        figures = ()

    pages = tuple(_page(number) for number in range(1, 13))
    assert _house_pages(Outline(takeaway="t", pages=pages), NoTemplate()) == []


class _Deck:
    """A project whose materials say this much, without one on disk."""

    def __init__(self, chars: int = 0) -> None:
        self.chars = chars
        self.ingest_dir = Path("/nowhere")


def test_the_says_field_names_no_character_count() -> None:
    """A count here is a different page in each language.

    Measured in one body box at one size, Chinese fills it at 400 characters and
    English at 1168. A budget stated to the planner therefore asks one language for
    a full page and the other for a third of one -- so the field says what a page
    has to carry, and whether it fits is settled by measuring the render.
    """
    from pathlib import Path

    from raven_ppt.tools.outline import PptOutlineTool

    said = PptOutlineTool(workspace=Path("/tmp")).parameters["properties"]["pages"]["items"]["properties"]["says"]

    assert not re.search(r"\d{3,} character", said["description"])
    assert "the claim is carried" in said["description"]


class _Bare:
    """A project state with no template bound, so no page counts as structural."""

    template = None


def _planned(number: int, says: tuple[str, ...], **kw) -> PagePlan:
    return PagePlan(page=number, claim=f"claim {number}", says=says, **kw)


def test_a_page_whose_whole_plan_is_one_line_is_reported() -> None:
    outline = Outline(takeaway="t", pages=(_planned(1, ("Only this.",)),))

    found = _thin_pages(outline, _Bare())

    assert [finding.kind for finding in found] == ["thin_page"]
    assert found[0].page == 1
    assert found[0].detail == {"says": 1}


def test_a_page_planning_nothing_at_all_is_reported() -> None:
    assert [f.kind for f in _thin_pages(Outline(takeaway="t", pages=(_planned(1, ()),)), _Bare())] == ["thin_page"]


def test_two_points_are_a_plan_however_short_they_are() -> None:
    """No character floor, deliberately.

    It used to take five points or 140 characters, measured off one deck. A count
    is a different page in each language -- in one body box at one size Chinese
    fills it at 400 characters and English at 1168 -- so the floor asked one
    language for a page and the other for a third of one. How full a page comes out
    is settled by measuring the render, not guessed from the plan.
    """
    outline = Outline(takeaway="t", pages=(_planned(1, ("Up.", "Down.")),))

    assert _thin_pages(outline, _Bare()) == []


def test_a_page_carrying_a_figure_plans_its_copy_in_the_figure() -> None:
    outline = Outline(takeaway="t", pages=(_planned(1, ("One line.",), figures=("fig-1",)),))

    assert _thin_pages(outline, _Bare()) == []


def test_a_page_with_an_errand_is_going_to_get_something_to_show() -> None:
    outline = Outline(takeaway="t", pages=(_planned(1, ("One line.",), needs="draw the six-stage path"),))

    assert _thin_pages(outline, _Bare()) == []


# The table plan, in three parts: the shape its rows come in, how the author reads it
# back, and whether a page has the room for the grid it describes.


_BENCHMARKS = {
    "columns": ["Benchmark", "Ours", "Baseline A", "Baseline B", "Baseline C", "Delta"],
    "rows": [
        ["DAVIS J&F", "82.4", "79.1", "78.6", "77.2", "+3.3"],
        ["YouTube-VOS", "85.1", "82.0", "81.4", "80.9", "+3.1"],
        ["MOSE", "71.9", "68.2", "67.5", "66.4", "+3.7"],
    ],
    "reading": "one model against three task-specific ones",
}
# Twelve columns of phrases rather than of figures, which is the distinction the
# deleted column cap could not draw: this needs 20.1in of text and the page carries
# 19.0in of it with every cell wrapped onto a second line.
_TWELVE_PHRASES = {
    "columns": [
        "Deployment model",
        "Annual licence",
        "p99 latency",
        "Languages",
        "On-premise",
        "Certification",
        "Migration effort",
        "Lock-in risk",
        "Support tier",
        "SLA",
        "Data residency",
        "Audit logging",
    ],
    "rows": [
        [
            "Self-hosted cluster",
            "$248,000",
            "38 milliseconds",
            "17 supported",
            "Available",
            "SOC 2 Type II",
            "Six engineer weeks",
            "Low overall",
            "Platinum 24/7",
            "99.95 percent",
            "EU and US regions",
            "Full retention",
        ]
    ],
}


def test_the_page_structure_a_plan_declares_survives_the_round_trip(tmp_path) -> None:
    """`section` was written on every call and read back as "" on every one for a
    release, so a field added to the plan without a line in `load_outline` reaches
    nothing. This is that line, asserted."""
    from raven_ppt.contracts import load_outline, outline_path, write_outline
    from raven_ppt.contracts.outline import Outline, PagePlan
    from raven_ppt.contracts.project import Project

    path = outline_path(Project(workspace=tmp_path, slug="deck"))
    plan = PagePlan(
        page=1,
        claim="c",
        layout="P14",
        layers=("M4", "M11"),
        anti_pattern="a lane that would read the same with the chart removed",
    )
    write_outline(Outline(takeaway="t", pages=(plan,)), path)

    outline = load_outline(path)

    assert outline is not None
    assert outline.pages[0].layout == "P14"
    assert outline.pages[0].layers == ("M4", "M11")
    assert outline.pages[0].anti_pattern == plan.anti_pattern

    # And the shape this field used to have, which is on disk in every outline written
    # before the split: the ids come apart into the two fields rather than the structure
    # column reading `P14 + M4 + M11` for one page and `P14` for the next, which is not a
    # spread of structures however it is counted.
    write_outline(Outline(takeaway="t", pages=(PagePlan(page=1, claim="c", layout="P14 + M4 + M11"),)), path)
    legacy = load_outline(path)

    assert legacy is not None
    assert (legacy.pages[0].layout, legacy.pages[0].layers) == ("P14", ("M4", "M11"))


# --- the sources' own citations -------------------------------------------
#
# The live failure: a markdown competitive analysis citing 56 URLs, ingest extracting
# nothing visual out of the markdown itself, and an author that answered the sweep
# errand in its second turn -- "no figures to fetch, it's tables, findings and
# recommendations" -- and never opened one of them. Twenty pages, no image on any.

_CITING = """# Competitive analysis

Mem0 documents its pipeline at https://docs.mem0.ai/core-concepts/how-it-works.
Zep publishes its benchmarks in [one post](https://blog.getzep.com/state-of-the-art).
"""

_MEM0 = "https://docs.mem0.ai/core-concepts/how-it-works"
_ZEP = "https://blog.getzep.com/state-of-the-art"


def _cited_deck(tmp_path: Path, *, materials: str = _CITING, figures: int = 0):
    """A deck with a brief recorded and the given text as its ingested materials."""
    import json

    from raven_ppt.contracts import DeckBrief, PageBudget, Project, brief_path, write_brief
    from raven_ppt.services.ingest import CATALOGUE_FILE, MATERIALS_FILE

    deck = Project(workspace=tmp_path, slug="deck")
    write_brief(DeckBrief(language="English", audience="a review", pages=PageBudget(1, 4)), brief_path(deck))
    deck.ingest_dir.mkdir(parents=True, exist_ok=True)
    (deck.ingest_dir / MATERIALS_FILE).write_text(materials, encoding="utf-8")
    if figures:
        assets = {f"fig_{index}": {"kind": "figure", "caption": "c"} for index in range(1, figures + 1)}
        (deck.ingest_dir / CATALOGUE_FILE).write_text(json.dumps({"assets": assets}), encoding="utf-8")
    return deck


async def _outlined(deck, **extra):
    import json

    from raven_ppt.tools.outline import PptOutlineTool

    return json.loads(
        await PptOutlineTool(deck.workspace).execute(
            project=deck.slug,
            takeaway="one memory layer beats four bolted together",
            pages=[
                {
                    "page": 1,
                    "claim": "One memory layer beats four bolted together",
                    "carries": "prose",
                    "says": ["Compare: the one against the four.", "Explain: why it holds."],
                }
            ],
            **extra,
        )
    )


def _refusal(body):
    return [entry for entry in body["measured"]["for_you"] if entry["kind"] == "unswept_citations"]


@pytest.fixture(autouse=True)
def _offline(monkeypatch):
    """Nothing in this file opens a socket.

    The sweep check opens the pages an entry calls empty, and the URLs below are
    real ones -- so without this the suite would reach docs.mem0.ai and decide
    whether a test passes by what is served there today. Every test gets a page
    holding nothing, which is the answer that leaves the older cases meaning what
    they meant; the ones about the check hand over a page holding something.
    """
    from raven_ppt.services import citations

    monkeypatch.setattr(citations, "_page_pictures", lambda url: ())


async def test_an_outline_is_refused_while_the_sources_own_citations_are_unopened(tmp_path) -> None:
    from raven_ppt.contracts import load_outline, outline_path

    deck = _cited_deck(tmp_path)

    body = await _outlined(deck)

    assert body["ok"] is False
    assert body["recorded"] is False
    assert load_outline(outline_path(deck)) is None
    refused = _refusal(body)
    assert len(refused) == 1
    assert refused[0]["severity"] == "blocking"
    assert refused[0]["detail"] == {"cited": 2, "unswept": 2, "outstanding": [_MEM0, _ZEP]}
    # A refusal has to name what to call, not only what is wrong.
    assert "web_fetch" in refused[0]["problem"] and "ppt_fetch" in refused[0]["problem"]
    assert "ppt_fetch" in refused[0]["problem"]
    assert "`swept`" in refused[0]["problem"]
    assert _MEM0 in refused[0]["problem"]


async def test_material_that_cites_nothing_is_not_held(tmp_path) -> None:
    """The case that must not be caught: a source that genuinely is only text."""
    deck = _cited_deck(tmp_path, materials="# Notes\n\nThe model reaches 44 FPS on one A100.\n")

    body = await _outlined(deck)

    assert body["ok"] is True
    assert body["recorded"] is True


async def test_a_deck_that_already_has_figures_is_not_held(tmp_path) -> None:
    """Whether these particular figures carry these pages is the author's judgement."""
    deck = _cited_deck(tmp_path, figures=1)

    body = await _outlined(deck)

    assert body["recorded"] is True


async def test_naming_what_came_back_from_each_citation_clears_it(tmp_path) -> None:
    deck = _cited_deck(tmp_path)

    body = await _outlined(
        deck,
        swept=[
            {"url": _MEM0, "found": "prose and code samples, no figure"},
            {"url": _ZEP, "found": "404"},
        ],
    )

    assert body["ok"] is True
    assert body["recorded"] is True


async def test_a_partial_sweep_still_names_the_rest(tmp_path) -> None:
    deck = _cited_deck(tmp_path)

    body = await _outlined(deck, swept=[{"url": _MEM0, "found": "prose only"}])

    assert _refusal(body)[0]["detail"]["outstanding"] == [_ZEP]


async def test_the_sweep_is_kept_with_the_deck_and_not_restated(tmp_path) -> None:
    """`ppt_outline` runs again on every replan, and a record restated on every call is
    a record forgotten on one of them -- which would refuse a plan with nothing wrong."""
    deck = _cited_deck(tmp_path)
    await _outlined(deck, swept=[{"url": _MEM0, "found": "prose only"}])

    body = await _outlined(deck, swept=[{"url": _ZEP, "found": "404"}])

    assert body["recorded"] is True
    assert (await _outlined(deck))["recorded"] is True


async def test_a_swept_entry_that_says_nothing_came_back_is_refused(tmp_path) -> None:
    """What keeps the way out from being a checkbox: a checkbox is cheaper to tick
    than the fetch it stands for, so an entry has to state what the page held."""
    deck = _cited_deck(tmp_path)

    body = await _outlined(deck, swept=[{"url": _MEM0}, {"url": _ZEP}])

    assert body["ok"] is False
    assert "does not say what came back" in body["error"]
    assert (await _outlined(deck))["recorded"] is False


async def test_a_sentence_instead_of_a_url_is_refused(tmp_path) -> None:
    """The sentence the live run actually offered, in the place that would take it."""
    deck = _cited_deck(tmp_path)

    body = await _outlined(deck, swept=[{"url": "no figures to fetch", "found": "it is a text analysis"}])

    assert body["ok"] is False
    assert "not an http(s) URL" in body["error"]


async def test_a_url_the_materials_do_not_cite_clears_nothing(tmp_path) -> None:
    deck = _cited_deck(tmp_path)

    body = await _outlined(deck, swept=[{"url": "https://example.com/somewhere-else", "found": "nothing"}])

    assert _refusal(body)[0]["detail"]["outstanding"] == [_MEM0, _ZEP]


async def test_a_citation_fetched_into_the_deck_needs_no_declaration(tmp_path) -> None:
    """Evidence rather than assertion: whoever fetched it recorded the URL."""
    from raven_ppt.services.ingest import sources

    deck = _cited_deck(tmp_path)
    sources.write_source(deck, "how-it-works.md", "the page", origin=f"fetch:{_MEM0}")

    body = await _outlined(deck, swept=[{"url": _ZEP, "found": "404"}])

    assert body["recorded"] is True


def test_a_namespace_declaration_is_not_a_citation() -> None:
    """A .docx or .pptx read as XML carries a dozen of these and none is a page."""
    from raven_ppt.services import citations

    declared = (
        "http://schemas.openxmlformats.org/drawingml/2006/main\n"
        "http://www.w3.org/2000/svg\n"
        "http://localhost:8000/report\n"
    )

    assert citations.of_text(declared) == ()


def test_two_spellings_of_one_page_are_one_citation() -> None:
    """Forgiving in one direction: over-matching costs a URL swept once instead of
    twice, and under-matching costs a refusal the author cannot clear by doing the
    work."""
    from raven_ppt.services import citations

    same = "http://www.getzep.com/post/ and https://getzep.com/post#results\n"

    assert citations.of_text(same) == ("http://www.getzep.com/post/",)


def test_a_markdown_link_does_not_take_the_bracket_with_it() -> None:
    from raven_ppt.services import citations

    text = "see [it](https://docs.mem0.ai/core-concepts/how-it-works), and https://en.wikipedia.org/wiki/Foo_(bar)."

    assert citations.of_text(text) == (_MEM0, "https://en.wikipedia.org/wiki/Foo_(bar)")


# --- and whether the sweep record was a look ------------------------------
#
# The second live failure, one gate later: thirteen URLs declared swept, thirteen
# plausible sentences about what each held, and zero `fetch:` lines in that deck's
# source manifest. `https://docs.mem0.ai/core-concepts/how-it-works` was recorded as
# "Core concepts guide, text and code snippets only"; one request to it returns an
# architecture diagram under its author's own caption.

_MEM0_PAGE = """
<html><head><meta property="og:image" content="https://cdn.example/social-card.png"></head>
<body>
<nav><img src="/nav/menu-open.png" alt="menu"></nav>
<header><img src="https://cdn.example/brand/wordmark.png" alt="mem0"></header>
<main>
  <p>Text and code snippets.</p>
  <figure>
    <img src="https://mintcdn.com/mem0/images/memory-extraction.png" alt="extraction pipeline">
    <figcaption>Memory extraction: Mem0 turns messages into stored facts.</figcaption>
  </figure>
  <img src="/img/spacer.gif" width="1" height="1">
  <img src="/assets/inline-tick.png" width="16" height="16" alt="ok">
  <img src="/images/architecture-overview.png" alt="the architecture">
</main>
<footer><img src="/f/logo-small.png"></footer>
</body></html>
"""

_DIAGRAM = "https://mintcdn.com/mem0/images/memory-extraction.png"


def _serving(**pages):
    """A stand-in for the look, answering from markup rather than from the web."""
    from raven_ppt.services import citations

    def look(url):
        return citations._pictures(pages.get(citations.key(url), ""), url)

    return look


async def test_a_page_that_holds_a_picture_refuses_the_entry_calling_it_empty(tmp_path, monkeypatch) -> None:
    """The sentence the live run actually wrote, against the page it wrote it about."""
    from raven_ppt.services import citations

    deck = _cited_deck(tmp_path)
    monkeypatch.setattr(citations, "_page_pictures", _serving(**{citations.key(_MEM0): _MEM0_PAGE}))

    body = await _outlined(
        deck,
        swept=[
            {"url": _MEM0, "found": "Core concepts guide, text and code snippets only"},
            {"url": _ZEP, "found": "404"},
        ],
    )

    assert body["ok"] is False
    # A refusal is only useful if it hands over what the author said it could not
    # find: which URL, what it claimed, what is actually there, and what to call.
    assert _MEM0 in body["error"]
    assert "Core concepts guide, text and code snippets only" in body["error"]
    assert _DIAGRAM in body["error"]
    assert "Memory extraction: Mem0 turns messages into stored facts." in body["error"]
    assert "ppt_fetch" in body["error"]


async def test_a_refused_sweep_records_nothing_and_the_next_call_still_asks(tmp_path, monkeypatch) -> None:
    from raven_ppt.services import citations

    deck = _cited_deck(tmp_path)
    monkeypatch.setattr(citations, "_page_pictures", _serving(**{citations.key(_MEM0): _MEM0_PAGE}))

    await _outlined(deck, swept=[{"url": _MEM0, "found": "text only, no figures"}])

    assert citations.load_sweep(deck) == ()
    assert (await _outlined(deck))["recorded"] is False


async def test_a_page_that_really_holds_nothing_still_clears_its_url(tmp_path, monkeypatch) -> None:
    """The case that must not be caught: the author looked, and there was nothing."""
    from raven_ppt.services import citations

    deck = _cited_deck(tmp_path)
    bare = "<html><body><main><p>Just prose.</p></main></body></html>"
    monkeypatch.setattr(citations, "_page_pictures", _serving(**{citations.key(_MEM0): bare}))

    body = await _outlined(
        deck,
        swept=[{"url": _MEM0, "found": "prose only"}, {"url": _ZEP, "found": "404"}],
    )

    assert body["recorded"] is True


async def test_a_page_that_cannot_be_reached_lets_the_entry_stand(tmp_path, monkeypatch) -> None:
    """The one this must never get wrong. Timeouts, 403s and dead hosts are the
    ordinary weather of the open web, and a gate that read one as a lie would let a
    bad minute of network stop a deck with nothing wrong with it."""
    from raven_ppt.services import citations

    deck = _cited_deck(tmp_path)

    def unreachable(url):
        raise TimeoutError(url)

    monkeypatch.setattr(citations, "_page_pictures", unreachable)

    body = await _outlined(
        deck,
        swept=[{"url": _MEM0, "found": "text only"}, {"url": _ZEP, "found": "404"}],
    )

    assert body["recorded"] is True


def test_only_the_urls_claimed_empty_are_opened() -> None:
    """A page the author granted has pictures is its judgement to make, and a page
    it fetched is already in the manifest -- neither is worth a request."""
    from raven_ppt.services import citations

    opened = []

    def look(url):
        opened.append(url)
        return ()

    citations.read(
        [
            {"url": _MEM0, "found": "three screenshots, all UI chrome"},
            {"url": _ZEP, "found": "text and code snippets only"},
        ],
        look=look,
    )

    assert opened == [_ZEP]


def test_a_sweep_longer_than_the_budget_costs_the_budget() -> None:
    """Sampled rather than truncated: the first N is a rule an author can pad past
    by listing the page it did not open ninth."""
    from raven_ppt.services import citations

    opened = []

    def look(url):
        opened.append(url)
        return ()

    entries = [{"url": f"https://example.com/page-{index}", "found": "no figures"} for index in range(20)]
    citations.read(entries, look=look)

    assert len(opened) == citations.MAX_CHECKED
    assert len(set(opened)) == citations.MAX_CHECKED


def test_a_look_counts_a_captioned_figure_and_not_the_page_furniture() -> None:
    """Stricter than the image extractor on the three counts that only drop things:
    no `og:image`, no tiny declared side, and no exemption for a `<figure>`."""
    from raven_ppt.services import citations

    pictures = citations._pictures(_MEM0_PAGE, _MEM0)

    assert [picture.url for picture in pictures] == [
        _DIAGRAM,
        "https://docs.mem0.ai/images/architecture-overview.png",
    ]
    assert pictures[0].where == "figure with caption"
    assert pictures[0].caption == "Memory extraction: Mem0 turns messages into stored facts."


def test_a_page_of_nothing_but_furniture_holds_no_picture() -> None:
    from raven_ppt.services import citations

    chrome = """
    <html><body>
    <nav><img src="/nav/burger.png" alt="menu"></nav>
    <main><p>Just prose.</p>
      <img src="/i/favicon-32.png">
      <img src="/i/tracking-beacon.gif" width="1" height="1">
      <img src="/i/logo.svg" alt="brand">
    </main></body></html>
    """

    assert citations._pictures(chrome, "https://example.com/x") == ()


def test_an_entry_that_grants_the_page_had_pictures_is_not_a_claim_of_nothing() -> None:
    """The words survive inside a denial -- "text only, no figures" names one and
    holds none -- so a denial anywhere puts the entry back in the queue, and an
    unreadable sentence is checked rather than skipped."""
    from raven_ppt.services import citations

    granted = ("three screenshots, all UI chrome", "two diagrams, both marketing", "a bar chart")
    denied = ("text only, no figures", "Core concepts guide, text and code snippets only", "404", "prose")

    assert [citations.claims_nothing(found) for found in granted] == [False, False, False]
    assert [citations.claims_nothing(found) for found in denied] == [True, True, True, True]


# ---------------------------------------------------------------------------
# The layout ids a page declares


def _composed(number: int, layout: str) -> PagePlan:
    return PagePlan(page=number, claim=f"claim {number}", layout=layout)


def test_the_catalogue_ids_are_read_out_of_the_catalogue() -> None:
    """Taken from the file, not written here, and counted without naming a count.

    The file is what the author reads, so it is what an id has to be in; two lists of
    the same ids drift. An earlier version of this test asserted twenty-two and eleven,
    and went red the day the catalogue grew -- which is the same staleness it exists to
    prevent, so what it asserts now is the shape: each family numbered from 1 with no
    gaps, because a gap is an id a page can name and the file cannot answer.
    """
    ids = catalogue_ids()
    structures = sorted(int(one[1:]) for one in ids if one.startswith("P"))
    modifiers = sorted(int(one[1:]) for one in ids if one.startswith("M"))

    assert structures and modifiers
    assert structures == list(range(1, len(structures) + 1))
    assert modifiers == list(range(1, len(modifiers) + 1))


def test_a_layout_id_the_catalogue_does_not_carry_is_refused() -> None:
    """The ids one live run wrote were zero-padded, and the catalogue has no P01.

    Its transcript names layouts.md once, in the output of a directory listing: it saw
    the filename, never opened the file, and invented eleven ids that read down the
    column like a deck with a range of structures in it. Nine of its pages came out as
    the same tinted panel with copy in it.
    """
    outline = Outline(
        takeaway="one claim", pages=(_composed(1, "P01"), _composed(2, "P04 + M2"), _composed(3, "P7 + M4"))
    )

    findings = _invented_layouts(outline)

    assert [finding.kind for finding in findings] == ["invented_layout"]
    assert findings[0].severity.value == "blocking"
    assert findings[0].detail["invented"] == {1: ["P01"], 2: ["P04"]}
    # The range comes off the catalogue, so the message cannot promise a bound the
    # file has stopped having.
    structures = findings[0].detail["structures"]
    assert f"{structures[0]} to {structures[-1]}" in findings[0].message


def test_the_field_offers_the_catalogue_and_nothing_else() -> None:
    """A closed list rather than free text, and the list is the catalogue's own.

    Free text is what the field had when a live run filled it with `P01`, `P04`, `P07`.
    The refusal caught it after the fact; an enum means the ids are the only thing that
    can be written -- and built off `catalogue_ids` so the list a caller picks from and
    the list it is checked against cannot come apart when layouts.md gains a structure.
    """
    from raven_ppt.tools.outline import _layer_enum, _structure_enum, catalogue_ids

    known = catalogue_ids()
    structures = {one for one in known if one.startswith("P")}
    layers = {one for one in known if one.startswith("M")}

    assert set(_structure_enum()) == structures | {""}, "the structure list is not the catalogue's"
    assert set(_layer_enum()) == layers, "the layer list is not the catalogue's"
    assert "P01" not in _structure_enum() and "P99" not in _structure_enum()
    # Part 1 folded forty-one ids into eleven skeletons and every one of them is still
    # writable: the convergence is what you read, not what you may say.
    assert len(structures) == 41, f"{len(structures)} structures, and the passages are for 41"


def test_the_guard_refuses_only_what_claims_to_be_an_id() -> None:
    """Both directions, because "the guard is too tight" is filed as often as too loose.

    The first version read every `letters+digits` token as an id claim, so `P11 grid
    2x2` was refused over its `2x2` and `p14` over its case -- each of them naming a
    structure the catalogue carries.
    """
    accepted = ("P14 + M4 + M11", "p14", "P11 grid 2x2", "16x9 free", "prototype 9", "")
    refused = ("P01", "P99", "M0")

    for layout in accepted:
        outline = Outline(takeaway="one claim", pages=(_composed(1, layout),))
        assert _invented_layouts(outline) == [], layout
    for layout in refused:
        outline = Outline(takeaway="one claim", pages=(_composed(1, layout),))
        assert [f.kind for f in _invented_layouts(outline)] == ["invented_layout"], layout


def test_ids_that_are_in_the_catalogue_and_an_empty_field_both_pass() -> None:
    """Leaving it out says "this page is a clone, or I have not decided", which is fine.

    Only a filled-in id that is not in the file says something false.
    """
    outline = Outline(
        takeaway="one claim", pages=(_composed(1, "P14 + M4 + M11"), _composed(2, ""), _composed(3, "P22"))
    )

    assert _invented_layouts(outline) == []


def test_the_composed_pages_are_told_to_open_the_catalogue() -> None:
    """Named on its own rather than inside the list of five, and for the pages that need it.

    The five were per-page -- a page carrying no chart needs no charts.md -- and this
    one is not: every page not cloned from a template example is composed out of it.
    Inside that list it was opened by one model in three.
    """
    outline = Outline(takeaway="one claim", pages=(_page(1, prototype=4), _composed(2, "P7"), _composed(3, "P14")))

    said = _references(outline)

    assert "layouts.md" in said
    assert "page(s) 2, 3" in said
    assert "1" not in said.split("page(s) 2, 3")[1][:20]


# --- the layout column, read down the deck ---------------------------------
#
# The gap these close: `page.layout` was read in exactly one place, and that place
# (`_invented_layouts`) only asks whether the ids exist. A plan writing `P3` on twelve
# pages cleared every plan-stage gate, and the uniformity was discovered off the render
# a whole build later, where `layout_variety` reads it.


def _read_spread(*layouts: str, roles: dict[int, str] | None = None, prototypes: dict[int, int] | None = None):
    """`_layout_spread` over one page per layout, carrying the roles or prototypes given."""
    named = prototypes or {}
    outline = Outline(
        takeaway="one claim",
        pages=tuple(
            PagePlan(page=number, claim=f"claim {number}", layout=layout, prototype=named.get(number))
            for number, layout in enumerate(layouts, start=1)
        ),
    )
    return _layout_spread(outline, roles or {})


def test_a_plan_that_composes_nine_of_twelve_pages_the_same_way_is_reported() -> None:
    """The whole point of the column, and nothing read it: `_invented_layouts` looks at
    the same field and only asks whether the ids exist, so this plan was clean."""
    findings = _read_spread(*["P3"] * 9, "P7 + M2", "P14", "P22")

    assert [f.kind for f in findings] == ["layout_spread"]
    assert findings[0].severity.value == "warning"
    assert findings[0].page is None
    assert findings[0].detail["reading"] == "concentrated"
    assert findings[0].detail["repeated"] == [1, 2, 3, 4, 5, 6, 7, 8, 9]
    assert findings[0].detail["composition"] == "P3"
    assert findings[0].detail["declared"] == 12
    assert findings[0].detail["distinct"] == 4
    assert "pages 1, 2, 3, 4, 5, 6, 7, 8, 9" in findings[0].message
    assert "9 of the 12 page(s) that declare one" in findings[0].message
    assert "4 distinct compositions" in findings[0].message
    # Said as a reading, because refusing here would refuse a design judgement the
    # measurement is not entitled to make -- the reason the built-deck row gives.
    assert "not a verdict" in findings[0].message


def test_a_plan_with_a_spread_of_compositions_is_not() -> None:
    assert _read_spread("P1", "P3 + M2", "P7", "P14 + M4", "P22 + M11", "P9", "P17 + M1", "P5") == []


def test_half_a_deck_on_one_composition_is_under_the_share() -> None:
    """The share is there because the count alone misreads: five pages of one shape and
    five of five others is a deck with a spread in it."""
    assert _read_spread(*["P3"] * 5, "P1", "P7", "P14", "P22", "P9") == []


def test_the_page_floor_is_the_one_the_built_deck_is_read_by() -> None:
    """Below the floor a share means nothing, and both ends of the build have to agree
    about where the floor is -- two numbers with one meaning between them is how they
    come to disagree about the same deck."""
    from raven_ppt.services.measure.variety import MIN_PAGES

    assert _read_spread(*["P3"] * (MIN_PAGES - 1)) == []
    assert [f.kind for f in _read_spread(*["P3"] * MIN_PAGES)] == ["layout_spread"]
    assert _read_spread(*[""] * (MIN_PAGES - 1)) == []


def test_the_templates_own_furniture_does_not_make_a_deck_uniform() -> None:
    """A cover, an index, a divider and a closing page are meant to be alike, so
    counting them reports a correct deck for the pages it was told to clone.

    Here the four of them are what carries `P1` over the share: without the skip the
    plan is reported for eleven pages of which seven declare `P1`, and its seven
    content pages have a spread.
    """
    layouts = ("P1", "P1", "P1", "P3", "P7 + M2", "P14", "P1", "P22", "P1", "P1", "P1")
    roles = {1: "cover", 2: "agenda", 10: "section", 11: "closing"}

    assert _read_spread(*layouts, roles=roles) == []
    # The same eleven pages with nothing said about their roles do report, so it is the
    # skip doing the work here rather than the shape of the deck.
    assert [f.kind for f in _read_spread(*layouts)] == ["layout_spread"]


def test_a_page_cloned_from_a_template_example_is_not_read_here() -> None:
    """Its structure is that example's, and the field asks it to leave `layout` out --
    so reading its silence as an undecided composition reports every template deck."""
    assert _read_spread(*[""] * 8, prototypes={number: 3 for number in range(1, 9)}) == []
    assert [f.kind for f in _read_spread(*[""] * 8)] == ["layout_spread"]


def test_a_column_nobody_filled_in_is_said_as_that_and_not_as_one_shape() -> None:
    """An empty `layout` means "template clone, or not decided", so five of them are not
    five pages of one composition. The two ask for different edits, and a message that
    ran them together would ask for the wrong one."""
    findings = _read_spread("", "", "", "", "", "P3", "P7 + M2", "P14")

    assert [f.kind for f in findings] == ["layout_spread"]
    assert findings[0].severity.value == "warning"
    assert findings[0].detail == {"reading": "undeclared", "pages": [1, 2, 3, 4, 5], "of": 8}
    assert "name no layout id" in findings[0].message
    assert "the same composition" not in findings[0].message
    assert "not a verdict" in findings[0].message


def test_the_pages_that_declare_nothing_are_left_out_of_the_count() -> None:
    """The denominator is the pages that declared, not the deck -- four silent pages
    are not four pages of one shape, and they cannot swell one either."""
    findings = _read_spread("", "", "", "", "P3", "P3", "P3", "P3", "P3", "P7", "P14", "P22")

    assert [f.kind for f in findings] == ["layout_spread"]
    assert findings[0].detail["reading"] == "concentrated"
    assert findings[0].detail["repeated"] == [5, 6, 7, 8, 9]
    assert findings[0].detail["declared"] == 8
    assert findings[0].detail["composition"] == "P3"


def test_reordered_and_recased_modifiers_are_one_composition() -> None:
    """The modifiers stack, so their order in the field is not a fact about the page,
    and `p14` is `P14` -- the lesson `_invented_layouts` already learned when it refused
    `P11 grid 2x2` over its `x2`."""
    assert _composition("P14 + M11 + M4") == _composition("P14 + M4 + M11") == "P14 + M4 + M11"
    assert _composition("p14 + m4 + m11") == "P14 + M4 + M11"
    assert _composition("P11 grid 2x2") == "P11"
    assert _composition("two columns, laid out by eye") == ""

    findings = _read_spread(
        "P14 + M4 + M11", "P14 + M11 + M4", "p14+m11+m4", "M4 + M11 + P14", "P14 + M4 + M11", "P14+M11+M4"
    )

    assert [f.kind for f in findings] == ["layout_spread"]
    assert findings[0].detail["composition"] == "P14 + M4 + M11"
    assert findings[0].detail["distinct"] == 1
    assert "1 distinct composition in it" in findings[0].message


async def test_a_uniform_layout_column_comes_back_from_the_tool(tmp_path) -> None:
    """A warning that rides along with the recorded outline: which way to answer it is
    the author's, and the plan is still what the next stage is written against."""
    import json

    from raven_ppt.contracts import DeckBrief, PageBudget, Project, brief_path, write_brief
    from raven_ppt.tools.outline import PptOutlineTool

    structures = sorted((one for one in catalogue_ids() if one.startswith("P")), key=lambda one: int(one[1:]))
    if not structures:
        pytest.skip("the layout catalogue is not in this checkout")
    deck = Project(workspace=tmp_path, slug="deck")
    write_brief(DeckBrief(language="English", audience="a review", pages=PageBudget(10, 14)), brief_path(deck))

    body = json.loads(
        await PptOutlineTool(tmp_path).execute(
            project="deck",
            takeaway="one model matches four",
            pages=[
                {
                    "page": number,
                    "claim": f"Page {number} says something of its own",
                    "says": [f"Compare: what page {number} holds.", "Explain: why it holds."],
                    "layout": structures[0],
                }
                for number in range(1, 13)
            ],
        )
    )

    assert body["ok"] is True
    assert body["recorded"] is True
    reported = [entry for entry in body["measured"]["for_you"] if entry["kind"] == "layout_spread"]
    assert len(reported) == 1
    assert reported[0]["severity"] == "warning"
    assert reported[0]["detail"]["composition"] == structures[0]


def _borrowing_plan(**fields) -> Outline:
    return Outline(takeaway="t", pages=(PagePlan(page=1, claim="c", says=("a", "b", "c"), **fields),))


def test_a_borrowed_page_from_a_content_page_of_another_bundled_template_passes() -> None:
    from raven_ppt.services.template.defaults import bundled_path
    from raven_ppt.tools.outline import _borrowed_pages

    bound = bundled_path("amber_wave_quarterly_summary")
    lender = bundled_path("gold_panel_year_end_summary")
    if bound is None or lender is None:
        pytest.skip("the bundled templates this reads are not on this checkout")

    plan = _borrowing_plan(prototype=13, borrowed="gold_panel_year_end_summary")

    assert _borrowed_pages(plan, _Bound(bound)) == []


def test_a_borrowed_name_that_ships_no_template_is_refused_with_the_ones_that_do() -> None:
    from raven_ppt.services.template.defaults import bundled_path
    from raven_ppt.tools.outline import _borrowed_pages

    bound = bundled_path("amber_wave_quarterly_summary")
    if bound is None:
        pytest.skip("the bundled templates this reads are not on this checkout")

    found = _borrowed_pages(_borrowing_plan(prototype=3, borrowed="gold"), _Bound(bound))

    assert [(f.kind, f.severity.value, f.page) for f in found] == [("borrowed", "blocking", 1)]
    assert "no bundled template is called 'gold'" in found[0].message
    assert "gold_panel_year_end_summary" in found[0].message


def test_borrowing_a_cover_or_the_bound_template_itself_is_refused() -> None:
    from raven_ppt.services.template.defaults import bundled_path
    from raven_ppt.tools.outline import _borrowed_pages

    bound = bundled_path("amber_wave_quarterly_summary")
    if bound is None or bundled_path("gold_panel_year_end_summary") is None:
        pytest.skip("the bundled templates this reads are not on this checkout")

    cover = _borrowed_pages(_borrowing_plan(prototype=1, borrowed="gold_panel_year_end_summary"), _Bound(bound))
    assert len(cover) == 1 and "cover" in cover[0].message and "13" in cover[0].message

    own = _borrowed_pages(_borrowing_plan(prototype=6, borrowed="amber_wave_quarterly_summary"), _Bound(bound))
    assert len(own) == 1 and "built in" in own[0].message

    unnumbered = _borrowed_pages(_borrowing_plan(prototype=None, borrowed="gold_panel_year_end_summary"), _Bound(bound))
    assert len(unnumbered) == 1 and "no `prototype`" in unnumbered[0].message


def test_a_borrowed_page_outside_the_curated_reference_list_is_refused() -> None:
    """`reference_pages` is not "any content page": it holds the pages measured to carry
    into another template's theme and master. A content page outside it may look fine in
    its own deck and arrive with source-specific styling, which is what the list was cut
    to exclude, so a plan naming one is refused at the plan."""
    from raven_ppt.services.template.defaults import REFERENCE_PAGES, bundled_path
    from raven_ppt.tools.outline import _borrowed_pages

    bound = bundled_path("amber_wave_quarterly_summary")
    lender = bundled_path("gold_panel_year_end_summary")
    if bound is None or lender is None:
        pytest.skip("the bundled templates this reads are not on this checkout")

    curated = REFERENCE_PAGES["gold_panel_year_end_summary"]
    assert 10 not in curated, "this case needs a content page the curated list leaves out"

    found = _borrowed_pages(_borrowing_plan(prototype=10, borrowed="gold_panel_year_end_summary"), _Bound(bound))

    assert [(f.kind, f.severity.value, f.page) for f in found] == [("borrowed", "blocking", 1)]
    assert "verified to carry" in found[0].message
    assert str(curated[0]) in found[0].message, "the reply names the pages that are offered"


def _content_template(path):
    """A template with a real content example: cover, agenda, one page of four points, closing."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for heading in ("A deck about something", "Agenda", "Four findings", "谢谢观看"):
        page = presentation.slides.add_slide(presentation.slide_layouts[6])
        page.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = heading
        if heading == "Four findings":
            for index in range(4):
                box = page.shapes.add_textbox(Inches(1 + 3 * index), Inches(3), Inches(2.6), Inches(2))
                box.text_frame.text = (
                    f"Finding {index + 1}: a sentence long enough to be copy rather than a label on it."
                )
    presentation.save(str(path))
    return path


def test_a_plan_composing_most_of_its_content_pages_inside_a_template_is_refused(tmp_path) -> None:
    """Two decks in one template family: 4 of 17 content pages composed read as the
    template's own, 17 of 17 read as nothing but self-drawn layouts, and the ask that
    listed the pages had been read twice. The share is a refusal, with the examples."""
    from raven_ppt.tools.outline import COMPOSED_SHARE, _composed_pages

    template = _content_template(tmp_path / "house.pptx")
    pages = (_page(1, prototype=1), _page(2), _page(3), _page(4), _page(5), _page(6, prototype=4))

    findings = _composed_pages(Outline(takeaway="t", pages=pages), _Bound(template))

    assert COMPOSED_SHARE == 1 / 4
    assert [f.kind for f in findings] == ["composed_pages"]
    assert findings[0].detail == {"composed": [2, 3, 4, 5], "allowed": 1, "content_pages": 4, "examples": [3]}
    assert "pages 3" in findings[0].message and "at most 1" in findings[0].message


def test_a_quarter_of_the_content_pages_may_still_be_composed(tmp_path) -> None:
    from raven_ppt.tools.outline import _composed_pages

    template = _content_template(tmp_path / "house.pptx")
    pages = (
        _page(1, prototype=1),
        _page(2, prototype=3),
        _page(3, prototype=3),
        _page(4, prototype=3),
        _page(5),
        _page(6, prototype=4),
    )

    assert _composed_pages(Outline(takeaway="t", pages=pages), _Bound(template)) == []


def test_a_borrowed_page_counts_as_a_designed_start(tmp_path) -> None:
    from dataclasses import replace

    from raven_ppt.tools.outline import _composed_pages

    template = _content_template(tmp_path / "house.pptx")
    borrowed = replace(_page(2), borrowed="gold_panel_year_end_summary", prototype=5)
    pages = (
        _page(1, prototype=1),
        borrowed,
        _page(3, prototype=3),
        _page(4),
        _page(5, prototype=3),
        _page(6, prototype=4),
    )
    # Four content pages, one composed: within the quarter only because the borrowed
    # page counts as a designed start.

    assert _composed_pages(Outline(takeaway="t", pages=pages), _Bound(template)) == []


def test_without_a_template_nothing_is_asked_to_start_from_one() -> None:
    from raven_ppt.tools.outline import _composed_pages

    class _Free:
        template = None
        figures = ()

    assert _composed_pages(Outline(takeaway="t", pages=(_page(1), _page(2), _page(3))), _Free()) == []


def test_a_page_without_a_prototype_is_offered_the_examples_that_fit_what_it_carries(tmp_path) -> None:
    """The ask that said "pick the nearest example" was acted on by nobody until a page
    number stood beside each page; the suggestion is by materials, and it stays a
    suggestion."""
    from dataclasses import replace

    from raven_ppt.services.template.menu import menu
    from raven_ppt.tools.outline import _suggested_examples

    template = _content_template(tmp_path / "house.pptx")
    examples = [entry for entry in menu(template) if not entry.role and not entry.hidden]
    four_points = replace(_page(2, says=4), carries="four cards")
    a_table = replace(_page(3), carries="a table")

    assert _suggested_examples(four_points, examples) == [3], "four slots carry four points"
    assert _suggested_examples(a_table, examples) == [], "no example holds a table, so nothing is pretended to"
