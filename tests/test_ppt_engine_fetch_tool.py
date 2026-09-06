"""Bringing something from the web into a deck's materials.

The tests that matter are about the bytes. A URL ending in `.png` that returns
HTML, and a `Content-Type: image/png` on something that is not an image, are both
ordinary on the open web -- and writing either into the materials directory would
fail later somewhere that reads like a bug in the ingest.
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import httpx
import pytest

from raven.security.network import validate_url_target as _real_validate
from raven_ppt.contracts import Project
from raven_ppt.tools.fetch import MAX_BYTES, PptFetchTool

pytest.importorskip("PIL")


def _png(width: int = 8, height: int = 8) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, format="PNG")
    return buffer.getvalue()


def _described_png(description: str, *, key: str = "Description") -> bytes:
    """A PNG carrying its own description in a text chunk, as an export tool writes it."""
    from PIL import Image
    from PIL.PngImagePlugin import PngInfo

    chunks = PngInfo()
    chunks.add_text(key, description)
    buffer = io.BytesIO()
    Image.new("RGB", (320, 180), "white").save(buffer, format="PNG", pnginfo=chunks)
    return buffer.getvalue()


def _described_jpeg(description: str) -> bytes:
    """A JPEG carrying EXIF `ImageDescription`, NUL-padded the way cameras write it."""
    from PIL import Image

    image = Image.new("RGB", (320, 180), "white")
    exif = image.getexif()
    exif[0x010E] = f"{description}\x00\x00"
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", exif=exif.tobytes())
    return buffer.getvalue()


def _pptx() -> bytes:
    """A real presentation, because a downloaded template is now prepared as it
    lands: the two ways one arrives -- by URL and by path -- have to end in the
    same state, and a zip that only looks like a deck no longer reaches it."""
    from pptx import Presentation

    buffer = io.BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(buffer)
    return buffer.getvalue()


def _not_a_deck() -> bytes:
    """A zip shaped like a .pptx that python-pptx will not open."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("ppt/presentation.xml", "<presentation/>")
    return buffer.getvalue()


@pytest.fixture()
def fetch(monkeypatch, tmp_path: Path):
    """A tool whose HTTP layer is a MockTransport.

    Patched at httpx rather than at the tool, so the code under test is the code
    that ships -- including the streaming size guard, which a fake download method
    would have skipped.
    """

    def make(body: bytes, *, headers: dict | None = None, status: int = 200):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(status, content=body, headers=headers or {})

        transport = httpx.MockTransport(handler)
        real = httpx.AsyncClient

        def factory(*args, **kwargs):
            kwargs.pop("proxy", None)
            return real(*args, transport=transport, **kwargs)

        monkeypatch.setattr(httpx, "AsyncClient", factory)
        monkeypatch.setattr("raven.security.network.validate_url_target", lambda url: (True, ""))
        return PptFetchTool(tmp_path)

    return make


async def _run(tool: PptFetchTool, **kw) -> dict:
    args = {"project": "tarvis", "url": "https://example.com/thing"}
    return json.loads(await tool.execute(**{**args, **kw}))


@pytest.mark.asyncio
async def test_an_image_lands_in_the_decks_own_sources(fetch, tmp_path: Path) -> None:
    """One place, always. A fetch that could land elsewhere started a second pile, and
    the ingest that followed replaced the index rather than adding to it."""
    body = await _run(fetch(_png()), url="https://example.com/figure")

    assert body["ok"] is True and body["kind"] == "image"
    landed = Path(body["path"])
    assert landed.parent == Project(workspace=tmp_path, slug="tarvis").sources_dir
    assert landed.suffix == ".png" and landed.read_bytes()


@pytest.mark.asyncio
async def test_the_url_is_recorded_beside_what_it_fetched(fetch, tmp_path: Path) -> None:
    """The attribution `documents.source_urls` exists to provide and never had:
    nothing wrote the manifest it reads, so every fetched figure reached a slide with
    no source at all."""
    from raven_ppt.services.ingest import fetched

    await _run(fetch(_png()), url="https://example.com/figure")

    held = fetched(Project(workspace=tmp_path, slug="tarvis"))
    assert [source.url for source in held] == ["https://example.com/figure"]


@pytest.mark.asyncio
async def test_the_caption_the_caller_gives_reaches_the_figure_catalogue(fetch, tmp_path: Path) -> None:
    """The whole point of the parameter. A picture off a web page carries no caption in
    its bytes -- the words are in the HTML beside it -- so before this every fetched
    figure entered the catalogue with `caption: null` while a paper's figure arrived
    with the line printed under it. One live deck captioned a marketing banner as an
    architecture diagram with nothing to hold it to."""
    from raven_ppt.services.ingest import ingest_materials, load_catalogue

    tool = fetch(_png(320, 180))
    tool.ingest = ingest_materials
    caption = "Figure 2: Architectural overview of the Mem0 system"

    body = await _run(tool, url="https://example.com/figure.png", caption=caption)

    assert body["ok"] is True and body["caption"] == caption
    deck = Project(workspace=tmp_path, slug="tarvis")
    catalogue = load_catalogue(deck.ingest_dir / "figures.json", figures_dir=deck.ingest_dir / "figures")
    (asset,) = catalogue.values()
    assert asset.caption == caption
    assert asset.source_url == "https://example.com/figure.png"


@pytest.mark.asyncio
async def test_the_caption_is_recorded_beside_the_url_it_came_with(fetch, tmp_path: Path) -> None:
    from raven_ppt.services.ingest import fetched

    await _run(fetch(_png()), url="https://example.com/figure", caption="  Fig. 3\n  Retrieval latency  ")

    (held,) = fetched(Project(workspace=tmp_path, slug="tarvis"))
    assert held.caption == "Fig. 3 Retrieval latency"


@pytest.mark.asyncio
async def test_a_picture_that_describes_itself_needs_no_caption_from_the_caller(fetch, tmp_path: Path) -> None:
    """What is reachable from the bytes is taken rather than asked for twice."""
    from raven_ppt.services.ingest import fetched

    await _run(fetch(_described_png("Figure 4: token cost per turn")))

    (held,) = fetched(Project(workspace=tmp_path, slug="tarvis"))
    assert held.caption == "Figure 4: token cost per turn"


@pytest.mark.asyncio
async def test_an_exif_image_description_is_read_and_its_padding_dropped(fetch, tmp_path: Path) -> None:
    """EXIF strings are NUL-padded, and a NUL in the catalogue makes the file a
    search tool treats as binary."""
    from raven_ppt.services.ingest import fetched

    await _run(fetch(_described_jpeg("Mem0 memory lifecycle")))

    (held,) = fetched(Project(workspace=tmp_path, slug="tarvis"))
    assert held.caption == "Mem0 memory lifecycle"


@pytest.mark.asyncio
async def test_the_pages_words_beat_the_files_own_metadata(fetch, tmp_path: Path) -> None:
    """The caller's caption is what the page printed next to this picture; an
    `ImageDescription` field is whatever the file was exported with."""
    from raven_ppt.services.ingest import fetched

    await _run(fetch(_described_png("stock photo 118372")), caption="Figure 1: the extraction phase")

    (held,) = fetched(Project(workspace=tmp_path, slug="tarvis"))
    assert held.caption == "Figure 1: the extraction phase"


@pytest.mark.asyncio
async def test_a_second_fetch_without_a_caption_does_not_blank_the_first(fetch, tmp_path: Path) -> None:
    """Same URL, twice: a retry, or the same picture wanted for a second page. That is
    not new information about these bytes, so it must not take the words away."""
    from raven_ppt.services.ingest import fetched

    # One tool for both calls: the factory wraps whatever `httpx.AsyncClient` is when
    # it runs, so calling it twice in one test nests the transport into itself.
    tool = fetch(_png(320, 180))
    await _run(tool, caption="Figure 5: end-to-end latency")
    await _run(tool)

    (held,) = fetched(Project(workspace=tmp_path, slug="tarvis"))
    assert held.caption == "Figure 5: end-to-end latency"


@pytest.mark.asyncio
async def test_a_caption_on_a_document_is_refused_out_loud_rather_than_dropped(fetch, tmp_path: Path) -> None:
    """A caption belongs to one picture and a paper holds many, so it is not recorded
    for a document -- and saying so is the difference between this and the silence the
    parameter exists to end. The paper's figures carry the captions its pages print."""
    from raven_ppt.services.ingest import fetched

    body = await _run(fetch(b"%PDF-1.7\n trailer"), caption="Figure 2: architectural overview")

    assert body["ok"] is True and "caption" not in body
    assert "describes one picture" in body["note"]
    assert [held.caption for held in fetched(Project(workspace=tmp_path, slug="tarvis"))] == [""]


@pytest.mark.asyncio
async def test_the_suffix_comes_from_the_bytes_not_the_url(fetch, tmp_path: Path) -> None:
    """A URL ending in .png that returns HTML is ordinary on the open web."""
    body = await _run(fetch(b"<html><body>not an image</body></html>"), url="https://example.com/photo.png")

    assert body["ok"] is True and body["kind"] == "document"
    assert Path(body["path"]).suffix == ".html"


@pytest.mark.asyncio
async def test_a_lying_content_type_does_not_get_a_pass(fetch) -> None:
    body = await _run(fetch(b"\x00\x01\x02 not an image", headers={"content-type": "image/png"}))
    assert body["ok"] is False
    assert "not a PDF, a PowerPoint file, an image, or text" in body["error"]


@pytest.mark.asyncio
async def test_a_pdf_is_recognised_by_its_header(fetch) -> None:
    body = await _run(fetch(b"%PDF-1.7\n trailer"))
    assert body["kind"] == "document" and Path(body["path"]).suffix == ".pdf"


@pytest.mark.asyncio
async def test_a_pptx_lands_in_the_template_slot_not_the_materials(fetch, tmp_path: Path) -> None:
    """A deck is not a source to quote; it is a house style to work inside."""
    body = await _run(fetch(_pptx()), url="https://example.com/house-style.pptx")

    assert body["ok"] is True and body["kind"] == "template"
    project = Project(workspace=tmp_path, slug="tarvis")
    assert Path(body["path"]).parent == project.root / "template"
    assert "house style" in body["next_step"]


@pytest.mark.asyncio
async def test_a_downloaded_template_is_prepared_where_it_lands(fetch, tmp_path: Path) -> None:
    """The two ways a template arrives have to end in the same state. Downloading
    one and leaving it in the slot was the earlier behaviour, and it meant a
    template that came by URL was never prepared, so nothing downstream saw one."""
    from raven_ppt.services.template import bound

    body = await _run(fetch(_pptx()), url="https://example.com/house-style.pptx")

    assert body["ok"] is True
    assert bound(Project(workspace=tmp_path, slug="tarvis")) is not None
    assert body["template"] == "house-style.pptx"


@pytest.mark.asyncio
async def test_a_pptx_that_will_not_open_is_refused_rather_than_stored(fetch, tmp_path: Path) -> None:
    body = await _run(fetch(_not_a_deck()))

    assert body["ok"] is False
    assert "does not open as a presentation" in body["error"]
    assert not list((Project(workspace=tmp_path, slug="tarvis").root / "template").glob("*.pptx"))


@pytest.mark.asyncio
async def test_a_zip_that_is_not_a_presentation_is_refused(fetch) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hello")
    body = await _run(fetch(buffer.getvalue()))
    assert body["ok"] is False and "not a PowerPoint file" in body["error"]


@pytest.mark.asyncio
async def test_a_declared_size_over_the_limit_is_refused_before_the_body_arrives(fetch) -> None:
    body = await _run(fetch(b"%PDF-1.7", headers={"content-length": str(MAX_BYTES + 1)}))
    assert body["ok"] is False and "32MB" in body["error"]


@pytest.mark.asyncio
async def test_an_undeclared_oversized_body_is_refused_as_it_arrives(fetch) -> None:
    """A server may not declare a length, so the bytes are counted too."""
    body = await _run(fetch(b"%PDF-" + b"x" * (MAX_BYTES + 1)))
    assert body["ok"] is False and "32MB" in body["error"]


@pytest.mark.asyncio
async def test_a_failed_request_says_so_rather_than_writing_nothing_quietly(fetch) -> None:
    body = await _run(fetch(b"nope", status=404))
    assert body["ok"] is False and "download failed" in body["error"]


@pytest.mark.asyncio
async def test_there_is_no_materials_path_to_point_anywhere(fetch) -> None:
    """The parameter is gone, not defaulted: it existed to let a fetch land somewhere
    other than the deck's sources, and there was never a good reason to."""
    from raven_ppt.tools.fetch import PptFetchTool

    assert "materials_dir" not in PptFetchTool(Path("/tmp")).parameters["properties"]


@pytest.mark.asyncio
async def test_a_bad_project_name_is_refused_before_the_download(fetch) -> None:
    body = await _run(fetch(_png()), project="../etc")
    assert body["ok"] is False and "usable project name" in body["error"]


@pytest.mark.asyncio
async def test_a_hostile_filename_is_reduced_to_something_safe(fetch, tmp_path: Path) -> None:
    body = await _run(fetch(_png()), filename="../../etc/passwd")
    assert Path(body["path"]).parent == Project(workspace=tmp_path, slug="tarvis").sources_dir
    assert Path(body["path"]).name == "passwd.png"


@pytest.mark.asyncio
async def test_a_body_of_control_bytes_is_binary_however_well_it_decodes(fetch) -> None:
    """UTF-8 accepts control characters, so decoding is not a test for text.

    Without this, a truncated video saved as `.md` becomes the ingest's problem
    to explain.
    """
    body = await _run(fetch(bytes(range(1, 32)) * 40))
    assert body["ok"] is False and "not a PDF" in body["error"]


@pytest.mark.asyncio
async def test_an_empty_body_is_not_a_document(fetch) -> None:
    body = await _run(fetch(b"   \n\n  "))
    assert body["ok"] is False


@pytest.mark.asyncio
async def test_ordinary_prose_with_punctuation_and_cjk_is_still_text(fetch) -> None:
    body = await _run(fetch("# 标题\n\n正文 — with an em dash, 90% coverage.\n".encode()))
    assert body["ok"] is True and Path(body["path"]).suffix == ".md"


SVG = (
    b'<svg xmlns="http://www.w3.org/2000/svg" width="240" height="60" viewBox="0 0 240 60">'
    b'<rect width="240" height="60" fill="#155FFD"/><text x="12" y="40" fill="#fff">mem0</text></svg>'
)


def test_a_brand_mark_arrives_as_svg_and_is_kept_as_a_picture() -> None:
    """A logo saved as `.md` is a competitor analysis with no competitor's mark on it.

    Rasterising needs cairo's native library, which `cairosvg` cannot supply itself, so
    this asserts the conversion only where the conversion is possible. The degraded path
    is a case of its own below rather than a looser assertion here.
    """
    try:
        # Not `importorskip`: it catches ImportError only, and a missing libcairo
        # surfaces as OSError from cairocffi -- which would error this test rather
        # than skip it, the same trap the guard under test exists for.
        import cairosvg  # noqa: F401
    except (ImportError, OSError) as exc:
        pytest.skip(f"rasterising an SVG needs the native cairo library: {exc}")
    from raven_ppt.tools.fetch import _sniff

    payload, suffix, kind = _sniff(SVG)

    assert (suffix, kind) == (".png", "image")
    assert payload.startswith(b"\x89PNG"), "python-pptx places rasters only"


def test_an_svg_that_will_not_draw_is_still_the_text_it_is() -> None:
    from raven_ppt.tools.fetch import _sniff

    payload, suffix, kind = _sniff(b"<svg this is not really an svg at all")

    assert kind == "document"
    assert payload == b"<svg this is not really an svg at all"


def test_the_other_three_kinds_come_back_byte_for_byte() -> None:
    from raven_ppt.tools.fetch import _sniff

    for blob, expected in (
        (b"<html><body>hi</body></html>", ".html"),
        (b"# a markdown note", ".md"),
        (b"%PDF-1.4 trailer", ".pdf"),
    ):
        payload, suffix, _ = _sniff(blob)
        assert suffix == expected
        assert payload == blob


def test_a_machine_without_libcairo_keeps_the_svg_instead_of_failing(monkeypatch) -> None:
    """`import cairosvg` pulls in cairocffi, which raises OSError -- not ImportError --
    when the native library is absent, so the guard has to catch that too."""
    import builtins
    import sys

    from raven_ppt.tools.fetch import _sniff

    real_import = builtins.__import__

    def no_libcairo(name, *args, **kwargs):
        if name == "cairosvg":
            raise OSError("no library called 'cairo-2' was found")
        return real_import(name, *args, **kwargs)

    monkeypatch.delitem(sys.modules, "cairosvg", raising=False)
    monkeypatch.setattr(builtins, "__import__", no_libcairo)

    payload, suffix, kind = _sniff(SVG)

    assert kind == "document", "an SVG that cannot be drawn is still the text it is"
    assert payload == SVG


@pytest.mark.asyncio
async def test_a_host_never_reached_names_the_setting_that_would_reach_it(fetch, monkeypatch, tmp_path: Path) -> None:
    """The refusal a run actually got was `the download failed: ` and nothing else.

    `httpx.ConnectTimeout("")` renders empty, so the class was the only thing there
    was to say and it was not said. And the cause is knowable: every client in this
    tool is built `trust_env=False`, so a machine that reaches the internet through
    a proxy goes direct here however its environment is set -- measured at 50s per
    timeout while curl through that same proxy answered in 0.86s.
    """
    from raven_ppt.tools.fetch import PptFetchTool

    tool = PptFetchTool(tmp_path)

    async def _timeout(url):
        raise httpx.ConnectTimeout("")

    monkeypatch.setattr(tool, "_download", _timeout)
    monkeypatch.setattr("raven.security.network.validate_url_target", lambda url: (True, ""))
    said = await _run(tool, url="https://example.com/a.png")

    assert said["ok"] is False
    assert "ConnectTimeout" in said["error"], said["error"]
    assert "tools.web.proxy" in said["hint"]
    assert "HTTPS_PROXY" in said["hint"]

    # A configured proxy means the network is already being routed, so the hint
    # would be wrong; and a status is the server talking, which needs no hint.
    routed = PptFetchTool(tmp_path, proxy="http://127.0.0.1:7890")
    monkeypatch.setattr(routed, "_download", _timeout)
    assert "hint" not in await _run(routed, url="https://example.com/a.png")

    async def _refused(url):
        raise httpx.HTTPStatusError("403", request=None, response=None)

    monkeypatch.setattr(tool, "_download", _refused)
    assert "hint" not in await _run(tool, url="https://example.com/a.png")


@pytest.mark.asyncio
async def test_an_address_the_network_guard_refuses_is_not_downloaded(fetch, monkeypatch, tmp_path: Path) -> None:
    """`validate_url_target` answers, it does not throw.

    The call read it by catching an exception, so every refused address went
    through: loopback, the link-local range, and the cloud metadata endpoint that
    hands out credentials to whoever asks from inside the host.
    """
    tool = fetch(_png())
    # After the factory, which installs the permissive stub the other cases need.
    monkeypatch.setattr("raven.security.network.validate_url_target", _real_validate)
    out = await _run(tool, url="http://169.254.169.254/latest/meta-data/iam/security-credentials/")

    assert out.get("ok") is False, "the metadata endpoint must not be fetched"
    assert "cannot be fetched" in json.dumps(out), out
    landed = list((tmp_path / "tarvis").rglob("*")) if (tmp_path / "tarvis").exists() else []
    assert not [p for p in landed if p.is_file()], "nothing may be written for a refused address"


@pytest.mark.asyncio
async def test_a_fetched_image_hands_back_the_id_the_catalogue_keys_it_under(fetch, tmp_path: Path) -> None:
    """Nothing else in the reply can be turned into one.

    `ppt_figure_inspect` takes catalogue ids and nothing else, and the catalogue keys a
    picture as `<stem>-<digest>`. A measured run fetched ten images, called inspect with
    the filenames it had just passed to this tool, was told none of them was in the
    catalogue, then guessed at the digests and got their length wrong -- two iterations
    spent recovering a string the fetch already had on disk.
    """
    from raven_ppt.services.ingest import ingest_materials, load_catalogue

    tool = fetch(_png(320, 180))
    tool.ingest = ingest_materials

    body = await _run(tool, url="https://example.com/figure.png", filename="diagram.png")

    deck = Project(workspace=tmp_path, slug="tarvis")
    catalogue = load_catalogue(deck.ingest_dir / "figures.json", figures_dir=deck.ingest_dir / "figures")
    assert body["figure_id"] in catalogue, "the id has to be one inspect will accept"
    assert body["figure_id"].startswith("diagram-")


@pytest.mark.asyncio
async def test_the_id_a_fetch_hands_back_is_one_figure_inspect_accepts(fetch, tmp_path: Path) -> None:
    """End to end across the two tools, because the ids agreeing is the whole claim and
    each tool builds its own view of the catalogue."""
    from raven_ppt.services.ingest import ingest_materials
    from raven_ppt.tools.inspect import PptFigureInspectTool

    tool = fetch(_png(320, 180))
    tool.ingest = ingest_materials
    body = await _run(tool, url="https://example.com/figure.png", filename="diagram.png")

    inspect = PptFigureInspectTool(tmp_path, _Views())
    seen = await inspect.execute(project="tarvis", figures=[body["figure_id"]])
    reply = json.loads(seen.model_text if hasattr(seen, "model_text") else seen)

    assert reply["ok"] is True, reply
    assert [f["figure_id"] for f in reply["figures"]] == [body["figure_id"]]


class _Views:
    """Just the surface `ppt_figure_inspect` asks of the render service."""

    async def pages(self, *_args, **_kwargs):
        return {}

    def data_uri(self, png, budget=None):
        return f"data:image/png;base64,{Path(png).stem}"


def test_a_webp_picture_is_kept_as_a_png_the_deck_can_place() -> None:
    """Image hosts serve WebP; python-pptx places none of them, and the ingest
    refused the file as unreadable -- a run spent three fetches on one picture."""
    import io

    from PIL import Image

    from raven_ppt.tools.fetch import _sniff

    raw = io.BytesIO()
    Image.new("RGB", (12, 8), (200, 40, 40)).save(raw, format="WEBP")

    payload, suffix, kind = _sniff(raw.getvalue())

    assert (suffix, kind) == (".png", "image")
    assert payload.startswith(b"\x89PNG")
    with Image.open(io.BytesIO(payload)) as png:
        assert png.size == (12, 8)


def test_a_webp_keeps_the_description_it_carried_across_the_conversion() -> None:
    """A fetch without a caption falls back to what the file says about itself, and a
    re-encode that dropped the EXIF block would silently cost that provenance."""
    import io

    from PIL import Image

    from raven_ppt.tools.fetch import _embedded_caption, _sniff

    image = Image.new("RGB", (12, 8), (40, 40, 200))
    exif = image.getexif()
    exif[0x010E] = "A red brand mark"
    raw = io.BytesIO()
    image.save(raw, format="WEBP", exif=exif.tobytes())
    assert _embedded_caption(raw.getvalue()) == "A red brand mark"

    payload, _suffix, _kind = _sniff(raw.getvalue())

    assert _embedded_caption(payload) == "A red brand mark"
