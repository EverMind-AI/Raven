"""How much copy a template's box holds, said before the deck is built.

The reactive half reports a box that could not show its copy at the end of the
program. This is the same measurement on the template's own boxes, in the roster
reply, where the author is still deciding what to write into them -- so the tests
here are about what the band says, which boxes it refuses to speak for, and that it
never contradicts the copy the template's own designer put in the box.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from raven_ppt.services.template.capacity import (
    LEGEND,
    held,
    page_band,
    row_band,
)

pytest.importorskip("pptx")

_TEMPLATES = Path(__file__).resolve().parents[1] / "plugins-dist" / "ppt-engine" / "raven_ppt" / "assets" / "templates"
_needs_templates = pytest.mark.skipif(
    not any(_TEMPLATES.glob("*.pptx")),
    reason="the template payload is fetched, not tracked; see plugins-dist/ppt-engine/templates.manifest.json",
)

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _page():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    return presentation, presentation.slides.add_slide(presentation.slide_layouts[6])


def _box(slide, *, left=0.5, top=1.0, width=3.0, height=0.6, size=18, text="模板自己的一句话", bold=False):
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.text = text
    run = frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold = Pt(size), bold
    return shape


def _body_pr(shape):
    return shape.text_frame._txBody.find(f"{{{_A}}}bodyPr")


def _autofit(shape, kind: str | None) -> None:
    """Put one autofit element on the shape's own `bodyPr`, or none at all.

    python-pptx writes `spAutoFit` into every text box it creates, so a test that
    wants the spilling case has to take it out rather than leave the default alone.
    """
    from lxml import etree

    body = _body_pr(shape)
    for child in list(body):
        if etree.QName(child).localname in ("normAutofit", "spAutoFit", "noAutofit"):
            body.remove(child)
    if kind is not None:
        body.append(body.makeelement(f"{{{_A}}}{kind}", {}))


def _anchor(shape, anchor: str) -> None:
    _body_pr(shape).set("anchor", anchor)


def _same_placeholder(slide, shape):
    """The layout placeholder this slide placeholder inherits from, by index."""
    index = shape.placeholder_format.idx
    return next(candidate for candidate in slide.slide_layout.placeholders if candidate.placeholder_format.idx == index)


def test_the_band_is_a_range_in_both_scripts() -> None:
    """Two bounds, never one: the measurer reads wider than the render draws, so a
    single number would be reporting the measurer."""
    _, slide = _page()
    found = held(_box(slide))
    assert found is not None
    latin_low, latin_high = found.latin
    han_low, han_high = found.han
    assert 0 < latin_low < latin_high
    assert 0 < han_low < han_high
    # A CJK glyph is a full em and a Latin character about half of one, which is the
    # whole reason an English headline does not fit a Chinese template's box.
    assert han_high < latin_low


def test_a_taller_box_holds_proportionally_more() -> None:
    """The band is per box and not per line: the height decides how many lines show."""
    _, slide = _page()
    one_line = held(_box(slide, height=0.4))
    three_lines = held(_box(slide, top=3.0, height=1.4))
    assert one_line is not None and three_lines is not None
    assert one_line.lines == 1
    assert three_lines.lines >= 3
    assert three_lines.latin[0] >= 3 * one_line.latin[0]


def test_the_band_says_which_way_the_overflow_goes() -> None:
    """Top is the assumed direction and stays unsaid; the two that surprise are named."""
    _, slide = _page()
    top, centred, bottom = _box(slide), _box(slide, top=2.0), _box(slide, top=3.0)
    for shape, anchor in ((top, "t"), (centred, "ctr"), (bottom, "b")):
        _anchor(shape, anchor)
    said = {anchor: held(shape).said() for shape, anchor in ((top, "t"), (centred, "ctr"), (bottom, "b"))}
    assert "centred" not in said["t"] and "anchored" not in said["t"]
    assert "centred" in said["ctr"]
    assert "bottom-anchored (grows up)" in said["b"]


def test_the_band_names_all_three_outcomes() -> None:
    """Shrink, the frame growing itself, and spill -- the third is 16% of the bundled
    templates' boxes and the majority behaviour on one of them."""
    _, slide = _page()
    shrinking, growing, spilling = _box(slide), _box(slide, top=2.0), _box(slide, top=3.0)
    _autofit(shrinking, "normAutofit")
    _autofit(growing, "spAutoFit")
    _autofit(spilling, None)
    assert held(shrinking).behaviour == "shrinks"
    assert held(growing).behaviour == "frame grows"
    assert held(spilling).behaviour == "spills"


def test_an_inherited_autofit_is_not_a_shrink() -> None:
    """The shape's own `bodyPr` decides, as `_shrinks_to_fit` reads it: a live deck's
    closing headline inherited `normAutofit` from its layout and was drawn at full
    size and clipped."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    title.text_frame.text = "标题"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(44)
    _autofit(title, None)
    _autofit(_same_placeholder(slide, title), "normAutofit")
    assert held(title).behaviour == "spills"


def test_no_band_for_a_box_that_states_no_size_anywhere() -> None:
    """Rule 4: a band computed off a size nobody declared is a band that is wrong."""
    from pptx.util import Inches

    _, slide = _page()
    shape = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(3.0), Inches(0.6))
    shape.text_frame.word_wrap = True
    shape.text_frame.text = "一句没有声明字号的话"
    assert held(shape) is None


def test_no_band_for_a_box_with_wrapping_off() -> None:
    """Its copy runs on one line as far as it likes, so characters per line is not
    what limits it."""
    _, slide = _page()
    shape = _box(slide)
    shape.text_frame.word_wrap = False
    assert held(shape) is None


def test_an_inherited_size_is_said_to_be_inherited() -> None:
    """Resolved rather than skipped -- 323 of the bundled templates' silent boxes are
    reachable this way -- and labelled, because a layout's size and the size the page
    renders are not always the same number."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    title.text_frame.word_wrap = True
    title.text_frame.text = "继承字号的标题"
    layout_body = _same_placeholder(slide, title).text_frame._txBody
    paragraph = layout_body.findall(f"{{{_A}}}p")[0]
    properties = paragraph.makeelement(f"{{{_A}}}pPr", {})
    properties.append(properties.makeelement(f"{{{_A}}}defRPr", {"sz": "2800"}))
    paragraph.insert(0, properties)
    found = held(title)
    assert found is not None
    assert found.inherited is True
    assert found.size_pt == 28.0
    assert "28pt inherited holds" in found.said()


def test_a_stated_size_beats_the_chain() -> None:
    """A size on a run is a fact and is not reported as inherited."""
    _, slide = _page()
    found = held(_box(slide, size=20))
    assert found is not None and found.inherited is False
    assert found.size_pt == 20.0


def test_a_render_measured_size_overrides_the_file() -> None:
    """`house_style` has a render and the renderer has resolved every inheritance, so
    where it read a size off the page that is the size to measure against."""
    _, slide = _page()
    shape = _box(slide, size=18)
    assert held(shape, 36).size_pt == 36.0
    assert held(shape, 36).latin[1] < held(shape, 18).latin[1]


def test_the_page_band_names_the_tightest_copy_box_and_the_row() -> None:
    """The two an author gets wrong in different ways: what breaks first, and where a
    card coming back a step smaller than its neighbours is visible."""
    _, slide = _page()
    _box(slide, left=0.5, top=0.4, width=11.0, height=0.9, size=28, text="页面标题在这里")
    for index in range(4):
        card = _box(slide, left=0.5 + index * 2.8, top=2.0, width=2.4, height=0.5, size=14, text=f"卡片{index}的标题")
        _autofit(card, "normAutofit")
    band = page_band(list(slide.shapes))
    assert band.startswith("tightest of 4 alike at 14pt hold ")
    assert " en / " in band and " zh" in band
    assert band.endswith("shrinks")


def test_the_page_band_says_the_row_separately_when_it_is_not_the_tightest() -> None:
    _, slide = _page()
    _box(slide, left=0.5, top=0.4, width=1.4, height=0.4, size=24, text="很窄的标题")
    for index in range(3):
        _box(slide, left=0.5 + index * 3.6, top=2.0, width=3.4, height=1.2, size=12, text=f"第{index}段说明文字")
    band = page_band(list(slide.shapes))
    assert band.startswith("tightest 24pt holds ")
    assert "; 3 alike at 12pt hold " in band


def test_a_mark_is_not_the_pages_copy_limit() -> None:
    """A box the designer filled with one glyph is a step number or a bracket, not a
    copy slot: `gold_panel` page 2's 0.031in angle brackets would otherwise be that
    page's reported capacity."""
    _, slide = _page()
    _box(slide, left=0.5, top=0.4, width=0.04, height=0.14, size=14, text="<")
    _box(slide, left=1.0, top=1.0, width=4.0, height=0.8, size=18, text="这是页面上真正的正文，长度足够")
    band = page_band(list(slide.shapes))
    assert "18pt holds" in band
    assert "holds 0-0" not in band


def test_the_boxes_it_cannot_speak_for_are_counted_not_hidden() -> None:
    """A page whose band covers four of its nine boxes is a different offer from one
    whose band covers all four of its four."""
    _, slide = _page()
    _box(slide, top=1.0, width=4.0, height=0.8, size=18, text="有字号的正文，够长算作文案")
    silent = _box(slide, top=2.2, width=4.0, height=0.8, text="没有字号")
    silent.text_frame.paragraphs[0].runs[0].font.size = None
    unwrapped = _box(slide, top=3.4, width=4.0, height=0.8, size=18, text="不折行的一句话")
    unwrapped.text_frame.word_wrap = False
    band = page_band(list(slide.shapes))
    assert "1 does not wrap" in band
    assert "1 states no size" in band


def test_a_page_nothing_can_be_said_about_says_so() -> None:
    """`mint_memphis_thesis_defense` page 2: four numbers and a title with wrapping
    off, four bullets whose size is stated nowhere. A blank band there reads as room."""
    _, slide = _page()
    for index in range(2):
        shape = _box(slide, top=1.0 + index, width=4.0, height=0.8, size=72, text="01")
        shape.text_frame.word_wrap = False
    band = page_band(list(slide.shapes))
    assert band == "copy capacity unmeasured: 2 do not wrap"


def test_an_empty_page_has_no_band() -> None:
    _, slide = _page()
    assert page_band(list(slide.shapes)) == ""


def test_the_row_band_is_the_geometry_only() -> None:
    """A house-style row is a box the author will draw with `write`, which sets no
    autofit, so the template's own autofit does not carry over -- only the numbers."""
    _, slide = _page()
    shape = _box(slide, width=11.88, height=0.98, size=28)
    _autofit(shape, "normAutofit")
    said = row_band(shape)
    assert said.startswith("holds ")
    assert "shrinks" not in said and "pt" not in said


def test_the_row_band_is_empty_for_a_box_it_cannot_measure() -> None:
    from pptx.util import Inches

    _, slide = _page()
    shape = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(3.0), Inches(0.6))
    shape.text_frame.text = "没有字号"
    assert row_band(shape) == ""


def test_the_legend_defines_every_word_the_band_uses() -> None:
    """The band is six numbers and three verbs, and the reply has to say what they
    mean somewhere: the vocabulary rides once beside the lines that use it."""
    for word in ("holds", "en /", "zh", "centred", "bottom-anchored", "shrinks", "frame grows", "spills"):
        assert word in LEGEND


@_needs_templates
def test_the_templates_own_copy_is_never_contradicted() -> None:
    """The calibration, and the one result a wrong band cannot survive.

    The designer's own copy is the ground truth: over the ten bundled templates'
    1,116 boxes with a size on a run, 993 sit inside the conservative bound and 121
    between the bounds. The only two past the generous bound are `gold_panel` page 2's
    0.031in decorative brackets, which carry one glyph each and are marks rather than
    copy. A band that contradicts a real copy box is wrong by construction, because
    that copy is on the page in the file.
    """
    from pptx import Presentation

    from raven_ppt.services.measure.geometry import iter_shapes
    from raven_ppt.services.measure.width import is_cjk_char
    from raven_ppt.services.template.capacity import _stated_size

    inside = between = 0
    contradicted = []
    for path in sorted(_TEMPLATES.glob("*.pptx")):
        for number, slide in enumerate(Presentation(str(path)).slides, start=1):
            for shape in iter_shapes(slide.shapes):
                frame = getattr(shape, "text_frame", None)
                if frame is None or not frame.text.strip() or _stated_size(frame) is None:
                    continue
                found = held(shape)
                if found is None:
                    continue
                text = frame.text.strip()
                glyphs = sum(1 for char in text if is_cjk_char(char))
                letters = sum(1 for char in text if char.isalpha() and not is_cjk_char(char))
                han = bool(glyphs and glyphs >= letters)
                low, high = found.han if han else found.latin
                count = glyphs if han else len(text)
                if count <= low:
                    inside += 1
                elif count <= high:
                    between += 1
                else:
                    contradicted.append((path.stem, number, text[:12], count, low, high, found.is_copy))
    assert (inside, between) == (993, 121)
    assert [name for name, _, _, _, _, _, is_copy in contradicted if is_copy] == []
    assert len(contradicted) == 2


@_needs_templates
def test_the_page_that_started_this_would_have_said_so() -> None:
    """`green_aurora_tech_trends` pages 20 and 22 carry four 2.165x0.464in cards at
    20pt holding 6 or 7 glyphs each. A live run wrote 'CATL-HyperStrong: 60 GWh' into
    one of them -- 24 characters against a generous bound of 17."""
    from raven_ppt.services.template.menu import menu

    entries = {entry.number: entry for entry in menu(_TEMPLATES / "green_aurora_tech_trends.pptx")}
    for number in (20, 22):
        band = entries[number].capacity
        assert "tightest of 4 alike at 20pt hold 12-17 en / 7-9 zh, centred, spills" in band
        assert len("CATL-HyperStrong: 60 GWh") > 17


@_needs_templates
def test_the_band_costs_what_it_was_budgeted() -> None:
    """One clause a page, not a paragraph: the roster is read in full on every bind."""
    from raven_ppt.services.template.menu import menu

    pages = bands = 0
    for path in sorted(_TEMPLATES.glob("*.pptx")):
        entries = menu(path)
        pages += len(entries)
        bands += sum(len(entry.capacity) for entry in entries)
    assert pages == 211
    assert bands / pages < 130
