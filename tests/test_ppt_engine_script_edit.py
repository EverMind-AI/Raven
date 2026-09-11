"""Rewriting one page without breaking the other seventeen.

Every case here is a round that was lost. The prelude checks come from an edit
that took the deck from finished to unbuildable; the whole-write verification
comes from a round that built eighteen slides from seventeen mapped blocks.
"""

from __future__ import annotations

import textwrap

from raven_ppt.backends.script import blocks_rejection, page_blocks

PRELUDE = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches
    from ppt_theme import THEMES, rgb

    TH = THEMES["ink-graphite"]
    ACC = rgb(TH["accent"])
    MUT = rgb(TH["muted"])
    FNT = "Source Serif 4"

    prs = Presentation()


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    def title(slide, text):
        slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1)).text_frame.text = text


    def footnote(slide, text):
        slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11), Inches(0.4)).text_frame.text = text
    """
).lstrip()

SCRIPT = PRELUDE + textwrap.dedent(
    """

    # SLIDE 1
    one = new_slide()
    title(one, "Unified video segmentation")

    # SLIDE 2
    two = new_slide()
    title(two, "Target queries")

    prs.save(os.environ["PPT_OUTPUT"])
    """
)

BLOCKS_TEXT = "one = new_slide()\ntitle(one, 'x')\ntwo = new_slide()\ntitle(two, 'y')\n"


def test_a_page_cloned_from_the_template_counts_as_a_page() -> None:
    """`clone_page` creates a slide -- through `add_slide`, inside -- and the text scan
    could not see that. With the structural pages cloned, the block that draws slide 1
    holds a `clone_page` call and no `add_slide`, so the whole design pass refused:
    measured across two live runs, every finished build, which is the pass the deck was
    counting on for its layout."""
    lines = (
        PRELUDE + "\n# SLIDE 1\nslide = clone_page(prs, prototype(tpl, 1))\n" + "\n# SLIDE 2\nslide = new_slide()\n"
    ).splitlines(keepends=True)
    blocks = page_blocks(lines)

    assert blocks_rejection(lines, blocks, slide_count=2) is None


def test_the_ordinary_way_to_make_a_page_is_a_method_call() -> None:
    """`prs.slides.add_slide(LAY)` is how every drawn page starts, dots and all.

    The first fix for the cloned-page case excluded a leading dot -- to keep
    `ppt_template.prototype(tpl, 1)` from counting as a creator -- and that turned every
    drawn page into a block that draws nothing. Both real scripts measured refused on
    their first drawn page.
    """
    lines = (
        PRELUDE
        + "\n# SLIDE 1\nslide = prs.slides.add_slide(LAY)\ntitle(slide, 'drawn')\n"
        + "\n# SLIDE 2\nslide = clone_page(prs, prototype(tpl, 2))\n"
    ).splitlines(keepends=True)
    blocks = page_blocks(lines)

    assert blocks_rejection(lines, blocks, slide_count=2) is None


def test_a_block_count_that_disagrees_with_the_deck_is_refused() -> None:
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    reason = blocks_rejection(lines, blocks, slide_count=3)
    assert reason is not None and "2 page block(s)" in reason and "3 slide(s)" in reason


def test_the_advice_differs_by_where_the_mapping_came_from() -> None:
    """Execution is already exact, so a mismatch means the script has no per-page block."""
    lines = SCRIPT.splitlines(keepends=True)
    blocks = page_blocks(lines)
    from_comments = blocks_rejection(lines, blocks, 3, derived="comments")
    from_execution = blocks_rejection(lines, blocks, 3, derived="execution")
    assert "# SLIDE <n>" in (from_comments or "")
    assert "shared loop" in (from_execution or "")


def test_execution_mapping_accepts_a_page_created_through_nested_helpers() -> None:
    script = PRELUDE + "\n# SLIDE 1\nslide, body = content_page('title')\n"
    lines = script.splitlines(keepends=True)
    blocks = page_blocks(lines)

    assert "creates 0 slides" in (blocks_rejection(lines, blocks, 1, derived="comments") or "")
    assert blocks_rejection(lines, blocks, 1, derived="execution") is None


def test_pages_numbered_out_of_range_are_refused() -> None:
    script = SCRIPT.replace("# SLIDE 2", "# SLIDE 5")
    lines = script.splitlines(keepends=True)
    reason = blocks_rejection(lines, page_blocks(lines), slide_count=2)
    assert reason is not None and "not 1..2" in reason


def test_a_block_that_draws_two_slides_is_refused() -> None:
    script = SCRIPT.replace('title(two, "Target queries")', 'three = new_slide()\ntitle(three, "x")')
    lines = script.splitlines(keepends=True)
    reason = blocks_rejection(lines, page_blocks(lines), slide_count=2)
    assert reason is not None and "creates 2 slides" in reason
