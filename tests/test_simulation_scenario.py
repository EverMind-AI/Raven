"""The scenario reads cleanly, the owner plays drill cards, judges the drills and speaks to the Curator."""

import hashlib
import json
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest

from experimental.curator.raven_adapter.exploration import _REPOSITORY, _SOURCE_PATHS, Exploration
from experimental.curator.raven_adapter.worker import Execution
from experimental.iteration.exchange import ExchangeError
from experimental.iteration.protocols import Exchange
from experimental.simulation.__main__ import CHAINS, settings, starting_harness
from experimental.simulation.agency import NAME as REVIEW
from experimental.simulation.agency import Agency, partition
from experimental.simulation.employee import HOME, Fresh, hire, housed_files, refresh, skills, withheld
from experimental.simulation.record import PLACEHOLDERS
from experimental.simulation.scenario import BUNDLED, Scenario
from experimental.simulation.traveller import Traveller
from raven.config.schema import Config
from raven.contracts.llm_provider import LLMResponse, ToolCallRequest

TRAVEL = BUNDLED / "travel_agency"


class Provider:
    def __init__(self, *responses):
        self.responses = list(responses)
        self.requests = []

    async def chat_with_retry(self, **kwargs):
        self.requests.append(kwargs)
        return self.responses.pop(0)


def response(name, arguments):
    return LLMResponse(content=None, tool_calls=[ToolCallRequest(name, name, arguments)])


def exchange(user, reply):
    records = [{"kind": "runner.event", "event_type": "Text", "event": {"content": reply}}]
    return Exchange(user, Execution("t1", [], records, {}, "artifact-1"))


def verdicts(scenario, **overrides):
    rows = [{"id": criterion.id, "result": "pass"} for criterion in scenario.criteria]
    for row in rows:
        row.update(overrides.get(row["id"], {}))
    return rows


def shortfall(*criteria, cause="not_held", material=None, strength="should"):
    return {
        "criteria": list(criteria),
        "cause": cause,
        "material": material,
        "behavior": "Quote only from the price list, for any party.",
        "observed": "In the student drill it guessed a price.",
        "evidence": ["student: About 300 each."],
        "acceptance": "Every price quoted appears in the price list for that party and season.",
        "strength": strength,
    }


def test_bundled_scenario_reads_profile_materials_personas_and_criteria():
    scenario = Scenario.load(TRAVEL)
    assert scenario.profile and "#D9531E" in scenario.text("brand-design-guide")
    assert (scenario.materials["brand-design-guide"] / "brand-kit.pptx").is_file()
    assert (scenario.materials["plan-deck-template"] / "template.pptx").is_file()
    assert set(scenario.materials) == {
        "service-sop",
        "price-list",
        "value-package",
        "comfort-package",
        "premium-package",
        "booking-policy",
        "consultation-scripts",
        "plan-deck-spec",
        "brand-design-guide",
        "plan-deck-template",
        "plan-deck-sample",
        "handover-ticket",
    }
    assert [persona.name for persona in scenario.personas] == [
        "family",
        "premium",
        "professional",
        "returning",
        "student",
    ]
    assert all(persona.text for persona in scenario.personas)
    assert len({criterion.id for criterion in scenario.criteria}) == len(scenario.criteria) == 17
    assert {"intake-and-confirmation", "deck-facts-researched"} <= {
        criterion.id for criterion in scenario.criteria if criterion.severity == "red_line"
    }
    assert "brand-design-guide" not in scenario.initial and len(scenario.initial) == 11
    assert "19,000" in scenario.text("price-list")
    assert Scenario.load(TRAVEL.parent.parent / "nowhere" / "travel_agency").root == TRAVEL


def test_load_rejects_repeated_ids_and_unknown_materials(tmp_path):
    root = tmp_path / "scenario"
    shutil.copytree(TRAVEL, root)
    spec = json.loads((root / "scenario.json").read_text())
    spec["initial"] = ["brochure"]
    (root / "scenario.json").write_text(json.dumps(spec))
    with pytest.raises(ValueError, match="brochure"):
        Scenario.load(root)
    spec["initial"] = []
    spec["criteria"][1]["id"] = spec["criteria"][0]["id"]
    (root / "scenario.json").write_text(json.dumps(spec))
    with pytest.raises(ValueError, match="repeat"):
        Scenario.load(root)


async def test_traveller_speaks_in_plain_text_and_leaves_with_the_word_leave():
    scenario = Scenario.load(TRAVEL)
    student = next(persona for persona in scenario.personas if persona.name == "student")
    provider = Provider(
        LLMResponse(content="Hi, three of us want a cheap trip to Edinburgh in June."),
        LLMResponse(content="LEAVE."),
        LLMResponse(content=""),
        LLMResponse(content="Thanks, that is all.\nLEAVE"),
        LLMResponse(content="Great, see you LEAVE."),
        LLMResponse(content="Error: boom", finish_reason="error"),
    )
    traveller = Traveller(student, provider, model="sim")
    assert traveller.name == "student"
    assert await traveller.speak([]) == "Hi, three of us want a cheap trip to Edinburgh in June."
    exchanges = [exchange("Hi", "Welcome to Harbourlight. What is your budget?")]
    assert await traveller.speak(exchanges) is None
    assert await traveller.speak(exchanges) is None and len(provider.requests) == 2
    first = traveller.card
    assert await traveller.speak([]) is None
    assert await traveller.speak([]) == "Thanks, that is all."
    assert await traveller.speak(exchanges) is None and len(provider.requests) == 4
    assert await traveller.speak([]) == "Great, see you"
    with pytest.raises(ExchangeError, match="traveller:student: provider error"):
        await traveller.speak([])
    packet = json.loads(provider.requests[1]["messages"][1]["content"])
    assert packet["persona"] == first.text and first.trip.phone in packet["persona"]
    assert packet["conversation"] == [{"traveller": "Hi", "assistant": "Welcome to Harbourlight. What is your budget?"}]
    assert provider.requests[1]["model"] == "sim" and "tools" not in provider.requests[1]


async def test_agency_judges_every_criterion_and_hands_over_the_material_of_a_failed_check(tmp_path):
    scenario = Scenario.load(TRAVEL)
    skills = tmp_path / "skills"
    (skills / "brand-design-guide").mkdir(parents=True)
    (skills / "unrelated").mkdir()
    failing = verdicts(
        scenario,
        **{
            "quote-sheet-correct": {
                "result": "fail",
                "session": "student",
                "actual": "About 300 each.",
                "note": "Read the list.",
            }
        },
    )
    provider = Provider(
        response(
            REVIEW,
            {
                "verdicts": failing,
                "shortfalls": [shortfall("quote-sheet-correct")],
                "remark": "It made a price up for Mia.",
                "handover": ["brand-design-guide"],
            },
        ),
        response(REVIEW, {"verdicts": verdicts(scenario), "remark": "All good today."}),
    )
    plan = SimpleNamespace(
        understanding="Quotes must come from the price list.",
        changes=[
            SimpleNamespace(target="action.tool_gates", reason="Stop invented prices.", expected="List prices only.")
        ],
    )
    agency = Agency(
        scenario,
        provider,
        skills,
        workdir=tmp_path,
        plan="staged",
        deliver="pool",
        analysis="owner",
        reply=lambda: plan,
        model="sim",
    )
    agency.prepare()
    assert sorted(path.name for path in skills.iterdir()) == sorted([*scenario.initial, "unrelated"])
    sessions = {"student": [exchange("Cheap trip?", "About 300 each.")]}

    signal = await agency.evaluate(sessions)
    assert signal.source == "agency" and signal.satisfied is False
    assert signal.text == "It made a price up for Mia."
    assert (skills / "brand-design-guide" / "SKILL.md").is_file()
    assert {item.id for item in signal.items} == {criterion.id for criterion in scenario.criteria}
    failed = next(item for item in signal.items if item.result == "fail")
    assert failed.session == "student" and failed.actual == "About 300 each." and "price list" in failed.expected
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["conversations"] == {"student": [{"traveller": "Cheap trip?", "assistant": "About 300 each."}]}
    assert packet["your_earlier_reviews"] == [] and packet["research"] == {}
    assert packet["curator_reply"] == {"understanding": "Quotes must come from the price list.", "is_new": True}
    assert "action.tool_gates" not in json.dumps(packet, ensure_ascii=False)
    assert {row["severity"] for row in packet["criteria"]} == {"red_line", "standard"}
    assert set(packet["materials"]) == set(scenario.materials) and packet["deliverables"] == {}
    assert packet["given_to_the_assistant"] == list(scenario.initial) and packet["withheld"] == ["brand-design-guide"]

    signal = await agency.evaluate(sessions)
    assert signal.satisfied is True and signal.text == "All good today."
    assert agency.released == [*scenario.initial, "brand-design-guide"]
    earlier = json.loads(provider.requests[1]["messages"][1]["content"])["your_earlier_reviews"]
    assert earlier == [
        {
            "round": 1,
            "remark": "It made a price up for Mia.",
            "failed": ["quote-sheet-correct"],
            "raised": ["quote-sheet-correct"],
            "waiting_on_material": [],
            "handed_over": ["brand-design-guide"],
        }
    ]


async def test_in_a_dialog_the_owner_uploads_its_opening_materials_and_pastes_later_ones(tmp_path):
    scenario = Scenario.load(TRAVEL)
    skills, uploads = tmp_path / "home" / "skills", tmp_path / "home" / "uploads"
    failing = verdicts(scenario, **{"deck-aesthetics": {"result": "fail", "session": "student", "note": "Too busy."}})
    provider = Provider(
        response(
            REVIEW,
            {
                "verdicts": failing,
                "shortfalls": [shortfall("deck-aesthetics", cause="material_missing", material="brand-design-guide")],
                "remark": "Slides are too busy.",
                "handover": ["brand-design-guide"],
            },
        )
    )
    shared = tmp_path / "work" / "uploads"
    agency = Agency(scenario, provider, skills, workdir=tmp_path, uploads=uploads, shared=shared, analysis="owner")
    agency.prepare()
    assert not skills.exists() or not any(skills.iterdir())
    for place in (uploads, shared):
        assert (place / "service-sop" / "service-sop.md").read_text().startswith("# ")
        assert (place / "plan-deck-template" / "template.pptx").is_file()
        assert not (place / "brand-design-guide").exists()
    (opening,) = agency.opening()
    assert "uploads/service-sop/service-sop.md" in opening.attachments
    assert "uploads/plan-deck-template/template.pptx" in opening.attachments
    assert "- uploads/price-list/price-list.md\n" in opening.text and "{files}" not in opening.text
    signal = await agency.evaluate({"student": [exchange("Cheap trip?", "Here is the deck.")]})
    assert (
        signal.text.startswith("Slides are too busy.")
        and scenario.document("brand-design-guide").strip() in signal.text
    )
    assert "name: brand-design-guide" not in signal.text and not (skills / "brand-design-guide").exists()
    assert agency.released[-1] == "brand-design-guide"
    assert Agency(scenario, provider, skills, workdir=tmp_path, deliver="pool").opening() == ()
    with pytest.raises(ValueError, match="uploads folder"):
        Agency(scenario, provider, skills, workdir=tmp_path)


def test_a_partition_gives_every_material_exactly_once_and_finishes_before_the_last_round():
    scenario = Scenario.load(TRAVEL)
    (stage,) = scenario.plans.values()
    assert partition(stage, scenario.materials, rounds=4) == stage
    everything = tuple(scenario.materials)
    with pytest.raises(ValueError, match="exactly once"):
        partition((everything[:-1],), scenario.materials)
    with pytest.raises(ValueError, match="exactly once"):
        partition((everything, everything[:1]), scenario.materials)
    with pytest.raises(ValueError, match="onboarding"):
        partition(((), everything), scenario.materials)
    with pytest.raises(ValueError, match="before the last"):
        partition(tuple((name,) for name in everything), scenario.materials, rounds=4)


async def test_on_a_fixed_partition_the_plan_hands_materials_over_whatever_the_owner_picks(tmp_path):
    scenario = Scenario.load(TRAVEL)
    everything = list(scenario.materials)
    steps = (tuple(everything[:8]), tuple(everything[8:10]), tuple(everything[10:]))
    skills, uploads = tmp_path / "home" / "skills", tmp_path / "home" / "uploads"
    reviews = [
        {"verdicts": verdicts(scenario), "remark": "Here is more.", "handover": [everything[10]]},
        {"verdicts": verdicts(scenario), "remark": "And the rest."},
    ]
    provider = Provider(*(response(REVIEW, review) for review in reviews))
    agency = Agency(
        scenario, provider, skills, workdir=tmp_path, plan=steps, deliver="dialog", uploads=uploads, rounds=4
    )
    agency.prepare()
    assert agency.released == list(steps[0]) and not agency.chooses
    sessions = {"student": [exchange("Cheap trip?", "Here.")]}
    first = await agency.evaluate(sessions)
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["handing_over_now"] == list(steps[1]) and packet["you_choose_handover"] is False
    assert agency.released == [*steps[0], *steps[1]]
    assert all(scenario.document(name).strip() in first.text for name in steps[1])
    for name in steps[1]:
        assert (uploads / name / f"{name}.md").is_file() and f"uploads/{name}/{name}.md" in first.attachments
        assert f"- uploads/{name}/{name}.md" in first.text
    assert "{files}" not in first.text
    await agency.evaluate(sessions)
    assert agency.released == everything and agency.withheld == []


async def test_a_deck_template_handed_over_after_a_round_reaches_the_uploads_folder(tmp_path):
    scenario = Scenario.load(TRAVEL)
    stage = scenario.plans["by-stage"]
    assert "plan-deck-template" in stage[1]
    uploads = tmp_path / "home" / "uploads"
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Here is the deck material."}))
    agency = Agency(
        scenario, provider, tmp_path / "skills", workdir=tmp_path, plan=stage, deliver="dialog", uploads=uploads
    )
    agency.prepare()
    assert not (uploads / "plan-deck-template").exists()
    signal = await agency.evaluate({"student": [exchange("Cheap trip?", "Here.")]})
    assert (uploads / "plan-deck-template" / "template.pptx").is_file()
    assert "uploads/plan-deck-template/template.pptx" in signal.attachments
    assert "- uploads/plan-deck-template/template.pptx" in signal.text


async def test_by_need_the_owner_chooses_and_the_rest_comes_before_the_last_round(tmp_path):
    scenario = Scenario.load(TRAVEL)
    skills = tmp_path / "skills"
    remaining = [name for name in scenario.materials if name not in scenario.initial]
    provider = Provider(*(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "ok"}) for _ in range(2)))
    agency = Agency(scenario, provider, skills, workdir=tmp_path, plan="staged", deliver="pool", rounds=3)
    agency.prepare()
    sessions = {"student": [exchange("Cheap trip?", "Here.")]}
    await agency.evaluate(sessions)
    assert agency.withheld == remaining and agency.chooses
    await agency.evaluate(sessions)
    packets = [json.loads(request["messages"][1]["content"]) for request in provider.requests]
    assert packets[0]["handing_over_now"] == [] and packets[1]["handing_over_now"] == remaining
    assert agency.withheld == [] and all((skills / name / "SKILL.md").is_file() for name in remaining)


def test_the_starting_harness_fingerprint_changes_only_when_a_products_files_do(tmp_path):
    product = tmp_path / "agents" / "raven-ppt"
    (product / "__pycache__").mkdir(parents=True)
    (product / "run.py").write_text("print('deck')")
    first = starting_harness(tmp_path)
    (product / "__pycache__" / "run.cpython.pyc").write_bytes(b"cache")
    assert starting_harness(tmp_path)["products"] == first["products"]
    (product / "run.py").write_text("print('deck v2')")
    assert starting_harness(tmp_path)["products"]["raven-ppt"] != first["products"]["raven-ppt"]
    assert first["raven_commit"] is None and first["raven_modified"] is False
    assert first["raven_diff"] == hashlib.sha256(b"").hexdigest()


def test_every_chain_names_a_valid_way_to_cultivate_and_settings_carry_no_credential(tmp_path):
    for chain in CHAINS.values():
        assert chain["deliver"] == "dialog" and chain["disclose"] in ("all", "staged", *Scenario.load(TRAVEL).plans)
    config = tmp_path / "config.json"
    args = SimpleNamespace(
        argv=[f"--config={config}", "--chain", "staged", "--home", str(tmp_path / "home"), "--seed", "7"],
        chain="staged",
        scenario=TRAVEL,
        deliver="dialog",
        disclose="staged",
        curator="improve",
        targets="all",
        analysis="analyst",
        rounds=4,
        turns=14,
        repeats=1,
        cards=["family"],
        concurrent_drills=3,
        seed=7,
        without=["plan-deck-sample"],
        curator_model="z-ai/glm-5.3",
        analyst_model=None,
        simulation_model=None,
        traveller_model="deepseek/deepseek-flash",
        subagent_model="z-ai/glm-5.3-flashx",
        curator_effort="high",
        traveller_effort="low",
        config=config,
        home=tmp_path / "home",
    )
    (tmp_path / "home" / "playbooks" / "plan").mkdir(parents=True)
    (tmp_path / "home" / "playbooks" / "plan" / "playbook.md").write_text("spec")
    (tmp_path / "home" / "sessions").mkdir()
    (tmp_path / "home" / "sessions" / "old.jsonl").write_text("{}")
    written = settings(args, "z-ai/glm-5.3-flashx", "low", "medium")
    assert written["efforts"] == {"employee": "low", "employee_tier": "medium", "curator": "high", "traveller": "low"}
    assert written["baseline"] == {"playbooks/plan/playbook.md": hashlib.sha256(b"spec").hexdigest()}
    start = written["starting_harness"]
    assert {"raven-research", "raven-ppt"} <= set(start["products"]) and start["raven_commit"]
    assert written["chain"] == "staged" and written["scenario"] == "travel_agency"
    assert written["analysis"] == "analyst"
    assert written["argv"] == [
        f"--config={PLACEHOLDERS['--config']}",
        "--chain",
        "staged",
        "--home",
        str(tmp_path / "home"),
        "--seed",
        "7",
    ]
    assert written["seed"] == 7 and written["without"] == ["plan-deck-sample"] and written["concurrent_drills"] == 3
    assert written["models"] == {
        "employee": "z-ai/glm-5.3-flashx",
        "curator": "z-ai/glm-5.3",
        "analyst": "z-ai/glm-5.3",
        "simulation": "z-ai/glm-5.3",
        "traveller": "deepseek/deepseek-flash",
        "subagents": "z-ai/glm-5.3-flashx",
    }
    assert str(config) not in json.dumps(written)


async def test_agency_sends_back_incomplete_misattributed_or_overreaching_reviews(tmp_path):
    scenario = Scenario.load(TRAVEL)
    partial = verdicts(scenario)[:-1]
    wrong = verdicts(scenario, **{"deck-delivered-and-consistent": {"result": "fail", "session": "ghost"}})
    provider = Provider(
        response(REVIEW, {"verdicts": partial, "remark": "Fine."}),
        response(REVIEW, {"verdicts": wrong, "remark": "Fine."}),
        response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine.", "handover": ["price-list"]}),
        response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine."}),
    )
    agency = Agency(scenario, provider, tmp_path / "skills", workdir=tmp_path, plan="all", deliver="pool")
    agency.prepare()
    assert sorted(agency.released) == sorted(scenario.materials) and agency.withheld == []
    signal = await agency.evaluate({"student": [exchange("Hi", "Hello")]})
    assert signal.satisfied is True
    errors = [json.loads(m["content"])["error"] for m in provider.requests[3]["messages"] if m["role"] == "tool"]
    assert "one verdict per criterion" in errors[0] and "ghost" in errors[1]
    assert "withheld" in errors[2]


async def test_agency_gives_up_within_its_call_budget(tmp_path):
    scenario = Scenario.load(TRAVEL)
    provider = Provider(*(response(REVIEW, {"verdicts": [], "remark": ""}) for _ in range(2)))
    agency = Agency(scenario, provider, tmp_path / "skills", workdir=tmp_path, deliver="pool", max_calls=2)
    with pytest.raises(ExchangeError, match="budget"):
        await agency.evaluate({})
    with pytest.raises(ValueError, match="plan"):
        Agency(scenario, provider, tmp_path / "skills", workdir=tmp_path, plan="later")


def test_the_curator_snapshot_does_not_carry_the_scenario_or_the_experiment_docs():
    mounted = [_REPOSITORY / relative for relative in _SOURCE_PATHS]
    hidden = [TRAVEL, _REPOSITORY / "experimental" / "docs"]
    assert not any(path.is_relative_to(mount) for path in hidden for mount in mounted)
    assert (_REPOSITORY / "experimental" / "curator") in mounted


def test_the_employees_curator_is_kept_from_the_simulation_its_judges_and_the_scenario(tmp_path):
    repository = tmp_path / "repository"
    files = {
        "raven/helper.py": "VALUE = 1\n",
        "tests/test_raven_helper.py": "def test_value(): pass\n",
        "tests/test_simulation_cards.py": "def test_cards(): pass\n",
        "tests/test_analyst_run.py": "def test_run(): pass\n",
        "tests/test_cards_again.py": "from experimental.simulation.cards import draw\n",
        "tests/test_loop.py": "from experimental.iteration.run import run\n",
        "tests/test_named.py": f"SCENARIO = '{TRAVEL.name}'\n",
    }
    for relative, text in files.items():
        (repository / relative).parent.mkdir(parents=True, exist_ok=True)
        (repository / relative).write_text(text)
    inspection = SimpleNamespace(facts={}, sources={})
    exploration = Exploration(
        Config(),
        inspection,
        repository=repository,
        source_paths=("raven", "tests"),
        withheld=withheld(Scenario.load(TRAVEL)),
        root=tmp_path / "x",
    )
    assert {path.name for path in (exploration.root / "source" / "tests").iterdir()} == {"test_raven_helper.py"}


def test_the_employee_works_from_its_own_copy_of_the_home(tmp_path):
    scenario = Scenario.load(TRAVEL)
    home = tmp_path / "shared-home"
    (home / "skills" / "brand-design-guide").mkdir(parents=True)
    (home / "skills" / "brand-design-guide" / "personal.txt").write_text("mine")
    (home / "sessions").mkdir()
    (home / "sessions" / "old.jsonl").write_text("{}")
    config = tmp_path / "config.json"
    config.write_text(
        json.dumps(
            {
                "providers": {"deepseek": {"api_key": "k"}},
                "agents": {"defaults": {"model": "deepseek/x", "provider": "deepseek"}},
            }
        )
    )
    root = tmp_path / "run"
    worker = hire(scenario, config, workdir=tmp_path, root=root, home=home)
    assert worker.withheld == withheld(scenario)
    pool = skills(worker)
    assert pool == root / HOME / "skills" and (pool / "brand-design-guide" / "personal.txt").is_file()
    assert not (root / HOME / "sessions").exists()
    Agency(scenario, Provider(), pool, workdir=tmp_path, plan="staged", deliver="pool").prepare()
    assert not (pool / "brand-design-guide").exists() and (pool / "service-sop" / "SKILL.md").is_file()
    assert (home / "skills" / "brand-design-guide" / "personal.txt").read_text() == "mine"


def test_the_employee_houses_its_external_subagents_homes_beside_its_own_harness(tmp_path, monkeypatch):
    from experimental.simulation import employee
    from raven.config.schema import ThirdPartyAcpSubagentConfig

    def row(name, enabled=True):
        return ThirdPartyAcpSubagentConfig(name=name, kind="acp", command="run", enabled=enabled, env={"A": "1"})

    monkeypatch.setattr(employee, "discover_product_rows", lambda: [row("Raven-PPT"), row("Raven-Research", False)])
    monkeypatch.setenv("PYTHONTZPATH", "/zones")
    config = tmp_path / "config.json"
    config.write_text(
        json.dumps(
            {
                "providers": {"deepseek": {"api_key": "k"}},
                "agents": {"defaults": {"model": "deepseek/x", "provider": "deepseek"}},
            }
        )
    )
    worker = hire(Scenario.load(TRAVEL), config, workdir=tmp_path, root=tmp_path / "run")
    rows = {item.name: item for item in worker.baseline.config.subagents.agents}
    housed = tmp_path / "run" / HOME / "subagents" / "Raven-PPT"
    assert rows["Raven-PPT"].env == {"A": "1", "PYTHONTZPATH": "/zones", "PPT_ACP_HOME": str(housed)}
    assert housed.is_dir() and "Raven-Research" not in rows


def test_the_hired_subagent_homes_are_kept_without_their_session_logs(tmp_path):
    root = tmp_path / "subagents"
    for relative, text in {
        "Raven-PPT/agent_memory/profile/agent.md": "deck method",
        "Raven-PPT/user_memory/episodic/episodes.md": "",
        "Raven-PPT/sessions/one.jsonl": "{}",
        "Raven-PPT/memory/.curator/trace.jsonl": "{}",
    }.items():
        (root / relative).parent.mkdir(parents=True, exist_ok=True)
        (root / relative).write_text(text)
    assert housed_files(tmp_path) == {
        "Raven-PPT/agent_memory/profile/agent.md": b"deck method",
        "Raven-PPT/user_memory/episodic/episodes.md": b"",
    }
    assert housed_files(tmp_path / "elsewhere") == {}


def test_each_drill_starts_from_the_hired_subagent_homes_with_the_installed_revision_on_top(tmp_path):
    seed = {"Raven-PPT/user_memory/episodic/episodes.md": b"", "Raven-PPT/user_memory/profile/user.md": b"default"}
    root = tmp_path / "subagents" / "Raven-PPT"
    (root / "user_memory" / "episodic").mkdir(parents=True)
    (root / "user_memory" / "episodic" / "episodes.md").write_text("met Mr Chen on the last drill")
    (root / "sessions").mkdir()
    (root / "sessions" / "one.jsonl").write_text("{}")
    authored = {"Raven-PPT/user_memory/profile/user.md": "house style"}
    assert sorted(refresh(tmp_path, seed, authored)) == sorted(seed)
    assert (root / "user_memory" / "episodic" / "episodes.md").read_text() == ""
    assert (root / "user_memory" / "profile" / "user.md").read_text() == "house style"
    assert (root / "sessions" / "one.jsonl").is_file()
    assert refresh(tmp_path, seed, authored) == []


async def test_a_fresh_trial_refreshes_the_subagent_homes_before_it_runs(tmp_path):
    from experimental.curator.harness import Artifact

    seen = []

    class Drill:
        async def run(self, worker):
            housed = tmp_path / "subagents" / "Raven-PPT"
            seen.append({name: (housed / name).read_text() for name in ("user.md", "USER.md")})
            return {"drill": []}

    child = SimpleNamespace(
        baseline=SimpleNamespace(config=SimpleNamespace(workspace_path=tmp_path / "subagents" / "Raven-PPT")),
        artifact=Artifact(values={"memory.prompt": {"USER.md": "authored"}}),
    )
    worker = SimpleNamespace(
        baseline=SimpleNamespace(config=SimpleNamespace(workspace_path=tmp_path)), children={"Raven-PPT": child}
    )
    (tmp_path / "subagents" / "Raven-PPT").mkdir(parents=True)
    (tmp_path / "subagents" / "Raven-PPT" / "user.md").write_text("remembered from the last drill")
    assert await Fresh(Drill(), {"Raven-PPT/user.md": b"seeded"}).run(worker) == {"drill": []}
    assert seen == [{"user.md": "seeded", "USER.md": "authored"}]


def test_the_housed_subagents_run_on_a_model_of_their_own_through_the_employees_provider(tmp_path, monkeypatch):
    from experimental.simulation import employee
    from raven.config.schema import ThirdPartyAcpSubagentConfig

    def row(name):
        return ThirdPartyAcpSubagentConfig(name=name, kind="acp", command=f"py /agents/{name}/run.py", enabled=True)

    folder = tmp_path / "product"
    folder.mkdir()
    shipped = {
        "agents": {"defaults": {"model": "vendor/other", "provider": "ppt", "maxTokens": 9, "reasoningEffort": "max"}},
        "providers": {"ppt": {"apiBase": "https://gateway.example/v1"}},
        "acp": {
            "modes": {"medium": {"reasoningEffort": "low"}, "high": {"reasoningEffort": "high"}},
            "defaultMode": "high",
        },
        "tools": {"disabledTools": ["spawn"]},
        "plugins": {"config": {"research-flow": {"verify": {"reasoningEffort": "low", "maxTokens": 16384}}}},
    }
    (folder / "config.json").write_text(json.dumps(shipped))
    monkeypatch.setattr(employee, "discover_product_rows", lambda: [row("Raven-PPT"), row("Raven-Research")])
    monkeypatch.setattr(employee, "product_folder", lambda name: folder)
    settings = {
        "providers": {"deepseek": {"apiKey": "ds-key"}},
        "permissions": {"mode": "full", "judgeTimeoutSeconds": 60},
        "acp": {"defaultMode": "medium"},
        "agents": {
            "defaults": {
                "model": "deepseek/deepseek-flash",
                "provider": "deepseek",
                "reasoningEffort": "low",
                "contextWindowTokens": 1048576,
            }
        },
    }
    config = tmp_path / "config.json"
    config.write_text(json.dumps(settings))
    worker = hire(
        Scenario.load(TRAVEL), config, workdir=tmp_path, root=tmp_path / "run", subagent_model="deepseek/deepseek-flash"
    )
    rows = {item.name: item for item in worker.baseline.config.subagents.agents}
    for name, secret in (("Raven-PPT", "PPT_API_KEY"), ("Raven-Research", "RESEARCH_API_KEY")):
        copy = tmp_path / "run" / "deployment" / f"{name.lower()}.json"
        written = json.loads(copy.read_text())
        assert rows[name].env[secret] == "ds-key" and rows[name].command.endswith(f"--config {copy}")
        assert written["agents"]["defaults"] == {
            "model": "deepseek/deepseek-flash",
            "provider": "deepseek",
            "maxTokens": 9,
            "reasoningEffort": "low",
            "contextWindowTokens": 1048576,
        }
        assert written["providers"]["deepseek"] == {"apiKey": "ds-key"} and copy.stat().st_mode & 0o777 == 0o600
        assert written["permissions"] == {"mode": "full", "judgeTimeoutSeconds": 60.0}
        # The shipped high tier would run every deck call at high effort over the pinned low one.
        assert written["acp"]["defaultMode"] == "medium"
        # A product's shell reaches past the run's bounds, so products run without one.
        assert written["tools"]["disabledTools"] == ["spawn", "ask_user", "exec"]
        # A thinking DeepSeek model refuses the verify gate's forced verdict call, and the gate then passes every draft.
        verify = written["plugins"]["config"]["research-flow"]["verify"]
        assert verify == {"reasoningEffort": "none", "maxTokens": 16384}
    assert "PPT_ACP_HOME" in rows["Raven-PPT"].env
    settings["agents"]["defaults"].pop("reasoningEffort")
    config.write_text(json.dumps(settings))
    with pytest.raises(ValueError, match="reasoning effort"):
        hire(Scenario.load(TRAVEL), config, workdir=tmp_path, root=tmp_path / "again", subagent_model="deepseek/x")


def test_a_component_root_above_the_declared_source_paths_does_not_widen_the_snapshot(tmp_path):
    repo = tmp_path / "repo"
    for relative in ("pkg/__init__.py", "pkg/curator/__init__.py", "pkg/curator/a.py", "pkg/hidden/secret.md"):
        (repo / relative).parent.mkdir(parents=True, exist_ok=True)
        (repo / relative).write_text("x")
    skill = tmp_path / "home" / "skills" / "sop"
    skill.mkdir(parents=True)
    (skill / "SKILL.md").write_text("x")
    digest = hashlib.sha256(b"x").hexdigest()
    inspection = SimpleNamespace(
        facts={},
        sources={
            "component.a": {
                "path": str(repo / "pkg" / "curator" / "a.py"),
                "root": str(repo / "pkg"),
                "digest": digest,
            },
            "skill.workspace/sop": {"path": str(skill / "SKILL.md"), "root": str(skill), "digest": digest},
        },
    )
    exploration = Exploration(Config(), inspection, repository=repo, source_paths=("pkg/curator",), root=tmp_path / "x")
    assert set(exploration.mounts) == {
        "source/pkg/curator",
        f"materials/{list(exploration.mounts)[-1].split('/')[1]}/sop",
    }
    assert repo / "pkg" not in exploration.mounts.values()


async def test_the_agency_reads_the_last_version_of_each_delivered_file(tmp_path, monkeypatch):
    from experimental.curator.raven_adapter.worker import Execution
    from experimental.simulation.traveller import opened, transcript

    if shutil.which("soffice") is None:
        from pptx import Presentation

        from experimental.simulation import files
        from experimental.simulation.render import cached

        def draw(file, out, width):
            for number in range(1, len(Presentation(file).slides) + 1):
                (out / f"page-{number:02d}.png").write_bytes(b"png")

        monkeypatch.setattr(files, "cached", lambda file, cache, width: cached(file, cache, width, draw))
    scenario = Scenario.load(TRAVEL)
    first, second = tmp_path / "t1" / "trip.html", tmp_path / "t2" / "trip.html"
    final = '<style>h1{color:red}</style><h1>final</h1><img src="data:image/png;base64,' + "A" * 300 + '">'
    for path, text in ((first, "<h1>draft</h1>"), (second, final)):
        path.parent.mkdir()
        path.write_text(text)
    deck = tmp_path / "t3" / "plan.pptx"
    deck.parent.mkdir()
    shutil.copy(scenario.materials["plan-deck-template"] / "template.pptx", deck)
    records = [{"kind": "runner.event", "event_type": "Text", "event": {"content": "Here it is."}}]
    exchanges = [
        Exchange("Plan it", Execution("t1", [], records, {}, "a", (str(first),))),
        Exchange("Fix it", Execution("t2", [], records, {}, "a", (str(second),))),
        Exchange("The deck?", Execution("t3", [], records, {}, "a", (str(deck),))),
    ]
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Good page."}))
    await Agency(scenario, provider, tmp_path / "skills", workdir=tmp_path, deliver="pool").evaluate(
        {"student": exchanges}
    )
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    delivered = packet["deliverables"]["student"]
    assert packet["references"]["student"]["deck"]["template_placeholders_left"]
    assert delivered["trip.html"] == "final"
    assert (
        delivered["plan.pptx"].startswith("19 slides, aspect 1.78") and "[style] background" in delivered["plan.pptx"]
    )
    assert packet["deck_pages"] == {"student": {"plan.pptx": 19}}
    pictures = provider.requests[0]["messages"][2]["content"]
    images = [part for part in pictures if part["type"] == "image_url"]
    assert len(images) == 19 and images[0]["image_url"]["url"].startswith("data:image/png;base64,")
    assert (deck.parent / "plan.pptx.thumbs" / "page-01.png").is_file()
    assert transcript(exchanges)[0]["delivered"] == ["trip.html"]
    seen = opened(exchanges)
    assert seen["trip.html"] == "final" and "--- slide 1\n" in seen["plan.pptx"] and "[style]" not in seen["plan.pptx"]


def test_the_agency_reads_what_the_research_colleague_reported_in_each_playbook_run(tmp_path):
    from experimental.curator.raven_adapter.worker import Execution
    from experimental.simulation.agency import research

    found, drafted = tmp_path / "research.out.md", tmp_path / "brief.out.md"
    found.write_text("G7311 Shanghai to Huangshan, 3h (source: 12306)")
    drafted.write_text("brief")
    manifest = {
        "files": [
            {"node": "p-research", "subagent": "Raven-Research", "output_file": str(found)},
            {"node": "p-brief", "subagent": "Raven", "output_file": str(drafted)},
        ]
    }
    records = [{"kind": "dag.progress", "name": "dag_run_completed", "payload": {"manifest": manifest}}]
    exchanges = [Exchange("Deck please", Execution("t1", [], records, {}, "a"))]
    assert research(exchanges) == {"p-research": "G7311 Shanghai to Huangshan, 3h (source: 12306)"}


async def test_fresh_trial_preserves_content_owned_by_the_current_child_strategy(tmp_path):
    from experimental.curator.harness import Artifact

    home = tmp_path / "subagents/Hosted"
    path = home / "TOOLS.md"
    path.parent.mkdir(parents=True)
    path.write_text("runtime edit")

    class Drill:
        async def run(self, worker):
            assert path.read_text() == "curated SOP"
            return {}

    worker = SimpleNamespace(
        baseline=SimpleNamespace(config=SimpleNamespace(workspace_path=tmp_path)),
        artifact=Artifact(values={}),
        children={
            "Hosted": SimpleNamespace(
                baseline=SimpleNamespace(config=SimpleNamespace(workspace_path=home)),
                artifact=Artifact(values={"memory.prompt": {"TOOLS.md": "curated SOP"}}),
            )
        },
    )
    await Fresh(Drill(), {"Hosted/TOOLS.md": b"original"}).run(worker)


async def test_the_owner_gets_the_played_card_and_its_figures_and_still_decides_every_verdict(tmp_path):
    from experimental.simulation.cards import Drawn, Trip

    scenario = Scenario.load(TRAVEL)
    sheet = scenario.text("service-sop").split("```")[1].strip()
    guest = Trip("guest", "origin", "somewhere", (12, 2), (12, 6), 4, 4, (), (), 20_000, "139 0571 6628")
    card = Drawn("student", "card text", guest)
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine."}))
    agency = Agency(
        scenario,
        provider,
        tmp_path / "skills",
        workdir=tmp_path,
        plan="all",
        deliver="pool",
        cards=lambda: {"student": card},
        records=tmp_path / "records",
    )
    agency.prepare()
    signal = await agency.evaluate({"student": [exchange("Quote please", sheet)]})
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["cards"] == {"student": "card text"}
    comfort = list(packet["references"]["student"]["prices"]["by_product"].values())[1]
    assert comfort["party_total"] == 12_800 and "deck" not in packet["references"]["student"]
    assert signal.satisfied is True and len(provider.requests) == 1
    rows = [json.loads(line) for line in (tmp_path / "records" / "references.jsonl").read_text().splitlines()]
    assert rows[0]["drill"] == "student" and rows[0]["trip"]["adults"] == 4 and rows[0]["card"] == "card text"
    assert rows[0]["references"]["prices"]["persons"] == 4


def test_a_scenario_without_a_material_loses_it_from_its_materials_opening_set_and_plans():
    scenario = Scenario.load(TRAVEL)
    control = scenario.without(["plan-deck-sample"])
    assert "plan-deck-sample" in scenario.materials and "plan-deck-sample" not in control.materials
    assert "plan-deck-sample" not in control.initial and len(control.initial) == len(scenario.initial) - 1
    assert all("plan-deck-sample" not in step for steps in control.plans.values() for step in steps)
    assert partition(control.plans["by-stage"], control.materials)
    with pytest.raises(ValueError, match="brochure"):
        scenario.without(["brochure"])


def test_the_employees_workdir_is_new_and_never_the_repository(tmp_path):
    from experimental.simulation.__main__ import workplace

    repository = Path(__file__).resolve().parents[1]
    for bad in (repository, repository / "experimental", repository.parent):
        with pytest.raises(ValueError, match="repository"):
            workplace(bad)
    (tmp_path / "used").mkdir()
    (tmp_path / "used" / "notes.md").write_text("x")
    with pytest.raises(ValueError, match="new or empty"):
        workplace(tmp_path / "used")
    assert workplace(tmp_path / "fresh").is_dir() and not any(workplace(None).iterdir())
