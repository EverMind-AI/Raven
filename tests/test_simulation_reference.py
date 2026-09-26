"""Figures and deck facts computed for the owner from the scenario's own materials and the played card."""

import shutil

from pptx import Presentation

from experimental.curator.raven_adapter.worker import Execution
from experimental.iteration.protocols import Exchange
from experimental.simulation.cards import Trip
from experimental.simulation.reference import Rules, deck, prices, quoted, references
from experimental.simulation.scenario import BUNDLED, Scenario

TRAVEL = BUNDLED / "travel_agency"


def said(user, reply, files=()):
    rows = [{"kind": "runner.event", "event_type": "Text", "event": {"content": reply}}]
    return Exchange(user, Execution("t1", [], rows, {}, "a", tuple(str(path) for path in files)))


def trip(start, nights, adults, children=(), budget=99_000):
    end = (start[0], start[1] + nights)
    return Trip(
        "guest", "origin", "somewhere", start, end, nights, adults, tuple(children), (), budget, "139 0571 6628"
    )


def loaded():
    scenario = Scenario.load(TRAVEL)
    return scenario, Rules.load(scenario)


def test_rules_read_prices_seasons_policy_deck_name_and_template_from_the_materials():
    _, rules = loaded()
    value, comfort, premium = rules.products.values()
    assert [product.base for product in (value, comfort, premium)] == [3, 4, 5]
    assert [product.extra for product in (value, comfort, premium)] == [260, 800, 2600]
    assert [value.child, comfort.child, premium.child] == [True, True, False]
    assert comfort.bands == ((1, 1), (2, 3), (4, 8)) and premium.bands == ((1, 1), (2, 2), (3, 6))
    assert [prices[1] for prices in comfort.prices.values()] == [3500, 3900, 4600]
    off, shoulder, peak = (rules.season(day) for day in ((12, 1), (10, 8), (10, 1)))
    assert len({off, shoulder, peak}) == 3
    assert rules.season((2, 28)) == off and rules.season((11, 15)) == shoulder and rules.season((10, 7)) == peak
    assert (rules.insurance, rules.child_age, rules.child_ratio) == (60, 12, 70)
    assert rules.fonts == {"Noto Sans CJK SC", "Noto Serif CJK SC"} and rules.cover != rules.contact
    assert rules.quote.fullmatch("HL-Q-1027-4P") and not rules.quote.fullmatch("HL-Q-127-4P")


def test_the_card_gets_every_products_price_as_the_price_list_and_policy_give_it():
    scenario, rules = loaded()
    sop = trip((12, 2), 4, 3)
    comfort = list(rules.products)[1]
    assert prices(rules, sop)["by_product"][comfort]["party_total"] == 10_500
    assert "10,500" in scenario.text("service-sop")
    family = prices(rules, trip((10, 22), 4, 2, (9, 14), budget=7000))
    value_row, _, premium_row = family["by_product"].values()
    assert (value_row["per_person"], value_row["child_price"], value_row["extra_nights_total"]) == (800, 560, 1040)
    assert value_row["children_at_child_price"] == 1 and value_row["party_total"] == 4000
    assert premium_row["nights_short_of_base"] == 1 and "child_price" not in premium_row
    assert (family["budget_per_person_per_night"], family["insurance_total"]) == (438, 240)
    assert family["quote_number"] == "HL-Q-1022-4P"
    crowd = prices(rules, trip((10, 22), 4, 9))["by_product"]
    assert all(isinstance(row, str) for row in crowd.values())


def test_the_agencys_sample_deck_reads_clean_and_the_bare_template_shows_its_placeholders(tmp_path):
    scenario, rules = loaded()
    number = "HL-Q-0512-2P"
    sample = tmp_path / f"{rules.deck_name}-{number}.pptx"
    shutil.copy(scenario.materials["plan-deck-sample"] / "sample.pptx", sample)
    exchanges = [said("Quote", f"Quote {number}"), said("Deck?", "Here it is.", [sample])]
    facts = deck(rules, exchanges)
    assert facts["file"] in facts["names_per_spec"] and facts["opens"] and facts["aspect"] == 1.778
    assert facts["fonts_outside_template"] == [] and facts["slides_with_motion"] == 0
    assert (facts["first_slide_layout"], facts["last_slide_layout"]) == (rules.cover, rules.contact)
    assert facts["template_placeholders_left"] == [] and facts["quote_numbers_in_deck"] == [number]
    template = tmp_path / "t" / sample.name
    template.parent.mkdir()
    shutil.copy(scenario.materials["plan-deck-template"] / "template.pptx", template)
    exchanges[-1] = said("Deck?", "Here it is.", [template])
    assert deck(rules, exchanges)["template_placeholders_left"] == list(rules.words["placeholders"])
    blank = Presentation()
    blank.slides.add_slide(blank.slide_layouts[0]).shapes.title.text = number
    fresh = tmp_path / "fresh.pptx"
    blank.save(fresh)
    exchanges[-1] = said("Deck?", "Here it is.", [fresh])
    facts = deck(rules, exchanges)
    assert facts["aspect"] == 1.333 and facts["fonts_outside_template"] and facts["first_slide_layout"] != rules.cover
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a deck")
    exchanges[-1] = said("Deck?", "Here it is.", [broken])
    assert deck(rules, exchanges)["opens"] is False


def test_references_hold_prices_only_for_a_card_with_a_trip_and_deck_facts_only_after_a_deck():
    _, rules = loaded()
    exchanges = [said("Hi", "Quote HL-Q-0925-4P, then HL-Q-0925-4P again and HL-Q-1022-4P.")]
    assert quoted(exchanges, rules.quote) == ["HL-Q-0925-4P", "HL-Q-1022-4P"]
    assert set(references(rules, exchanges, trip((10, 22), 4, 4))) == {"prices"}
    assert references(rules, exchanges, None) == {}
