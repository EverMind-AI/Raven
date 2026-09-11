"""The script route: running a program a model wrote, and surviving it.

Most of these are regressions. Each names the run it came from, because the
reason a guard exists is the only thing that stops it being simplified away.
"""

from __future__ import annotations

import textwrap
from pathlib import Path

import pytest

from raven_ppt.backends.script import (
    HelperSources,
    broken_page,
    carries_a_program,
    page_blocks,
    page_sources,
    provision,
    run_script,
    script_path,
    submission_refusal,
)
from raven_ppt.contracts import Project
from tests._ppt_engine_fixtures import (  # noqa: F401
    COMMENT_ABOVE_FIRST_BANNER,
    deck,
    image,
    noise_image,
    noise_png,
    product_page,
    template_file,
)

pytest.importorskip("pptx")

DECK = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    def title(slide, text):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        box.text_frame.text = text


    # SLIDE 1
    one = new_slide()
    title(one, "Unified video segmentation")

    # SLIDE 2
    two = new_slide()
    title(two, "Target queries")

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.fixture()
def project(tmp_path: Path) -> Project:
    return Project(workspace=tmp_path, slug="tarvis")


async def _build(project: Project, script: str | None = None):
    return await run_script(project, script, timeout_s=120.0)


@pytest.mark.asyncio
async def test_a_program_builds_a_deck_and_says_which_lines_drew_each_page(project: Project) -> None:
    outcome = await _build(project, DECK)

    assert outcome.ok, outcome.stderr
    assert outcome.pages == 2
    assert outcome.pptx_path is not None and outcome.pptx_path.is_file()
    # Derived from execution, not from the banners: the pages are told apart by
    # the line that created each slide.
    assert [s.page for s in outcome.sources] == [1, 2]
    first, second = outcome.sources
    assert first.last_line == second.first_line, "page spans must abut, not overlap"
    assert outcome.source_digest


@pytest.mark.asyncio
async def test_the_helpers_are_waiting_beside_the_script(project: Project) -> None:
    provision(project, HelperSources(theme="THEMES = {}\n", icons="def add_icon(*a, **k):\n    pass\n"))
    assert (project.build_dir / "ppt_theme.py").is_file()
    assert (project.build_dir / "ppt_icons.py").is_file()


def test_the_real_assets_land_as_modules_the_author_can_import(project: Project, tmp_path: Path) -> None:
    """The names are the asset service's, not the backend's.

    `from ppt_theme import THEMES, rgb` is what the author is told to write, so
    the filenames come from the service that owns the contents. This runs the
    generated modules in a subprocess because they read their data file relative
    to their own location, which is the whole point of writing them to disk -- and
    because `ppt_shapes` imports two of its neighbours, which only works if the
    build directory really is one importable place.
    """
    import subprocess
    import sys

    from raven_ppt.backends.script import asset_helpers

    build = provision(project, asset_helpers())
    assert {p.name for p in build.iterdir()} == {
        "ppt_theme.py",
        "ppt_icons.py",
        "ppt_layout.py",
        "ppt_charts.py",
        "ppt_shapes.py",
        "themes.json",
        "icons.json",
        "icon_keywords.json",
        "shapes.json",
        # The skill's reference documents, which reach the author no other way:
        # the pool reads SKILL.md and nothing beside it, and the fence then
        # refuses the path its `references/...` links resolve to.
        "references",
    }
    assert {p.name for p in (build / "references").iterdir()} == {
        "charts.md",
        # What a finished page owes, which `ppt_review` judges the render by. It lands
        # here rather than only reaching the reviewer because a requirement the author
        # never saw is a requirement nobody agreed to.
        "design-requirements.md",
        "formulas.md",
        "icons.md",
        "layouts.md",
        # The layout registry's passages, one file per family, so a page loads the
        # worked code it needs and not the other nine hundred lines.
        "layouts-data.md",
        "layouts-figures.md",
        "layouts-multiples.md",
        "layouts-primitives.md",
        "layouts-type.md",
        "shapes.md",
        "tables.md",
    }
    probe = build / "probe.py"
    probe.write_text(
        "from ppt_theme import THEMES, rgb\n"
        "from ppt_icons import ICON_NAMES, add_icon\n"
        "from ppt_shapes import PRESET_NAMES, timeline\n"
        "print(len(THEMES), len(ICON_NAMES), len(PRESET_NAMES), rgb('#FFFFFF') is not None)\n",
        encoding="utf-8",
    )
    done = subprocess.run([sys.executable, "probe.py"], cwd=build, capture_output=True, text=True, timeout=60)
    assert done.returncode == 0, done.stderr
    themes, icons, presets, colour = done.stdout.split()
    assert int(themes) == 10 and int(icons) > 100 and int(presets) == 109 and colour == "True"


@pytest.mark.asyncio
async def test_a_helper_the_author_edited_is_restored_on_the_next_build(project: Project) -> None:
    """A quietly modified helper fails in a way that reads as a deck bug."""
    provision(project, HelperSources(theme="THEMES = {'a': 1}\n"))
    (project.build_dir / "ppt_theme.py").write_text("broken\n", encoding="utf-8")
    provision(project, HelperSources(theme="THEMES = {'a': 1}\n"))
    assert (project.build_dir / "ppt_theme.py").read_text() == "THEMES = {'a': 1}\n"


@pytest.mark.asyncio
async def test_a_restored_helper_is_named_in_the_build_note(project: Project) -> None:
    """The restore above is right, and it was also silent.

    A run wrote its own themes.json so a theme name of its own would resolve,
    the next build put the engine's file back, and the run then read that file
    127 times over an hour trying to find where its edit had gone -- every read
    succeeding, so nothing counted it as failure. The restore stands; what
    changes is that the build says it happened.
    """
    helpers = HelperSources(theme="THEMES = {'a': 1}\n")
    provision(project, helpers)
    (project.build_dir / "ppt_theme.py").write_text("THEMES = {'mine': 2}\n", encoding="utf-8")

    outcome = await run_script(project, DECK, helpers=helpers, timeout_s=120.0)
    assert outcome.ok, outcome.stderr
    assert "ppt_theme.py" in (outcome.note or ""), "the build names the helper it put back"
    assert (project.build_dir / "ppt_theme.py").read_text() == "THEMES = {'a': 1}\n"

    again = await run_script(project, DECK, helpers=helpers, timeout_s=120.0)
    assert "ppt_theme.py" not in (again.note or ""), "and says nothing when nothing was edited"


@pytest.mark.asyncio
@pytest.mark.parametrize("submission", [" ", "\n", "# use existing build.py"])
async def test_a_submission_with_no_program_runs_what_is_there(project: Project, submission: str) -> None:
    """Both halves matter, and each cost a run.

    The file must survive: `script="."` once replaced a working forty-kilobyte
    program and every later call failed on the wreckage. And the build must still
    happen: refusing all three of these in one run is what drove an author out of
    the tool and into exec, delivering a deck no gate had seen.
    """
    await _build(project, DECK)
    outcome = await _build(project, submission)

    assert outcome.ok, outcome.stderr
    assert outcome.pages == 2
    assert script_path(project).read_text(encoding="utf-8") == DECK
    assert "left as it was" in outcome.note


@pytest.mark.asyncio
async def test_text_that_is_not_python_is_refused_by_name(project: Project) -> None:
    await _build(project, DECK)
    outcome = await _build(project, ".")

    assert not outcome.ok
    assert "not valid python" in outcome.stderr
    assert script_path(project).read_text(encoding="utf-8") == DECK


@pytest.mark.asyncio
async def test_a_shim_that_runs_another_file_is_refused_with_a_working_skeleton(project: Project) -> None:
    outcome = await _build(project, "import runpy\nrunpy.run_path('real_build.py')\n")

    assert not outcome.ok
    assert "hands the work to another file" in outcome.stderr
    # Refusing without an alternative is what pushed one author out of the tool.
    assert "# SLIDE 1" in outcome.stderr


@pytest.mark.asyncio
async def test_a_fragment_much_smaller_than_the_program_is_refused(project: Project) -> None:
    await _build(project, DECK)
    outcome = await _build(project, "prs.save(os.environ['PPT_OUTPUT'])\n")

    assert not outcome.ok
    assert "fragment sent by mistake" in outcome.stderr
    assert script_path(project).read_text(encoding="utf-8") == DECK


@pytest.mark.asyncio
async def test_a_build_that_dies_leaves_the_last_deck_alone(project: Project) -> None:
    """That deck is what a failed edit gets repaired against."""
    good = await _build(project, DECK)
    assert good.ok and good.pptx_path is not None
    before = good.pptx_path.read_bytes()

    # In the prelude, before a deck exists: a raise inside a page's block is that page's
    # failure now and the deck is rebuilt around it (see the isolation tests below).
    broken = await _build(project, DECK.replace("prs = Presentation()", "raise RuntimeError('boom')"))

    assert not broken.ok
    assert "boom" in broken.stderr
    assert good.pptx_path.read_bytes() == before
    failure = project.review_dir / "build_failures" / "failure-001"
    assert (failure / "build.py").is_file()
    assert "boom" in (failure / "stderr.txt").read_text(encoding="utf-8")
    assert (failure / "failure.json").is_file()


@pytest.mark.asyncio
async def test_a_failed_build_discards_the_line_record_it_can_no_longer_describe(project: Project) -> None:
    await _build(project, DECK)
    await _build(project, DECK.replace("prs.save", "raise RuntimeError('boom')  # prs.save"))
    assert not (project.build_dir / ".slide_lines.json").is_file()


@pytest.mark.asyncio
async def test_a_script_that_writes_nothing_is_told_where_to_save(project: Project) -> None:
    outcome = await _build(project, "from pptx import Presentation\nprs = Presentation()\n")

    assert not outcome.ok
    assert "PPT_OUTPUT" in outcome.stderr
    failure = project.review_dir / "build_failures" / "failure-001"
    assert (failure / "build.py").is_file()
    assert "PPT_OUTPUT" in (failure / "stderr.txt").read_text(encoding="utf-8")


@pytest.mark.asyncio
async def test_with_no_script_yet_the_error_is_the_instruction(project: Project) -> None:
    """And the path in it is one the file tools can take.

    Absolute is unambiguous and unusable: `write_file` resolves against the
    workspace, so an author handed an absolute path writes it verbatim and the file
    lands outside the project. A live run did exactly that -- it wrote
    `build/build.py` from the description, got this message, and spent four calls
    recovering.
    """
    outcome = await _build(project)

    assert not outcome.ok
    assert str(script_path(project).relative_to(project.workspace)) in outcome.stderr
    assert str(project.workspace) not in outcome.stderr, "an absolute path is one write_file cannot take"
    assert "write_file" in outcome.stderr


@pytest.mark.asyncio
async def test_pages_drawn_by_a_loop_cannot_be_told_apart(project: Project) -> None:
    """One call site for every page, so there is no per-page block to hand over.

    Reported as an absence rather than guessed at: pairing a render with the
    wrong code is worse than saying the pages cannot be separated.
    """
    looped = textwrap.dedent(
        """
        import os
        from pptx import Presentation
        prs = Presentation()
        for text in ("a", "b", "c"):
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(os.environ["PPT_OUTPUT"])
        """
    ).lstrip()
    outcome = await _build(project, looped)

    assert outcome.ok, outcome.stderr
    assert outcome.pages == 3
    assert outcome.sources == ()


def test_carries_a_program_separates_empty_from_wrong() -> None:
    assert carries_a_program("x = 1") is True
    assert carries_a_program("") is False
    assert carries_a_program("   \n\n") is False
    assert carries_a_program("# only a comment") is False
    # Not empty, only wrong -- so it goes on to be refused by name.
    assert carries_a_program(".") is True


def test_a_page_keeps_its_banner_when_the_block_is_taken_from_execution() -> None:
    lines = DECK.splitlines(keepends=True)
    created = [next(i for i, line in enumerate(lines, start=1) if "one = new_slide()" in line)]
    created.append(next(i for i, line in enumerate(lines, start=1) if "two = new_slide()" in line))
    sources = page_sources(lines, created)

    assert "# SLIDE 1" in "".join(lines[sources[0].first_line : sources[0].last_line])
    assert "# SLIDE 2" in "".join(lines[sources[1].first_line : sources[1].last_line])


def test_the_last_page_stops_at_the_save_call() -> None:
    lines = DECK.splitlines(keepends=True)
    blocks = page_blocks(lines)
    _, end = blocks[2]
    assert "prs.save" not in "".join(lines[blocks[2][0] : end])


def test_banner_blocks_do_not_overlap() -> None:
    lines = DECK.splitlines(keepends=True)
    blocks = page_blocks(lines)
    assert blocks[1][1] <= blocks[2][0]


def test_a_crash_inside_a_shared_helper_is_attributed_to_the_calling_page() -> None:
    """The frame that names the page is the call site, not the deepest frame."""
    stderr = (
        '  File "/w/build/_run_build.py", line 30, in <module>\n'
        '  File "/w/build/build.py", line 15, in title\n'
        '  File "/w/build/build.py", line 25, in <module>\n'
        "RuntimeError: boom\n"
    )
    lines = DECK.splitlines(keepends=True)
    marker = next(i for i, line in enumerate(lines) if "# SLIDE 2" in line)
    stderr = stderr.replace("line 25", f"line {marker + 3}")
    assert broken_page(DECK, stderr) == 2


def test_a_crash_in_the_shared_prelude_is_not_blamed_on_a_page() -> None:
    stderr = '  File "/w/build/build.py", line 2, in <module>\nImportError: no pptx\n'
    assert broken_page(DECK, stderr) is None


def test_the_runners_own_frames_are_never_mistaken_for_the_script() -> None:
    stderr = '  File "/w/build/_run_build.py", line 40, in <module>\nRuntimeError: boom\n'
    assert broken_page(DECK, stderr) is None


def test_refusal_returns_none_for_a_real_program(tmp_path: Path) -> None:
    assert submission_refusal(DECK, tmp_path / "build.py") is None


# --- building inside a template -------------------------------------------

TEMPLATE_DECK = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches
    from ppt_template import clone_page, drop_shape, replace_text

    prs = Presentation(os.environ["PPT_TEMPLATE"])
    original = Presentation(os.environ["PPT_TEMPLATE_SOURCE"])

    # SLIDE 1
    one = prs.slides.add_slide(prs.slide_layouts[6])
    box = one.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
    box.text_frame.text = "Written against the template"

    # SLIDE 2
    two = clone_page(prs, original.slides[0])
    heading = next(s for s in two.shapes if getattr(s, "has_text_frame", False))
    replace_text(heading, "Cloned out of the template")
    drop_shape(two.shapes[-1])

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


async def test_a_bound_template_reaches_the_program(tmp_path: Path, template_file):
    """The whole wiring, end to end: the deck is built in the prepared copy, a page
    is cloned out of the original, and the operations python-pptx does not have
    arrive as a module the program imports."""
    from pptx import Presentation

    from raven_ppt.services.template import bind

    project = Project(workspace=tmp_path, slug="talk")
    assert bind(template_file(), project) is not None

    outcome = await run_script(project, TEMPLATE_DECK)

    assert outcome.ok, outcome.stderr
    assert outcome.pages == 2
    assert (project.build_dir / "ppt_template.py").is_file()
    built = Presentation(str(outcome.pptx_path))
    assert built.slide_width == Presentation(str(template_file())).slide_width
    assert "Cloned out of the template" in "\n".join(
        shape.text_frame.text for shape in built.slides[1].shapes if getattr(shape, "has_text_frame", False)
    )


async def test_a_deck_without_a_template_gets_neither_the_paths_nor_the_module(tmp_path: Path):
    """A build directory holding an importable `ppt_template` for a deck with no
    template is an invitation to import it, and the failure that follows is about
    a file the author did not write."""
    project = Project(workspace=tmp_path, slug="talk")

    outcome = await run_script(
        project,
        textwrap.dedent(
            """
            import os
            from pptx import Presentation

            prs = Presentation()
            # SLIDE 1
            prs.slides.add_slide(prs.slide_layouts[6])
            print("template:", os.environ.get("PPT_TEMPLATE", "none"))
            prs.save(os.environ["PPT_OUTPUT"])
            """
        ).lstrip(),
    )

    assert outcome.ok, outcome.stderr
    assert "template: none" in outcome.stdout
    assert not (project.build_dir / "ppt_template.py").exists()


# A program of the shape the two defects came in: pages started from a layout that
# carries placeholders, and a table drawn with the layout helper.
TIDY_DECK = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches
    from ppt_layout import Box, table, write
    from ppt_theme import THEMES

    TH = THEMES["ink-graphite"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # SLIDE 1
    one = prs.slides.add_slide(prs.slide_layouts[1])
    one.placeholders[0].text_frame.text = "实验：单一共享模型全面超过专用模型"
    table(one, Box(0.7, 1.6, 8.0, 4.4), [["方法", "YTVIS AP"], ["TarViS", "48.3"], ["VITA", "45.7"]], TH)

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.mark.asyncio
async def test_a_built_deck_carries_neither_defect_a_render_cannot_show(project: Project) -> None:
    """The guard, rather than the fix: `tidy` runs on the way out of every build.

    Both defects it corrects were in every deck this route delivered and neither is
    visible in a render, so a change that moved the deck past `tidy` -- another build
    path, a reordered runner -- would be invisible in exactly the same way. This test
    is the thing that would not be.
    """
    from pptx import Presentation

    from raven_ppt.backends.script.workspace import asset_helpers
    from raven_ppt.services.tidy import EDGES, GALLERY_STYLE, tidy

    outcome = await run_script(project, TIDY_DECK, helpers=asset_helpers(), timeout_s=120.0)
    assert outcome.ok, outcome.stderr

    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    slide = Presentation(str(outcome.pptx_path)).slides[0]
    empty = [
        shape.name
        for shape in slide.shapes
        if shape.is_placeholder and not (getattr(shape, "text_frame", None) and shape.text_frame.text.strip())
    ]
    assert empty == [], f"'Click to add' placeholders reached the delivered deck: {empty}"

    graphic = next(shape for shape in slide.shapes if getattr(shape, "has_table", False))
    styles = [(e.text or "").upper() for e in graphic.table._tbl.iter(f"{namespace}tableStyleId")]
    assert GALLERY_STYLE not in styles, "the Office gallery style reached the delivered deck"
    properties = graphic.table.cell(0, 0)._tc.find(f"{namespace}tcPr")
    order = [element.tag.replace(namespace, "") for element in properties]
    assert order[: len(EDGES)] == list(EDGES), f"cell borders out of the schema's order: {order}"

    # And it is idempotent, which is what makes running it on every build safe.
    assert tidy(outcome.pptx_path) == ()


FAILING_MIDDLE = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    def title(slide, text):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        box.text_frame.text = text


    # SLIDE 1
    one = new_slide()
    title(one, "First")

    # SLIDE 2
    two = new_slide()
    title(two, "Second")
    rows = [("a", 1), ("b", 2, 3)]
    heights = [len(label) for label, count in rows]

    # SLIDE 3
    three = new_slide()
    title(three, "Third")

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.mark.asyncio
async def test_a_page_that_raises_loses_only_itself(project: Project) -> None:
    """Measured across four runs, one build in five died in one page's block and the
    author got that traceback and nothing else. The block is run on its own: its half
    page is taken back, a page saying what went wrong stands in its place, and the
    pages after it are still drawn and still mapped to their code."""
    from pptx import Presentation

    from raven_ppt.backends.script import page_failures

    outcome = await _build(project, FAILING_MIDDLE)

    assert outcome.ok and outcome.pages == 3
    assert [source.page for source in outcome.sources] == [1, 2, 3], "the placeholder is page 2 and maps to block 2"
    deck = Presentation(str(outcome.pptx_path))
    texts = [" ".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
    assert "First" in texts[0] and "Third" in texts[2]
    assert "Page 2 did not draw" in texts[1] and "too many values to unpack" in texts[1]
    assert "Second" not in texts[1], "the half-drawn page was taken back"
    failed = page_failures(project)
    assert [entry["page"] for entry in failed] == [2]
    assert "ValueError: too many values to unpack" in failed[0]["error"]
    assert 'build.py", line' in failed[0]["traceback"]
    assert "ValueError" in outcome.stderr
    kept = project.review_dir / "build_failures" / "failure-001"
    assert "ValueError" in (kept / "stderr.txt").read_text(encoding="utf-8"), "the failure is on record like any other"


HELPER_EXITS = textwrap.dedent(
    """
    import os

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def sub1(slide, prefix, new):
        '''the shape of helper an author writes: a loud stop when nothing matches'''
        found = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith(prefix)]
        if not found:
            raise SystemExit(f"sub1: nothing says {prefix!r}")
        found[0].text_frame.text = new


    def new_slide(text):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        box.text_frame.text = text
        return slide


    # SLIDE 1
    new_slide("First")

    # SLIDE 2
    two = new_slide("Second")
    sub1(two, "00%", "22.0%")

    # SLIDE 3
    new_slide("Third")

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.mark.asyncio
async def test_a_helper_that_exits_loses_only_its_page(project: Project) -> None:
    """`SystemExit` is a BaseException, so `except Exception` did not hold it, and the
    helpers authors write raise exactly that: a live run's own `sub1` ended
    `raise SystemExit(f"sub1: nothing says {prefix!r}")`, guessed a placeholder string
    the template did not have, and cost the whole build -- two pages already drawn and
    14.6 minutes of writing thrown away for one page's bad guess.

    A helper saying "this cannot work" is that page's failure, which is what the
    isolation branch is for."""
    from pptx import Presentation

    from raven_ppt.backends.script import page_failures

    outcome = await _build(project, HELPER_EXITS)

    assert outcome.ok and outcome.pages == 3
    deck = Presentation(str(outcome.pptx_path))
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
    assert "First" in texts[0] and "Third" in texts[2]
    assert "Page 2 did not draw" in texts[1] and "nothing says '00%'" in texts[1]
    assert [entry["page"] for entry in page_failures(project)] == [2]


@pytest.mark.asyncio
async def test_a_script_saved_with_a_byte_order_mark_builds_like_one_without(project: Project) -> None:
    """Editors on one platform put U+FEFF in front of the file the author edits in place
    (a submitted string is checked by compile() first and refused by name). `python
    build.py` skips the mark; handed to compile() as text it is a SyntaxError on line 1,
    and neither `\\s` nor `str.lstrip()` sees past it, so the first banner and the comment
    above it were lost to every reader that splits the file into pages. The mark is
    dropped once, where the file is read, so every reader sees the same text."""
    import codecs

    from raven_ppt.backends.script import page_failures, read_script
    from raven_ppt.backends.script.runner import _isolable_blocks
    from raven_ppt.backends.script.workspace import script_path

    project.build_dir.mkdir(parents=True, exist_ok=True)
    plain = project.build_dir / "plain.py"
    plain.write_text(COMMENT_ABOVE_FIRST_BANNER, encoding="utf-8")
    script_path(project).write_bytes(codecs.BOM_UTF8 + COMMENT_ABOVE_FIRST_BANNER.encode("utf-8"))

    assert read_script(project) == COMMENT_ABOVE_FIRST_BANNER
    assert _isolable_blocks(script_path(project)) == _isolable_blocks(plain) == [[1, 0, 20], [2, 20, 23]], (
        "the comment above the first banner belongs to page 1's block, mark or no mark"
    )

    script_path(project).write_bytes(codecs.BOM_UTF8 + HELPER_EXITS.encode("utf-8"))
    outcome = await _build(project, None)

    assert outcome.ok and outcome.pages == 3, outcome.stderr
    assert [entry["page"] for entry in page_failures(project)] == [2]


PRELUDE_RAISES = textwrap.dedent(
    """
    import os

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def new_slide(text):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        box.text_frame.text = text
        return slide


    # SLIDE 1
    new_slide("First")

    # SLIDE 2
    new_slide("Second")
    raise KeyboardInterrupt("the caller stopped it")

    # SLIDE 3
    new_slide("Third")

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.mark.asyncio
async def test_the_pages_already_drawn_are_saved_even_when_nothing_catches(project: Project) -> None:
    """`prs.save` is the last line of the author's file, so anything the loop does not
    catch skips it and a mostly finished deck is reported as "wrote no deck". The runner
    saves what was drawn on the way out, so the build has pages to show and measure.

    KeyboardInterrupt is used here because it is the one thing the page loop must not
    swallow -- it is the caller stopping the build -- and it is therefore the sharpest
    test that the rescue is in `finally` rather than in the catch."""
    from pptx import Presentation

    outcome = await _build(project, PRELUDE_RAISES)

    assert not outcome.ok, "an interrupted build is not a deck"
    kept = project.review_dir / "build_failures" / "failure-001" / "deck.pptx.building"
    assert kept.is_file(), "nothing was saved, so the writing is gone with the round"
    texts = [
        " ".join(s.text_frame.text for s in Presentation(str(kept)).slides[n].shapes if s.has_text_frame)
        for n in range(len(Presentation(str(kept)).slides))
    ]
    assert "First" in texts[0] and "Second" in texts[1], "the two pages that were drawn are on disk"
    assert len(texts) == 2, "and the page after the interruption is not"
    assert "deck.pptx.building" in outcome.stderr, "and the author is told where to look"


NO_BANNERS = textwrap.dedent(
    """
    import os

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def grab(shapes, pred):
        found = [s for s in shapes if pred(s)]
        if not found:
            raise RuntimeError("grab: nothing matched")
        return found[0]


    def page(text):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        box.text_frame.text = text
        return slide


    page("First")
    page("Second")
    grab(page("Third").shapes)

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.mark.asyncio
async def test_a_script_whose_pages_cannot_be_separated_still_keeps_what_it_drew(project: Project) -> None:
    """Isolation needs `# SLIDE n` banners numbering the pages in file order, and a
    rewrite can drop them without saying so: three builds of one live run ended with
    "the build script did not produce a deck" over five pages that were drawn, each for
    one helper call that raised. Nothing there told the author that the banners were
    what they had lost.

    So the failure names it, and the pages are kept either way."""
    from pptx import Presentation

    outcome = await _build(project, NO_BANNERS)

    assert not outcome.ok
    assert "pages are not separable" in outcome.stderr, outcome.stderr
    kept = project.review_dir / "build_failures" / "failure-001" / "deck.pptx.building"
    assert kept.is_file(), "the pages drawn before the raise are gone"
    assert len(Presentation(str(kept)).slides) == 3
    assert "deck.pptx.building" in outcome.stderr


@pytest.mark.asyncio
async def test_a_clean_build_clears_the_failures_of_the_last_one(project: Project) -> None:
    from raven_ppt.backends.script import page_failures

    await _build(project, FAILING_MIDDLE)
    assert page_failures(project)

    outcome = await _build(project, DECK)

    assert outcome.ok and outcome.pages == 2
    assert page_failures(project) == []


@pytest.mark.asyncio
async def test_a_crash_in_the_prelude_is_still_the_whole_builds(project: Project) -> None:
    """Nothing to isolate: no page has been drawn and there is no deck to stand a page in."""
    outcome = await _build(project, DECK.replace("prs = Presentation()", "raise RuntimeError('no prs yet')"))

    assert not outcome.ok
    assert "no prs yet" in outcome.stderr


SAVES_HALFWAY = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


    def new_slide(text):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        box.text_frame.text = text
        return slide


    # SLIDE 1
    new_slide("First")
    prs.save(os.environ["PPT_OUTPUT"])

    # SLIDE 2
    new_slide("Second")

    # SLIDE 3
    new_slide("Third")
    raise KeyboardInterrupt("the caller stopped it")

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.mark.asyncio
async def test_a_script_that_saved_once_already_still_keeps_every_page_it_drew(project: Project) -> None:
    """The rescue used to stand down whenever a file was already at PPT_OUTPUT, and a
    script grown page by page keeps the `prs.save` it had when it was shorter. A live
    16-page run had exactly that line after its eighth page: the file existed from page
    8 on, so the rescue skipped, and the eight-page deck it wrote was kept as the record
    of a run that had drawn ten. Two pages thrown away by the line meant to save them.

    Now the rescue asks whether the script reached its own last save, not whether a file
    is there."""
    from pptx import Presentation

    outcome = await _build(project, SAVES_HALFWAY)

    assert not outcome.ok, "an interrupted build is not a deck"
    kept = project.review_dir / "build_failures" / "failure-001" / "deck.pptx.building"
    assert kept.is_file()
    deck = Presentation(str(kept))
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
    assert len(texts) == 3, f"the stale one-page save was kept instead of the three pages that drew: {texts}"
    assert "First" in texts[0] and "Second" in texts[1] and "Third" in texts[2]


REWRITES_PAST_THE_SAVE = DECK.replace(
    'prs.save(os.environ["PPT_OUTPUT"])',
    'prs.save(os.environ["PPT_OUTPUT"])\ntitle(one, "rewritten past the save")',
)


@pytest.mark.asyncio
async def test_a_build_that_finished_is_the_file_the_script_saved(project: Project) -> None:
    """The other half of the rescue's condition. It now writes over whatever is at
    PPT_OUTPUT, so it must not run at all on a build that finished -- a deck saved and
    then edited in memory would otherwise be delivered as the edit, and no author writing
    `prs.save` last expects the bytes on disk to be anything but what they saved.

    The page is rewritten after the save here because that is the only difference a
    second save would show."""
    from pptx import Presentation

    outcome = await _build(project, REWRITES_PAST_THE_SAVE)

    assert outcome.ok and outcome.pages == 2, outcome.stderr
    deck = Presentation(str(outcome.pptx_path))
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
    assert "Unified video segmentation" in texts[0], f"the deck was written again after the script saved it: {texts}"
    assert "rewritten past the save" not in texts[0]


STAND_IN_CANNOT_WRITE = textwrap.dedent(
    """
    import os
    from pptx import Presentation
    from pptx.shapes.shapetree import SlideShapes
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    _add_textbox = SlideShapes.add_textbox
    _refused = []


    def _refuse_once(self, *args, **kwargs):
        if not _refused:
            _refused.append(1)
            raise RuntimeError("this deck cannot take a textbox")
        return _add_textbox(self, *args, **kwargs)


    def new_slide(text):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        box.text_frame.text = text
        return slide


    # SLIDE 1
    new_slide("First")

    # SLIDE 2
    new_slide("Second")
    SlideShapes.add_textbox = _refuse_once
    raise ValueError("page two")

    # SLIDE 3
    new_slide("Third")

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.mark.asyncio
async def test_a_stand_in_that_cannot_be_written_on_still_holds_its_place(project: Project) -> None:
    """The placeholder is drawn inside the `except` that caught the page, so a raise
    while drawing it leaves the loop and every page after it goes with the one that
    failed -- the whole thing the loop exists to prevent. What can raise in there is a
    deck with no slide size or no layout the placeholder can use, which a script cannot
    construct; refusing the textbox once is the same code path.

    The page is added before anything is written on it, because the page is what keeps
    page and block paired by position. The failure is named in the record either way."""
    from pptx import Presentation

    from raven_ppt.backends.script import page_failures

    outcome = await _build(project, STAND_IN_CANNOT_WRITE)

    assert outcome.ok and outcome.pages == 3, outcome.stderr
    deck = Presentation(str(outcome.pptx_path))
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
    assert "First" in texts[0] and "Third" in texts[2], f"page 2's placeholder cost the pages around it: {texts}"
    assert not texts[1].strip(), "the placeholder could not be written on, so it says nothing"
    assert [entry["page"] for entry in page_failures(project)] == [2]
    assert "ValueError: page two" in page_failures(project)[0]["error"]
