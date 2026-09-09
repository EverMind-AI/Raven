"""What `measure.figures` reports, and the crops it must not report.

The crop cases carry the weight. A .pptx says twice over which part of an image is
shown and where it lands, and a check that reads only `shape.image.size` calls a
correctly cropped photograph distorted: on 250 pictures across the twelve bundled
templates and eight built decks, the raw reading reports 81 of them (up to 4.26x) and
every one is false. Reading `a:srcRect` alone still reports 37. So the fixtures below
reproduce both spellings from the decks they were measured on.
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest

from raven_ppt.contracts.findings import Severity
from raven_ppt.contracts.masters import Bands
from raven_ppt.services.measure.bands import band_of
from raven_ppt.services.measure.figures import (
    DISTORTION_LIMIT,
    SUPPORT_MIN_IN,
    Figure,
    cropped_figures,
    distorted_figures,
    figure_findings,
    figure_roles,
    read_figures,
    title_band_figures,
)
from raven_ppt.services.measure.figures import _pixels as pixels_of
from raven_ppt.services.measure.geometry import Rect
from tests._ppt_engine_fixtures import deck, image, noise_image, noise_png, product_page, template_file  # noqa: F401

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# The grid measured off the delivered deck the module was calibrated on: kicker at
# 0.34in, title to 1.22in, its rule to 1.33in, the footer rule at 6.62in.
GRID = Bands(0.0, 1.45, 6.55, 7.5, 13.333, 7.5)


@pytest.fixture
def png(tmp_path: Path):
    """A solid PNG at an exact pixel size, so a frame's proportions can be aimed at it."""
    from PIL import Image

    def build(width: int, height: int, colour: tuple[int, int, int] = (200, 60, 40)) -> Path:
        path = tmp_path / f"{width}x{height}_{colour[0]}.png"
        Image.new("RGB", (width, height), colour).save(path)
        return path

    return build


def crop(shape: Any, *, left: float = 0.0, right: float = 0.0, top: float = 0.0, bottom: float = 0.0) -> None:
    shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom = left, right, top, bottom


def fill_rect(shape: Any, **sides: float) -> None:
    """Write `a:stretch/a:fillRect`, which is how a designer's tool spells cover."""
    rect = shape._element.blipFill.find(f"{{{_A}}}stretch/{{{_A}}}fillRect")
    for side, share in sides.items():
        rect.set(side, str(int(round(share * 100000))))


def tile(shape: Any) -> None:
    from lxml import etree

    fill = shape._element.blipFill
    fill.remove(fill.find(f"{{{_A}}}stretch"))
    etree.SubElement(fill, f"{{{_A}}}tile")


def kinds(findings: list[Any]) -> list[str]:
    return [finding.kind for finding in findings]


def test_a_crop_that_makes_a_wide_photograph_fit_a_narrow_frame_is_not_distortion(deck, png):
    """The delivered deck's page 3, reproduced: 1280x720 taken 14.75% off each side."""
    picture = deck.picture(deck.page(), png(1280, 720), left=7.16, top=1.68, width=5.39, height=4.30)
    crop(picture, left=0.1475, right=0.1475)

    figure = read_figures(deck.save())[0]

    assert figure.crop == pytest.approx((0.1475, 0.1475, 0.0, 0.0))
    assert figure.source_aspect == pytest.approx(1.253, abs=0.002)
    assert figure.distortion == pytest.approx(1.0, abs=0.002)
    assert distorted_figures([figure]) == []


def test_a_frame_that_keeps_under_half_its_image_is_reported(deck, png):
    """A measured deck cover-fitted a 16:9 generation into an 11.9x1.9in band and kept 28%
    of it: not a distortion -- the proportions are right -- but what the reader saw was a
    strip of sky. Under half of the image shown is a frame asking for a picture of
    another shape, and the finding says which shape."""
    band = deck.picture(deck.page(), png(1536, 864), left=0.72, top=1.6, width=11.88, height=1.88)
    crop(band, top=0.3589, bottom=0.3589)

    found = cropped_figures(read_figures(deck.save()))

    assert [one.kind for one in found] == ["figure_crop"]
    assert found[0].detail["shown_share"] == pytest.approx(0.282, abs=0.002)
    assert "aspect_ratio nearest 6.32" in found[0].message


def test_a_photograph_trimmed_to_its_frame_is_not_a_crop_finding(deck, png):
    """A 3:2 photograph in a 16:9 frame keeps 84% of itself, which is a fit, not a loss."""
    photo = deck.picture(deck.page(), png(3000, 2000), left=1.0, top=1.0, width=8.0, height=4.5)
    crop(photo, top=0.0788, bottom=0.0788)

    assert cropped_figures(read_figures(deck.save())) == []


def test_the_same_frame_without_the_crop_is_reported(deck, png):
    deck.picture(deck.page(), png(1280, 720), left=7.16, top=1.68, width=5.39, height=4.30)

    found = distorted_figures(read_figures(deck.save()))

    assert kinds(found) == ["figure_distortion"]
    assert found[0].detail["distortion"] == pytest.approx(1.42, abs=0.01)


def test_a_stretch_names_the_height_the_frame_needs(deck, png):
    deck.picture(deck.page(), png(1600, 900), left=1.0, top=2.0, width=6.0, height=4.5)

    found = distorted_figures(read_figures(deck.save()))

    assert found[0].page == 1
    assert found[0].detail["fitted_height_in"] == pytest.approx(3.38, abs=0.01)
    assert "3.38in tall" in found[0].message


def test_cover_written_as_a_negative_fill_rect_is_not_distortion(deck, png):
    """A cover crop written as a negative `fillRect`: frame 1.06, image 0.67, correct.

    The shape a template gives a portrait photograph in a squarer frame -- the
    renderer scales to cover and the negative top and bottom insets are what falls
    outside. Reading the frame alone calls it a 1.58x stretch, which the next test
    pins down."""
    picture = deck.picture(deck.page(), png(1280, 1920), left=0.73, top=2.52, width=2.26, height=2.14)
    fill_rect(picture, t=-0.28984, b=-0.28984)

    figure = read_figures(deck.save())[0]

    assert figure.box.width / figure.box.height == pytest.approx(1.056, abs=0.005)
    assert figure.display_box.height == pytest.approx(3.38, abs=0.02)
    assert figure.distortion == pytest.approx(1.0, abs=0.01)


def test_reading_only_the_source_rectangle_would_have_reported_that_cover(deck, png):
    picture = deck.picture(deck.page(), png(1280, 1920), left=0.73, top=2.52, width=2.26, height=2.14)
    fill_rect(picture, t=-0.28984, b=-0.28984)
    figure = read_figures(deck.save())[0]

    frame_only = replace(figure, inset=(0.0, 0.0, 0.0, 0.0))

    assert frame_only.distortion > DISTORTION_LIMIT


def test_a_fit_stated_twice_is_read_as_the_one_fit_it_is(deck, png):
    """`warm_bauhaus_quarterly_review` page 4, rebuilt: the template's inset outlived
    the crop `_fill_crop` wrote for the replacement image, and the pair read as 1.64x."""
    picture = deck.picture(deck.page(), png(800, 458), left=0.0, top=1.49, width=5.49, height=6.01)
    crop(picture, left=0.239, right=0.239)
    fill_rect(picture, l=-0.322, r=-0.320)

    figure = read_figures(deck.save())[0]

    assert figure.source_aspect == pytest.approx(0.913, abs=0.005)
    assert figure.display_box.width == pytest.approx(9.01, abs=0.02)
    assert figure.fitted_box.width == pytest.approx(5.49, abs=0.02)
    assert figure.distortion == pytest.approx(1.0, abs=0.01)
    assert distorted_figures([figure]) == []


def test_a_fit_stated_twice_and_wrong_both_ways_is_still_reported(deck, png):
    picture = deck.picture(deck.page(), png(1600, 900), left=1.0, top=2.0, width=6.0, height=4.5)
    crop(picture, top=0.3, bottom=0.3)
    fill_rect(picture, l=-0.2, r=-0.2)

    found = distorted_figures(read_figures(deck.save()))

    assert kinds(found) == ["figure_distortion"]


def test_a_tiled_fill_is_not_measured(deck, png):
    picture = deck.picture(deck.page(), png(1600, 900), left=1.0, top=2.0, width=6.0, height=4.5)
    tile(picture)

    figure = read_figures(deck.save())[0]

    assert figure.tiled
    assert figure.distortion is None
    assert distorted_figures([figure]) == []


def test_an_image_whose_pixels_cannot_be_read_is_skipped_rather_than_guessed_at():
    assert pixels_of(b"\x01\x00\x00\x00 not a raster") is None
    blind = Figure(
        page=1,
        name="Picture 1",
        box=Rect(1.0, 1.0, 9.0, 2.0),
        digest="0" * 40,
        pixels=None,
        crop=(0.0, 0.0, 0.0, 0.0),
        inset=(0.0, 0.0, 0.0, 0.0),
        tiled=False,
    )

    assert blind.source_aspect is None
    assert blind.distortion is None
    assert distorted_figures([blind]) == []


def test_role_is_graded_by_share_of_the_content_band(deck, png):
    page = deck.page()
    deck.picture(page, png(1600, 800, (10, 10, 10)), left=0.5, top=1.6, width=8.0, height=4.0)
    deck.picture(page, png(1600, 1440, (20, 90, 30)), left=9.0, top=1.6, width=2.0, height=1.8)
    deck.picture(page, png(64, 64, (90, 20, 90)), left=12.5, top=6.8, width=0.5, height=0.5)

    roles = [figure.role(GRID) for figure in read_figures(deck.save())]

    assert roles == ["hero", "support", "mark"]


def test_a_figure_with_a_figures_area_and_a_slivers_height_is_reported(deck, png):
    deck.picture(deck.page(), png(1600, 120), left=1.0, top=2.0, width=8.0, height=0.6)

    found = figure_roles(read_figures(deck.save()), GRID)

    assert kinds(found) == ["figure_undersized"]
    assert found[0].detail["box_in"] == [8.0, 0.6]
    assert "not tall enough" in found[0].message


def test_a_support_figure_that_clears_the_floor_is_not_reported(deck, png):
    deck.picture(deck.page(), png(1600, 1440), left=1.0, top=2.0, width=2.0, height=1.8)

    figure = read_figures(deck.save())[0]

    assert figure.role(GRID) == "support"
    assert figure.box.width >= SUPPORT_MIN_IN[0] and figure.box.height >= SUPPORT_MIN_IN[1]
    assert figure_roles([figure], GRID) == []


def test_without_bands_the_role_and_title_families_say_nothing_and_distortion_still_speaks(deck, png):
    deck.picture(deck.page(), png(1600, 900), left=1.0, top=2.0, width=8.0, height=0.6)
    built = deck.save()
    figures = read_figures(built)

    assert figure_roles(figures, None) == []
    assert title_band_figures(figures, None) == []
    assert kinds(figure_findings(built, None)) == ["figure_distortion"]


def test_a_repeated_mark_that_lands_somewhere_different_is_reported_once(deck, png):
    mark = png(64, 64, (30, 30, 200))
    for left, top in ((12.0, 0.2), (12.4, 0.2), (12.0, 0.9)):
        deck.picture(deck.page(), mark, left=left, top=top, width=0.45, height=0.45)

    found = figure_roles(read_figures(deck.save()), GRID)

    assert kinds(found) == ["figure_mark_drift"]
    assert found[0].page is None
    assert found[0].detail["pages"] == [1, 2, 3]
    assert found[0].detail["spread_in"] == [0.4, 0.7]


def test_a_repeated_mark_that_stays_put_is_not(deck, png):
    mark = png(64, 64, (30, 30, 200))
    for _ in range(3):
        deck.picture(deck.page(), mark, left=12.0, top=0.2, width=0.45, height=0.45)

    assert figure_roles(read_figures(deck.save()), GRID) == []


def test_a_drifting_mark_is_not_also_reported_as_title_band_decoration(deck, png):
    mark = png(64, 64, (30, 30, 200))
    for left, top in ((12.0, 0.2), (12.4, 0.2), (12.0, 0.9)):
        deck.picture(deck.page(), mark, left=left, top=top, width=0.45, height=0.45)

    assert kinds(figure_findings(deck.save(), GRID)) == ["figure_mark_drift"]


def test_a_small_picture_inside_the_title_band_is_reported(deck, png):
    deck.picture(deck.page(), png(240, 100), left=9.0, top=0.6, width=1.2, height=0.5)

    found = title_band_figures(read_figures(deck.save()), GRID)

    assert kinds(found) == ["title_band_figure"]
    assert found[0].page == 1
    assert found[0].detail["size_in"] == [1.2, 0.5]


def test_a_corner_ornament_drawn_to_the_trim_is_not(deck, png):
    """`red_chinese_traditional_culture` page 6: flush to the top and right trim."""
    deck.picture(deck.page(), png(240, 210), left=12.57, top=0.0, width=0.76, height=0.68)

    assert title_band_figures(read_figures(deck.save()), GRID) == []


def test_a_picture_that_only_reaches_into_the_title_band_is_not(deck, png):
    deck.picture(deck.page(), png(240, 120), left=9.0, top=0.95, width=1.2, height=0.6)
    figures = read_figures(deck.save())

    assert band_of(GRID, figures[0].box.y0, figures[0].box.y1) == "title"
    assert title_band_figures(figures, GRID) == []


def test_a_masthead_the_deck_repeats_is_not(deck, png):
    logo = png(240, 100, (10, 120, 90))
    for _ in range(3):
        deck.picture(deck.page(), logo, left=9.0, top=0.6, width=1.2, height=0.5)

    assert title_band_figures(read_figures(deck.save()), GRID) == []


def test_a_deck_with_no_pictures_measures_to_nothing(deck):
    deck.page()

    assert figure_findings(deck.save(), GRID) == []


def test_an_unreadable_deck_is_not_a_measurement(tmp_path: Path):
    broken = tmp_path / "not-a-deck.pptx"
    broken.write_bytes(b"PK\x03\x04 truncated")

    assert read_figures(broken) == []
    assert figure_findings(broken, GRID) == []


def test_every_finding_this_module_makes_is_a_warning(deck, png):
    page = deck.page()
    deck.picture(page, png(1600, 900), left=1.0, top=2.0, width=8.0, height=0.6)
    deck.picture(page, png(240, 100), left=9.0, top=0.6, width=1.2, height=0.5)
    mark = png(64, 64, (30, 30, 200))
    for left in (12.0, 12.4, 12.9):
        deck.picture(deck.page(), mark, left=left, top=6.9, width=0.4, height=0.4)

    found = figure_findings(deck.save(), GRID)

    assert sorted(set(kinds(found))) == [
        "figure_distortion",
        "figure_mark_drift",
        "figure_undersized",
        "title_band_figure",
    ]
    assert {finding.severity for finding in found} == {Severity.WARNING}


def test_a_picture_filled_into_a_rounded_panel_is_read_like_a_picture_frame(deck, png, tmp_path: Path):
    """A template's photograph is as often a shape with a `blipFill` as it is a frame."""
    from pptx.util import Inches

    slide = deck.page()
    panel = slide.shapes.add_shape(5, Inches(1.0), Inches(2.0), Inches(6.0), Inches(4.5))
    image = png(1600, 900)
    _, relationship = panel.part.get_or_add_image_part(str(image))
    from lxml import etree

    from raven_ppt.services.measure.figures import _P

    properties = panel._element.find(f"{{{_P}}}spPr")
    fill = etree.SubElement(properties, f"{{{_A}}}blipFill")
    blip = etree.SubElement(fill, f"{{{_A}}}blip")
    blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", relationship)
    etree.SubElement(etree.SubElement(fill, f"{{{_A}}}stretch"), f"{{{_A}}}fillRect")

    figures = read_figures(deck.save())

    assert len(figures) == 1
    assert figures[0].pixels == (1600, 900)
    assert kinds(distorted_figures(figures)) == ["figure_distortion"]


def test_a_cover_spelled_as_a_negative_fill_rect_is_the_same_crop(deck, png):
    """The other spelling of the same loss: a designer's template photograph is fitted
    with `a:stretch/a:fillRect` and no `srcRect`. The frame is then a window onto a
    larger display box, and the share it shows is the frame's area over that box's."""
    from raven_ppt.services.measure.figures import Figure
    from raven_ppt.services.measure.geometry import Rect

    figure = Figure(
        page=3,
        name="sky",
        box=Rect(0.72, 1.6, 0.72 + 11.88, 1.6 + 1.88),
        digest="x",
        pixels=(1536, 864),
        crop=(0.0, 0.0, 0.0, 0.0),
        inset=(0.0, 0.0, -1.2766, -1.2766),
        tiled=False,
    )

    assert figure.shown_share == pytest.approx(0.281, abs=0.002)
    found = cropped_figures([figure])
    assert [one.kind for one in found] == ["figure_crop"]
    assert found[0].detail["shown_share"] == pytest.approx(0.281, abs=0.002)


def test_a_fit_stated_twice_is_counted_once():
    """`compose._fill_crop` restates a template's fillRect fit as an srcRect against the
    frame; read as two crops the pair says a quarter is shown when the page shows over
    half. The share follows the reading `fitted_box` chose."""
    from raven_ppt.services.measure.figures import Figure
    from raven_ppt.services.measure.geometry import Rect

    twice = Figure(
        page=1,
        name="photo",
        box=Rect(0.0, 0.0, 8.0, 4.5),
        digest="x",
        pixels=(3000, 2000),
        crop=(0.0, 0.0, 0.0788, 0.0788),
        inset=(0.0, 0.0, -0.0935, -0.0935),
        tiled=False,
    )

    assert twice.shown_share == pytest.approx(0.842, abs=0.01)
    assert cropped_figures([twice]) == []


def _page_with_picture(tmp_path: Path, alpha: float, *, share: float = 1.0) -> Path:
    """One 16:9 page with a picture over `share` of it, washed to `alpha`."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.template.compose import wash

    photo = tmp_path / f"photo-{alpha}-{share}.png"
    Image.new("RGB", (800, 450), (30, 60, 90)).save(photo)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture = slide.shapes.add_picture(str(photo), 0, 0, Inches(13.333 * share), Inches(7.5))
    wash(picture, alpha)
    out = tmp_path / f"deck-{alpha}-{share}.pptx"
    presentation.save(str(out))
    return out


def test_a_page_sized_picture_washed_between_texture_and_full_strength_reads_as_fog(tmp_path: Path) -> None:
    """The live cover: a night skyline at 30% on a white page under black type. Texture
    (0.1) and full strength (1.0, the scrim form) are the two ends that read; a washed
    picture that is not the page (half of it) is a figure, not a backdrop."""
    from raven_ppt.services.measure.figures import washed_backdrops

    fog = washed_backdrops(_page_with_picture(tmp_path, 0.3))
    assert [f.kind for f in fog] == ["washed_backdrop"]
    assert fog[0].page == 1 and fog[0].severity is Severity.WARNING
    assert "30%" in fog[0].message and "plane of ink" in fog[0].message

    assert washed_backdrops(_page_with_picture(tmp_path, 0.1)) == []
    assert washed_backdrops(_page_with_picture(tmp_path, 1.0)) == []
    assert washed_backdrops(_page_with_picture(tmp_path, 0.3, share=0.5)) == []
