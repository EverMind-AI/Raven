"""Edges that were meant to line up: what the check reports, and what it refuses to.

Half of these tests are the refusals. The measurement's whole risk is reporting a
page a reader would call fine (design doc D17), so every discriminator it uses --
the file's declared x, the paragraph's alignment, wrapping, the run it sits in,
the glyph's own side bearing, a bar drawn in the gap -- is pinned here by a case
either side of it.

The rendered word boxes are stated rather than rendered: a unit test has no
LibreOffice, and what a check does with a given render is exactly what has to be
pinned. The numbers in the fixtures are the ones measured on the deck this came
from -- a 7.1pt offset on 15.6pt copy -- so a threshold that stops catching it
fails here.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from raven_ppt.contracts.findings import Severity
from raven_ppt.services.measure.alignment import (
    BEARING_SHARE,
    DRIFTS_PER_PAGE,
    INDENT_PT,
    RUN_CEILING,
    SIBLING_DRIFT_PT,
    alignment_findings,
    flush_drift,
    sibling_drift,
)
from raven_ppt.services.measure.words import WordBox
from tests._ppt_engine_fixtures import (  # noqa: F401
    DeckBuilder,
    deck,
    image,
    noise_image,
    noise_png,
    product_page,
    template_file,
)

pytest.importorskip("pptx")

PT = 72.0
# The deck this check was built from: 12pt CJK body, whose rendered word boxes come
# back 15.6pt tall, set one inherited 0.1in inset out of line.
BODY_PT = 15.6
SLIP_PT = 7.1


def _word(text: str, left: float, top: float, width: float, height: float = BODY_PT, page: int = 1) -> WordBox:
    return WordBox(page=page, text=text, x0=left, y0=top, x1=left + width, y1=top + height)


def _left_aligned(box: Any) -> Any:
    from pptx.enum.text import PP_ALIGN

    for para in box.text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT
    return box


def _inset(box: Any, inches: float) -> Any:
    from pptx.util import Inches

    box.text_frame.margin_left = Inches(inches)
    return box


def _stack(deck: DeckBuilder, *, at: float, tops: tuple[float, ...], width: float = 3.0) -> Any:
    """A column of one-line text boxes, all declared at the same left edge."""
    page = deck.page()
    for top in tops:
        _left_aligned(deck.text(page, "line", left=at, top=top / PT, width=width, height=0.3, wrap=True))
    return page


# --- the offset the file cannot see -----------------------------------------


def test_a_block_the_file_declares_flush_and_the_render_sets_out_is_reported(deck: DeckBuilder) -> None:
    """The case this exists for: same declared x, two different rendered edges."""
    _stack(deck, at=1.0, tops=(100.0, 130.0))
    words = [_word("aligned", PT, 100.0, 90.0), _word("adrift", PT + SLIP_PT, 130.0, 90.0)]

    found = flush_drift(deck.save(), words)

    assert [finding.kind for finding in found] == ["flush_drift"]
    assert found[0].severity is Severity.WARNING
    assert found[0].page == 1
    assert found[0].detail["drift_pt"] == pytest.approx(SLIP_PT, abs=0.01)
    assert "'adrift'" in found[0].message and "'aligned'" in found[0].message
    assert "to the right of" in found[0].message
    assert "0.10in left" in found[0].message


def test_the_same_rendered_offset_is_silent_when_the_file_declares_it(deck: DeckBuilder) -> None:
    """The line between an indent and a slip: the file said so, or it did not.

    Identical render, and the only difference is that the author moved the box
    rather than leaving an inset to move the copy.
    """
    page = deck.page()
    _left_aligned(deck.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True))
    _left_aligned(deck.text(page, "line", left=1.0 + SLIP_PT / PT, top=130.0 / PT, width=3.0, height=0.3, wrap=True))
    words = [_word("aligned", PT, 100.0, 90.0), _word("indented", PT + SLIP_PT, 130.0, 90.0)]

    assert flush_drift(deck.save(), words) == []


def test_the_left_inset_behind_the_offset_is_named_in_the_message(deck: DeckBuilder) -> None:
    """What to change, not just what is wrong -- the frames differ, not the boxes."""
    page = deck.page()
    _inset(_left_aligned(deck.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True)), 0.0)
    _inset(_left_aligned(deck.text(page, "line", left=1.0, top=130.0 / PT, width=3.0, height=0.3, wrap=True)), 0.1)
    words = [_word("aligned", PT, 100.0, 90.0), _word("adrift", PT + 7.2, 130.0, 90.0)]

    found = flush_drift(deck.save(), words)

    assert "left insets (0.10in against 0.00in)" in found[0].message
    assert found[0].detail["below_inset_in"] == pytest.approx(0.1)


def test_one_inherited_inset_reports_as_one_offset_and_counts_the_rest(deck: DeckBuilder) -> None:
    """Four blocks, one cause, one finding -- the page is the unit, not the block."""
    page = deck.page()
    words: list[WordBox] = []
    for index, at in enumerate((1.0, 4.0, 7.0, 10.0)):
        _left_aligned(deck.text(page, "line", left=at, top=100.0 / PT, width=2.5, height=0.3, wrap=True))
        _left_aligned(deck.text(page, "line", left=at, top=130.0 / PT, width=2.5, height=0.3, wrap=True))
        words += [
            _word(f"head{index}", at * PT, 100.0, 90.0),
            _word(f"tail{index}", at * PT + SLIP_PT, 130.0, 90.0),
        ]

    found = flush_drift(deck.save(), words)

    assert len(found) == 1
    assert "3 more blocks on this page sit the same amount out" in found[0].message
    assert found[0].detail["also"] == ["tail1", "tail2", "tail3"]


def test_a_page_reports_at_most_the_offsets_it_is_allowed(deck: DeckBuilder) -> None:
    page = deck.page()
    words: list[WordBox] = []
    for index, (at, offset) in enumerate(((1.0, 4.0), (5.0, 8.0), (9.0, 12.0))):
        _left_aligned(deck.text(page, "line", left=at, top=100.0 / PT, width=2.5, height=0.3, wrap=True))
        _left_aligned(deck.text(page, "line", left=at, top=130.0 / PT, width=2.5, height=0.3, wrap=True))
        words += [
            _word(f"head{index}", at * PT, 100.0, 90.0),
            _word(f"tail{index}", at * PT + offset, 130.0, 90.0),
        ]

    found = flush_drift(deck.save(), words, per_page=DRIFTS_PER_PAGE)

    assert len(found) == DRIFTS_PER_PAGE
    assert [round(finding.detail["drift_pt"]) for finding in found] == [12, 8]


# --- what the glyph, the frame and the layout explain -------------------------


def test_an_offset_the_glyph_itself_could_account_for_is_left_alone(deck: DeckBuilder) -> None:
    """A word's box starts at its ink, and big type has a big side bearing.

    The same 5pt offset, on 15.6pt copy and on 58.5pt copy. On the first it is a
    third of the line's height and no font does that; on the second it is inside
    what a leading figure's own side bearing measured on this deck.
    """
    small = DeckBuilder(deck.tmp_path)
    _stack(small, at=1.0, tops=(100.0, 130.0))
    body = [_word("aligned", PT, 100.0, 90.0), _word("adrift", PT + 5.0, 130.0, 90.0)]

    assert 5.0 > BEARING_SHARE * BODY_PT
    assert flush_drift(small.save("small.pptx"), body) != []

    large = DeckBuilder(deck.tmp_path)
    _stack(large, at=1.0, tops=(100.0, 200.0))
    display = [
        _word("100", PT, 100.0, 120.0, height=58.5),
        _word("10", PT + 5.0, 200.0, 120.0, height=58.5),
    ]

    assert 5.0 < BEARING_SHARE * 58.5
    assert flush_drift(large.save("large.pptx"), display) == []


def test_an_indent_deeper_than_a_quarter_inch_is_a_decision(deck: DeckBuilder) -> None:
    _stack(deck, at=1.0, tops=(100.0, 130.0))
    words = [_word("aligned", PT, 100.0, 90.0), _word("indented", PT + INDENT_PT + 1, 130.0, 90.0)]

    assert flush_drift(deck.save(), words) == []


def test_a_frame_with_wrapping_off_is_placed_by_the_renderer_not_by_the_file(deck: DeckBuilder) -> None:
    """Its declared x says nothing about where its copy starts, so it is not read."""
    page = deck.page()
    _left_aligned(deck.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True))
    _left_aligned(deck.text(page, "line", left=1.0, top=130.0 / PT, width=3.0, height=0.3, wrap=False))
    words = [_word("aligned", PT, 100.0, 90.0), _word("autofit", PT + SLIP_PT, 130.0, 90.0)]

    assert flush_drift(deck.save(), words) == []


def test_copy_the_file_centres_has_no_left_edge_to_compare(deck: DeckBuilder) -> None:
    from pptx.enum.text import PP_ALIGN

    page = deck.page()
    _left_aligned(deck.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True))
    centred = deck.text(page, "line", left=1.0, top=130.0 / PT, width=3.0, height=0.3, wrap=True)
    for para in centred.text_frame.paragraphs:
        para.alignment = PP_ALIGN.CENTER
    words = [_word("aligned", PT, 100.0, 90.0), _word("centred", PT + SLIP_PT, 130.0, 90.0)]

    assert flush_drift(deck.save(), words) == []


def test_copy_with_no_declared_alignment_is_read_only_when_the_render_settles_it(deck: DeckBuilder) -> None:
    """Undeclared and plainly left-set is read; undeclared and ambiguous is not."""
    plain = DeckBuilder(deck.tmp_path)
    page = plain.page()
    plain.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True)
    plain.text(page, "line", left=1.0, top=130.0 / PT, width=3.0, height=0.3, wrap=True)
    roomy = [_word("aligned", PT, 100.0, 40.0), _word("adrift", PT + SLIP_PT, 130.0, 40.0)]

    assert flush_drift(plain.save("plain.pptx"), roomy) != []

    tight = DeckBuilder(deck.tmp_path)
    page = tight.page()
    tight.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True)
    tight.text(page, "line", left=1.0, top=130.0 / PT, width=3.0, height=0.3, wrap=True)
    filling = [_word("aligned", PT, 100.0, 210.0), _word("adrift", PT + SLIP_PT, 130.0, 200.0)]

    assert flush_drift(tight.save("tight.pptx"), filling) == []


def test_a_bar_drawn_in_the_gap_is_why_the_copy_moved_in(deck: DeckBuilder) -> None:
    """An accent rule at a block's left is the reason its copy sits further right."""
    page = deck.page()
    _left_aligned(deck.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True))
    _left_aligned(deck.text(page, "line", left=1.0, top=130.0 / PT, width=3.0, height=0.3, wrap=True))
    deck.panel(page, left=1.0, top=129.0 / PT, width=SLIP_PT / PT, height=0.25)
    words = [_word("aligned", PT, 100.0, 90.0), _word("beside a rule", PT + SLIP_PT, 130.0, 90.0)]

    assert flush_drift(deck.save(), words) == []


def test_a_band_that_runs_under_the_copy_explains_nothing(deck: DeckBuilder) -> None:
    """The other side of the same rule: a row band spans the row, so it is not a reason.

    Without this the check would excuse every misaligned row of a striped table --
    which is one of the pages it was built to report.
    """
    page = deck.page()
    _left_aligned(deck.text(page, "line", left=1.0, top=100.0 / PT, width=3.0, height=0.3, wrap=True))
    _left_aligned(deck.text(page, "line", left=1.0, top=130.0 / PT, width=3.0, height=0.3, wrap=True))
    deck.panel(page, left=1.0, top=129.0 / PT, width=3.0, height=0.25)
    words = [_word("aligned", PT, 100.0, 90.0), _word("on a band", PT + SLIP_PT, 130.0, 90.0)]

    assert flush_drift(deck.save(), words) != []


def test_blocks_too_far_apart_to_read_as_one_run_are_not_compared(deck: DeckBuilder) -> None:
    """A card's footer answers to the footers beside it, not to the bullets above it."""
    apart = RUN_CEILING * BODY_PT + 10
    _stack(deck, at=1.0, tops=(100.0, 100.0 + apart))
    words = [_word("bullet", PT, 100.0, 90.0), _word("footer", PT + SLIP_PT, 100.0 + apart, 90.0)]

    assert flush_drift(deck.save(), words) == []


def test_a_run_is_judged_against_its_own_rhythm(deck: DeckBuilder) -> None:
    """Four blocks 30pt apart and a fifth 90pt below: the fifth is a separate thing."""
    tops = (100.0, 130.0, 160.0, 190.0, 280.0)
    _stack(deck, at=1.0, tops=tops)
    words = [_word(f"line{index}", PT, top, 90.0) for index, top in enumerate(tops[:-1])]
    words.append(_word("elsewhere", PT + SLIP_PT, tops[-1], 90.0))

    assert flush_drift(deck.save(), words) == []


# --- containers standing side by side -----------------------------------------


def _row(deck: DeckBuilder, tops: tuple[tuple[float, ...], ...]) -> tuple[Path, list[WordBox]]:
    """Two cards side by side, each holding one block per top given for it."""
    page = deck.page()
    words: list[WordBox] = []
    for card, (left, column) in enumerate(zip((1.0, 6.0), tops)):
        deck.panel(page, left=left, top=1.0, width=4.0, height=3.0)
        for index, top in enumerate(column):
            _left_aligned(deck.text(page, "line", left=left + 0.2, top=top / PT, width=3.0, height=0.3, wrap=True))
            words.append(_word(f"card{card}block{index}", (left + 0.2) * PT, top, 90.0))
    return deck.save(), words


def test_the_same_block_of_two_side_by_side_containers_must_share_a_height(deck: DeckBuilder) -> None:
    path, words = _row(deck, ((100.0, 150.0), (100.0, 150.0 + SIBLING_DRIFT_PT + 4)))

    found = sibling_drift(path, words)

    assert [finding.kind for finding in found] == ["sibling_drift"]
    assert found[0].severity is Severity.WARNING
    assert found[0].detail["block"] == 2
    assert found[0].detail["of"] == 2
    assert "last block in each" in found[0].message
    assert "'card1block1'" in found[0].message


def test_a_difference_a_reader_would_not_pick_up_is_left_alone(deck: DeckBuilder) -> None:
    path, words = _row(deck, ((100.0, 150.0), (100.0, 150.0 + SIBLING_DRIFT_PT - 1)))

    assert sibling_drift(path, words) == []


def test_containers_holding_different_numbers_of_blocks_have_nothing_to_correspond(
    deck: DeckBuilder,
) -> None:
    """A card with one block more is not a card whose tail is out of line."""
    path, words = _row(deck, ((100.0, 150.0), (100.0, 150.0, 200.0)))

    assert sibling_drift(path, words) == []


def test_containers_stacked_rather_than_beside_each_other_are_not_a_row(deck: DeckBuilder) -> None:
    """Two containers over each other hold the same blocks, and correspond in nothing.

    The lower one sets its second block 20pt further down its own card than the
    upper one does. Side by side that is the finding above; stacked it is two
    cards a reader never compares.
    """
    page = deck.page()
    words: list[WordBox] = []
    for card, (panel_top, tops) in enumerate(((72.0, (100.0, 150.0)), (288.0, (316.0, 386.0)))):
        deck.panel(page, left=1.0, top=panel_top / PT, width=4.0, height=2.0)
        for index, top in enumerate(tops):
            _left_aligned(deck.text(page, "line", left=1.2, top=top / PT, width=3.0, height=0.3, wrap=True))
            words.append(_word(f"card{card}block{index}", 1.2 * PT, top, 90.0))

    assert sibling_drift(deck.save(), words) == []


# --- no render, nothing to say ------------------------------------------------


@pytest.mark.parametrize("check", [flush_drift, sibling_drift, alignment_findings])
def test_no_rendered_words_is_no_signal_rather_than_a_clean_deck(deck: DeckBuilder, check: Any) -> None:
    _stack(deck, at=1.0, tops=(100.0, 130.0))
    path = deck.save()

    assert check(path, None) == []
    assert check(path, []) == []


def test_both_readings_come_back_from_one_call(deck: DeckBuilder) -> None:
    """One card's tail sits low and starts wide: one page, one call, two kinds."""
    page = deck.page()
    words: list[WordBox] = []
    for card, left in enumerate((1.0, 6.0)):
        deck.panel(page, left=left, top=1.0, width=4.0, height=3.0)
        for index, top in enumerate((100.0, 130.0)):
            drop = SIBLING_DRIFT_PT + 4 if card and index else 0.0
            slip = SLIP_PT if card and index else 0.0
            _left_aligned(
                deck.text(page, "line", left=left + 0.2, top=(top + drop) / PT, width=3.0, height=0.3, wrap=True)
            )
            words.append(_word(f"card{card}block{index}", (left + 0.2) * PT + slip, top + drop, 90.0))

    found = alignment_findings(deck.save(), words)

    assert {finding.kind for finding in found} == {"flush_drift", "sibling_drift"}
    assert all(finding.severity is Severity.WARNING for finding in found)
