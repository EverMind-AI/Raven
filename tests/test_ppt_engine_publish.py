"""Delivery, and the three ways the predecessor lost control of it."""

from __future__ import annotations

from pathlib import Path

import pytest

from raven_ppt.contracts import Finding, Project, Severity
from raven_ppt.services.publish import PublishRefusedError, publish, stage
from raven_ppt.services.publish.deliver import (
    delivery_report,
    published_digests,
    tampered_delivery,
    unrecorded_deliveries,
)


@pytest.fixture()
def project(tmp_path: Path) -> Project:
    p = Project(workspace=tmp_path, slug="tarvis")
    p.build_dir.mkdir(parents=True)
    p.state_dir.mkdir(parents=True)
    return p


def _deck(project: Project, body: bytes = b"PK\x03\x04 deck") -> Path:
    path = project.build_dir / "deck.pptx"
    path.write_bytes(body)
    return path


def _fact() -> Finding:
    return Finding(kind="fact", severity=Severity.BLOCKING, message="48.3 is not in the sources", page=11)


def _warning() -> Finding:
    return Finding(kind="type_floor", severity=Severity.WARNING, message="11.5pt body", page=3)


def test_a_deck_is_delivered_atomically_to_the_named_path(project: Project) -> None:
    staged = stage(project, _deck(project), pages=18)
    out = publish(project, staged, project.exports_dir / "TarViS.pptx", findings=[], blocking_kinds=frozenset())
    assert out.is_file() and out.read_bytes() == b"PK\x03\x04 deck"
    assert not list(out.parent.glob(".*tmp"))


def test_a_warning_does_not_stop_delivery(project: Project) -> None:
    """Measurements of the rendered page warn and ride along -- see design D2."""
    staged = stage(project, _deck(project))
    out = publish(
        project, staged, project.exports_dir / "d.pptx", findings=[_warning()], blocking_kinds=frozenset({"band"})
    )
    assert out.is_file()


def test_the_gate_cannot_be_forgotten_because_it_is_an_argument(project: Project) -> None:
    """Fail-closed that depends on being remembered is not fail-closed.

    The route that forgot it shipped a deck with seventeen colour bars and an
    unanchored number, because its publish step was a bare file copy.
    """
    staged = stage(project, _deck(project))
    with pytest.raises(TypeError):
        publish(project, staged, project.exports_dir / "d.pptx")  # type: ignore[call-arg]


def test_a_blocking_finding_refuses_and_names_what_stands(project: Project) -> None:
    staged = stage(project, _deck(project))
    with pytest.raises(PublishRefusedError, match="1 fact"):
        publish(project, staged, project.exports_dir / "d.pptx", findings=[_fact()], blocking_kinds=frozenset())
    assert not (project.exports_dir / "d.pptx").exists()


def test_a_kind_the_route_calls_fatal_refuses_even_at_warning_severity(project: Project) -> None:
    """A route's own kind, which the shared checks do not rank at all.

    `unmapped_page` is built inside the route that cares whether a page traces to
    its own code, so severity is not where its weight lives. The example used to be
    `band`, and that was the bug: the gate reported it, one route still called it
    fatal, and this test froze the contradiction as intended behaviour.
    """
    orphan = Finding(kind="unmapped_page", severity=Severity.WARNING, message="page 4 maps to no block", page=4)
    staged = stage(project, _deck(project))
    with pytest.raises(PublishRefusedError, match="1 unmapped_page"):
        publish(
            project,
            staged,
            project.exports_dir / "d.pptx",
            findings=[orphan],
            blocking_kinds=frozenset({"unmapped_page"}),
        )


def test_a_deck_that_changed_after_it_was_checked_is_refused(project: Project) -> None:
    """The findings describe a file that no longer exists."""
    staged = stage(project, _deck(project))
    staged.path.write_bytes(b"PK\x03\x04 something else")
    with pytest.raises(PublishRefusedError, match="changed after it was checked"):
        publish(project, staged, project.exports_dir / "d.pptx", findings=[], blocking_kinds=frozenset())


def test_staging_keeps_the_build_directory_copy(project: Project) -> None:
    """It is what a failed edit gets repaired against."""
    built = _deck(project)
    stage(project, built)
    assert built.is_file()


@pytest.mark.parametrize(
    "destination",
    ["../outside.pptx", "/etc/deck.pptx", "exports/deck.txt", "exports"],
)
def test_delivery_outside_the_workspace_or_of_the_wrong_kind_is_refused(project: Project, destination: str) -> None:
    (project.workspace / "exports").mkdir(exist_ok=True)
    staged = stage(project, _deck(project))
    with pytest.raises(PublishRefusedError):
        publish(project, staged, Path(destination), findings=[], blocking_kinds=frozenset())


def test_an_empty_build_is_not_stageable(project: Project) -> None:
    empty = project.build_dir / "deck.pptx"
    empty.write_bytes(b"")
    with pytest.raises(PublishRefusedError, match="empty"):
        stage(project, empty)


def test_a_missing_build_says_where_it_looked(project: Project) -> None:
    with pytest.raises(PublishRefusedError, match="no deck at"):
        stage(project, project.build_dir / "nope.pptx")


# --- the record against the file it names ---------------------------------------


def test_the_record_names_the_bytes_that_reached_the_path(project: Project) -> None:
    """Read back off the delivered file, not taken from the copy it was made from.

    The record's whole promise is "this sha256 is the file at this path", and the only
    way to keep it is to hash the file at that path after writing it.
    """
    import hashlib

    staged = stage(project, _deck(project), pages=4)
    out = publish(project, staged, project.exports_dir / "deck.pptx", findings=[], blocking_kinds=frozenset())

    assert published_digests(project.state_dir) == {hashlib.sha256(out.read_bytes()).hexdigest()}
    assert tampered_delivery(project.state_dir, out) is None


def test_a_delivery_edited_where_it_lies_is_named_and_not_believed(project: Project) -> None:
    """The live failure this reading exists for.

    A run holding a deck it had already delivered ran `python3 - # Apply the same fix to
    the published deck (out/deck.pptx) in place` through `exec` and rewrote the delivery
    where it lay. Nothing was copied anywhere -- the build directory and the staged copy
    both still held the recorded bytes -- so every record in the deck folder went on
    describing a file that no longer existed, and the user held a deck no gate had seen.
    """
    staged = stage(project, _deck(project), pages=4)
    out = publish(project, staged, project.exports_dir / "deck.pptx", findings=[], blocking_kinds=frozenset())

    out.write_bytes(b"PK\x03\x04 edited in place")

    report = tampered_delivery(project.state_dir, out)
    assert report is not None
    assert "changed after it was published" in report
    assert "never in the delivered file" in report
    assert unrecorded_deliveries(project.state_dir) == [report]
    # And the same answer through the entry point the stage uses, which resolves the
    # destination the way `publish` does.
    assert delivery_report(project, project.exports_dir / "deck.pptx") == report


def test_the_next_publish_puts_the_measured_deck_back_and_the_record_agrees(project: Project) -> None:
    """The repair is the publish itself: the gated deck is written over the edited one.

    Which is why this reads and never refuses -- the user must still get a deck, and the
    deck they should get is the one that passed.
    """
    import hashlib

    staged = stage(project, _deck(project), pages=4)
    out = publish(project, staged, project.exports_dir / "deck.pptx", findings=[], blocking_kinds=frozenset())
    out.write_bytes(b"PK\x03\x04 edited in place")

    again = stage(project, _deck(project, b"PK\x03\x04 rebuilt"), pages=4)
    out = publish(project, again, project.exports_dir / "deck.pptx", findings=[], blocking_kinds=frozenset())

    assert out.read_bytes() == b"PK\x03\x04 rebuilt"
    assert unrecorded_deliveries(project.state_dir) == []
    assert hashlib.sha256(out.read_bytes()).hexdigest() in published_digests(project.state_dir)


def test_a_record_that_cannot_be_written_is_not_a_silent_delivery(project: Project) -> None:
    """A delivery with no record is one the harness reads as a copy the model made.

    The write is small and the deck is megabytes, so this is the unlikely half of the
    pair -- and it is the half that used to raise a bare OSError past a caller catching
    only `PublishRefusedError`, leaving the file in place and the record describing the
    build before it.
    """
    from raven_ppt.services.publish.deliver import PUBLISHED_RECORD

    staged = stage(project, _deck(project), pages=4)
    record = project.state_dir / PUBLISHED_RECORD
    record.mkdir()  # a directory where the record goes: os.replace onto it fails

    with pytest.raises(PublishRefusedError, match="its record could not be"):
        publish(project, staged, project.exports_dir / "deck.pptx", findings=[], blocking_kinds=frozenset())


def test_a_deck_with_no_record_is_not_reported_as_changed(project: Project) -> None:
    """Nothing recorded for a path is not a claim about it, so there is nothing to break."""
    loose = project.exports_dir / "someone_elses.pptx"
    loose.parent.mkdir(parents=True, exist_ok=True)
    loose.write_bytes(b"PK\x03\x04 not ours")

    assert tampered_delivery(project.state_dir, loose) is None
    assert unrecorded_deliveries(project.state_dir) == []


# -- the destination the user named ------------------------------------------


def _delivered(project: Project, destination: Path, body: bytes = b"PK\x03\x04 deck"):
    from raven_ppt.services.publish import deliver

    staged = stage(project, _deck(project, body), pages=7)
    out = publish(project, staged, project.exports_dir / "deck.pptx", findings=[], blocking_kinds=frozenset())
    preview = out.with_suffix(".pdf")
    preview.write_bytes(b"%PDF-1.4 " + body)
    return deliver(project, out, staged.digest, staged.pages, destination=destination, preview=preview)


def test_a_destination_is_stated_once_and_kept_for_the_deck(project: Project, tmp_path: Path) -> None:
    from raven_ppt.services.publish import as_destination, read_destination, write_destination

    wanted = tmp_path / "handoff" / "intro.pptx"
    kept = write_destination(project, as_destination(str(wanted), project, default_name="deck.pptx"))

    assert kept == wanted and wanted.parent.is_dir(), "the directory is made when the destination is stated"
    assert read_destination(project) == wanted
    assert read_destination(Project(workspace=tmp_path / "other", slug="none")) is None


def test_a_directory_takes_the_decks_own_name(project: Project, tmp_path: Path) -> None:
    from raven_ppt.services.publish import as_destination

    existing = tmp_path / "deliverables"
    existing.mkdir()
    assert as_destination(str(existing), project, default_name="deck.pptx") == existing / "deck.pptx"
    trailing = tmp_path / "not-yet" / "here" / ""
    assert (
        as_destination(str(trailing) + "/", project, default_name="deck.pptx")
        == tmp_path / "not-yet" / "here" / "deck.pptx"
    )


@pytest.mark.parametrize(
    ("given", "said"),
    [
        ("decks/intro.pptx", "not an absolute path"),
        ("/{ws}/handoff/intro.txt", "not a .pptx path"),
        ("/{ws}/handoff/intro", "not a .pptx path"),
        ("/{ws}/deck/build/intro.pptx", "inside the engine's own deck/"),
        ("/{ws}/out/intro.pptx", "inside the engine's own out/"),
        ("", "is empty"),
    ],
)
def test_a_destination_the_publish_step_cannot_promise_is_refused(project: Project, given: str, said: str) -> None:
    from raven_ppt.services.publish import DestinationError, as_destination

    given = given.replace("/{ws}", str(project.workspace))
    with pytest.raises(DestinationError, match=said):
        as_destination(given, project, default_name="deck.pptx")


def test_a_destination_whose_directory_cannot_be_made_is_refused(project: Project, tmp_path: Path) -> None:
    from raven_ppt.services.publish import DestinationError, as_destination

    blocker = tmp_path / "a-file"
    blocker.write_text("not a directory", encoding="utf-8")
    with pytest.raises(DestinationError, match="cannot be created"):
        as_destination(str(blocker / "intro.pptx"), project, default_name="deck.pptx")


def test_the_delivery_holds_the_published_bytes_and_is_on_the_record(project: Project, tmp_path: Path) -> None:
    from raven_ppt.services.publish import delivered_decks
    from raven_ppt.services.publish.deliver import published_digests

    destination = tmp_path / "handoff" / "intro.pptx"
    copy = _delivered(project, destination)

    assert copy.path == destination and destination.read_bytes() == b"PK\x03\x04 deck"
    assert copy.preview == destination.with_suffix(".pdf") and copy.preview.read_bytes().startswith(b"%PDF")
    assert copy.pages == 7
    assert delivered_decks(project.state_dir) == [destination]
    assert published_digests(project.state_dir) == {copy.digest}
    assert not list(destination.parent.glob(".*tmp"))


def test_a_pdf_the_user_already_has_is_not_replaced_by_the_preview(project: Project, tmp_path: Path) -> None:
    """The deliverable is the user's request; the preview beside it is our convenience.

    `deliver` wrote `<stem>.pdf` unconditionally, and that path went through none of
    the checks the `.pptx` did: a delivery of `intro.pptx` into a directory holding
    the user's own `intro.pdf` replaced their file, and `Delivered.preview` then named
    it as this deck's preview. Nothing on the record says we wrote an existing `.pdf`,
    and this module's rule is that nothing here deletes a file the user has -- so the
    name being taken is enough, and the reply says where the preview stayed.
    """
    destination = tmp_path / "deliverables" / "intro.pptx"
    destination.parent.mkdir(parents=True)
    theirs = destination.with_suffix(".pdf")
    theirs.write_bytes(b"%PDF-1.7 the user's own quarterly report")

    copy = _delivered(project, destination)

    assert theirs.read_bytes() == b"%PDF-1.7 the user's own quarterly report", "their file was replaced"
    assert copy.path == destination and destination.read_bytes() == b"PK\x03\x04 deck", "the deck still delivered"
    assert copy.preview is None, "preview names what this delivery wrote, and it wrote none"
    assert str(theirs) in copy.preview_kept_back and "already holds a file" in copy.preview_kept_back
    assert str(project.exports_dir / "deck.pdf") in copy.preview_kept_back, "and where it stayed instead"
    assert not list(destination.parent.glob(".*tmp"))


def test_the_preview_still_rides_along_where_that_name_is_free(project: Project, tmp_path: Path) -> None:
    """The control: the same call into a directory with no such file writes it."""
    destination = tmp_path / "empty" / "intro.pptx"

    copy = _delivered(project, destination)

    assert copy.preview == destination.with_suffix(".pdf")
    assert copy.preview.read_bytes().startswith(b"%PDF")
    assert copy.preview_kept_back == ""


def test_a_later_publish_overwrites_the_delivery(project: Project, tmp_path: Path) -> None:
    from raven_ppt.services.publish import delivered_decks

    destination = tmp_path / "handoff" / "intro.pptx"
    _delivered(project, destination, b"PK\x03\x04 first")
    _delivered(project, destination, b"PK\x03\x04 second")

    assert destination.read_bytes() == b"PK\x03\x04 second"
    assert delivered_decks(project.state_dir) == [destination], "one record per path, the latest digest"


def test_a_delivery_of_bytes_the_gates_never_saw_is_refused(project: Project, tmp_path: Path) -> None:
    from raven_ppt.services.publish import DeliveryError, deliver

    staged = stage(project, _deck(project), pages=1)
    out = publish(project, staged, project.exports_dir / "deck.pptx", findings=[], blocking_kinds=frozenset())
    out.write_bytes(b"PK\x03\x04 swapped after the checks")
    destination = tmp_path / "handoff" / "intro.pptx"
    with pytest.raises(DeliveryError, match="changed before it could be copied"):
        deliver(project, out, staged.digest, staged.pages, destination=destination, preview=None)
    assert not destination.exists()


def test_a_preview_that_cannot_lose_its_pages_is_not_offered(project: Project, monkeypatch) -> None:
    """`pdf_without_pages` is what keeps the preview and the delivered deck the same
    length when the build cap leaves the runner's stand-in pages out. It refuses rather
    than falling back to the whole render: a preview one page longer than the file it
    sits beside previews a deck nobody has, and the caller drops it instead."""
    from raven_ppt.services.publish import pdf_without_pages

    monkeypatch.setattr("raven_ppt.services.render.capabilities.pdfium", lambda: None)
    source = project.build_dir / "render.pdf"
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_bytes(b"%PDF-1.4 not really")
    target = project.build_dir / "short.pdf"

    assert pdf_without_pages(source, [2], target) is False
    assert not target.exists(), "a preview was written from a render it could not trim"


def test_the_pages_left_out_are_counted_from_the_deck_not_from_each_other(project: Project) -> None:
    """Removing page 2 makes the old page 3 the new page 2, so a loop that removes in
    ascending order takes out the wrong second page. They come out from the back."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    from raven_ppt.services.publish import without_pages

    presentation = Presentation()
    for number in range(1, 6):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = f"page {number}"
    built = project.build_dir / "deck.pptx"
    built.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(built))
    target = project.build_dir / "released.pptx"

    assert without_pages(built, [2, 4, 99], target) == 3
    kept = Presentation(str(target))
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in kept.slides]
    assert texts == ["page 1", "page 3", "page 5"], texts
    assert len(Presentation(str(built)).slides) == 5, "the built deck is not the one edited"
