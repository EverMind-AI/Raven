"""The order of build, measure, deliver -- and why it is that order."""

from __future__ import annotations

import json
import os
import textwrap
from pathlib import Path

import pytest

from raven_ppt.backends.script import script_path
from raven_ppt.contracts import BuildOutcome, Finding, PageSource, Project, Severity
from raven_ppt.profiles import registry
from raven_ppt.stages.build import BuildStage
from tests._ppt_engine_fixtures import COMMENT_ABOVE_FIRST_BANNER

SCRIPT = textwrap.dedent(
    """
    import os
    from pptx import Presentation

    prs = Presentation()


    def new_slide():
        return prs.slides.add_slide(prs.slide_layouts[6])


    # SLIDE 1
    one = new_slide()

    # SLIDE 2
    two = new_slide()

    prs.save(os.environ["PPT_OUTPUT"])
    """
).lstrip()


@pytest.fixture()
def project(tmp_path: Path) -> Project:
    p = Project(workspace=tmp_path, slug="tarvis")
    p.build_dir.mkdir(parents=True)
    p.state_dir.mkdir(parents=True)
    script_path(p).write_text(SCRIPT, encoding="utf-8")
    return p


def test_the_page_mapping_diagnosis_reads_a_marked_file_as_the_unmarked_one(project: Project) -> None:
    """The one door for build.py is only a door if every reader uses it. This is the reader
    the review named: a UTF-8 byte-order mark in front of a comment above the first banner
    put that comment in the prelude instead of page 1's block, and the author was told
    `# SLIDE 2` creates 0 slides where the unmarked file says `# SLIDE 1` creates 2. Same
    file, same answer, whichever way an editor saved it."""
    import codecs

    from raven_ppt.stages.build import _mapping_findings

    project.build_dir.mkdir(parents=True, exist_ok=True)
    unmapped = BuildOutcome(ok=True, pages=2, sources=())

    script_path(project).write_text(COMMENT_ABOVE_FIRST_BANNER, encoding="utf-8")
    (plain,) = _mapping_findings(project, unmapped)
    script_path(project).write_bytes(codecs.BOM_UTF8 + COMMENT_ABOVE_FIRST_BANNER.encode("utf-8"))
    (marked,) = _mapping_findings(project, unmapped)

    assert plain.kind == "unmapped_page" and "`# SLIDE 1` creates 2 slides" in plain.message
    assert marked.message == plain.message


def _outcome(project: Project, pages: int = 2, mapped: bool = True) -> BuildOutcome:
    deck = project.build_dir / "deck.pptx"
    deck.write_bytes(b"PK deck")
    lines = SCRIPT.splitlines(keepends=True)
    first = next(i for i, line in enumerate(lines) if "# SLIDE 1" in line)
    second = next(i for i, line in enumerate(lines) if "# SLIDE 2" in line)
    sources = (
        (
            PageSource(page=1, first_line=first, last_line=second),
            PageSource(page=2, first_line=second, last_line=len(lines) - 2),
        )
        if mapped
        else ()
    )
    return BuildOutcome(ok=True, pptx_path=deck, pages=pages, sources=sources, source_digest="x")


def _drawn(project: Project, pages: int = 3) -> BuildOutcome:
    """A real deck on disk, because dropping a page means opening the file."""
    from pptx import Presentation
    from pptx.util import Inches

    outcome = _outcome(project, pages=pages)
    presentation = Presentation()
    for number in range(1, pages + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"page {number}"
    # After _outcome, which writes a stub over this same path.
    presentation.save(str(project.build_dir / "deck.pptx"))
    return outcome


def _measure(findings):
    async def measure(_project: Project, _pptx: Path, _outcome=None, _changed: str = "delivery") -> list[Finding]:
        return list(findings)

    return measure


def _stage(project: Project, *, findings=(), outcome=None, profile="script_author") -> BuildStage:
    measured = list(findings)

    async def backend(_project: Project, _script: str | None) -> BuildOutcome:
        return outcome or _outcome(_project)

    return BuildStage(
        backend=backend,
        measure=_measure(measured),
        profile=registry.get(profile),
    )


def _fact() -> Finding:
    return Finding(kind="fact", severity=Severity.BLOCKING, message="48.3 is not in the sources", page=11)


def _warning(kind: str = "type_floor") -> Finding:
    return Finding(kind=kind, severity=Severity.WARNING, message="11.5pt body", page=3)


@pytest.mark.asyncio
async def test_a_clean_deck_is_measured_and_delivered(project: Project) -> None:
    result = await _stage(project).run(project)

    assert result.ok, result.note
    assert Path(result.data["pptx_path"]).is_file()


@pytest.mark.asyncio
async def test_a_blocking_finding_refuses_the_export_whatever_it_is_about(project: Project) -> None:
    """Both kinds of refusal, because they used to be two different axes.

    A claim the materials do not make and a word painted over another word were once
    sorted into separate lists, and one waited on the other: a deck with 29 unreplaced
    placeholders also had 25 pairs of overlapping words, and nothing looked at the
    second list until the first cleared. There is one list now, and either entry on it
    stops the file being written.
    """
    collision = Finding(
        kind="word_collision",
        severity=Severity.BLOCKING,
        message="'31.1' is painted over 'VPS-only'",
        page=16,
    )
    for finding in (_fact(), collision):
        result = await _stage(project, findings=[finding]).run(project)

        assert not result.ok, finding.kind
        assert not (project.exports_dir / "deck.pptx").exists()
        # Said, not implied: the tool's reply carries this note as `not_delivered`, and
        # a live run got the bare fallback "the deck was not published" with no reason.
        assert result.note and result.note.startswith("not published: 1 blocking finding(s)"), result.note


@pytest.mark.asyncio
async def test_a_refused_build_leaves_its_reason_where_the_hook_reads_it(project: Project) -> None:
    """The reason reaches the model once, in the tool reply; the hook that sends a
    "delivered" reply about a copied build back needs it again, from the directory. A
    publish clears it: what was refused before no longer stands in the way."""
    from raven_ppt.services.publish.deliver import last_refusal

    refused = await _stage(project, findings=[_fact()]).run(project)
    assert not refused.ok
    reason = last_refusal(project.state_dir)
    assert reason and reason.startswith(refused.note), reason
    assert "page 11: fact -- 48.3 is not in the sources" in reason

    delivered = await _stage(project).run(project)
    assert delivered.ok, delivered.note
    assert last_refusal(project.state_dir) is None


@pytest.mark.asyncio
async def test_a_warning_never_stops_the_deck(project: Project) -> None:
    """Shrinking the copy would satisfy the measurement and make the page worse."""
    result = await _stage(project, findings=[_warning()]).run(project)

    assert result.ok
    assert [f.kind for f in result.findings] == ["type_floor"]


@pytest.mark.asyncio
async def test_the_deck_is_measured_once(project: Project) -> None:
    """Measuring renders the whole deck, so a second pass over it costs minutes.

    It was measured twice, because a stage between the two rewrote the pages it had
    just looked at. Nothing rewrites them now, and the second render bought nothing.
    """
    seen: list[int] = []

    async def backend(_project: Project, _script: str | None) -> BuildOutcome:
        return _outcome(_project)

    async def measure(_p: Project, _pptx: Path, _outcome=None, _changed: str = "delivery") -> list[Finding]:
        seen.append(1)
        return []

    stage = BuildStage(backend=backend, measure=measure, profile=registry.get("script_author"))
    await stage.run(project)
    assert len(seen) == 1


@pytest.mark.asyncio
async def test_pages_that_cannot_be_told_apart_are_blocking_and_carry_a_skeleton(project: Project) -> None:
    """Refused without one, an author rewrote the same shim three times and left."""
    looped = textwrap.dedent(
        """
        import os
        from pptx import Presentation
        prs = Presentation()
        for _ in range(5):
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(os.environ["PPT_OUTPUT"])
        """
    ).lstrip()
    script_path(project).write_text(looped, encoding="utf-8")
    result = await _stage(project, outcome=_outcome(project, pages=5, mapped=False)).run(project)

    assert not result.ok
    unmapped = [f for f in result.findings if f.kind == "unmapped_page"]
    assert "# SLIDE 1" in unmapped[0].message


@pytest.mark.asyncio
async def test_a_mapped_deck_produces_no_mapping_finding(project: Project) -> None:
    result = await _stage(project).run(project)
    assert not any(f.kind == "unmapped_page" for f in result.findings)


@pytest.mark.asyncio
async def test_a_build_that_failed_stops_before_anything_else(project: Project) -> None:
    result = await _stage(project, outcome=BuildOutcome(ok=False, stderr="boom")).run(project)

    assert not result.ok
    assert result.data["outcome"].stderr == "boom"


@pytest.mark.asyncio
async def test_the_route_decides_which_warning_kinds_refuse_delivery(project: Project) -> None:
    """A page that traces to no block of its own code, so no render can be matched to it.

    The example was `band` until the gate was downgraded to a report and this test was
    one of the places still holding it fatal.
    """
    orphan = Finding(kind="unmapped_page", severity=Severity.WARNING, message="page 4 maps to no block", page=4)
    result = await _stage(project, findings=[orphan]).run(project)
    assert not result.ok

    other = Finding(kind="wrapped_label", severity=Severity.WARNING, message="a label wrapped", page=4)
    assert (await _stage(project, findings=[other]).run(project)).ok


@pytest.mark.asyncio
async def test_a_draft_is_measured_and_not_published(project: Project) -> None:
    """A half-written program is an accepted intermediate state: it is measured and
    handed back, and it is not held to a length nobody has finished writing to."""
    result = await _stage(project, findings=[_warning("page_budget")]).run(project, draft=True)

    assert result.ok and result.data["draft"] is True
    assert "pptx_path" not in result.data
    assert not (project.exports_dir / "deck.pptx").exists()
    assert [f.kind for f in result.findings] == [], "a draft is not held to the agreed length"


def _planned(project: Project, figures: dict[int, list[str]]) -> None:
    """An outline that gives these pages these figures."""
    import json

    from raven_ppt.contracts import outline_path

    outline_path(project).parent.mkdir(parents=True, exist_ok=True)
    outline_path(project).write_text(
        json.dumps(
            {
                "pages": [
                    {
                        "page": n,
                        "claim": f"page {n}",
                        "carries": "prose",
                        "says": ["a", "b"],
                        "figures": figures.get(n, []),
                    }
                    for n in (1, 2)
                ]
            }
        ),
        encoding="utf-8",
    )


def _deck_with(path: Path, pictures: dict[int, bool], layouts: tuple[int, int] = (6, 6)) -> None:
    """Two slides, each with a picture or without one, and each on the named layout."""
    from pptx import Presentation
    from pptx.util import Inches

    png = path.parent / "dot.png"
    png.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d494844520000000100000001080600000"
            "01f15c4890000000a49444154789c6300010000050001-0d0a2db40000000049454e44ae426082".replace("-", "")
        )
    )
    prs = Presentation()
    for n in (1, 2):
        slide = prs.slides.add_slide(prs.slide_layouts[layouts[n - 1]])
        if pictures.get(n):
            slide.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(1), Inches(1))
    prs.save(path)


def test_a_page_the_plan_gave_a_figure_and_the_build_left_bare_is_reported(project: Project) -> None:
    """Twenty pages were built carrying no picture while the plan named two."""
    from raven_ppt.stages.build import _unplaced_figure_findings

    deck = project.build_dir / "deck.pptx"
    _deck_with(deck, {1: False, 2: True})
    _planned(project, {1: ["mem0-logo-2afc7b9e29"], 2: ["arch-7979644ab7"]})
    found = _unplaced_figure_findings(project, BuildOutcome(ok=True, pptx_path=deck, pages=2, source_digest="x"))

    assert [f.page for f in found] == [1], "only the page built without one"
    assert "mem0-logo-2afc7b9e29" in found[0].message


def _onto_layout(deck: Path, image: Path) -> None:
    """Put `image` on the layout of the deck's first page, where template cover art lives."""
    from pptx import Presentation
    from pptx.oxml.shapes.picture import CT_Picture
    from pptx.util import Inches

    presentation = Presentation(str(deck))
    layout = presentation.slides[0].slide_layout
    _, rId = layout.part.get_or_add_image_part(str(image))
    pic = CT_Picture.new_pic(2, image.stem, image.name, rId, Inches(1), Inches(1), Inches(2), Inches(2))
    layout.shapes._spTree.append(pic)
    presentation.save(str(deck))


def test_a_figure_swapped_into_the_layout_counts_as_placed(project: Project) -> None:
    """Several bundled templates draw their cover art on the layout, and the helper an
    author writes swaps it there. Reading the slide alone called such a page bare: a
    live run sat on this finding across several builds while page 1's layout held the
    ingested cover art byte for byte and the slide held no picture at all."""
    from raven_ppt.stages.build import _unplaced_figure_findings

    deck = project.build_dir / "deck.pptx"
    # Each page on its own layout: pages sharing one would share the swap, and the
    # second page is here to show a page whose layout carries nothing is still bare.
    _deck_with(deck, {1: False, 2: False}, layouts=(6, 5))
    figure = project.figures_dir / "cover-2afc7b9e29.png"
    figure.parent.mkdir(parents=True, exist_ok=True)
    figure.write_bytes((project.build_dir / "dot.png").read_bytes())

    _onto_layout(deck, figure)
    _planned(project, {1: ["cover-2afc7b9e29"], 2: ["arch-7979644ab7"]})

    found = _unplaced_figure_findings(project, BuildOutcome(ok=True, pptx_path=deck, pages=2, source_digest="x"))

    assert [f.page for f in found] == [2], "page 1 has its figure, on the shape that draws it"


def test_a_template_picture_on_the_layout_does_not_answer_a_promised_figure(project: Project) -> None:
    """Only the deck's own figures count on a layout. Any layout picture would let a
    template's own photograph answer every figure a plan promised, which is the check
    saying nothing at all."""
    from raven_ppt.stages.build import _unplaced_figure_findings

    deck = project.build_dir / "deck.pptx"
    _deck_with(deck, {1: False, 2: False})
    project.figures_dir.mkdir(parents=True, exist_ok=True)
    (project.figures_dir / "other-9de1a0.png").write_bytes(b"\x89PNG\r\n\x1a\n not the picture on the layout")

    _onto_layout(deck, project.build_dir / "dot.png")
    _planned(project, {1: ["cover-2afc7b9e29"]})

    found = _unplaced_figure_findings(project, BuildOutcome(ok=True, pptx_path=deck, pages=2, source_digest="x"))

    assert [f.page for f in found] == [1]


class _LayoutPicture:
    """A picture on the page's layout, as `picture_blob` reads one."""

    def __init__(self, blob: bytes) -> None:
        self.image = type("Image", (), {"blob": blob})()


def test_a_figure_swapped_onto_the_layout_counts_as_placed(project: Project, monkeypatch) -> None:
    """Where this check was a false positive. Several templates keep the cover's, the
    section page's and the closing page's picture on the *layout*, and `replace_picture`
    there is then the only place a page can put one -- `pictures={...}` on the cloned
    page never reaches it. A live run spent six builds on its cover with the layout
    holding the deck's own cover art byte for byte, while this said the page had been
    built without a picture on it.
    """
    from raven_ppt.stages import build as build_stage
    from raven_ppt.stages.build import _unplaced_figure_findings

    deck = project.build_dir / "deck.pptx"
    _deck_with(deck, {1: False, 2: False})
    _planned(project, {1: ["cover-art-1a2b3c"]})
    project.figures_dir.mkdir(parents=True, exist_ok=True)
    mine = b"\x89PNG the cover this deck ingested"
    (project.figures_dir / "cover-art-1a2b3c.png").write_bytes(mine)
    outcome = BuildOutcome(ok=True, pptx_path=deck, pages=2, source_digest="x")

    monkeypatch.setattr(build_stage, "_layout_pictures", lambda slide: [_LayoutPicture(mine)])

    assert _unplaced_figure_findings(project, outcome) == [], "the figure is placed, on the only slot there is"


def test_a_template_picture_on_the_layout_does_not_answer_for_a_promised_figure(project: Project, monkeypatch):
    """The other half, and why the first is by bytes. If any layout picture counted, a
    template's stock photograph would satisfy every figure every page ever promised --
    which is the failure this check exists to catch. It still refuses, and it now says
    what the author is looking at, because a render that plainly shows a picture makes
    "built without a picture on it" read as false."""
    from raven_ppt.stages import build as build_stage
    from raven_ppt.stages.build import _unplaced_figure_findings

    deck = project.build_dir / "deck.pptx"
    _deck_with(deck, {1: False, 2: False})
    _planned(project, {1: ["cover-art-1a2b3c"]})
    project.figures_dir.mkdir(parents=True, exist_ok=True)
    (project.figures_dir / "cover-art-1a2b3c.png").write_bytes(b"\x89PNG the cover this deck ingested")
    outcome = BuildOutcome(ok=True, pptx_path=deck, pages=2, source_digest="x")

    monkeypatch.setattr(build_stage, "_layout_pictures", lambda slide: [_LayoutPicture(b"the template own photo")])
    found = _unplaced_figure_findings(project, outcome)

    assert [f.page for f in found] == [1], "a picture that is not this deck's answers for nothing"
    assert "The picture you can see on it is the template layout's" in found[0].message
    assert "cover-art-1a2b3c" in found[0].message, "and it still names what was promised"

    monkeypatch.setattr(build_stage, "_layout_pictures", lambda slide: [])
    bare = _unplaced_figure_findings(project, outcome)
    assert "the template layout's" not in bare[0].message, "a page showing nothing is not told it shows something"


def test_a_plan_that_promises_no_figure_is_not_asked_for_one(project: Project) -> None:
    from raven_ppt.stages.build import _unplaced_figure_findings

    deck = project.build_dir / "deck.pptx"
    _deck_with(deck, {1: False, 2: False})
    _planned(project, {})
    assert _unplaced_figure_findings(project, BuildOutcome(ok=True, pptx_path=deck, pages=2, source_digest="x")) == []


def _housed_deck(path: Path, content_pages: int, *, title_left: float = 0.72, wander: int = 0) -> None:
    """A cover and `content_pages` pages that put their title in the same box.

    `wander` moves that many of them somewhere else, which is how a deck that has not
    settled on a title row is built. The body copy sits below the title band and names
    its own size, so the record has both a row to report and a size to report.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    cover.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1)).text_frame.text = "A deck about deck houses"
    for number in range(content_pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        left = title_left if number >= wander else title_left + 1.5 + number * 0.3
        title = slide.shapes.add_textbox(Inches(left), Inches(0.4), Inches(8.5), Inches(0.6))
        run = title.text_frame.paragraphs[0].add_run()
        run.text = f"What page {number + 2} argues"
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.name = "Arial"
        body = slide.shapes.add_textbox(Inches(title_left), Inches(3), Inches(8.5), Inches(3))
        line = body.text_frame.paragraphs[0].add_run()
        line.text = "A sentence long enough that nothing reads it as a divider or a label."
        line.font.size = Pt(18)
        line.font.name = "Arial"
    prs.save(path)


def _house_of(project: Project, path: Path, pages: int) -> str:
    from raven_ppt.stages.build import _deck_house

    return _deck_house(project, BuildOutcome(ok=True, pptx_path=path, pages=pages, source_digest="x"))


def test_a_row_only_half_the_pages_use_is_not_named_as_the_decks(project: Project) -> None:
    """Half is not most, and half is where the measurement was measurably wrong.

    Three of six pages sharing a box named a right-hand panel on a real deck, and four
    of nine did on the same one -- a tie broken by insertion order. The record says the
    pages have not settled rather than naming a box they did not settle on, and rather
    than saying nothing: a deck whose pages disagree about where the title goes is a
    deck to look at.
    """
    deck = project.build_dir / "deck.pptx"
    _housed_deck(deck, 6, wander=3)
    said = _house_of(project, deck, 7)

    assert "have not settled on one" in said, said
    assert "0.60in" not in said, "the box itself is not named when it is not the deck's"


def test_a_deck_whose_pages_agree_carries_what_they_agreed(project: Project) -> None:
    deck = project.build_dir / "deck.pptx"
    _housed_deck(deck, 9)
    said = _house_of(project, deck, 10)

    assert "the title row 9 of them agree on" in said, said
    assert "(0.72, 0.40) 8.50x0.60in" in said, said
    assert "24pt bold" in said, said
    assert "18pt on 9 pages" in said, said
    # A record, not a finding: nothing in it asks for anything or refuses anything.
    assert "Nothing here is a finding" in said


def test_pages_that_have_not_settled_on_a_title_row_are_not_told_they_have(project: Project) -> None:
    """Nine content pages of which three share a box: a plurality, and not the deck's.

    This is the shape the relation exists for. On a real deck, three of six pages
    sharing a box was a right-hand panel rather than the title row those pages actually
    use, and naming it would have sent the author to move the titles. Three of nine is
    the same plurality at a different length, which is why the test is a comparison and
    not a count.
    """
    deck = project.build_dir / "deck.pptx"
    _housed_deck(deck, 9, wander=6)
    said = _house_of(project, deck, 10)

    assert "3 of them put a title row in the same place" in said, said
    assert "have not settled on one" in said, said
    assert "agree on is" not in said, "a plurality is not an agreement"


def test_a_deck_built_in_a_template_reports_no_house_of_its_own(project: Project) -> None:
    """The house is the template's then, and `title_row` already reports the drift per
    page against it. A second, deck-level house would read as permission to keep it."""
    from raven_ppt.services.template import prepared_path

    deck = project.build_dir / "deck.pptx"
    _housed_deck(deck, 9)
    assert _house_of(project, deck, 10) != ""

    prepared_path(project).parent.mkdir(parents=True, exist_ok=True)
    _housed_deck(prepared_path(project), 2)
    assert _house_of(project, deck, 10) == ""


def test_the_record_travels_on_the_note_the_reply_already_carries(project: Project) -> None:
    """`ppt_build` reads `outcome.note` into the reply; nothing else the stage returns
    reaches the author as prose while the build succeeded."""
    from raven_ppt.stages.build import _with_deck_house

    deck = project.build_dir / "deck.pptx"
    _housed_deck(deck, 9)
    outcome = BuildOutcome(ok=True, pptx_path=deck, pages=10, note="the script held no program", source_digest="x")
    carried = _with_deck_house(project, outcome)

    assert carried.note.startswith("the script held no program\n"), "the harness's own note is not displaced"
    assert "this deck's own house" in carried.note
    assert carried.pptx_path == outcome.pptx_path and carried.pages == outcome.pages


def test_an_unreadable_deck_costs_the_record_and_not_the_build(project: Project) -> None:
    from raven_ppt.stages.build import _with_deck_house

    deck = project.build_dir / "deck.pptx"
    deck.write_bytes(b"not a deck")
    outcome = BuildOutcome(ok=True, pptx_path=deck, pages=10, source_digest="x")
    assert _with_deck_house(project, outcome) == outcome


@pytest.mark.asyncio
async def test_a_built_deck_hands_back_its_own_house(project: Project) -> None:
    """The whole way through `run`, because the record is only worth anything in a reply.

    `ppt_build` reads `outcome` out of the stage's data and its note into the payload,
    so the record has to be on the outcome the stage returns rather than beside it. Ten
    pages and three renders is also the deck that is refused for the seven nobody has
    looked at, which is the round the record matters in: the author builds again.
    """
    deck = project.build_dir / "deck.pptx"
    _housed_deck(deck, 9)
    outcome = BuildOutcome(
        ok=True,
        pptx_path=deck,
        pages=10,
        sources=tuple(PageSource(page=n, first_line=n, last_line=n + 1) for n in range(1, 11)),
        source_digest="x",
    )
    result = await _stage(project, outcome=outcome).run(project)

    assert [f.kind for f in result.findings] == ["unseen_page"]
    assert "this deck's own house" in result.data["outcome"].note


@pytest.mark.asyncio
async def test_a_draft_hands_back_its_house_too(project: Project) -> None:
    """A half-written program is exactly when knowing what the written half decided is
    worth something, so the record is taken before the draft returns."""
    deck = project.build_dir / "deck.pptx"
    _housed_deck(deck, 9)
    outcome = BuildOutcome(ok=True, pptx_path=deck, pages=10, source_digest="x")
    result = await _stage(project, outcome=outcome).run(project, draft=True)

    assert result.data["draft"] is True
    assert "this deck's own house" in result.data["outcome"].note


def _ten_pages(project: Project) -> BuildOutcome:
    deck = project.build_dir / "deck.pptx"
    deck.write_bytes(b"PK")
    return BuildOutcome(
        ok=True,
        pptx_path=deck,
        pages=10,
        sources=tuple(PageSource(page=n, first_line=n, last_line=n + 1) for n in range(1, 11)),
        source_digest="x",
    )


@pytest.mark.asyncio
async def test_a_draft_render_counts_as_a_page_seen(project: Project) -> None:
    """The record is about pages put in front of the author, and a draft render is that.

    Recording only on publication made the two halves disagree, and the disagreement
    is unwinnable rather than merely wasteful: a measured run walked its nineteen
    pages three times as drafts, was told on the first real build that eighteen had
    never been seen, and spent the rest of its budget walking them again one non-draft
    build at a time. It reasoned its way to the cause and said so -- "draft renders do
    not count towards the seen record" -- which is a fact about this code and not
    about the deck.
    """
    from raven_ppt.services import seen

    outcome = _ten_pages(project)
    stage = _stage(project, outcome=outcome)

    await stage.run(project, slides=[1, 2, 3], draft=True)

    recorded = set(json.loads(seen.seen_path(project).read_text())["pages"])
    assert recorded == {"1", "2", "3"}, "a draft render is a page seen"

    # And the deck built straight afterwards does not ask for them again.
    result = await stage.run(project, slides=[1, 2, 3])
    unseen = [f for f in result.findings if f.kind == "unseen_page"]
    assert unseen, "the seven nobody looked at still refuse"
    assert unseen[0].detail["pages"] == [4, 5, 6, 7, 8, 9, 10]


@pytest.mark.asyncio
async def test_a_draft_walk_leaves_nothing_for_the_real_build_to_refuse(project: Project) -> None:
    """The whole point, as a run actually does it: draft through the deck, then publish.

    Before this, the same walk left the record empty and the publishing build refused
    every page -- so the author had to repeat the entire walk with `draft` off.
    """
    outcome = _ten_pages(project)
    stage = _stage(project, outcome=outcome)

    for first in (1, 4, 7, 10):
        await stage.run(project, page_from=first, draft=True)

    result = await stage.run(project, page_from=1)

    assert [f.kind for f in result.findings if f.kind == "unseen_page"] == []


@pytest.mark.asyncio
async def test_a_render_that_did_not_happen_is_not_a_page_seen(project: Project) -> None:
    """The guard that survives the move. A page whose render failed was never put in
    front of anyone, and recording it would say the author saw code that drew nothing."""
    from raven_ppt.services import seen

    outcome = _ten_pages(project)
    stage = _stage(project, outcome=outcome, findings=[_warning("unrendered")])

    await stage.run(project, slides=[1, 2, 3], draft=True)

    assert not seen.seen_path(project).exists()


@pytest.mark.asyncio
async def test_a_build_without_slides_shows_the_unseen_pages_first(project: Project) -> None:
    """The tail this removes: a run reached twenty pages and then spent ten builds on
    nothing but looking, because seven pages had changed since they were shown and a
    walk from page 1 handed back three at a time, seen pages included. The pages that
    refuse publication are the ones a plain build comes back with, and more of them
    than a walk carries."""
    from raven_ppt.stages.build import CATCH_UP_VIEWS

    outcome = _ten_pages(project)
    stage = _stage(project, outcome=outcome)
    await stage.run(project, slides=[1, 2, 3], draft=True)

    result = await stage.run(project)

    assert result.data["showing"] == [4, 5, 6, 7, 8, 9], "the unseen pages, up to the catch-up batch"
    assert CATCH_UP_VIEWS == 6
    unseen = [f for f in result.findings if f.kind == "unseen_page"]
    assert unseen[0].detail["pages"] == [10] == result.data["unseen_after"]
    assert "without `slides`" in unseen[0].message

    delivered = await stage.run(project)
    assert delivered.data["showing"] == [10]
    assert delivered.ok and delivered.data["unseen_after"] == []


@pytest.mark.asyncio
async def test_page_from_says_where_the_catch_up_starts(project: Project) -> None:
    """A number past every unseen page does not lose them: the batch wraps to the first."""
    stage = _stage(project, outcome=_ten_pages(project))

    late = await stage.run(project, page_from=8, draft=True)
    assert late.data["showing"] == [8, 9, 10]

    wrapped = await stage.run(project, page_from=11, draft=True)
    assert wrapped.data["showing"] == [1, 2, 3, 4, 5, 6]


@pytest.mark.asyncio
async def test_a_deck_that_has_been_seen_whole_is_walked_by_number(project: Project) -> None:
    """Once nothing is unseen the plain build is the old walk: three from page_from."""
    stage = _stage(project, outcome=_ten_pages(project))
    await stage.run(project, page_from=1, draft=True)
    await stage.run(project, page_from=7, draft=True)

    result = await stage.run(project, page_from=4)

    assert result.data["showing"] == [4, 5, 6]
    assert result.ok


@pytest.mark.asyncio
async def test_a_named_page_is_still_the_one_that_comes_back(project: Project) -> None:
    stage = _stage(project, outcome=_ten_pages(project))

    result = await stage.run(project, slides=[7])

    assert result.data["showing"] == [7]
    assert result.data["unseen_after"] == [1, 2, 3, 4, 5, 6, 8, 9, 10]


@pytest.mark.asyncio
async def test_a_deck_without_a_mapping_keeps_no_unseen_record(project: Project) -> None:
    stage = _stage(project, outcome=_outcome(project, mapped=False))

    result = await stage.run(project)

    assert result.data["showing"] == [1, 2]
    assert result.data["unseen_after"] is None


@pytest.mark.asyncio
async def test_a_page_the_runner_stood_in_for_refuses_the_deck_and_names_the_error(project: Project) -> None:
    from raven_ppt.backends.script.workspace import page_failures_path

    page_failures_path(project).parent.mkdir(parents=True, exist_ok=True)
    page_failures_path(project).write_text(
        json.dumps(
            {"pages": [{"page": 2, "error": "NameError: name 'plane' is not defined", "traceback": "Traceback..."}]}
        ),
        encoding="utf-8",
    )
    stage = _stage(project, outcome=_outcome(project))

    result = await stage.run(project, slides=[1, 2])

    failed = [f for f in result.findings if f.kind == "page_failed"]
    assert not result.ok
    assert [f.page for f in failed] == [2]
    assert "NameError: name 'plane' is not defined" in failed[0].message
    assert failed[0].detail["traceback"] == "Traceback..."


@pytest.mark.asyncio
async def test_a_delivered_deck_takes_its_render_along_as_a_pdf(project: Project) -> None:
    """The web surface previews PDFs and not .pptx, so a user there could not look at a
    delivered deck. The render the measurement made is copied beside the deck; a render
    older than the deck is not, because it would preview a file that no longer exists."""
    import os
    import time

    project.review_dir.mkdir(parents=True, exist_ok=True)
    outcome = _outcome(project)
    rendered = project.review_dir / "deck.pdf"
    rendered.write_bytes(b"%PDF-1.4 render")
    later = time.time() + 5
    os.utime(rendered, (later, later))

    result = await _stage(project, outcome=outcome).run(project)

    assert result.ok, result.note
    pdf = Path(result.data["pptx_path"]).with_suffix(".pdf")
    assert pdf.is_file() and pdf.read_bytes() == b"%PDF-1.4 render"
    assert result.data["pdf_path"] == str(pdf)

    stale = time.time() - 600
    os.utime(rendered, (stale, stale))
    pdf.unlink()
    again = await _stage(project, outcome=_outcome(project)).run(project)
    assert again.ok and "pdf_path" not in again.data and not pdf.exists()


@pytest.mark.asyncio
async def test_a_second_delivery_is_marked_as_a_republish(project: Project) -> None:
    first = await _stage(project, outcome=_outcome(project)).run(project)
    assert first.ok and "republished" not in first.data

    again = await _stage(project, outcome=_outcome(project)).run(project)
    assert again.ok and again.data["republished"] is True


@pytest.mark.asyncio
async def test_a_released_build_publishes_past_its_blocking_findings_and_names_them(project: Project) -> None:
    """The capped tiers' last word: at the build cap the deck goes out as it stands.
    The findings are not lost -- they ride in the result for the reply to list -- but
    they no longer hold the file back."""
    collision = Finding(
        kind="word_collision", severity=Severity.BLOCKING, message="'31.1' is painted over 'VPS-only'", page=16
    )
    result = await _stage(project, findings=[collision]).run(project, release=True)

    assert result.ok, result.note
    assert Path(result.data["pptx_path"]).is_file()
    assert result.data["released"] == ["word_collision"]
    assert any(f.kind == "word_collision" for f in result.findings), "released, not forgotten"


@pytest.mark.asyncio
async def test_a_delivery_edited_in_place_is_reported_by_the_build_that_replaces_it(project: Project) -> None:
    """Asked before the publish, because afterwards there is nothing left to compare.

    A live run's author edited `out/deck.pptx` where it lay through `exec`; the record
    and the gates went on describing the build, and the deck the user held had been
    through neither. The repair is the publish itself -- the measured deck written over
    the edited one -- so this reads and never withholds anything.
    """
    stage = _stage(project)
    first = await stage.run(project)
    assert "delivery_changed" not in first.data

    delivered = Path(first.data["pptx_path"])
    delivered.write_bytes(delivered.read_bytes() + b"EDITED-IN-PLACE")

    second = await stage.run(project)

    assert "was not the deck this route delivered" in second.data["delivery_changed"]
    assert "never in the delivered file" in second.data["delivery_changed"]
    assert Path(second.data["pptx_path"]).read_bytes() == (project.build_dir / "deck.pptx").read_bytes()
    # And it is said once, by the build that repaired it, not on every build after.
    assert "delivery_changed" not in (await stage.run(project)).data


# -- the destination the user named ------------------------------------------


def _built(project: Project, body: bytes, pages: int = 2) -> BuildOutcome:
    deck = project.build_dir / "deck.pptx"
    deck.write_bytes(body)
    lines = SCRIPT.splitlines(keepends=True)
    first = next(i for i, line in enumerate(lines) if "# SLIDE 1" in line)
    second = next(i for i, line in enumerate(lines) if "# SLIDE 2" in line)
    sources = (
        PageSource(page=1, first_line=first, last_line=second),
        PageSource(page=2, first_line=second, last_line=len(lines) - 2),
    )
    return BuildOutcome(ok=True, pptx_path=deck, pages=pages, sources=sources, source_digest="x")


def _stated(project: Project, tmp_path: Path) -> Path:
    from raven_ppt.services.publish import as_destination, write_destination

    wanted = tmp_path / "handoff" / "ravenx-intro.pptx"
    return write_destination(project, as_destination(str(wanted), project, default_name="deck.pptx"))


@pytest.mark.asyncio
async def test_every_publish_writes_the_stated_destination_as_well(project: Project, tmp_path: Path) -> None:
    """Stated once; the first delivery and every revision after land on it."""
    from raven_ppt.services.publish import delivered_decks

    destination = _stated(project, tmp_path)

    first = await _stage(project, outcome=_built(project, b"PK first")).run(project)
    assert first.ok, first.note
    assert first.data["delivered_to"] == str(destination)
    assert destination.read_bytes() == b"PK first"
    assert (project.exports_dir / "deck.pptx").read_bytes() == b"PK first", "out/ is written as before"

    second = await _stage(project, outcome=_built(project, b"PK second")).run(project)
    assert second.ok and second.data.get("republished") is True
    assert destination.read_bytes() == b"PK second"
    assert delivered_decks(project.state_dir) == [destination]
    record = json.loads((project.state_dir / "published.json").read_text(encoding="utf-8"))["published"]
    assert [entry.get("role") for entry in record] == [None, "delivery"]
    assert {entry["sha256"] for entry in record} == {record[0]["sha256"]}, "one digest, two paths"


@pytest.mark.asyncio
async def test_a_refused_build_leaves_the_delivery_as_it_was(project: Project, tmp_path: Path) -> None:
    """The destination holds the last deck that passed, never a half-built one."""
    destination = _stated(project, tmp_path)
    good = await _stage(project, outcome=_built(project, b"PK good")).run(project)
    assert good.ok, good.note

    refused = await _stage(project, findings=[_fact()], outcome=_built(project, b"PK broken")).run(project)
    assert not refused.ok and "delivered_to" not in refused.data
    assert destination.read_bytes() == b"PK good"

    drafted = await _stage(project, outcome=_built(project, b"PK draft")).run(project, draft=True)
    assert drafted.ok and "delivered_to" not in drafted.data
    assert destination.read_bytes() == b"PK good"


@pytest.mark.asyncio
async def test_a_deck_with_no_stated_destination_is_published_as_before(project: Project) -> None:
    result = await _stage(project).run(project)
    assert result.ok and "delivered_to" not in result.data and "delivery_failed" not in result.data


@pytest.mark.asyncio
async def test_a_destination_that_cannot_be_written_is_said_not_swallowed(project: Project, tmp_path: Path) -> None:
    """The deck is published under out/; that the named path did not get it is the author's to report."""
    from raven_ppt.services.publish import write_destination

    blocker = tmp_path / "handoff"
    blocker.write_text("a file where the directory should be", encoding="utf-8")
    write_destination(project, blocker / "intro.pptx")

    result = await _stage(project, outcome=_built(project, b"PK deck")).run(project)
    assert result.ok, result.note
    assert Path(result.data["pptx_path"]).read_bytes() == b"PK deck"
    assert "delivered_to" not in result.data
    assert str(blocker / "intro.pptx") in result.data["delivery_failed"]


def _stood_in(project: Project, page: int = 2) -> None:
    from raven_ppt.backends.script.workspace import page_failures_path

    page_failures_path(project).parent.mkdir(parents=True, exist_ok=True)
    page_failures_path(project).write_text(
        json.dumps({"pages": [{"page": page, "error": "KeyError: no shape says 'Outlook'", "traceback": "T"}]}),
        encoding="utf-8",
    )


@pytest.mark.asyncio
async def test_the_build_cap_delivers_the_pages_that_drew_without_the_stand_ins(project: Project) -> None:
    """The cap is the one thing that stops buying quality: past it a deck goes out with
    whatever the gates still say, because a deck that exists beats another round. That
    includes a page the runner stood in for -- but not the stand-in itself. "Page 2 did
    not draw" was never written, so the delivered file is the two pages that were, and
    the built deck keeps all three for the next edit and for the findings' numbering."""
    from pptx import Presentation

    _stood_in(project)
    result = await _stage(project, outcome=_drawn(project, 3)).run(project, release=True)

    assert result.ok, result.note
    assert result.data["dropped_pages"] == [2]
    assert "page_failed" in result.data["released"]
    delivered = Presentation(result.data["pptx_path"])
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in delivered.slides]
    assert texts == ["page 1", "page 3"], texts
    assert len(Presentation(str(project.build_dir / "deck.pptx")).slides) == 3, "the built deck keeps every page"


@pytest.mark.asyncio
async def test_the_preview_beside_a_short_deck_is_short_too(project: Project) -> None:
    """The PDF the web surface previews is the render of the built deck, which has the
    stand-in in it. One page longer than the file it sits beside, it previews a deck
    nobody has -- so the same pages come out of the copy, and if they cannot come out
    there is no preview rather than a wrong one."""
    pdfium = pytest.importorskip("pypdfium2")

    _stood_in(project)
    outcome = _drawn(project, 3)
    project.review_dir.mkdir(parents=True, exist_ok=True)
    render = pdfium.PdfDocument.new()
    for _ in range(3):
        render.new_page(600, 400)
    render.save(str(project.review_dir / "deck.pdf"))
    later = outcome.pptx_path.stat().st_mtime + 5
    os.utime(project.review_dir / "deck.pdf", (later, later))

    result = await _stage(project, outcome=outcome).run(project, release=True)

    assert result.ok, result.note
    assert result.data["dropped_pages"] == [2]
    preview = Path(result.data["pdf_path"])
    assert len(pdfium.PdfDocument(str(preview))) == 2, "the preview still shows the page the deck does not have"


@pytest.mark.asyncio
async def test_a_draft_at_the_cap_delivers_instead_of_returning_unpublished(project: Project) -> None:
    """The stage's draft branch returns before the publish path, so a run that only ever
    passed `draft: true` could not reach delivery however many builds it took -- and
    three of four measured runs are exactly that run. The cap is what publishes a deck
    with problem pages rather than leaving the author with nothing, so it binds a draft:
    at the cap the build delivers whatever the call asked for."""
    stage = _stage(project, findings=[_fact()], outcome=_drawn(project, 3))

    held = await stage.run(project, draft=True)
    assert held.ok and held.data["draft"] is True
    assert "pptx_path" not in held.data, "a draft under the cap publishes nothing"

    released = await _stage(project, findings=[_fact()], outcome=_drawn(project, 3)).run(
        project, draft=True, release=True
    )

    assert released.ok, released.note
    assert Path(released.data["pptx_path"]).is_file(), "the draft that reached the cap delivered nothing"
    assert released.data["released"] == ["fact"]
    assert "draft" not in released.data, "it took the delivery path, not the draft one"
