"""Reading the horizontal grid off a template, and refusing to invent one.

Two edges, both by agreement across the template's own content pages: the bottom
of the title row they share, and the top of the furniture they keep along their
bottom edge. A template whose pages do not agree has no grid to report, and the
answer there is None rather than a plausible default -- every check built on this
returns nothing without it, which is the honest reading of "nobody measured".
"""

from __future__ import annotations

from pathlib import Path

import pytest

from raven_ppt.contracts.masters import Bands
from raven_ppt.contracts.project import Project
from raven_ppt.services.measure.bands import band_of
from raven_ppt.services.template.bands import bands_of, bands_path, read_bands, write_bands

pytest.importorskip("pptx")

TEMPLATES = Path(__file__).resolve().parents[1] / "plugins-dist" / "ppt-engine" / "raven_ppt" / "assets" / "templates"
_needs_templates = pytest.mark.skipif(
    not any(TEMPLATES.glob("*.pptx")),
    reason="the template payload is fetched, not tracked; see plugins-dist/ppt-engine/templates.manifest.json",
)
_FOOTER_ROLES = ("FOOTER", "DATE", "SLIDE_NUMBER")

# Built once per module: four saves of a five-page deck per test is most of the runtime.
_TEMPLATE_CACHE: dict[str, Path] = {}


@pytest.fixture(scope="module", autouse=True)
def _templates(tmp_path_factory: pytest.TempPathFactory) -> None:
    folder = tmp_path_factory.mktemp("templates")
    _TEMPLATE_CACHE["plain"] = _template(folder / "plain.pptx")
    _TEMPLATE_CACHE["drift"] = _template(folder / "drift.pptx", drift=0.6)
    _TEMPLATE_CACHE["bare"] = _template(folder / "bare.pptx", footers=False)
    _TEMPLATE_CACHE["marks"] = _template(folder / "marks.pptx", footers=False, marks=True)


def _template(path: Path, *, drift: float = 0.0, footers: bool = True, marks: bool = False) -> Path:
    """A cover, a divider and three content pages, on a 13.333x7.5in canvas.

    `drift` moves each content page's title row down by that much, which is a
    template whose pages agree on nothing. `footers` keeps the layout's reserved
    bottom strip; `marks` draws a hairline and a page number instead, which is how
    a template that reserves nothing still says where its footer is.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    if not footers:
        _strip_footers(presentation)
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = "封面标题"
    divider = presentation.slides.add_slide(presentation.slide_layouts[5])
    divider.shapes.title.text = "章节标题"
    for index in range(3):
        page = presentation.slides.add_slide(presentation.slide_layouts[5])
        heading = page.shapes.add_textbox(Inches(0.72), Inches(0.4 + index * drift), Inches(11.9), Inches(0.9))
        run = heading.text_frame.paragraphs[0].add_run()
        run.text = "单击此处添加页面标题"
        run.font.size, run.font.bold, run.font.name = Pt(28), True, "Arial"
        body = page.shapes.add_textbox(Inches(0.72), Inches(1.8), Inches(11.9), Inches(4.2))
        body_run = body.text_frame.paragraphs[0].add_run()
        body_run.text = f"单击此处添加文本，第 {index + 1} 页的示例正文，足够长以算作正文而不是标记"
        body_run.font.size, body_run.font.name = Pt(18), "Arial"
        if marks:
            page.shapes.add_textbox(Inches(0.72), Inches(6.9), Inches(11.9), Inches(0.02))
            number = page.shapes.add_textbox(Inches(12.0), Inches(7.1), Inches(0.6), Inches(0.3))
            number.text_frame.paragraphs[0].add_run().text = f"{index + 3:02d}"
    presentation.save(str(path))
    return path


def _strip_footers(presentation) -> None:
    for layout in presentation.slide_layouts:
        for shape in list(layout.shapes):
            if not shape.is_placeholder:
                continue
            kind = str(getattr(getattr(shape, "placeholder_format", None), "type", ""))
            if any(role in kind for role in _FOOTER_ROLES):
                shape._element.getparent().remove(shape._element)  # noqa: SLF001 -- no API removes a placeholder


def test_the_title_band_ends_where_the_content_pages_put_their_title_row() -> None:
    bands = bands_of(_TEMPLATE_CACHE["plain"])

    assert bands is not None
    assert bands.title_top == 0.0
    assert bands.title_bottom == pytest.approx(1.3), "the row at 0.4in, 0.9in tall, on all three pages"


def test_the_footer_band_starts_at_the_strip_the_layout_reserves() -> None:
    """The reservation lives on the layout and nothing copies it onto a slide, so a
    reading taken off the pages alone finds no footer on any real template."""
    bands = bands_of(_TEMPLATE_CACHE["plain"])

    assert bands is not None
    assert bands.body_bottom == pytest.approx(6.95)
    assert bands.footer_bottom == pytest.approx(7.5), "the footer band runs to the bottom of the page"


def test_a_template_that_reserves_nothing_gets_an_empty_footer_band() -> None:
    """Not a guess: nothing sits along the bottom edge, so the content band runs to
    the page edge and no shape can ever be in the footer."""
    bands = bands_of(_TEMPLATE_CACHE["bare"])

    assert bands is not None
    assert bands.body_bottom == pytest.approx(7.5)
    assert bands.body_bottom == bands.footer_bottom
    assert band_of(bands, 6.5, 7.4) == "body"


def test_a_hairline_and_a_page_number_name_the_footer_when_no_placeholder_does() -> None:
    """The topmost of the agreed marks wins, so the band starts above the rule rather
    than between the rule and the number it sits over."""
    bands = bands_of(_TEMPLATE_CACHE["marks"])

    assert bands is not None
    assert bands.body_bottom == pytest.approx(6.9)


def test_pages_that_agree_on_nothing_get_no_grid() -> None:
    """Each content page puts its title row 0.6in below the last one, so no box is on
    more than one page and there is no title band to report."""
    assert bands_of(_TEMPLATE_CACHE["drift"]) is None


def test_a_file_that_is_not_a_template_gets_no_grid(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a package")

    assert bands_of(broken) is None
    assert bands_of(tmp_path / "missing.pptx") is None


@_needs_templates
def test_the_real_template_reads_as_three_bands() -> None:
    """Measured off `beige_geometric_general_report.pptx`, which ships 25 example
    pages of which 19 are content pages.

    The numbers below were read off that file: its title placeholder is at 0.14in
    and 0.98in tall on 16 of the 19, and its content layouts ('Title Only' and
    'Blank') reserve the footer strip from 7.01in of a 13.333x7.5in canvas. All
    eight templates that ship here read the same 1.12in title row and 7.01in body
    bottom, so the grid below is the bundled grid rather than one file's.
    """
    bands = bands_of(TEMPLATES / "beige_geometric_general_report.pptx")

    assert bands is not None
    assert (bands.canvas_w, bands.canvas_h) == pytest.approx((13.33, 7.5), abs=0.01)
    assert bands.title_top == 0.0
    assert bands.title_bottom == pytest.approx(1.12, abs=0.02)
    assert bands.body_bottom == pytest.approx(7.01, abs=0.02)
    assert bands.footer_bottom == pytest.approx(7.5, abs=0.01)
    assert bands.body_area == pytest.approx(78.5, abs=0.5)


@_needs_templates
def test_the_real_template_puts_its_own_title_row_in_the_title_band() -> None:
    """The reading has to hold the page it was measured from: the title placeholder
    at 0.14in..1.12in is the title band, and the body placeholder under it is not."""
    bands = bands_of(TEMPLATES / "beige_geometric_general_report.pptx")

    assert bands is not None
    assert band_of(bands, 0.14, 1.12) == "title"
    assert band_of(bands, 1.24, 6.71) == "body"
    assert band_of(bands, 7.01, 7.31) == "footer"


def test_the_grid_is_kept_beside_the_template(tmp_path: Path) -> None:
    project = Project(workspace=tmp_path, slug="deck")
    bands = Bands(title_top=0.0, title_bottom=1.12, body_bottom=7.01, footer_bottom=7.5, canvas_w=13.33, canvas_h=7.5)

    path = write_bands(project, bands)

    assert path == bands_path(project)
    assert path.name == "bands.json"
    assert path.parent == project.root / "template", "beside palette.json and ground.txt"
    assert read_bands(project) == bands


def test_a_deck_that_measured_no_grid_reads_none(tmp_path: Path) -> None:
    project = Project(workspace=tmp_path, slug="deck")

    assert read_bands(project) is None


def test_a_grid_file_that_cannot_be_read_is_a_deck_without_one(tmp_path: Path) -> None:
    project = Project(workspace=tmp_path, slug="deck")
    path = bands_path(project)
    path.parent.mkdir(parents=True, exist_ok=True)

    path.write_text("{ not json", encoding="utf-8")
    assert read_bands(project) is None

    path.write_text('{"title_top": 0.0, "title_bottom": 1.0}', encoding="utf-8")
    assert read_bands(project) is None
