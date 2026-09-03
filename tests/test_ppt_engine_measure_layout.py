"""The two things the declared geometry alone can decide -- and, for the label
check, what the render is allowed to overrule it about."""

from __future__ import annotations

import pytest

from raven_ppt.contracts.findings import Severity
from raven_ppt.services.measure.layout import (
    BOX_SIDE_MARGINS_PT,
    EDGE_SLACK_EMU,
    LABEL_MAX_CHARS,
    LABEL_MAX_SPACES,
    LABEL_MIN_CHARS,
    LABEL_SLACK,
    off_page_shapes,
    spilled_copy,
    wrapped_labels,
)
from raven_ppt.services.measure.width import DEFAULT_MEASURER
from tests._ppt_engine_fixtures import (  # noqa: F401
    DeckBuilder,
    deck,
    image,
    noise_image,
    noise_png,
    product_page,
    template_file,
)

pytest.importorskip("pptx")


def test_the_edge_slack_is_half_a_point() -> None:
    """A rule drawn exactly on the margin rounds either way."""
    assert EDGE_SLACK_EMU == 6350
    assert EDGE_SLACK_EMU == int(0.5 * 12700)


def test_a_shape_over_the_right_edge_is_reported(deck: DeckBuilder) -> None:
    page = deck.page()
    deck.panel(page, left=13.0, top=1.0, width=1.0, height=1.0)
    deck.panel(page, left=1.0, top=1.0, width=2.0, height=1.0)

    findings = off_page_shapes(deck.save())

    assert [finding.detail["edges"] for finding in findings] == [("right",)]
    assert findings[0].page == 1
    assert findings[0].severity is Severity.WARNING


def test_a_shape_over_two_edges_names_both(deck: DeckBuilder) -> None:
    page = deck.page()
    deck.panel(page, left=13.0, top=7.0, width=1.0, height=1.0, text="Results")

    finding = off_page_shapes(deck.save())[0]

    assert finding.detail["edges"] == ("right", "bottom")
    assert "right and bottom" in finding.message
    assert "'Results'" in finding.message


def test_the_slack_is_what_decides_a_shape_on_the_margin(deck: DeckBuilder) -> None:
    """Inside the slack is inside the page; a hair past it is not."""
    from pptx.util import Emu

    findings = []
    for past in (EDGE_SLACK_EMU - 1000, EDGE_SLACK_EMU + 1000):
        builder = DeckBuilder(deck.tmp_path)
        page = builder.page()
        shape = builder.panel(page, left=1.0, top=1.0, width=1.0, height=1.0)
        shape.left = Emu(0)
        shape.width = Emu(builder.presentation.slide_width + past)
        findings.append(off_page_shapes(builder.save(f"edge{past}.pptx")))

    assert [len(found) for found in findings] == [0, 1]


def test_the_label_bounds_are_short_text_and_few_spaces() -> None:
    """Single characters cannot wrap and prose is meant to; labels are between."""
    assert (LABEL_MIN_CHARS, LABEL_MAX_CHARS, LABEL_MAX_SPACES) == (2, 24, 2)
    assert BOX_SIDE_MARGINS_PT == 14.4  # python-pptx's own left+right inset


def test_a_label_in_a_box_too_narrow_for_it_is_reported(deck: DeckBuilder) -> None:
    """'01' in a box guessed at 0.3in wraps into a stacked 0 over 1.

    A render check cannot see it: nothing overlaps, the label is just broken.
    """
    page = deck.page()
    deck.text(page, ("01", 18.0), left=1.0, top=1.0, width=0.3, height=0.4, wrap=True)
    deck.text(
        page,
        ("a full sentence of body copy that is meant to wrap over lines", 18.0),
        left=1.0,
        top=2.0,
        width=2.0,
        height=0.4,
        wrap=True,
    )

    findings = wrapped_labels(deck.save())

    assert len(findings) == 1  # the prose box wraps by design and is not held to this
    assert findings[0].detail["label"] == "01"
    assert findings[0].kind == "wrapped_label"


def test_the_slack_is_five_per_cent_of_the_boxs_own_width(deck: DeckBuilder) -> None:
    """The measurer's font is not the renderer's, so only a clear miss reports.

    Both boxes below are too narrow for the label on the measurer's own numbers.
    The one that misses by 3% is inside the slack and stays quiet; the one that
    misses by 7% does not.
    """
    assert LABEL_SLACK == 1.05
    needed = DEFAULT_MEASURER.width("MM", 18, False)

    findings = []
    for ratio in (1.03, 1.07):
        builder = DeckBuilder(deck.tmp_path)
        width_in = (needed / ratio + BOX_SIDE_MARGINS_PT) / 72
        builder.text(builder.page(), ("MM", 18.0), width=width_in, height=0.4, wrap=True)
        # The estimator explicitly, because the box was sized off its numbers: the
        # default is now FreeType on the bundled faces, which reads 'MM' wider and
        # would report both boxes. This test is about the slack, not the measurer.
        findings.append(wrapped_labels(builder.save(f"slack{ratio}.pptx"), measurer=DEFAULT_MEASURER))

    assert [len(found) for found in findings] == [0, 1]


def test_prose_and_single_characters_are_left_alone(deck: DeckBuilder) -> None:
    page = deck.page()
    for text in ("0", "x" * 25, "one two three four"):
        deck.text(page, (text, 18.0), width=0.3, height=0.4, wrap=True)

    assert wrapped_labels(deck.save()) == []


def test_a_box_told_not_to_wrap_cannot_wrap_mid_label(deck: DeckBuilder) -> None:
    deck.text(deck.page(), ("01", 18.0), width=0.3, height=0.4, wrap=False)

    assert wrapped_labels(deck.save()) == []


def test_a_run_with_no_declared_size_is_not_measured(deck: DeckBuilder) -> None:
    """Without a size there is no width to compare against the box."""
    page = deck.page()
    box = page.shapes.add_textbox(*_inches(1, 1, 0.3, 0.4))
    box.text_frame.word_wrap = True
    box.text_frame.paragraphs[0].add_run().text = "01"

    assert wrapped_labels(deck.save()) == []


def _inches(*values: float):
    from pptx.util import Inches

    return tuple(Inches(value) for value in values)


def test_chinese_body_copy_in_a_column_is_not_a_broken_label(deck) -> None:
    """The length window and the space count are both about English.

    A twenty-character Chinese sentence has no spaces at all, so it passed the
    label test and was held to "must fit on one line": the check reported
    '解码器结构沿用，改动只在查询定义与时序颈' wrapping to two lines in a three-line
    column, on a page a reviewer had just called clean. A box tall enough for a
    second line expects to wrap.
    """
    page = deck.page()
    deck.text(
        page,
        ("解码器结构沿用，改动只在查询定义与时序颈", 16.0),
        left=1.0,
        top=2.0,
        width=3.2,
        height=1.2,
        wrap=True,
    )

    assert wrapped_labels(deck.save()) == []


def test_a_label_in_a_box_that_holds_one_line_still_reports(deck) -> None:
    """The case this check exists for: '01' in a box guessed at a quarter inch."""
    page = deck.page()
    # `wrap=True` because a python-pptx textbox is created with wrap="none", and a box
    # that does not wrap cannot wrap a label -- it overflows, which other checks own.
    deck.text(page, ("01", 18.0), left=1.0, top=2.0, width=0.18, height=0.32, wrap=True)

    findings = wrapped_labels(deck.save())
    assert [f.detail["label"] for f in findings] == ["01"]


def test_copy_that_paints_off_the_page_is_reported(deck) -> None:
    """A box with wrapping off does not clip: the renderer centres the line on the box
    and paints straight out of both sides. Two page titles of a delivered deck ran off
    the canvas that way -- 14.05in of copy in an 11.48in box on a 13.33in page -- and
    every other check passed them, because the box is inside the page, the words collide
    with nothing and the type is the right size.
    """
    page = deck.page()
    deck.text(
        page,
        ("TarViS：把四类视频分割统一成一个模型的完整技术评审与后续工程建议", 30.0),
        left=0.5,
        top=0.3,
        width=3.0,
        height=0.8,
        wrap=False,
    )

    findings = spilled_copy(deck.save())
    assert [f.kind for f in findings] == ["spilled_copy"]
    assert findings[0].detail["off_page"] is True
    assert "off the edge of the page" in findings[0].message


def test_a_label_centred_in_its_own_anchor_box_is_not(deck) -> None:
    """The idiom this check must not touch: 'VIS' in a 0.12in box inside a coloured
    circle, rendered dead centre and perfectly legible. Across six real decks 38 lines
    spill out of their box and every one of them is this."""
    page = deck.page()
    deck.text(page, ("VIS", 12.0), left=6.0, top=3.0, width=0.12, height=0.2, wrap=False)

    assert spilled_copy(deck.save()) == []


def _word(page: int, text: str, x0: float, y0: float, *, w: float = 20.0, h: float = 12.0):
    """One word as the render reported it, in points with the origin top left."""
    from raven_ppt.contracts import WordBox

    return WordBox(page=page, text=text, x0=x0, y0=y0, x1=x0 + w, y1=y0 + h)


def test_the_render_overrules_a_label_it_shows_on_one_line(deck: DeckBuilder) -> None:
    """`_em_width` is a class average per character, not a font metric, and on chart
    labels it reads high: eight digit-and-percent labels on one delivered page were
    over-measured by about 13% where `LABEL_SLACK` allows 5, and every one of them was
    on a single line in the render. Raising the slack trades that for a number guessed
    the other way, so the prediction finds candidates and the render decides.
    """
    page = deck.page()
    deck.text(page, ("15.3%", 18.0), left=1.0, top=1.0, width=0.6, height=0.4, wrap=True)
    path = deck.save()

    assert [f.detail["label"] for f in wrapped_labels(path)] == ["15.3%"]

    # The render put it on one line: one word box inside the shape's own rectangle.
    on_one_line = [_word(1, "15.3%", 74.0, 74.0)]
    assert wrapped_labels(path, None, on_one_line) == []


def test_the_render_does_not_overrule_a_label_it_shows_broken(deck: DeckBuilder) -> None:
    """The case the check exists for has to survive the veto. Measured on a real
    render: '01' in a 0.25in box came back as '0' at y=76.9 and '1' at y=110.5.
    """
    page = deck.page()
    deck.text(page, ("01", 28.0), left=1.0, top=1.0, width=0.25, height=1.0, wrap=True)
    path = deck.save()

    broken = [_word(1, "0", 79.0, 76.9, w=9.0), _word(1, "1", 79.0, 110.5, w=9.0)]
    assert [f.detail["label"] for f in wrapped_labels(path, None, broken)] == ["01"]


def test_a_label_the_render_says_nothing_about_keeps_the_file_s_answer(deck: DeckBuilder) -> None:
    """No render, a render that dropped the glyph, a transcription that does not match:
    all the same case, and none of them is evidence the label was fine. The veto acts
    only on positive evidence, or a host without pdftotext silently loses the check.
    """
    page = deck.page()
    deck.text(page, ("01", 28.0), left=1.0, top=1.0, width=0.25, height=1.0, wrap=True)
    path = deck.save()

    assert [f.detail["label"] for f in wrapped_labels(path, None, None)] == ["01"]
    assert [f.detail["label"] for f in wrapped_labels(path, None, [])] == ["01"]
    # A word box nowhere near the shape says nothing about it either.
    elsewhere = [_word(1, "01", 700.0, 400.0)]
    assert [f.detail["label"] for f in wrapped_labels(path, None, elsewhere)] == ["01"]


def test_the_veto_reads_only_the_page_the_label_is_on(deck: DeckBuilder) -> None:
    """Word boxes carry their page, and a label set on one line on page 2 says nothing
    about the same label broken on page 1."""
    first = deck.page()
    deck.text(first, ("01", 28.0), left=1.0, top=1.0, width=0.25, height=1.0, wrap=True)
    second = deck.page()
    deck.text(second, ("01", 28.0), left=1.0, top=1.0, width=0.25, height=1.0, wrap=True)
    path = deck.save()

    only_page_two_is_fine = [_word(2, "01", 79.0, 76.9)]
    kept = wrapped_labels(path, None, only_page_two_is_fine)

    assert [f.page for f in kept] == [1]
