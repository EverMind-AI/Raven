"""A built deck leaves without the template vendor's marks, and with its own title."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from raven_ppt.services.publish import strip_vendor_marks
from raven_ppt.services.publish.provenance import CREATOR, THEME_NAME

VENDOR = re.compile(rb"islide", re.IGNORECASE)

TAGS_PART = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    b'<p:tagLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    b'<p:tag name="ISLIDE.GUIDESSETTING" val="{&quot;Id&quot;:&quot;GuidesStyle_Normal&quot;}"/>'
    b'<p:tag name="OWN.MARK" val="kept"/></p:tagLst>'
)

CUSTOM_PART = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    b'<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="ISLIDE.THEME">'
    b"<vt:lpwstr>8da5d54c-8099-496d-baf7-4f02a6040b7d</vt:lpwstr></property>"
    b'<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="OWN.PROP">'
    b"<vt:lpwstr>kept</vt:lpwstr></property></Properties>"
)

LAYOUT_CREDIT = (
    b'<p:sp><p:nvSpPr><p:cNvPr id="90" name="Credit"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
    b'<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
    b'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/>'
    b'<a:p><a:r><a:rPr lang="en-US"/><a:t>Template by </a:t></a:r><a:r><a:rPr lang="en-US"/><a:t>iSlide</a:t></a:r>'
    b'<a:r><a:rPr lang="en-US"/><a:t> studio, all rights reserved</a:t></a:r></a:p></p:txBody></p:sp>'
)
"""A sentence on a layout with the vendor's name in its own run, as the kits write it."""


def _signed_deck(path: Path) -> Path:
    """A two-page deck carrying every kind of mark the bundled templates carry."""
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    box = cover.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Night markets"
    credit = cover.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    credit.text_frame.text = "iSlide"
    credit.name = "iSlide credit"
    footer = cover.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    footer.text_frame.text = "www.islide.cc"
    layout_part = cover.slide_layout.part.partname.lstrip("/")
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.core_properties.title = "iSlide PowerPoint Template"
    prs.core_properties.author = "iSlide"
    prs.core_properties.last_modified_by = "xin li"
    prs.core_properties.keywords = "iSlide, template"
    prs.save(path)

    # The theme's own names and an add-in's tag records are below python-pptx's
    # API, so they go in the way the kit wrote them: straight into the package.
    signed = path.with_name("signed.pptx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(signed, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "ppt/theme/theme1.xml":
                data = data.replace(b'<a:fontScheme name="Office"', b'<a:fontScheme name="iSlide"', 1)
                data = re.sub(rb'(<a:theme [^>]*?name=")[^"]*(")', rb"\1Designed by iSlide\2", data, count=1)
            if info.filename == "docProps/app.xml":
                data = data.replace(b"</Properties>", b"<Company>iSlide</Company></Properties>")
            if info.filename == layout_part:
                data = data.replace(b"</p:spTree>", LAYOUT_CREDIT + b"</p:spTree>", 1)
            if info.filename == "_rels/.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    b'relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>',
                )
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/ppt/tags/tag1.xml" ContentType="application/vnd.openxmlformats-'
                    b'officedocument.presentationml.tags+xml"/>'
                    b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-'
                    b'officedocument.custom-properties+xml"/></Types>',
                )
            dst.writestr(info, data)
        dst.writestr("ppt/tags/tag1.xml", TAGS_PART)
        dst.writestr("docProps/custom.xml", CUSTOM_PART)
    return signed


LAYOUT_PART = "ppt/slideLayouts/slideLayout7.xml"
"""Where python-pptx's blank layout lives; `_signed_deck` signs the layout the cover uses."""


def _marked_parts(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path) as z:
        return {
            n: len(VENDOR.findall(z.read(n)))
            for n in z.namelist()
            if n.endswith((".xml", ".rels")) and VENDOR.search(z.read(n))
        }


def test_the_fixture_carries_the_marks_the_templates_carry(tmp_path: Path) -> None:
    signed = _signed_deck(tmp_path / "deck.pptx")
    marked = _marked_parts(signed)
    assert {
        "docProps/core.xml",
        "docProps/app.xml",
        "docProps/custom.xml",
        "ppt/theme/theme1.xml",
        "ppt/tags/tag1.xml",
        "ppt/slides/slide1.xml",
        LAYOUT_PART,
    } <= set(marked)


def test_every_vendor_mark_leaves_and_the_deck_still_opens(tmp_path: Path) -> None:
    signed = _signed_deck(tmp_path / "deck.pptx")

    changed = strip_vendor_marks(signed, title="Night markets: how a street stays open")

    assert _marked_parts(signed) == {}
    assert set(changed) == {
        "docProps/core.xml",
        "docProps/app.xml",
        "docProps/custom.xml",
        "ppt/theme/theme1.xml",
        "ppt/tags/tag1.xml",
        "ppt/slides/slide1.xml",
        LAYOUT_PART,
    }
    prs = Presentation(signed)
    assert len(prs.slides) == 2
    assert prs.core_properties.title == "Night markets: how a street stays open"
    assert prs.core_properties.author == CREATOR
    assert prs.core_properties.last_modified_by == CREATOR
    assert prs.core_properties.keywords == ""
    # The name and the address the vendor left are gone whole; the deck's own words stay.
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert texts == ["Night markets", "", ""]
    # On a layout the vendor's sentence goes whole, not just the run with the name in it.
    layout_texts = [sh.text_frame.text for sh in prs.slides[0].slide_layout.shapes if sh.has_text_frame]
    assert "" in layout_texts and not any("rights reserved" in t for t in layout_texts)
    with zipfile.ZipFile(signed) as z:
        theme = z.read("ppt/theme/theme1.xml")
        assert f'<a:fontScheme name="{THEME_NAME}"'.encode() in theme
        # The theme keeps everything the template was chosen for.
        assert b"<a:clrScheme" in theme and b"<a:majorFont>" in theme
        tags = z.read("ppt/tags/tag1.xml")
        assert b"OWN.MARK" in tags and b"ISLIDE" not in tags
        app = z.read("docProps/app.xml")
        assert b"<Company></Company>" in app or b"<Company/>" in app
        custom = z.read("docProps/custom.xml")
        assert b"OWN.PROP" in custom and b"ISLIDE" not in custom


@pytest.mark.parametrize(
    "sentence",
    [
        "iSlide market analysis",
        "PowerPoint with iSlide",
        "iSlide PowerPoint templates",
        "PPT by iSlide",
        "Designed by iSlide",
        "Great Presentation Made Easy with iSlide",
    ],
)
def test_a_slide_that_speaks_of_the_vendor_keeps_its_words(tmp_path: Path, sentence: str) -> None:
    """A deck about the vendor is the deck's business: only the bare name or address leaves a slide.

    That holds for a heading as short as "PPT by iSlide", and even for "Designed by
    iSlide" typed on a slide: a credit the template phrased that way sits on its
    layout, where the whole paragraph goes.
    """
    deck = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = sentence
    slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = "www.islide.cc"
    prs.save(deck)

    changed = strip_vendor_marks(deck, title="Software landscape")

    assert set(changed) == {"docProps/core.xml", "ppt/slides/slide1.xml"}
    texts = [sh.text_frame.text for sh in Presentation(deck).slides[0].shapes if sh.has_text_frame]
    assert texts == [sentence, ""]


@pytest.mark.parametrize("credit", ["iSlide", "www.islide.cc", "https://www.islide.cc/", "islide.cn", " iSlide -- "])
def test_a_bare_name_or_address_is_a_credit(credit: str) -> None:
    from raven_ppt.services.publish.provenance import _is_credit

    assert _is_credit(credit)
    assert not _is_credit(credit + " for the board")
    assert not _is_credit("PPT by " + credit)


def test_a_deck_without_marks_is_left_byte_for_byte(tmp_path: Path) -> None:
    """A user's own template keeps its author; only the last writer is this build."""
    clean = tmp_path / "clean.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.core_properties.title = "Own title"
    prs.core_properties.author = "Their design office"
    prs.core_properties.last_modified_by = CREATOR
    prs.save(clean)
    before = clean.read_bytes()

    assert strip_vendor_marks(clean) == []
    assert clean.read_bytes() == before


def test_a_title_is_written_even_when_the_template_left_none(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.core_properties.title = ""
    prs.save(deck)

    assert strip_vendor_marks(deck, title="Shanghai 2035") == ["docProps/core.xml"]
    assert Presentation(deck).core_properties.title == "Shanghai 2035"


def test_something_that_is_not_a_package_is_left_alone(tmp_path: Path) -> None:
    stub = tmp_path / "deck.pptx"
    stub.write_bytes(b"PK deck")
    assert strip_vendor_marks(stub, title="x") == []
    assert stub.read_bytes() == b"PK deck"


@pytest.mark.parametrize("name", ["docProps/core.xml", "ppt/theme/theme1.xml"])
def test_the_rewrite_keeps_the_part_well_formed(tmp_path: Path, name: str) -> None:
    from lxml import etree

    signed = _signed_deck(tmp_path / "deck.pptx")
    strip_vendor_marks(signed, title="t")
    with zipfile.ZipFile(signed) as z:
        etree.fromstring(z.read(name))


@pytest.mark.asyncio
async def test_the_build_stage_delivers_the_deck_without_the_marks_and_with_the_topic(tmp_path: Path) -> None:
    """Placed before measuring, so the file the gates saw is the file delivered."""
    import json

    from raven_ppt.contracts import BuildOutcome, PageSource, Project
    from raven_ppt.profiles import registry
    from raven_ppt.stages.build import BuildStage

    project = Project(workspace=tmp_path, slug="night")
    project.build_dir.mkdir(parents=True)
    project.state_dir.mkdir(parents=True)
    (project.state_dir / "intake.json").write_text(json.dumps({"topic": "夜间市集运营要点"}), encoding="utf-8")
    signed = _signed_deck(project.build_dir / "deck.pptx")
    seen_by_measure: list[dict[str, int]] = []

    async def backend(_project: Project, _script: str | None) -> BuildOutcome:
        return BuildOutcome(
            ok=True,
            pptx_path=signed,
            pages=2,
            source_digest="x",
            sources=(PageSource(page=1, first_line=0, last_line=1), PageSource(page=2, first_line=1, last_line=2)),
        )

    async def measure(_project: Project, pptx: Path, _outcome=None, _changed: str = "delivery") -> list:
        seen_by_measure.append(_marked_parts(pptx))
        return []

    result = await BuildStage(backend=backend, measure=measure, profile=registry.get("script_author")).run(project)

    assert result.ok, result.note
    assert seen_by_measure == [{}]
    delivered = Path(result.data["pptx_path"])
    assert _marked_parts(delivered) == {}
    assert Presentation(delivered).core_properties.title == "夜间市集运营要点"
