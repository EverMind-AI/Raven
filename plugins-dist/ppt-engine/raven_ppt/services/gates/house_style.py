"""Whether the deck was actually built inside the template it was given.

Nothing checked this. A user handed over their house style, `ppt_template` bound
it, the build directory got `PPT_TEMPLATE`, the tool description explained how to
open it -- and an author that wrote `Presentation()` instead produced a white deck
on a default canvas in Calibri, which published clean. Every instruction in that
chain was prose, and the one guarantee that could have caught it did not exist.

Read off the theme's colour scheme, because that is the part a template cannot
lose and a default cannot fake: `Presentation()` carries Office's own scheme, and
any template worth binding has replaced it. Comparing masters or layout names
would not do -- a template can ship one master called "Office Theme" like everyone
else -- and comparing the canvas would not either, since an author that sets
13.333in by hand gets that right while getting everything else wrong.

Reported. A deck in somebody else's colours is not a deck the user asked for and
is not answerable by rearranging a page, which argued for a refusal and got one --
but nothing measured this check against a template and a deck really built inside
it, and the whole predicate is exact equality of every theme slot, so an author who
repainted the theme itself would be told it never opened the template (D52).
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from raven_ppt.contracts import Finding, Severity

# How many differing colours the message shows. Enough to recognise which deck is
# which; a full twelve-slot diff is a table nobody reads to decide a one-line fix.
_SHOWN = 3


def house_style_findings(pptx_path: Path, template: Path | None) -> list[Finding]:
    """One finding when the deck's theme is not the bound template's."""
    if template is None or not Path(template).is_file():
        return []
    wanted = _scheme(Path(template))
    if not wanted:
        # A template with no colour scheme cannot be checked against, and saying
        # nothing is the honest answer -- refusing here would refuse a deck for a
        # property of the user's own file.
        return []
    built = _scheme(Path(pptx_path))
    if not built or built == wanted:
        return []
    differing = [
        f"{slot} is {built.get(slot, 'unset')} and the template's is {value}"
        for slot, value in wanted.items()
        if built.get(slot) != value
    ]
    return [
        Finding(
            kind="house_style",
            severity=Severity.WARNING,
            message=(
                f"this deck's theme is not {Path(template).name}'s, so it was not built inside the template "
                f"the user gave: {'; '.join(differing[:_SHOWN])}. Open the deck with "
                "`Presentation(os.environ['PPT_TEMPLATE'])` rather than `Presentation()` -- every page added "
                "to it then inherits the master, the theme and the canvas"
            ),
            detail={"differing": len(differing), "template": Path(template).name},
        )
    ]


def _scheme(path: Path) -> dict[str, str]:
    """The theme's colour scheme by slot, or {} when there is none to read."""
    try:
        from pptx import Presentation

        from raven_ppt.services.template.inventory import _theme_colours, _theme_root
    except ImportError:  # pragma: no cover - python-pptx ships with the extra
        return {}
    try:
        presentation = Presentation(str(path))
    except Exception:  # noqa: BLE001 -- a file that will not open is not comparable
        return {}
    return dict(_theme_colours(_theme_root(presentation)))


# How far a page's title row may sit from the template's own before a reader sees two
# grids. A tenth of an inch is the nudge a designer makes; a quarter is a different
# decision, and the deck that produced this had its titles at 0.04in from the left edge
# against the template's 0.72in, with a 9.92in box where the template's is 11.88in --
# which is also why two of those titles painted off the page.
TITLE_DRIFT_IN = 0.2
TITLE_WIDTH_DRIFT_IN = 0.6


def _title_of(rows, house_width):
    """Which of the shapes in the title band is the title.

    Not the topmost one. A page that sets a kicker over its title -- "01 问题" above
    "四类任务的差别" -- puts the kicker first, and taking it as the title measured a
    row 0.29in above where the template's title sits: eight pages of a deck reported
    for a title row that was in exactly the right place, and an author that could not
    fix it because nothing was wrong.

    A title runs the width of the safe area and is set larger than anything else up
    there, so among the shapes that are as wide as the template's own title row, it is
    the tallest. When none is that wide -- which is itself what `narrow` reports --
    fall back to the whole band so the finding still lands.
    """
    from raven_ppt.services.measure.geometry import EMU_PER_INCH

    wide = [shape for shape in rows if house_width - (shape.width or 0) / EMU_PER_INCH <= TITLE_WIDTH_DRIFT_IN]
    return max(wide or rows, key=lambda shape: ((shape.height or 0), -(shape.top or 0)))


def title_row_findings(pptx_path: Path, prototypes: Path | None, structural: Sequence[int] = ()) -> list[Finding]:
    """Content pages whose title row is not where the template puts one.

    The one element every page of a deck shares. `type_drift` measures the sizes and
    this measures the place: a deck whose titles start at three different left edges
    reads as three decks, and a title box narrower than the template's is what sends
    copy off the canvas when wrapping is off.

    The place is the box and where the copy sits in it. A row anchored to the bottom of
    a 0.98in box and one anchored to its top are the same rectangle and 0.375in apart on
    the render, which is what a deck alternating cloned and composed pages came back as:
    cloned titles starting at y=53px and composed ones at y=23px, same glyph height,
    every second page. Nothing about the box says which of the two a page chose, so the
    anchor is read separately and compared to the template's own.

    Only the pages the deck composed. A cloned page carries the template's own title row
    by construction -- the closing page of one deck puts its title mid-page, which is the
    template's design and was this check's first false positive. Two readings find those
    pages and both are needed: `_cloned_pages` matches geometry against the template's
    structural prototypes, and `structural` is what the author itself declared. The
    declaration was taken as an argument here and never read, and geometry alone missed
    every page the author composed *as* a cover or a divider rather than cloning one:
    across four live decks this reported nine pages, and eight of them were a contents,
    a section or a closing page the outline had already named.
    """
    if prototypes is None or not Path(prototypes).is_file():
        return []
    from raven_ppt.services.measure.geometry import EMU_PER_INCH, iter_shapes, open_deck
    from raven_ppt.services.template.house import house_style, title_anchor
    from raven_ppt.services.template.menu import menu, roles

    house = house_style(Path(prototypes))
    if house is None or house.title is None or house.title.pages < 2:
        return []  # no row the template itself agrees on, so nothing to hold a page to
    wanted = house.title.box
    # Measured, not promised: naming a prototype is now the default for a content
    # page, so trusting the promise skipped every page in the deck. `_cloned_pages`
    # already restricts itself to the template's structural pages, which is the set
    # whose title row is the template's own and not this page's to answer for.
    cloned = _cloned_pages(pptx_path, Path(prototypes), roles(menu(Path(prototypes))))
    cloned |= set(structural)
    findings: list[Finding] = []
    presentation = open_deck(pptx_path)
    band = (presentation.slide_height or 0) / EMU_PER_INCH * 0.35
    for number, slide in enumerate(presentation.slides, start=1):
        if number in cloned:
            continue
        rows = [
            shape
            for shape in iter_shapes(slide.shapes)
            if getattr(shape, "has_text_frame", False)
            and shape.top is not None
            and shape.text_frame.text.strip()
            and shape.top / EMU_PER_INCH <= band
        ]
        if not rows:
            continue
        title = _title_of(rows, wanted[2])
        left, top = title.left / EMU_PER_INCH, title.top / EMU_PER_INCH
        width = (title.width or 0) / EMU_PER_INCH
        height = (title.height or 0) / EMU_PER_INCH
        off = max(abs(left - wanted[0]), abs(top - wanted[1]))
        narrow = wanted[2] - width
        anchored = title_anchor(slide, title)
        drifted = off > TITLE_DRIFT_IN or narrow > TITLE_WIDTH_DRIFT_IN
        unanchored = house.title.anchor is not None and anchored != house.title.anchor
        if not drifted and not unanchored:
            continue
        said = []
        if drifted:
            said.append(
                f"this page's title sits at ({left:.2f}, {top:.2f}) {width:.2f}in wide and the template puts "
                f"its own at ({wanted[0]:.2f}, {wanted[1]:.2f}) {wanted[2]:.2f}in, on {house.title.pages} of "
                f"its pages. Put the title in that box: it is the one element every page of the deck shares, "
                f"and a box narrower than the template's is what sends a long title off the canvas"
            )
        if unanchored:
            said.append(
                f"this page's title is anchored {anchored} in its box and the template anchors its own "
                f"{house.title.anchor}, so the two sets of pages put the same line up to {height:.2f}in apart "
                f'while every box agrees. Pass anchor="{house.title.anchor}" to the call that writes it '
                f"-- `title_row_as_code` in the house style is that call with the anchor already in it"
            )
        findings.append(
            Finding(
                kind="title_row",
                severity=Severity.WARNING,
                page=number,
                message=". ".join(said),
                detail={
                    "at": [round(left, 2), round(top, 2), round(width, 2)],
                    "house": list(wanted),
                    "anchor": anchored,
                    "house_anchor": house.title.anchor,
                },
            )
        )
    return findings


def _cloned_pages(pptx_path: Path, prototypes: Path, named: dict[str, int]) -> set[int]:
    """Pages built on one of the template's structural pages, which own their title row."""
    from raven_ppt.services.measure.adherence import FROM_PROTOTYPE, MIN_SHAPES, _matches, _pages

    if not named:
        return set()
    shipped = _pages(prototypes)
    cloned: set[int] = set()
    for number, shapes in _pages(pptx_path).items():
        # Half the page's shapes and at least four of them. The loose version -- a fifth
        # of them, any count -- called a two-shape page cloned because one of its two
        # shapes was the layout's own title placeholder, which sits exactly where the
        # template's title row does. That is the page this check is for.
        if len(shapes) < MIN_SHAPES:
            continue
        for page in named.values():
            boxes = shipped.get(page) or set()
            if boxes and sum(1 for box in shapes if _matches(box, boxes)) / len(shapes) >= FROM_PROTOTYPE:
                cloned.add(number)
                break
    return cloned
