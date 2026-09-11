"""The deck as it is delivered when the build cap releases it.

The cap (services/tier) exists to stop buying quality: past it a deck goes out
with whatever the gates still say about it, because a deck that exists is worth
more than another round of fixes. What it must not do is deliver the runner's
own furniture. A page reading "Page 11 did not draw" was never written, and
"publish the pages that are there" does not mean it.

So the deck published at the cap is a copy with those pages taken out, and the
PDF that previews it has the same pages taken out, so the two agree. The built
deck keeps them all: it is what the next edit is repaired against, what every
finding's page number describes, and what the second reader reads.
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path


def without_pages(built: Path, pages: Sequence[int], target: Path) -> int:
    """Write `built` to `target` without the 1-based `pages`; returns pages written.

    Both halves of removing a slide are required. Dropping the id from the slide
    list without dropping the relationship leaves an orphaned part PowerPoint
    offers to repair; the reverse leaves a reference that will not open at all.
    """
    from pptx import Presentation

    presentation = Presentation(str(built))
    listing = presentation.slides._sldIdLst
    entries = list(listing)
    for index in sorted({int(page) for page in pages}, reverse=True):
        if not 1 <= index <= len(entries):
            continue
        entry = entries[index - 1]
        presentation.part.drop_rel(entry.rId)
        listing.remove(entry)
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(target))
    return len(listing)


def pdf_without_pages(rendered: Path, pages: Sequence[int], target: Path) -> bool:
    """Copy `rendered` to `target` without the 1-based `pages`. False when it cannot.

    False rather than a copy of the whole render: the preview stands for the file
    delivered beside it, and one page longer than that file is a preview of a deck
    nobody has. The caller drops the preview instead of showing a wrong one.
    """
    from raven_ppt.services.render.capabilities import PDFIUM_LOCK, pdfium

    module = pdfium()
    if module is None:
        return False
    with PDFIUM_LOCK:
        try:
            document = module.PdfDocument(str(rendered), autoclose=True)
            try:
                count = len(document)
                for index in sorted({int(page) for page in pages}, reverse=True):
                    if 1 <= index <= count:
                        document.del_page(index - 1)
                target.parent.mkdir(parents=True, exist_ok=True)
                document.save(str(target))
            finally:
                document.close()
        except Exception:  # noqa: BLE001 -- no preview is better than the wrong one
            return False
    return target.is_file()
