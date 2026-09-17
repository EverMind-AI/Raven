"""A formula typeset as a picture, for a build script that is not the engine's.

Below the top tier the deck is written by hand in python-pptx on the interpreter
raven runs on, so the script imports this rather than the generated ``ppt_layout``
module and gets the same answer that module's ``formula`` gives a TeX expression:
matplotlib's own typesetter (mathtext, no TeX installation) renders the expression
to a transparent PNG in the deck's ink, and the picture is placed at true size --
at 300 dpi a point is a point, so a formula asked for at the body size sits level
with the body copy beside it.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from pptx.util import Inches

DPI = 300
SANS_FONTSET = "dejavusans"
SERIF_FONTSET = "stix"
# STIX's capitals stand 0.89 of DejaVu's and its x-height 0.81, measured at 18pt, so
# a serif formula set at the body size read a size smaller than the body.
SERIF_SCALE = 1.2
_CJK = re.compile(r"[　-ヿ一-鿿]")


@dataclass(frozen=True)
class Formula:
    """The rendered expression: its PNG, the inches it takes at true size, the size it was set at."""

    png: bytes
    width_in: float
    height_in: float
    size_pt: float


def display_style(tex: str) -> str:
    """``\\frac`` set as ``\\dfrac``.

    mathtext sizes an inline fraction's halves down to script size, which on a
    slide turned ``\\frac{\\lambda}{2}`` into two specks around a bar; a formula on
    its own line is display mathematics and gets the full-size halves.
    """
    return tex.replace("\\frac{", "\\dfrac{")


def formula_png(tex: str, *, size_pt: float = 18.0, colour: str = "1F2A44", serif: bool = False) -> Formula:
    """Render ``tex`` at ``size_pt`` in ``colour`` (RRGGBB, with or without the ``#``)."""
    body = str(tex).strip().strip("$").strip()
    if _CJK.search(body):
        raise ValueError(
            "a TeX expression cannot carry CJK text -- the typesetter has no glyphs for it. Keep the "
            "words outside: write the sentence as text, and give the formula the mathematics alone"
        )
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    from PIL import Image

    size = size_pt * SERIF_SCALE if serif else size_pt
    figure = Figure(figsize=(0.1, 0.1))
    FigureCanvasAgg(figure)
    with matplotlib.rc_context({"mathtext.fontset": SERIF_FONTSET if serif else SANS_FONTSET}):
        figure.text(0, 0, f"${display_style(body)}$", fontsize=size, color=f"#{colour.lstrip('#')}")
        buffer = io.BytesIO()
        try:
            figure.savefig(buffer, dpi=DPI, transparent=True, bbox_inches="tight", pad_inches=0.02, format="png")
        except ValueError as exc:
            raise ValueError(
                f"the expression {body!r} did not typeset: {exc}. Write it in the TeX mathtext knows -- "
                "\\frac{}{}, \\sqrt{}, \\sum_{}^{}, \\mathrm{} for a word set upright, _{} and ^{} -- "
                "no \\text{}, no align or matrix environments, and keep the prose around it outside"
            ) from exc
    png = buffer.getvalue()
    with Image.open(io.BytesIO(png)) as image:
        return Formula(png, image.width / DPI, image.height / DPI, size)


def add_formula(
    slide,
    tex: str,
    left_in: float,
    top_in: float,
    *,
    size_pt: float = 18.0,
    colour: str = "1F2A44",
    serif: bool = False,
    max_width_in: float | None = None,
):
    """Place ``tex`` as a picture with its top-left at (``left_in``, ``top_in``) inches.

    ``size_pt`` is the size of the body copy the formula sits beside; the picture
    lands at true size, scaled down only when ``max_width_in`` is narrower than the
    expression. Hands back the picture shape, whose ``width`` and ``height`` are
    what the page now has to make room for.
    """
    rendered = formula_png(tex, size_pt=size_pt, colour=colour, serif=serif)
    width = rendered.width_in
    if max_width_in is not None and width > max_width_in > 0:
        width = max_width_in
    shape = slide.shapes.add_picture(io.BytesIO(rendered.png), Inches(left_in), Inches(top_in), width=Inches(width))
    shape.name = "formula"
    return shape


_SCRIPT = re.compile(r"([_^])(\{[^}]*\}|[^\s_^{}])")
_STAR = re.compile(r"(?<=[A-Za-z\u0370-\u03ff])\*")
_WORDS = re.compile(r"[A-Za-z]+|[^A-Za-z]+")
_MATH_NEIGHBOUR = re.compile(r"[(){}\[\]_^*|=+\-\u2212\u00d7\u00b7,/<>\u2264\u2265\u226a\u226b0-9]")
SUBSCRIPT_BASELINE = "-25000"
SUPERSCRIPT_BASELINE = "30000"


def _pieces(text: str):
    """(piece, level) with level -1 for a subscript, 1 for a superscript, 0 for the line."""
    text = _STAR.sub("^*", text)
    at = 0
    for match in _SCRIPT.finditer(text):
        if match.start() > at:
            yield text[at : match.start()], 0
        body = match.group(2)
        yield (body[1:-1] if body.startswith("{") else body), (-1 if match.group(1) == "_" else 1)
        at = match.end()
    if at < len(text):
        yield text[at:], 0


def math_runs(paragraph, text: str, *, size_pt: float = 18.0, colour: str = "1F2A44", font: str | None = None):
    """Set a sentence that carries symbols as text runs with real sub- and superscripts.

    ``_x`` and ``_{xyz}`` subscript, ``^x`` and ``^{xyz}`` superscript, and a ``*`` right
    after a letter is the superscript star of an optimum. A lone latin letter is italic
    when it is a variable: in a script, beside a bracket, operator or digit, or anywhere
    in a CJK sentence; the English words ``a`` and ``I`` in Latin prose stay upright, as
    does a word of two or more letters. Every run is set at ``size_pt``, the scripts
    included. An underscore is always a subscript marker here, so a file name or an
    identifier does not belong in the text. For anything that stacks (a fraction, a
    root, a sum with limits) use ``add_formula``.
    """
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    runs = []
    cjk_sentence = bool(_CJK.search(text))
    for piece, level in _pieces(text):
        words = _WORDS.findall(piece)
        for index, word in enumerate(words):
            before = words[index - 1][-1:] if index else ""
            after = words[index + 1][:1] if index + 1 < len(words) else ""
            variable = level != 0 or cjk_sentence or bool(_MATH_NEIGHBOUR.search(before + after))
            run = paragraph.add_run()
            run.text = word
            run.font.size = Pt(size_pt)
            run.font.italic = len(word) == 1 and word.isalpha() and word.isascii() and variable
            run.font.color.rgb = RGBColor.from_string(colour.lstrip("#"))
            if font:
                run.font.name = font
            if level:
                run._r.get_or_add_rPr().set("baseline", SUBSCRIPT_BASELINE if level < 0 else SUPERSCRIPT_BASELINE)
            runs.append(run)
    return runs
