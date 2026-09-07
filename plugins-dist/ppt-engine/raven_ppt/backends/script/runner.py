"""Running the author's program, and surviving it failing.

Everything here exists because the program is written by a model and may do
anything: not compile, compile and crash, crash halfway through after writing a
partial file, hang, or finish without writing anything at all. The deck that last
built is what a failed edit gets repaired against, so none of those outcomes may
destroy it.
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import os
import shutil
import sys
from pathlib import Path

from raven_ppt.backends.script.blocks import level_that_separates_pages, page_blocks, page_sources
from raven_ppt.backends.script.submission import carries_a_program, submission_refusal
from raven_ppt.backends.script.workspace import (
    SCRIPT_ENCODING,
    HelperSources,
    deck_path,
    page_failures_path,
    provision,
    script_path,
    slide_lines_path,
    with_template_helpers,
)
from raven_ppt.contracts import BuildOutcome, DeckPlan, Project
from raven_ppt.services.template import bound
from raven_ppt.services.tidy import tidy

log = logging.getLogger(__name__)

BUILD_TIMEOUT_S = 300.0
MAX_OUTPUT_CHARS = 20_000

# Runs the script and records which of its lines created each slide. Reading that
# from a `# SLIDE n` comment would make annotation the author's job and let a
# stale number pair one page's render with another page's code; execution knows
# the answer exactly.
_RUNNER = """import builtins, json, os, re, sys, traceback

script = sys.argv[1]
record = os.environ.get("PPT_SLIDE_LINES")
blocks = json.loads(os.environ.get("PPT_SLIDE_BLOCKS") or "[]")
sys.argv = [script]
target = os.path.realpath(script)

created = []
failed = []
slides_seen = [None]
standing_in = [None]
if record or blocks:
    import inspect

    from pptx.slide import Slides

    original = Slides.add_slide

    def add_slide(self, *args, **kwargs):
        slide = original(self, *args, **kwargs)
        slides_seen[0] = self
        if standing_in[0] is not None:
            # A placeholder stands where a failed block's page would be, so it is
            # attributed to that block's first line and the mapping still holds.
            created.append([standing_in[0]])
            return slide
        # Every frame in the script, outermost first. Which level identifies the
        # page depends on how the script is shaped, so record them all and let
        # the caller pick the level whose lines differ per slide.
        stack = [
            frame.lineno for frame in reversed(inspect.stack()) if os.path.realpath(frame.filename) == target
        ]
        created.append(stack)
        return slide

    Slides.add_slide = add_slide


def _drop_last(slides, count):
    listing = slides._sldIdLst
    for entry in list(listing)[-count:] if count else []:
        slides.part.drop_rel(entry.rId)
        listing.remove(entry)


def _stand_in(slides, number, error):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = slides.part.presentation
    layouts = prs.slide_layouts
    layout = layouts[6] if len(layouts) > 6 else layouts[len(layouts) - 1]
    slide = slides.add_slide(layout)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.6), prs.slide_width - Inches(1.2), prs.slide_height - Inches(1.2))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = f"Page {number} did not draw"
    head = frame.paragraphs[0].runs[0].font
    head.size, head.bold, head.color.rgb = Pt(28), True, RGBColor(0xB0, 0x00, 0x20)
    body = frame.add_paragraph()
    body.text = error
    body.runs[0].font.size = Pt(14)
    body.runs[0].font.color.rgb = RGBColor(0x40, 0x40, 0x40)


def _slides():
    if slides_seen[0] is not None:
        return slides_seen[0]
    from pptx.presentation import Presentation

    for value in list(namespace.values()):
        if isinstance(value, Presentation):
            return value.slides
    return None


with open(script, "rb") as handle:
    body = handle.read()
# utf-8-sig, as workspace.SCRIPT_ENCODING has it: a byte-order mark the interpreter
# would skip on `python build.py` is a SyntaxError once the text reaches compile().
lines = body.decode("utf-8-sig", "replace").splitlines(keepends=True)
namespace = {"__name__": "__main__", "__file__": target, "__builtins__": builtins}


def run(start, end):
    exec(compile("\\n" * start + "".join(lines[start:end]), target, "exec"), namespace)


try:
    if not blocks:
        # Through `run` rather than `runpy.run_path`: a script whose pages cannot be
        # isolated still has its pages, and the rescue below finds a presentation only
        # in this namespace. runpy builds its own, so a crash on this path threw away
        # everything drawn -- which is the path a crash takes, since a script the
        # banners cannot separate is also the one with no page-level catch.
        run(0, len(lines))
    else:
        # One block at a time, so a page that raises is that page's failure and not the
        # deck's: the slides it managed to add are taken back, a page saying what went
        # wrong stands in their place, and the pages after it are still drawn.
        run(0, blocks[0][1])
        for number, start, end in blocks:
            before = len(created)
            try:
                run(start, end)
            # `SystemExit` as well as `Exception`, because the helpers an author writes
            # raise it: a live run's own `sub1` ended `raise SystemExit(f"sub1: nothing
            # says {prefix!r}")`, which is a BaseException, escaped this catch, unwound
            # past the save at the end of the file and cost the whole build -- two pages
            # already drawn and 14.6 minutes of writing, for one page's bad guess at a
            # placeholder string. A helper saying "this cannot work" is that page's
            # failure, which is exactly what this branch is for. KeyboardInterrupt is
            # not caught: that one is the caller stopping the build.
            except (Exception, SystemExit):
                text = traceback.format_exc()
                sys.stderr.write(text)
                slides = _slides()
                if slides is None:
                    raise
                error = text.strip().splitlines()[-1]
                _drop_last(slides, len(created) - before)
                del created[before:]
                standing_in[0] = start + 1
                try:
                    _stand_in(slides, number, error)
                finally:
                    standing_in[0] = None
                failed.append({"page": number, "error": error, "traceback": text[-4000:]})
        run(blocks[-1][2], len(lines))
finally:
    # The pages that were drawn, whatever happened to the rest. `prs.save` is the last
    # line of the author's file, so anything that escapes the loop above -- a prelude
    # that raises, a tail that raises, a KeyboardInterrupt -- skips it and the build
    # reports "wrote no deck" over a deck that was mostly finished. Saving here costs a
    # file write on a path that already had to be written for the build to succeed.
    if not os.path.exists(os.environ.get("PPT_OUTPUT", "")):
        try:
            from pptx.presentation import Presentation as _Presentation

            for _value in list(namespace.values()):
                if isinstance(_value, _Presentation) and len(_value.slides):
                    _value.save(os.environ["PPT_OUTPUT"])
                    break
        except Exception:  # noqa: BLE001 -- a rescue that fails leaves the original failure
            pass
    if record:
        import hashlib

        with open(record, "w", encoding="utf-8") as handle:
            # The digest of the text these line numbers were read from. Edit the
            # script and they point at whatever now sits on those lines.
            json.dump({"lines": created, "script_sha256": hashlib.sha256(body).hexdigest(), "failed": failed}, handle)
"""


class ScriptBackend:
    """Backend protocol over `run_script`, for a pipeline that holds a backend."""

    name = "script"

    def __init__(self, helpers: HelperSources | None = None) -> None:
        self.helpers = helpers

    async def compose(self, project: Project, plan: DeckPlan | None = None) -> BuildOutcome:
        del plan  # the program is the plan on this route
        return await run_script(project, helpers=self.helpers)


def _relative(project: Project, path: Path) -> str:
    """A path the file tools can take. Absolute is unambiguous and unusable: the
    file tools resolve against the workspace, so an author handed an absolute path
    writes it verbatim and the file lands outside the project -- which a live run
    did, then spent four calls recovering."""
    try:
        return str(path.relative_to(project.workspace))
    except ValueError:
        return str(path)


async def run_script(
    project: Project,
    script: str | None = None,
    *,
    helpers: HelperSources | None = None,
    timeout_s: float = BUILD_TIMEOUT_S,
) -> BuildOutcome:
    """Execute the author's program with the project's figures on hand.

    Passing `script` submits a whole program inline. Omitting it -- or passing
    something with no program in it -- runs the build.py already in the build
    directory, and that is the path meant for a real deck: a script long enough
    to lay out twenty pages runs to hundreds of lines, inline submission makes
    every revision cost a full regeneration of all of them, and a generation that
    long is long enough for the connection under it to drop, which returns
    nothing at all rather than a partial script.
    """
    # Asked once per build rather than passed in: whether this deck has a
    # template is a fact about the project on disk, and a caller that had to
    # remember to pass it would forget on the path that matters -- the design
    # pass rebuilding after an edit.
    template = bound(project)
    workdir = provision(project, with_template_helpers(helpers, template) if template else helpers)
    source = script_path(project)
    note = ""

    if script is not None and carries_a_program(script):
        refusal = submission_refusal(script, source)
        if refusal:
            return BuildOutcome(ok=False, stderr=refusal)
        source.write_text(script, encoding="utf-8")
    elif script is not None and source.is_file():
        note = (
            f"the submitted script held no program, so {source.name} was left as it was and that is what "
            "ran. Omitting script does the same thing; pass one only to replace the program."
        )
    elif not source.is_file():
        return BuildOutcome(
            ok=False,
            stderr=(
                f"no build script yet: write the program to {_relative(project, source)} with write_file -- "
                "that whole path, relative to the workspace, because a bare build/build.py lands somewhere "
                "this does not look. Then run the build again with just the project. ppt_layout.py, "
                "ppt_charts.py, ppt_shapes.py, ppt_theme.py and ppt_icons.py are already beside it."
                + (
                    "\nThis deck has a template: open it with "
                    "`Presentation(os.environ['PPT_TEMPLATE'])` instead of `Presentation()`, and take its "
                    "palette and type from ppt_theme, which holds the template's own and nothing else. "
                    "ppt_template.py is beside the others, to clone a page out of PPT_TEMPLATE_SOURCE."
                    if template
                    else ""
                )
            ),
        )

    target = deck_path(project)
    # The script writes to a staging file and the deck is replaced only on
    # success. A build that dies mid-script must not also destroy the last deck
    # that built, because that deck is what a failed edit gets repaired against.
    staging = target.with_name(target.name + ".building")
    lines_map = slide_lines_path(project)
    staging.unlink(missing_ok=True)
    lines_map.unlink(missing_ok=True)

    env = dict(os.environ)
    # The two paths the program needs, so it never has to guess or reach outside
    # the project.
    env["PPT_FIGURES_DIR"] = str(project.figures_dir)
    env["PPT_OUTPUT"] = str(staging)
    env["PPT_SLIDE_LINES"] = str(lines_map)
    # No MPLBACKEND on purpose: matplotlib is here for the formula typesetter
    # (services/assets/layout), not for the program. Charts on this route are drawn
    # with python-pptx so they stay editable and on the deck's palette, and a
    # backend set here would read as an invitation to plot into a picture instead.
    if template:
        # Two paths, because they answer different questions. The prepared copy is
        # what the deck is built in -- the template with its example pages removed,
        # so a page added to it inherits the master, the theme and the canvas. The
        # original still holds those pages, and it is the only place the two thirds
        # of them that python-pptx cannot redraw can be reached at all: a page is
        # cloned out of it and edited.
        env["PPT_TEMPLATE"] = str(template.prepared)
        env["PPT_TEMPLATE_SOURCE"] = str(template.source)
    # Set with or without a template bound: `bundled()` in the projected ppt_template
    # answers by this, and a program that borrows a page should learn a name is wrong
    # rather than that nothing can be opened.
    from raven_ppt.services.template.defaults import templates_dir

    env["PPT_BUNDLED_TEMPLATES"] = str(templates_dir())
    isolable = _isolable_blocks(source)
    env["PPT_SLIDE_BLOCKS"] = json.dumps(isolable)

    runner = workdir / "_run_build.py"
    runner.write_text(_RUNNER, encoding="utf-8")
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        str(runner),
        str(source),
        cwd=str(workdir),
        env=env,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    try:
        raw_out, raw_err = await asyncio.wait_for(process.communicate(), timeout=timeout_s)
    except TimeoutError:
        process.kill()
        await process.wait()
        _save_failure(
            project,
            source,
            staging,
            "timeout",
            "",
            f"the build script exceeded {timeout_s:.0f}s and was stopped",
        )
        _discard(staging, lines_map)
        return BuildOutcome(ok=False, stderr=f"the build script exceeded {timeout_s:.0f}s and was stopped")

    stdout = raw_out.decode("utf-8", "replace")[-MAX_OUTPUT_CHARS:]
    stderr = raw_err.decode("utf-8", "replace")[-MAX_OUTPUT_CHARS:]

    if process.returncode != 0:
        # The line record describes the script that failed, not the deck on disk, so
        # neither is kept as this deck's. The pages that were drawn are: the runner
        # saves them on its way out, and the failure folder keeps that file, so a crash
        # past the page loop costs the round rather than the writing.
        kept = _save_failure(project, source, staging, "crash", stdout, stderr or f"exit code {process.returncode}")
        drawn = (kept / "deck.pptx.building").is_file()
        _discard(staging, lines_map)
        said = stderr or f"exit code {process.returncode}"
        if not isolable:
            # Worth saying on the crash rather than only in the mapping warning a
            # successful build reports: without the banners nothing was isolated, so
            # this error cost every page, and the author cannot see that from the
            # traceback. A live run's rewrite dropped them and the next error took
            # five drawn pages with it.
            said += (
                "\n\nThis script's pages are not separable: `# SLIDE n` banners have to number the "
                "pages 1..N in file order for a page that raises to fail alone. Without them one "
                "error costs the whole build, as it did here."
            )
        if drawn:
            said += (
                f"\n\nThe pages drawn before this are kept at {_relative(project, kept / 'deck.pptx.building')} -- "
                "open it to see how far the program got."
            )
        return BuildOutcome(ok=False, stdout=stdout, stderr=said, note=note)
    if not staging.is_file():
        # The runner recorded line numbers under this script's hash, but the deck
        # those numbers would be paired with is still the old one -- keeping the
        # record would create exactly the mismatch the hash guards against.
        message = f"the script finished but wrote no deck at {target.name}; save to os.environ['PPT_OUTPUT']"
        _save_failure(project, source, staging, "no_output", stdout, message)
        _discard(staging, lines_map)
        return BuildOutcome(
            ok=False,
            stdout=stdout,
            note=note,
            stderr=message,
        )

    # The defects every deck arrives with, corrected before anything measures it:
    # the empty placeholders a template's layout leaves on a page ("Click to add
    # title" in Office), the Office gallery style python-pptx stamps on every table,
    # and the theme drop shadow it references on every shape it draws. Only the last
    # of those shows up in a render -- see raven_ppt/services/tidy.py.
    for line in tidy(staging):
        log.debug("tidy: %s", line)
    failed = _failed_pages(lines_map)
    if failed:
        _save_failure(project, source, staging, "page", stdout, stderr)
        page_failures_path(project).parent.mkdir(parents=True, exist_ok=True)
        page_failures_path(project).write_text(
            json.dumps({"pages": failed}, ensure_ascii=False, indent=1), encoding="utf-8"
        )
    else:
        page_failures_path(project).unlink(missing_ok=True)
    staging.replace(target)
    sources, digest = _sources(source, lines_map)
    return BuildOutcome(
        ok=True,
        pptx_path=target,
        pages=_count_slides(target),
        stdout=stdout,
        stderr=stderr,
        note=note,
        sources=sources,
        source_digest=digest,
    )


def _isolable_blocks(source: Path) -> list[list[int]]:
    """The `# SLIDE n` blocks the runner may execute one at a time, or nothing.

    Nothing when the banners do not number the pages 1..N in file order: the isolation
    stands a placeholder page where a failed block's page would be, and that only keeps
    page and block paired when the blocks come in page order.
    """
    try:
        lines = source.read_text(encoding=SCRIPT_ENCODING).splitlines(keepends=True)
    except OSError:
        return []
    blocks = page_blocks(lines)
    ordered = sorted(blocks.items(), key=lambda item: item[1][0])
    if [number for number, _ in ordered] != list(range(1, len(ordered) + 1)):
        return []
    return [[number, start, end] for number, (start, end) in ordered]


def _failed_pages(lines_map: Path) -> list[dict]:
    try:
        payload = json.loads(lines_map.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    failed = payload.get("failed")
    return (
        [entry for entry in failed if isinstance(entry, dict) and "page" in entry] if isinstance(failed, list) else []
    )


def page_failures(project: Project) -> list[dict]:
    """The pages the last build stood in for, from the record it wrote: page, error, traceback."""
    try:
        payload = json.loads(page_failures_path(project).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    pages = payload.get("pages") if isinstance(payload, dict) else None
    return [entry for entry in pages if isinstance(entry, dict) and "page" in entry] if isinstance(pages, list) else []


def _discard(*paths: Path) -> None:
    for path in paths:
        path.unlink(missing_ok=True)


def _save_failure(
    project: Project,
    source: Path,
    staging: Path,
    kind: str,
    stdout: str,
    stderr: str,
) -> Path:
    """Keep the failed script and logs so the next edit can continue from it."""
    root = project.review_dir / "build_failures"
    index = 1
    while (root / f"failure-{index:03d}").exists():
        index += 1
    destination = root / f"failure-{index:03d}"
    destination.mkdir(parents=True, exist_ok=True)
    if source.is_file():
        shutil.copy2(source, destination / "build.py")
    if staging.is_file():
        shutil.copy2(staging, destination / "deck.pptx.building")
    (destination / "stdout.txt").write_text(stdout, encoding="utf-8")
    (destination / "stderr.txt").write_text(stderr, encoding="utf-8")
    (destination / "failure.json").write_text(
        json.dumps({"kind": kind, "script": str(source), "staging": str(staging)}, indent=2) + "\n",
        encoding="utf-8",
    )
    return destination


def _sources(source: Path, lines_map: Path):
    """Page spans from the execution record, or nothing if it cannot be trusted.

    Line numbers are only true of the text they were read from, so a record whose
    digest no longer matches the script is discarded rather than believed. A
    caller getting an empty tuple is being told the pages cannot be told apart,
    which is a finding -- not a reason to guess.
    """
    try:
        payload = json.loads(lines_map.read_text(encoding="utf-8"))
        body = source.read_bytes()
    except (OSError, ValueError):
        return (), ""
    digest = hashlib.sha256(body).hexdigest()
    if payload.get("script_sha256") != digest:
        return (), ""
    stacks = payload.get("lines")
    if not isinstance(stacks, list) or not stacks:
        return (), digest
    chains = [[int(v) for v in stack] for stack in stacks if isinstance(stack, list) and stack]
    if len(chains) != len(stacks):
        return (), digest
    created = level_that_separates_pages(chains)
    if not created:
        return (), digest
    lines = body.decode(SCRIPT_ENCODING, "replace").splitlines(keepends=True)
    return page_sources(lines, created), digest


def _count_slides(pptx_path: Path) -> int:
    from pptx import Presentation

    return len(Presentation(str(pptx_path)).slides)
