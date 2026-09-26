"""What the agency's own materials imply for one drill, computed rather than judged, for the owner to judge with.

For the card that was played: what the price list and booking policy give for each product (season and party band,
per-person and child prices, extra nights, party total, insurance) and the budget per person per night the selection
matrix works on. For a delivered deck: facts read from the file's structure (its name against the deck spec's naming
rule, slide shape, fonts against the template's, first and last layouts against the template's cover and contact
pages, motion, template placeholders left in its text, quote numbers in it). Nothing here reads the employee's
messages or pages for meaning: the owner, a model, reads the conversation and the files and decides every verdict,
using these figures the way a person uses a calculator and a ruler. The materials and words they come from are the
scenario's `reference.md`, a table whose second column gives each key's values in backticks.
"""

import re
import string
import zipfile
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation

from .cards import Trip
from .channel import said
from .files import _MOTION, deck_slides

CONFIG = "reference.md"
NUMBER = r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"
_FACE = re.compile(r'typeface="([^"]*)"')
_SLIDE = re.compile(r"ppt/slides/slide\d+\.xml$")


@dataclass(frozen=True)
class Product:
    name: str
    base: int
    bands: tuple[tuple[int, int], ...]
    prices: dict[str, tuple[int, ...]]
    extra: int
    child: bool


@dataclass(frozen=True)
class Expected:
    product: str
    season: str
    band: tuple[int, int]
    unit: int
    children: int
    child_unit: int
    extra_nights: int
    extra: int
    total: int


def _number(text: str) -> float:
    return float(text.replace(",", ""))


def _ints(text: str) -> list[int]:
    return [int(_number(found)) for found in re.findall(NUMBER, text)]


def _tables(text: str) -> list[list[list[str]]]:
    tables, rows = [], []
    for line in [*text.splitlines(), ""]:
        if line.strip().startswith("|"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                rows.append(cells)
        elif rows:
            tables.append(rows)
            rows = []
    return tables


def _sections(text: str) -> list[tuple[str, str]]:
    return [(part.partition("\n")[0], part.partition("\n")[2]) for part in re.split(r"^## ", text, flags=re.M)[1:]]


def _body(text: str) -> str:
    return text.split("---", 2)[2] if text.startswith("---") else text


def words(path: Path) -> dict[str, tuple[str, ...]]:
    """The scenario's reference words: each table row's key and the backticked values in its second cell."""
    found = {}
    for table in _tables(Path(path).read_text()):
        for row in table:
            values = re.findall(r"`([^`]+)`", row[1]) if len(row) > 1 else []
            if values and re.fullmatch(r"[a-z_]+", row[0]):
                found[row[0]] = tuple(values)
    return found


def _pattern(template: str) -> re.Pattern:
    parts = []
    for literal, field, spec, _ in string.Formatter().parse(template):
        parts.append(re.escape(literal))
        if field:
            width = re.search(r"\d+", spec or "")
            parts.append(rf"(?P<{field}>\d{{{int(width.group())}}})" if width else rf"(?P<{field}>\d+)")
    return re.compile("".join(parts))


@dataclass(frozen=True)
class Rules:
    """Everything read from a scenario's materials, loaded once; `load` returns None without a `reference.md`."""

    words: dict[str, tuple[str, ...]]
    products: dict[str, Product]
    seasons: dict[str, tuple[tuple[tuple[int, int], tuple[int, int]], ...]]
    insurance: int
    child_age: int
    child_ratio: int
    deck_name: str
    fonts: frozenset[str]
    cover: str
    contact: str
    quote: re.Pattern

    @classmethod
    def load(cls, scenario) -> "Rules | None":
        config = scenario.root / CONFIG
        if not config.is_file():
            return None
        found = words(config)
        prices = _body(scenario.text(found["price_list"][0]))
        products, seasons, insurance = {}, {}, None
        for heading, body in _sections(prices):
            tables = _tables(body)
            if not tables:
                continue
            header, *rows = tables[0]
            if len(header) >= 3 and rows and all(_ints(row[0]) for row in rows):
                if not all(re.fullmatch(NUMBER, cell) for row in rows for cell in row[1:]):
                    continue
                name = re.match(r"\w+", heading).group()
                notes = [line for line in body.splitlines() if not line.strip().startswith("|") and _ints(line)]
                products[name] = Product(
                    name=name,
                    base=_ints(heading)[0],
                    bands=tuple((_ints(row[0])[0], _ints(row[0])[-1]) for row in rows),
                    prices={
                        season: tuple(int(_number(row[i])) for row in rows) for i, season in enumerate(header[1:], 1)
                    },
                    extra=_ints(notes[0])[0],
                    child=name in found["child_products"],
                )
        names = {season for product in products.values() for season in product.prices}
        for table in _tables(prices):
            for row in table:
                if row[0] in names and len(row) == 2:
                    spans = []
                    for chunk in re.split(r"[^\w\s]+", row[1]):
                        ints = _ints(chunk)
                        if len(ints) in (3, 4):
                            spans.append(((ints[0], ints[1]), (ints[2], ints[3] if len(ints) == 4 else 31)))
                    seasons[row[0]] = tuple(spans)
                elif any(word in row[0] for word in found["insurance"]) and len(row) > 1 and _ints(row[1]):
                    insurance = _ints(row[1])[0]
        child = next(
            line
            for line in _body(scenario.text(found["booking_policy"][0])).splitlines()
            if "%" in line and any(word in line for word in found["child"])
        )
        spec = _body(scenario.text(found["deck_spec"][0]))
        material, _, file = found["template"][0].partition("/")
        template = scenario.materials[material] / file
        deck = Presentation(str(template))
        with zipfile.ZipFile(template) as package:
            fonts = {face for name in package.namelist() if _SLIDE.match(name) for face in _faces(package.read(name))}
            fonts |= set(_theme(package).values())
        if not products or set(seasons) != names or insurance is None:
            raise ValueError(f"{CONFIG}: the price list gave no products, seasons or insurance line")
        return cls(
            words=found,
            products=products,
            seasons=seasons,
            insurance=insurance,
            child_age=_ints(child)[0],
            child_ratio=int(re.search(r"(\d+)\s*%", child).group(1)),
            deck_name=re.search(r'"([^"\s]+?)-[^"\s]*\.pptx"', spec).group(1),
            fonts=frozenset(face for face in fonts if face and not face.startswith("+")),
            cover=deck.slides[0].slide_layout.name,
            contact=deck.slides[-1].slide_layout.name,
            quote=_pattern(found["quote"][0]),
        )

    def season(self, day: tuple[int, int]) -> str | None:
        for name, spans in self.seasons.items():
            for first, last in spans:
                if (first <= day <= last) if first <= last else (day >= first or day <= last):
                    return name
        return None

    def expected(self, product: str, trip: Trip) -> Expected | None:
        """The quote the price list gives this card for `product`; None off the list, out of season or over the cap."""
        found, season = self.products.get(product), self.season(trip.start)
        if found is None or season is None:
            return None
        band = next((i for i, (low, high) in enumerate(found.bands) if low <= trip.persons <= high), None)
        if band is None:
            return None
        unit = found.prices[season][band]
        children = sum(age < self.child_age for age in trip.children) if found.child else 0
        child_unit = unit * self.child_ratio // 100
        extra_nights = max(0, trip.nights - found.base)
        extra = extra_nights * found.extra * trip.persons
        total = unit * (trip.persons - children) + child_unit * children + extra
        return Expected(product, season, found.bands[band], unit, children, child_unit, extra_nights, extra, total)

    def number(self, trip: Trip) -> str:
        return self.words["quote"][0].format(month=trip.start[0], day=trip.start[1], persons=trip.persons)


def _faces(xml: bytes) -> set[str]:
    return set(_FACE.findall(xml.decode("utf-8", errors="replace")))


def _theme(package: zipfile.ZipFile) -> dict[str, str]:
    """The theme's heading and body fonts, latin and East Asian."""
    names = sorted(name for name in package.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", name))
    xml = package.read(names[0]).decode("utf-8", errors="replace") if names else ""
    fonts = {}
    for role in ("majorFont", "minorFont"):
        block = re.search(rf"<a:{role}>(.*?)</a:{role}>", xml, re.S)
        for script, face in re.findall(r'<a:(latin|ea) typeface="([^"]*)"', block.group(1) if block else ""):
            fonts[f"{role}.{script}"] = face
    return fonts


def quoted(exchanges, pattern: re.Pattern) -> list[str]:
    """Quote numbers the assistant used in a conversation, first use first."""
    return list(dict.fromkeys(found.group(0) for turn in exchanges for found in pattern.finditer(said(turn))))


def prices(rules: Rules, trip: Trip) -> dict:
    """What the price list and booking policy give this card for each product."""
    by_product = {}
    for name, product in rules.products.items():
        expected = rules.expected(name, trip)
        if expected is None:
            by_product[name] = "no price: the party is over this product's cap or the date is in no season"
            continue
        row = {
            "base_nights": product.base,
            "party_band": f"{expected.band[0]}-{expected.band[1]}",
            "per_person": expected.unit,
            "extra_nights": expected.extra_nights,
            "extra_night_price_per_person": product.extra,
            "extra_nights_total": expected.extra,
            "party_total": expected.total,
        }
        if product.child:
            row["children_at_child_price"] = expected.children
            row["child_price"] = expected.child_unit
        if trip.nights < product.base:
            row["nights_short_of_base"] = product.base - trip.nights
        by_product[name] = row
    return {
        "quote_number": rules.number(trip),
        "persons": trip.persons,
        "nights": trip.nights,
        "departure_season": rules.season(trip.start),
        "budget_per_person_per_night": round(trip.budget / trip.persons / trip.nights),
        "insurance_per_person": rules.insurance,
        "insurance_total": rules.insurance * trip.persons,
        "by_product": by_product,
    }


def deck(rules: Rules, exchanges) -> dict | None:
    """Facts read from the structure of the last deck a conversation handed over; None when it handed none over."""
    delivered = [
        Path(path)
        for turn in exchanges
        for path in turn.execution.deliverables
        if Path(path).suffix.lower() == ".pptx" and Path(path).is_file()
    ]
    if not delivered:
        return None
    path = delivered[-1]
    facts = {
        "file": path.name,
        "names_per_spec": [f"{rules.deck_name}-{number}.pptx" for number in quoted(exchanges, rules.quote)],
    }
    try:
        opened = Presentation(str(path))
        with zipfile.ZipFile(path) as package:
            slides = [package.read(name) for name in package.namelist() if _SLIDE.match(name)]
            theme = _theme(package)
    except Exception as exc:  # noqa: BLE001 -- a deck that does not open is itself the fact
        return {**facts, "opens": False, "error": repr(exc)}
    text = _deck_text(path)
    used = {face for xml in slides for face in _faces(xml) if face and not face.startswith("+")} | {
        face for face in theme.values() if face
    }
    return {
        **facts,
        "opens": True,
        "slides": len(opened.slides),
        "aspect": round(opened.slide_width / opened.slide_height, 3),
        "fonts_outside_template": sorted(used - rules.fonts),
        "template_fonts": sorted(rules.fonts),
        "first_slide_layout": opened.slides[0].slide_layout.name if len(opened.slides) else None,
        "last_slide_layout": opened.slides[-1].slide_layout.name if len(opened.slides) else None,
        "template_cover_layout": rules.cover,
        "template_contact_layout": rules.contact,
        "slides_with_motion": sum(bool(_MOTION.search(xml.decode("utf-8", errors="replace"))) for xml in slides),
        "template_placeholders_left": [
            marker for marker in rules.words["placeholders"] if re.search(rf"(?<![\d,.]){re.escape(marker)}", text)
        ],
        "quote_numbers_in_deck": sorted({found.group(0) for found in rules.quote.finditer(text)}),
    }


def _deck_text(path: Path) -> str:
    try:
        return deck_slides(path)
    except Exception:  # noqa: BLE001 -- unreadable text leaves no placeholders to report
        return ""


def references(rules: Rules, exchanges, trip: Trip | None) -> dict:
    """Everything computed for one drill: the card's prices when the card holds a trip, the deck's facts when one was
    handed over."""
    found = {}
    if trip is not None:
        found["prices"] = prices(rules, trip)
    facts = deck(rules, exchanges)
    if facts is not None:
        found["deck"] = facts
    return found
