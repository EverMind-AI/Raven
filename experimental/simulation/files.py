"""How a person reads a delivered file: a web page as its visible text, a deck slide by slide.

A traveller reads what the file says; the owner, reviewing against the agency's design norms, also reads
how each slide is laid out and styled, and looks at every page as rendered.
"""

import re
from html.parser import HTMLParser
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .render import cached

_HIDDEN = ("script", "style", "template")
_TEXT = (".md", ".txt", ".json", ".csv", ".yaml", ".yml")
_MOTION = re.compile(r"<p:(transition|timing)\b")
PAGE_WIDTH = 960


class _Visible(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self._hidden = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in _HIDDEN:
            self._hidden += 1

    def handle_endtag(self, tag):
        if tag in _HIDDEN and self._hidden:
            self._hidden -= 1

    def handle_data(self, data):
        if not self._hidden and data.strip():
            self.parts.append(data.strip())


def visible_text(path: Path) -> str:
    parser = _Visible()
    parser.feed(Path(path).read_text(errors="replace"))
    return "\n".join(parser.parts)


def _shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _colour(font) -> str | None:
    try:
        return f"#{font.color.rgb}" if font.color and font.color.type is not None else None
    except (AttributeError, TypeError):
        return None


def deck_slides(path: Path, *, style: bool = False) -> str:
    """Each slide's text in reading order; with `style`, also its layout, fonts, colours, pictures and motion."""
    deck = Presentation(str(path))
    area = deck.slide_width * deck.slide_height
    lines = [f"{len(deck.slides)} slides, aspect {deck.slide_width / deck.slide_height:.2f}"]
    for number, slide in enumerate(deck.slides, 1):
        lines.append(f"--- slide {number}" + (f" · layout {slide.slide_layout.name}" if style else ""))
        fonts, pictures = set(), []
        for shape in _shapes(slide.shapes):
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    words = "".join(run.text for run in paragraph.runs).strip()
                    if words:
                        lines.append(words)
                    for run in paragraph.runs:
                        size = run.font.size.pt if run.font.size else None
                        fonts.add((run.font.name, size, _colour(run.font)))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width and shape.height:
                pictures.append(round(100 * shape.width * shape.height / area))
        if style:
            fill = slide.background.fill
            try:
                background = f"#{fill.fore_color.rgb}" if fill.type == MSO_FILL.SOLID else "template default"
            except (AttributeError, TypeError):
                background = "template default"
            used = sorted(
                f"{name or 'inherited'} {size or '?'}pt {colour or 'inherited'}" for name, size, colour in fonts
            )
            lines.append(f"[style] background {background}; fonts: {', '.join(used) or 'inherited'}")
            if pictures:
                lines.append(f"[style] pictures covering {', '.join(f'{p}%' for p in pictures)} of the slide")
            if _MOTION.search(slide.part.blob.decode("utf-8", errors="replace")):
                lines.append("[style] has a transition or animation")
    return "\n".join(lines)


def read_delivered(path: Path, *, style: bool = False) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".html", ".htm"):
        return visible_text(path)
    if suffix == ".pptx":
        return deck_slides(path, style=style)
    if suffix in _TEXT:
        return path.read_text(errors="replace")
    return f"(a {suffix or 'binary'} file of {path.stat().st_size} bytes, not readable as text)"


def page_images(path: Path) -> list[Path]:
    """How each page of a delivered deck looks, rendered once beside it as `<file>.thumbs/` (the record keeps them)."""
    file = Path(path)
    return cached(file, file.with_name(f"{file.name}.thumbs"), PAGE_WIDTH)
